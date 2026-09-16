"""投标承诺书：按固定模板排版，只填致/投标人/签字人/日期。"""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from api.services.tenders.schema import BidBrief

_SONG = "宋体"

# 按常见「附件五：投标承诺书」条款整理；项目名/招标人走表单，正文不经 LLM。
_CLAUSES: tuple[str, ...] = (
    "保证投标文件无虚假内容，一旦发现并查实，同意取消投标资格，并同意被没收投标保证金。",
    "保证不恶意低于工程成本价报价进行投标。",
    "保证报价已综合考虑本工程招标范围内全部工作内容及合同价款调整原则，不恶意追求变更。",
    "保证除法律法规另有规定外，不以任何方式向招标人打听中标、落标原因。",
    "保证按照投标文件配备项目负责人、主要管理人员及施工机具设备。",
    "保证全面履行合同范围内的供货、安装、调试及质保期内保修责任。",
    "保证不将本工程主体及关键工作违法分包或转包。",
    "保证已充分理解并接受招标文件及合同条款的全部内容。",
    "保证无条件执行技术交底及项目管理各项制度要求。",
    "保证即使出现暂时资金不到位等情况，也不采取停工、上访等不当行为。",
)


def append_commitment_letter(doc: Document, brief: BidBrief) -> None:
    """在文末追加一页投标承诺书。"""
    doc.add_page_break()
    _write_annex_label(doc)
    _write_title(doc)
    write_commitment_body(doc, brief)


def write_commitment_body(doc: Document, brief: BidBrief) -> None:
    _write_to_line(doc, (brief.tenderer or "").strip())
    _write_intro(doc)
    for i, text in enumerate(_CLAUSES, start=1):
        _write_clause(doc, i, text)
    _write_sign_block(doc, brief)


def _font(run, size_pt: float, *, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = _SONG
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    for tag in (qn("w:b"), qn("w:bCs")):
        if not bold:
            for old in list(rpr.findall(tag)):
                rpr.remove(old)
    if bold:
        if rpr.find(qn("w:b")) is None:
            rpr.append(OxmlElement("w:b"))
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), _SONG)
    rfonts.set(qn("w:hAnsi"), _SONG)
    rfonts.set(qn("w:eastAsia"), _SONG)


def _clear_bold(run) -> None:
    run.bold = False
    rpr = run._element.get_or_add_rPr()
    for tag in (qn("w:b"), qn("w:bCs")):
        for old in list(rpr.findall(tag)):
            rpr.remove(old)


def _underline_fill(run, text: str, *, size_pt: float = 12) -> None:
    """横线填空：下划线、不加粗、不放大。"""
    run.text = f" {(text or '').strip() or '　'} "
    _font(run, size_pt, bold=False)
    _clear_bold(run)
    run.underline = True


def _write_annex_label(doc: Document) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run("附件五：")
    _font(run, 12, bold=True)


def _write_title(doc: Document) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(18)
    run = para.add_run("投标承诺书")
    _font(run, 16, bold=True)


def _write_to_line(doc: Document, tenderer: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(10)
    label = para.add_run("致：")
    _font(label, 12, bold=False)
    fill = para.add_run()
    _underline_fill(fill, tenderer, size_pt=12)


def _write_intro(doc: Document) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(8)
    run = para.add_run("我公司现做出如下承诺：")
    _font(run, 12, bold=False)


def _write_clause(doc: Document, index: int, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = para.paragraph_format
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(2)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.first_line_indent = Pt(0)
    run = para.add_run(f"{index}.{text}")
    _font(run, 12, bold=False)


def _ymd(iso: str) -> tuple[str, str, str]:
    raw = (iso or "").strip()
    parts = raw.replace("/", "-").split("-")
    if len(parts) >= 3:
        month = (parts[1].lstrip("0") or "1").zfill(2)
        day = (parts[2].lstrip("0") or "1").zfill(2)
        return parts[0], month, day
    return "　　", "　", "　"


def _write_sign_block(doc: Document, brief: BidBrief) -> None:
    for _ in range(3):
        doc.add_paragraph()

    bidder = (brief.bidderName or "").strip() or "河南伟泰光电科技有限公司"
    signer = ""
    year, month, day = _ymd(brief.bidDate)

    line1 = doc.add_paragraph()
    line1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    line1.paragraph_format.space_before = Pt(8)
    a = line1.add_run("投 标 人：")
    _font(a, 12, bold=False)
    b = line1.add_run()
    _underline_fill(b, bidder, size_pt=12)
    c = line1.add_run("（盖章）")
    _font(c, 12, bold=False)

    line2 = doc.add_paragraph()
    line2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    line2.paragraph_format.space_before = Pt(12)
    d = line2.add_run("法定代表人或委托代理人：")
    _font(d, 12, bold=False)
    e = line2.add_run()
    _underline_fill(e, signer, size_pt=12)
    f = line2.add_run("（签字并盖章）")
    _font(f, 12, bold=False)

    line3 = doc.add_paragraph()
    line3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    line3.paragraph_format.space_before = Pt(12)
    g = line3.add_run("日    期：")
    _font(g, 12, bold=False)
    y = line3.add_run()
    _underline_fill(y, year, size_pt=12)
    line3.add_run("年")
    _font(line3.runs[-1], 12, bold=False)
    m = line3.add_run()
    _underline_fill(m, month, size_pt=12)
    line3.add_run("月")
    _font(line3.runs[-1], 12, bold=False)
    d2 = line3.add_run()
    _underline_fill(d2, day, size_pt=12)
    line3.add_run("日")
    _font(line3.runs[-1], 12, bold=False)
