"""按 purpose 写出造价 / 报价 / 预算 Excel。"""

from __future__ import annotations

import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from uuid import uuid4

from api.services.quotes.catalog import money
from api.services.quotes.contract import EXPORT_SHEETS, normalize_purpose
from api.services.quotes.library import quotes_output_dir
from api.services.quotes.schema import QuoteLineIn
from common.errors import AppError, ErrorCode

_TWO = Decimal("0.01")
_FILE_RE = re.compile(r"^[a-f0-9]{12}\.xlsx$", re.I)
_UNSAFE = re.compile(r'[\\/:*?"<>|\s]+')
# 列宽；规格型号(C)留宽，方便多行技术参数换行
_COL_WIDTHS = (8, 22, 36, 8, 10, 14, 14, 14, 18, 22)
_LINE_PT = 14.5
_ROW_H_MIN = 22.0
_ROW_H_MAX = 360.0


def _norm_cell_text(val: Any) -> Any:
    if not isinstance(val, str):
        return val
    return val.replace("\r\n", "\n").replace("\r", "\n").strip()


def _wrapped_lines(text: str, col_width: float) -> int:
    t = _norm_cell_text(text)
    if not isinstance(t, str) or not t:
        return 1
    # Excel 列宽约等于半角字符数；中文按约 2 个宽度算
    usable = max(4, int(float(col_width) * 0.92))
    total = 0
    for para in t.split("\n"):
        if not para:
            total += 1
            continue
        lines = 1
        cur = 0
        for ch in para:
            w = 2 if ord(ch) > 0x7F else 1
            if cur + w > usable and cur > 0:
                lines += 1
                cur = w
            else:
                cur += w
        total += lines
    return max(1, total)


def _row_height_for(values: list[Any], widths: tuple[float, ...] | list[float]) -> float:
    lines = 1
    for i, val in enumerate(values):
        if not isinstance(val, str) or not val.strip():
            continue
        w = float(widths[i]) if i < len(widths) else 12.0
        lines = max(lines, _wrapped_lines(val, w))
    h = 4.0 + lines * _LINE_PT
    return min(max(_ROW_H_MIN, h), _ROW_H_MAX)


def _col_letter(i: int) -> str:
    # 1-based；本表最多十列
    return chr(64 + i) if 1 <= i <= 26 else "A"


def _body_alignment(col: int, val: Any, widths: tuple[float, ...] | list[float]):
    """规格等多行顶对齐；序号/名称/单位/数量/金额/来源等字少列垂直居中。"""
    from openpyxl.styles import Alignment

    # 3=规格型号
    if col == 3:
        return Alignment(vertical="top", horizontal="left", wrap_text=True)
    # 8=备注：多行仍顶对齐，短文居中
    if col == 8 and isinstance(val, str) and val.strip():
        w = float(widths[col - 1]) if col - 1 < len(widths) else 14.0
        if "\n" in val or _wrapped_lines(val, w) > 2:
            return Alignment(vertical="top", horizontal="left", wrap_text=True)
    # 序号、单位、数量、单价、合价：水平+垂直居中
    if col in (1, 4, 5, 6, 7) or isinstance(val, (int, float, Decimal)):
        return Alignment(vertical="center", horizontal="center", wrap_text=True)
    # 名称、来源等
    return Alignment(vertical="center", horizontal="center", wrap_text=True)


def safe_download_name(project_name: str, *, purpose: str = "quote") -> str:
    stem = _UNSAFE.sub("-", (project_name or "").strip())[:40].strip("-") or "quote"
    suffix = {"cost": "造价测算表", "budget": "控制预算表", "quote": "报价单"}.get(
        normalize_purpose(purpose), "报价单"
    )
    return f"{stem}-{suffix}.xlsx"


def resolve_quote_file(file_name: str) -> Path:
    name = (file_name or "").strip().replace("\\", "/").split("/")[-1]
    if not _FILE_RE.match(name):
        raise AppError(ErrorCode.BAD_REQUEST, "报价文件名无效", status_code=400)
    path = (quotes_output_dir() / name).resolve()
    root = quotes_output_dir().resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise AppError(ErrorCode.NOT_FOUND, "报价文件不存在", status_code=404)
    return path


def write_quote_xlsx(
    *,
    project_name: str,
    note: str,
    tax_rate: Decimal,
    lines: list[QuoteLineIn],
    purpose: str = "quote",
    cost_summary: dict[str, Any] | None = None,
    schemes: list[dict[str, Any]] | None = None,
    instruction: str = "",
    verify: dict[str, Any] | None = None,
) -> tuple[Path, str, Decimal, Decimal]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    purpose = normalize_purpose(purpose)
    rows = [ln for ln in lines if str(ln.name or "").strip()]
    if not rows:
        raise AppError(ErrorCode.VALIDATION, "没有明细行，无法生成", status_code=422)

    thin = Border(
        left=Side(style="thin", color="C5CDD6"),
        right=Side(style="thin", color="C5CDD6"),
        top=Side(style="thin", color="C5CDD6"),
        bottom=Side(style="thin", color="C5CDD6"),
    )
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, name="微软雅黑", size=10)
    title_font = Font(bold=True, name="微软雅黑", size=16)
    body_font = Font(name="微软雅黑", size=10)
    money_fmt = "#,##0.00"

    wb = Workbook()
    sheet_title = EXPORT_SHEETS.get(purpose, "报价单")
    ws = wb.active
    ws.title = sheet_title[:31]
    title = (project_name or "").strip() or sheet_title
    headers = _headers_for(purpose)
    col_n = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_n)
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_n)
    extra = (note or "").strip()
    ws["A2"] = extra if extra else _default_note(purpose)
    ws["A2"].font = Font(name="微软雅黑", size=9, color="666666")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = _row_height_for([ws["A2"].value or ""], (sum(_COL_WIDTHS[:col_n]),))

    for i, h in enumerate(headers, start=1):
        cell = ws.cell(3, i, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[3].height = 32

    widths = _COL_WIDTHS[:col_n]
    for i in range(1, col_n + 1):
        ws.column_dimensions[_col_letter(i)].width = widths[i - 1]

    total = Decimal("0")
    for i, ln in enumerate(rows, start=1):
        qty = ln.qty if ln.qty > 0 else Decimal("0")
        cost = ln.costPrice if ln.costPrice > 0 else Decimal("0")
        sell = ln.sellPrice if ln.sellPrice > 0 else (ln.unitPrice if ln.unitPrice > 0 else Decimal("0"))
        if cost <= 0 and ln.unitPrice > 0:
            cost = ln.unitPrice
        if sell <= 0 and cost > 0:
            sell = cost
        price = _line_price(purpose, cost, sell, ln.unitPrice)
        amount = money(qty, price)
        total += amount
        values = [_norm_cell_text(v) for v in _row_values(purpose, i, ln, qty, cost, sell, price, amount)]
        r = 3 + i
        row_h = _row_height_for(values, widths)
        for c, val in enumerate(values, start=1):
            cell = ws.cell(r, c, val)
            cell.font = body_font
            cell.border = thin
            cell.alignment = _body_alignment(c, val, widths)
            if isinstance(val, (int, float, Decimal)) and c >= 5:
                cell.number_format = money_fmt if c != 5 else "0.####"
        ws.row_dimensions[r].height = row_h

    tax = (total * tax_rate).quantize(_TWO, rounding=ROUND_HALF_UP)
    inc = (total + tax).quantize(_TWO, rounding=ROUND_HALF_UP)
    foot = 4 + len(rows)
    amount_col = len(headers)
    for label_i, (label, val) in enumerate(
        [
            ("不含税合计", total),
            (f"增值税 {tax_rate * 100:.0f}%", tax),
            ("含税合计", inc),
        ]
    ):
        r = foot + label_i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(1, amount_col - 1))
        ws.cell(r, 1, label).font = Font(name="微软雅黑", size=10, bold=True)
        ws.cell(r, 1).border = thin
        cell = ws.cell(r, amount_col, val)
        cell.number_format = money_fmt
        cell.font = Font(name="微软雅黑", size=10, bold=True)
        cell.border = thin

    if cost_summary:
        _write_summary_sheet(wb, cost_summary, purpose)
    if verify and (verify.get("issues") or []):
        _write_verify_sheet(wb, verify)
    if purpose == "quote" and schemes:
        _write_schemes_sheet(wb, schemes, instruction or "")

    stem = uuid4().hex[:12]
    dest = quotes_output_dir() / f"{stem}.xlsx"
    wb.save(dest)
    # 导出合计：造价用成本、预算用预算、报价用售价合计
    if cost_summary:
        if purpose == "cost":
            total = Decimal(str(cost_summary.get("costExTax") or total))
            inc = Decimal(str(cost_summary.get("costIncTax") or inc))
        elif purpose == "budget":
            total = Decimal(str(cost_summary.get("budgetExTax") or total))
            inc = Decimal(str(cost_summary.get("budgetIncTax") or inc))
        elif purpose == "quote":
            total = Decimal(str(cost_summary.get("quoteExTax") or total))
            inc = Decimal(str(cost_summary.get("quoteIncTax") or inc))
    return dest, safe_download_name(title, purpose=purpose), total, inc


def _headers_for(purpose: str) -> tuple[str, ...]:
    if purpose == "cost":
        return ("序号", "名称", "规格型号", "单位", "数量", "成本单价", "成本合价", "备注", "来源")
    if purpose == "budget":
        return ("序号", "名称", "规格型号", "单位", "数量", "预算单价", "预算合价", "备注", "来源")
    return ("序号", "名称", "规格型号", "单位", "数量", "不含税单价（元）", "合价", "备注", "来源")


def _default_note(purpose: str) -> str:
    if purpose == "cost":
        return "造价测算：单价取价目成本价，另含措施/管理等费率，请人工核对。"
    if purpose == "budget":
        return "控制预算：在造价基础上套预算系数与不可预见费，请人工核对。"
    return "由场地规划图/工程量识别，单价来自所选知识库价目表；请人工核对后使用。"


def _line_price(purpose: str, cost: Decimal, sell: Decimal, unit: Decimal) -> Decimal:
    if purpose == "cost":
        return cost if cost > 0 else (unit if unit > 0 else sell)
    if purpose == "budget":
        # 明细行仍用成本价展示，汇总用预算；单价展示成本以便核对
        return cost if cost > 0 else (unit if unit > 0 else sell)
    return sell if sell > 0 else (unit if unit > 0 else cost)


def _row_values(purpose, i, ln, qty, cost, sell, price, amount):
    return [
        i,
        ln.name,
        ln.spec,
        ln.unit or "项",
        qty,
        price if price > 0 else None,
        amount if price > 0 else None,
        ln.note,
        _source_label(ln.source, ln.matchName),
    ]


def _source_label(source: str, match_name: str) -> str:
    s = (source or "").strip()
    hit = (match_name or "").strip()
    labels = {
        "vision": "读图",
        "rule": "规则估算",
        "catalog": "价目表",
        "boq": "工程量表",
        "manual": "人工",
    }
    base = labels.get(s, s or "人工")
    if hit and s == "catalog":
        return f"{base}：{hit}"
    return base


def _write_summary_sheet(wb, summary: dict[str, Any], purpose: str) -> None:
    ws = wb.create_sheet("费用汇总")
    ws.append(["科目", "金额（元）"])
    rows = [
        ("材料费", summary.get("materialCost")),
        ("人工费", summary.get("laborCost")),
        ("措施费", summary.get("measureFee")),
        ("管理费", summary.get("manageFee")),
        ("造价（不含税）", summary.get("costExTax")),
        ("利润", summary.get("profit")),
        ("报价（不含税）", summary.get("quoteExTax")),
        ("不可预见费", summary.get("contingency")),
        ("控制预算（不含税）", summary.get("budgetExTax")),
    ]
    for label, val in rows:
        ws.append([label, val])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16


def _write_verify_sheet(wb, verify: dict[str, Any]) -> None:
    ws = wb.create_sheet("核算报告")
    ws.append(["级别", "行号", "代码", "说明"])
    for it in verify.get("issues") or []:
        ws.append([it.get("level"), it.get("row"), it.get("code"), it.get("message")])
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["D"].width = 60


def _write_schemes_sheet(wb, schemes: list[dict[str, Any]], instruction: str) -> None:
    from openpyxl.styles import Alignment

    ws = wb.create_sheet("报价方案")
    headers = ("方案", "不含税合计", "毛利率", "风险", "说明", "推荐")
    widths = (12, 14, 10, 36, 36, 8)
    ws.append(list(headers))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[_col_letter(i)].width = w
    for s in schemes:
        values = [
            _norm_cell_text(s.get("name")),
            s.get("totalExTax"),
            s.get("margin"),
            _norm_cell_text("；".join(s.get("risks") or [])),
            _norm_cell_text(s.get("tip")),
            "是" if s.get("recommended") else "",
        ]
        ws.append(values)
        r = ws.max_row
        h = _row_height_for(values, widths)
        ws.row_dimensions[r].height = h
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top" if h > _ROW_H_MIN + 1 else "center",
            )
    if instruction:
        ws2 = wb.create_sheet("编制说明")
        ws2["A1"] = _norm_cell_text(instruction)
        ws2["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        ws2.column_dimensions["A"].width = 100
        ws2.row_dimensions[1].height = _row_height_for([instruction], (100,))
