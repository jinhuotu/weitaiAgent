"""Drafts fill body headings even before the source TOC has PAGEREF fields."""
import pytest
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from api.services.tenders.assemble import (
    _fill_business_format_sections,
    _fill_drafted_format_sections,
    assemble_bid_docx,
)
from api.services.tenders.document import _bookmark_paragraph, _pageref_bookmark, _write_toc_line
from api.services.tenders.pdf_format import (
    _link_toc_body,
    _repair_format_table_toc,
    toc_block_ids,
    unify_full_docx,
)
from api.services.tenders.schema import BidBrief, OutlineItem


def _signature(path):
    doc = Document(path)
    marks = tuple((n.get(qn('w:name')), n.get(qn('w:id')))
                  for n in doc.element.body.iter(qn('w:bookmarkStart')))
    ends = tuple(n.get(qn('w:id')) for n in doc.element.body.iter(qn('w:bookmarkEnd')))
    assert len(marks) == len(set(marks)) == len({name for name, _ in marks})
    assert len(marks) == len({ident for _, ident in marks})
    assert len(ends) == len(set(ends))
    assert {ident for _, ident in marks} == set(ends)
    fields = tuple(n.text for n in doc.element.body.iter(qn('w:instrText')))
    return (tuple(p.text for p in doc.paragraphs), fields, marks, ends,
            tuple(sec._sectPr.xml for sec in doc.sections),
            tuple(sec.footer._element.xml for sec in doc.sections))


def _tab(para, location):
    if location == 'ppr':
        para.paragraph_format.tab_stops.add_tab_stop(Cm(15), WD_TAB_ALIGNMENT.RIGHT)
    elif location == 'hyperlink':
        link = OxmlElement('w:hyperlink')
        run = OxmlElement('w:r')
        run.append(OxmlElement('w:tab'))
        link.append(run)
        para._p.append(link)
    else:
        para.add_run().add_tab()


@pytest.mark.parametrize('kind,fill', [
    ('commitment_copy', _fill_business_format_sections),
    ('tech_plan', _fill_drafted_format_sections),
])
@pytest.mark.parametrize('location', ['run', 'ppr', 'hyperlink'])
@pytest.mark.parametrize('with_toc_heading', [False, True])
def test_unlinked_toc_never_receives_draft(tmp_path, kind, fill, location, with_toc_heading):
    doc = Document()
    if with_toc_heading:
        doc.add_paragraph('目录')
    title = '响应承诺书' if kind == 'commitment_copy' else '施工组织设计'
    entry = doc.add_paragraph(title)
    _tab(entry, location)
    original = entry._p.xml
    doc.add_page_break()
    doc.add_paragraph(title)
    doc.add_paragraph('保留原正文。')
    path = tmp_path / 'draft.docx'
    doc.save(path)
    brief = BidBrief(outlineItems=[OutlineItem(
        id='draft', kind=kind, title=title, draftBody='本项目已生成正文。',
    )])
    assert fill(path, brief) == (1, [])
    out = Document(path)
    # The first matching title belongs to the untouched source TOC.
    assert out.paragraphs[1 if with_toc_heading else 0]._p.xml == original
    texts = [p.text for p in out.paragraphs]
    body_heading = max(i for i, text in enumerate(texts) if text == title)
    assert texts.index('本项目已生成正文。') > body_heading


@pytest.mark.parametrize('kind,fill', [
    ('commitment_copy', _fill_business_format_sections),
    ('tech_plan', _fill_drafted_format_sections),
])
def test_missing_body_heading_does_not_fall_back_to_toc(tmp_path, kind, fill):
    doc = Document()
    doc.add_paragraph('目录')
    _tab(doc.add_paragraph('响应承诺书'), 'run')
    doc.add_paragraph('施工组织设计')
    doc.add_page_break()
    doc.add_paragraph('其他章节')
    doc.add_paragraph('仅保留这段正文。')
    path = tmp_path / 'missing.docx'
    doc.save(path)
    original = path.read_bytes()
    title = '响应承诺书' if kind == 'commitment_copy' else '施工组织设计'
    brief = BidBrief(outlineItems=[OutlineItem(
        id='missing', kind=kind, title=title, draftBody='不应写进目录。',
    )])
    inserted, warnings = fill(path, brief)
    assert inserted == 0 and warnings
    assert path.read_bytes() == original


@pytest.mark.parametrize('volume', ['full', 'business', 'technical'])
@pytest.mark.parametrize('location', ['run', 'ppr'])
def test_assembled_drafts_stay_after_toc_and_fields_are_stable(
    tmp_path, monkeypatch, volume, location,
):
    doc = Document()
    doc.add_paragraph('投标文件')
    doc.add_page_break()
    doc.add_paragraph('目录')
    titles = ['一、响应承诺书', '二、施工组织设计']
    for title in titles:
        _tab(doc.add_paragraph(title), location)
    doc.add_page_break()
    doc.add_paragraph(titles[0])
    doc.add_paragraph('原商务内容。')
    doc.add_page_break()
    doc.add_paragraph(titles[1])
    doc.add_paragraph('原技术内容。')
    source = tmp_path / 'source.docx'
    doc.save(source)
    source.with_name(source.name + '.source').write_text('docx-raw', encoding='utf-8')
    original = source.read_bytes()
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: source)
    brief = BidBrief(
        layoutMode='outline', invitationId='abc123def456', generateVolume=volume,
        splitSlice=True, autoDraftBizResp=True, autoDraftTechPlan=True,
        attachQualifications=False, includePlaceholders=False, includeCommitment=False,
        outlineItems=[
            OutlineItem(id='biz', title='响应承诺书', rawTitle=titles[0],
                        kind='commitment_copy', draftBody='一、投标保证金\n商务正文已填充。'),
            OutlineItem(id='tech', title='施工组织设计', rawTitle=titles[1],
                        kind='tech_plan', draftBody='二、施工组织格式\n技术正文已填充。'),
        ],
    )
    path, warnings = assemble_bid_docx(
        brief, tmp_path / f'{volume}.docx', volume=volume,
        catalog_slots=[], catalog_media={}, catalog_facts=[],
    )
    assert not any('未找到格式章正文落位' in warning for warning in warnings)
    out = Document(path)
    block = toc_block_ids(out)
    toc_text = '\n'.join(Paragraph(el, out).text for el in block)
    assert '商务正文已填充' not in toc_text and '技术正文已填充' not in toc_text
    assert '一、投标保证金' not in toc_text and '二、施工组织格式' not in toc_text
    entries = [Paragraph(el, out) for el in block[1:] if _pageref_bookmark(Paragraph(el, out))]
    assert len(entries) == (2 if volume == 'full' else 1)
    assert all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)
    text = '\n'.join(p.text for p in out.paragraphs)
    assert ('商务正文已填充' in text) == (volume != 'technical')
    assert ('技术正文已填充' in text) == (volume != 'business')
    body_start = list(out.element.body).index(block[-1]) + 1
    hosts = {n.get(qn('w:name')) for el in list(out.element.body)[body_start:]
             for n in el.iter(qn('w:bookmarkStart'))}
    assert all(_pageref_bookmark(p) in hosts for p in entries)
    assert source.read_bytes() == original
    if volume == 'full':
        before = _signature(path)
        for _ in range(2):
            unify_full_docx(path, brief)
            assert _signature(path) == before
    names = [n.get(qn('w:name')) for n in Document(path).element.body.iter(qn('w:bookmarkStart'))]
    assert len(names) == len(set(names))


@pytest.mark.parametrize('operation', ['link', 'repair', 'unify'])
def test_toc_repair_does_not_promote_unanchored_draft_lines(tmp_path, operation):
    doc = Document()
    doc.add_paragraph('目录')
    entry = doc.add_paragraph()
    _write_toc_line(entry, 0, '响应承诺书', 'toc_biz', '1')
    stray = doc.add_paragraph('一、投标保证金')
    stray_xml = stray._p.xml
    doc.add_page_break()
    _bookmark_paragraph(doc.add_paragraph('响应承诺书'), 'toc_biz')
    # This looks like a heading, but has no body bookmark / original TOC reference.
    doc.add_paragraph('一、投标保证金')
    doc.add_paragraph('正文内的分项说明。')
    brief = BidBrief(outlineItems=[OutlineItem(id='biz', title='响应承诺书')])
    if operation == 'link':
        assert _link_toc_body(doc, brief) == 1
    elif operation == 'repair':
        _repair_format_table_toc(doc, brief)
    else:
        path = tmp_path / 'unify.docx'
        doc.save(path)
        unify_full_docx(path, brief)
        doc = Document(path)
    assert doc.paragraphs[2]._p.xml == stray_xml
    assert sum('PAGEREF' in (n.text or '') for n in doc.element.body.iter(qn('w:instrText'))) == 1
    assert len(list(doc.element.body.iter(qn('w:bookmarkStart')))) == 1


def test_repair_does_not_renumber_unreferenced_unbookmarked_body_heading():
    doc = Document()
    doc.add_paragraph('目录')
    entry = doc.add_paragraph('一、投标保证金')
    _tab(entry, 'run')
    before = entry._p.xml
    doc.add_page_break()
    doc.add_paragraph('一、投标保证金')
    brief = BidBrief(outlineItems=[OutlineItem(id='deposit', title='投标保证金')])
    assert _repair_format_table_toc(doc, brief) == 0
    assert entry._p.xml == before


def test_mixed_skeleton_toc_keeps_plain_original_outline_entry(tmp_path):
    doc = Document()
    doc.add_paragraph('目录')
    _tab(doc.add_paragraph('响应承诺书'), 'run')
    doc.add_paragraph('施工组织设计')  # Some original entries have no tab stops at all.
    doc.add_page_break()
    doc.add_paragraph('响应承诺书')
    doc.add_paragraph('商务原正文。')
    doc.add_page_break()
    doc.add_paragraph('施工组织设计')
    doc.add_paragraph('技术原正文。')
    brief = BidBrief(outlineItems=[
        OutlineItem(id='biz', title='响应承诺书'),
        OutlineItem(id='tech', title='施工组织设计'),
    ])
    path = tmp_path / 'mixed.docx'
    doc.save(path)
    unify_full_docx(path, brief)
    entries = [p for p in Document(path).paragraphs if _pageref_bookmark(p)]
    assert len(entries) == 2
    assert [p.text.rsplit('\t', 1)[-1] for p in entries] == ['1', '2']
    before = _signature(path)
    unify_full_docx(path, brief)
    assert _signature(path) == before


def test_repeated_linking_prefers_existing_body_anchor_over_draft_lookalike():
    doc = Document()
    doc.add_paragraph('目录')
    entry = doc.add_paragraph()
    _write_toc_line(entry, 0, '响应承诺书', 'toc_biz', '2')
    doc.add_page_break()
    lookalike = doc.add_paragraph('响应承诺书')
    doc.add_paragraph('前一章草稿中的同名说明。')
    doc.add_page_break()
    heading = doc.add_paragraph('响应承诺书')
    _bookmark_paragraph(heading, 'toc_biz')
    doc.add_paragraph('实际章节正文。')
    brief = BidBrief(outlineItems=[OutlineItem(id='biz', title='响应承诺书')])
    for _ in range(2):
        assert _link_toc_body(doc, brief, preserve_existing=True) == 1
        assert _pageref_bookmark(entry) == 'toc_biz'
        assert not list(lookalike._p.iter(qn('w:bookmarkStart')))
        assert len(list(doc.element.body.iter(qn('w:bookmarkStart')))) == 1
