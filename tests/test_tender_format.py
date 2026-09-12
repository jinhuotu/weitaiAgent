"""招标书排版条款：页边距 / 封面 / 目录页码 / 页脚 PAGE 域。"""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.format_rules import (
    extract_document_format,
    extract_tender_no,
    toc_item_label,
)
from api.services.tenders.schema import DocumentFormat, OutlineItem, default_brief


def test_extract_skips_unspecified_invitation() -> None:
    fmt = extract_document_format("第一章 投标邀请\n项目名称：某某充电站\n投标人须知")
    assert fmt.specified is False
    assert fmt.pageNumberPos == "bottom-center"
    assert fmt.pageNumberStart == "toc"
    assert fmt.tocNumbering == "cn"
    assert extract_tender_no("项目名称：某某") == ""


def test_extract_full_format_chapter() -> None:
    text = """
投标文件格式要求：
1. 封面（必填）项目全称、招标编号、投标人全称、“正本 / 副本” 字样、投标日期；封面加盖公章。
招标编号：WT-2026-088。
2. 目录层级一、1、(1)，目录页码和正文页码一一对应，不能错页。
3. 页码连续编码，不能跳号。页底居中。封面、扉页不编页码，从目录起编。
页边距均为 2.5cm。正文用宋体小四，一级标题用三号。
"""
    fmt = extract_document_format(text)
    assert fmt.specified is True
    assert fmt.marginLeftCm == 2.5
    assert fmt.marginTopCm == 2.5
    assert fmt.fontName == "宋体"
    assert fmt.bodySizePt == 12
    assert fmt.headingSizePt == 16
    assert fmt.coverShowTenderNo is True
    assert fmt.coverShowCopyMark is True
    assert fmt.coverNeedSeal is True
    assert fmt.tocNeedPageNos is True
    assert fmt.tocNumbering == "cn"
    assert fmt.pageNumberPos == "bottom-center"
    assert fmt.pageNumberStart == "toc"
    assert extract_tender_no(text) == "WT-2026-088"
    assert "封面含" in "；".join(fmt.notes)


def test_extract_bottom_right_and_named_margins() -> None:
    text = (
        "装订要求：上边距 3cm，下边距 2.5cm，左边距 2.8cm，右边距 2.6cm。"
        "页码置于右下角，从正文起编。"
    )
    fmt = extract_document_format(text)
    assert fmt.marginTopCm == 3.0
    assert fmt.marginBottomCm == 2.5
    assert fmt.pageNumberPos == "bottom-right"
    assert fmt.pageNumberStart == "body"


def test_toc_item_label_schemes() -> None:
    assert toc_item_label(0, "cn") == "一、"
    assert toc_item_label(1, "arabic") == "2."
    assert toc_item_label(0, "paren") == "(1)"
    assert toc_item_label(0, "attach") == "附件一"


def test_assemble_applies_format_cover_and_page_fields(tmp_path) -> None:
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "厂内充电桩采购项目"
    brief.tenderer = "某市热电有限公司"
    brief.tenderNo = "WT-2026-088"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.documentFormat = DocumentFormat(
        specified=True,
        marginLeftCm=2.5,
        marginRightCm=2.5,
        marginTopCm=2.5,
        marginBottomCm=2.5,
        coverShowTenderNo=True,
        coverShowCopyMark=True,
        coverCopyMark="正本",
        coverNeedSeal=True,
        tocNeedPageNos=True,
        pageNumberPos="bottom-right",
        pageNumberStart="toc",
        fontName="宋体",
    )
    brief.outlineItems = [
        OutlineItem(id="o01", title="授权委托书", kind="auth", source="generate"),
        OutlineItem(id="o02", title="法定代表人身份证明", kind="legal_id", source="generate"),
    ]
    dest = tmp_path / "fmt.docx"
    _, warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    assert any("封面加盖公章" in w for w in warnings)
    doc = Document(str(dest))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    blob = "\n".join(texts)
    assert "厂内充电桩采购项目" in blob
    assert "招标编号：WT-2026-088" in blob
    assert "正本" in blob
    assert "（封面加盖公章）" in blob
    assert any(t.startswith("一、授权委托书") for t in texts)
    cover = doc.sections[0]
    rest = doc.sections[1]
    assert abs(float(cover.left_margin.cm) - 2.5) < 0.05
    assert abs(float(rest.left_margin.cm) - 2.5) < 0.05
    cover_footer = cover.footer.paragraphs[0]._p.xml
    assert "PAGE" not in cover_footer
    rest_xml = rest.footer.paragraphs[0]._p.xml
    assert "PAGE" in rest_xml
    assert rest.footer.paragraphs[0].alignment is not None
    bookmarks = {el.get(qn("w:name")) for el in doc.element.iter(qn("w:bookmarkStart"))}
    assert "toc_o01" in bookmarks
    pagerefs = [
        "".join(t.text or "" for t in el.iter(qn("w:instrText")))
        for el in doc.element.iter(qn("w:instrText"))
    ]
    assert any("PAGEREF toc_o01" in x for x in pagerefs)


def test_assemble_cover_meta_sits_in_lower_third(tmp_path) -> None:
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "厂内新能源充电桩采购项目"
    brief.tenderer = "二连浩特市联源热电有限公司"
    brief.bidderName = "河南伟泰光电科技有限公司"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.documentFormat = DocumentFormat(specified=False)
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate"),
    ]
    dest = tmp_path / "cover-meta.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    body = doc.element.body
    events: list[tuple[str, str]] = []
    for child in body:
        tag = child.tag
        if tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t"))).replace("\u200b", "").strip()
            p_pr = child.find(qn("w:pPr"))
            line_twips = 0
            if p_pr is not None:
                sp = p_pr.find(qn("w:spacing"))
                if sp is not None and (sp.get(qn("w:lineRule")) or "") == "exact":
                    line_twips = int(sp.get(qn("w:line")) or 0)
            if line_twips >= 240 and not text:
                events.append(("gap", str(line_twips)))
            elif text:
                events.append(("p", text))
        if any(e[0] == "p" and e[1] == "目录" for e in events):
            break
    texts = [t for k, t in events if k == "p"]
    assert any("厂内新能源充电桩采购项目" in t for t in texts)
    assert any(t.replace(" ", "") == "投标文件" for t in texts)
    title_i = next(i for i, (k, t) in enumerate(events) if k == "p" and "投 标 文 件" in t)
    meta_i = next(i for i, (k, t) in enumerate(events) if k == "p" and t.startswith("招标人："))
    spacer_i = next(i for i, (k, t) in enumerate(events) if k == "gap" and title_i < i < meta_i)
    spacer_twips = int(events[spacer_i][1])
    assert spacer_twips >= int(4.0 * 567)
    assert any(t.startswith("投标人：") for t in texts)


def test_assemble_default_skips_cover_page_number(tmp_path) -> None:
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate"),
    ]
    dest = tmp_path / "default-fmt.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    assert len(doc.sections) >= 2
    assert "PAGE" not in doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in doc.sections[1].footer.paragraphs[0]._p.xml
    assert abs(float(doc.sections[0].left_margin.cm) - 2.8) < 0.05
    assert abs(float(doc.sections[0].page_width.cm) - 21.0) < 0.05


def test_toc_item_label_uses_chinese_past_ten() -> None:
    assert toc_item_label(0, "cn") == "一、"
    assert toc_item_label(9, "cn") == "十、"
    assert toc_item_label(10, "cn") == "十一、"
    assert toc_item_label(16, "cn") == "十七、"
