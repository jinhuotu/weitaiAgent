"""资料事实：存取和填空。不连模型。"""

from docx import Document

from api.services.tenders.facts import (
    apply_correction,
    dump_facts,
    fact_fill,
    load_facts,
    view_facts,
)
from api.services.tenders.id_card import dump_id_meta, load_id_meta
from api.services.tenders.performance import PerformanceLine, dump_perf_meta, load_perf_meta
from api.services.tenders.schema import default_brief


def _company(code: str) -> dict:
    return {
        "role": "company",
        "fields": [{"name": "统一社会信用代码", "value": code, "conf": 0.9}],
    }


def test_facts_roundtrip_drops_low_conf_and_price() -> None:
    raw = dump_facts(
        {
            "role": "company",
            "fields": [
                {"name": "统一社会信用代码", "value": "91410000MA0000000X", "conf": 0.91},
                {"name": "注册资本", "value": "1000万元", "conf": 0.4},
                {"name": "投标总价", "value": "999", "conf": 0.99},
            ],
        }
    )
    assert raw.startswith("FACTS:")
    hit = load_facts(raw)
    assert hit is not None
    names = [item["name"] for item in hit["fields"]]
    assert names == ["统一社会信用代码"]
    assert view_facts(raw)["source"] == "llm"


def test_view_id_and_perf_without_rewriting() -> None:
    card_raw = dump_id_meta(
        {
            "side": "front",
            "name": "郭志伟",
            "gender": "男",
            "idNo": "410521199802104053",
            "age": "28",
            "personKey": "p1",
        }
    )
    viewed = view_facts(card_raw)
    assert viewed["role"] == "person/legal"
    assert viewed["source"] == "idcard"
    assert viewed["fields"][0]["value"] == "郭志伟"
    assert load_id_meta(card_raw)["personKey"] == "p1"
    assert card_raw.startswith("idcard:")

    line = PerformanceLine(projectName="某充电站", client="甲公司", amountYuan=200000)
    perf_raw = dump_perf_meta(line)
    perf_view = view_facts(perf_raw)
    assert perf_view["role"] == "performance"
    assert perf_view["source"] == "perf"
    assert load_perf_meta(perf_raw).projectName == "某充电站"
    assert perf_raw.startswith("PERFJSON:")


def test_correction_keeps_prefix() -> None:
    card_raw = dump_id_meta({"side": "front", "name": "郭志伟", "gender": "男", "personKey": "p1"})
    nxt = apply_correction(
        card_raw,
        role="company",
        fields=[{"name": "姓名", "value": "李四"}, {"name": "性别", "value": "女"}],
    )
    assert nxt.startswith("idcard:")
    assert load_id_meta(nxt)["name"] == "李四"
    assert load_id_meta(nxt)["personKey"] == "p1"

    facts_raw = dump_facts(_company("91410000MA0000000X"))
    edited = apply_correction(
        facts_raw,
        role="company",
        fields=[{"name": "统一社会信用代码", "value": "91410000MA1111111Y"}],
    )
    assert edited.startswith("FACTS:")
    assert load_facts(edited)["fields"][0]["value"] == "91410000MA1111111Y"


def test_one_license_fills_credit_code() -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "统一社会信用代码"
    from api.services.tenders.pdf_format import _fill_empty_label_cells

    notes = _fill_empty_label_cells(table._tbl, default_brief(), [_company("91410000MA0000000X")])
    assert table.rows[0].cells[1].text.strip() == "91410000MA0000000X"
    assert notes == []


def test_two_licenses_stay_blank() -> None:
    value, note = fact_fill(
        "统一社会信用代码",
        "company",
        [_company("91410000MA0000000X"), _company("91410000MA1111111Y")],
    )
    assert value == ""
    assert note == "多份未选定，未填统一社会信用代码"


def test_legal_fact_does_not_fill_pm_name() -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "拟在本工程任职"
    table.rows[0].cells[1].text = "项目经理"
    table.rows[1].cells[0].text = "姓名"
    facts = [
        {
            "role": "person/legal",
            "fields": [{"name": "姓名", "value": "郭志伟", "conf": 1}],
        }
    ]
    from api.services.tenders.pdf_format import _fill_empty_label_cells

    _fill_empty_label_cells(table._tbl, default_brief(), facts)
    assert table.rows[1].cells[1].text.strip() == ""

    facts.append(
        {"role": "person/pm", "fields": [{"name": "姓名", "value": "张三", "conf": 0.95}]}
    )
    _fill_empty_label_cells(table._tbl, default_brief(), facts)
    assert table.rows[1].cells[1].text.strip() == "张三"


def test_schedule_not_from_library() -> None:
    facts = [{"role": "company", "fields": [{"name": "工期", "value": "30天", "conf": 1}]}]
    assert fact_fill("工期", "company", facts) == ("", "")


def test_appendix_pm_from_facts() -> None:
    from api.services.tenders.document import _fill_appendix_table

    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "项目经理"
    facts = [{"role": "person/pm", "fields": [{"name": "姓名", "value": "张三", "conf": 1}]}]
    _fill_appendix_table(table, default_brief(), "零元", "0", facts)
    labels = [row.cells[0].text.strip() for row in table.rows]
    hit = next(row for row in table.rows if "项目经理" in row.cells[0].text)
    assert "张三" in hit.cells[-1].text
    assert "郭志伟" not in "".join(labels)
