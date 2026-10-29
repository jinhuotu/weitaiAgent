"""Requirement-driven tenders retain complete structure without a source TOC."""
import asyncio
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders import requirement_bid
from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.document import _pageref_bookmark, build_bid_docx
from api.services.tenders.generate import generate_pair
from api.services.tenders.records import _prepare_and_generate, require_format_confirm
from api.services.tenders.schema import BidBrief, BidRequirement, OutlineItem
from common.errors import AppError


def _brief():
    return BidBrief(
        invitationId='deadbeef1234', layoutMode='outline', generateVolume='full',
        projectName='要求型采购项目', tenderer='采购单位', bidPriceYuan=100,
        bidRequirements=[BidRequirement(
            id='q1', text='投标人须提供营业执照及有效的财务、税收和社保证明',
            category='qualification', targetKind='scan',
        )],
    )


@pytest.fixture
def missing_format(monkeypatch):
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: None)
    monkeypatch.setattr(requirement_bid, 'invitation_text', lambda _: '')


@pytest.mark.parametrize('cover,toc', [(False, False), (False, True), (True, False), (True, True)])
def test_format_completeness_requires_cover_and_actual_toc(tmp_path, monkeypatch, cover, toc):
    doc = Document()
    if cover:
        doc.add_paragraph('要求型采购项目 投标文件')
        doc.add_paragraph('项目编号：0001 投标人：供应商（盖章）')
    if toc:
        doc.add_paragraph('目录')
        doc.add_paragraph('投标函').add_run().add_tab()
    doc.add_page_break()
    doc.add_paragraph('投标函')
    # A body tab alone is not a TOC skeleton.
    doc.add_paragraph('签字：').add_run().add_tab()
    path = tmp_path / 'source.docx'
    doc.save(path)
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: path)
    complete, _ = requirement_bid.format_completeness(_brief())
    assert complete == (cover and toc)
    brief = _brief()
    notes = requirement_bid.prepare_requirements_bid(brief)
    assert bool(notes) == (not complete)
    assert brief.requirementsDriven == (not complete)


def test_synthesis_preserves_edits_skips_ids_and_uses_full_text(missing_format):
    brief = _brief()
    brief.outlineItems = [
        OutlineItem(id='manual-quality', title='质量保证方案', kind='tech_chapter',
                    draftBody='用户确认的质量管理正文。', draftStatus='edited'),
        OutlineItem(id='skip-auth', title='授权委托书', kind='auth', skipped=True),
    ]
    text = '\n'.join(['普通通知内容' * 10] * 60 + [
        '产品须取得医疗器械注册证；投标人提供医疗器械经营许可证。',
        '供应商须提供救护车公告目录证明。',
        '商务标得分为评委打分之和，不得作为单独提交章节。',
    ])
    notes = requirement_bid.prepare_requirements_bid(brief, text=text)
    assert notes and '补齐章节' in notes[1]
    assert any(i.title == '医疗器械注册及备案证明' for i in brief.outlineItems)
    assert any(i.title == '救护车公告目录证明' for i in brief.outlineItems)
    assert not any('评委打分' in i.title for i in brief.outlineItems)
    quality = next(i for i in brief.outlineItems if i.id == 'manual-quality')
    assert quality.draftBody == '用户确认的质量管理正文。'
    assert next(i for i in brief.outlineItems if i.id == 'skip-auth').skipped
    before = [i.model_dump() for i in brief.outlineItems]
    requirement_bid.prepare_requirements_bid(brief, text=text)
    assert [i.model_dump() for i in brief.outlineItems] == before
    assert len({i.id for i in brief.outlineItems}) == len(brief.outlineItems)


def test_classified_target_table_is_reflected_in_outline(missing_format):
    brief = _brief()
    brief.bidRequirements.append(BidRequirement(
        text='投标人应提交设备配置明细表', targetTitle='设备配置明细表', targetKind='company',
        category='material',
    ))
    requirement_bid.prepare_requirements_bid(brief)
    assert any(i.title == '设备配置明细表' and i.kind == 'company' for i in brief.outlineItems)


def _assert_complete(path, *, expect_technical=True, expect_business=True):
    doc = Document(path)
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert entries
    assert all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)
    starts = list(doc.element.body.iter(qn('w:bookmarkStart')))
    names = [n.get(qn('w:name')) for n in starts]
    ids = [n.get(qn('w:id')) for n in starts]
    assert len(names) == len(set(names)) and len(ids) == len(set(ids))
    assert set(ids) == {n.get(qn('w:id')) for n in doc.element.body.iter(qn('w:bookmarkEnd'))}
    assert all(_pageref_bookmark(p) in names for p in entries)
    lines = [p.text for p in doc.paragraphs]
    # Every directory row links into the body after all TOC entries.
    tail = max(list(doc.element.body).index(p._p) for p in entries)
    hosts = {n.get(qn('w:name')) for el in list(doc.element.body)[tail + 1:]
             for n in el.iter(qn('w:bookmarkStart'))}
    assert all(_pageref_bookmark(p) in hosts for p in entries)
    assert '目录' in lines
    assert any('PAGE' in (n.text or '') for sec in doc.sections
               for n in sec.footer._element.iter(qn('w:instrText')))
    assert ('质量保证方案' in lines) == expect_technical
    assert ('营业执照' in lines) == expect_business
    text = '\n'.join(lines)
    assert any(word in text for word in ('盖章', '公章')) and '签字' in text
    return doc


@pytest.mark.parametrize('entry', ['assemble', 'build'])
def test_requirement_generation_owns_toc_and_both_drafts(
    tmp_path, monkeypatch, missing_format, entry,
):
    brief = _brief()
    brief.outlineItems = [
        OutlineItem(id='quality', title='质量保证方案', kind='tech_chapter',
                    draftBody='已确认的技术方案正文。', draftStatus='edited'),
        OutlineItem(id='commit', title='投标承诺函', kind='commitment_copy',
                    draftBody='已确认的商务响应正文。', draftStatus='edited'),
    ]
    # Requirement rendering must not use either a fixed template or the incomplete source.
    monkeypatch.setattr('api.services.tenders.document.find_chapter5_template', lambda: None)
    def unexpected(*args, **kwargs):
        pytest.fail('Incomplete format must bypass copying and native TOC rebuilding')
    for name in ('assemble_from_word_format', '_link_toc_body', '_repair_format_table_toc'):
        monkeypatch.setattr(f'api.services.tenders.pdf_format.{name}', unexpected)
    if entry == 'build':
        requirement_bid.prepare_requirements_bid(brief)
    function = build_bid_docx if entry == 'build' else assemble_bid_docx
    path, notes = function(brief, tmp_path / 'full.docx', volume='full',
                           catalog_slots=[], catalog_media={}, catalog_facts=[])
    doc = _assert_complete(path)
    lines = [p.text for p in doc.paragraphs]
    groups = [title for _, title in requirement_bid._GROUPS]
    positions = [lines.index(title) for title in groups]
    assert positions == sorted(positions)
    first_body = lines.index(groups[0])
    assert lines.index('已确认的商务响应正文。') > first_body
    assert lines.index('已确认的技术方案正文。') > first_body
    assert any('完整性结构自检' in note for note in notes)
    assert any('未提供分项报价清单' in note for note in notes)
    assert '【待补充】' in '\n'.join(lines)
    assert any('投标总价' in p.text and '100' in p.text for p in doc.paragraphs)


def test_requirement_split_has_two_separate_volumes(tmp_path, monkeypatch, missing_format):
    brief = _brief()
    requirement_bid.prepare_requirements_bid(brief)
    monkeypatch.setattr('api.services.tenders.generate.tenders_output_dir', lambda: tmp_path)
    result = generate_pair(brief, catalog_slots=[], catalog_media={}, catalog_facts=[])
    assert result['docxFile'] and result['techDocxFile']
    _assert_complete(tmp_path / result['docxFile'], expect_technical=False)
    _assert_complete(tmp_path / result['techDocxFile'], expect_business=False)


def test_legacy_missing_format_without_requirements_remains_explicit_error(
    missing_format, tmp_path,
):
    brief = _brief()
    brief.bidRequirements = []
    with pytest.raises(AppError, match='未截到投标文件格式章'):
        assemble_bid_docx(brief, tmp_path / 'missing.docx', volume='full')


def test_requirement_preparation_precedes_record_validation(missing_format, monkeypatch):
    brief = _brief()
    def validate(prepared):
        assert prepared is brief
        assert prepared.requirementsDriven
        assert any(i.kind == 'tech_chapter' for i in prepared.outlineItems)
        require_format_confirm(prepared)
        return ['测试在访问数据库前结束']
    monkeypatch.setattr('api.services.tenders.records._brief_generate_issues', validate)
    with pytest.raises(AppError, match='测试在访问数据库前结束'):
        asyncio.run(_prepare_and_generate(None, brief))


def test_real_hospital_source_preserves_native_structure_and_supplements(tmp_path):
    source = Path(__file__).parents[1] / 'storage/tender-invitations/ba636a37878c.txt'
    if not source.is_file():
        pytest.skip('Local hospital acceptance fixture is not installed')
    brief = BidBrief(
        invitationId=source.stem, layoutMode='outline', generateVolume='full',
        projectName='河南省胸科医院车辆采购（救护车）项目（二次）', tenderer='河南省胸科医院',
    )
    original = source.with_suffix('.format.docx').read_bytes()
    path, notes = assemble_bid_docx(brief, tmp_path / 'hospital-full.docx', volume='full',
                                   catalog_slots=[], catalog_media={}, catalog_facts=[])
    doc = Document(path)
    assert brief.layoutMode == 'outline' and not brief.requirementsDriven
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert entries and all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)
    names = [n.get(qn('w:name')) for n in doc.element.body.iter(qn('w:bookmarkStart'))]
    ids = [n.get(qn('w:id')) for n in doc.element.body.iter(qn('w:bookmarkStart'))]
    assert len(names) == len(set(names)) and len(ids) == len(set(ids))
    assert all(_pageref_bookmark(p) in names for p in entries)
    assert len(doc.tables) >= len(Document(source.with_suffix('.format.docx')).tables)
    assert any('原版' in note for note in notes)
    assert any('缺少' in note and '补齐' in note for note in notes)
    assert any('投标总价未填写' in note for note in notes)
    assert source.with_suffix('.format.docx').read_bytes() == original


def test_native_content_control_toc_is_a_complete_format(tmp_path, monkeypatch):
    doc = Document()
    doc.add_paragraph('要求型采购项目 投标文件')
    doc.add_paragraph('项目编号：0001 投标人：供应商（盖章）')
    heading = doc.add_paragraph('目录')
    entry = doc.add_paragraph('投标函')
    entry.add_run().add_tab()
    control = OxmlElement('w:sdt')
    content = OxmlElement('w:sdtContent')
    control.append(content)
    heading._p.addprevious(control)
    content.append(heading._p)
    content.append(entry._p)
    doc.add_page_break()
    doc.add_paragraph('投标函')
    path = tmp_path / 'native.docx'
    doc.save(path)
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: path)
    assert requirement_bid.format_completeness(_brief()) == (True, [])


def test_direct_builder_can_prepare_an_empty_requirement_outline(tmp_path, missing_format):
    brief = _brief()
    brief.requirementsDriven = True
    brief.layoutMode = 'chapter5'
    path, notes = build_bid_docx(brief, tmp_path / 'direct.docx', volume='full',
                                catalog_slots=[], catalog_media={}, catalog_facts=[])
    _assert_complete(path)
    assert any('补齐章节' in note for note in notes)


def test_synthesis_drops_old_volume_labels_and_preserves_template_commitment(missing_format):
    brief = _brief()
    brief.outlineItems = [
        OutlineItem(id='biz-label', title='商务标'),
        OutlineItem(id='tech-label', title='技术标'),
        OutlineItem(id='commitment', title='投标承诺函', kind='commitment_copy',
                    body='投标承诺函\n投标人：________（盖章）\n日期：____年__月__日'),
    ]
    requirement_bid.prepare_requirements_bid(brief)
    assert not any(item.title in {'商务标', '技术标'} for item in brief.outlineItems)
    assert next(item for item in brief.outlineItems if item.id == 'commitment').body


def test_completeness_reports_explicitly_skipped_required_section(tmp_path, missing_format):
    brief = _brief()
    brief.outlineItems = [OutlineItem(id='auth', title='授权委托书', kind='auth', skipped=True)]
    path, notes = assemble_bid_docx(brief, tmp_path / 'skipped.docx', volume='full',
                                   catalog_slots=[], catalog_media={}, catalog_facts=[])
    _assert_complete(path)
    assert any('已手动跳过必需章节' in note and '授权委托书' in note for note in notes)


def test_named_policy_forms_after_requirement_cap_are_pending_sections(missing_format, tmp_path):
    brief = _brief()
    text = '\n'.join(['评审通知内容。'] * 60 + [
        '适用时提交《中小企业声明函》。',
        '提供《关于符合本国产品标准的声明函》，内容须经投标人核实。',
        '《政府采购管理办法》是规则文件，不是提交章节。',
    ])
    requirement_bid.prepare_requirements_bid(brief, text=text)
    for title in ('中小企业声明函', '关于符合本国产品标准的声明函'):
        item = next(i for i in brief.outlineItems if i.title == title)
        assert item.kind == 'commitment_copy' and item.rawTitle == title
        assert not item.body and not item.draftBody  # Never fabricate a policy eligibility claim.
    assert not any('管理办法' in i.title for i in brief.outlineItems)
    path, _ = assemble_bid_docx(brief, tmp_path / 'policy.docx', volume='full',
                                catalog_slots=[], catalog_media={}, catalog_facts=[])
    _assert_complete(path)
