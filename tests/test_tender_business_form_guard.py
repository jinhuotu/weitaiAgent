import asyncio
import json

import pytest
from docx import Document

from api.services.tenders.assemble import _fill_business_format_sections
from api.services.tenders.business_draft_policy import chapter_paragraphs
from api.services.tenders.gen_biz_resp import draft_business_response
from api.services.tenders.generate import generate_bid
from api.services.tenders.native_format import validate_native_toc
from api.services.tenders.schema import BidBrief, BidRequirement, OutlineItem

ORIGINAL = "我方在此声明，我方拟派项目经理现阶段没有担任其他在建工程建设项目的项目经理。"
BAD = "本承诺书章节汇总本批需由投标人出具的书面承诺。现阶段按招标文件要求作概括响应"
LEGITIMATE = "按本项目合同约定组织交付验收，落实履约衔接安排。"
REQUIREMENT = "付款方式必须按本项目合同约定执行"


@pytest.fixture
def native_form(tmp_path, monkeypatch):
    doc = Document()
    doc.add_paragraph("投标文件")
    doc.add_paragraph("目录")
    for title in ("一、承诺书", "二、响应承诺书", "三、下一章"):
        doc.add_paragraph(title).add_run().add_tab()
    doc.add_page_break()
    doc.add_paragraph("一、承诺书", style="Heading 1")
    doc.add_paragraph(ORIGINAL)
    doc.add_paragraph("特此承诺")
    doc.add_paragraph("投标人（盖章）：")
    doc.add_paragraph("年 月 日")
    doc.add_paragraph("二、响应承诺书", style="Heading 1")
    doc.add_paragraph("请说明本项目商务响应。")
    doc.add_paragraph("三、下一章", style="Heading 1")
    doc.add_paragraph("下一章原文。")
    path = tmp_path / "source.format.docx"
    doc.save(path)
    path.with_name(path.name + ".source").write_text("docx-raw", encoding="utf-8")
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: path)
    monkeypatch.setattr("api.services.tenders.generate.tenders_output_dir", lambda: tmp_path)
    monkeypatch.setattr(
        "api.services.tenders.gen_biz_resp.load_invitation_text", lambda _: REQUIREMENT
    )
    return path


def brief():
    return BidBrief(
        invitationId="form-audit", layoutMode="outline", outlineChapter="第八章 投标文件格式",
        autoDraftBizResp=True, attachQualifications=False,
        outlineItems=[
            OutlineItem(id="fixed", title="承诺书", rawTitle="一、承诺书", kind="commitment_copy",
                        body="承诺书", draftBody=BAD, draftStatus="drafted"),
            OutlineItem(id="response", title="响应承诺书", rawTitle="二、响应承诺书",
                        kind="commitment_copy"),
            OutlineItem(id="next", title="下一章", rawTitle="三、下一章"),
        ],
        bidRequirements=[BidRequirement(text=REQUIREMENT, volume="business")],
    )


def test_fixed_form_clears_stale_automatic_prose_and_bypasses_model(native_form):
    data = brief()
    data.outlineItems = data.outlineItems[:1]
    before = data.model_dump()

    class Unexpected:
        async def complete(self, *_args, **_kwargs):
            pytest.fail("A fixed project-manager declaration must not invoke the LLM")

    result, notes = asyncio.run(draft_business_response(None, data, client=Unexpected()))
    assert result.outlineItems[0].draftBody == ""
    assert result.outlineItems[0].body == "承诺书"
    assert data.model_dump() == before
    assert any("原版固定承诺" in note for note in notes)


@pytest.mark.parametrize("volume", ["business", "full"])
def test_fixed_form_and_explicit_response_are_separate_in_final_bid(native_form, volume):
    original = native_form.read_bytes()

    class Client:
        async def complete(self, messages, **_kwargs):
            prompt = messages[1]["content"]
            titles = prompt.split("【仅允许填充的既有章节】\n", 1)[1].split("\n【", 1)[0]
            assert titles == "响应承诺书"
            assert "请说明本项目商务响应。" in prompt
            return json.dumps({"clauses": [], "chapters": [
                {"title": "承诺书", "body": BAD},
                {"title": "响应承诺书", "body": LEGITIMATE},
            ]}, ensure_ascii=False)

    data, _ = asyncio.run(draft_business_response(None, brief(), client=Client()))
    data.generateVolume = volume
    result = generate_bid(data, catalog_slots=[], catalog_media={})
    doc = Document(native_form.parent / result["docxFile"])
    texts = [p.text for p in doc.paragraphs]
    assert texts.count(ORIGINAL) == 1
    assert not any("本批" in text or "概括响应" in text for text in texts)
    assert texts.index(ORIGINAL) < texts.index("二、响应承诺书") < texts.index(LEGITIMATE)
    assert validate_native_toc(doc)["entries"] == 3
    assert native_form.read_bytes() == original


def test_complete_explicit_form_is_protected_using_actual_word_body(native_form):
    doc = Document(native_form)
    for paragraph in doc.paragraphs:
        if paragraph.text == "请说明本项目商务响应。":
            paragraph.text = (
                "我方承诺按照招标文件和合同约定履行本项目全部义务，"
                "并承担相应责任，特此承诺。"
            )
    doc.save(native_form)
    data = brief()
    data.outlineItems[1].body = "响应承诺书"
    data.outlineItems[1].draftBody = BAD
    data.outlineItems[1].draftStatus = "drafted"

    class Unexpected:
        async def complete(self, *_args, **_kwargs):
            pytest.fail("Existing full Word form must override title-only outline metadata")

    drafted, _ = asyncio.run(draft_business_response(None, data, client=Unexpected()))
    assert all(not item.draftBody for item in drafted.outlineItems)
    before = native_form.read_bytes()
    assert _fill_business_format_sections(native_form, data) == (0, [])
    assert native_form.read_bytes() == before


def test_writer_blocks_stale_fixed_form_draft_without_running_drafter(native_form):
    data = brief()
    before = native_form.read_bytes()
    assert _fill_business_format_sections(native_form, data) == (0, [])
    assert native_form.read_bytes() == before


def test_instructions_requesting_a_response_are_not_a_completed_declaration(native_form):
    doc = Document(native_form)
    instruction = "投标人应对付款、交付和验收作出响应承诺，并说明本项目商务服务安排。"
    for paragraph in doc.paragraphs:
        if paragraph.text == "请说明本项目商务响应。":
            paragraph.text = instruction
    doc.save(native_form)

    class Client:
        async def complete(self, messages, **_kwargs):
            assert instruction in messages[1]["content"]
            return json.dumps({"chapters": [{"title": "响应承诺书", "body": LEGITIMATE}]},
                              ensure_ascii=False)

    data, _ = asyncio.run(draft_business_response(None, brief(), client=Client()))
    assert data.outlineItems[1].draftBody == LEGITIMATE


def test_manual_specialized_supplement_is_retained(native_form):
    data = brief()
    data.outlineItems[0].draftStatus = "manual"
    data.outlineItems[0].draftBody = "人工确认的项目经理专项补充内容。"
    inserted, _ = _fill_business_format_sections(native_form, data)
    assert inserted == 1
    paragraphs = [p.text for p in Document(native_form).paragraphs]
    assert paragraphs.count(data.outlineItems[0].draftBody) == 1
    assert _fill_business_format_sections(native_form, data) == (0, [])


@pytest.mark.parametrize("merge_ok", [True, False])
def test_multiple_batches_are_composed_once_instead_of_concatenated(monkeypatch, merge_ok):
    monkeypatch.setattr("api.services.tenders.gen_biz_resp.load_invitation_text", lambda _: "")
    data = BidBrief(
        outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")],
        bidRequirements=[BidRequirement(text=f"付款要求{i}必须按合同执行", volume="business")
                         for i in range(21)],
    )
    candidates = ["按合同约定衔接付款流程。", "依照采购要求组织交付验收。"]
    merged = "按照合同约定衔接付款、交付和验收流程，落实本项目履约安排。"

    class Client:
        calls = 0

        async def complete(self, messages, **_kwargs):
            self.calls += 1
            if self.calls <= 2:
                body = candidates[self.calls - 1]
            else:
                assert "章节统稿" in messages[0]["content"]
                assert all(text in messages[1]["content"] for text in candidates)
                if not merge_ok:
                    return "{}"
                body = merged
            return json.dumps({"chapters": [{"title": "响应承诺书", "body": body}]},
                              ensure_ascii=False)

    client = Client()
    drafted, notes = asyncio.run(draft_business_response(None, data, client=client))
    assert client.calls == 3
    assert len(drafted.bizDeviationLines) == 21
    assert drafted.outlineItems[0].draftBody == (merged if merge_ok else candidates[0])
    assert any("统稿" in note or "统一整理" in note for note in notes)


def test_generation_instructions_are_filtered_from_formal_prose():
    body = "\n\n".join([
        BAD,
        "我方不在本承诺书中作出无依据的资质、业绩、金额、优惠或合同条件承诺。",
        "我方已阅读本批招标商务合同条款，拟履行承包人义务。",
        LEGITIMATE,
    ])
    assert chapter_paragraphs(body) == [LEGITIMATE]
