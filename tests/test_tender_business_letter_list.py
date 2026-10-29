import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders.native_format import ensure_toc_rows
from api.services.tenders.pdf_format import _drop_technical
from api.services.tenders.schema import BidBrief, OutlineItem

PARTS = [
    "投标函",
    "投标函附录",
    "法定代表人身份证明",
    "授权委托书",
    "投标保证金",
    "施工组织设计",
    "项目管理机构",
    "拟分包项目情况表",
    "资格审查资料",
]


@pytest.mark.parametrize("prefix", ["({n})", "（{n}）"])
@pytest.mark.parametrize("style", [None, "Heading 2"])
@pytest.mark.parametrize("native", [False, True])
@pytest.mark.parametrize("signoff", [False, True])
def test_business_slice_preserves_letter_composition_list(prefix, style, native, signoff):
    doc = Document()
    doc.add_paragraph("目录")
    for title in ("一、投标函", "二、施工组织设计", "三、资格审查资料"):
        doc.add_paragraph(title).add_run().add_tab()
    doc.add_page_break()
    doc.add_paragraph("一、投标函", style="Heading 1")
    doc.add_paragraph("致：采购人")
    doc.add_paragraph("2. 投标文件组成如下：")
    doc.add_paragraph("随附施工组织设计")
    items = []
    for n, title in enumerate(PARTS, 1):
        text = prefix.format(n=n) + title
        items.append(text)
        doc.add_paragraph(text, style=style)
    if signoff:
        doc.add_paragraph("投标人：（盖章）")
    doc.add_paragraph("二、施工组织设计", style="Heading 1")
    doc.add_paragraph("独立技术正文应移出商务卷。")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "技术资源表应移出"
    doc.add_paragraph("三、资格审查资料", style="Heading 1")
    doc.add_paragraph("资格证明正文必须保留。")
    brief = BidBrief(
        layoutMode="outline",
        outlineChapter="第八章 投标文件格式",
        outlineItems=[
            OutlineItem(id="letter", title="投标函", kind="letter"),
            OutlineItem(id="tech", title="施工组织设计", kind="tech_plan"),
            OutlineItem(id="qual", title="资格审查资料"),
        ],
    )
    _drop_technical(doc, brief=brief if native else None)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(item in text for item in items)
    assert "随附施工组织设计" in text
    assert "独立技术正文" not in text and not doc.tables
    assert "资格证明正文必须保留" in text
    once = doc.element.xml
    _drop_technical(doc, brief=brief if native else None)
    assert doc.element.xml == once
    if native:
        ensure_toc_rows(doc, brief)
        assert sum(p.text == "目录" for p in doc.paragraphs) == 1
        assert all(item in "\n".join(p.text for p in doc.paragraphs) for item in items)


def test_letter_list_alone_is_not_a_technical_chapter():
    doc = Document()
    doc.add_paragraph("投标函")
    doc.add_paragraph("2. 投标文件组成如下：")
    for n, title in enumerate(PARTS, 1):
        doc.add_paragraph(f"({n}){title}", style="Heading 2")
    doc.add_paragraph("函件最后一行必须保留。")
    before = doc.element.xml
    note = _drop_technical(doc)
    assert "没有单独" in note
    assert doc.element.xml == before


def test_incomplete_toc_does_not_hide_plain_letter_boundary():
    doc = Document()
    doc.add_paragraph("目录")
    doc.add_paragraph("五、已标价工程量清单").add_run().add_tab()
    doc.add_paragraph("一、投标函")
    doc.add_paragraph("致：采购人")
    for n, title in enumerate(PARTS, 1):
        doc.add_paragraph(f"({n}){title}", style="Heading 2")
    doc.add_paragraph("二、施工组织设计", style="Heading 1")
    doc.add_paragraph("独立技术正文。")
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    _drop_technical(doc, brief=brief)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(f"({n}){title}" in text for n, title in enumerate(PARTS, 1))
    assert "独立技术正文。" not in text


@pytest.mark.parametrize(
    "signature",
    ["法定代表人：（签字）", "投标人：（盖单位公章）", "法定代表人或其委托代理人：（签字）"],
)
def test_word_automatic_numbering_inside_letter_is_protected(signature):
    doc = Document()
    doc.add_paragraph("投标函", style="Heading 1")
    for text in PARTS:
        p = doc.add_paragraph(text)
        num = OxmlElement("w:numPr")
        ident = OxmlElement("w:numId")
        ident.set(qn("w:val"), "1")
        num.append(ident)
        p._p.get_or_add_pPr().append(num)
    doc.add_paragraph(signature)
    doc.add_paragraph("施工组织设计")
    doc.add_paragraph("独立技术正文。")
    _drop_technical(doc)
    assert [p.text for p in doc.paragraphs][1:10] == PARTS
    assert "独立技术正文。" not in [p.text for p in doc.paragraphs]
