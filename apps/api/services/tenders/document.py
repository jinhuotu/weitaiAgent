"""克隆第五章空白 Word，按表单填空。LLM 只抽邀请书字段，不直接写 docx。"""

from __future__ import annotations

import io
import logging
import re
import shutil
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from api.services.tenders.assets import find_chapter5_template
from api.services.tenders.categories import tech_plan_body, tech_plan_text
from api.services.tenders.commitment import append_commitment_letter
from api.services.tenders.format_rules import (
    apply_footer_page_number,
    apply_section_page,
    flatten_toc_tree,
    format_brief_notes,
    section_page_flags,
    toc_page_cache,
)
from api.services.tenders.outline import tech_chapter_titles
from api.services.tenders.money import rmb_lowercase, rmb_uppercase
from api.services.tenders.placeholders import (
    TECH_DRAWING_KEY,
    TECH_DRAWING_SLOT,
    _set_row_height,
    collect_slots,
    draw_placeholder_box,
    fill_perf_placeholders,
    id_slot,
    inline_id_scans,
)
from api.services.tenders.quote import (
    SECOND_ROUND_YUAN,
    QuoteSheet,
    apply_traffic_note,
    is_template_sheet,
    money_text,
    qty_text,
    resolve_quote_sheet,
    scale_quote,
)
from api.services.tenders.schema import BidBrief, DeviationLine, PerformanceLine
from api.services.tenders.slots import attachments_for_slots
from api.services.tenders.tables import (
    CHARGER_QUOTE_HEADERS,
    classify_quote_headers,
    name_amount_indexes,
    quote_role,
    width_ratios,
)

logger = logging.getLogger("api.tenders")

_SONG = "宋体"

# title, bookmark, 旧占位（现按 documentFormat.pageNumberStart 重算）
_TOC_ITEMS: tuple[tuple[str, str, str], ...] = (
    ("投标函及投标函附录", "toc_letter", "3"),
    ("分项报价表", "toc_quote", "5"),
    ("法定代表人身份证明", "toc_legal", "6"),
    ("授权委托书", "toc_auth", "7"),
    ("技术偏差表", "toc_dev", "8"),
    ("企业业绩", "toc_perf", "9"),
    ("原厂生产承诺", "toc_factory", "11"),
    ("技术标（实施方案）", "toc_other", "12"),
    ("投标承诺书", "toc_commit", "13"),
)
_HEADING_BOOKMARKS: tuple[tuple[str, str], ...] = (
    ("投标函及投标函附录", "toc_letter"),
    ("分项报价表", "toc_quote"),
    ("法定代表人身份证明", "toc_legal"),
    ("授权委托书", "toc_auth"),
    ("技术偏差表", "toc_dev"),
    ("原厂生产承诺", "toc_factory"),
    ("其他材料", "toc_other"),
    ("技术标（实施方案）", "toc_other"),
    ("投标函附录", "toc_letter_app"),
    ("支付条件", "toc_pay"),
)
_CN_NUM = "一二三四五六七八九"


def _compact(text: str) -> str:
    """去掉空白与括号，便于「技术标（实施方案）」与「技术标实施方案」互认。"""
    return re.sub(r"[\s/（）()]+", "", text or "")


def _font(run, size_pt: float | None = None, *, bold: bool | None = None) -> None:
    if bold is not None:
        run.bold = bold
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.font.name = _SONG
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), _SONG)
    rfonts.set(qn("w:hAnsi"), _SONG)
    rfonts.set(qn("w:eastAsia"), _SONG)
    for old in list(rpr.findall(qn("w:lang"))):
        rpr.remove(old)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rpr.append(lang)
    for old in list(rpr.findall(qn("w:noProof"))):
        rpr.remove(old)
    rpr.append(OxmlElement("w:noProof"))


def _clear_bold(rpr) -> None:
    for tag in (qn("w:b"), qn("w:bCs")):
        for old in list(rpr.findall(tag)):
            rpr.remove(old)


def _clone_fill_rpr(source_rpr, *, underline: bool = True):
    """从相邻正文复制字号/字体/字距，填空只加下划线、不加粗、不放大。"""
    rpr = OxmlElement("w:rPr")
    keep = {
        qn("w:rFonts"),
        qn("w:sz"),
        qn("w:szCs"),
        qn("w:spacing"),
        qn("w:kern"),
        qn("w:color"),
    }
    if source_rpr is not None:
        for child in list(source_rpr):
            if child.tag in keep:
                rpr.append(deepcopy(child))
    _clear_bold(rpr)
    if underline:
        for old in list(rpr.findall(qn("w:u"))):
            rpr.remove(old)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    return rpr


def _neighbor_rpr(para: Paragraph, skip_run) -> object | None:
    for run in para.runs:
        if run._element is skip_run._element:
            continue
        text = (run.text or "").strip()
        if not text:
            continue
        if _run_underlined(run) and (
            run._element.find(qn("w:tab")) is not None or not text
        ):
            continue
        rpr = run._element.find(qn("w:rPr"))
        if rpr is not None:
            return rpr
    p_pr = para._p.find(qn("w:pPr"))
    if p_pr is not None:
        return p_pr.find(qn("w:rPr"))
    return None


def _clear_runs(para: Paragraph) -> None:
    for run in list(para.runs):
        parent = run._element.getparent()
        if parent is not None:
            parent.remove(run._element)


def _clear_tab_stops(para: Paragraph) -> None:
    p_pr = para._p.find(qn("w:pPr"))
    if p_pr is None:
        return
    for old in p_pr.findall(qn("w:tabs")):
        p_pr.remove(old)


def _run_underlined(run) -> bool:
    rpr = run._element.find(qn("w:rPr"))
    return rpr is not None and rpr.find(qn("w:u")) is not None


def _blank_runs(para: Paragraph) -> list:
    """模板空白栏：带下划线、且几乎只有制表符的 run。"""
    blanks = []
    for run in para.runs:
        if not _run_underlined(run):
            continue
        text = (run.text or "").strip()
        if text and text not in {"元"}:
            continue
        if run._element.find(qn("w:tab")) is None and text:
            continue
        blanks.append(run)
    return blanks


def _remove_line_drawings(para: Paragraph) -> int:
    """去掉模板里的形状/VML 手绘横线（比字符下划线更粗）。"""
    removed = 0
    for run in list(para.runs):
        r = run._element
        for child in list(r):
            tag = child.tag
            if (
                tag == qn("w:drawing")
                or tag == qn("w:pict")
                or tag.endswith("}AlternateContent")
            ):
                r.remove(child)
                removed += 1
        if any(c.tag in (qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")) for c in r):
            continue
        parent = r.getparent()
        if parent is not None:
            parent.remove(r)
    return removed


def _ensure_blank_underline(para: Paragraph) -> None:
    """若段落有制表空位但无下划线，补上与其它填空一致的单线、不加粗。"""
    if _blank_runs(para):
        return
    for run in para.runs:
        if run._element.find(qn("w:tab")) is None:
            continue
        if (run.text or "").strip():
            continue
        rpr = run._element.get_or_add_rPr()
        _clear_bold(rpr)
        if rpr.find(qn("w:u")) is None:
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rpr.append(u)
        return


def _normalize_blank_underlines(doc: Document) -> None:
    """签字空位等未填下划线：去掉加粗，避免比其它横线更粗。

    只扫正文段落，不进表格：报价/偏差表单元格很长时 row.cells/text 极慢。
    """
    for para in doc.paragraphs:
        for run in _blank_runs(para):
            rpr = run._element.find(qn("w:rPr"))
            if rpr is None:
                continue
            _clear_bold(rpr)


def _put_on_line(run, text: str, *, para: Paragraph | None = None) -> None:
    """把文字写进原有下划线空白：跟相邻正文字体字号一致，不加粗、不放大。"""
    r = run._element
    ref = _neighbor_rpr(para, run) if para is not None else r.find(qn("w:rPr"))
    new_rpr = _clone_fill_rpr(ref, underline=True)
    old_rpr = r.find(qn("w:rPr"))
    if old_rpr is not None:
        r.remove(old_rpr)
    r.insert(0, new_rpr)
    for child in list(r):
        if child.tag in (qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")):
            r.remove(child)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = f" {(text or '').strip() or '　'} "
    r.append(t)


def _fill_blanks(para: Paragraph, *values: str) -> int:
    blanks = _blank_runs(para)
    filled = 0
    for run, val in zip(blanks, values):
        if not (val or "").strip():
            continue
        _put_on_line(run, val, para=para)
        filled += 1
    return filled


def _style_fill_run(run, para: Paragraph | None = None, *, underline: bool = True) -> None:
    """新建填空 run：样式对齐段落正文，仅下划线。"""
    ref = None
    if para is not None:
        ref = _neighbor_rpr(para, run)
        if ref is None:
            p_pr = para._p.find(qn("w:pPr"))
            if p_pr is not None:
                ref = p_pr.find(qn("w:rPr"))
    r = run._element
    old = r.find(qn("w:rPr"))
    if old is not None:
        r.remove(old)
    r.insert(0, _clone_fill_rpr(ref, underline=underline))


def _write_tenderer_line(para: Paragraph, name: str) -> None:
    """公司名下划线填空，其后保留完整「（招标人名称）：」。与正文字体一致。"""
    align = para.alignment
    _clear_runs(para)
    name_run = para.add_run((name or "").strip())
    _style_fill_run(name_run, para, underline=True)
    hint_run = para.add_run("（招标人名称）：")
    _style_fill_run(hint_run, para, underline=False)
    if align is not None:
        para.alignment = align


def _use_body_style(para: Paragraph) -> None:
    """封面落款在模板里是 Heading1，OnlyOffice 会当大纲标题转换，又慢还容易「下载失败」。"""
    p_pr = para._p.get_or_add_pPr()
    for old in list(p_pr.findall(qn("w:pStyle"))):
        p_pr.remove(old)
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "BodyText")
    p_pr.insert(0, style)
    rpr = p_pr.find(qn("w:rPr"))
    if rpr is not None:
        p_pr.remove(rpr)


def _em_len(text: str) -> float:
    """宋体 12pt 下汉字≈1em、西文≈0.55em，用来估算填空线长度。"""
    n = 0.0
    for ch in text or "":
        code = ord(ch)
        if ch in {"\u3000", "\u2003"}:
            n += 1.0
        elif code <= 32:
            continue
        elif code < 127:
            n += 0.55
        else:
            n += 1.0
    return n


def _add_form_pad(para: Paragraph, ems: float) -> None:
    n = int(round(max(0.0, ems)))
    if n <= 0:
        return
    _form_run(para, "\u3000" * n, underline=True)


def _set_word_wrap(para: Paragraph, *, enabled: bool) -> None:
    p_pr = para._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:wordWrap")):
        p_pr.remove(old)
    if enabled:
        return
    el = OxmlElement("w:wordWrap")
    el.set(qn("w:val"), "off")
    p_pr.append(el)


def _set_fill_tab_stop(para: Paragraph, pos_twips: int) -> None:
    """图三空栏：无前导符的左制表位，下划线 tab 拉成贴字底的长实线。"""
    p_pr = para._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:tabs")):
        p_pr.remove(old)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:leader"), "none")
    tab.set(qn("w:pos"), str(int(pos_twips)))
    tabs.append(tab)
    rpr = p_pr.find(qn("w:rPr"))
    if rpr is not None:
        rpr.addprevious(tabs)
    else:
        p_pr.append(tabs)


def _form_run(para: Paragraph, text: str, *, underline: bool = False) -> None:
    """与图三身份证明同款：继承正文，只加单实线下划线。"""
    run = para.add_run(text)
    _style_fill_run(run, para, underline=underline)
    _clear_bold(run._element.get_or_add_rPr())


def _add_form_blank(para: Paragraph) -> None:
    """空栏：全角空格实线下划线。不用制表符，避免预览把 tab 拉出页边。"""
    _add_form_pad(para, 8)


def _underline_value(para: Paragraph, value: str, target_em: int) -> None:
    """有字则字下划线，再补全角空格拉到固定长度；空栏整段下划线。超宽时不再强行补线，避免换行。"""
    target = max(1, int(target_em or 8))
    text = (value or "").strip()
    if text:
        filled = f" {text} "
        if _em_len(filled) > target + 0.4:
            filled = text
        _form_run(para, filled, underline=True)
        extra = target - _em_len(filled)
        if extra >= 0.5:
            _add_form_pad(para, extra)
        return
    _add_form_pad(para, max(4, target))


def _align_sign_label(label: str, width: int = 5) -> str:
    """短标签左右撑满与长标签同宽：空格插在字中间，冒号对齐。"""
    raw = (label or "").strip()
    if raw.endswith("：") or raw.endswith(":"):
        raw = raw[:-1]
    compact = "".join(ch for ch in raw if ch not in {"\u3000", " ", "\t"})
    if not compact:
        return "："
    n = len(compact)
    if n >= width:
        return compact + "："
    extra = width - n
    if n == 1:
        return compact + ("\u3000" * extra) + "："
    gaps = n - 1
    base, rem = divmod(extra, gaps)
    parts: list[str] = []
    for i, ch in enumerate(compact):
        parts.append(ch)
        if i < gaps:
            parts.append("\u3000" * (base + (1 if i < rem else 0)))
    return "".join(parts) + "："


def _line_em_budget(para: Paragraph, label: str, suffix: str, *, size_pt: float = 12.0) -> int:
    """当前段左右缩进后还能排下多少个汉字，用来截断填空线。"""
    doc = para.part.document
    sec = doc.sections[-1]
    usable = float(sec.page_width.cm) - float(sec.left_margin.cm) - float(sec.right_margin.cm)
    pf = para.paragraph_format
    if pf.left_indent:
        usable -= float(pf.left_indent.cm)
    if pf.right_indent:
        usable -= float(pf.right_indent.cm)
    em_cm = max(0.32, size_pt * 2.54 / 72.0)
    remain = max(12.0, usable / em_cm) - _em_len(label) - _em_len(suffix) - 0.4
    return max(4, int(remain))


_LETTER_CONTACT_KW: dict = {
    "align": WD_ALIGN_PARAGRAPH.LEFT,
    "left_indent_cm": 3.2,
    "line_spacing": 1.5,
    "space_before": 6,
    "space_after": 4,
    "line_em": 8,
}

# 投标函网址/电话/传真/邮编：四字标签对齐，下划线拉到同一右缘
_LETTER_FIELD_KW: dict = {
    **_LETTER_CONTACT_KW,
    "line_em": 22,
    "nowrap": True,
}

_COVER_SIGN_KW: dict = {
    **_LETTER_CONTACT_KW,
    "left_indent_cm": 3.6,
    "space_before": 8,
    "space_after": 6,
    "line_em": 8,
}

_LETTER_ADDR_KW: dict = {
    **_LETTER_CONTACT_KW,
    "line_spacing": 2.0,
    "space_before": 8,
    "space_after": 8,
    "line_em": 8,
}

# 身份证明 / 授权委托书落款：标签冒号对齐、单行不换行
_SIGN_LABEL_WIDTH = 5
_SIGN_OFF_KW: dict = {
    "align": WD_ALIGN_PARAGRAPH.LEFT,
    "left_indent_cm": 3.2,
    "right_indent_cm": 0.25,
    "line_spacing": 1.5,
    "space_before": 4,
    "space_after": 4,
    "line_em": 8,
    "nowrap": True,
    "label_width": _SIGN_LABEL_WIDTH,
}

_SIGN_BOTTOM_PAD_CM = 2.2
_SIGN_MIN_GAP_CM = 1.5
_TWIPS_PER_CM = 567
_EMU_PER_CM = 360000.0


def _rewrite_labeled_underline(
    para: Paragraph,
    rows: list[tuple[str, str, str]],
    *,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
    line_spacing: float = 1.5,
    space_before: float | None = 8,
    space_after: float | None = 4,
    left_indent_cm: float | None = None,
    right_indent_cm: float | None = 0.55,
    blank_width: int = 16,
    line_em: int = 8,
    nowrap: bool = False,
    label_width: int | None = None,
) -> None:
    """标签 + 底部实线下划线 + 后缀。填空线不超过当前行宽，避免「公章）」掉到下一行。"""
    del blank_width
    _clear_runs(para)
    _clear_tab_stops(para)
    if align is not None:
        para.alignment = align
    pf = para.paragraph_format
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    pf.first_line_indent = Cm(0)
    if left_indent_cm is not None:
        pf.left_indent = Cm(left_indent_cm)
    pf.right_indent = Cm(0.0 if right_indent_cm is None else right_indent_cm)
    _use_body_style(para)
    if nowrap:
        _set_word_wrap(para, enabled=False)
    target = max(4, int(line_em or 8))
    for i, (label, value, suffix) in enumerate(rows):
        if i:
            para.add_run().add_break()
        display = _align_sign_label(label, label_width) if label_width else label
        _form_run(para, display, underline=False)
        fill = "" if any(k in f"{label}{suffix}{value}" for k in ("签字", "签名", "手签")) else value
        _underline_value(para, fill, min(target, _line_em_budget(para, display, suffix)))
        if suffix:
            _form_run(para, suffix, underline=False)


def _rewrite_date_line(
    para: Paragraph,
    year: str,
    month: str,
    day: str,
    **kw,
) -> None:
    """日期与上方落款同一左缘；年/月/日数字下划线，不用制表虚线。"""
    align = kw.get("align", WD_ALIGN_PARAGRAPH.LEFT)
    line_spacing = kw.get("line_spacing", 1.75)
    space_before = kw.get("space_before", 6)
    space_after = kw.get("space_after", 6)
    left_indent_cm = kw.get("left_indent_cm")
    _clear_runs(para)
    _clear_tab_stops(para)
    if align is not None:
        para.alignment = align
    pf = para.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    pf.first_line_indent = Cm(0)
    if left_indent_cm is not None:
        pf.left_indent = Cm(left_indent_cm)
    right_indent_cm = kw.get("right_indent_cm", 0.55)
    pf.right_indent = Cm(0.0 if right_indent_cm is None else right_indent_cm)
    _use_body_style(para)
    if kw.get("nowrap"):
        _set_word_wrap(para, enabled=False)
    label = str(kw.get("label") or "日期：")
    label_width = kw.get("label_width")
    if label_width:
        label = _align_sign_label(label, int(label_width))
    target = int(kw.get("line_em") or 0)
    month = _two_digit_md(month) or month
    day = _two_digit_md(day) or day
    _form_run(para, label, underline=False)
    if target >= 12:
        _underline_value(para, f"{year}年{month}月{day}日", target)
        return
    for val, tail in ((year, " 年 "), (month, " 月 "), (day, " 日")):
        _form_run(para, f" {val} ", underline=True)
        _form_run(para, tail, underline=False)


def _scrub_hints(para: Paragraph, hints: tuple[str, ...]) -> None:
    for run in para.runs:
        raw = run.text
        if not raw:
            continue
        text = raw
        for hint in hints:
            text = text.replace(hint, "")
        text = text.replace("（）：", "").replace("（）", "")
        if text.strip() in {"）", "（", "："}:
            text = ""
        if text != raw:
            run.text = text


def _write_para(
    para: Paragraph,
    text: str,
    *,
    size: float | None = None,
    bold: bool | None = False,
) -> None:
    """整段替换：与正文同字号，仅下划线、不加粗。"""
    del size, bold
    align = para.alignment
    _clear_runs(para)
    run = para.add_run(text)
    _style_fill_run(run, para, underline=True)
    if align is not None:
        para.alignment = align


def _insert_paragraph_after(para: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    para._element.addnext(new_p)
    return Paragraph(new_p, para._parent)


def _paragraph_has_page_break(p_el) -> bool:
    for br in p_el.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _el_plain_text(el) -> str:
    return "".join(node.text or "" for node in el.iter(qn("w:t")))


def _page_body_cm(doc: Document) -> float:
    sec = doc.sections[-1]
    return max(
        16.0,
        float(sec.page_height.cm) - float(sec.top_margin.cm) - float(sec.bottom_margin.cm) - 0.6,
    )


def _drawing_height_cm(el) -> float:
    h = 0.0
    for ext in el.findall(".//" + qn("wp:extent")):
        cy = int(ext.get("cy") or 0)
        if cy > 0:
            h += cy / _EMU_PER_CM
    return h


def _tbl_height_cm(tbl) -> float:
    h = 0.0
    rows = tbl.findall(qn("w:tr"))
    for tr in rows:
        val = 0
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is not None:
            th = tr_pr.find(qn("w:trHeight"))
            if th is not None:
                val = int(th.get(qn("w:val")) or 0)
        h += val / _TWIPS_PER_CM if val > 0 else 0.72
        h += _drawing_height_cm(tr)
    return max(h, 0.72)


def _p_spacing_cm(child) -> float:
    p_pr = child.find(qn("w:pPr"))
    if p_pr is None:
        return 0.0
    sp = p_pr.find(qn("w:spacing"))
    if sp is None:
        return 0.0
    before = int(sp.get(qn("w:before")) or 0)
    after = int(sp.get(qn("w:after")) or 0)
    return (before + after) / _TWIPS_PER_CM


def _estimate_cm_on_current_page(doc: Document, stop_el=None) -> float:
    """从本页开头估正文高度，用于把落款推到接近页脚且不翻页。"""
    used = 0.0
    body = doc.element.body
    for child in body:
        if stop_el is not None and child is stop_el:
            break
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:tbl"):
            used += _tbl_height_cm(child)
            continue
        if child.tag != qn("w:p"):
            continue
        if _paragraph_has_page_break(child):
            used = 0.35
            continue
        p_pr = child.find(qn("w:pPr"))
        if p_pr is not None and (
            p_pr.find(qn("w:pageBreakBefore")) is not None or p_pr.find(qn("w:sectPr")) is not None
        ):
            used = 0.35
            continue
        draw_h = _drawing_height_cm(child)
        if draw_h > 0:
            used += draw_h + _p_spacing_cm(child) + 0.15
            continue
        text = _el_plain_text(child).strip()
        if not text:
            if p_pr is not None:
                sp = p_pr.find(qn("w:spacing"))
                if sp is not None and (sp.get(qn("w:lineRule")) or "") == "exact":
                    line = int(sp.get(qn("w:line")) or 0)
                    if line:
                        used += line / _TWIPS_PER_CM
                        continue
            used += 0.22
            continue
        lines = max(1, (len(text) + 29) // 30)
        used += 0.48 * lines + 0.22
        if len(text) <= 18:
            used += 0.2
    return used


def _sign_spacer_cm(
    doc: Document,
    *,
    sign_cm: float,
    stop_el=None,
    bottom_pad_cm: float | None = None,
) -> float:
    used = _estimate_cm_on_current_page(doc, stop_el=stop_el)
    pad = _SIGN_BOTTOM_PAD_CM if bottom_pad_cm is None else bottom_pad_cm
    page = _page_body_cm(doc)
    remain = page - used
    # 本页已经放不下落后落款时不要再垫空白，否则整块签字会翻到下一页独页
    if remain < sign_cm + _SIGN_MIN_GAP_CM:
        return 0.0
    budget = remain - sign_cm - pad
    if used > page * 0.62 or budget < 0.35:
        return 0.0
    return min(budget, 8.0)


def _glue_sign_off(doc: Document, spacer: Paragraph | None, lines: list[Paragraph]) -> None:
    """落款跟前一段钉在一起，避免投标人/签字/日期单独成页。"""
    del doc
    first = spacer or (lines[0] if lines else None)
    if first is None:
        return
    prev = first._element.getprevious()
    while prev is not None and prev.tag not in {qn("w:p"), qn("w:tbl")}:
        prev = prev.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        Paragraph(prev, first._parent).paragraph_format.keep_with_next = True
    chain = ([spacer] if spacer is not None else []) + lines
    for p in chain[:-1]:
        if p is None:
            continue
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    if chain:
        chain[-1].paragraph_format.keep_together = True


def _apply_sign_spacer_paragraph(para: Paragraph, height_cm: float) -> None:
    """用精确行距撑开空白。不用空表，避免预览把无边框表画成虚线方框。"""
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    p_pr = para._p.get_or_add_pPr()
    for old in list(p_pr.findall(qn("w:spacing"))):
        p_pr.remove(old)
    spacing = OxmlElement("w:spacing")
    twips = str(max(240, int(height_cm * _TWIPS_PER_CM)))
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), twips)
    spacing.set(qn("w:lineRule"), "exact")
    p_pr.append(spacing)
    _clear_runs(para)
    run = para.add_run("\u200b")
    rpr = run._element.get_or_add_rPr()
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rpr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), "2")
    rpr.append(sz_cs)


def _add_bottom_sign_spacer(
    doc: Document,
    *,
    sign_cm: float,
    bottom_pad_cm: float | None = None,
) -> Paragraph:
    """在当前页正文后插入弹性空白，把随后落款压到接近页脚、且尽量不翻页。"""
    height = _sign_spacer_cm(doc, sign_cm=sign_cm, bottom_pad_cm=bottom_pad_cm)
    if height < 0.35:
        return None
    para = doc.add_paragraph()
    _apply_sign_spacer_paragraph(para, height)
    para.paragraph_format.keep_with_next = True
    return para


def _insert_bottom_sign_spacer_before(para: Paragraph, *, sign_cm: float) -> Paragraph | None:
    doc = para.part.document
    height = _sign_spacer_cm(doc, sign_cm=sign_cm, stop_el=para._element)
    if height < 0.35:
        return None
    spacer = doc.add_paragraph()
    _apply_sign_spacer_paragraph(spacer, height)
    para._element.addprevious(spacer._element)
    spacer.paragraph_format.keep_with_next = True
    return spacer


def _insert_bottom_sign_spacer_after(para: Paragraph, *, sign_cm: float) -> Paragraph | None:
    doc = para.part.document
    height = _sign_spacer_cm(doc, sign_cm=sign_cm, stop_el=para._element.getnext())
    if height < 0.35:
        return None
    spacer = doc.add_paragraph()
    _apply_sign_spacer_paragraph(spacer, height)
    para._element.addnext(spacer._element)
    spacer.paragraph_format.keep_with_next = True
    return spacer


def _write_plain(para: Paragraph, text: str, *, size: float = 12, bold: bool = False) -> None:
    _clear_runs(para)
    run = para.add_run(text)
    _font(run, size, bold=bold)


def _tech_drawing_files(catalog_media: dict | None) -> list[Path]:
    if isinstance(catalog_media, dict):
        return list(catalog_media.get(TECH_DRAWING_KEY) or [])
    return []


def _insert_picture_after(para: Paragraph, source, *, width_cm: float = 15.5) -> Paragraph:
    pic = _insert_paragraph_after(para)
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(4)
    pic.paragraph_format.space_after = Pt(4)
    pic.add_run().add_picture(source, width=Cm(width_cm))
    return pic


_DRAWING_MAX_FILES = 12
_DRAWING_PDF_PAGES = 4


def _insert_drawings_after(
    cursor: Paragraph, files: list[Path], *, max_pages: int = _DRAWING_PDF_PAGES
) -> tuple[Paragraph, int]:
    """压缩后嵌入实施方案图纸。"""
    import io

    from api.services.tenders.placeholders import _JPEG_QUALITY, _render_pdf_page_jpeg, _shrink_pil

    inserted = 0
    for path in files[:_DRAWING_MAX_FILES]:
        suf = path.suffix.lower()
        try:
            if suf in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"} and path.is_file():
                from PIL import Image

                pic = _insert_paragraph_after(cursor)
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.paragraph_format.space_before = Pt(4)
                pic.paragraph_format.space_after = Pt(4)
                with Image.open(path) as raw:
                    image = _shrink_pil(raw.convert("RGB"))
                    buf = io.BytesIO()
                    image.save(buf, format="JPEG", quality=max(40, _JPEG_QUALITY - 8), optimize=False)
                    buf.seek(0)
                    pic.add_run().add_picture(buf, width=Cm(15.5))
                cursor = pic
                inserted += 1
                continue
            if suf != ".pdf" or not path.is_file():
                continue
            try:
                import pypdfium2 as pdfium
            except ImportError:
                logger.warning("pypdfium2 missing, skip drawing pdf %s", path.name)
                continue
            pdf = pdfium.PdfDocument(str(path))
            count = min(len(pdf), max_pages)
            for i in range(count):
                page = pdf[i]
                buf = _render_pdf_page_jpeg(page)
                cursor = _insert_picture_after(cursor, buf)
                inserted += 1
        except Exception:
            logger.exception("insert tech drawing failed: %s", path)
    return cursor, inserted


def _embed_tech_drawings(doc: Document, after: Paragraph, catalog_media: dict | None) -> list[str]:
    files = [p for p in _tech_drawing_files(catalog_media) if p.is_file()]
    n = 0
    if files:
        _, n = _insert_drawings_after(after, files)
    if n > 0:
        notes = [f"已将 {n} 张实施方案图纸写入技术标"]
        skipped = max(0, len(files) - _DRAWING_MAX_FILES)
        if skipped:
            notes.append(f"另有 {skipped} 个图纸文件未写入，装订时可另附")
        return notes
    draw_placeholder_box(doc, TECH_DRAWING_SLOT, after=after)
    if files:
        return ["实施方案图纸未能写入 Word，已用虚线框占位"]
    return ["技术标实施方案尚未上传图纸，已用虚线框占位"]


def _is_toc_list_line(para: Paragraph) -> bool:
    """目录条目：一、 / 1.1 / 第一章。不要把正文里 List Paragraph 标题当成目录。"""
    style = para.style.name if para.style else ""
    if style in {"toc 1", "TOC 1", "toc 2", "TOC 2"}:
        return True
    raw = (para.text or "").strip()
    if re.match(r"^[一二三四五六七八九十]+[、.．]", raw):
        return True
    if re.match(r"^\d+(?:\.\d+)+", raw):
        return True
    if re.match(r"^第[一二三四五六七八九十0-9]+章", raw):
        return True
    return False


def _find_technical_heading(doc: Document) -> Paragraph | None:
    """定位正文「其他材料 / 技术标」标题，跳过目录里的同名条目。"""
    candidates: list[Paragraph] = []
    for para in doc.paragraphs:
        if _compact(para.text) not in {"其他材料", "技术标实施方案"}:
            continue
        if _is_toc_list_line(para):
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading") or style in {"Title", "Body Text"}:
            return para
        candidates.append(para)
    return candidates[0] if candidates else None


def _fill_technical_section(doc: Document, brief: BidBrief, catalog_media: dict | None = None) -> list[str]:
    heading = _find_technical_heading(doc)
    if heading is None:
        return []
    _write_plain(heading, "技术标（实施方案）", size=16, bold=True)
    _bookmark_paragraph(heading, "toc_other")
    cursor = heading
    intro = _insert_paragraph_after(cursor)
    _write_plain(
        intro,
        "本部分为技术标。评审看充电站如何建成：文字说明加本项目图纸。"
        "不响应技术要求会大量扣分。公司资格、业绩、报价与函件见商务标及附件。",
        size=12,
        bold=False,
    )
    cursor = intro
    sub = _insert_paragraph_after(cursor)
    _write_plain(sub, "一、文字描述", size=14, bold=True)
    _bookmark_paragraph(sub, "toc_other_text")
    body = _insert_paragraph_after(sub)
    _write_plain(body, tech_plan_body(brief), size=12, bold=False)
    cursor = body
    draw_title = _insert_paragraph_after(cursor)
    _write_plain(draw_title, "二、图纸", size=14, bold=True)
    _bookmark_paragraph(draw_title, "toc_other_draw")
    cursor = draw_title
    warnings = _embed_tech_drawings(doc, draw_title, catalog_media)
    if not tech_plan_text(brief):
        warnings.append("技术标实施方案文字说明待补，已在 Word 中标注或按供货期写入工期")
    return warnings


def _drop_para(para: Paragraph) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _strip_para_bottom(para: Paragraph) -> None:
    p_pr = para._p.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def _no_underline(run) -> None:
    run.underline = False
    rpr = run._element.get_or_add_rPr()
    for uel in list(rpr.findall(qn("w:u"))):
        rpr.remove(uel)
    uel = OxmlElement("w:u")
    uel.set(qn("w:val"), "none")
    rpr.append(uel)


def _keep_one_para(cell: _Cell) -> None:
    tc = cell._tc
    paras = [child for child in list(tc) if child.tag == qn("w:p")]
    for extra in paras[1:]:
        tc.remove(extra)


def _write_cell(
    cell: _Cell,
    text: str,
    *,
    size: float = 10.5,
    bold: bool = False,
    underline: bool = True,
    center: bool = False,
    bottom_line: bool = False,
    tight: bool = False,
) -> None:
    """单元格填空：默认下划线、不加粗；表头等需加粗时显式传 bold=True。"""
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para = cell.paragraphs[0]
    _strip_para_bottom(para)
    pad = Pt(0) if tight else Pt(2)
    para.paragraph_format.space_before = pad
    para.paragraph_format.space_after = pad
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    lines = (text or "").split("\n")
    for i, line in enumerate(lines):
        if i:
            para.add_run().add_break()
        run = para.add_run(line)
        # 填空内容不加粗；仅表头等显式 bold=True 时加粗。字号不大于正文常用值。
        _font(run, size, bold=bool(bold))
        if underline:
            run.underline = True
        else:
            run.underline = False
            rpr_u = run._element.get_or_add_rPr()
            for uel in list(rpr_u.findall(qn("w:u"))):
                rpr_u.remove(uel)
        if not bold:
            rpr = run._element.get_or_add_rPr()
            _clear_bold(rpr)
    if bottom_line:
        _cell_bottom_line(para)


def _cell_bottom_line(para: Paragraph) -> None:
    p_pr = para._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for old in list(p_bdr):
        if old.tag == qn("w:bottom"):
            p_bdr.remove(old)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)


def _two_digit_md(part: str) -> str:
    """月、日个位数补 0，如 9 → 09。"""
    digits = "".join(ch for ch in (part or "") if ch.isdigit())
    if not digits:
        return part or ""
    try:
        n = int(digits)
    except ValueError:
        return part
    if 1 <= n <= 31:
        return f"{n:02d}"
    return digits


def _ymd(iso: str) -> tuple[str, str, str]:
    raw = (iso or "").strip()
    parts = raw.replace("/", "-").split("-")
    if len(parts) >= 3:
        month = _two_digit_md(parts[1]) or "01"
        day = _two_digit_md(parts[2]) or "01"
        return parts[0], month, day
    return "　　", "　", "　"


def _date_line(iso: str) -> str:
    year, month, day = _ymd(iso)
    return f"{year}    年    {month}    月    {day}    日"


def _set_dot_tab(paragraph: Paragraph, pos_cm: float) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:tabs")):
        p_pr.remove(old)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(int(Cm(pos_cm).twips)))
    tabs.append(tab)
    p_pr.append(tabs)


def _toc_tab_pos_cm(para: Paragraph) -> float:
    for section in para.part.document.sections:
        usable = float(section.page_width.cm) - float(section.left_margin.cm) - float(section.right_margin.cm)
        if usable > 10:
            return max(usable - 0.15, 12.0)
    return 16.8


def _strip_toc_numbering(para: Paragraph) -> None:
    try:
        para.style = "Normal"
    except (KeyError, ValueError):
        pass
    p_pr = para._p.get_or_add_pPr()
    for tag in ("w:numPr", "w:ind", "w:pBdr"):
        for el in p_pr.findall(qn(tag)):
            p_pr.remove(el)
    rpr = p_pr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        p_pr.append(rpr)
    for uel in list(rpr.findall(qn("w:u"))):
        rpr.remove(uel)
    none = OxmlElement("w:u")
    none.set(qn("w:val"), "none")
    rpr.append(none)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = para.paragraph_format
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(2)


def _set_dot_leader_font(run) -> None:
    """点线必须用西文字体，宋体会被 WPS 画成实线。"""
    run.bold = False
    run.underline = False
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    for uel in list(rpr.findall(qn("w:u"))):
        rpr.remove(uel)
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")


def _next_bookmark_id(doc: Document) -> int:
    # 缓存在 document part，避免每次打书签都全表扫描
    cache_attr = "_weitai_next_bookmark_id"
    cached = getattr(doc, cache_attr, None)
    if isinstance(cached, int):
        setattr(doc, cache_attr, cached + 1)
        return cached
    used: list[int] = []
    for el in doc.element.iter(qn("w:bookmarkStart")):
        raw = el.get(qn("w:id"))
        if raw is None:
            continue
        try:
            used.append(int(raw))
        except ValueError:
            continue
    nxt = (max(used) + 1) if used else 1
    setattr(doc, cache_attr, nxt + 1)
    return nxt


def _bookmark_paragraph(para: Paragraph, name: str) -> None:
    """在标题段打书签，供目录 PAGEREF 定位。"""
    body = para._p.getparent()
    if body is None:
        return
    # 只清理本段上的同名书签，避免每次全文档 iter（分页书签多时极慢）
    for start in list(para._p.findall(qn("w:bookmarkStart"))):
        if start.get(qn("w:name")) != name:
            continue
        bid = start.get(qn("w:id"))
        para._p.remove(start)
        for end in list(para._p.findall(qn("w:bookmarkEnd"))):
            if end.get(qn("w:id")) == bid:
                para._p.remove(end)
                break
    bookmark_id = str(_next_bookmark_id(para.part.document))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    para._p.insert(0, start)
    para._p.append(end)


def _add_pageref_run(para: Paragraph, bookmark: str, cache: str) -> None:
    """插入 PAGEREF 域，打开 Word/WPS 后按实际页码更新。"""
    begin = para.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    begin._r.append(fld)
    _font(begin, 10.5)
    _no_underline(begin)

    instr = para.add_run()
    text = OxmlElement("w:instrText")
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = f" PAGEREF {bookmark} \\h "
    instr._r.append(text)
    _font(instr, 10.5)
    _no_underline(instr)

    sep = para.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    sep._r.append(fld_sep)
    _font(sep, 10.5)
    _no_underline(sep)

    cache_run = para.add_run(cache)
    _font(cache_run, 10.5)
    _no_underline(cache_run)

    end = para.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end._r.append(fld_end)
    _font(end, 10.5)
    _no_underline(end)


def _clear_toc_tabs(para: Paragraph) -> None:
    p_pr = para._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:tabs")):
        p_pr.remove(old)


def _apply_toc_indent(para: Paragraph, level: int) -> None:
    pf = para.paragraph_format
    pf.left_indent = Cm(0.74 * max(0, int(level) - 1))
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(3 if level <= 1 else 1)
    pf.space_after = Pt(3 if level <= 1 else 1)


def _write_toc_line(
    para: Paragraph,
    index: int,
    title: str,
    bookmark: str,
    cache: str,
    *,
    label: str | None = None,
    with_pages: bool = True,
    level: int = 1,
) -> None:
    _strip_toc_numbering(para)
    _apply_toc_indent(para, level)
    prefix = label if label is not None else f"{_CN_NUM[index]}、"
    size = 12.0 if level <= 1 else 10.5
    bold = level <= 1
    if not with_pages:
        _clear_toc_tabs(para)
        _clear_runs(para)
        title_run = para.add_run(f"{prefix}{title}")
        _font(title_run, size, bold=bold)
        _no_underline(title_run)
        return
    _set_dot_tab(para, _toc_tab_pos_cm(para))
    _clear_runs(para)
    title_run = para.add_run(f"{prefix}{title}")
    _font(title_run, size, bold=bold)
    _no_underline(title_run)
    tab_run = para.add_run()
    _set_dot_leader_font(tab_run)
    _no_underline(tab_run)
    tab_run.add_tab()
    _add_pageref_run(para, bookmark, cache)


def _chapter5_toc_tree(brief: BidBrief) -> list[dict]:
    note = tech_plan_text(brief)
    chaps = tech_chapter_titles(note)
    if len(chaps) >= 2:
        tech_kids = [{"title": t, "bookmark": "toc_other"} for t in chaps]
    else:
        tech_kids = [
            {"title": "文字描述", "bookmark": "toc_other_text"},
            {"title": "图纸", "bookmark": "toc_other_draw"},
        ]
    tree = [
        {
            "title": "投标函及投标函附录",
            "bookmark": "toc_letter",
            "children": [
                {"title": "投标函", "bookmark": "toc_letter"},
                {
                    "title": "投标函附录",
                    "bookmark": "toc_letter_app",
                    "children": [
                        {"title": "分项报价表", "bookmark": "toc_quote"},
                        {"title": "支付条件", "bookmark": "toc_pay"},
                    ],
                },
            ],
        },
        {
            "title": "法定代表人身份证明及授权委托书",
            "bookmark": "toc_legal",
            "children": [
                {"title": "法定代表人身份证明", "bookmark": "toc_legal"},
                {"title": "授权委托书", "bookmark": "toc_auth"},
            ],
        },
        {"title": "技术偏差表", "bookmark": "toc_dev"},
        {"title": "企业业绩", "bookmark": "toc_perf"},
        {"title": "原厂生产承诺", "bookmark": "toc_factory"},
        {
            "title": "技术标（实施方案）",
            "bookmark": "toc_other",
            "children": tech_kids,
        },
    ]
    if brief.includeCommitment:
        tree.append({"title": "投标承诺书", "bookmark": "toc_commit"})
    return tree


def _toc_region(doc: Document) -> tuple[Paragraph | None, list[Paragraph]]:
    toc_title: Paragraph | None = None
    lines: list[Paragraph] = []
    for para in doc.paragraphs:
        n = _compact(para.text)
        if n == "目录":
            toc_title = para
            continue
        if toc_title is None:
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading") and n and n != "目录":
            break
        if _p_sectpr(para._p) is not None:
            break
        lines.append(para)
    return toc_title, lines


def _scrub_toc_title_bookmarks(toc_title: Paragraph) -> None:
    for start in list(toc_title._p.findall(qn("w:bookmarkStart"))):
        name = start.get(qn("w:name")) or ""
        if name.startswith("toc_"):
            continue
        bid = start.get(qn("w:id"))
        toc_title._p.remove(start)
        for end in list(toc_title._p.findall(qn("w:bookmarkEnd"))):
            if end.get(qn("w:id")) == bid:
                toc_title._p.remove(end)


def _rebuild_toc(doc: Document, brief: BidBrief, *, with_pages: bool, fmt=None) -> None:
    toc_title, old = _toc_region(doc)
    if toc_title is None:
        return
    _scrub_toc_title_bookmarks(toc_title)
    scheme = fmt.tocNumbering if fmt is not None else "cn"
    rows = flatten_toc_tree(_chapter5_toc_tree(brief), scheme)
    for para in old:
        _drop_para(para)
    anchor = toc_title
    for i, (label, title, level, bm) in enumerate(rows):
        cache = toc_page_cache(i, fmt) if fmt is not None else "1"
        neo = _insert_paragraph_after(anchor)
        _write_toc_line(
            neo,
            i,
            title,
            bm,
            cache,
            label=label,
            with_pages=with_pages,
            level=level,
        )
        anchor = neo


def _ensure_toc_entries(
    doc: Document,
    *,
    include_commitment: bool = True,
    with_pages: bool = True,
    fmt=None,
    brief: BidBrief | None = None,
) -> None:
    if brief is None:
        brief = BidBrief(includeCommitment=include_commitment)
    _rebuild_toc(doc, brief, with_pages=with_pages, fmt=fmt)


# 模板第 9、12 节 header=0，页眉贴顶，OnlyOffice/部分 Word 会裁掉字头只剩横线。
_HEADER_DISTANCE = Cm(1.75)
_HEADER_TOP_MIN = Cm(2.6)


def _header_bottom_border(para: Paragraph) -> None:
    p_pr = para._p.get_or_add_pPr()
    for old in list(p_pr.findall(qn("w:pBdr"))):
        p_pr.remove(old)
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    # space 过大会把文字往页顶推，裁切更明显；1pt 让字紧贴横线
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "4F4F4F")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _reset_header_para(para: Paragraph) -> None:
    fmt = para.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _write_header_para(para: Paragraph, project_name: str) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _reset_header_para(para)
    title = (project_name or "").strip()
    if len(title) > 28:
        title = title[:28] + "…"
    _clear_runs(para)
    if title:
        run = para.add_run(title)
        _font(run, 10.5)
    _header_bottom_border(para)


def _ensure_header(section, project_name: str, *, cover: bool = False) -> None:
    """每一节都写页眉。封面节首页留空，其后各页显示项目名。"""
    section.header_distance = _HEADER_DISTANCE
    try:
        top = int(section.top_margin or 0)
    except (TypeError, ValueError):
        top = 0
    if top < int(_HEADER_TOP_MIN):
        section.top_margin = _HEADER_TOP_MIN
    header = section.header
    header.is_linked_to_previous = False
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    _write_header_para(para, project_name)
    if not cover:
        section.different_first_page_header_footer = False
        return
    section.different_first_page_header_footer = True
    first = section.first_page_header
    first.is_linked_to_previous = False
    first_para = first.paragraphs[0] if first.paragraphs else first.add_paragraph()
    _reset_header_para(first_para)
    _clear_runs(first_para)


def _deviation_rows(brief: BidBrief) -> list[tuple[str, str, str, str]]:
    lines = list(brief.deviationLines or [])
    if not lines:
        from api.services.tenders.extract import deviation_lines_from_quote

        quote_lines = list(brief.quoteLines or [])
        lines = deviation_lines_from_quote(quote_lines)
    if not lines:
        text = (brief.bidContent or "").strip() or (brief.projectName or "").strip() or "按招标文件要求供货"
        lines = [
            DeviationLine(seq="1", requirement=text, response=text, deviation="无偏差"),
        ]
    rows: list[tuple[str, str, str, str]] = []
    for i, line in enumerate(lines, start=1):
        req = (line.requirement or "").strip()
        if not req:
            continue
        rows.append(
            (
                (line.seq or str(i)).strip() or str(i),
                req,
                (line.response or req).strip(),
                (line.deviation or "无偏差").strip() or "无偏差",
            )
        )
    return rows or [("1", "按招标文件要求供货", "按招标文件要求供货", "无偏差")]


def _commercial_dev_rows(brief: BidBrief) -> list[tuple[str, str, str, str]]:
    """商务偏离只写交货/质保/付款等条款，不用报价单技术参数顶。"""
    rows: list[tuple[str, str, str, str]] = []

    def add(title: str, req: str, resp: str) -> None:
        req = (req or "").strip()
        if not req:
            return
        rows.append((str(len(rows) + 1), f"{title}：{req}", (resp or req).strip(), "无偏差"))

    if brief.deliveryDays:
        add("交货期", f"{brief.deliveryDays}天内交货", f"响应：中标通知后{brief.deliveryDays}天内交货")
    if brief.warrantyYears:
        add("质保期", f"{brief.warrantyYears}年", f"响应：质保期{brief.warrantyYears}年")
    if brief.bidValidityDays:
        add("投标有效期", f"{brief.bidValidityDays}日历天", f"响应：投标有效期{brief.bidValidityDays}日历天")
    quality = (brief.quality or "").strip()
    if quality:
        add("质量要求", quality, f"响应：{quality}")
    pay = []
    if brief.prepaidPct:
        pay.append(f"预付款{brief.prepaidPct}%")
    if brief.arrivalPct:
        pay.append(f"到货款{brief.arrivalPct}%")
    if brief.settlementPct:
        pay.append(f"验收款{brief.settlementPct}%")
    if brief.warrantyPct:
        pay.append(f"质保金{brief.warrantyPct}%")
    if pay:
        add("付款条件", "、".join(pay), "响应：" + "、".join(pay))
    return rows or [("1", "按招标文件商务条款执行", "完全响应招标文件商务条款，无偏差", "无偏差")]


def infer_factory_role(brief: BidBrief) -> str:
    role = (brief.factoryRole or "").strip()
    if role:
        return role
    blob = " ".join(
        [
            brief.bidContent or "",
            brief.projectName or "",
            brief.quoteTitle or "",
            *[ln.name for ln in (brief.quoteLines or [])],
        ]
    )
    if "充电" in blob:
        return "充电设备生产厂商"
    return "投标产品生产厂商"


def _factory_addressee(brief: BidBrief) -> str:
    """致送对象：优先招标人，空则用项目名称（投标项目名称）。"""
    return (brief.tenderer or "").strip() or (brief.projectName or "").strip()


def _is_factory_body(n: str) -> bool:
    if not n:
        return False
    if any(mark in n for mark in ("见模", "充电设备生产厂商", "投标产品生产厂商", "承担原厂责任", "联源热电")):
        return True
    if n.endswith("：") and any(k in n for k in ("公司", "局", "中心", "院")):
        return True
    return False


def _add_factory_run(para: Paragraph, text: str, *, underline: bool) -> None:
    if not text:
        return
    run = para.add_run(text)
    _style_fill_run(run, para, underline=underline)


def _write_factory_commitment(para: Paragraph, brief: BidBrief) -> Paragraph:
    """原厂承诺：致送单位单独一行，正文缩进，仅填空处下划线。返回正文段，便于后续落款。"""
    addressee = _factory_addressee(brief) or "　　　　"
    project = (brief.projectName or "").strip()
    _clear_runs(para)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = para.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.5
    _add_factory_run(para, addressee, underline=True)
    _add_factory_run(para, "：", underline=False)

    body = _insert_paragraph_after(para)
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bf = body.paragraph_format
    bf.first_line_indent = Cm(0.74)
    bf.line_spacing = 1.5
    bf.space_before = Pt(6)
    bf.space_after = Pt(12)
    _add_factory_run(body, "我单位 ", underline=False)
    _add_factory_run(body, brief.bidderName, underline=True)
    _add_factory_run(body, " 作为", underline=False)
    _add_factory_run(body, infer_factory_role(brief), underline=True)
    _add_factory_run(
        body,
        "，承诺具备加工生产条件，"
        "在人员、设备、资金等方面具备相应的供货能力，对",
        underline=False,
    )
    if project:
        _add_factory_run(body, f"「{project}」", underline=True)
    else:
        _add_factory_run(body, "本次", underline=False)
    _add_factory_run(body, "投标产品的质量、供货期、质保期及售后服务承担原厂责任。本项目供货期 ", underline=False)
    _add_factory_run(body, str(brief.deliveryDays), underline=True)
    _add_factory_run(body, " 日历天内，质保期 ", underline=False)
    _add_factory_run(body, str(brief.warrantyYears), underline=True)
    _add_factory_run(body, " 年。特此承诺。", underline=False)
    return body


def _append_right_sign_block(after: Paragraph, brief: BidBrief, *, unit: str) -> None:
    signer = (brief.agentName or "").strip() or (brief.legalPersonName or "").strip()
    year, month, day = _ymd(brief.bidDate)
    spacer = _insert_bottom_sign_spacer_after(after, sign_cm=4.2)
    p1 = _insert_paragraph_after(spacer or after)
    _rewrite_labeled_underline(
        p1,
        [(unit, (brief.bidderName or "").strip(), "")],
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        space_before=4,
        space_after=6,
    )
    p2 = _insert_paragraph_after(p1)
    _rewrite_labeled_underline(
        p2,
        [("法定代表人或其委托代理人（签字或盖章）：", signer, "")],
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        space_before=8,
        space_after=6,
    )
    p3 = _insert_paragraph_after(p2)
    _rewrite_labeled_underline(
        p3,
        [("日期：", f"{year}年{month}月{day}日", "")],
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        space_before=8,
        space_after=4,
    )
    _glue_sign_off(after.part.document, spacer, [p1, p2, p3])


def _auth_text(brief: BidBrief) -> str:
    legal = (brief.legalPersonName or "").strip() or "　　　"
    agent = (brief.agentName or "").strip()
    if not agent:
        return (
            "本次投标文件由法定代表人直接签署、递交，不另行委托代理人。"
            "如需委托，请填写委托代理人姓名后重新生成。"
        )
    until = (brief.agentAuthUntil or "").strip() or "至本项目投标有效期届满"
    return (
        f"本人 {legal} （姓名）系 {brief.bidderName} （投标人名称）的法定代表人，"
        f"现委托 {agent} （姓名）为我方代理人。代理人根据授权，以我方名义签署、澄清、说明、"
        f"补正、递交、撤回、修改 {brief.projectName} 投标文件、签订合同和处理有关事宜，"
        f"其法律后果由我方承担。委托期限：{until}。代理人无转委托权。"
    )


def _fill_paragraphs(
    doc: Document,
    brief: BidBrief,
    media: dict | None = None,
    inlined: set[str] | None = None,
) -> Paragraph | None:
    price_cn = rmb_uppercase(brief.bidPriceYuan)
    price_en = rmb_lowercase(brief.bidPriceYuan)
    extra = (brief.extraNote or "").strip() or "（无）"
    year, month, day = _ymd(brief.bidDate)
    section = "cover"
    drop: list[Paragraph] = []
    quote_anchor: Paragraph | None = None
    letter_date_done = False
    letter_sign_gapped = False
    legal_sign_gapped = False
    auth_legal_id_done = False
    auth_agent_id_done = False
    auth_agent_line_done = False
    auth_sign_started = False
    hints_project = ("（项目名称）", "（项目名称")
    hints_bidder = ("（投标人名称）",)

    for para in doc.paragraphs:
        raw = para.text
        n = _compact(raw)
        style = para.style.name if para.style else ""
        heading = style.startswith("Heading")

        if n.startswith("第五章"):
            drop.append(para)
            continue
        if n.startswith("目录") or n == "目录":
            section = "toc"
            continue
        if section == "toc" and not heading:
            if "备注" in n or "应编制" in n or "页码及内容" in n:
                drop.append(para)
            continue
        if heading:
            for title, bookmark in _HEADING_BOOKMARKS:
                if n == _compact(title):
                    _bookmark_paragraph(para, bookmark)
                    break
        if n == "投标函及投标函附录" and heading:
            section = "letter"
            continue
        if n == "投标函附录":
            section = "appendix"
            _bookmark_paragraph(para, "toc_letter_app")
            continue
        if n == "投标函" and section == "letter":
            _bookmark_paragraph(para, "toc_letter")
        if n == "支付条件" and heading:
            section = "pay"
            continue
        if n == "分项报价表" and heading:
            section = "quote"
            continue
        if n == "法定代表人身份证明" and heading:
            section = "legal"
            continue
        if n == "授权委托书" and heading:
            section = "auth"
            continue
        if n == "技术偏差表" and heading:
            section = "dev"
            continue
        if "近年完成的类似项目" in n:
            section = "perf"
            _bookmark_paragraph(para, "toc_perf")
            continue
        if "正在供货和新承接" in n:
            section = "perf"
            continue
        if n == "原厂生产承诺" and heading:
            section = "factory"
            continue
        if n == "其他材料" and heading:
            section = "other"
            continue

        if section == "cover":
            if "项目名称" in n:
                _fill_blanks(para, brief.projectName)
                _scrub_hints(para, hints_project)
            elif n.startswith("投标人") and "盖单位公章" in n:
                mashed_legal = "法定代表人" in n
                _remove_line_drawings(para)
                _rewrite_labeled_underline(
                    para,
                    [("投标人：", brief.bidderName, "（盖单位公章）")],
                    **_COVER_SIGN_KW,
                )
                if mashed_legal:
                    legal_p = _insert_paragraph_after(para)
                    _rewrite_labeled_underline(
                        legal_p,
                        [("法定代表人或其委托代理人：", "", "（签字）")],
                        **_COVER_SIGN_KW,
                    )
            elif "法定代表人或其委托代理人" in n and "盖单位公章" not in n:
                _rewrite_labeled_underline(
                    para,
                    [("法定代表人或其委托代理人：", "", "（签字）")],
                    **_COVER_SIGN_KW,
                )
            elif n.startswith("日期") or n == "年月日":
                _rewrite_date_line(para, year, month, day, **_COVER_SIGN_KW)
            continue

        if section == "letter":
            if "招标人名称" in n:
                _write_tenderer_line(para, brief.tenderer)
            elif "愿意以人民币投标总报价" in n or "我方已认真研究了" in n:
                _fill_blanks(
                    para,
                    brief.projectName,
                    price_cn,
                    price_en,
                    f"{brief.deliveryDays}天内",
                    brief.quality,
                )
                _scrub_hints(para, hints_project)
            elif "其他补充说明" in n:
                _remove_line_drawings(para)
                _ensure_blank_underline(para)
                _fill_blanks(para, extra)
                if not _blank_runs(para):
                    _write_para(para, extra, size=12)
            elif n.startswith("投标人") and "盖单位公章" in n:
                if not letter_sign_gapped:
                    _insert_bottom_sign_spacer_before(para, sign_cm=7.4)
                    letter_sign_gapped = True
                _rewrite_labeled_underline(
                    para,
                    [("投　标　人：", brief.bidderName, "（盖单位公章）")],
                    **{**_LETTER_CONTACT_KW, "space_before": 4, "nowrap": True},
                )
            elif "法定代表人或其委托代理人" in n:
                mashed_addr = "地址" in n
                _rewrite_labeled_underline(
                    para,
                    [
                        (
                            "法定代表人或其委托代理人：",
                            (brief.agentName or "").strip(),
                            "（签字或盖章）",
                        )
                    ],
                    **_LETTER_CONTACT_KW,
                )
                if mashed_addr:
                    addr_para = _insert_paragraph_after(para)
                    _rewrite_labeled_underline(
                        addr_para,
                        [("地　址：", (brief.bidderAddress or "").strip(), "")],
                        **_LETTER_ADDR_KW,
                    )
            elif n.startswith("地址"):
                _rewrite_labeled_underline(
                    para,
                    [("地　址：", (brief.bidderAddress or "").strip(), "")],
                    **_LETTER_ADDR_KW,
                )
            elif n.startswith("网址"):
                _rewrite_labeled_underline(
                    para,
                    [("网　　址：", (brief.bidderWebsite or "").strip(), "")],
                    **_LETTER_FIELD_KW,
                )
            elif n.startswith("电话"):
                _rewrite_labeled_underline(
                    para,
                    [("电　　话：", (brief.bidderPhone or "").strip(), "")],
                    **_LETTER_FIELD_KW,
                )
            elif "传真" in n:
                mashed_post = "邮政编码" in n
                _rewrite_labeled_underline(
                    para,
                    [("传　　真：", (brief.bidderFax or "").strip(), "")],
                    **_LETTER_FIELD_KW,
                )
                cursor = para
                if mashed_post:
                    post_para = _insert_paragraph_after(para)
                    _rewrite_labeled_underline(
                        post_para,
                        [("邮政编码：", (brief.bidderPostcode or "").strip(), "")],
                        **_LETTER_FIELD_KW,
                    )
                    cursor = post_para
                _rewrite_date_line(
                    _insert_paragraph_after(cursor),
                    year,
                    month,
                    day,
                    **{**_LETTER_FIELD_KW, "label": "日　　期："},
                )
                letter_date_done = True
            elif n.startswith("邮政编码"):
                _rewrite_labeled_underline(
                    para,
                    [("邮政编码：", (brief.bidderPostcode or "").strip(), "")],
                    **_LETTER_FIELD_KW,
                )
                _rewrite_date_line(
                    _insert_paragraph_after(para),
                    year,
                    month,
                    day,
                    **{**_LETTER_FIELD_KW, "label": "日　　期："},
                )
                letter_date_done = True
            elif n.startswith("日期") or n == "年月日":
                if letter_date_done:
                    drop.append(para)
                else:
                    _rewrite_date_line(
                        para,
                        year,
                        month,
                        day,
                        **{**_LETTER_FIELD_KW, "label": "日　　期："},
                    )
                    letter_date_done = True
            continue

        if section in {"appendix", "pay"}:
            if "盖单位公章" in n or "盖单位章" in n or (n.startswith("投标人") and "签字" in n):
                mashed_legal = "法定代表人" in n
                seal = "（盖单位公章）" if "公章" in n else "（盖单位章）"
                label = "投标人名称：" if "名称" in n else "投　标　人："
                _rewrite_labeled_underline(
                    para,
                    [(label, brief.bidderName, seal)],
                    **_LETTER_CONTACT_KW,
                )
                if mashed_legal:
                    legal_p = _insert_paragraph_after(para)
                    sign_hint = "（签字）" if "签字" in n else "（签字或盖章）"
                    _rewrite_labeled_underline(
                        legal_p,
                        [("法定代表人或委托代理人：", "", sign_hint)],
                        **_LETTER_CONTACT_KW,
                    )
            elif n == "年月日" or n.startswith("日期"):
                _rewrite_date_line(para, year, month, day, **_LETTER_CONTACT_KW)
            continue

        if section == "quote":
            if "附表" in n:
                quote_anchor = para
                align = para.alignment
                _clear_runs(para)
                run = para.add_run(
                    "备注：详见附表。工程量与招标清单一致，综合单价按投标总价折算。"
                )
                _font(run, 10.5)
                run.underline = False
                if align is not None:
                    para.alignment = align
            continue

        if section == "legal":
            if n.startswith("投标人名称"):
                _fill_blanks(para, brief.bidderName)
            elif n.startswith("单位性质"):
                _fill_blanks(para, brief.bidderNature)
            elif n.startswith("地址"):
                _fill_blanks(para, (brief.bidderAddress or "").strip())
                pf = para.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                pf.line_spacing = 2.0
                pf.space_before = Pt(12)
                pf.space_after = Pt(12)
            elif "成立时间" in n:
                fy, fm, fd = _ymd(brief.foundedDate) if (brief.foundedDate or "").strip() else ("", "", "")
                _fill_blanks(para, fy, fm, fd, (brief.businessTerm or "").strip())
            elif n.startswith("姓名"):
                _fill_blanks(
                    para,
                    (brief.legalPersonName or "").strip(),
                    (brief.legalPersonGender or "").strip(),
                    (brief.legalPersonAge or "").strip(),
                    (brief.legalPersonTitle or "").strip(),
                )
            elif "的法定代表人" in n:
                _fill_blanks(para, brief.bidderName)
                _scrub_hints(para, hints_bidder)
            elif n.startswith("投标人：") and "盖单位公章" in n:
                if not legal_sign_gapped:
                    inline_id_scans(
                        doc,
                        (media or {}).get("id_legal") if media else None,
                        empty_slot=id_slot("id_legal"),
                        before=para,
                    )
                    if inlined is not None:
                        inlined.add("id_legal")
                    _insert_bottom_sign_spacer_before(para, sign_cm=4.4)
                    legal_sign_gapped = True
                para.paragraph_format.keep_with_next = True
                para.paragraph_format.keep_together = True
                _rewrite_labeled_underline(
                    para,
                    [("投标人：", brief.bidderName, "（盖单位公章）")],
                    **_SIGN_OFF_KW,
                )
            elif n == "年月日":
                _rewrite_date_line(para, year, month, day, **_SIGN_OFF_KW)
            continue

        if section == "auth":
            if n.startswith("本人") and "法定代表人" in n:
                agent = (brief.agentName or "").strip() or "（不委托）"
                _fill_blanks(
                    para,
                    (brief.legalPersonName or "").strip(),
                    brief.bidderName,
                    agent,
                    brief.projectName,
                )
                _scrub_hints(para, hints_project + hints_bidder + ("（姓名）",))
            elif n.startswith("委托期限"):
                until = (brief.agentAuthUntil or "").strip() or "至本项目投标有效期届满"
                _rewrite_labeled_underline(
                    para,
                    [("委托期限：", until, "。代理人无转委托权。")],
                    align=WD_ALIGN_PARAGRAPH.LEFT,
                    line_spacing=1.5,
                    space_before=6,
                    space_after=6,
                )
            elif n.startswith("投标人") and (
                "法定代表人" in n or "盖单位公章" in n or "盖单位章" in n
            ):
                if not auth_sign_started:
                    if (brief.agentName or "").strip():
                        inline_id_scans(
                            doc,
                            (media or {}).get("id_agent") if media else None,
                            empty_slot=id_slot("id_agent"),
                            before=para,
                        )
                        if inlined is not None:
                            inlined.add("id_agent")
                    _insert_bottom_sign_spacer_before(para, sign_cm=6.6)
                auth_sign_started = True
                para.paragraph_format.keep_with_next = True
                para.paragraph_format.keep_together = True
                mashed_legal = "法定代表人" in n
                _rewrite_labeled_underline(
                    para,
                    [("投标人：", brief.bidderName, "（盖单位公章）")],
                    **_SIGN_OFF_KW,
                )
                if mashed_legal:
                    legal_p = _insert_paragraph_after(para)
                    _rewrite_labeled_underline(
                        legal_p,
                        [
                            (
                                "法定代表人：",
                                "",
                                "（签字）",
                            )
                        ],
                        **_SIGN_OFF_KW,
                    )
            elif n.startswith("身份证号码") and "委托代理人" not in n:
                # 委托期限后、落款前的身份证行与签字区重复，丢掉。
                if not auth_sign_started:
                    drop.append(para)
                elif not auth_legal_id_done:
                    _rewrite_labeled_underline(
                        para,
                        [("身份证号码：", (brief.legalPersonIdNo or "").strip(), "")],
                        **{**_SIGN_OFF_KW, "line_em": 18},
                    )
                    auth_legal_id_done = True
                elif auth_agent_id_done:
                    drop.append(para)
                else:
                    _rewrite_labeled_underline(
                        para,
                        [("身份证号码：", (brief.agentIdNo or "").strip(), "")],
                        **{**_SIGN_OFF_KW, "line_em": 18},
                    )
                    auth_agent_id_done = True
            elif "委托代理人" in n and "法定代表人" not in n:
                mashed_id = "身份证" in n
                if not auth_sign_started or auth_agent_line_done:
                    drop.append(para)
                else:
                    _rewrite_labeled_underline(
                        para,
                        [
                            (
                                "委托代理人：",
                                "",
                                "（签字）",
                            )
                        ],
                        **{**_SIGN_OFF_KW, "line_em": 18},
                    )
                    auth_agent_line_done = True
                    if mashed_id:
                        id_p = _insert_paragraph_after(para)
                        _rewrite_labeled_underline(
                            id_p,
                            [("身份证号码：", (brief.agentIdNo or "").strip(), "")],
                            **{**_SIGN_OFF_KW, "line_em": 18},
                        )
                        auth_agent_id_done = True
            elif n == "年月日":
                _rewrite_date_line(para, year, month, day, **_SIGN_OFF_KW)
            continue

        if section == "factory":
            if _is_factory_body(n):
                last = _write_factory_commitment(para, brief)
                _append_right_sign_block(last, brief, unit="承诺单位（盖章）：")
            continue

        if section == "other":
            # 去掉模板里「本部分附企业介绍…」说明段，待补资料改用后文小标题+方框
            if n:
                drop.append(para)
            continue

    for para in drop:
        _drop_para(para)
    return quote_anchor


def _fill_tables(doc: Document, brief: BidBrief) -> None:
    price_cn = rmb_uppercase(brief.bidPriceYuan)
    price_en = rmb_lowercase(brief.bidPriceYuan)
    perf_iter = iter(_ordered_perf_lines(brief))
    for table in doc.tables:
        kind = _table_kind(table)
        if kind == "appendix":
            _fill_appendix_table(table, brief, price_cn, price_en)
        elif kind == "pay":
            _fill_payment_table(table, brief)
        elif kind == "dev":
            _fill_deviation_table(table, brief)
        elif kind == "perf":
            _fill_perf_table(table, next(perf_iter, None))


def _table_kind(table: Table) -> str:
    if not table.rows or not table.rows[0].cells:
        return ""
    joined = "".join(_compact(c.text) for c in table.rows[0].cells)
    if "招标文件要求" in joined:
        return "dev"
    if "支付阶段" in joined or "支付比例" in joined:
        return "pay"
    if "不含税综合单价" in joined or "技术参数要求" in joined:
        return "quote"
    first = _compact(table.rows[0].cells[0].text)
    if first.startswith("项目名称"):
        if len(table.rows[0].cells) >= 3:
            return "appendix"
        return "perf"
    return ""


def _usable_width_twips(doc: Document) -> int:
    """版心宽度（页宽减左右边距），再留一点余量避免表格画出页边。"""
    for section in doc.sections:
        usable = (
            int(section.page_width.twips)
            - int(section.left_margin.twips)
            - int(section.right_margin.twips)
        )
        if usable > 5000:
            return max(usable - 120, 5000)
    return int(Cm(16.0).twips)


def _distribute_twips(total: int, ratios: tuple[int, ...]) -> list[int]:
    weight = sum(ratios)
    parts = [total * r // weight for r in ratios]
    parts[-1] += total - sum(parts)
    return parts


def _apply_fixed_table_widths(table: Table, widths_twips: list[int]) -> None:
    """固定列宽并限制表宽，避免 OnlyOffice / Word 按内容撑出页边。"""
    total = sum(widths_twips)
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for old in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    for old in tbl_pr.findall(qn("w:tblLayout")):
        tbl_pr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl_pr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips, strict=False):
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            for old in tc_pr.findall(qn("w:noWrap")):
                tc_pr.remove(old)


def _pin_cjk_fonts(doc: Document) -> None:
    """去掉主题西文字体映射，正文固定 SimSun/宋体，避免 OnlyOffice 用 Liberation 撑高分页。"""
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        return
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]
    rfonts.set(qn("w:ascii"), "SimSun")
    rfonts.set(qn("w:hAnsi"), "SimSun")
    rfonts.set(qn("w:eastAsia"), "SimSun")
    rfonts.set(qn("w:cs"), "SimSun")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def _hide_proofing_marks(doc: Document) -> None:
    """关闭拼写/语法红波浪线，避免中文技术用语被标红。

    注意：不要开启 w:updateFields。分页改造后目录使用 PAGEREF，
    若强制打开时更新全部域，OnlyOffice/WPS 预览会卡住数秒到数十秒。
    目录页码已写入域缓存；需要精确页码时可在 Word 里手动更新域。
    """
    root = doc.settings.element
    for tag in ("hideSpellingErrors", "hideGrammaticalErrors"):
        el_tag = qn(f"w:{tag}")
        for old in list(root.findall(el_tag)):
            root.remove(old)
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "true")
        root.append(el)
    for old in list(root.findall(qn("w:updateFields"))):
        root.remove(old)


def _set_tbl_borders(table: Table) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def _insert_table_after(doc: Document, para: Paragraph, rows: int, cols: int) -> Table:
    table = doc.add_table(rows=rows, cols=cols)
    para._p.addnext(table._tbl)
    return table


def _quote_warnings(sheet: QuoteSheet, target: Decimal, *, origin: str) -> list[str]:
    notes: list[str] = []
    n = len(sheet.lines)
    if origin == "template" or is_template_sheet(sheet):
        notes.append(
            f"未解析到本次工程量清单，分项暂用内置模板《{sheet.title}》共 {n} 项。"
            "请上传招标工程量 Excel/Word，或在页面改分项后再生成"
        )
        if target > 0 and target != SECOND_ROUND_YUAN:
            notes.append(
                f"二次报价函曾报 {money_text(SECOND_ROUND_YUAN)} 元，当前投标总价为 {money_text(target)} 元，请确认使用哪一档"
            )
    else:
        source_name = origin if origin in {"file", "llm", "text", "form"} else "清单"
        label = {
            "file": "上传的工程量/邀请书表格",
            "llm": "邀请书正文抽取",
            "text": "邀请书表格文本",
            "form": "页面分项明细",
        }.get(source_name, "工程量清单")
        notes.append(f"分项报价来自{label}《{sheet.title}》，共 {n} 项")
    if target <= 0:
        notes.append(
            f"分项报价表暂按清单含税 {money_text(sheet.source_inc_tax)} 元填写，生成前请填写投标总价"
        )
    elif target != sheet.source_inc_tax and sheet.source_inc_tax > 0:
        notes.append(
            f"分项单价已按投标总价从清单含税 {money_text(sheet.source_inc_tax)} 元折算"
        )
    if sheet.source_inc_tax <= 0 and any(ln.unit_price <= 0 for ln in sheet.lines):
        notes.append("清单单价为空，无法按总价折算，请补单价或确认投标总价")
    return notes


def _fill_quote_section(
    doc: Document,
    brief: BidBrief,
    anchor: Paragraph | None,
    *,
    headers: list[str] | tuple[str, ...] | None = None,
) -> list[str]:
    sheet, origin = resolve_quote_sheet(brief)
    if sheet is None or not sheet.lines:
        return ["未识别到报价清单，未插入报价表。请上传工程量清单后重新识别再生成"]
    target = Decimal(str(brief.bidPriceYuan or 0)).quantize(Decimal("0.01"))
    filled = scale_quote(sheet, target)
    notes = _quote_warnings(sheet, target, origin=origin)
    if anchor is None:
        notes.append("未找到「分项报价表」备注段，未能插入报价明细")
        return notes

    charger = headers is None
    if charger:
        titles = list(CHARGER_QUOTE_HEADERS)
        roles = [quote_role(title) or "extra" for title in titles]
        ratios = (11, 20, 58, 10, 12, 26, 23)
    else:
        titles = [str(h).strip() for h in headers if str(h).strip()]
        layout = classify_quote_headers(titles) if titles else None
        if layout is not None:
            titles = list(layout.titles)
            roles = list(layout.roles)
        else:
            titles = list(CHARGER_QUOTE_HEADERS) if not titles else titles
            roles = [quote_role(title) or "extra" for title in titles]
        ratios = width_ratios(roles)

    extra = 3
    cols = len(titles)
    table = _insert_table_after(doc, anchor, rows=1 + len(filled.lines) + extra, cols=cols)
    _set_tbl_borders(table)
    table.autofit = False
    col_twips = _distribute_twips(_usable_width_twips(doc), ratios)
    for i, title in enumerate(titles):
        _write_cell(
            table.rows[0].cells[i],
            title,
            size=9,
            bold=True,
            underline=False,
            center=True,
            bottom_line=False,
            tight=True,
        )
    for i, line in enumerate(filled.lines, start=1):
        values = _quote_row_values(line, roles, traffic_note=brief.trafficFeeNote)
        for j, val in enumerate(values):
            if j >= cols:
                break
            center = roles[j] != "spec"
            _write_cell(
                table.rows[i].cells[j],
                val,
                size=8 if roles[j] == "spec" else 9,
                bold=False,
                underline=False,
                center=center,
                bottom_line=False,
                tight=True,
            )
    summaries = (
        ("不含税合计", money_text(filled.total_ex_tax)),
        ("税率", f"{(filled.tax_rate * 100).quantize(Decimal('1'))}%"),
        ("含税合计", money_text(filled.total_inc_tax)),
    )
    base = 1 + len(filled.lines)
    start_seq = len(filled.lines) + 1
    name_i, amount_i = name_amount_indexes(roles)
    if amount_i >= cols:
        amount_i = cols - 1
    if name_i >= cols:
        name_i = 1 if cols > 1 else 0
    for offset, (label, value) in enumerate(summaries):
        row = table.rows[base + offset]
        if amount_i > name_i + 1:
            row.cells[name_i].merge(row.cells[amount_i - 1])
            _keep_one_para(row.cells[name_i])
        if "seq" in roles:
            _write_cell(
                row.cells[0],
                str(start_seq + offset),
                size=9,
                bold=True,
                underline=False,
                center=True,
                bottom_line=False,
                tight=True,
            )
        _write_cell(
            row.cells[name_i],
            label,
            size=9,
            bold=True,
            underline=False,
            center=False,
            bottom_line=False,
            tight=True,
        )
        _write_cell(
            row.cells[amount_i],
            value,
            size=9,
            bold=True,
            underline=False,
            center=True,
            bottom_line=False,
            tight=True,
        )
        _set_row_height(row, 0.62)
    _apply_fixed_table_widths(table, col_twips)
    return notes


def _quote_row_values(line, roles: list[str], *, traffic_note: str) -> list[str]:
    spec = apply_traffic_note(line.spec, traffic_note)
    groups = list(line.groups or ())
    gi = 0
    values: list[str] = []
    for role in roles:
        if role == "seq":
            values.append(line.seq)
        elif role == "group":
            values.append(groups[gi] if gi < len(groups) else "")
            gi += 1
        elif role == "name":
            values.append(line.name)
        elif role == "spec":
            values.append(spec)
        elif role == "unit":
            values.append(line.unit)
        elif role == "qty":
            values.append(qty_text(line.qty))
        elif role == "price":
            values.append(money_text(line.unit_price))
        elif role == "amount":
            values.append(money_text(line.amount))
        else:
            values.append("")
    return values


def _fill_appendix_table(table: Table, brief: BidBrief, price_cn: str, price_en: str) -> None:
    mapping = {
        "项目名称": brief.projectName,
        "投标人": brief.bidderName,
        "投标总报价": f"（大写）{price_cn}\n（小写）{price_en}元",
        "投标内容": brief.bidContent,
        "投标质量": brief.quality,
        "交货期": f"{brief.deliveryDays}日历天内",
        "投标工期": f"{brief.deliveryDays}日历天内",
        "质保期": f"基础要求{brief.warrantyYears}年",
        "投标有效期": f"{brief.bidValidityDays} 日历天（从投标截止之日算起）",
        "需要说明的问题": (brief.extraNote or "").strip() or "投标人自行填写。（可另加附页）",
    }
    for row in table.rows:
        label = _compact(row.cells[0].text)
        for key, value in mapping.items():
            if key in label:
                _write_cell(row.cells[-1], value, underline=False)
                break


def _fill_payment_table(table: Table, brief: BidBrief) -> None:
    pcts = {
        "预付款": brief.prepaidPct,
        "到货款": brief.arrivalPct,
        "结算款": brief.settlementPct,
        "质保金": brief.warrantyPct,
    }
    template_p = None
    for row in table.rows[1:]:
        if len(row.cells) < 3:
            continue
        for para in row.cells[2].paragraphs:
            if _blank_runs(para) and "支付" in (para.text or ""):
                template_p = deepcopy(para._p)
                break
        if template_p is not None:
            break

    for row in table.rows[1:]:
        stage = _compact(row.cells[1].text) if len(row.cells) > 1 else ""
        for key, pct in pcts.items():
            if key not in stage:
                continue
            cell = row.cells[2]
            if template_p is not None:
                para = _replace_cell_paragraph(cell, template_p)
                _fill_blanks(para, str(pct))
            else:
                target = next((p for p in cell.paragraphs if _blank_runs(p)), None)
                if target is not None:
                    _fill_blanks(target, str(pct))
                else:
                    _write_payment_pct(cell, pct)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            break


def _replace_cell_paragraph(cell: _Cell, p_element) -> Paragraph:
    tc = cell._tc
    for child in list(tc):
        if child.tag == qn("w:p"):
            tc.remove(child)
    tc.append(deepcopy(p_element))
    return cell.paragraphs[0]


def _write_payment_pct(cell: _Cell, pct: int | str) -> None:
    """无模板空位时，按预付款格：支付至（常规）+ 下划线数字（不加粗）+ %。"""
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para = cell.paragraphs[0]
    _strip_para_bottom(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head = para.add_run("支付至")
    _font(head, 10.5, bold=False)
    num = para.add_run(f" {pct} ")
    _font(num, 10.5, bold=False)
    num.underline = True
    tail = para.add_run("%")
    _font(tail, 10.5, bold=False)


def _strip_cell_underline(cell: _Cell) -> None:
    for para in cell.paragraphs:
        p_pr = para._p.find(qn("w:pPr"))
        if p_pr is not None:
            p_bdr = p_pr.find(qn("w:pBdr"))
            if p_bdr is not None:
                p_pr.remove(p_bdr)
        for run in para.runs:
            run.underline = False
            rpr = run._element.find(qn("w:rPr"))
            if rpr is None:
                continue
            for uel in rpr.findall(qn("w:u")):
                rpr.remove(uel)


def _center_cell(cell: _Cell, *, horizontal: bool) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    align = WD_ALIGN_PARAGRAPH.CENTER if horizontal else WD_ALIGN_PARAGRAPH.LEFT
    for para in cell.paragraphs:
        para.alignment = align


def _is_short_cell(text: str) -> bool:
    compact = _compact(text)
    if not compact:
        return True
    if compact in {"无偏差", "……", "...", "序号"}:
        return True
    return len(compact) <= 12


def _fill_deviation_table(table: Table, brief: BidBrief) -> None:
    rows = _deviation_rows(brief)
    _ensure_deviation_rows(table, len(rows))
    data_rows = [row for row in table.rows[1:] if "…" not in _compact(row.cells[0].text)]
    for i, item in enumerate(rows):
        if i >= len(data_rows):
            break
        row = data_rows[i]
        for j, val in enumerate(item):
            if j >= len(row.cells):
                continue
            center = j in (0, 3) or _is_short_cell(val)
            _write_cell(
                row.cells[j],
                val,
                size=9,
                bold=(j == 0),
                underline=False,
                center=center,
                bottom_line=False,
            )
    for extra in data_rows[len(rows) :]:
        for j, cell in enumerate(extra.cells):
            if j == 0:
                continue
            _write_cell(cell, "", size=9, bold=False, underline=False, bottom_line=False)
    # 对齐已在 _write_cell 完成；不再二次遍历 cell.text（长技术参数时很慢）


def _ensure_deviation_rows(table: Table, needed: int) -> None:
    """保证偏差表有足够数据行（最后一行省略号保留）。"""
    if needed <= 0 or len(table.rows) < 2:
        return
    last = table.rows[-1]
    ellipsis = "…" in _compact(last.cells[0].text)
    data_count = len(table.rows) - 1 - (1 if ellipsis else 0)
    while data_count < needed:
        src = table.rows[-2] if ellipsis and len(table.rows) > 2 else table.rows[-1]
        new_tr = deepcopy(src._tr)
        if ellipsis:
            last._tr.addprevious(new_tr)
        else:
            table._tbl.append(new_tr)
        data_count += 1


def _ordered_perf_lines(brief: BidBrief) -> list[PerformanceLine]:
    from api.services.tenders.performance import bid_performance_lines, rank_performance_lines

    lines = bid_performance_lines(brief.performanceLines)
    req = brief.performanceRequirement
    done = rank_performance_lines([item for item in lines if not item.ongoing], requirement=req)
    doing = rank_performance_lines([item for item in lines if item.ongoing], requirement=req)
    return done + doing


def _perf_amount_text(amount: float) -> str:
    if amount <= 0:
        return ""
    if amount >= 10000:
        wan = amount / 10000
        if abs(wan - round(wan)) < 0.005:
            return f"{int(round(wan))}万元"
        compact = f"{wan:.2f}".rstrip("0").rstrip(".")
        return f"{compact}万元"
    return f"{amount:,.2f}元"


def _fill_perf_table(table: Table, line: PerformanceLine | None) -> None:
    if line is None or not (line.projectName or "").strip():
        fill_perf_placeholders(table)
        return
    amount = _perf_amount_text(float(line.amountYuan or 0))
    name_hits = 0
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        label = _compact(row.cells[0].text)
        value = ""
        if "项目概况" in label or "履约情况" in label:
            value = (line.summary or "").strip()
        elif "规格" in label:
            value = (line.spec or "").strip()
        elif "买方联系人" in label:
            value = (line.contact or "").strip()
        elif "买方名称" in label:
            value = (line.client or "").strip()
        elif "合同价格" in label or "签约合同价" in label:
            value = amount
        elif "备注" in label:
            value = (line.note or "").strip()
        elif "项目名称" in label:
            name_hits += 1
            value = (line.projectName or "").strip() if name_hits == 1 else (
                line.location or line.projectName or ""
            ).strip()
        if value:
            _write_cell(row.cells[-1], value, size=10.5, bold=False, underline=False)


def _append_pdf_pages(
    doc: Document,
    pdf_path: Path,
    *,
    max_pages: int = 6,
    title: bool = True,
    leading_break: bool = True,
) -> int:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        logger.warning("pypdfium2 missing, skip qualification pages")
        return 0

    try:
        pdf = pdfium.PdfDocument(str(pdf_path))
    except Exception:
        logger.exception("open qualification pdf failed")
        return 0

    try:
        count = min(len(pdf), max_pages)
        if count <= 0:
            return 0
        from api.services.tenders.placeholders import _render_pdf_page_jpeg

        if leading_break:
            doc.add_page_break()
        if title:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run("附件：企业资质文件扫描件")
            _font(run, 16, bold=True)
        inserted = 0
        for i in range(count):
            try:
                page = pdf[i]
                buf = _render_pdf_page_jpeg(page)
                if i > 0:
                    doc.add_page_break()
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = cap.add_run(f"资质文件 第 {i + 1} 页")
                _font(cap_run, 10.5)
                pic = doc.add_paragraph()
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(buf, width=Cm(15.5))
                inserted += 1
            except Exception:
                logger.exception("render qualification page %s failed", i + 1)
        return inserted
    finally:
        try:
            pdf.close()
        except Exception:
            pass


def qualification_attach_notes(
    doc: Document,
    pdf_path: Path | None,
    *,
    title: bool = True,
    leading_break: bool = True,
) -> list[str]:
    if pdf_path and pdf_path.is_file():
        pages = _append_pdf_pages(
            doc, pdf_path, title=title, leading_break=leading_break
        )
        if pages:
            return [
                f"已插入资质文件 {pages} 页扫描件",
                "资质彩页/扫描件未单独编页，装订时请按招标书对图纸、彩页的约定处理",
            ]
        return ["资质 PDF 未能渲染为图片，Word 中未插入扫描件"]
    return [
        "未找到资质 PDF。可将《伟泰科技资质文件最终版.pdf》复制到 "
        "storage/tender-assets/weitai-qualifications.pdf 后重新生成"
    ]


_PAGE_BODY_CM = 22.0
# 与目录条目对应的章标题；投标函附录 / 支付条件也另起一页，避免挤在上一节签字盖章区。
_CHAPTER_PAGE_TITLES = {
    "目录",
    "投标函及投标函附录",
    "投标函附录",
    "支付条件",
    "分项报价表",
    "法定代表人身份证明",
    "授权委托书",
    "技术偏差表",
    "原厂生产承诺",
    "其他材料",
    "技术标实施方案",
    "投标承诺书",
}

# 签字盖章块起始：在其前插入空段，给盖章留出空间
_SIGN_BLOCK_MARKS = ("盖单位公章", "盖单位章")


def _p_has_page_br(p_el) -> bool:
    return any(br.get(qn("w:type")) == "page" for br in p_el.iter(qn("w:br")))


def _p_sectpr(p_el):
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is None:
        return None
    return p_pr.find(qn("w:sectPr"))


def _sect_starts_new_page(sect_pr) -> bool:
    if sect_pr is None:
        return False
    types = [el.get(qn("w:val")) for el in sect_pr.findall(qn("w:type"))]
    if not types:
        return True
    return types[0] not in {"continuous", "evenPage", "oddPage"}


def _set_section_type(sect_pr, value: str) -> None:
    if sect_pr is None:
        return
    for old in list(sect_pr.findall(qn("w:type"))):
        sect_pr.remove(old)
    typ = OxmlElement("w:type")
    typ.set(qn("w:val"), value)
    sect_pr.append(typ)


def _set_section_next_page(sect_pr) -> None:
    _set_section_type(sect_pr, "nextPage")


def _set_section_continuous(sect_pr) -> None:
    _set_section_type(sect_pr, "continuous")


def _normalize_section_page_numbers(doc: Document) -> None:
    """封面节原先从第 11 页起算（第五章残稿），改为全书从 1 连续编号。"""
    for i, section in enumerate(doc.sections):
        sp = section._sectPr
        for old in list(sp.findall(qn("w:pgNumType"))):
            sp.remove(old)
        if i == 0:
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:start"), "1")
            sp.append(pg)


def _normalize_body_section_breaks(doc: Document) -> None:
    """
    节间一律 continuous：模板里附录/支付条件等默认 nextPage 会多占页，
    导致目录 PAGEREF 与正文分页对不上。真正换页改由显式分页符完成。
    """
    for child in list(doc.element.body):
        if child.tag == qn("w:sectPr"):
            _set_section_continuous(child)
            continue
        if child.tag != qn("w:p"):
            continue
        sect = _p_sectpr(child)
        if sect is not None:
            _set_section_continuous(sect)


def _insert_page_break_paragraph_before(para: Paragraph) -> None:
    """插入仅含分页符的空段。比 pageBreakBefore / 分节 nextPage 更兼容各编辑器。"""
    if _prev_is_explicit_page_break(para):
        return
    br_p = OxmlElement("w:p")
    br_r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    br_r.append(br)
    br_p.append(br_r)
    para._p.addprevious(br_p)


def _plain_sectpr(doc: Document):
    sp = deepcopy(doc.sections[0]._sectPr)
    for tag in ("w:headerReference", "w:footerReference", "w:pgNumType", "w:titlePg"):
        for el in list(sp.findall(qn(tag))):
            sp.remove(el)
    _set_section_continuous(sp)
    return sp


def _end_section_at_p(p_el, sect_pr) -> None:
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        p_el.insert(0, p_pr)
    for old in list(p_pr.findall(qn("w:sectPr"))):
        p_pr.remove(old)
    p_pr.append(deepcopy(sect_pr))


def _para_section_index(doc: Document, para: Paragraph) -> int:
    idx = 0
    target = para._p
    for child in doc.element.body:
        if child is target:
            return idx
        if child.tag == qn("w:p") and _p_sectpr(child) is not None:
            idx += 1
    return idx


def _first_body_section_index(doc: Document) -> int:
    for para in doc.paragraphs:
        n = _compact(para.text)
        if n != "投标函及投标函附录":
            continue
        if _is_toc_list_line(para):
            continue
        return _para_section_index(doc, para)
    return 2


def _isolate_toc_section(doc: Document) -> None:
    """目录独占一页：目录前、后各插显式分页符，并拆成封面 / 目录 / 正文三节。"""
    toc_title: Paragraph | None = None
    first_after_toc: Paragraph | None = None
    seen_toc = False
    for para in doc.paragraphs:
        n = _compact(para.text)
        if n == "目录":
            toc_title = para
            seen_toc = True
            _set_page_break_before(para, False)
            continue
        if not seen_toc:
            continue
        if _p_sectpr(para._p) is not None:
            continue
        if _is_toc_list_line(para):
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading") and n and n != "目录":
            first_after_toc = para
            break
        if n and _is_chapter_start(para.text) and n != "目录":
            first_after_toc = para
            break

    if toc_title is not None:
        # 模板常为「目\t录」，规范成「目录」便于检索与排版
        if (toc_title.text or "").replace("\t", "").strip() == "目录":
            _write_plain(toc_title, "目录", size=16, bold=True)
        _insert_page_break_paragraph_before(toc_title)
        prev = toc_title._p.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            _end_section_at_p(prev, _plain_sectpr(doc))
    if first_after_toc is not None:
        _insert_page_break_paragraph_before(first_after_toc)
        prev = first_after_toc._p.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            _end_section_at_p(prev, _plain_sectpr(doc))


def _set_page_break_before(para: Paragraph, on: bool = True) -> None:
    p_pr = para._p.get_or_add_pPr()
    for old in list(p_pr.findall(qn("w:pageBreakBefore"))):
        p_pr.remove(old)
    if on:
        el = OxmlElement("w:pageBreakBefore")
        el.set(qn("w:val"), "true")
        p_pr.append(el)


def _is_chapter_start(text: str) -> bool:
    n = _compact(text)
    if not n:
        return False
    if n == "目录":
        return True
    if n in _CHAPTER_PAGE_TITLES:
        return True
    # 「（三）支付条件」「（二）投标函附录」等带序号标题
    if n.endswith("支付条件") and "投标函及" not in n:
        return True
    if n.endswith("投标函附录") and "及" not in n:
        return True
    if "近年完成的类似项目" in n:
        return True
    # 「正在供货」是业绩子表，不单独强制另起一页
    if n.startswith("附件：企业资质") or n.startswith("附件企业资质"):
        return True
    return False


def _prev_is_explicit_page_break(para: Paragraph) -> bool:
    el = para._p.getprevious()
    while el is not None:
        if el.tag == qn("w:tbl"):
            return False
        if el.tag == qn("w:p"):
            if _p_has_page_br(el) or _has_page_break_before(el):
                return True
            if _p_sectpr(el) is not None and _sect_starts_new_page(_p_sectpr(el)):
                return True
            text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
            if text:
                return False
            el = el.getprevious()
            continue
        el = el.getprevious()
    return False


def _mark_chapter_page_starts(doc: Document) -> None:
    """各章标题另起一页。优先显式分页符，已有换页则不重复。"""
    for para in doc.paragraphs:
        n = _compact(para.text)
        if n == "目录":
            continue
        if _is_toc_list_line(para):
            continue
        if not _is_chapter_start(para.text):
            continue
        if _prev_is_explicit_page_break(para):
            _set_page_break_before(para, False)
            continue
        _insert_page_break_paragraph_before(para)
        _set_page_break_before(para, False)


def _is_signature_block_start(text: str) -> bool:
    raw = text or ""
    n = _compact(raw)
    if not any(mark in n for mark in _SIGN_BLOCK_MARKS):
        return False
    return (
        n.startswith("投标人")
        or n.startswith("投标人名称")
        or "投标人" in n[:12]
        or "投 标 人" in raw
        or "投　标　人" in raw
    )


def _count_blank_paras_before(para: Paragraph, *, limit: int = 6) -> int:
    n = 0
    el = para._p.getprevious()
    while el is not None and n < limit:
        if el.tag != qn("w:p"):
            break
        if _p_has_page_br(el) or _has_page_break_before(el):
            break
        if _p_sectpr(el) is not None:
            break
        text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
        if text:
            break
        n += 1
        el = el.getprevious()
    return n


def _pad_signature_blocks(doc: Document, *, blanks: int = 3) -> None:
    """签字盖章行前插入空段，避免正文顶住签章区。"""
    for para in list(doc.paragraphs):
        if not _is_signature_block_start(para.text):
            continue
        have = _count_blank_paras_before(para)
        for _ in range(max(0, blanks - have)):
            blank = OxmlElement("w:p")
            para._p.addprevious(blank)


def _drawing_height_cm(el) -> float:
    total = 0.0
    for child in el.iter():
        if child.tag.split("}")[-1] != "extent":
            continue
        cy = child.get("cy")
        if cy:
            total += int(cy) / 914400 * 2.54
    return total


def _para_height_cm(para: Paragraph) -> float:
    drawn = _drawing_height_cm(para._p)
    if drawn:
        return min(drawn + 0.6, _PAGE_BODY_CM)
    text = (para.text or "").strip()
    if not text:
        return 0.28
    size = 12.0
    for run in para.runs:
        if run.font.size:
            size = max(size, float(run.font.size.pt))
            break
    line_h = max(0.55, size / 72 * 2.54 * 1.55)
    lines = max(1, (len(text) + 31) // 32)
    return min(line_h * lines, _PAGE_BODY_CM)


def _table_height_cm(table: Table) -> float:
    """估算表高。不读 cell.text（长技术参数/合并格在 python-docx 里很慢）。"""
    total = 0.35
    for row in table.rows:
        tr_pr = row._tr.find(qn("w:trPr"))
        row_cm = 0.0
        if tr_pr is not None:
            tr_h = tr_pr.find(qn("w:trHeight"))
            if tr_h is not None:
                raw = tr_h.get(qn("w:val"))
                if raw:
                    row_cm = int(raw) / 1440 * 2.54
        if row_cm <= 0:
            # 按 XML 文本长度粗估，避免 Table.cell.text 展开合并单元格
            longest = 0
            for tc in row._tr.findall(qn("w:tc")):
                text_len = sum(len(t.text or "") for t in tc.iter(qn("w:t")))
                longest = max(longest, text_len)
            lines = max(1, (longest + 17) // 18)
            row_cm = min(2.2, 0.52 * lines)
        total += max(row_cm, 0.45)
    return total


def _bookmark_names(el) -> list[str]:
    names: list[str] = []
    for start in el.iter(qn("w:bookmarkStart")):
        name = start.get(qn("w:name"))
        if name and not name.startswith("_"):
            names.append(name)
    return names


def _has_page_break_before(p_el) -> bool:
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is None:
        return False
    el = p_pr.find(qn("w:pageBreakBefore"))
    if el is None:
        return False
    return el.get(qn("w:val")) not in {"0", "false"}


def _next_page_section_starts(doc: Document) -> set[int]:
    """各节起始子节点下标。第一节是封面，其后 type=nextPage 的节另起一页。"""
    body = list(doc.element.body)
    starts: set[int] = set()
    first = 0
    section_i = 0
    for i, child in enumerate(body):
        sect = _p_sectpr(child) if child.tag == qn("w:p") else child if child.tag == qn("w:sectPr") else None
        if sect is None:
            continue
        if section_i > 0 and _sect_starts_new_page(sect):
            starts.add(first)
        section_i += 1
        first = i + 1
    return starts


def _estimate_bookmark_pages(doc: Document) -> dict[str, int]:
    """按显式分页符、nextPage 分节和下一块高度估算书签所在页，写入目录缓存。"""
    page = 1
    used = 0.0
    pages: dict[str, int] = {}
    new_section_at = _next_page_section_starts(doc)

    def force_new_page(*, hard: bool = False) -> None:
        """soft：仅在本页已有内容时换页；hard：正文中的强制分页一律 +1。"""
        nonlocal page, used
        if hard:
            page += 1
            used = 0.0
            return
        if used > 0.05:
            page += 1
            used = 0.0

    def consume(cm: float) -> None:
        nonlocal page, used
        remain = max(cm, 0.0)
        while remain > 0:
            space = _PAGE_BODY_CM - used
            if space <= 0.35:
                page += 1
                used = 0.0
                space = _PAGE_BODY_CM
            take = min(remain, space)
            used += take
            remain -= take

    for index, child in enumerate(list(doc.element.body)):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:tbl"):
            if index in new_section_at:
                force_new_page(hard=index > 0)
            table = Table(child, doc)
            for name in _bookmark_names(child):
                pages[name] = page
            consume(_table_height_cm(table))
            continue
        if child.tag != qn("w:p"):
            continue
        para = Paragraph(child, doc)
        bare_break = _p_has_page_br(child) and not (para.text or "").strip() and _drawing_height_cm(child) <= 0
        hard_break = index in new_section_at or _has_page_break_before(child) or bare_break
        if hard_break:
            if index > 0:
                force_new_page(hard=True)
            if bare_break:
                continue
        for name in _bookmark_names(child):
            if name.startswith("toc_"):
                pages[name] = page
        if not bare_break:
            consume(_para_height_cm(para))
    return pages


def _pageref_bookmark(para: Paragraph) -> str | None:
    for el in para._p.iter(qn("w:instrText")):
        parts = (el.text or "").split()
        if len(parts) >= 2 and parts[0].upper() == "PAGEREF":
            return parts[1]
    return None


def _set_pageref_cache(para: Paragraph, number: str) -> None:
    in_result = False
    for run in para._p.iter(qn("w:r")):
        fld = run.find(qn("w:fldChar"))
        if fld is not None:
            kind = fld.get(qn("w:fldCharType"))
            if kind == "separate":
                in_result = True
            elif kind == "end":
                in_result = False
            continue
        if not in_result:
            continue
        text_el = run.find(qn("w:t"))
        if text_el is not None:
            text_el.text = number
            return


def _apply_toc_page_numbers(doc: Document, pages: dict[str, int]) -> None:
    for para in doc.paragraphs:
        bookmark = _pageref_bookmark(para)
        if bookmark:
            _set_pageref_cache(para, str(pages.get(bookmark, 1)))


def _finalize_pagination(doc: Document, fmt=None) -> dict[str, int]:
    """分页：分节 continuous + 章标题显式分页。目录页码仅在招标书要求时预填。"""
    _normalize_section_page_numbers(doc)
    _normalize_body_section_breaks(doc)
    _isolate_toc_section(doc)
    _mark_chapter_page_starts(doc)
    _pad_signature_blocks(doc)
    if fmt is not None and not fmt.tocNeedPageNos:
        return {}
    pages: dict[str, int] = {}
    for i, (_title, bookmark, _old) in enumerate(_TOC_ITEMS):
        pages[bookmark] = int(toc_page_cache(i, fmt)) if fmt is not None else int(_old)
    pages.setdefault("toc_perf", 9)
    pages.setdefault("toc_commit", 13)
    _apply_toc_page_numbers(doc, pages)
    return pages


def _apply_template_format(doc: Document, brief: BidBrief) -> list[str]:
    """公司模板组卷：页脚按起编位置处理；页边距仅在招标书写明时改。"""
    fmt = brief.documentFormat
    notes = format_brief_notes(fmt)
    if fmt.coverNeedSeal:
        notes.append("招标书要求封面加盖公章，请在打印后于封面预留处盖章")
    body_si = _first_body_section_index(doc)
    for i, section in enumerate(doc.sections):
        if fmt.specified:
            apply_section_page(section, fmt)
        numbered, restart = section_page_flags(i, fmt, body_section=body_si)
        apply_footer_page_number(section, fmt, numbered=numbered, restart=restart)
    return notes


def build_bid_docx(
    brief: BidBrief,
    dest: Path,
    *,
    qualification_pdf: Path | None = None,
    catalog_slots: list | None = None,
    catalog_media: dict | None = None,
) -> tuple[Path, list[str]]:
    warnings: list[str] = []
    if brief.bidPriceYuan <= 0:
        warnings.append("投标总价为 0，请确认是否已填写报价")
    if not (brief.legalPersonName or "").strip():
        warnings.append("未填写法定代表人姓名，身份证明中该栏为空白")
    pct = brief.prepaidPct + brief.arrivalPct + brief.settlementPct + brief.warrantyPct
    if pct != 100:
        warnings.append(f"支付比例合计为 {pct}%，建议为 100%")

    template = find_chapter5_template()
    if template is None:
        raise FileNotFoundError(
            "未找到第五章投标文件格式模板。请将空白稿保存为 "
            "apps/api/services/tenders/templates/chapter5.docx"
        )

    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, dest)
    doc = Document(str(dest))
    _pin_cjk_fonts(doc)
    _hide_proofing_marks(doc)
    slots = catalog_slots if catalog_slots is not None else collect_slots(brief.extraPlaceholders)
    media = catalog_media if catalog_media is not None else attachments_for_slots(slots)
    inlined_ids: set[str] = set()
    quote_anchor = _fill_paragraphs(doc, brief, media=media, inlined=inlined_ids)
    _ensure_toc_entries(
        doc,
        include_commitment=bool(brief.includeCommitment),
        with_pages=bool(brief.documentFormat.tocNeedPageNos),
        fmt=brief.documentFormat,
        brief=brief,
    )
    warnings.extend(_fill_technical_section(doc, brief, media))
    _fill_tables(doc, brief)
    warnings.extend(_fill_quote_section(doc, brief, quote_anchor))
    _normalize_blank_underlines(doc)

    if brief.includeCommitment:
        append_commitment_letter(doc, brief)
        for para in reversed(doc.paragraphs):
            if _compact(para.text) == "投标承诺书":
                _bookmark_paragraph(para, "toc_commit")
                break
        warnings.append("已附「投标承诺书」（附件五），请核对后签字盖章")

    if brief.attachQualifications:
        warnings.extend(qualification_attach_notes(doc, qualification_pdf))

    _finalize_pagination(doc, brief.documentFormat)
    for i, section in enumerate(doc.sections):
        _ensure_header(section, brief.projectName, cover=(i == 0))
    warnings.extend(_apply_template_format(doc, brief))
    doc.save(str(dest))
    tail = "封面与目录不编页码，从正文首页起编。目录独占一页；投标函附录/支付条件等签字节另起一页，签章前已留白。"
    if brief.documentFormat.tocNeedPageNos:
        tail = (
            "目录独占一页；投标函附录/支付条件等签字节另起一页，签章前已留白。"
            "目录页码已预填；若与正文不符，可在 Word/WPS 中右键目录域「更新域」。"
        )
    warnings.append(tail)
    return dest, warnings
