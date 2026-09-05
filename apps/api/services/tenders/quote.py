"""招标报价清单：只用本次邀请书/工程量文件或页面分项。工程量固定，单价可随投标总价折算。"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, replace
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from api.services.tenders.schema import BidBrief

_TWO = Decimal("0.01")
SHEET_INC_TAX = Decimal("518800")
SECOND_ROUND_YUAN = Decimal("439700")
_MAX_LINES = 80

_SKIP_NAME = ("不含税合计", "含税合计", "税率", "合计", "小计", "总计", "备注")


def _q(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("￥", "").replace("¥", "").replace("元", "")
        if not text or text.startswith("="):
            return Decimal("0")
        text = text.replace("%", "")
        try:
            return Decimal(text).quantize(_TWO, rounding=ROUND_HALF_UP)
        except Exception:
            return Decimal("0")
    try:
        return Decimal(str(value)).quantize(_TWO, rounding=ROUND_HALF_UP)
    except Exception:
        return Decimal("0")


def _norm(text: object) -> str:
    return re.sub(r"[\s/（）()【】\[\]:：]", "", str(text or "")).lower()


@dataclass(frozen=True)
class QuoteLine:
    seq: str
    name: str
    spec: str
    unit: str
    qty: Decimal
    unit_price: Decimal
    amount: Decimal


@dataclass(frozen=True)
class QuoteSheet:
    title: str
    lines: tuple[QuoteLine, ...]
    tax_rate: Decimal
    total_ex_tax: Decimal
    total_inc_tax: Decimal
    note: str
    source_inc_tax: Decimal


def _qty(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text or text.startswith("="):
            return Decimal("0")
        try:
            return Decimal(text)
        except Exception:
            return Decimal("0")
    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0")


def is_template_sheet(sheet: QuoteSheet) -> bool:
    return "高途智成港" in (sheet.title or "")


def _col_kind(cell: object) -> str | None:
    n = _norm(cell)
    if not n:
        return None
    if n in {"序号", "编号", "no", "num"} or n.startswith("序号"):
        return "seq"
    # 技术参数/明细列优先于「单价」等，避免表头含「参数」时误判
    if any(
        k in n
        for k in (
            "技术参数",
            "技术要求",
            "参数要求",
            "规格型号",
            "规格参数",
            "特征描述",
            "工作内容",
            "设备参数",
            "配置要求",
        )
    ):
        return "spec"
    if n in {"规格", "参数", "明细"} or n.endswith("明细"):
        return "spec"
    if n in {"单位", "计量单位"} or n.endswith("单位"):
        return "unit"
    if any(k in n for k in ("工程量", "数量")):
        return "qty"
    if any(k in n for k in ("综合单价", "不含税单价", "单价")):
        return "price"
    if n in {"合价", "金额"} or n.endswith("合价") or "不含税合价" in n:
        return "amount"
    if any(k in n for k in ("设备名称", "项目名称", "货物名称", "物料名称", "品名")):
        return "name"
    if n in {"设备", "名称", "项目", "货物", "物料"}:
        return "name"
    # 宽松兜底：「规格」单独出现已在上；含「参数」但不像单价列
    if "参数" in n and "单价" not in n and "合价" not in n:
        return "spec"
    return None


def _cell_text(value: object) -> str:
    """把单元格转成可读文本，保留技术参数多行换行。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    # openpyxl CellRichText / TextBlock 序列
    try:
        from openpyxl.cell.rich_text import CellRichText

        if isinstance(value, CellRichText):
            return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    except Exception:
        pass
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _header_map(cells: list[object]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for i, cell in enumerate(cells):
        kind = _col_kind(cell)
        if kind and kind not in mapping:
            mapping[kind] = i
    if "name" in mapping and ("qty" in mapping or "price" in mapping or "amount" in mapping):
        return mapping
    return None


def _skip_name(name: str) -> bool:
    compact = re.sub(r"\s+", "", name or "")
    return any(compact.startswith(k) or compact == k for k in _SKIP_NAME)


def _tax_rate(value: object) -> Decimal:
    raw = _q(value)
    if raw <= 0:
        return Decimal("0")
    if raw > 1:
        return (raw / Decimal("100")).quantize(Decimal("0.0001"))
    return raw


def _normalize_cell(value: object) -> object:
    if value is None or isinstance(value, (int, float, Decimal, bool)):
        return value
    return _cell_text(value)


def sheet_from_rows(rows: list[list[object]], *, title_hint: str = "") -> QuoteSheet | None:
    """从二维表抽出报价行。表头需能识别名称+数量/单价。"""
    header_i = -1
    mapping: dict[str, int] = {}
    title = (title_hint or "").strip()
    normalized_rows: list[list[object]] = [[_normalize_cell(c) for c in row] for row in rows]
    for i, row in enumerate(normalized_rows[:20]):
        found = _header_map(row)
        if found:
            header_i = i
            mapping = found
            if not title:
                for prev in reversed(normalized_rows[:i]):
                    head = _cell_text(prev[0] if prev else "")
                    if head and not _header_map(prev):
                        title = head
                        break
            break
    if header_i < 0:
        return None

    lines: list[QuoteLine] = []
    tax_rate = Decimal("0.13")
    total_ex = Decimal("0")
    total_inc = Decimal("0")
    note = ""
    for row in normalized_rows[header_i + 1 :]:
        cells = list(row) + [None] * 8
        name = _cell_text(cells[mapping["name"]]) if "name" in mapping else ""
        seq_raw = cells[mapping["seq"]] if "seq" in mapping else None
        label = name or _cell_text(seq_raw)
        if str(seq_raw or "").startswith("备注") or (label.startswith("备注") and not name):
            note = _cell_text(seq_raw or name)
            continue
        if _skip_name(name) or _skip_name(label):
            amount_cell = cells[mapping["amount"]] if "amount" in mapping else None
            qty_cell = cells[mapping["qty"]] if "qty" in mapping else None
            if name.startswith("不含税合计") or label.startswith("不含税合计"):
                total_ex = _q(amount_cell)
            elif name.startswith("含税合计") or label.startswith("含税合计"):
                total_inc = _q(amount_cell) or _q(qty_cell)
            elif "税率" in name or "税率" in label:
                total_inc_or_rate = amount_cell if amount_cell not in (None, "") else qty_cell
                parsed = _tax_rate(total_inc_or_rate)
                if parsed > 0:
                    tax_rate = parsed
            continue
        if not name:
            continue
        qty = _qty(cells[mapping["qty"]]) if "qty" in mapping else Decimal("0")
        price = _q(cells[mapping["price"]]) if "price" in mapping else Decimal("0")
        amount = _q(cells[mapping["amount"]]) if "amount" in mapping else Decimal("0")
        if amount <= 0 and qty and price:
            amount = (price * qty).quantize(_TWO, rounding=ROUND_HALF_UP)
        if qty <= 0 and amount <= 0 and price <= 0:
            continue
        spec = _cell_text(cells[mapping["spec"]]) if "spec" in mapping else ""
        unit = _cell_text(cells[mapping["unit"]]) if "unit" in mapping else ""
        seq = _cell_text(seq_raw) if seq_raw not in (None, "") else str(len(lines) + 1)
        if seq.endswith(".0") and seq[:-2].isdigit():
            seq = seq[:-2]
        lines.append(
            QuoteLine(
                seq=seq,
                name=name,
                spec=spec,
                unit=unit,
                qty=qty,
                unit_price=price,
                amount=amount,
            )
        )
        if len(lines) >= _MAX_LINES:
            break
    if not lines:
        return None
    if total_ex <= 0:
        total_ex = sum((ln.amount for ln in lines), Decimal("0"))
    if tax_rate <= 0:
        tax_rate = Decimal("0.13")
    if total_inc <= 0:
        total_inc = (total_ex * (1 + tax_rate)).quantize(_TWO, rounding=ROUND_HALF_UP)
    return QuoteSheet(
        title=title or "工程量清单",
        lines=tuple(lines),
        tax_rate=tax_rate,
        total_ex_tax=total_ex,
        total_inc_tax=total_inc,
        note=note,
        source_inc_tax=total_inc,
    )


def parse_quote_path(path: Path | None) -> QuoteSheet | None:
    if path is None or not path.is_file():
        return None
    ext = path.suffix.lower().lstrip(".")
    if ext in {"xlsx", "xlsm"}:
        return _parse_xlsx(path)
    if ext == "csv":
        return _parse_csv(path)
    if ext == "docx":
        return _parse_docx(path)
    return None


def parse_quote_from_text(text: str, *, title_hint: str = "") -> QuoteSheet | None:
    """从抽字结果里找用 | 拼出的表格（Word/PDF 入库文本）。"""
    best: QuoteSheet | None = None
    block: list[str] = []

    def flush() -> None:
        nonlocal best, block
        if len(block) >= 2:
            rows = [[c.strip() for c in line.split("|")] for line in block]
            sheet = sheet_from_rows(rows, title_hint=title_hint)
            if sheet is not None and (best is None or len(sheet.lines) > len(best.lines)):
                best = sheet
        block = []

    for raw in (text or "").splitlines():
        line = raw.strip()
        if line.count("|") >= 2:
            block.append(line.strip("|"))
        else:
            flush()
    flush()
    return best


def load_quote_sheet(path: Path | None) -> QuoteSheet:
    parsed = parse_quote_path(path)
    if parsed is not None:
        return parsed
    raise ValueError("未能解析工程量清单")


def sheet_from_brief(brief: BidBrief) -> QuoteSheet | None:
    items = list(brief.quoteLines or [])
    if not items:
        return None
    lines: list[QuoteLine] = []
    for i, item in enumerate(items, 1):
        name = (item.name or "").strip()
        if not name or _skip_name(name):
            continue
        qty = Decimal(str(item.qty or 0))
        price = _q(item.unitPrice)
        amount = _q(item.amount)
        if amount <= 0 and qty and price:
            amount = (price * qty).quantize(_TWO, rounding=ROUND_HALF_UP)
        if qty <= 0 and amount <= 0 and price <= 0:
            continue
        lines.append(
            QuoteLine(
                seq=(item.seq or str(i)).strip() or str(i),
                name=name,
                spec=(item.spec or "").strip(),
                unit=(item.unit or "").strip(),
                qty=qty,
                unit_price=price,
                amount=amount,
            )
        )
    if not lines:
        return None
    rate = Decimal(str(brief.quoteTaxRate or 0.13))
    if rate > 1:
        rate = rate / Decimal("100")
    if rate <= 0:
        rate = Decimal("0.13")
    total_ex = sum((ln.amount for ln in lines), Decimal("0"))
    stored_inc = _q(brief.quoteSourceIncTax)
    total_inc = stored_inc if stored_inc > 0 else (total_ex * (1 + rate)).quantize(_TWO, rounding=ROUND_HALF_UP)
    return QuoteSheet(
        title=(brief.quoteTitle or "").strip() or "工程量清单",
        lines=tuple(lines),
        tax_rate=rate,
        total_ex_tax=total_ex,
        total_inc_tax=total_inc,
        note="",
        source_inc_tax=total_inc,
    )


def resolve_quote_sheet(brief: BidBrief, template_path: Path | None = None) -> tuple[QuoteSheet | None, str]:
    del template_path
    custom = sheet_from_brief(brief)
    if custom is not None:
        return custom, (brief.quoteSource or "form")
    return None, "none"


def attach_sheet(brief: BidBrief, sheet: QuoteSheet, *, source: str) -> BidBrief:
    data = brief.model_dump()
    data["quoteLines"] = [
        {
            "seq": ln.seq,
            "name": ln.name,
            "spec": ln.spec,
            "unit": ln.unit,
            "qty": float(ln.qty),
            "unitPrice": float(ln.unit_price),
            "amount": float(ln.amount),
        }
        for ln in sheet.lines
    ]
    data["quoteTitle"] = sheet.title
    data["quoteTaxRate"] = float(sheet.tax_rate)
    data["quoteSourceIncTax"] = float(sheet.source_inc_tax)
    data["quoteSource"] = source
    if float(data.get("bidPriceYuan") or 0) <= 0 and sheet.total_inc_tax > 0:
        data["bidPriceYuan"] = float(sheet.total_inc_tax)
    return BidBrief.model_validate(data)


def quote_payload_from_patch(patch: dict[str, Any]) -> QuoteSheet | None:
    """把模型 JSON 清单转成表。识别邀请书不再调用，避免用其它项目台数顶替。"""
    raw = patch.get("quoteLines")
    if not isinstance(raw, list) or not raw:
        return None
    rows: list[list[object]] = [["序号", "设备", "技术参数要求", "单位", "数量", "单价", "合价"]]
    for i, item in enumerate(raw, 1):
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        rows.append(
            [
                item.get("seq") or i,
                name,
                item.get("spec") or "",
                item.get("unit") or "",
                item.get("qty") or 0,
                item.get("unitPrice") or item.get("unit_price") or 0,
                item.get("amount") or 0,
            ]
        )
    title = str(patch.get("quoteTitle") or "").strip()
    sheet = sheet_from_rows(rows, title_hint=title)
    if sheet is None:
        return None
    tax = _tax_rate(patch.get("quoteTaxRate"))
    if tax > 0:
        total_ex = sheet.total_ex_tax
        total_inc = (total_ex * (1 + tax)).quantize(_TWO, rounding=ROUND_HALF_UP)
        sheet = replace(sheet, tax_rate=tax, total_inc_tax=total_inc, source_inc_tax=total_inc)
    return sheet


def scale_quote(sheet: QuoteSheet, target_inc_tax: Decimal) -> QuoteSheet:
    """工程量不变，单价/合价按含税总价比例折算，并使含税合计等于目标总价。"""
    target = _q(target_inc_tax)
    source = sheet.total_inc_tax if sheet.total_inc_tax > 0 else sheet.source_inc_tax
    if target <= 0 or source <= 0:
        return sheet
    if target == _q(source):
        return sheet
    factor = target / source
    scaled: list[QuoteLine] = []
    for line in sheet.lines:
        unit = (line.unit_price * factor).quantize(_TWO, rounding=ROUND_HALF_UP)
        amt = (unit * line.qty).quantize(_TWO, rounding=ROUND_HALF_UP)
        scaled.append(replace(line, unit_price=unit, amount=amt))
    rate = sheet.tax_rate if sheet.tax_rate > 0 else Decimal("0.13")
    target_ex = (target / (1 + rate)).quantize(_TWO, rounding=ROUND_HALF_UP)
    drift = target_ex - sum((ln.amount for ln in scaled), Decimal("0"))
    if scaled and drift != 0:
        last = scaled[-1]
        new_amt = last.amount + drift
        new_unit = (new_amt / last.qty).quantize(_TWO, rounding=ROUND_HALF_UP) if last.qty else last.unit_price
        scaled[-1] = replace(last, unit_price=new_unit, amount=new_amt)
    return replace(sheet, lines=tuple(scaled), total_ex_tax=target_ex, total_inc_tax=target)


def apply_traffic_note(spec: str, note: str) -> str:
    text = (note or "").strip()
    if not text or "流量" not in (spec or ""):
        return spec
    return (
        spec.replace("（备注质保期外的流量年收费标准？）", text)
        .replace("（备注质保期外的流量年收费标准?）", text)
    )


def money_text(value: Decimal) -> str:
    return f"{_q(value):,.2f}"


def qty_text(value: Decimal) -> str:
    if value == value.to_integral_value():
        return str(int(value))
    return f"{value.normalize()}"


def _worksheet_matrix(ws, *, max_col: int = 16, max_row: int = 300) -> list[list[object]]:
    """读取工作表为二维表；合并单元格取左上角值，避免技术参数列落空。"""
    merge_values: dict[tuple[int, int], object] = {}
    for rng in getattr(ws, "merged_cells", None).ranges if getattr(ws, "merged_cells", None) else []:
        top = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merge_values[(r, c)] = top

    last_row = min(max_row, int(ws.max_row or 0) or max_row)
    last_col = min(max_col, int(ws.max_column or 0) or max_col)
    rows: list[list[object]] = []
    for r in range(1, last_row + 1):
        row: list[object] = []
        for c in range(1, last_col + 1):
            val = ws.cell(r, c).value
            if val is None and (r, c) in merge_values:
                val = merge_values[(r, c)]
            row.append(val)
        rows.append(row)
    return rows


def _parse_xlsx(path: Path) -> QuoteSheet | None:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=False, rich_text=True)
    except TypeError:
        wb = load_workbook(path, data_only=False)
    try:
        best: QuoteSheet | None = None
        for ws in wb.worksheets:
            rows = _worksheet_matrix(ws)
            title_hint = _cell_text(ws.cell(1, 1).value)
            sheet = sheet_from_rows(rows, title_hint=title_hint)
            # 优先保留带技术参数明细的表；行数相同则取有 spec 更多的
            if sheet is None:
                continue
            if best is None:
                best = sheet
                continue
            best_specs = sum(1 for ln in best.lines if (ln.spec or "").strip())
            cur_specs = sum(1 for ln in sheet.lines if (ln.spec or "").strip())
            if len(sheet.lines) > len(best.lines) or (
                len(sheet.lines) == len(best.lines) and cur_specs > best_specs
            ):
                best = sheet
        return best
    finally:
        wb.close()


def _parse_csv(path: Path) -> QuoteSheet | None:
    raw = path.read_bytes()
    text = ""
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        return None
    reader = csv.reader(io.StringIO(text))
    rows = [list(row) for row in reader]
    return sheet_from_rows(rows)


def _parse_docx(path: Path) -> QuoteSheet | None:
    from docx import Document

    doc = Document(str(path))
    best: QuoteSheet | None = None
    for table in doc.tables:
        rows = [[c.text for c in row.cells] for row in table.rows]
        sheet = sheet_from_rows(rows)
        if sheet is not None and (best is None or len(sheet.lines) > len(best.lines)):
            best = sheet
    return best


__all__ = [
    "QuoteLine",
    "QuoteSheet",
    "SECOND_ROUND_YUAN",
    "SHEET_INC_TAX",
    "apply_traffic_note",
    "attach_sheet",
    "is_template_sheet",
    "load_quote_sheet",
    "money_text",
    "parse_quote_from_text",
    "parse_quote_path",
    "qty_text",
    "quote_payload_from_patch",
    "resolve_quote_sheet",
    "scale_quote",
    "sheet_from_brief",
    "sheet_from_rows",
]
