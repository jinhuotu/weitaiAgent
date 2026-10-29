"""上传后台调度、持久状态与章节查询反馈；不连接生产库/向量服务。"""

import asyncio
from contextlib import asynccontextmanager
from io import BytesIO
from threading import Event
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document

from api.services.tenders import bid_archive as archive
from api.services.tenders import bid_marks
from db.models.tender import TenderRecord


class MemorySession:
    def __init__(self, row=None):
        self.row = row
        self.commits = 0

    def add(self, row):
        self.row = row

    async def commit(self):
        self.commits += 1
        if self.row and not self.row.id:
            self.row.id = 41

    async def refresh(self, row):
        pass

    async def get(self, model, pk):
        return self.row


def document_bytes():
    doc = Document()
    doc.add_paragraph("一、售后服务")
    doc.add_paragraph("售后服务响应时间为两小时，备件供应周期不超过十五个工作日，"
                      "质保期内免费更换损坏部件，并保留巡检记录。")
    data = BytesIO()
    doc.save(data)
    return data.getvalue()


def job_session(monkeypatch, session):
    @asynccontextmanager
    async def factory():
        yield session
    monkeypatch.setattr(archive, "_job_session", factory)


@pytest.mark.asyncio
@pytest.mark.parametrize("outcome", ["won", "lost", ""])
async def test_docx_upload_persists_pending_before_schedule(monkeypatch, outcome):
    db = MemorySession()

    def scheduled(**kwargs):
        assert db.commits >= 2
        assert archive.archive_state(db.row)["status"] == "processing"
        assert kwargs["record_id"] == 41
        assert "record:41" in kwargs["chapters"][0]["tags"]

    monkeypatch.setattr(archive, "schedule_upload", scheduled)
    out = await archive.accept_history_upload(
        db, data=document_bytes(), filename="历史商务标.docx", user_id=3,
        username="tester", outcome=outcome, record_public_id="", admin=False,
    )
    assert out["recordId"] == db.row.public_id
    assert out["status"] == "processing"
    assert out["queued"] == 1
    assert db.row.status == (outcome or "processing")


def test_spawn_without_loop_executes_in_thread():
    completed = Event()
    async def work():
        assert asyncio.get_running_loop().is_running()
        assert archive._worker_loop.get()
        completed.set()
    archive._spawn(work())
    assert completed.wait(3), "无事件循环时也必须执行任务"


@pytest.mark.asyncio
async def test_spawn_with_loop_retains_task():
    completed = asyncio.Event()
    async def work():
        completed.set()
    archive._spawn(work())
    await asyncio.wait_for(completed.wait(), 1)


@pytest.mark.asyncio
@pytest.mark.parametrize("failure_stage", ["index", "bind", "none"])
async def test_job_status_and_reason_survive_partial_failure(monkeypatch, failure_stage):
    row = TenderRecord(id=41, public_id="rec", brief_json={"projectName": "保留项目"},
                       status="won", warnings=["保留警告"])
    archive._set_archive_state(row, "processing")
    db = MemorySession(row)
    job_session(monkeypatch, db)
    index = AsyncMock(return_value=[{"volume": "business", "title": "售后", "docId": "doc"}])
    bind = AsyncMock()
    if failure_stage != "none":
        (index if failure_stage == "index" else bind).side_effect = RuntimeError("测试入库失败原因")
    monkeypatch.setattr(archive, "_index", index)
    monkeypatch.setattr(bid_marks, "bind_archived", bind)
    await archive._upload_job(3, "tester", [], 41)
    state = archive.archive_state(row)
    assert state["status"] == ("completed" if failure_stage == "none" else "failed")
    assert state["archiveNote"] == ("" if failure_stage == "none" else "测试入库失败原因")
    assert row.brief_json["projectName"] == "保留项目"
    assert row.status == "won"
    assert "保留警告" in row.warnings
    assert db.commits


@pytest.mark.asyncio
async def test_schedule_failure_returns_failed_state(monkeypatch):
    def rejected(**kwargs):
        raise RuntimeError("调度器不可用")
    monkeypatch.setattr(archive, "schedule_upload", rejected)
    db = MemorySession()
    out = await archive.accept_history_upload(
        db, data=document_bytes(), filename="历史标书.docx", user_id=3,
        username="tester", outcome="lost", record_public_id="", admin=False,
    )
    assert out["queued"] == 0
    assert out["status"] == "failed"
    assert out["archiveNote"] == "调度器不可用"
    assert archive.archive_state(db.row)["status"] == "failed"


@pytest.mark.asyncio
@pytest.mark.parametrize("status", ["processing", "completed", "failed"])
async def test_chapters_exposes_persisted_status_after_owner_check(monkeypatch, status):
    from api.services.tenders import records
    row = TenderRecord(id=41, public_id="rec", user_id=3, brief_json={})
    archive._set_archive_state(row, status, "失败原因" if status == "failed" else "")
    monkeypatch.setattr(records, "_get_row", AsyncMock(return_value=row))
    monkeypatch.setattr(bid_marks, "chapters_for_record", AsyncMock(return_value=[]))
    monkeypatch.setattr(bid_marks, "_marks_of", AsyncMock(return_value=[]))
    data = await bid_marks.list_history_chapters(
        None, "rec", user=SimpleNamespace(id=3), admin=False,
    )
    assert data["status"] == status
    assert data["chapters"] == []
    assert data["archiveNote"] == ("失败原因" if status == "failed" else "")
    with pytest.raises(Exception) as denied:
        await bid_marks.list_history_chapters(None, "rec", user=SimpleNamespace(id=4), admin=False)
    assert denied.value.status_code == 403


@pytest.mark.asyncio
@pytest.mark.parametrize("failed", [False, True])
async def test_status_and_marks_persist_across_database_sessions(monkeypatch, tmp_path, failed):
    """真实 SQLite 事务验证：任务结束后用新会话查询，不依赖内存中的 row。"""
    from sqlalchemy import Integer, MetaData, create_engine
    from sqlalchemy.orm import Session

    from db.models.tender import TenderApprovalLog, TenderBidMark

    engine = create_engine(f"sqlite:///{tmp_path / 'archive.sqlite'}")
    metadata = MetaData()
    for model in (TenderRecord, TenderApprovalLog, TenderBidMark):
        table = model.__table__.to_metadata(metadata)
        table.c.id.type = Integer()  # SQLite 主键自增要求 INTEGER。
    metadata.create_all(engine)

    class SessionAdapter:
        def __init__(self, session):
            self.session = session

        def add(self, row):
            self.session.add(row)

        async def get(self, *args):
            return self.session.get(*args)

        async def execute(self, statement):
            return self.session.execute(statement)

        async def commit(self):
            self.session.commit()

        async def rollback(self):
            self.session.rollback()

    @asynccontextmanager
    async def factory():
        with Session(engine, expire_on_commit=False) as session:
            yield SessionAdapter(session)

    monkeypatch.setattr(archive, "_job_session", factory)
    saved = [{"volume": "business", "title": "售后服务", "docId": "doc1", "outcome": "won"}]
    index = AsyncMock(return_value=saved)
    if failed:
        index.side_effect = RuntimeError("持久化验证失败原因")
    monkeypatch.setattr(archive, "_index", index)
    monkeypatch.setattr(bid_marks, "chapters_for_record", AsyncMock(return_value=[]))
    try:
        with Session(engine) as session:
            row = TenderRecord(public_id="persistent", user_id=3, username="tester",
                               docx_file="", status="won", brief_json={})
            archive._set_archive_state(row, "processing")
            session.add(row)
            session.commit()
            pk = row.id
        await archive._upload_job(3, "tester", [], pk)
        async with factory() as db:
            data = await bid_marks.list_history_chapters(
                db, "persistent", user=SimpleNamespace(id=3), admin=False,
            )
        assert data["status"] == ("failed" if failed else "completed")
        assert data["archiveNote"] == ("持久化验证失败原因" if failed else "")
        if not failed:
            assert data["chapters"][0]["verdict"] == "good"
            assert data["chapters"][0]["docId"] == "doc1"
    finally:
        engine.dispose()
