from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image

from api.services.tenders.pdf_format import _attach_library
from api.services.tenders.scan_placement import place_library_scans
from api.services.tenders.schema import BidBrief, PlaceholderItem

CASES = [
    ("id_legal", "法定代表人身份证"),
    ("id_agent", "委托代理人身份证"),
    ("license", "营业执照"),
    ("finance", "纳税证明"),
    ("social", "社保证明"),
    ("perf", "业绩合同"),
    ("bank_permit", "开户许可证"),
    ("iso", "ISO质量管理体系证书"),
    ("product", "产品检测报告"),
    ("bond", "投标保证金回单"),
    ("credit", "信用中国查询页"),
    ("qualification", "企业资质证书"),
    ("statement", "中小企业声明函"),
]


def images(tmp_path, key, count=3):
    paths = []
    for n in range(count):
        path = tmp_path / f"{key}-{n}.png"
        Image.new("RGB", (120, 90), (sum(key.encode()) % 170 + n * 20, 100, 160)).save(path)
        paths.append(path)
    return paths


def drawings(el):
    return list(el.iter(qn("w:drawing")))


@pytest.mark.parametrize("key,title", CASES)
def test_every_type_is_anchored_after_explicit_holder_not_on_heading(tmp_path, key, title):
    doc = Document()
    doc.add_paragraph("目录")
    toc = doc.add_paragraph("一、" + title)
    toc.add_run().add_tab()
    heading = doc.add_paragraph("一、" + title, style="Heading 1")
    doc.add_paragraph("本节正文保持原样。")
    holder = doc.add_paragraph("此处粘贴" + title + "复印件（在此粘贴扫描件）")
    end = doc.add_paragraph("二、下一章", style="Heading 1")
    original = heading._p.xml
    files = images(tmp_path, key)
    slots = [PlaceholderItem(key=key, title=title)]
    _attach_library(doc, BidBrief(), {key: files}, slots=slots)
    assert heading._p.xml == original
    assert not drawings(toc._p) and not drawings(holder._p)
    assert len(drawings(doc.element.body)) == (
        (len(files) + 1) // 2 if key.startswith("id_") else len(files)
    )
    children = list(doc.element.body)
    for p in children:
        if drawings(p):
            assert children.index(holder._p) < children.index(p) < children.index(end._p)
            assert not "".join(t.text or "" for t in p.iter(qn("w:t")))
    once = doc.element.xml
    _attach_library(doc, BidBrief(), {key: files}, slots=slots)
    assert doc.element.xml == once


def test_multiple_materials_never_use_each_others_media(tmp_path):
    doc = Document()
    holders = []
    media = {}
    for key, title in CASES[:7]:
        doc.add_paragraph("一、" + title, style="Heading 1")
        holders.append((key, doc.add_paragraph("附：" + title + "扫描件")))
        media[key] = images(tmp_path, key, 1)
    _attach_library(doc, BidBrief(), media)
    for key, holder in holders:
        p = holder._p.getnext()
        assert drawings(p)
        rid = next(p.iter(qn("a:blip"))).get(qn("r:embed"))
        with Image.open(BytesIO(doc.part.related_parts[rid].blob)) as scan:
            assert abs(scan.getpixel((20, 20))[0] - sum(key.encode()) % 170) <= 3
        assert not drawings(holder._p)
    assert len(drawings(doc.element.body)) == len(holders)


def test_generic_bordered_blank_is_typed_by_preceding_placeholder(tmp_path):
    doc = Document()
    doc.add_paragraph("一、营业执照", style="Heading 1")
    doc.add_paragraph("此处粘贴营业执照扫描件")
    holder = doc.add_paragraph()
    holder._p.get_or_add_pPr().append(OxmlElement("w:pBdr"))
    files = images(tmp_path, "license", 1)
    _attach_library(doc, BidBrief(), {"license": files})
    # Adjacent text and a blank frame represent a single physical placeholder.
    assert len(drawings(doc.element.body)) == 1
    assert drawings(holder._p)


def test_blank_cell_frame_expands_and_images_fit_cell(tmp_path):
    doc = Document()
    doc.add_paragraph("一、社保证明", style="Heading 1")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).width = Cm(5)
    table.rows[0].height = Cm(1)
    height = table.rows[0]._tr.find("w:trPr/w:trHeight", table.rows[0]._tr.nsmap)
    height.set(qn("w:hRule"), "exact")
    cell = table.cell(0, 0)
    _attach_library(doc, BidBrief(), {"social": images(tmp_path, "social")})
    assert len(drawings(cell._tc)) == 3
    assert height.get(qn("w:hRule")) == "atLeast"
    assert all(int(e.get("cx")) <= int(Cm(5)) for e in cell._tc.iter(qn("wp:extent")))
    assert not list(cell._tc.iter(qn("wp:anchor")))


def test_no_holder_puts_images_after_body_and_table_not_after_title(tmp_path):
    doc = Document()
    heading = doc.add_paragraph("一、营业执照", style="Heading 1")
    body = doc.add_paragraph("这是原版正文。")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "必须保留的原表"
    end = doc.add_paragraph("二、下一章", style="Heading 1")
    _attach_library(doc, BidBrief(), {"license": images(tmp_path, "license", 1)})
    assert not drawings(heading._p) and not drawings(body._p) and not drawings(table._tbl)
    assert drawings(table._tbl.getnext())
    assert table._tbl.getnext().getnext() is end._p


def test_pdf_all_pages_and_missing_material_holder_survive(tmp_path):
    import fitz

    path = tmp_path / "financial.pdf"
    pdf = fitz.open()
    for _ in range(5):
        pdf.new_page()
    pdf.save(path)
    pdf.close()
    doc = Document()
    holder = doc.add_paragraph("此处粘贴财务报表扫描件")
    missing = doc.add_paragraph("此处粘贴营业执照扫描件")
    notes = _attach_library(doc, BidBrief(), {"finance": [path]})
    assert len(drawings(doc.element.body)) == 5
    assert not drawings(holder._p) and not drawings(missing._p)
    assert missing.text == "此处粘贴营业执照扫描件"
    assert any("缺营业执照" in n for n in notes)


def test_untyped_frame_does_not_get_arbitrary_library_file(tmp_path):
    doc = Document()
    holder = doc.add_paragraph()
    holder._p.get_or_add_pPr().append(OxmlElement("w:pBdr"))
    before = doc.element.xml
    notes, _ = place_library_scans(doc, media={"license": images(tmp_path, "license")})
    assert doc.element.xml == before
    assert any("未找到明确占位" in n for n in notes)


def test_scan_stays_before_original_section_boundary(tmp_path):
    doc = Document()
    doc.add_paragraph("一、营业执照", style="Heading 1")
    body = doc.add_paragraph("原版正文。")
    section = OxmlElement("w:sectPr")
    body._p.get_or_add_pPr().append(section)
    end = doc.add_paragraph("二、下一章", style="Heading 1")
    _attach_library(doc, BidBrief(), {"license": images(tmp_path, "license", 1)})
    picture = body._p.getnext()
    assert drawings(picture)
    assert picture.getnext().find("w:pPr/w:sectPr", picture.nsmap) is section
    assert picture.getnext().getnext() is end._p


def test_tax_and_audit_files_in_shared_finance_slot_are_not_interchanged(tmp_path):
    tax = tmp_path / "纳税证明.png"
    audit = tmp_path / "财务审计报告.png"
    Image.new("RGB", (100, 100), "red").save(tax)
    Image.new("RGB", (100, 100), "blue").save(audit)
    doc = Document()
    holders = [
        doc.add_paragraph("此处粘贴纳税证明扫描件"),
        doc.add_paragraph("此处粘贴财务报表扫描件"),
    ]
    _attach_library(doc, BidBrief(), {"finance": [tax, audit]})
    assert len(drawings(doc.element.body)) == 2
    for holder, channel in zip(holders, (0, 2), strict=True):
        rid = next(holder._p.getnext().iter(qn("a:blip"))).get(qn("r:embed"))
        with Image.open(BytesIO(doc.part.related_parts[rid].blob)) as image:
            assert image.getpixel((10, 10))[channel] > 240


def test_each_identity_face_uses_its_named_placeholder(tmp_path):
    front = tmp_path / "front.png"
    back = tmp_path / "back.png"
    Image.new("RGB", (100, 60), "red").save(front)
    Image.new("RGB", (100, 60), "green").save(back)
    doc = Document()
    doc.add_paragraph("附：法定代表人身份证正面复印件")
    doc.add_paragraph("附：法定代表人身份证反面复印件")
    _attach_library(doc, BidBrief(), {"id_legal": [front, back]})
    assert len(drawings(doc.element.body)) == 2


@pytest.mark.parametrize("label", ["正反面", "正、反面"])
def test_combined_identity_holder_keeps_both_faces(tmp_path, label):
    front = tmp_path / "front.png"
    back = tmp_path / "back.png"
    Image.new("RGB", (100, 60), "red").save(front)
    Image.new("RGB", (100, 60), "green").save(back)
    doc = Document()
    doc.add_paragraph("附：法定代表人身份证" + label + "复印件")
    _attach_library(doc, BidBrief(), {"id_legal": [front, back]})
    assert len(drawings(doc.element.body)) == 1


def test_untyped_identity_holder_in_authorization_uses_agent_not_legal_person(tmp_path):
    legal = tmp_path / "legal.png"
    agent = tmp_path / "agent.png"
    Image.new("RGB", (100, 60), "red").save(legal)
    Image.new("RGB", (100, 60), "blue").save(agent)
    doc = Document()
    doc.add_paragraph("法定代表人授权委托书")
    holder = doc.add_paragraph("此处粘贴身份证正反面")
    _attach_library(doc, BidBrief(), {"id_legal": [legal], "id_agent": [agent]})
    rid = next(holder._p.getnext().iter(qn("a:blip"))).get(qn("r:embed"))
    with Image.open(BytesIO(doc.part.related_parts[rid].blob)) as image:
        assert image.getpixel((10, 10))[2] > 240
    assert len(drawings(doc.element.body)) == 1


@pytest.mark.parametrize("volume", ["business", "technical", "full"])
def test_native_volume_assembly_uses_the_same_scan_anchors(tmp_path, monkeypatch, volume):
    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.schema import OutlineItem

    path = tmp_path / "native.format.docx"
    source = Document()
    source.add_paragraph("投 标 文 件")
    source.add_paragraph("一、营业执照", style="Heading 1")
    source.add_paragraph("此处粘贴营业执照扫描件")
    source.add_paragraph("二、技术证明材料", style="Heading 1")
    source.add_paragraph("此处粘贴产品检测报告扫描件")
    source.save(path)
    path.with_name(path.name + ".source").write_text("docx-raw", encoding="utf-8")
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: path)
    brief = BidBrief(
        layoutMode="outline",
        outlineChapter="第八章 投标文件格式",
        outlineItems=[
            OutlineItem(id="lic", title="营业执照", rawTitle="一、营业执照", kind="scan"),
            OutlineItem(id="tech", title="技术证明材料", rawTitle="二、技术证明材料", kind="scan"),
        ],
    )
    dest, _ = assemble_bid_docx(
        brief,
        tmp_path / (volume + ".docx"),
        volume=volume,
        catalog_slots=[
            PlaceholderItem(key="license", title="营业执照"),
            PlaceholderItem(key="product", title="产品检测报告"),
        ],
        catalog_media={
            "license": images(tmp_path, "license", 1),
            "product": images(tmp_path, "product", 1),
        },
    )
    doc = Document(dest)
    assert len(drawings(doc.element.body)) == (2 if volume == "full" else 1)
    assert all(not drawings(p._p) for p in doc.paragraphs if p.text.startswith(("一、", "二、")))
