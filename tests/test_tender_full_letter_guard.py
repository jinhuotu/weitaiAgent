import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders.assemble import (
    _fill_business_format_sections,
    _fill_drafted_format_sections,
)
from api.services.tenders.business_render import finalize_business_responses
from api.services.tenders.full_bid_sections import letter_region
from api.services.tenders.generate import generate_bid
from api.services.tenders.native_format import validate_native_toc
from api.services.tenders.schema import BidBrief, DeviationLine, OutlineItem

PENDING = "【待补充】「施工组织设计」技术方案正文未完整生成，请结合本项目招标要求补充。"
BODY = "本项目已确认的施工组织设计正文。"


def source(tmp_path, *, automatic=False, real_section=True):
    doc = Document()
    doc.add_paragraph("投标文件")
    doc.add_paragraph("目录")
    doc.add_paragraph("一、投标函及投标函附录").add_run().add_tab()
    if real_section:
        doc.add_paragraph("二、施工组织设计").add_run().add_tab()
    doc.add_page_break()
    doc.add_paragraph("一、投标函及投标函附录", style="Heading 1")
    doc.add_paragraph("（一）投标函", style="Heading 2")
    doc.add_paragraph("我方的投标文件包括下列内容：")
    composition = []
    for n in range(1, 10):
        title = "施工组织设计；" if n == 6 else f"原版组成清单第{n}项；"
        p = doc.add_paragraph(title if automatic else f"（{n}）{title}", style="Heading 2")
        if automatic:
            num = OxmlElement("w:numPr")
            for name, value in (("ilvl", "2"), ("numId", "1")):
                node = OxmlElement("w:" + name)
                node.set(qn("w:val"), value)
                num.append(node)
            p._p.get_or_add_pPr().append(num)
        composition.append(p._p)
    doc.add_paragraph("我方承诺在投标有效期内不修改或撤销投标文件。")
    doc.add_paragraph("投标人：（盖章）")
    doc.add_paragraph("年 月 日")
    heading = None
    if real_section:
        heading = doc.add_paragraph("二、施工组织设计", style="Heading 1")
        doc.add_paragraph("原技术章节内容。")
    path = tmp_path / "source.format.docx"
    doc.save(path)
    path.with_name(path.name + ".source").write_text("docx-raw", encoding="utf-8")
    return doc, path, composition, heading


def brief(*, automatic=False, body=PENDING):
    return BidBrief(
        invitationId="full-audit", layoutMode="outline", outlineChapter="第八章 投标文件格式",
        generateVolume="full", autoDraftTechPlan=True, attachQualifications=False,
        outlineItems=[
            OutlineItem(id="letter", title="投标函及投标函附录", kind="letter"),
            OutlineItem(id="plan", title="施工组织设计", kind="tech_plan",
                        rawTitle="施工组织设计；" if automatic else "（6）施工组织设计；",
                        draftBody=body, draftStatus="drafted"),
        ],
    )


@pytest.mark.parametrize("automatic", [False, True])
def test_full_draft_skips_letter_list_and_fills_real_technical_section(tmp_path, automatic):
    original, path, composition, heading = source(tmp_path, automatic=automatic)
    data = brief(automatic=automatic)
    protected = letter_region(original, data)
    assert all(p in protected for p in composition)
    assert heading._p not in protected
    assert _fill_drafted_format_sections(path, data, full_bid=True) == (1, [])
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    assert texts.index(PENDING) > texts.index("二、施工组织设计")
    assert not any(PENDING in "".join(t.text or "" for t in p.iter(qn("w:t")))
                   for p in letter_region(doc, data) if p.tag == qn("w:p"))
    before = path.read_bytes()
    assert _fill_drafted_format_sections(path, data, full_bid=True) == (0, [])
    assert path.read_bytes() == before


@pytest.mark.parametrize("pending_first", [False, True])
def test_full_prefers_real_body_over_pending_alias_for_same_heading(tmp_path, pending_first):
    _, path, _, _ = source(tmp_path)
    data = brief()
    real = OutlineItem(id="real", title="施工组织格式", rawTitle="二、施工组织设计",
                       kind="tech_chapter", draftBody=BODY)
    if pending_first:
        data.outlineItems.append(real)
    else:
        data.outlineItems.insert(1, real)
    inserted, notes = _fill_drafted_format_sections(path, data, full_bid=True)
    assert inserted == 1
    texts = [p.text for p in Document(path).paragraphs]
    assert texts.count(BODY) == 1 and PENDING not in texts
    assert any("未追加重复" in note for note in notes)


def test_full_does_not_use_letter_list_when_real_section_is_missing(tmp_path):
    _, path, _, _ = source(tmp_path, real_section=False)
    before = path.read_bytes()
    inserted, notes = _fill_drafted_format_sections(path, brief(), full_bid=True)
    assert inserted == 0 and notes
    assert path.read_bytes() == before


def test_full_markdown_table_draft_is_inserted_only_once(tmp_path):
    _, path, _, _ = source(tmp_path)
    data = brief(body="技术响应说明\n| 项目 | 响应 |\n|---|---|\n| 质量 | 按要求执行 |")
    assert _fill_drafted_format_sections(path, data, full_bid=True)[0] == 1
    before = path.read_bytes()
    assert _fill_drafted_format_sections(path, data, full_bid=True)[0] == 0
    assert path.read_bytes() == before


@pytest.mark.parametrize("automatic", [False, True])
def test_complete_full_generation_keeps_letter_composition_intact(tmp_path, monkeypatch, automatic):
    _, path, _, _ = source(tmp_path, automatic=automatic)
    original = path.read_bytes()
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: path)
    monkeypatch.setattr("api.services.tenders.generate.tenders_output_dir", lambda: tmp_path)
    data = brief(automatic=automatic)
    result = generate_bid(data, catalog_slots=[], catalog_media={})
    doc = Document(tmp_path / result["docxFile"])
    protected = letter_region(doc, data)
    protected_text = "\n".join(
        p.text for p in doc.paragraphs if p._p in protected
    )
    assert "【待补充】" not in protected_text
    assert all(f"原版组成清单第{n}项" in protected_text for n in range(1, 10) if n != 6)
    assert "施工组织设计；" in protected_text
    assert PENDING in [p.text for p in doc.paragraphs]
    assert validate_native_toc(doc)["entries"] == 2
    assert path.read_bytes() == original


def test_full_business_response_does_not_route_to_letter_composition(tmp_path):
    doc = Document()
    doc.add_paragraph("投标函", style="Heading 1")
    doc.add_paragraph("投标文件组成如下：")
    doc.add_paragraph("（7）营业执照", style="Heading 2")
    doc.add_paragraph("投标人：（盖章）")
    doc.add_paragraph("一、其他资料", style="Heading 1")
    data = BidBrief(autoDraftBizResp=True, outlineItems=[OutlineItem(
        title="营业执照", rawTitle="（7）营业执照", kind="scan",
    )], bizDeviationLines=[DeviationLine(
        requirement="须提供营业执照", response="已提供有效营业执照扫描件", deviation="无偏离",
    )])
    notes = finalize_business_responses(doc, data, full_bid=True)
    assert any("文末核对表" in note for note in notes)
    assert "已提供有效营业执照扫描件" not in [p.text for p in doc.paragraphs]


def test_full_business_draft_uses_actual_response_chapter(tmp_path):
    doc = Document()
    doc.add_paragraph("投标函", style="Heading 1")
    doc.add_paragraph("投标文件组成如下：")
    doc.add_paragraph("（6）商务响应承诺书", style="Heading 2")
    doc.add_paragraph("投标人：（盖章）")
    doc.add_paragraph("一、商务响应承诺书", style="Heading 1")
    path = tmp_path / "response.docx"
    doc.save(path)
    data = BidBrief(outlineItems=[OutlineItem(
        title="商务响应承诺书", rawTitle="（6）商务响应承诺书", kind="commitment_copy",
        draftBody="本项目按照合同约定组织履约。", draftStatus="drafted",
    )])
    inserted, _ = _fill_business_format_sections(path, data, full_bid=True)
    assert inserted == 1
    texts = [p.text for p in Document(path).paragraphs]
    assert texts.index(data.outlineItems[0].draftBody) > texts.index("一、商务响应承诺书")
