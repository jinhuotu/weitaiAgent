"""知识库解析单测（不连云 OCR / MySQL / Qdrant）。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from api.services.knowledge.parsers import (
    assert_supported,
    extract_text_from_file,
    sniff_extension,
)
from common.errors import AppError


def test_sniff_and_image_ext() -> None:
    assert sniff_extension("手册.PDF") == "pdf"
    assert sniff_extension("a.PNG") == "png"
    assert sniff_extension("汇报.PPTX") == "pptx"
    assert sniff_extension(None, "image/jpeg") == "jpg"
    assert sniff_extension(
        None, "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ) == "pptx"
    assert assert_supported("webp") == "webp"
    assert assert_supported("pptx") == "pptx"


@pytest.mark.asyncio
async def test_extract_txt(tmp_path: Path) -> None:
    p = tmp_path / "note.txt"
    p.write_text("冷却水温度过高时停机。", encoding="utf-8")
    result = await extract_text_from_file(p)
    assert "冷却水" in result.text


@pytest.mark.asyncio
async def test_extract_docx_xlsx(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("禁止带电作业。")
    path = tmp_path / "s.docx"
    doc.save(path)
    text = (await extract_text_from_file(path)).text
    assert "禁止带电作业" in text

    wb = Workbook()
    ws = wb.active
    ws.title = "冷却"
    ws["A1"] = "水温"
    ws["B1"] = 42
    xlsx = tmp_path / "t.xlsx"
    wb.save(xlsx)
    sheet = await extract_text_from_file(xlsx)
    assert "## 冷却" in sheet.text
    assert "水温" in sheet.text
    assert sheet.formula_fallback is False


@pytest.mark.asyncio
async def test_xlsx_formula_fallback(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "公式"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=A1+A2"
    xlsx = tmp_path / "f.xlsx"
    wb.save(xlsx)
    result = await extract_text_from_file(xlsx)
    assert result.formula_fallback is True
    assert "=A1+A2" in result.text
    assert "未经 Excel 计算" in result.text


@pytest.mark.asyncio
async def test_extract_pptx(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "窑炉升温曲线"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "注意保温段时长"
    path = tmp_path / "brief.pptx"
    prs.save(path)
    result = await extract_text_from_file(path)
    assert result.page_count == 1
    assert "窑炉升温曲线" in result.text
    assert "保温段" in result.text


@pytest.mark.asyncio
async def test_image_without_ocr_fails(tmp_path: Path) -> None:
    from PIL import Image

    img = tmp_path / "a.png"
    Image.new("RGB", (32, 32), color=(20, 20, 20)).save(img)
    with pytest.raises(AppError) as ei:
        await extract_text_from_file(img)
    assert "OCR" in ei.value.msg or "正文过短" in ei.value.msg


def test_doc_rejected() -> None:
    with pytest.raises(AppError) as ei:
        assert_supported("doc")
    assert "docx" in ei.value.msg


def test_collapse_cjk_spaces() -> None:
    from api.services.knowledge.parsers import collapse_cjk_spaces

    assert collapse_cjk_spaces("投 标 文 件") == "投标文件"
    assert "投标文件" in collapse_cjk_spaces("关于 投 标 文 件 的说明")


@pytest.mark.asyncio
async def test_extract_json_yaml_xml(tmp_path: Path) -> None:
    jp = tmp_path / "a.json"
    jp.write_text('{"title": "投标文件", "n": 1}', encoding="utf-8")
    jtxt = (await extract_text_from_file(jp)).text
    assert "投标文件" in jtxt

    yp = tmp_path / "a.yaml"
    yp.write_text("title: 冷却水\nlimit: 42\n", encoding="utf-8")
    ytxt = (await extract_text_from_file(yp)).text
    assert "冷却水" in ytxt

    xp = tmp_path / "a.xml"
    xp.write_text("<root><item>窑炉升温</item></root>", encoding="utf-8")
    xtxt = (await extract_text_from_file(xp)).text
    assert "窑炉升温" in xtxt


def test_rapidocr_result_parse() -> None:
    from api.services.knowledge.ocr.rapidocr import _texts_from_result

    assert _texts_from_result([[(0, 0), "冷却水", 0.9]]) == ["冷却水"]
    assert _texts_from_result(([[(0, 0), "停机", 0.8]], 0.1)) == ["停机"]

    class Box:
        txts = ["投标", "文件"]

    assert _texts_from_result(Box()) == ["投标", "文件"]


@pytest.mark.asyncio
async def test_extract_xls(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    xlrd = pytest.importorskip("xlrd")

    class Sheet:
        name = "冷却"
        nrows = 1
        ncols = 2

        def cell_value(self, r: int, c: int) -> object:
            return [["水温", 42]][r][c]

    class Book:
        def sheets(self) -> list[Sheet]:
            return [Sheet()]

    monkeypatch.setattr(xlrd, "open_workbook", lambda _p: Book())
    path = tmp_path / "t.xls"
    path.write_bytes(b"fake-xls")
    result = await extract_text_from_file(path)
    assert "## 冷却" in result.text
    assert "水温" in result.text
    assert result.page_count == 1

