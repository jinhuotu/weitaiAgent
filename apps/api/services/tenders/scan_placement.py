"""Place selected library scans at typed placeholders, using inline blank paragraphs."""

from __future__ import annotations

import io
import re
from hashlib import sha1
from pathlib import Path

from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ALIASES = {
    "id_legal": ("法定代表人身份证", "法人身份证", "法定代表人身份证明"),
    "id_agent": ("委托代理人身份证", "代理人身份证", "授权代理人身份证", "授权委托书"),
    "license": ("营业执照",),
    "finance": ("纳税", "完税", "税收", "财务", "审计报告"),
    "social": ("社保", "社会保障"),
    "perf": ("业绩", "类似项目合同", "类似项目情况", "合同及发票"),
    "bank_permit": ("开户许可", "基本账户", "基本存款账户"),
    "iso": ("ISO", "质量管理体系", "体系认证"),
    "product": ("检测报告", "3C", "产品认证", "产品证明"),
    "bond": ("投标保证金", "保证金缴存", "保函"),
    "credit": ("信用中国", "企业信用信息", "信用查询"),
    "seal": ("签章页", "印章扫描件"),
    "qualification": ("企业资质", "资质证书", "安全生产许可"),
}
_CUE = re.compile(r"复印件|扫描件|粘贴处|[此这]处粘贴|在此粘贴|附[：:]")


def _text(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t")))


def _compact(text):
    return re.sub(r"\s+", "", text or "").upper()


def _prefix(key):
    return "scan_slot_" + sha1(key.encode()).hexdigest()[:12] + "_"


def _matching_files(key, anchor, files):
    """Honor explicit subtype labels when filenames distinguish a shared slot."""
    text = _text(anchor)
    if not text.strip() and anchor.getprevious() is not None:
        text = _text(anchor.getprevious())
    if key == "finance":
        tax = re.compile(r"纳税|完税|税收")
        audit = re.compile(r"财务|审计|报表")
        desired = (
            tax
            if tax.search(text) and not audit.search(text)
            else (audit if audit.search(text) and not tax.search(text) else None)
        )
        if desired is not None and any(tax.search(p.name) or audit.search(p.name) for p in files):
            return [p for p in files if desired.search(p.name)]
    if key in {"id_legal", "id_agent"}:
        if re.search(r"正.{0,3}反.{0,2}面", text):
            return files
        front = "正面" in text and "反面" not in text
        back = "反面" in text and "正面" not in text
        if front or back:
            kind = re.compile("正面|FRONT" if front else "反面|BACK", re.I)
            if any(re.search("正面|反面|FRONT|BACK", p.name, re.I) for p in files):
                return [p for p in files if kind.search(p.name)]
    return files


def _frame(el):
    return el.find("w:pPr/w:pBdr", el.nsmap) is not None


def _cell(el):
    return next((n for n in el.iterancestors(qn("w:tc"))), None)


def _cell_frame(el, *, explicit=False):
    cell = _cell(el)
    if cell is None:
        return False
    table = next(cell.iterancestors(qn("w:tbl")), None)
    return cell.find("w:tcPr/w:tcBorders", cell.nsmap) is not None or (
        table is not None
        and (
            table.find("w:tblPr/w:tblBorders", table.nsmap) is not None
            or (not explicit and len(list(table.iter(qn("w:tc")))) == 1)
        )
    )


def _heading(el, doc):
    from api.services.tenders.native_format import _physical_chapter

    text = _text(el)
    return not _CUE.search(text) and _physical_chapter(el, doc)


def _width(doc, anchor):
    sec = doc.sections[-1]
    width = (sec.page_width - sec.left_margin - sec.right_margin) / 360000 - 0.4
    cell = _cell(anchor)
    if cell is not None:
        w = cell.find("w:tcPr/w:tcW", cell.nsmap)
        if w is not None and w.get(qn("w:type")) == "dxa":
            width = min(width, int(w.get(qn("w:w"), "0")) * 2.54 / 1440 - 0.4)
    return max(0.8, width)


def _pictures(doc, anchor, files, key=""):
    """Create inline paragraphs in this document (retaining its image relations)."""
    from PIL import Image, ImageOps

    from api.services.tenders.placeholders import (
        _compose_id_strip,
        _fit_cm,
        _normalize_id_image,
        _render_pdf_page_jpeg,
        _scan_height_cm,
    )

    sec = doc.sections[-1]
    max_h = min(
        _scan_height_cm(doc), (sec.page_height - sec.top_margin - sec.bottom_margin) / 360000 - 1
    )
    paragraphs = []
    errors = []
    cards = []

    def picture(stream, size):
        if key in {"id_legal", "id_agent"}:
            with Image.open(stream) as raw:
                card = _normalize_id_image(raw, force_landscape=True)
                buffer = io.BytesIO()
                card.save(buffer, format="JPEG", quality=90)
                buffer.seek(0)
                cards.append(buffer)
            return
        insert(stream, size, _width(doc, anchor), max_h)

    def insert(stream, size, width, height):
        w, h = _fit_cm(*size, max_w=width, max_h=height)
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.keep_together = True
        p.alignment = 1
        try:
            p.add_run().add_picture(stream, width=Cm(w), height=Cm(h))
        except Exception:
            p._p.getparent().remove(p._p)
            raise
        paragraphs.append(p._p)

    for path in files:
        try:
            if path.suffix.lower() == ".pdf":
                import pypdfium2 as pdfium

                pdf = pdfium.PdfDocument(str(path))
                try:
                    for n in range(len(pdf)):
                        page = pdf[n]
                        try:
                            stream = _render_pdf_page_jpeg(page)
                            with Image.open(stream) as img:
                                size = img.size
                            stream.seek(0)
                            picture(stream, size)
                        finally:
                            page.close()
                finally:
                    pdf.close()
            else:
                with Image.open(path) as raw:
                    image = ImageOps.exif_transpose(raw).convert("RGB")
                    stream = io.BytesIO()
                    image.save(stream, format="JPEG", quality=90)
                    stream.seek(0)
                    picture(stream, image.size)
        except Exception:  # A bad scan must not erase its placeholder or other valid pages.
            errors.append(path.name)
    for n in range(0, len(cards), 2):
        strip = _compose_id_strip(cards[n : n + 2])
        if strip is not None:
            stream, width, height = strip
            with Image.open(stream) as image:
                size = image.size
            stream.seek(0)
            insert(stream, size, min(_width(doc, anchor), width), min(5.4, height))
    return paragraphs, errors


def _defer_boundary(doc, anchor, cursor):
    """Keep attachments before an end-of-chapter page/section boundary."""
    if anchor.tag != qn("w:p"):
        return
    section = anchor.find("w:pPr/w:sectPr", anchor.nsmap)
    nodes = list(anchor.iter())
    breaks = [
        node
        for n, node in enumerate(nodes)
        if node.tag == qn("w:br")
        and node.get(qn("w:type")) == "page"
        and not any(
            other.tag == qn("w:t") and (other.text or "").strip() for other in nodes[n + 1 :]
        )
    ]
    if section is None and not breaks:
        return
    carrier = doc.add_paragraph(style="Normal")
    if section is not None:
        section.getparent().remove(section)
        carrier._p.get_or_add_pPr().append(section)
    for node in breaks:
        node.getparent().remove(node)
        carrier.add_run()._r.append(node)
    cursor.addnext(carrier._p)


def place_library_scans(doc, brief=None, media=None, slots=None, skip_keys=None, global_pasted: set[str] | None = None):
    from api.services.tenders.document import _bookmark_paragraph
    from api.services.tenders.material_identity import (
        file_identity,
        material_key_aliases,
        unique_material_files,
    )
    from api.services.tenders.native_format import _body_toc_block
    from api.services.tenders.pdf_format import toc_block_ids
    from api.services.tenders.placeholders import DEFAULT_SLOTS

    catalog = {s.key: s for s in DEFAULT_SLOTS}
    catalog.update({s.key: s for s in slots or []})
    keys = set(catalog) | set(media or {})
    keys.discard("tech_drawings")
    key_aliases = material_key_aliases(catalog.values(), preferred_keys=(s.key for s in DEFAULT_SLOTS))
    groups = {}
    for key in sorted(keys):
        groups.setdefault(key_aliases.get(key, key), []).append(key)
    key_aliases.update((key, key) for key in keys if key not in key_aliases)
    names = {}
    for key in keys:
        aliases = list(ALIASES.get(key, ()))
        if key in catalog:
            # Catalog titles describe that key; hints are never matching evidence.
            aliases += [
                part.strip()
                for part in re.split(r"[/／]|及", catalog[key].title)
                if len(part.strip()) >= 3
            ]
        names[key] = [_compact(a) for a in aliases]
    block = set(_body_toc_block(doc, brief) if brief is not None else toc_block_ids(doc))
    paras = [
        p
        for p in doc.element.body.iter(qn("w:p"))
        if p not in block and not any(n in block for n in p.iterancestors())
    ]
    sections = []
    owner = {}
    current = None
    material_titles = {title for aliases in names.values() for title in aliases}
    material_titles.update(
        _compact(slot.attachToSection) for slot in catalog.values() if slot.attachToSection
    )
    material_titles.update({"法定代表人授权委托书", "法定代表人身份证明书"})
    for p in paras:
        if p.getparent() is doc.element.body and (
            _heading(p, doc) or _compact(_text(p)) in material_titles
        ):
            current = p
            sections.append(p)
        owner[p] = current
    skipped = set(skip_keys or ())
    # 合并全局已贴集合，避免重复贴图
    if global_pasted is not None:
        skipped |= global_pasted
    skipped = {key_aliases.get(key, key) for key in skipped}
    existing = {n.get(qn("w:name"), "") for n in doc.element.iter(qn("w:bookmarkStart"))}
    handled = {
        key for key, members in groups.items()
        if any(n.startswith(_prefix(member)) for member in members for n in existing)
    }
    skipped |= handled
    notes = []

    def matches(text, section=False, context=""):
        normalized = _compact(text)
        result = {key for key, aliases in names.items() if any(a in normalized for a in aliases)}
        result = {key_aliases.get(key, key) for key in result}
        # A social-insurance placeholder is never an identity-card placeholder.
        if "身份证" not in normalized:
            result -= {"id_legal", "id_agent"}
            if section and "授权委托书" in normalized:
                result.add("id_agent")
        elif not any(w in normalized for w in ("代理", "委托")):
            result.discard("id_agent")
            if "身份证" in normalized:
                if not any(w in normalized for w in ("法定", "法人")) and any(
                    w in context for w in ("代理", "委托")
                ):
                    result.discard("id_legal")
                    result.add("id_agent")
                else:
                    result.add("id_legal")
        else:
            result.discard("id_legal")
        return result

    targets = {}
    for p in paras:
        text = _text(p)
        if not (
            _CUE.search(text)
            or (_frame(p) and not text.strip())
            or (_cell_frame(p) and not text.strip())
        ):
            continue
        if next(p.iter(qn("w:drawing")), None) is not None:
            continue
        section_context = _text(owner[p]) if owner[p] is not None else ""
        found = matches(text, context=section_context)
        if not found:
            cell = _cell(p)
            context = _text(cell) if cell is not None else ""
            found = matches(context, context=section_context)
        if not found and not text.strip():
            prev = p.getprevious()
            while prev is not None and not _text(prev).strip():
                prev = prev.getprevious()
            if prev is not None and _CUE.search(_text(prev)):
                found = matches(_text(prev), context=section_context)
        if not found and owner[p] is not None:
            found = matches(_text(owner[p]), section=True)
        # An untyped frame in a broad multi-material chapter is ambiguous.
        if not _text(p).strip() and len(found) != 1:
            continue
        for key in found - skipped:
            if not text.strip() and (_frame(p) or _cell_frame(p)):
                host = p
                while host.getparent() is not doc.element.body and host.getparent() is not None:
                    host = host.getparent()
                previous = host.getprevious()
                while previous is not None and not _text(previous).strip():
                    previous = previous.getprevious()
                if previous is not None and _CUE.search(_text(previous)):
                    targets[key] = [old for old in targets.get(key, []) if old is not previous]
            if _cell(p) is not None and any(_cell(old) is _cell(p) for old in targets.get(key, [])):
                continue
            targets.setdefault(key, []).append(p)

    hashes = {}
    placed_files = {}
    for key in sorted(set(groups) - skipped):
        files = []
        members = [key, *(member for member in groups[key] if member != key)]
        for value in (value for member in members for value in (media or {}).get(member, [])):
            path = Path(value)
            if path.is_file() and path.suffix.lower() in {
                ".png",
                ".jpg",
                ".jpeg",
                ".bmp",
                ".webp",
                ".gif",
                ".pdf",
            }:
                files.append(path)
        if (
            not files
            and key in catalog
            and any(
                other != key
                and other in catalog
                and values
                and _compact(catalog[other].title) == _compact(catalog[key].title)
                for other, values in (media or {}).items()
            )
        ):
            handled.add(key)
            continue
        anchors = targets.get(key, [])
        fallback = False
        if not anchors and files:
            candidates = [p for p in sections if key in matches(_text(p), section=True)]
            section_names = [
                catalog[member].attachToSection for member in members
                if member in catalog and catalog[member].attachToSection
            ]
            if not candidates and section_names:
                candidates = [
                    p
                    for p in sections
                    if any(_compact(title) in _compact(_text(p)) for title in section_names)
                ]
            if candidates:
                heading = candidates[0]
                at = paras.index(heading)
                end = next(
                    (n for n in range(at + 1, len(paras)) if paras[n] in sections), len(paras)
                )
                region = paras[at:end]
                # Stay before a trailing page/section-break carrier.
                anchor = next(
                    (p for p in reversed(region) if _text(p).strip() or _cell(p) is not None),
                    heading,
                )
                while anchor.getparent() is not doc.element.body:
                    anchor = anchor.getparent()
                anchors = [anchor]
                fallback = True
        label = catalog[key].title if key in catalog else key
        if key == "finance" and any(
            w in _text(p) for p in anchors for w in ("纳税", "完税", "税收")
        ):
            label = "纳税证明"
        elif key == "social":
            label = "社保证明"
        if not anchors:
            if files:
                handled.add(key)
                notes.append(f"「{label}」未找到明确占位或对应章节，未贴入，请指定落位")
            continue
        if not files:
            if any(
                _frame(p)
                or _cell_frame(p, explicit=True)
                or (not _text(p).strip() and _cell_frame(p))
                for p in anchors
            ):
                handled.add(key)
            notes.append(f"缺{label}，已保留原占位")
            if key in {"id_legal", "id_agent"} and brief is not None:
                who = getattr(brief, "legalPersonName" if key == "id_legal" else "agentName", "")
                notes.append(f"未贴上{who or label}的身份证扫描件，请在资料库对应项上传")
            continue
        handled.add(key)
        for index, anchor in enumerate(anchors):
            selected = _matching_files(key, anchor, files)
            if not selected:
                notes.append(f"「{label}」文件类型与该占位不符，已保留占位，请核对资料库分类")
                continue
            seen = placed_files.setdefault(anchor, set())
            selected = [
                path for path in unique_material_files(selected, cache=hashes)
                if file_identity(path, hashes) not in seen
            ]
            if not selected:
                continue
            pictures, errors = _pictures(doc, anchor, selected, key)
            if not pictures:
                notes.append(f"「{label}」扫描件无法读取，已保留原占位")
                if fallback:
                    from api.services.tenders.placeholders import _drop_after
                    from api.services.tenders.schema import PlaceholderItem

                    _drop_after(
                        doc, anchor, PlaceholderItem(key=key, title=label), [], complete=True
                    )
                continue
            seen.update(file_identity(path, hashes) for path in selected if path.name not in errors)
            cursor = anchor
            marker_host = pictures[0]
            if fallback and anchor.tag == qn("w:tbl"):
                from api.services.tenders.placeholders import _floating_table_cm

                clearance = _floating_table_cm(anchor)
                if clearance:
                    spacer = doc.add_paragraph(style="Normal")
                    spacer.paragraph_format.line_spacing = Cm(clearance)
                    spacer.add_run().font.size = Pt(1)
                    anchor.addnext(spacer._p)
                    cursor = spacer._p
            if not fallback and not _text(anchor).strip() and _frame(anchor):
                first = pictures.pop(0)
                for run in list(first.findall(qn("w:r"))):
                    anchor.append(run)
                first.getparent().remove(first)
                para = Paragraph(anchor, doc)
                para.style = "Normal"
                para.paragraph_format.line_spacing = 1
                para.paragraph_format.page_break_before = False
                marker_host = anchor
                for node in list(anchor.findall("w:pPr/w:framePr", anchor.nsmap)):
                    node.getparent().remove(node)
            for p in pictures:
                cursor.addnext(p)
                cursor = p
            _defer_boundary(doc, anchor, cursor)
            cell = _cell(anchor)
            if cell is not None:
                table = next(cell.iterancestors(qn("w:tbl")), None)
                if table is not None:
                    for pos in list(table.findall("w:tblPr/w:tblpPr", table.nsmap)):
                        pos.getparent().remove(pos)
                for row in cell.iterancestors(qn("w:tr")):
                    for height in row.findall("w:trPr/w:trHeight", row.nsmap):
                        height.set(qn("w:hRule"), "atLeast")
                    for no_split in list(row.findall("w:trPr/w:cantSplit", row.nsmap)):
                        no_split.getparent().remove(no_split)
            _bookmark_paragraph(Paragraph(marker_host, doc), _prefix(key) + str(index))
            if not errors:
                for node in anchor.iter(qn("w:t")):
                    node.text = (
                        (node.text or "")
                        .replace("（在此粘贴扫描件）", "")
                        .replace("在此粘贴扫描件", "")
                    )
            notes.append(f"已贴上{label}扫描件")
            notes.append(f"「{label}」已在对应占位贴上资料")
            if errors:
                notes.append(f"「{label}」部分扫描件无法读取，待补：{'、'.join(errors)}")
    handled = {member for key in handled for member in groups.get(key, [key])}
    # 更新全局已贴集合
    if global_pasted is not None:
        global_pasted.update(handled)
    return notes, handled
