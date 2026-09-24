"""投标文件填空与人民币大写。"""

from __future__ import annotations

import hashlib
from decimal import Decimal

from api.services.tenders.assets import find_chapter5_template
from api.services.tenders.document import build_bid_docx
from api.services.tenders.money import rmb_lowercase, rmb_uppercase
from api.services.tenders.schema import BidBrief, DeviationLine, QuoteLineIn


def _sample_quote_lines() -> list[QuoteLineIn]:
    return [
        QuoteLineIn(
            seq="1",
            name="交流充电桩",
            spec="7kW",
            unit="台",
            qty=10,
            unitPrice=1000,
            amount=10000,
        ),
        QuoteLineIn(
            seq="2",
            name="立柱",
            spec="配套",
            unit="个",
            qty=10,
            unitPrice=100,
            amount=1000,
        ),
    ]


def test_default_brief_project_fields_empty() -> None:
    from api.services.tenders.schema import default_brief

    brief = default_brief()
    assert brief.projectName == ""
    assert brief.tenderer == ""
    assert brief.bidContent == ""
    assert brief.bidderName == "河南伟泰光电科技有限公司"


def test_chapter5_template_exists() -> None:
    path = find_chapter5_template()
    assert path is not None
    assert path.is_file()
    assert path.stat().st_size > 10000


def test_custom_layout_template_overrides_bundled(tmp_path, monkeypatch) -> None:
    from api.services.tenders import assets as assets_mod

    monkeypatch.setattr(assets_mod, "tender_assets_dir", lambda: tmp_path)
    custom = tmp_path / "chapter5.docx"
    custom.write_bytes(b"PK" + b"x" * 2000)
    path = assets_mod.find_chapter5_template()
    assert path == custom
    status = assets_mod.chapter5_template_status()
    assert status["custom"] is True
    assert status["source"] == "custom"
    restored = assets_mod.restore_chapter5_template()
    assert restored["custom"] is False
    assert not custom.exists()


def test_rmb_uppercase_integer() -> None:
    assert rmb_uppercase(518800) == "伍拾壹万捌仟捌佰元整"
    assert rmb_uppercase(0) == "零元整"
    assert rmb_uppercase("1.23") == "壹元贰角叁分"


def test_rmb_lowercase() -> None:
    assert rmb_lowercase(518800) == "518,800.00"
    assert rmb_lowercase(Decimal("1000.5")) == "1,000.50"


def test_build_bid_docx_fills_company(tmp_path) -> None:
    brief = BidBrief(
        projectName="高途智成港（一期）汽车充电桩采购安装项目",
        tenderer="郑州有爱文化科技有限公司",
        bidContent="地下车库汽车充电桩供货、安装调试及技术服务",
        bidPriceYuan=518800,
        legalPersonName="测试",
        attachQualifications=False,
        quoteSource="form",
        quoteTitle="本标工程量",
        quoteSourceIncTax=518800,
        quoteLines=_sample_quote_lines(),
        deviationLines=[
            DeviationLine(
                seq="1",
                requirement="交流充电桩须满足招标文件第七章技术规格及安装调试要求",
                response="按招标文件供货、安装、调试",
                deviation="无偏差",
            )
        ],
    )
    dest = tmp_path / "bid.docx"
    path, warnings = build_bid_docx(
        brief, dest, qualification_pdf=None, catalog_media={}
    )
    assert path.is_file()
    assert path.stat().st_size > 2000
    assert not any("未找到资质" in w for w in warnings)

    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    all_text = text + "\n" + table_text
    assert "河南伟泰光电科技有限公司" in all_text
    assert "伍拾壹万捌仟捌佰元整" in all_text
    assert "高途智成港" in all_text
    assert "测试" in all_text
    assert "投  标  文  件" in text
    assert "录" in text
    assert "投标函及投标函附录" in text
    assert "技术标（实施方案）" in text
    assert "一、文字描述" in text
    assert "二、图纸" in text
    assert "法定代表人身份证明" in text
    assert "第五章" not in text
    assert "本部分附企业介绍" not in all_text
    assert "待补附件占位" not in all_text
    assert "（一）法定代表人身份证正反面" not in all_text
    assert "（在此粘贴扫描件）" in all_text
    assert "投标承诺书" in text
    assert "附件五：" in text
    assert "我公司现做出如下承诺" in text
    assert "保证投标文件无虚假内容" in text
    assert any('w:val="dashed"' in t._tbl.xml for t in doc.tables)

    toc = next(p for p in doc.paragraphs if "一、投标函及投标函附录" in p.text)
    toc_xml = toc._p.xml
    assert "w:numPr" not in toc_xml
    assert "PAGEREF" in toc_xml
    assert "toc_letter" in toc_xml
    assert 'w:leader="dot"' in toc_xml
    toc_sizes = {run.font.size.pt for run in toc.runs if run.font.size is not None}
    assert 12.0 in toc_sizes
    toc_blob = "\n".join(p.text for p in doc.paragraphs)
    assert "1.1 投标函" in toc_blob or "1.1投标函" in toc_blob.replace(" ", "")
    assert "1.2.1 分项报价表" in toc_blob or "1.2.1分项报价表" in toc_blob.replace(" ", "")
    assert "2.1 法定代表人身份证明" in toc_blob or "2.1法定代表人身份证明" in toc_blob.replace(" ", "")
    assert any("文字描述" in (p.text or "") and (p.text or "").strip().startswith("6.") for p in doc.paragraphs)

    assert "（招标人名称）" in text
    assert brief.tenderer in text
    assert f"{brief.tenderer}）：" not in text
    assert f"{brief.tenderer} )：" not in text

    underlined = any("w:u" in p._p.xml for p in doc.paragraphs)
    assert underlined

    # 横线填空字：有下划线、不加粗、不强制放大字号
    blank_filled = [
        run
        for p in doc.paragraphs
        for run in p.runs
        if (run.text or "").strip()
        and run.underline is True
        and "河南伟泰" in (run.text or "")
    ]
    assert blank_filled
    for run in blank_filled:
        assert run.bold is not True
        rpr = run._element.find(qn("w:rPr"))
        assert rpr is not None
        bold_el = rpr.find(qn("w:b"))
        assert bold_el is None
        sz = rpr.find(qn("w:sz"))
        if sz is not None:
            assert int(sz.get(qn("w:val")) or "0") <= 24  # ≤12pt
        # 不再强制把填空改成宋体放大观感，应跟邻行一致或无自设字号
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            assert rfonts.get(qn("w:eastAsia")) != "黑体"

    cover_bidder = next(
        p
        for p in doc.paragraphs
        if "投标人" in "".join((p.text or "").split())
        and "盖单位公章" in (p.text or "")
        and "河南伟泰" in (p.text or "")
        and "法定代表人" not in (p.text or "")
    )
    cover_legal = next(
        p
        for p in doc.paragraphs
        if "法定代表人或其委托代理人：" in (p.text or "")
        and "（签字）" in (p.text or "")
        and "盖单位公章" not in (p.text or "")
        and "签字或盖章" not in (p.text or "")
    )
    cover_date = next(
        p
        for p in doc.paragraphs
        if "".join((p.text or "").split()).startswith("日期：")
        and "2026" in (p.text or "")
    )
    assert "法定代表人" not in (cover_bidder.text or "")
    assert float(cover_date.paragraph_format.left_indent or 0) > 0
    cover_style = cover_bidder._p.find(qn("w:pPr"))
    cover_pstyle = cover_style.find(qn("w:pStyle")) if cover_style is not None else None
    assert cover_pstyle is None or cover_pstyle.get(qn("w:val")) != "Heading1"
    legal_label = next(r for r in cover_legal.runs if "法定代表人" in (r.text or ""))
    assert legal_label.bold is not True
    legal_fonts = None
    legal_rpr = legal_label._element.find(qn("w:rPr"))
    if legal_rpr is not None and legal_rpr.find(qn("w:rFonts")) is not None:
        legal_fonts = legal_rpr.find(qn("w:rFonts")).get(qn("w:eastAsia"))
    assert legal_fonts not in {"微软雅黑", "黑体", "Microsoft YaHei", "SimHei"}
    assert any(r.underline is True for r in cover_legal.runs)
    assert "（签字）" in (cover_legal.text or "")
    assert "郭志伟" not in (cover_legal.text or "")
    assert cover_legal._p.find(".//" + qn("w:tab")) is None
    assert "\u3000" in (cover_legal.text or "")
    assert "——" not in (cover_legal.text or "") and "---" not in (cover_legal.text or "")
    spaced = next(
        p for p in doc.paragraphs if "投　标　人" in (p.text or "") and "河南伟泰" in (p.text or "")
    )
    assert "w:drawing" not in spaced._p.xml
    assert "w:pict" not in spaced._p.xml
    assert "AlternateContent" not in spaced._p.xml

    east_asia = {
        run._element.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:eastAsia"))
        for p in doc.paragraphs
        for run in p.runs
        if run.text.strip()
        and run._element.find(qn("w:rPr")) is not None
        and run._element.find(qn("w:rPr")).find(qn("w:rFonts")) is not None
    }
    assert "宋体" in east_asia

    footer_xml = doc.sections[-1].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml
    assert "w:fldChar" in footer_xml
    header_text = doc.sections[-1].header.paragraphs[0].text
    assert "高途智成港" in header_text
    assert "投标文件" not in header_text
    assert "高途智成港" in (doc.sections[0].header.paragraphs[0].text or "")
    assert "投标文件" not in (doc.sections[0].header.paragraphs[0].text or "")
    # 分页改造后不再写 updateFields=true，避免 OnlyOffice/WPS 打开时全量更新域卡顿
    assert not doc.settings.element.findall(qn("w:updateFields"))
    bookmark_names = [
        el.get(qn("w:name")) for el in doc.element.iter(qn("w:bookmarkStart"))
    ]
    assert "toc_letter" in bookmark_names
    assert "toc_commit" in bookmark_names

    appendix = next(t for t in doc.tables if "项目名称" in t.rows[0].cells[0].text)
    assert brief.projectName in appendix.rows[0].cells[-1].text
    assert "w:pBdr" not in appendix.rows[0].cells[-1]._tc.xml
    for row in appendix.rows:
        assert "<w:u " not in row.cells[-1]._tc.xml

    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    deviation = next(
        t
        for t in doc.tables
        if len(t.rows[0].cells) >= 4 and "招标文件要求" in t.rows[0].cells[1].text
    )
    for row in deviation.rows:
        for cell in row.cells:
            assert "w:pBdr" not in cell._tc.xml
            for para in cell.paragraphs:
                for run in para.runs:
                    assert run.underline is not True
                    rpr = run._element.find(qn("w:rPr"))
                    if rpr is not None:
                        assert rpr.find(qn("w:u")) is None
    serial = deviation.rows[1].cells[0]
    assert serial.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert serial.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    note = deviation.rows[1].cells[3]
    assert "无偏差" in note.text
    assert note.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert note.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    long_req = deviation.rows[1].cells[1]
    assert long_req.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert long_req.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER

    quote = next(
        t
        for t in doc.tables
        if any("不含税综合单价" in (c.text or "") for c in t.rows[0].cells)
    )
    assert "交流充电桩" in quote.rows[1].cells[1].text
    assert "10" in quote.rows[1].cells[4].text
    assert "133" not in quote.rows[1].cells[4].text
    assert "56" not in quote.rows[2].cells[4].text
    assert "518,800.00" in quote.rows[-1].cells[-1].text
    assert not any("二次报价函" in w for w in warnings)
    assert not any("折算" in w for w in warnings)


def test_quote_table_fits_page_width(tmp_path) -> None:
    brief = BidBrief(
        projectName="宽表测试",
        tenderer="招标人",
        bidPriceYuan=518800,
        legalPersonName="测试",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
        quoteLines=_sample_quote_lines(),
        quoteSourceIncTax=518800,
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    quote = next(
        t
        for t in doc.tables
        if any("不含税综合单价" in (c.text or "") for c in t.rows[0].cells)
    )
    tbl_w = quote._tbl.tblPr.find(qn("w:tblW"))
    assert tbl_w is not None
    assert tbl_w.get(qn("w:type")) == "dxa"
    table_twips = int(tbl_w.get(qn("w:w")))
    section = doc.sections[0]
    usable = (
        int(section.page_width.twips)
        - int(section.left_margin.twips)
        - int(section.right_margin.twips)
    )
    assert table_twips <= usable
    layout = quote._tbl.tblPr.find(qn("w:tblLayout"))
    assert layout is not None
    assert layout.get(qn("w:type")) == "fixed"
    settings_xml = doc.settings.element.xml
    assert "hideSpellingErrors" in settings_xml
    assert "hideGrammaticalErrors" in settings_xml
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rfonts = rpr_default.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:eastAsia")) == "SimSun"
    assert qn("w:eastAsiaTheme") not in rfonts.attrib


def test_quote_sheet_scales_to_second_round() -> None:
    from api.services.tenders.assets import find_quote_xlsx
    from api.services.tenders.quote import SECOND_ROUND_YUAN, load_quote_sheet, scale_quote

    path = find_quote_xlsx()
    assert path is not None
    sheet = load_quote_sheet(path)
    assert sheet.lines[0].qty == 133
    scaled = scale_quote(sheet, SECOND_ROUND_YUAN)
    assert scaled.lines[0].qty == 133
    assert scaled.total_inc_tax == SECOND_ROUND_YUAN
    assert scaled.lines[0].unit_price != sheet.lines[0].unit_price


def test_build_bid_docx_quote_scaled(tmp_path) -> None:
    brief = BidBrief(
        bidPriceYuan=8000,
        legalPersonName="测试",
        attachQualifications=False,
        quoteLines=_sample_quote_lines(),
        quoteSourceIncTax=11000,
    )
    path, warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(str(path))
    quote = next(
        t
        for t in doc.tables
        if any("不含税综合单价" in (c.text or "") for c in t.rows[0].cells)
    )
    assert "交流充电桩" in quote.rows[1].cells[1].text
    assert "10" in quote.rows[1].cells[4].text
    assert "8,000.00" in quote.rows[-1].cells[-1].text
    assert any("折算" in w for w in warnings)
    appendix = next(t for t in doc.tables if "项目名称" in t.rows[0].cells[0].text and len(t.rows[0].cells) >= 3)
    assert brief.projectName in appendix.rows[0].cells[-1].text
    pay = next(
        t
        for t in doc.tables
        if len(t.rows[0].cells) >= 2 and "支付阶段" in t.rows[0].cells[1].text
    )
    for row in pay.rows[1:]:
        pct_cell = row.cells[2]
        assert "支付至" in pct_cell.text
        assert "%" in pct_cell.text
        para = next(p for p in pct_cell.paragraphs if "支付至" in p.text)
        assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert "w:u" in para._p.xml
        assert "w:tab" in para._p.xml
    assert "17" in pay.rows[3].cells[2].text
    assert "3" in pay.rows[4].cells[2].text
    settle_xml = next(p._p.xml for p in pay.rows[3].cells[2].paragraphs if "支付至" in p.text)
    prepay_xml = next(p._p.xml for p in pay.rows[1].cells[2].paragraphs if "支付至" in p.text)
    assert settle_xml.count("w:tab") == prepay_xml.count("w:tab")


def test_placeholders_can_be_skipped(tmp_path) -> None:
    brief = BidBrief(
        bidPriceYuan=518800,
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "（一）法定代表人身份证正反面" not in text
    assert "本部分附企业介绍" not in text
    assert "投标承诺书" not in text


def test_factory_commitment_follows_tender_name(tmp_path) -> None:
    brief = BidBrief(
        projectName="厂内新能源充电桩采购项目",
        tenderer="测试招标人新能源有限公司",
        bidPriceYuan=100000,
        legalPersonName="张三",
        deliveryDays=45,
        warrantyYears=3,
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document

    doc = Document(str(path))
    addr = next(p for p in doc.paragraphs if (p.text or "").startswith("测试招标人新能源有限公司"))
    assert "：" in addr.text
    factory = next(p for p in doc.paragraphs if "作为充电设备生产厂商" in (p.text or ""))
    assert "我单位" in factory.text
    assert "厂内新能源充电桩采购项目" in factory.text
    assert "联源热电" not in factory.text
    assert "45" in factory.text
    assert "3" in factory.text
    filled = [run.text for run in factory.runs if run.underline is True]
    assert any("厂内新能源充电桩采购项目" in t for t in filled)
    assert any("承诺单位" in (p.text or "") for p in doc.paragraphs)


def test_factory_commitment_falls_back_to_project_name(tmp_path) -> None:
    brief = BidBrief(
        projectName="备件采购投标项目",
        tenderer="",
        bidPriceYuan=100000,
        legalPersonName="张三",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document

    doc = Document(str(path))
    addr = next(p for p in doc.paragraphs if (p.text or "").startswith("备件采购投标项目"))
    assert "：" in addr.text
    factory = next(p for p in doc.paragraphs if "作为投标产品生产厂商" in (p.text or ""))
    assert "充电设备生产厂商" not in factory.text


def test_chapter5_letter_contact_auth_factory_layout(tmp_path) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    brief = BidBrief(
        projectName="厂内新能源充电桩采购项目",
        tenderer="测试招标人新能源有限公司",
        bidPriceYuan=518800,
        bidderAddress="河南省郑州市高新区开发区梧桐街与红松路交叉口东南角远大产业园区内6号楼三层",
        bidderPhone="17630567052",
        bidderWebsite="",
        bidderFax="",
        bidderPostcode="",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=True,
        bidDate="2026-09-09",
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    def _has_blank_u(para) -> bool:
        xml = para._p.xml
        if "w:tab" in xml and 'w:val="single"' in xml:
            return True
        return any(
            r.underline is True
            and (
                "\u3000" in (r.text or "")
                or r._element.find(qn("w:tab")) is not None
            )
            for r in para.runs
        )

    doc = Document(str(path))
    addr = next(
        p
        for p in doc.paragraphs
        if "址" in (p.text or "") and "梧桐街" in (p.text or "")
    )
    assert any(run.underline is True and "梧桐街" in (run.text or "") for run in addr.runs)
    legal_i = next(i for i, p in enumerate(doc.paragraphs) if "法定代表人身份证明" in (p.text or ""))
    legal_addr = next(
        p
        for p in doc.paragraphs[legal_i:]
        if "地址" in (p.text or "") and "梧桐街" in (p.text or "")
    )
    assert float(legal_addr.paragraph_format.line_spacing or 0) >= 1.9
    web = next(p for p in doc.paragraphs if "".join((p.text or "").split()).startswith("网址"))
    assert _has_blank_u(web)
    assert (web.text or "").count("\u3000") >= 18
    assert web._p.find(".//" + qn("w:tab")) is None
    assert "─" not in (web.text or "")
    assert "---" not in (web.text or "")
    phone = next(
        p
        for p in doc.paragraphs
        if "17630567052" in (p.text or "") and "电话" in "".join((p.text or "").split())
    )
    assert any(r.underline is True and "17630567052" in (r.text or "") for r in phone.runs)
    assert (phone.text or "").count("\u3000") >= 10
    letter_rep = next(
        p
        for p in doc.paragraphs
        if "法定代表人或其委托代理人" in (p.text or "") and "签字或盖章" in (p.text or "")
    )
    assert _has_blank_u(letter_rep) or any(r.underline is True for r in letter_rep.runs)
    assert "（签字或盖章）" in (letter_rep.text or "")
    assert "郭志伟" not in (letter_rep.text or "")
    assert "梧桐街" not in (letter_rep.text or "")
    letter_label = next(r for r in letter_rep.runs if (r.text or "").startswith("法定代表人"))
    assert letter_label.bold is not True
    letter_east = None
    letter_rpr = letter_label._element.find(qn("w:rPr"))
    if letter_rpr is not None and letter_rpr.find(qn("w:rFonts")) is not None:
        letter_east = letter_rpr.find(qn("w:rFonts")).get(qn("w:eastAsia"))
    assert letter_east not in {"微软雅黑", "黑体", "Microsoft YaHei", "SimHei"}
    fax = next(p for p in doc.paragraphs if "".join((p.text or "").split()).startswith("传真"))
    assert _has_blank_u(fax)
    assert "邮政编码" not in (fax.text or "")
    assert (fax.text or "").count("\u3000") >= 18
    post_i = next(i for i, p in enumerate(doc.paragraphs) if "邮政编码" in (p.text or ""))
    post = doc.paragraphs[post_i]
    assert (post.text or "").count("\u3000") >= 18
    letter_date = next(
        p
        for p in doc.paragraphs[post_i:]
        if "".join((p.text or "").split()).startswith("日期") and "2026" in (p.text or "")
    )
    assert "2026年" in (letter_date.text or "").replace(" ", "").replace("\u3000", "")
    assert (letter_date.text or "").count("\u3000") >= 10
    assert float(letter_date.paragraph_format.left_indent or 0) > 0
    pay_bidder = next(
        p for p in doc.paragraphs if "投标人名称" in (p.text or "") and "伟泰" in (p.text or "")
    )
    assert "法定代表人" not in (pay_bidder.text or "")
    assert "盖单位" in (pay_bidder.text or "") and "章）" in (pay_bidder.text or "")
    until = next(p for p in doc.paragraphs if "委托期限" in (p.text or ""))
    assert "至本项目投标有效期届满" in (until.text or "")
    assert "代理人无转委托权" in (until.text or "")
    legal_heads = [
        i for i, p in enumerate(doc.paragraphs) if (p.text or "").strip() == "法定代表人身份证明"
    ]
    legal_i = legal_heads[-1]
    legal_bidder_i = next(
        i
        for i, p in enumerate(doc.paragraphs[legal_i:], start=legal_i)
        if "".join((p.text or "").split()).startswith("投标人")
        and "盖单位公章" in (p.text or "")
        and "法定代表人" not in (p.text or "")
    )
    legal_bidder = doc.paragraphs[legal_bidder_i]
    legal_date = next(
        p
        for p in doc.paragraphs[legal_bidder_i:]
        if "".join((p.text or "").split()).startswith("日期：") and "2026" in (p.text or "")
    )
    assert "日期" in "".join((legal_date.text or "").split())
    assert abs(
        float(legal_bidder.paragraph_format.left_indent or 0)
        - float(legal_date.paragraph_format.left_indent or 0)
    ) < 50000
    auth_heads = [i for i, p in enumerate(doc.paragraphs) if (p.text or "").strip() == "授权委托书"]
    auth_i = auth_heads[-1]
    auth_bidder_i = next(
        i
        for i, p in enumerate(doc.paragraphs[auth_i:], start=auth_i)
        if "".join((p.text or "").split()).startswith("投标人") and "盖单位公章" in (p.text or "")
    )
    auth_bidder = doc.paragraphs[auth_bidder_i]
    assert "法定代表人" not in (auth_bidder.text or "")
    auth_date = next(
        p
        for p in doc.paragraphs[auth_bidder_i:]
        if "".join((p.text or "").split()).startswith("日期：") and "2026" in (p.text or "")
    )
    assert abs(
        float(auth_bidder.paragraph_format.left_indent or 0)
        - float(auth_date.paragraph_format.left_indent or 0)
    ) < 50000
    auth_end = next(
        i
        for i, p in enumerate(doc.paragraphs[auth_i:], start=auth_i)
        if (p.text or "").strip() in {"技术偏差表", "原厂生产承诺"}
    )
    agent_lines = [
        i
        for i, p in enumerate(doc.paragraphs[auth_i:auth_end], start=auth_i)
        if "".join((p.text or "").split()).startswith("委托代理人")
    ]
    agent_id_lines = [
        i
        for i, p in enumerate(doc.paragraphs[auth_i:auth_end], start=auth_i)
        if "".join((p.text or "").split()).startswith("身份证号码")
    ]
    assert len(agent_lines) == 1
    assert agent_lines[0] > auth_bidder_i
    assert len(agent_id_lines) == 2
    assert auth_bidder_i < agent_id_lines[0] < agent_lines[0] < agent_id_lines[1]
    factory_body = next(p for p in doc.paragraphs if "作为充电设备生产厂商" in (p.text or ""))
    assert factory_body.paragraph_format.first_line_indent
    assert any("承诺单位" in (p.text or "") for p in doc.paragraphs)
    commit_title = next(i for i, p in enumerate(doc.paragraphs) if (p.text or "").strip() == "投标承诺书")
    sign = next(
        p
        for p in doc.paragraphs[commit_title:]
        if "投 标 人" in (p.text or "") and "盖章" in (p.text or "")
    )
    assert sign.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_commitment_letter_fills_tenderer(tmp_path) -> None:
    brief = BidBrief(
        projectName="厂内新能源充电桩采购项目",
        tenderer="二连浩特市联源热电有限公司",
        bidPriceYuan=100000,
        legalPersonName="张三",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=True,
        bidDate="2026-08-25",
    )
    path, warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "附件五：" in text
    assert "投标承诺书" in text
    assert "二连浩特市联源热电有限公司" in text
    assert "河南伟泰光电科技有限公司" in text
    assert "张三" in text
    assert "2026" in text
    assert any("投标承诺书" in w for w in warnings)
    min_header = int(Cm(1.5))
    for i, section in enumerate(doc.sections):
        assert int(section.header_distance or 0) >= min_header, f"section {i}"
        header_text = "\n".join(p.text for p in section.header.paragraphs)
        assert "厂内新能源充电桩采购项目" in header_text
        assert "投标文件" not in header_text
    annex = next(p for p in doc.paragraphs if (p.text or "").strip() == "附件五：")
    prev = annex._p.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        assert prev.find(qn("w:pPr")) is None or prev.find(qn("w:pPr")).find(qn("w:pBdr")) is None
    # 填空不加粗
    filled = [
        run
        for p in doc.paragraphs
        for run in p.runs
        if run.underline is True and "联源热电" in (run.text or "")
    ]
    assert filled
    for run in filled:
        assert run.bold is not True


def test_extra_placeholder_box(tmp_path) -> None:
    from api.services.tenders.schema import PlaceholderItem

    brief = BidBrief(
        bidPriceYuan=518800,
        attachQualifications=False,
        extraPlaceholders=[PlaceholderItem(key="iso", title="ISO体系证书", hint="按邀请书第3.8条")],
    )
    path, warnings = build_bid_docx(
        brief, tmp_path / "bid.docx", qualification_pdf=None, catalog_media={}
    )
    from docx import Document

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ISO体系证书" not in text
    assert "附件：资料库扫描件" not in text
    assert "（一）法定代表人身份证正反面" not in text


def test_slot_image_fills_placeholder(tmp_path, monkeypatch) -> None:
    """法人身份证写入身份证明页，不再出现在文末附件小标题。"""
    from PIL import Image

    from api.services.tenders import slots as slots_mod
    from api.services.tenders.slots import clear_slot, save_slot_file

    root = tmp_path / "tender-assets"
    root.mkdir()
    monkeypatch.setattr(slots_mod, "tender_assets_dir", lambda: root)

    img = Image.new("RGB", (120, 80), color=(200, 200, 200))
    buf = tmp_path / "id.png"
    img.save(buf, format="PNG")
    save_slot_file("id_legal", filename="id.png", data=buf.read_bytes(), replace=True)

    brief = BidBrief(
        bidPriceYuan=100000,
        attachQualifications=False,
        includePlaceholders=True,
        includeCommitment=False,
    )
    path, warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    texts = [p.text or "" for p in doc.paragraphs]
    assert not any("附件：资料库扫描件" in t for t in texts)
    assert not any("（一）法定代表人身份证正反面" in t for t in texts)
    legal_i = next(i for i, t in enumerate(texts) if t.strip() == "法定代表人身份证明")
    legal_sign_i = next(
        i
        for i, t in enumerate(texts[legal_i:], start=legal_i)
        if "投标人" in "".join(t.split()) and "盖单位公章" in t and "法定代表人" not in t
    )
    xml = doc.element.body.xml
    assert "a:blip" in xml or "pic:blipFill" in xml
    heading = doc.paragraphs[legal_i]
    found = False
    el = heading._element.getnext()
    sign_el = doc.paragraphs[legal_sign_i]._element
    while el is not None and el is not sign_el:
        if el.findall(".//" + qn("a:blip")) or el.findall(".//" + qn("pic:blipFill")):
            found = True
            break
        el = el.getnext()
    assert found
    clear_slot("id_legal")


def test_id_legal_front_back_side_by_side_one_page(tmp_path) -> None:
    """身份证正反面合成一张横版小图，紧跟小标题，不各占一页。"""
    from PIL import Image
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm

    from api.services.tenders.placeholders import append_placeholder_section
    from api.services.tenders.schema import PlaceholderItem

    front = tmp_path / "id-front.png"
    back = tmp_path / "id-back.png"
    Image.new("RGB", (220, 350), (190, 170, 150)).save(front)
    Image.new("RGB", (220, 350), (150, 170, 190)).save(back)

    doc = Document()
    doc.add_paragraph("技术标实施方案")
    filled, boxes, _notes = append_placeholder_section(
        doc,
        [PlaceholderItem(key="id_legal", title="法定代表人身份证正反面")],
        {"id_legal": [front, back]},
    )
    assert filled == 1
    assert boxes == 0

    heading = next(p for p in doc.paragraphs if "法定代表人身份证正反面" in (p.text or ""))
    extents = []
    el = heading._element.getnext()
    while el is not None:
        if el.tag == qn("w:tbl"):
            break
        extents.extend(el.findall(".//" + qn("wp:extent")))
        el = el.getnext()
    assert len(extents) == 1
    cx = int(extents[0].get("cx") or 0)
    cy = int(extents[0].get("cy") or 0)
    assert cx > cy
    assert cx > int(Cm(15.8).emu)
    assert cx < int(Cm(17.2).emu)
    assert cy < int(Cm(6.5).emu)


def test_placeholder_skips_internal_notes_and_keeps_heading_with_body(tmp_path) -> None:
    from PIL import Image
    from docx import Document
    from docx.oxml.ns import qn

    from api.services.tenders.placeholders import append_placeholder_section
    from api.services.tenders.schema import PlaceholderItem

    tall = tmp_path / "report.png"
    Image.new("RGB", (400, 720), (200, 200, 200)).save(tall)
    skip_pdf = tmp_path / "finance.pdf"
    skip_pdf.write_bytes(b"%PDF-1.4\n")

    doc = Document()
    doc.add_paragraph("技术标实施方案")
    filled, boxes, notes = append_placeholder_section(
        doc,
        [
            PlaceholderItem(key="product", title="所投产品检测报告 / 3C / 对应功率桩型证明"),
            PlaceholderItem(key="finance", title="近三年财务审计报告"),
            PlaceholderItem(key="bond", title="投标保证金缴存回单"),
        ],
        {"product": [tall], "finance": [skip_pdf]},
    )
    assert filled >= 1
    blob = "\n".join(p.text or "" for p in doc.paragraphs)
    assert "资料库已有" not in blob
    assert "为加快生成" not in blob
    assert "576aa2fbf71a" not in blob
    assert "未写入扫描页" not in blob
    assert "finance.pdf" not in blob
    xml = doc.element.body.xml
    assert "a:blip" in xml or "pic:blipFill" in xml
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "（在此粘贴扫描件）" in table_text
    assert any("未嵌入 Word" in n or "装订时请" in n for n in notes) or boxes >= 1

    finance = next(p for p in doc.paragraphs if "近三年财务审计报告" in (p.text or ""))
    prev = finance._element.getprevious()
    assert prev is not None
    has_break = any(br.get(qn("w:type")) == "page" for br in prev.iter(qn("w:br")))
    assert has_break


def test_placeholder_embeds_credit_pdf_and_merges_slots(tmp_path) -> None:
    from PIL import Image
    from docx import Document

    from api.services.tenders.placeholders import append_placeholder_section
    from api.services.tenders.schema import PlaceholderItem

    pdf = tmp_path / "信用中国.pdf"
    Image.new("RGB", (320, 450), (240, 240, 240)).save(pdf, "PDF")
    skip_pdf = tmp_path / "skip.pdf"
    skip_pdf.write_bytes(b"%PDF-1.4\n")

    doc = Document()
    doc.add_paragraph("技术标实施方案")
    filled, boxes, notes = append_placeholder_section(
        doc,
        [
            PlaceholderItem(
                key="credit",
                title="信用中国查询页及国家企业信用信息公示（股东）",
            ),
            PlaceholderItem(key="m_credit_shot", title="信用截图"),
            PlaceholderItem(key="finance", title="近三年财务审计报告"),
        ],
        {"m_credit_shot": [pdf], "finance": [skip_pdf]},
    )
    assert filled >= 1
    assert boxes == 1
    blob = "\n".join(p.text or "" for p in doc.paragraphs)
    assert "信用中国查询页" in blob
    assert "信用截图" not in blob
    xml = doc.element.body.xml
    assert "a:blip" in xml or "pic:blipFill" in xml
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert table_text.count("（在此粘贴扫描件）") == 1
    assert not any("信用" in (n or "") and "未嵌入" in (n or "") for n in notes)


def test_credit_pdf_embed_keeps_readable_pixels(tmp_path) -> None:
    import zipfile

    from PIL import Image
    from docx import Document

    from api.services.tenders.placeholders import append_placeholder_section
    from api.services.tenders.schema import PlaceholderItem

    pdf = tmp_path / "信用中国.pdf"
    Image.new("RGB", (1240, 1754), (255, 255, 255)).save(pdf, "PDF")
    dest = tmp_path / "credit.docx"
    doc = Document()
    filled, boxes, _notes = append_placeholder_section(
        doc,
        [PlaceholderItem(key="credit", title="信用中国查询页及国家企业信用信息公示（股东）")],
        {"credit": [pdf]},
    )
    assert filled == 1
    assert boxes == 0
    doc.save(str(dest))
    with zipfile.ZipFile(dest) as zf:
        names = [n for n in zf.namelist() if n.startswith("word/media/")]
        assert names
        with zf.open(names[0]) as fh:
            im = Image.open(fh)
            assert im.format in {"PNG", "JPEG"}
            assert max(im.size) >= 1600


def test_collect_slots_status_keys() -> None:
    from api.services.tenders.slots import library_payload, list_slots_status

    rows = list_slots_status()
    keys = {r["key"] for r in rows}
    assert "id_legal" in keys
    assert "bond" in keys

    lib = library_payload()
    assert lib["totalCount"] == len(lib["slots"])
    assert lib["totalCount"] >= 11
    assert {s["key"] for s in lib["slots"]} >= {"license", "bank_permit", "iso", "perf"}
    assert all(s.get("title") for s in lib["slots"])
    assert "常备" in str(lib.get("hint") or "") or "扫描" in str(lib.get("hint") or "")


def test_extract_patch_keeps_bidder_and_legal() -> None:
    from api.services.tenders.extract import apply_extract_patch, parse_extract_payload, parse_kb_ids
    from api.services.tenders.placeholders import collect_slots
    from api.services.tenders.schema import PlaceholderItem

    base = BidBrief(
        projectName="旧项目",
        tenderer="旧招标人",
        legalPersonName="张三",
        bidderName="河南伟泰光电科技有限公司",
        bidPriceYuan=100,
    )
    blob = """<think>先读邀请书</think>
```json
{
  "projectName": "南京广场光伏充电桩项目",
  "tenderer": "南京某建设有限公司",
  "deliveryDays": 45,
  "bidPriceYuan": 0,
  "bidderName": "郑州容新新能源有限公司",
  "legalPersonName": "李四",
  "notes": ["核对保证金账户"],
  "missingMaterials": [{"key": "iso", "title": "ISO体系证书", "reason": "资格预审额外要求"}]
}
```
"""
    patch = parse_extract_payload(blob)
    brief, filled = apply_extract_patch(base, patch)
    assert brief.projectName == "南京广场光伏充电桩项目"
    assert brief.tenderer == "南京某建设有限公司"
    assert brief.deliveryDays == 45
    assert brief.bidPriceYuan == 100
    assert brief.bidderName == "河南伟泰光电科技有限公司"
    assert brief.legalPersonName == "张三"
    assert "projectName" in filled
    assert "bidPriceYuan" not in filled
    slots = collect_slots([PlaceholderItem(key="iso", title="ISO体系证书", hint="资格预审额外要求")])
    assert any(s.key == "iso" for s in slots)
    assert any(s.key == "id_legal" for s in slots)
    assert parse_kb_ids('["a","b"]') == ["a", "b"]
    assert parse_kb_ids("a, b") == ["a", "b"]


def test_fill_empty_company_from_defaults_and_profile() -> None:
    from api.services.tenders.extract import apply_extract_patch, fill_empty_company_fields

    empty = BidBrief(
        bidderName="",
        bidderAddress="",
        bidderEmail="",
        bidderPhone="",
        foundedDate="",
        legalPersonName="",
        legalPersonAge="",
        legalPersonIdNo="",
        legalPersonTitle="",
    )
    patched, filled = apply_extract_patch(
        empty,
        {"projectName": "某充电站", "tenderer": "某招标人", "bidderName": "郑州容新新能源有限公司"},
    )
    assert patched.projectName == "某充电站"
    assert patched.tenderer == "某招标人"
    assert patched.bidderName == "河南伟泰光电科技有限公司"
    assert patched.bidderAddress == ""
    assert patched.legalPersonName == ""

    brief, company_filled = fill_empty_company_fields(
        patched,
        empty,
        {"legalPersonName": "王五", "foundedDate": "2016年3月", "bidderPhone": "0371-0000000"},
    )
    assert brief.bidderName == "河南伟泰光电科技有限公司"
    assert "梧桐街" in brief.bidderAddress
    assert brief.bidderEmail == "gzwceo@163.com"
    assert brief.legalPersonName == "王五"
    assert brief.foundedDate == "2016年3月"
    assert brief.bidderPhone == "0371-0000000"
    assert "bidderAddress" in company_filled
    assert "legalPersonName" in company_filled

    kept, _ = fill_empty_company_fields(
        BidBrief(bidderAddress="用户手填地址", legalPersonName="张三"),
        {"bidderAddress": "历史地址", "legalPersonName": "历史法人"},
    )
    assert kept.bidderAddress == "用户手填地址"
    assert kept.legalPersonName == "张三"


def test_extract_payload_prefers_tender_json_over_think_fragment() -> None:
    from api.services.tenders.extract import parse_extract_payload

    blob = """好的，我先看邀请书 { "cars": 8
{"projectName": "南京广场光伏充电桩项目", "tenderer": "南京某建设有限公司", "quoteLines": []}
"""
    patch = parse_extract_payload(blob)
    assert patch["projectName"] == "南京广场光伏充电桩项目"
    assert patch["quoteLines"] == []


def test_parse_generic_xlsx_bom(tmp_path) -> None:
    from decimal import Decimal

    from openpyxl import Workbook

    from api.services.tenders.quote import parse_quote_path

    wb = Workbook()
    ws = wb.active
    ws.append(["南京广场光伏充电桩工程量清单"])
    ws.append(["序号", "项目名称", "规格", "单位", "数量", "单价", "合价"])
    ws.append([1, "7kW交流桩", "7kW", "台", 20, 1000, 20000])
    ws.append([2, "30kW直流桩", "30kW", "台", 8, 5000, 40000])
    ws.append([None, "不含税合计", None, None, None, None, 60000])
    ws.append([None, "税率", None, None, None, None, 0.13])
    ws.append([None, "含税合计", None, None, None, None, 67800])
    path = tmp_path / "bom.xlsx"
    wb.save(path)
    sheet = parse_quote_path(path)
    assert sheet is not None
    assert len(sheet.lines) == 2
    assert sheet.lines[0].qty == Decimal("20")
    assert sheet.lines[1].qty == Decimal("8")
    assert "南京广场" in sheet.title
    assert sheet.total_inc_tax == Decimal("67800")


def test_parse_xlsx_keeps_multiline_tech_specs(tmp_path) -> None:
    from openpyxl import Workbook

    from api.services.tenders.extract import deviation_lines_from_quote
    from api.services.tenders.quote import parse_quote_path

    spec = "(1)输入电压：220V\n(2)频率：50Hz\n(3)输出功率：7kW"
    wb = Workbook()
    ws = wb.active
    ws.append(["高途智成港（一期）汽车充电桩采购安装报价清单"])
    ws.append(["序号", "设备", "技术参数要求", "单位", "数量", "不含税综合单价（元）", "合价"])
    ws.append([1, "7KW 交流汽车充电桩", spec, "台", 133, 787.61, 104752.13])
    ws.append([2, "慢充立柱", "含供货安装", "个", 133, 97.35, 12947.55])
    ws.append([None, "不含税合计", None, None, None, None, 117699.68])
    ws.append([None, "税率", None, None, None, None, 0.13])
    ws.append([None, "含税合计", None, None, None, None, 133000])
    path = tmp_path / "quote_specs.xlsx"
    wb.save(path)
    sheet = parse_quote_path(path)
    assert sheet is not None
    assert sheet.lines[0].spec.startswith("(1)输入电压")
    assert "\n(2)" in sheet.lines[0].spec
    assert sheet.lines[1].spec == "含供货安装"
    dev = deviation_lines_from_quote(list(sheet.lines))
    assert "7KW" in dev[0].requirement
    assert "(1)输入电压" in dev[0].requirement


def test_parse_quote_from_pipe_text() -> None:
    from api.services.tenders.quote import parse_quote_from_text

    text = (
        "前言\n"
        "序号 | 设备 | 单位 | 数量 | 单价 | 合价\n"
        "1 | 充电桩 | 台 | 12 | 800 | 9600\n"
        "2 | 立柱 | 个 | 12 | 100 | 1200\n"
        "结束"
    )
    sheet = parse_quote_from_text(text)
    assert sheet is not None
    assert len(sheet.lines) == 2
    assert sheet.lines[0].name == "充电桩"
    assert int(sheet.lines[0].qty) == 12


def test_build_uses_brief_quote_lines(tmp_path) -> None:
    from api.services.tenders.schema import QuoteLineIn

    brief = BidBrief(
        projectName="南京广场光伏充电板安装项目",
        bidPriceYuan=22600,
        legalPersonName="测试",
        attachQualifications=False,
        includePlaceholders=False,
        quoteSource="file",
        quoteTitle="南京广场清单",
        quoteTaxRate=0.13,
        quoteSourceIncTax=22600,
        quoteLines=[
            QuoteLineIn(
                seq="1",
                name="7kW交流桩",
                spec="7kW",
                unit="台",
                qty=20,
                unitPrice=1000,
                amount=20000,
            )
        ],
    )
    path, warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document

    doc = Document(str(path))
    quote = next(
        t
        for t in doc.tables
        if any("不含税综合单价" in (c.text or "") for c in t.rows[0].cells)
    )
    assert "7kW交流桩" in quote.rows[1].cells[1].text
    assert "20" in quote.rows[1].cells[4].text
    assert "133" not in quote.rows[1].cells[4].text
    assert "22,600.00" in quote.rows[-1].cells[-1].text
    assert not any("二次报价函" in w for w in warnings)
    assert any("南京广场清单" in w or "分项报价来自" in w for w in warnings)


def test_quote_payload_from_llm_patch() -> None:
    from api.services.tenders.quote import quote_payload_from_patch

    sheet = quote_payload_from_patch(
        {
            "quoteTitle": "邀请书清单",
            "quoteTaxRate": 0.13,
            "quoteLines": [
                {"seq": "1", "name": "直流桩", "unit": "台", "qty": 6, "unitPrice": 4000, "amount": 24000}
            ],
        }
    )
    assert sheet is not None
    assert sheet.lines[0].name == "直流桩"
    assert int(sheet.lines[0].qty) == 6


def test_collect_slots_uses_catalog_instead_of_defaults() -> None:
    from api.services.tenders.placeholders import collect_slots
    from api.services.tenders.schema import PlaceholderItem

    catalog = [PlaceholderItem(key="iso_cert", title="ISO体系证书", hint="现行有效")]
    slots = collect_slots(
        [PlaceholderItem(key="iso_cert", title="忽略", hint="邀请书要求")],
        catalog=catalog,
    )
    assert [s.key for s in slots] == ["iso_cert"]
    assert slots[0].title == "ISO体系证书"
    assert slots[0].hint == "邀请书要求"

    extras = collect_slots(
        [PlaceholderItem(key="safety", title="安全生产许可证", hint="第3条")],
        catalog=catalog,
    )
    assert [s.key for s in extras] == ["iso_cert", "safety"]


def test_this_bid_keys_does_not_dump_catalog() -> None:
    from api.services.tenders.placeholders import collect_slots, this_bid_keys
    from api.services.tenders.records import (
        _brief_generate_issues,
        _missing_required_titles,
        _missing_slot_warnings,
    )
    from api.services.tenders.schema import BidBrief, PlaceholderItem

    catalog = [
        PlaceholderItem(key="id_legal", title="法定代表人身份证正反面"),
        PlaceholderItem(key="iso_cert", title="ISO体系证书"),
        PlaceholderItem(key="finance", title="近三年财务"),
    ]
    extra = [PlaceholderItem(key="iso_cert", title="ISO体系证书", hint="邀请书要求")]
    required, include = this_bid_keys(
        extra=extra,
        catalog=catalog,
        required_keys=["iso_cert"],
        include_keys=None,
        has_agent=False,
    )
    assert required == ["id_legal", "iso_cert"]
    assert include == ["id_legal", "iso_cert"]
    slots = collect_slots(extra, catalog=catalog, include_keys=include)
    assert [item.key for item in slots] == ["id_legal", "iso_cert"]

    required_agent, include_agent = this_bid_keys(
        extra=extra,
        catalog=[*catalog, PlaceholderItem(key="id_agent", title="授权代理人身份证")],
        required_keys=["iso_cert"],
        include_keys=["iso_cert", "finance"],
        has_agent=True,
    )
    assert required_agent == ["id_legal", "iso_cert", "id_agent"]
    assert include_agent == ["iso_cert", "finance", "id_legal", "id_agent"]

    required_no, include_no = this_bid_keys(
        extra=extra,
        catalog=[*catalog, PlaceholderItem(key="id_agent", title="授权代理人身份证")],
        required_keys=["iso_cert", "id_agent"],
        include_keys=["iso_cert", "id_agent", "finance"],
        has_agent=False,
    )
    assert "id_agent" not in required_no
    assert "id_agent" not in include_no

    issues = _brief_generate_issues(BidBrief(projectName="x"))
    assert "招标人" in issues
    assert "投标总价" in issues
    assert "报价清单" in issues
    assert "委托代理人（招标书要求授权委托）" not in issues
    assert "委托代理人（招标书要求授权委托）" in _brief_generate_issues(
        BidBrief(projectName="x", authNeed="required")
    )
    assert "代理人身份证号" in _brief_generate_issues(BidBrief(projectName="x", authNeed="required"))
    from api.services.tenders.schema import PerformanceRequirement

    perf_issues = _brief_generate_issues(
        BidBrief(
            projectName="x",
            performanceRequirement=PerformanceRequirement(keywords=["MES"], minCount=3),
        )
    )
    assert any("类似业绩" in x for x in perf_issues)
    missing = _missing_required_titles(["id_legal"], slots, {})
    assert missing == ["法定代表人身份证正反面"]
    notes = _missing_slot_warnings(missing)
    assert notes and "虚线框占位" in notes[0]
    assert _missing_slot_warnings([]) == []


def test_resolve_quote_sheet_skips_builtin_template() -> None:
    from api.services.tenders.quote import resolve_quote_sheet

    sheet, origin = resolve_quote_sheet(BidBrief(projectName="备件采购", bidPriceYuan=1000))
    assert sheet is None
    assert origin == "none"

    filled, form_origin = resolve_quote_sheet(
        BidBrief(quoteLines=_sample_quote_lines(), quoteSource="form")
    )
    assert filled is not None
    assert form_origin == "form"
    assert filled.lines[0].name == "交流充电桩"


def test_assets_only_use_bundled_or_storage() -> None:
    from pathlib import Path

    from api.services.tenders import assets as assets_mod

    text = Path(assets_mod.__file__).read_text(encoding="utf-8")
    assert "E:\\downLoad" not in text
    assert "F:\\伟泰" not in text


def test_slots_from_payload_maps_required_catalog_keys() -> None:
    from api.services.tenders.extract import slots_from_payload
    from api.services.tenders.schema import PlaceholderItem

    catalog = [
        PlaceholderItem(key="id_legal", title="法定代表人身份证正反面", hint="默认"),
        PlaceholderItem(key="iso_cert", title="ISO体系证书", hint="现行有效"),
    ]
    extra = slots_from_payload(
        {
            "requiredMaterials": [
                {"key": "iso_cert", "reason": "邀请书要求体系认证"},
                {"key": "unknown", "reason": "应忽略"},
            ],
            "missingMaterials": [{"key": "safety", "title": "安全生产许可证", "reason": "第3条"}],
        },
        catalog=catalog,
    )
    keys = [s.key for s in extra]
    assert "iso_cert" in keys
    assert "unknown" not in keys
    assert any(s.title == "安全生产许可证" for s in extra)


def test_invitation_materials_intersect_catalog_by_title() -> None:
    from api.services.tenders.match import (
        attachment_match_notes,
        build_attachment_match,
        split_invitation_materials,
    )
    from api.services.tenders.schema import PlaceholderItem

    catalog = [
        PlaceholderItem(key="id_legal", title="法定代表人身份证正反面", hint="默认"),
        PlaceholderItem(key="bond", title="投标保证金缴存回单"),
        PlaceholderItem(key="perf", title="类似项目合同及发票"),
    ]
    mapped, missing = split_invitation_materials(
        {
            "requiredMaterials": [
                {"key": "id_legal", "reason": "资格审查"},
                {"key": "invented_other_project", "reason": "应忽略无标题"},
                {"title": "投标保证金回单", "reason": "须缴纳"},
            ],
            "missingMaterials": [
                {"title": "法定代表人身份证", "reason": "重复点名"},
                {"title": "安全生产许可证", "reason": "资格预审额外要求"},
            ],
        },
        catalog,
    )
    keys = [item.key for item in mapped]
    assert keys.count("id_legal") == 1
    assert "bond" in keys
    assert "invented_other_project" not in keys
    assert all(item.key != "perf" for item in mapped)
    assert any(item.title == "安全生产许可证" and not item.key for item in missing)

    report = build_attachment_match(
        [
            {"key": "id_legal", "title": "法定代表人身份证正反面", "fileCount": 2},
            {"key": "bond", "title": "投标保证金缴存回单", "fileCount": 0},
            {"key": "mnew1", "title": "安全生产许可证", "fileCount": 0},
        ],
        required_keys=["id_legal", "bond", "mnew1"],
        include_keys=["id_legal", "bond", "mnew1"],
        created_keys=["mnew1"],
    )
    assert [x["key"] for x in report["matched"]] == ["id_legal"]
    assert [x["key"] for x in report["missingFiles"]] == ["bond"]
    assert [x["key"] for x in report["createdItems"]] == ["mnew1"]
    notes = attachment_match_notes(report)
    joined = " ".join(notes)
    assert "已匹配资料库扫描件" in joined
    assert "投标保证金缴存回单" in joined
    assert "不会用其他项目" in joined
    assert "安全生产许可证" in joined


def test_find_catalog_item_does_not_borrow_other_slot() -> None:
    from api.services.tenders.match import find_catalog_item
    from api.services.tenders.schema import PlaceholderItem

    catalog = [
        PlaceholderItem(key="id_legal", title="法定代表人身份证正反面"),
        PlaceholderItem(key="perf", title="类似项目合同及发票"),
        PlaceholderItem(key="bond", title="投标保证金缴存回单"),
    ]
    assert find_catalog_item(catalog, key="invented_other_project", title="") is None
    assert find_catalog_item(catalog, title="安全生产许可证") is None
    hit = find_catalog_item(catalog, title="法定代表人身份证")
    assert hit is not None and hit.key == "id_legal"
    borrowed = find_catalog_item(catalog, key="id_legal", title="营业执照")
    assert borrowed is None
    remapped = find_catalog_item(catalog, key="id_legal", title="类似项目合同")
    assert remapped is not None and remapped.key == "perf"


def test_find_catalog_item_maps_credit_screenshot_aliases() -> None:
    from api.services.tenders.match import find_catalog_item, split_invitation_materials
    from api.services.tenders.schema import PlaceholderItem

    catalog = [
        PlaceholderItem(
            key="credit",
            title="信用截图",
            hint="信用中国、政府采购网失信查询截图",
        ),
        PlaceholderItem(key="id_legal", title="法定代表人身份证正反面"),
        PlaceholderItem(key="bond", title="投标保证金缴存回单"),
    ]
    hit = find_catalog_item(
        catalog,
        title="信用中国、政府采购网、国家企业信用信息公示系统查询截图",
        hint="资格要求第（3）条强制要求提供三网站无不良记录截图",
    )
    assert hit is not None and hit.key == "credit"
    assert find_catalog_item(catalog, title="安全生产许可证") is None
    assert find_catalog_item(catalog, title="各类承诺书") is None
    mapped, missing = split_invitation_materials(
        {
            "missingMaterials": [
                {
                    "title": "信用中国、政府采购网、国家企业信用信息公示系统查询截图",
                    "reason": "资格要求第（3）条强制要求提供三网站无不良记录截图",
                }
            ],
        },
        catalog,
    )
    assert [item.key for item in mapped] == ["credit"]
    assert missing == []


def test_find_catalog_item_maps_default_slot_aliases() -> None:
    from api.services.tenders.match import find_catalog_item, split_invitation_materials
    from api.services.tenders.placeholders import DEFAULT_SLOTS
    from api.services.tenders.schema import PlaceholderItem

    catalog = [
        *DEFAULT_SLOTS,
        PlaceholderItem(key="commit", title="各类承诺书"),
        PlaceholderItem(key="iso_cert", title="ISO体系证书"),
        PlaceholderItem(key="safety", title="安全生产许可证"),
        PlaceholderItem(key="license", title="营业执照"),
    ]

    def key_of(title: str) -> str:
        hit = find_catalog_item(catalog, title=title)
        return hit.key if hit is not None else ""

    assert key_of("法人代表身份证复印件") == "id_legal"
    assert key_of("委托代理人身份证正反面") == "id_agent"
    assert key_of("近三年类似业绩合同") == "perf"
    assert key_of("财务报表") == "finance"
    assert key_of("近半年社保缴费证明") in {"finance", "id_agent"}
    assert key_of("型式试验报告") == "product"
    assert key_of("3C认证证书") == "product"
    assert key_of("投标承诺书") == "commit"
    assert key_of("ISO9001质量管理体系认证证书") in {"iso", "iso_cert"}
    assert key_of("企业法人营业执照副本") == "license"
    assert key_of("法定代表人身份证明") == ""
    assert key_of("基本账户开户许可证") == "bank_permit"
    mapped, missing = split_invitation_materials(
        {
            "missingMaterials": [
                {"title": "近三年类似业绩合同及发票", "reason": "资格要求"},
                {"title": "型式试验报告", "reason": "产品检测"},
            ]
        },
        list(DEFAULT_SLOTS),
    )
    keys = [item.key for item in mapped]
    assert "perf" in keys
    assert "product" in keys
    assert missing == []


def test_library_slug_and_prompt_lines() -> None:
    from api.services.tenders.library_kb import catalog_prompt_lines, key_from_tags, slug_key

    assert slug_key("ISO 9001", "abc123") == "iso_9001"
    assert slug_key("安全生产许可证", "abc123def456").startswith("m")
    assert key_from_tags(["投标资料", "slot:id_legal"]) == "id_legal"
    text = catalog_prompt_lines(
        [{"key": "iso_cert", "title": "ISO体系证书", "fileCount": 1, "hint": "现行有效"}]
    )
    assert "iso_cert" in text
    assert "已有扫描件" in text


def test_library_file_kind() -> None:
    from api.services.tenders.library_kb import library_file_kind
    from api.middleware.auth import _query_token_allowed

    assert library_file_kind("png") == "image"
    assert library_file_kind(".JPG") == "image"
    assert library_file_kind("pdf") == "pdf"
    assert library_file_kind("docx") == "file"
    assert _query_token_allowed("/api/v1/tenders/library/files/abc123") is True
    assert _query_token_allowed("/api/v1/tenders/library") is False
    assert _query_token_allowed("/api/v1/knowledge/documents/abc123/file") is True
    assert _query_token_allowed("/api/v1/knowledge/documents/abc123/download") is False


def test_delete_library_file_rejects_empty_id() -> None:
    import asyncio

    from api.services.tenders.library_kb import delete_library_file
    from common.errors import AppError

    async def run() -> None:
        try:
            await delete_library_file(None, "  ")  # type: ignore[arg-type]
        except AppError as exc:
            assert exc.status_code == 422
            return
        raise AssertionError("expected AppError")

    asyncio.run(run())


def test_legacy_copy_skips_same_content_hash(tmp_path) -> None:
    from api.services.tenders.library_kb import _already_has_file

    src = tmp_path / "图片20.png"
    src.write_bytes(b"credit-scan")
    digest = hashlib.sha256(b"credit-scan").hexdigest()
    assert _already_has_file({digest}, src) is True
    assert _already_has_file(set(), src) is False
    other = tmp_path / "other.png"
    other.write_bytes(b"different")
    assert _already_has_file({digest}, other) is False


def test_unlink_legacy_slot_copies_by_hash(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    from api.services.tenders import library_kb

    gone = tmp_path / "Snipaste_2026-08-27_10-45-26.png"
    gone.write_bytes(b"snip-bytes")
    keep = tmp_path / "keep.png"
    keep.write_bytes(b"keep-me")
    digest = hashlib.sha256(b"snip-bytes").hexdigest()
    monkeypatch.setattr(
        library_kb, "list_slot_files", lambda key: [gone, keep] if key == "perf" else []
    )
    monkeypatch.setattr(library_kb, "_doc_file_path", lambda _doc: None)
    doc = SimpleNamespace(
        content_hash=digest,
        name="Snipaste_2026-08-27_10-45-26",
        file_key=None,
        storage_path=None,
    )
    library_kb._unlink_legacy_slot_copies("perf", doc)
    assert not gone.exists()
    assert keep.exists()


def test_legacy_migrate_marker_skips_when_present(tmp_path, monkeypatch) -> None:
    import asyncio

    from api.services.tenders import library_kb

    marker = tmp_path / ".kb_migrated"
    marker.write_text("1", encoding="utf-8")
    monkeypatch.setattr(library_kb, "_legacy_migrate_marker", lambda: marker)
    called = {"n": 0}

    async def boom(_db, _base_id):
        called["n"] += 1
        return []

    monkeypatch.setattr(library_kb, "_list_docs", boom)

    async def run() -> None:
        await library_kb._migrate_legacy_files(None, None)  # type: ignore[arg-type]

    asyncio.run(run())
    assert called["n"] == 0


def test_infer_factory_role() -> None:
    from api.services.tenders.document import infer_factory_role

    assert infer_factory_role(BidBrief(factoryRole="供货单位")) == "供货单位"
    assert infer_factory_role(BidBrief(projectName="充电桩采购")) == "充电设备生产厂商"
    assert infer_factory_role(BidBrief(projectName="备件采购")) == "投标产品生产厂商"


def test_deviation_follows_custom_quote_not_charger_template(tmp_path) -> None:
    from api.services.tenders.schema import QuoteLineIn

    brief = BidBrief(
        projectName="备件采购",
        tenderer="某电厂",
        bidPriceYuan=80000,
        legalPersonName="张三",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
        quoteLines=[
            QuoteLineIn(
                seq="1",
                name="高压开关柜",
                spec="10kV",
                unit="面",
                qty=2,
                unitPrice=30000,
                amount=60000,
            ),
            QuoteLineIn(
                seq="2",
                name="电缆",
                spec="YJV",
                unit="米",
                qty=100,
                unitPrice=200,
                amount=20000,
            ),
        ],
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document

    doc = Document(str(path))
    deviation = next(
        t
        for t in doc.tables
        if len(t.rows[0].cells) >= 4 and "招标文件要求" in t.rows[0].cells[1].text
    )
    body = "\n".join(cell.text for row in deviation.rows for cell in row.cells)
    assert "高压开关柜" in body
    assert "电缆" in body
    assert "7kW" not in body
    assert "30kW" not in body


def test_perf_table_writes_project_name(tmp_path) -> None:
    from api.services.tenders.schema import PerformanceLine

    brief = BidBrief(
        projectName="备件采购投标项目",
        tenderer="招标人",
        bidPriceYuan=100000,
        legalPersonName="张三",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
        performanceLines=[
            PerformanceLine(
                projectName="某电厂充电桩供货合同",
                spec="7kW交流桩",
                client="某电厂",
                amountYuan=500000,
                summary="已验收",
                note="合同扫描件",
            ),
            PerformanceLine(
                projectName="在建变电站项目",
                ongoing=True,
                client="某局",
                amountYuan=200000,
            ),
        ],
    )
    path, _warnings = build_bid_docx(brief, tmp_path / "bid.docx", qualification_pdf=None)
    from docx import Document

    doc = Document(str(path))
    perf_tables = [
        table
        for table in doc.tables
        if table.rows
        and "项目名称" in (table.rows[0].cells[0].text or "")
        and len(table.rows[0].cells) < 3
    ]
    assert len(perf_tables) >= 2
    done_text = "\n".join(cell.text for row in perf_tables[0].rows for cell in row.cells)
    doing_text = "\n".join(cell.text for row in perf_tables[1].rows for cell in row.cells)
    assert "某电厂充电桩供货合同" in done_text
    assert "50万元" in done_text
    assert "在建变电站项目" in doing_text


def test_deviation_lines_from_payload() -> None:
    from api.services.tenders.extract import deviation_lines_from_payload

    lines = deviation_lines_from_payload(
        {
            "deviationLines": [
                {"requirement": "绝缘电阻", "response": "符合", "deviation": "无偏差"},
                {"req": "防护等级"},
                {"requirement": ""},
            ]
        }
    )
    assert len(lines) == 2
    assert lines[0].requirement == "绝缘电阻"
    assert lines[1].requirement == "防护等级"
    assert lines[1].response == "防护等级"


def test_merge_performance_lines_dedupes() -> None:
    from api.services.tenders.extract import merge_performance_lines
    from api.services.tenders.schema import PerformanceLine

    primary = [PerformanceLine(projectName="甲电厂合同")]
    extra = [
        PerformanceLine(projectName="甲 电厂合同"),
        PerformanceLine(projectName="乙电厂合同"),
    ]
    merged = merge_performance_lines(primary, extra)
    assert [item.projectName for item in merged] == ["甲电厂合同", "乙电厂合同"]


def test_amount_from_library_filename() -> None:
    from api.services.tenders.performance import amount_from_text

    assert amount_from_text("某某合同 80万元.pdf") == 800000.0
    assert amount_from_text("发票 5000元") == 5000.0
    assert amount_from_text("无金额") == 0.0


def test_performance_extract_skips_screenshot_name() -> None:
    from api.services.tenders.performance import is_weak_title, performance_from_dict

    assert is_weak_title("Snipaste_2026-08-27_1")
    assert performance_from_dict({"projectName": "Snipaste_2026-08-27_1"}, allow_client_title=False) is None
    line = performance_from_dict(
        {"projectName": "Snipaste_2026-08-27_1", "client": "某供电公司"},
        allow_client_title=True,
    )
    assert line is not None
    assert line.projectName == "某供电公司项目合同"


def test_rank_performance_prefers_completed_charger() -> None:
    from api.services.tenders.performance import rank_performance_lines
    from api.services.tenders.schema import PerformanceLine

    ranked = rank_performance_lines(
        [
            PerformanceLine(projectName="在建充电桩", ongoing=True, chargerRelated=True, amountYuan=900000),
            PerformanceLine(projectName="已完成电缆工程", ongoing=False, amountYuan=500000),
            PerformanceLine(projectName="已完成充电桩", ongoing=False, chargerRelated=True, amountYuan=200000),
            PerformanceLine(projectName="已完成充电桩大额", ongoing=False, chargerRelated=True, amountYuan=800000),
        ]
    )
    assert [item.projectName for item in ranked] == [
        "已完成充电桩大额",
        "已完成电缆工程",
        "已完成充电桩",
        "在建充电桩",
    ]


def test_merge_performance_drops_screenshot_names() -> None:
    from api.services.tenders.extract import merge_performance_lines
    from api.services.tenders.schema import PerformanceLine

    merged = merge_performance_lines(
        [PerformanceLine(projectName="Snipaste_2026-08-27_1")],
        [
            PerformanceLine(
                projectName="某充电站EPC合同",
                client="某供电公司",
                spec="120kW直流桩",
                amountYuan=800000,
                chargerRelated=True,
            )
        ],
    )
    assert [item.projectName for item in merged] == ["某充电站EPC合同"]


def test_fallback_performance_from_contract_text() -> None:
    from api.services.tenders.performance import fallback_from_text, dump_perf_meta, load_perf_meta

    line = fallback_from_text(
        "工程名称：城东充电站设备供货\n甲方：某供电公司\n合同额：80万元\n已竣工验收，直流桩 10 台 120kW",
        "Snipaste_2026-08-27_1.png",
    )
    assert line is not None
    assert line.projectName == "城东充电站设备供货"
    assert line.client == "某供电公司"
    assert line.amountYuan == 800000.0
    assert line.ongoing is False
    assert line.chargerRelated
    assert line.spec.lower() == "120kw"
    restored = load_perf_meta(dump_perf_meta(line))
    assert restored is not None
    assert restored.projectName == line.projectName


def test_ordered_perf_lines_completed_first() -> None:
    from api.services.tenders.document import _ordered_perf_lines
    from api.services.tenders.schema import BidBrief, PerformanceLine

    brief = BidBrief(
        performanceLines=[
            PerformanceLine(projectName="在建充电桩A", ongoing=True, chargerRelated=True, amountYuan=100),
            PerformanceLine(projectName="已完成充电桩B", ongoing=False, chargerRelated=True, amountYuan=50),
            PerformanceLine(projectName="Snipaste_xxx", ongoing=False),
        ]
    )
    ordered = _ordered_perf_lines(brief)
    assert [item.projectName for item in ordered] == ["已完成充电桩B", "在建充电桩A"]


def test_ordered_perf_lines_skips_unchecked() -> None:
    from api.services.tenders.document import _ordered_perf_lines
    from api.services.tenders.schema import BidBrief, PerformanceLine

    brief = BidBrief(
        performanceLines=[
            PerformanceLine(projectName="写入充电桩", chargerRelated=True, includeInBid=True),
            PerformanceLine(projectName="不写充电桩", chargerRelated=True, includeInBid=False),
        ]
    )
    ordered = _ordered_perf_lines(brief)
    assert [item.projectName for item in ordered] == ["写入充电桩"]


def test_technical_soft_issues_do_not_block() -> None:
    from api.services.tenders.categories import technical_soft_issues
    from api.services.tenders.schema import BidBrief

    notes = technical_soft_issues(BidBrief(deliveryDays=30), media={})
    assert any("技术偏差表" in item for item in notes)
    assert any("文字描述" in item for item in notes)
    assert any("图纸" in item for item in notes)
    assert any("类似业绩" in item for item in notes)


def test_technical_section_writes_construction_plan(tmp_path) -> None:
    from api.services.tenders.placeholders import TECH_DRAWING_KEY

    brief = BidBrief(
        projectName="充电站采购安装",
        tenderer="招标人",
        bidPriceYuan=100000,
        legalPersonName="张三",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
        constructionPlan="先做基础，再安装直流桩，最后调试送电。",
        layoutPlan="箱变靠场地北侧，充电区南向布置。",
        quoteLines=_sample_quote_lines(),
    )
    path, warnings = build_bid_docx(
        brief,
        tmp_path / "tech.docx",
        qualification_pdf=None,
        catalog_media={TECH_DRAWING_KEY: []},
    )
    from docx import Document

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "技术标（实施方案）" in text
    assert "一、文字描述" in text
    assert "二、图纸" in text
    assert "先做基础，再安装直流桩" in text
    assert "箱变靠场地北侧" in text
    assert "（在此粘贴扫描件）" in table_text
    assert any('w:val="dashed"' in t._tbl.xml for t in doc.tables)
    assert any("虚线框" in w for w in warnings)

    # 技术标必须写在正文「其他材料」位置，不能插进目录页
    texts = [(p.text or "").strip() for p in doc.paragraphs]
    toc_i = next(i for i, t in enumerate(texts) if t.replace("\t", "") == "目录" or t.startswith("目"))
    letter_i = next(i for i, t in enumerate(texts) if t == "投标函及投标函附录")
    tech_i = next(i for i, t in enumerate(texts) if t == "技术标（实施方案）")
    assert toc_i < letter_i < tech_i
    assert "先做基础，再安装直流桩" in texts[tech_i + 1] or any(
        "先做基础" in t for t in texts[tech_i:tech_i + 8]
    )


def test_technical_section_embeds_uploaded_drawings(tmp_path) -> None:
    from PIL import Image

    from api.services.tenders.placeholders import TECH_DRAWING_KEY

    png = tmp_path / "site.png"
    Image.new("RGB", (160, 90), (40, 120, 200)).save(png, "PNG")
    brief = BidBrief(
        projectName="充电站采购安装",
        tenderer="招标人",
        bidPriceYuan=100000,
        legalPersonName="张三",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=False,
        constructionPlan="先做基础，再安装直流桩。",
        quoteLines=_sample_quote_lines(),
    )
    path, warnings = build_bid_docx(
        brief,
        tmp_path / "tech.docx",
        qualification_pdf=None,
        catalog_media={TECH_DRAWING_KEY: [png]},
    )
    from docx import Document

    doc = Document(str(path))
    blob = doc.element.xml
    assert "a:blip" in blob
    assert any("图纸写入" in w for w in warnings)
    assert not any("为加快生成未写入" in w for w in warnings)


def test_toc_owns_own_page_and_pageref_matches_chapters(tmp_path) -> None:
    """目录独占一页；封面/目录不编页码，正文页脚从 1 起。"""
    from docx import Document
    from docx.oxml.ns import qn

    from api.services.tenders.document import (
        _estimate_bookmark_pages,
        _first_body_section_index,
        _p_has_page_br,
        _pageref_bookmark,
        _sect_starts_new_page,
    )

    brief = BidBrief(
        projectName="分页一致性项目",
        tenderer="测试招标人",
        bidPriceYuan=200000,
        legalPersonName="李四",
        attachQualifications=False,
        includePlaceholders=False,
        includeCommitment=True,
        quoteLines=_sample_quote_lines(),
        deviationLines=[
            DeviationLine(seq="1", requirement="要求", response="响应", deviation="无偏差")
        ],
    )
    path, _ = build_bid_docx(brief, tmp_path / "toc_page.docx", qualification_pdf=None)
    doc = Document(str(path))

    toc_title = next(p for p in doc.paragraphs if (p.text or "").replace("\t", "").strip() == "目录")
    # 目录前应有显式分页符（空段 w:br type=page）
    prev = toc_title._p.getprevious()
    assert prev is not None and prev.tag == qn("w:p")
    assert _p_has_page_br(prev)

    letter = next(p for p in doc.paragraphs if (p.text or "").strip() == "投标函及投标函附录")
    prev_letter = letter._p.getprevious()
    assert prev_letter is not None
    # 目录与投标函之间：显式分页或空段分页，保证目录独占一页
    found_br = False
    el = letter._p.getprevious()
    while el is not None:
        if el.tag == qn("w:p"):
            if _p_has_page_br(el):
                found_br = True
                break
            text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
            if text == "目录" or text.startswith("一、") or text.startswith("九、"):
                break
            if text:
                break
        el = el.getprevious()
    assert found_br

    # 节间均为 continuous，换页由显式分页符控制
    for section in doc.sections:
        types = [el.get(qn("w:val")) for el in section._sectPr.findall(qn("w:type"))]
        assert types == ["continuous"]
        assert not _sect_starts_new_page(section._sectPr)

    assert any(_pageref_bookmark(p) for p in doc.paragraphs)
    bookmarks = {
        el.get(qn("w:name")) for el in doc.element.iter(qn("w:bookmarkStart"))
    }
    assert "toc_letter" in bookmarks
    assert "toc_other" in bookmarks
    assert "toc_commit" in bookmarks

    # 目录条目与其后第一个正文标题之间不应夹技术标正文
    texts = [(p.text or "").strip() for p in doc.paragraphs]
    toc_i = next(i for i, t in enumerate(texts) if t.replace("\t", "") == "目录")
    letter_i = next(i for i, t in enumerate(texts) if t == "投标函及投标函附录")
    between = texts[toc_i + 1 : letter_i]
    assert not any("先做基础" in t or t.startswith("一、文字描述") for t in between)
    assert any("技术标（实施方案）" in t or "八、技术标" in t for t in between)

    pages = _estimate_bookmark_pages(doc)
    # 封面=1、目录=2、投标函从目录下一页起
    assert pages["toc_letter"] == 3
    body_si = _first_body_section_index(doc)
    assert body_si >= 2
    for i, section in enumerate(doc.sections):
        xml = section.footer.paragraphs[0]._p.xml if section.footer.paragraphs else ""
        if i < body_si:
            assert "PAGE" not in xml
        else:
            assert "PAGE" in xml
    assert pages["toc_other"] > pages["toc_letter"]
    assert pages["toc_commit"] > pages["toc_other"]
    assert pages["toc_quote"] > pages["toc_letter"]
    # 附录 / 支付条件另起一页，不挤在上一节签字盖章区
    appendix = next(p for p in doc.paragraphs if (p.text or "").strip() == "投标函附录")
    pay = next(p for p in doc.paragraphs if (p.text or "").strip() == "支付条件")
    prev_app = appendix._p.getprevious()
    prev_pay = pay._p.getprevious()
    assert prev_app is not None and _p_has_page_br(prev_app)
    assert prev_pay is not None and _p_has_page_br(prev_pay)



def test_sanitize_history_brief_strips_project_and_qty() -> None:
    from api.services.tenders.history import (
        contains_foreign_bidder,
        playbooks_prompt_block,
        sanitize_history_brief,
    )
    from api.services.tenders.schema import PerformanceLine

    book = sanitize_history_brief(
        BidBrief(
            projectName="高途智成港",
            tenderer="郑州有爱文化科技有限公司",
            bidderName="河南伟泰光电科技有限公司",
            bidPriceYuan=518800,
            factoryRole="充电设备生产厂商",
            requiredSlotKeys=["id_legal", "bond", "perf"],
            constructionPlan="先基础后安装",
            quoteLines=[QuoteLineIn(name="交流桩", qty=133, unit="台")],
            performanceLines=[PerformanceLine(projectName="旧合同", amountYuan=800000)],
        )
    )
    assert book is not None
    assert book.slot_keys == ("id_legal", "bond", "perf")
    assert book.factory_role == "充电设备生产厂商"
    assert "实施方案说明" in book.tech_written
    text = playbooks_prompt_block([book])
    assert "133" not in text
    assert "518800" not in text
    assert "高途智成港" not in text
    assert "旧合同" not in text
    assert "quoteLines 必须为 []" in text
    assert "id_legal" in text
    assert contains_foreign_bidder("摘录：容新新能源投标文件")
    assert (
        sanitize_history_brief(
            BidBrief(bidderName="郑州容新新能源有限公司", requiredSlotKeys=["bond"])
        )
        is None
    )


def test_user_prompt_drops_foreign_kb_and_history_qty() -> None:
    from api.services.tenders.extract import _user_prompt
    from api.services.tenders.history import sanitize_history_brief

    book = sanitize_history_brief(
        BidBrief(
            bidderName="河南伟泰光电科技有限公司",
            requiredSlotKeys=["perf"],
            quoteLines=[QuoteLineIn(name="桩", qty=133, unit="台")],
        )
    )
    prompt = _user_prompt(
        "本标采购交流充电桩，详见工程量清单。",
        [{"name": "容新新能源投标书", "content": "郑州容新报价 133 台"}],
        None,
        [book] if book else [],
    )
    assert "郑州容新" not in prompt
    assert "容新新能源投标书" not in prompt
    assert "133" not in prompt
    assert "perf" in prompt
    assert "禁止抄项目名" in prompt
    bid_prompt = _user_prompt(
        "本标采购交流充电桩。",
        [{"name": "容新商务标投标文件.docx", "content": "资格要求见前附表，投标人已盖章"}],
        None,
        [],
    )
    assert "容新商务标投标文件" not in bid_prompt
    assert "投标人已盖章" not in bid_prompt


def test_merge_quote_ignores_llm_patch_and_clears_old_qty(tmp_path) -> None:
    from api.services.tenders.extract import _merge_quote

    invite = tmp_path / "invite.txt"
    invite.write_text("某充电站设备采购，正文没有工程量表格。", encoding="utf-8")
    current = BidBrief(
        projectName="新项目",
        quoteLines=[QuoteLineIn(name="交流桩", qty=133, unit="台")],
    )
    patch = {
        "quoteTitle": "历史清单",
        "quoteLines": [{"name": "直流桩", "qty": 56, "unit": "台"}],
    }
    notes, brief, filled = _merge_quote(
        current,
        patch,
        "某充电站设备采购，正文没有工程量表格。",
        invite,
        None,
    )
    assert filled is False
    assert brief.quoteLines == []
    joined = " ".join(notes)
    assert "未采用模型或历史" in joined
    assert "已清空上次表单里的分项" in joined
    assert all(row.qty != 133 for row in brief.quoteLines)
    assert all(row.qty != 56 for row in brief.quoteLines)
