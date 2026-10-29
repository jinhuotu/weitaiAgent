"""投标 Word 中缺失扫描件的小标题 + 方框占位。程序排版，不经 LLM。"""

from __future__ import annotations

import io
import hashlib
import logging
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from api.services.tenders.schema import PlaceholderItem

logger = logging.getLogger("api.tenders")

_SONG = "宋体"
_CN_ORD = "一二三四五六七八九十"
_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}

# 证件类写入 Word；合同/财税等 PDF 仍占位，避免生成卡住。
_EMBED_SCAN_KEYS = frozenset({"id_legal", "id_agent"})
_EMBED_PDF_KEYS = frozenset({"id_legal", "id_agent", "credit"})
_SLOT_PDF_MAX_PAGES = 1
_CREDIT_PDF_MAX_PAGES = 4
_SLOT_MAX_FILES = 2
_CREDIT_MAX_FILES = 6
_PLACEHOLDER_MAX_IMAGES = 8
_RENDER_SCALE = 0.5
_JPEG_QUALITY = 45
_IMAGE_MAX_PX = 800
# 信用中国等文字页按约 220dpi 出 PNG，不再走 JPEG 45/800px
_CREDIT_DPI = 220
_CREDIT_MAX_PX = 2800
_CREDIT_MAX_H_CM = 12.0
# 身份证正反面：按实物尺寸上限并排，窄版心等比缩小
_ID_PAIR_MAX_IMAGES = 2
_ID_STRIP_WIDTH_CM = 17.4
_ID_SINGLE_WIDTH_CM = 8.6
_ID_STRIP_MAX_H_CM = 5.4
_ID_PREVIEW_DPI = 96
_ID_ASPECT_MIN = 1.22
_ID_ASPECT_MAX = 1.98
_HEADING_CM = 1.1
_BOX_CM = 5.4
_SLOT_GAP_STYLE = "TenderAttachmentGap"
_SLOT_CLEARANCE_STYLE = "TenderTableClearance"
_SLOT_GAP_CM = 0.5
_IMAGE_MAX_H_CM = 12.0
_TWIPS_PER_CM = 567
_EMU_PER_CM = 360000.0

DEFAULT_SLOTS: tuple[PlaceholderItem, ...] = (
    PlaceholderItem(
        key="id_legal",
        title="法定代表人身份证正反面",
        hint="可放多位法定代表人。每人上传正面和反面，正面识别姓名和证号；做标书时再选本标用谁。",
    ),
    PlaceholderItem(
        key="id_agent",
        title="授权代理人身份证及近半年社保缴纳证明",
        hint="法人亲自投标可划掉本项。须为本人证件。",
    ),
    PlaceholderItem(
        key="license",
        title="营业执照",
        hint="企业法人营业执照副本扫描件，不要用身份证或其他证件顶替。",
    ),
    PlaceholderItem(
        key="bank_permit",
        title="开户许可证",
        hint="基本账户开户许可证扫描件。",
    ),
    PlaceholderItem(
        key="iso",
        title="ISO/质量管理体系认证证书",
        hint="按邀请书资质标题一对一提供，不要拿其他证书顶。",
    ),
    PlaceholderItem(
        key="perf",
        title="类似项目合同及发票",
        hint="签约主体必须是河南伟泰光电科技有限公司。上传后识别项目/买方/合同额；生成本标时按邀请书业绩门槛（行业、MES/MOM 等）筛选，不再默认充电桩。",
    ),
    PlaceholderItem(
        key="finance",
        title="近三年财务审计报告 / 完税证明",
        hint="按招标文件资格审查条款提供。纳税完税放本项，不要和社保证明混放。",
    ),
    PlaceholderItem(
        key="social",
        title="社保缴纳证明",
        hint="近半年社保缴纳证明。与纳税完税、身份证分开上传。",
    ),
    PlaceholderItem(
        key="credit",
        title="信用中国查询页及国家企业信用信息公示（股东）",
        hint="截图须含查询日期。",
    ),
    PlaceholderItem(
        key="product",
        title="所投产品检测报告 / 3C / 对应功率桩型证明",
        hint="功率与招标要求一致；正偏离须另附厂家盖章说明。",
    ),
    PlaceholderItem(
        key="bond",
        title="投标保证金缴存回单",
        hint="按邀请书载明的金额与账户缴纳后替换。",
    ),
    PlaceholderItem(
        key="seal",
        title="投标文件签章页",
        hint="逐页盖投标人公章；法定代表人签字或盖章。",
    ),
)

CRITICAL_SLOT_KEYS = frozenset(item.key for item in DEFAULT_SLOTS if item.key != 'product')

# Evidence comes from this invitation, never from the library's own descriptions.
_MATERIAL_SECTION_MARKS = (
    ('id_legal', ('法定代表人身份证', '法人身份证', '法人代表身份证')),
    ('id_agent', ('代理人身份证', '被授权人身份证', '委托人身份证')),
    ('license', ('营业执照', '独立承担民事责任', '独立法人资格')),
    ('bank_permit', ('开户许可', '基本账户', '基本存款账户', '基本户')),
    ('finance', (
        '财务状况', '财务审计', '审计报告', '财务报告', '财务报表', '财务会计制度',
        '完税证明', '纳税证明', '纳税凭证', '依法缴纳税收', '税收及社会保障',
    )),
    ('social', ('社保缴纳', '社保缴费', '社保证明', '社会保障资金')),
    ('credit', ('信用中国', '信用记录', '信用查询', '失信查询', '企业信用信息公示')),
    ('perf', ('业绩', '类似项目', '合同及发票', '履约证明')),
    ('iso', ('资质证书', '资质证明', '资质材料', '质量体系', '体系认证', 'iso')),
    ('bond', ('投标保证金', '保证金回单', '保证金缴存', '保证金缴纳')),
    ('seal', ('签章页', '盖章页', '印鉴备案', '印鉴及签字')),
)


def _material_keys(text: str) -> list[str]:
    compact = re.sub(r'\s+', '', text or '').lower()
    keys = []
    for key, marks in _MATERIAL_SECTION_MARKS:
        for mark in marks:
            start = compact.find(mark)
            if start < 0:
                continue
            before, after = compact[max(0, start - 12):start], compact[start + len(mark):]
            if (re.search(r'(?:无需|无须|免交|免缴|不收取|不要求).{0,4}$', before)
                    or re.match(r'[：:]?(?:不收取|免交|免缴|无需提供|无须提供)', after)):
                continue
            keys.append(key)
            break
    keys.extend(_paste_box_keys(compact))
    if compact in {'资质', '资质要求'} and 'iso' not in keys:
        keys.append('iso')
    return keys


def _with_format_paste(invitation_id: str, source: str) -> str:
    invite = (invitation_id or "").strip()
    if not invite:
        return source or ""
    from api.services.tenders.pdf_format import format_paste_source

    extra = format_paste_source(invite)
    if not extra:
        return source or ""
    return ((source or "").rstrip() + "\n" + extra).strip()


def _paste_box_keys(compact: str) -> list[str]:
    if "此处粘贴" not in compact and "在此粘贴" not in compact:
        return []
    keys = []
    if "营业执照" in compact:
        keys.append("license")
    if any(word in compact for word in ("纳税", "完税", "税收", "财务", "审计")):
        keys.append("finance")
    if "社保" in compact or "社会保障" in compact:
        keys.append("social")
    if "身份证" in compact:
        keys.append("id_agent" if ("委托" in compact or "代理" in compact) else "id_legal")
    return keys


def _bid_material_keys(outline_items: list | None, requirements: list | None,
                       source_text: str) -> list[str]:
    keys = []
    for item in outline_items or []:
        if getattr(item, 'skipped', False):
            continue
        text = '\n'.join(str(getattr(item, attr, '') or '') for attr in
                         ('title', 'rawTitle', 'body', 'inviteExcerpt'))
        keys.extend(_material_keys(text))
        if getattr(item, 'kind', '') == 'performance':
            keys.append('perf')
    for req in requirements or []:
        key = str(getattr(req, 'slotKey', '') or '').strip()
        if key:
            keys.append(key)
        if getattr(req, 'category', '') not in {'qualification', 'material'} and not key:
            continue
        keys.extend(_material_keys('\n'.join(str(getattr(req, attr, '') or '') for attr in
                                              ('text', 'targetTitle', 'attachToSection'))))
    for line in (source_text or '').splitlines():
        keys.extend(_material_keys(line))
    return normalize_slot_keys(keys)

TECH_DRAWING_KEY = "tech_drawings"
TECH_DRAWING_SLOT = PlaceholderItem(
    key=TECH_DRAWING_KEY,
    title="本项目图纸",
    hint="本标平面图、系统图或施工图。换一份邀请书需重新上传，不要用其他项目图纸。",
)

_PERF_NOTE = (
    "【待补】请粘贴河南伟泰光电科技有限公司作为签约主体的合同及发票。"
    "按本标邀请书业绩门槛筛选，禁止使用其他公司业绩。"
)


def normalize_slot_keys(keys: list[str] | None) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for raw in keys or []:
        key = str(raw or "").strip()
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(key)
    return out


def this_bid_keys(
    *,
    extra: list[PlaceholderItem] | None = None,
    catalog: list[PlaceholderItem] | None = None,
    required_keys: list[str] | None = None,
    include_keys: list[str] | None = None,
    has_agent: bool = False,
    outline_items: list | None = None,
    requirements: list | None = None,
    source_text: str = '',
    invitation_id: str = '',
) -> tuple[list[str], list[str]]:
    """本标必填 key、写入 Word 的 key。不把整份资料库都塞进文件。"""
    source_text = _with_format_paste(invitation_id, source_text)
    extras = extra or []
    known = {item.key for item in (catalog if catalog is not None else DEFAULT_SLOTS) if item.key}
    known.update(item.key for item in DEFAULT_SLOTS)
    known.update(item.key for item in extras if item.key)

    required = normalize_slot_keys(required_keys)
    if not required:
        required = normalize_slot_keys([item.key for item in extras])
    required = normalize_slot_keys([
        *required, *_bid_material_keys(outline_items, requirements, source_text),
    ])
    if "id_legal" in known and "id_legal" not in required:
        required = ["id_legal", *required]
    if has_agent:
        if "id_agent" in known and "id_agent" not in required:
            required.append("id_agent")
    else:
        required = [key for key in required if key != "id_agent"]
    required = [key for key in required if key in known]

    include = normalize_slot_keys(include_keys)
    if not include:
        include = list(required)
    else:
        for key in required:
            if key not in include:
                include.append(key)
    include = [key for key in include if key in known]
    for item in catalog if catalog is not None else DEFAULT_SLOTS:
        key = (item.key or "").strip()
        if not key or key not in known or key in include:
            continue
        if _always_in_bid_scan(item):
            include.append(key)
    if not has_agent:
        include = [key for key in include if key != "id_agent"]
    return required, include


_ALWAYS_SCAN_MARKS = (
    ("license", ("营业执照",)),
    ("bank_permit", ("开户许可", "开户许可证", "基本账户", "基本户")),
    ("iso", ("ISO", "iso", "质量体系", "体系认证", "资质证书")),
)


def _always_in_bid_scan(item: PlaceholderItem) -> bool:
    key = (item.key or "").strip()
    if key in {"license", "bank_permit", "iso"}:
        return True
    blob = f"{item.title or ''}{item.hint or ''}"
    return any(any(m in blob for m in marks) for _, marks in _ALWAYS_SCAN_MARKS)


def collect_slots(
    extra: list[PlaceholderItem] | None = None,
    *,
    catalog: list[PlaceholderItem] | None = None,
    include_keys: list[str] | None = None,
    outline_items: list | None = None,
) -> list[PlaceholderItem]:
    """资料库清单 + 邀请书多出来的资料项（按 key 去重，后者可改说明）。

    include_keys 为 None 时返回合并后的全部项（资料库列表）；传入列表则只保留本标要写入 Word 的项。
    """
    base_items = list(catalog) if catalog is not None else list(DEFAULT_SLOTS)
    present = {item.key for item in base_items}
    base_items.extend(item for item in DEFAULT_SLOTS if item.key not in present)
    by_key: dict[str, PlaceholderItem] = {
        item.key: item.model_copy() for item in base_items if item.key
    }
    for item in extra or []:
        title = (item.title or "").strip()
        if not title:
            continue
        key = (item.key or "").strip() or f"extra_{len(by_key) + 1}"
        hint = (item.hint or "").strip()
        if key in by_key:
            current = by_key[key]
            by_key[key] = PlaceholderItem(
                key=key,
                title=current.title,
                hint=hint or current.hint,
                attachToSection=(item.attachToSection or current.attachToSection or ""),
            )
        else:
            by_key[key] = PlaceholderItem(
                key=key,
                title=title,
                hint=hint or "请按招标文件要求补附",
                attachToSection=item.attachToSection or "",
            )
    # Supply attachment routing metadata to the existing placement functions.
    # A broad qualification heading can own an explicit business-license request.
    for slot in by_key.values():
        if slot.attachToSection:
            continue
        candidates = [item for item in outline_items or [] if not getattr(item, 'skipped', False)]
        # Prefer a dedicated chapter over an incidental mention in a letter.
        for attrs in (('title', 'rawTitle'), ('body', 'inviteExcerpt')):
            hit = next((item for item in candidates
                        if (attrs[0] == 'title' or getattr(item, 'kind', '')
                            not in {'letter', 'legal_id', 'auth', 'quote'})
                        and slot.key in _material_keys('\n'.join(
                            str(getattr(item, attr, '') or '') for attr in attrs))), None)
            if hit is not None:
                slot.attachToSection = getattr(hit, 'rawTitle', '') or hit.title
                break
    if include_keys is None:
        return list(by_key.values())
    ordered: list[PlaceholderItem] = []
    seen: set[str] = set()
    for key in normalize_slot_keys(include_keys):
        item = by_key.get(key)
        if item is None or key in seen:
            continue
        seen.add(key)
        ordered.append(item)
    return ordered


def _is_credit_slot(slot: PlaceholderItem) -> bool:
    if (slot.key or "").strip() == "credit":
        return True
    n = f"{slot.title or ''}{slot.hint or ''}"
    if "信用" not in n:
        return False
    return any(k in n for k in ("截图", "查询", "公示", "信用中国", "失信"))


def _is_embed_pdf_slot(slot: PlaceholderItem) -> bool:
    key = (slot.key or "").strip()
    if key in _EMBED_PDF_KEYS or "身份证" in (slot.title or ""):
        return True
    return _is_credit_slot(slot)


def _merge_credit_slots(
    slots: list[PlaceholderItem],
    files_by_key: dict[str, list[Path]] | None,
) -> tuple[list[PlaceholderItem], dict[str, list[Path]]]:
    files_by_key = {k: list(v) for k, v in (files_by_key or {}).items()}
    credit = [s for s in slots if _is_credit_slot(s)]
    if len(credit) <= 1:
        return slots, files_by_key
    keep = next((s for s in credit if (s.key or "").strip() == "credit"), credit[0])
    merged: list[Path] = []
    seen: set[str] = set()
    for s in credit:
        for p in files_by_key.get(s.key) or []:
            mark = str(p)
            if mark in seen:
                continue
            seen.add(mark)
            merged.append(p)
    files_by_key[keep.key] = merged
    out: list[PlaceholderItem] = []
    used = False
    for s in slots:
        if not _is_credit_slot(s):
            out.append(s)
            continue
        if used:
            continue
        out.append(keep)
        used = True
    return out, files_by_key


def id_slot(key: str) -> PlaceholderItem:
    for item in DEFAULT_SLOTS:
        if item.key == key:
            return item
    title = "法定代表人身份证正反面" if key == "id_legal" else "授权代理人身份证正反面"
    return PlaceholderItem(key=key, title=title)


def inline_id_scans(
    doc: Document,
    files: list[Path] | None,
    *,
    empty_slot: PlaceholderItem | None = None,
    before: Paragraph | None = None,
) -> int:
    """把身份证正反面贴到当前页；before 有值时挪到该段之前。返回嵌入图数，0 为虚线框。"""
    last = _last_body_child(doc)
    media = [p for p in (files or []) if p.is_file()]
    n = 0
    if media:
        n = _insert_slot_media(doc, media, compact_pair=True)
    if n <= 0:
        _draw_box(doc, empty_slot or PlaceholderItem(key="id_scan", title="身份证扫描件"))
        n = 0
    if before is not None:
        _relocate_appended_before(doc, before=before, last_before=last)
    _merge_slot_gaps(doc)
    return n


def append_placeholder_section(
    doc: Document,
    slots: list[PlaceholderItem],
    attachments: dict[str, list[Path]] | None = None,
    *,
    before: Paragraph | None = None,
    heading: bool = True,
) -> tuple[int, int, list[str]]:
    """每个附件项：小标题 +（轻量扫描件 | 虚线框说明）。
    证件/图片写入 Word；财税/合同等 PDF 默认不渲进，避免生成卡住。
    before 为「投标承诺书」等标题时，整块插到该段之前（技术标之后）。
    返回 (已插入项数, 仍待补方框数, 限流提示)。
    """
    if not slots:
        return 0, 0, []
    slots = [item for item in slots if (item.key or "").strip() != TECH_DRAWING_KEY]
    if not slots:
        return 0, 0, []
    slots, files_by_key = _merge_credit_slots(slots, attachments)
    last_before = _last_body_child(doc)
    if _used_on_page_cm(doc) > 1.2:
        doc.add_page_break()
    if heading:
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(8)
        title.paragraph_format.space_after = Pt(8)
        title.paragraph_format.keep_with_next = True
        head = title.add_run("附件：资料库扫描件")
        _font(head, 16, bold=True)
    filled = 0
    boxes = 0
    notes: list[str] = []
    skipped_heavy = 0
    budget = {"left": _PLACEHOLDER_MAX_IMAGES}
    for i, slot in enumerate(slots):
        key = (slot.key or "").strip()
        media = files_by_key.get(key) or []
        images = [p for p in media if p.suffix.lower() in _IMAGE_EXT]
        is_id_scan = key in _EMBED_SCAN_KEYS or "身份证" in (slot.title or "")
        embed_pdf = _is_embed_pdf_slot(slot)
        embed_files = list(media) if (is_id_scan or embed_pdf) else images
        if not media:
            need_cm = _HEADING_CM + _BOX_CM
        elif is_id_scan:
            need_cm = _HEADING_CM + 6.4
        elif embed_files:
            max_h = _CREDIT_MAX_H_CM if (embed_pdf and not is_id_scan) else _IMAGE_MAX_H_CM
            need_cm = _HEADING_CM + _first_media_cm(embed_files, max_h=max_h) + 0.4
        else:
            need_cm = _HEADING_CM + _BOX_CM
        _break_before_slot(doc, need_cm)
        _write_subheading(doc, i, slot.title)
        if not media:
            _draw_box(doc, slot)
            boxes += 1
            continue
        if embed_files:
            max_pages = _CREDIT_PDF_MAX_PAGES if embed_pdf and not is_id_scan else _SLOT_PDF_MAX_PAGES
            n = _insert_slot_media(
                doc,
                embed_files[: _CREDIT_MAX_FILES if embed_pdf else _SLOT_MAX_FILES],
                budget=None if not is_id_scan else budget,
                max_pages=max_pages,
                compact_pair=is_id_scan,
                force_landscape=is_id_scan,
                sharp=embed_pdf and not is_id_scan,
            )
            if n > 0:
                filled += 1
                leftover_pdf = (
                    [p for p in media if p.suffix.lower() == ".pdf"] if not (is_id_scan or embed_pdf) else []
                )
                if leftover_pdf:
                    skipped_heavy += 1
                continue
        skipped_heavy += 1
        _draw_box(doc, slot)
        boxes += 1
        filled += 1

    if before is not None:
        _relocate_appended_before(doc, before=before, last_before=last_before)
    _merge_slot_gaps(doc)

    if skipped_heavy:
        notes.append(
            f"已有 {skipped_heavy} 类附件未嵌入 Word（合同/财税等仅占位提示），生成已加速；"
            "装订时请从资料库打印原件附上。证件类扫描件已写入「附件：资料库扫描件」。"
        )
    return filled, boxes, notes


def _last_body_child(doc: Document):
    body = doc.element.body
    last = None
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            break
        last = child
    return last


def _relocate_appended_before(doc: Document, *, before: Paragraph, last_before) -> None:
    body = doc.element.body
    to_move = []
    passed = last_before is None
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if not passed:
            if child is last_before:
                passed = True
            continue
        to_move.append(child)
    for el in to_move:
        before._element.addprevious(el)


_LEVEL1 = re.compile(
    r"^(?:第[一二三四五六七八九十0-9]+[章节]|[一二三四五六七八九十]{1,3}、|\d+[.、])"
)


def _compact_para(text: str) -> str:
    return "".join((text or "").split())


def _level1_core(text: str) -> str:
    n = _compact_para(text)
    if not n or len(n) > 36 or "如有" in n or "。" in n:
        return ""
    if n.startswith(("（", "(")):
        return ""
    if _LEVEL1.match(n) is None:
        return ""
    n = re.sub(r"^第[一二三四五六七八九十0-9]+[章节]", "", n)
    n = re.sub(r"^[一二三四五六七八九十]{1,3}、", "", n)
    n = re.sub(r"^\d+[.、]", "", n)
    return n


def _para_text(el) -> str:
    return "".join(node.text or "" for node in el.iter(qn("w:t")))


_NO_PASTE = {
    "投标函",
    "投标函附录",
    "目录",
    "目次",
    "法定代表人身份证明",
    "法定代表人身份证明书",
    "授权委托书",
    "中小企业声明函",
    "残疾人福利性单位声明函",
}
_CHAPTER_MARKS = (
    "投标保证金",
    "已标价工程量清单",
    "已标价的工程量清单",
    "工程量清单",
    "已标价",
    "营业执照",
    "开户许可",
    "基本账户",
    "基本户",
    "信用中国",
    "财务审计",
    "财务报表",
    "类似项目",
    "资质证明",
    "资质证书",
    "安全生产许可",
    "检测报告",
    "授权委托",
)
_KEY_CHAPTER = {
    "bond": "投标保证金",
    "license": "营业执照",
    "bank_permit": "开户许可",
    "perf": "类似项目",
    "finance": "财务审计",
    "credit": "信用中国",
    "iso": "资质证书",
    "quote": "工程量清单",
}


def _heading_core(text: str) -> str:
    n = _compact_para(text)
    if not n or len(n) > 36 or "如有" in n or "。" in n or n.startswith(("（", "(")):
        return ""
    core = _level1_core(text)
    if core:
        return core
    # 组成清单里的「身份证明或授权委托书」不是专章
    if any(x in n for x in ("或", "；", ";", "，", "须", "提供")):
        return ""
    if len(n) > 18:
        return ""
    for mark in _CHAPTER_MARKS:
        if n == mark or (n.endswith(mark) and len(n) - len(mark) <= 2):
            return n
    return ""


def _letter_block(doc: Document) -> list:
    """投标函标题到下一章之前。这段正文不能插附件。"""
    els = []
    inside = False
    for para in doc.paragraphs:
        n = _compact_para(para.text or "")
        core = _level1_core(_section_label(para._p))
        if core:
            inside = core == "投标函" or core.startswith("投标函")
            if inside:
                els.append(para._p)
            continue
        if n.startswith("投标函") or n.startswith("（一）投标函") or n.startswith("(一)投标函"):
            inside = True
            els.append(para._p)
            continue
        if inside and _is_other_title(para.text or ""):
            inside = False
            continue
        if inside and _heading_core(para.text or "") and not _paste_ban(para.text or ""):
            inside = False
            continue
        if inside:
            els.append(para._p)
    return els


def _chapter_last(heading_el):
    head = _para_text(heading_el)
    stop_sub = _is_other_title(head) or bool(_SUB_HEAD.match(_compact_para(head)))
    last = heading_el
    el = heading_el.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        if el.tag == qn("w:p"):
            text = _para_text(el)
            if _sec_head(_section_label(el)) or _paste_ban(text):
                break
            if stop_sub and _OTHER_PREFIX.match(_compact_para(text)):
                break
        if el.tag in {qn("w:p"), qn("w:tbl")}:
            last = el
        el = el.getnext()
    return last


def _move_after(doc: Document, last_el, marker) -> None:
    body = doc.element.body
    moved = []
    seen = marker is None
    for child in list(body):
        if not seen:
            if child is marker:
                seen = True
            continue
        if child.tag == qn("w:sectPr"):
            continue
        moved.append(child)
    anchor = _body_block(last_el)
    clearance = _floating_table_cm(anchor)
    if clearance:
        # wrapNone 的浮动原表不占正文高度；只给新增附件预留它的占位，原表 XML 不动。
        spacer = doc.add_paragraph()
        if _SLOT_CLEARANCE_STYLE not in doc.styles:
            doc.styles.add_style(_SLOT_CLEARANCE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        spacer.style = _SLOT_CLEARANCE_STYLE
        pf = spacer.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.keep_with_next = False
        pf.keep_together = True
        pf.line_spacing = Cm(min(clearance, _page_body_cm(doc) - 0.5))
        spacer.add_run().font.size = Pt(1)
        anchor.addnext(spacer._p)
        anchor = spacer._p
    for child in moved:
        parent = child.getparent()
        if parent is not None:
            parent.remove(child)
        anchor.addnext(child)
        anchor = child
    _merge_slot_gaps(doc)


def _body_block(anchor):
    """单元格、表格行或文本框内的锚点必须提升到正文直属块。"""
    while anchor.getparent() is not None and anchor.getparent().tag != qn("w:body"):
        anchor = anchor.getparent()
    return anchor


def _floating_table_cm(block) -> float:
    """浮动范本表的占用高度；Choice/Fallback 取最大值，不能重复累加。"""
    heights = []
    for anchor in block.iter(qn("wp:anchor")):
        if next(anchor.iter(qn("w:tbl")), None) is None or anchor.find(qn("wp:wrapNone")) is None:
            continue
        extent = anchor.find(qn("wp:extent"))
        position = anchor.find(qn("wp:positionV"))
        offset = position.find(qn("wp:posOffset")) if position is not None else None
        if extent is not None:
            heights.append((int(extent.get("cy", "0")) + max(0, int(offset.text or "0") if offset is not None else 0)) / _EMU_PER_CM)
    for shape in block.iter("{urn:schemas-microsoft-com:vml}shape"):
        if next(shape.iter(qn("w:tbl")), None) is None:
            continue
        style = shape.get("style", "")
        if "position:absolute" not in style.replace(" ", ""):
            continue
        values = {}
        for key, value, unit in re.findall(r"(?:^|;)\s*(height|margin-top)\s*:\s*(-?[\d.]+)(pt|cm|mm|in)", style):
            values[key] = float(value) * {"pt": 2.54 / 72, "cm": 1, "mm": 0.1, "in": 2.54}[unit]
        heights.append(values.get("height", 0) + max(0, values.get("margin-top", 0)))
    if block.tag == qn("w:tbl"):
        positioning = block.find("w:tblPr/w:tblpPr", block.nsmap)
        if positioning is not None:
            heights.append(_tbl_cm(block) + max(0, int(positioning.get(qn("w:tblpY"), "0"))) / _TWIPS_PER_CM)
    return max(heights, default=0.0)


def _slot_anchor(anchor):
    """附件放在本节分页符前；说明后紧接的原范本表格必须先保留。"""
    anchor = _body_block(anchor)
    # _span_last/_chapter_last 会包含末尾空白分节段；在它后面插入等于另起一页。
    while anchor.tag == qn("w:p") and not _para_text(anchor).strip():
        if any(next(anchor.iter(qn("w:" + name)), None) is not None for name in ("drawing", "pict", "object", "tbl")):
            break
        ppr = anchor.find(qn("w:pPr"))
        page_before = ppr.find(qn("w:pageBreakBefore")) if ppr is not None else None
        boundary = (
            next(anchor.iter(qn("w:sectPr")), None) is not None
            or any(br.get(qn("w:type")) == "page" for br in anchor.iter(qn("w:br")))
            or (page_before is not None and page_before.get(qn("w:val"), "true") not in {"0", "false", "off"})
        )
        previous = anchor.getprevious()
        if not boundary or previous is None:
            break
        anchor = previous
    # 不能把资料插在说明与原表之间；连续原表也作为同一个范本块保留。
    cursor = anchor.getnext()
    while cursor is not None and next(cursor.iter(qn("w:tbl")), None) is not None:
        text = _para_text(cursor)
        if "此处粘贴" in text or _INPLACE in text:
            break
        if _sec_head(_section_label(cursor)):
            break
        anchor = cursor
        cursor = cursor.getnext()
    return anchor


def _drop_after(doc: Document, anchor_el, slot: PlaceholderItem, files: list[Path], *, wording: str = "", complete: bool = False, at_marker: bool = False) -> bool:
    anchor_el = _body_block(anchor_el) if at_marker else _slot_anchor(anchor_el)
    marker = _last_body_child(doc)
    placed = False
    if files:
        key = (slot.key or "").strip()
        n = _insert_slot_media(
            doc,
            files if complete else files[:2],
            max_pages=2**31 - 1 if complete else _SLOT_PDF_MAX_PAGES,
            compact_pair=key in {"id_legal", "id_agent"},
        )
        placed = n > 0
    if not placed:
        _draw_box(doc, slot, wording=wording)
    _move_after(doc, anchor_el, marker)
    return placed


def _drop_on_chapter(doc: Document, heading, slot: PlaceholderItem, files: list[Path]) -> bool:
    return _drop_after(doc, _chapter_last(heading._element), slot, files, complete=True)


_QUAL_VETO_MARKS = ("资质证明", "资质证书", "企业资质", "资质文件")
_OTHER_CHAPTERS = {"其他资料", "其他材料", "其他有必要提供的材料", "其他有必要提供的资料"}
_OTHER_PREFIX = re.compile(
    r"^(?:[（(][一二三四五六七八九十0-9]+[)）]|[一二三四五六七八九十0-9]+[、.])"
)
_SUB_HEAD = re.compile(r"^[（(][一二三四五六七八九十0-9]+[)）]")
_CHAPTER_LEVEL = re.compile(
    r"^(?:第[一二三四五六七八九十0-9]+[章节]|[一二三四五六七八九十]{1,3}、)"
)
_INPLACE = "此处应粘贴对应的资料附件或图片"
_NONE_TITLE = "资格审查其他材料无"
_STRONG_CUES = (
    "应在此提供",
    "在此提供",
    "在此附",
    "本表后应附",
    "本表后附",
    "后应附",
    "应附",
    "须附",
    "需附",
    "应提供",
    "需提供",
    "须提供",
    "粘贴",
    "复印件或扫描件",
    "扫描件",
)
_NEAR_CUES = ("材料", "证明", "报表", "证书", "合同", "发票", "执照", "截图", "审计", "纳税", "社保")


def _wants_attach(text: str) -> bool:
    n = _compact_para(text)
    if not n or "此处粘贴" in n or _INPLACE in n:
        return False
    if n.startswith("注") and "附" in n:
        return True
    if any(k in n for k in _STRONG_CUES):
        return True
    i = 0
    while True:
        j = n.find("提供", i)
        if j < 0:
            return False
        window = n[max(0, j - 12) : j + 14]
        if any(k in window for k in _NEAR_CUES):
            return True
        i = j + 2


def _paste_ban(text: str) -> bool:
    n = _compact_para(text)
    if not n or len(n) > 36 or "。" in n:
        return False
    if any(x in n for x in ("或", "；", ";", "，", "须", "提供")):
        return False
    n = re.sub(r"[（(]如有[)）]", "", n)
    n = _OTHER_PREFIX.sub("", n)
    n = re.sub(r"^第[一二三四五六七八九十0-9]+[章节]", "", n)
    n = re.sub(r"^[一二三四五六七八九十]{1,3}、", "", n)
    if n in _NO_PASTE:
        return True
    return n.startswith("投标函") and len(n) <= 16


def _qual_mark(text: str) -> str:
    n = "".join((text or "").split())
    for mark in _QUAL_VETO_MARKS:
        if mark in n:
            return mark
    return ""


_SCATTER_CUES = (
    "应附",
    "须附",
    "需附",
    "应提供",
    "须提供",
    "需提供",
    "应包含",
    "须包含",
    "复印件",
    "扫描件",
)
_NAME_CUT = re.compile(
    r"(?:应附|须附|需附上|需附|应提供|须提供|需提供|应包含|须包含)"
    r"([^，。；;、（(]{2,18})"
)


def _clip_material_name(text: str, target: str) -> str:
    title = "".join((target or "").split())
    if 2 <= len(title) <= 18:
        return title
    n = "".join((text or "").split())
    hit = _NAME_CUT.search(n)
    if hit is None:
        return ""
    name = hit.group(1)
    for tail in ("复印件", "扫描件", "原件"):
        if name.endswith(tail) and len(name) > len(tail):
            name = name[: -len(tail)]
    if len(name) < 2 or len(name) > 18:
        return ""
    return name


def _merge_scattered_materials(
    by_key: dict[str, PlaceholderItem],
    keys: list[str],
    requirements,
    vol_of: dict[str, str],
    volume: str,
) -> list[str]:
    from api.services.tenders.match import (
        _clean_attach_section,
        find_catalog_item,
        same_library_material,
        scan_slot_key,
    )

    vol = (volume or "").strip().lower()
    have = set(keys)
    added: list[str] = []
    seen_names: set[str] = set()
    catalog = list(by_key.values())
    for item in requirements or []:
        text = getattr(item, "text", "") or ""
        n = "".join(text.split())
        if not any(cue in n for cue in _SCATTER_CUES):
            continue
        if n.startswith(("（", "(")) and "如有" in n:
            continue
        if (getattr(item, "matchStatus", "") or "") == "format":
            continue
        if (getattr(item, "targetKind", "") or "") == "format":
            continue
        if (getattr(item, "strength", "") or "") == "mandatory" and _qual_mark(text):
            continue
        item_vol = (getattr(item, "volume", "") or "").strip().lower()
        if vol in {"business", "technical"} and item_vol in {"business", "technical"} and item_vol != vol:
            continue
        name = _clip_material_name(text, getattr(item, "targetTitle", "") or "")
        if not name or name in seen_names:
            continue
        if any(same_library_material(name, (by_key[k].title or "")) for k in have if k in by_key):
            seen_names.add(name)
            continue
        sk = (getattr(item, "slotKey", "") or "").strip()
        found = None
        if sk and sk in by_key:
            found = by_key[sk]
        if found is None:
            found = find_catalog_item(catalog, key=sk, title=name, hint="")
        key = (found.key or "").strip() if found is not None else ""
        if not key:
            guessed = scan_slot_key(name, catalog)
            if guessed:
                key = guessed
                found = by_key.get(guessed)
        if key and key in have:
            seen_names.add(name)
            continue
        if not key:
            key = "scatter_" + name
        if key in have:
            seen_names.add(name)
            continue
        section = _clean_attach_section(getattr(item, "attachToSection", "") or "")
        if found is None or key not in by_key:
            by_key[key] = PlaceholderItem(
                key=key,
                title=(found.title if found is not None and found.title else name),
                hint="",
                attachToSection=section,
            )
            catalog.append(by_key[key])
        elif section and not (by_key[key].attachToSection or "").strip():
            by_key[key].attachToSection = section
        keys.append(key)
        have.add(key)
        seen_names.add(name)
        if item_vol in {"business", "technical"}:
            vol_of[key] = item_vol
        added.append(by_key[key].title or name)
    return added


def _merge_unmatched_qual(by_key: dict[str, PlaceholderItem], keys: list[str], requirements) -> None:
    """否决项里的资质证明没对上资料时，补一个挂在「资格审查其他材料」的空位。"""
    seen: set[str] = set()
    have = set(keys)
    for item in requirements or []:
        if (getattr(item, "strength", "") or "") != "mandatory":
            continue
        if (getattr(item, "matchStatus", "") or "") == "matched":
            continue
        if (getattr(item, "volume", "") or "") == "technical":
            continue
        mark = _qual_mark(getattr(item, "text", "") or "")
        if not mark or mark in seen:
            continue
        sk = (getattr(item, "slotKey", "") or "").strip()
        if sk and sk in have and sk in by_key:
            seen.add(mark)
            continue
        seen.add(mark)
        key = f"veto_{mark}"
        if key in by_key:
            if key not in have:
                keys.append(key)
            continue
        by_key[key] = PlaceholderItem(
            key=key,
            title=mark,
            hint="否决项要求提供，未匹配到扫描件，请在此补",
            attachToSection="资格审查其他材料",
        )
        keys.append(key)


def _is_other_title(text: str) -> bool:
    n = _compact_para(text)
    if not n or len(n) > 28 or "。" in n or "如有" in n:
        return False
    n = _OTHER_PREFIX.sub("", n)
    if n in _OTHER_CHAPTERS:
        return True
    return n.startswith("其他有必要提供") and len(n) <= 16


def _already_placed(doc: Document, slot: PlaceholderItem) -> bool:
    title = _compact_para(slot.title)
    if len(title) < 4:
        return False
    for para in doc.paragraphs:
        if _note_score(title, para.text or "") >= 1000 and _filled_slot_mark(para._p):
            return True
    for tbl in doc.element.body.iter(qn("w:tbl")):
        blob = "".join(t.text or "" for t in tbl.iter(qn("w:t")))
        if _note_score(title, blob) >= 1000 and _filled_slot_mark(tbl):
            return True
    if (slot.key or "") != "id_legal" and "身份证" not in title:
        return False
    # 身份证明页已经贴过图或虚线框，卷末和其他资料里不再来一份
    for para in doc.paragraphs:
        n = _compact_para(para.text or "")
        if "法定代表人身份证" not in n or "委托" in n:
            continue
        if not (n.startswith("附") or "正反面" in n):
            continue
        sib = para._p.getnext()
        for _ in range(6):
            if sib is None or (sib.tag == qn("w:p") and _heading_core(_para_text(sib))):
                break
            xml = sib.xml
            if "w:drawing" in xml or "w:pict" in xml or "此处粘贴" in xml:
                return True
            sib = sib.getnext()
    return False


def _filled_slot_mark(el) -> bool:
    if _is_slot_box(el) or next(el.iter(qn("w:drawing")), None) is not None:
        return True
    return el.tag == qn("w:tbl") and any(
        border.get(qn("w:val")) in {"dashed", "dotted"}
        for border in el.iter(qn("w:top"))
    )


def _other_heading(doc: Document, toc_els: list):
    for para in doc.paragraphs:
        if para._p in toc_els:
            continue
        if _is_other_title(para.text or ""):
            return para
    return None


def _slot_marks(slot: PlaceholderItem) -> list[str]:
    title = _compact_para(slot.title)
    hint = _compact_para(slot.hint)
    blob = title + hint
    marks: list[str] = []
    sec = _level1_core(slot.attachToSection) or _compact_para(slot.attachToSection)
    if (
        sec
        and sec not in _OTHER_CHAPTERS
        and sec not in _NO_PASTE
        and "如有" not in sec
        and not _paste_ban(slot.attachToSection or "")
    ):
        marks.append(sec)
        # 章节标题同义：邀请书可能写全称，原版章节只写简称。
        _SECTION_ALIASES = (
            ("法定代表人授权委托书", "授权委托书"),
            ("授权委托书", "法定代表人授权委托书"),
            ("财务审计报告", "财务审计"),
            ("财务报表", "财务审计"),
            ("财务状况", "财务审计"),
            ("业绩合同", "类似项目"),
            ("业绩", "类似项目"),
            ("资质证书", "资质证明"),
            ("资质证明", "资质证书"),
            ("营业执照", "营业执照"),
            ("开户许可", "开户许可证"),
        )
        for left, right in _SECTION_ALIASES:
            if left in sec and right not in marks:
                marks.append(right)
    key = (slot.key or "").strip()
    if key in _KEY_CHAPTER:
        marks.append(_KEY_CHAPTER[key])
    if key == "credit":
        marks.append("国家企业信用信息公示系统")
    elif key == "finance":
        marks.append("财务状况")
    for mark in _CHAPTER_MARKS:
        if mark in blob:
            marks.append(mark)
    if 4 <= len(title) <= 16:
        marks.append(title)
    marks.sort(key=len, reverse=True)
    out: list[str] = []
    for mark in marks:
        if mark and mark not in out:
            out.append(mark)
    return out


def _locate_chapter(doc: Document, slot: PlaceholderItem, toc_els: list, letter_els: list | None = None):
    """对上正文专章就用那一章。目录、投标函和「如有」列举处不算。"""
    banned = letter_els or []
    best = None
    best_n = 0
    label = ""
    for para in doc.paragraphs:
        if para._p in toc_els or para._p in banned or _paste_ban(para.text or ""):
            continue
        source = _section_label(para._p)
        core = _level1_core(source) or (_SUB_HEAD.sub("", source) if _SUB_HEAD.match(source) else _heading_core(source))
        if not core or core in _NO_PASTE or core in _OTHER_CHAPTERS or core.startswith("投标函"):
            continue
        if "工作经历" in core or "简历" in core:
            continue
        for mark in _slot_marks(slot):
            if mark in core or core in mark:
                if len(core) > best_n:
                    best = para
                    best_n = len(core)
                    label = _compact_para(para.text or "")
                break
    return best, label


def _real_chapter(text: str) -> bool:
    n = _compact_para(text)
    if not n or len(n) > 36 or "。" in n or "如有" in n or n.startswith(("（", "(")):
        return False
    if not _CHAPTER_LEVEL.match(n):
        return False
    if _paste_ban(text) or _is_other_title(text):
        return False
    return True


def _form_els(doc: Document) -> set:
    banned = set(_letter_block(doc))
    inside = False
    for child in doc.element.body:
        if child.tag == qn("w:sectPr"):
            break
        if child.tag == qn("w:p"):
            text = _para_text(child)
            if _paste_ban(text):
                inside = True
                banned.add(child)
                continue
            if inside and (
                _real_chapter(_section_label(child))
                or (_SUB_HEAD.match(_compact_para(_section_label(child))) and not _paste_ban(text))
                or (_heading_core(text) and not _paste_ban(text))
            ):
                inside = False
                continue
            if inside:
                banned.add(child)
        elif inside and child.tag == qn("w:tbl"):
            banned.add(child)
    return banned


def _note_score(title: str, text: str) -> int:
    name = _compact_para(title)
    n = _compact_para(text)
    if len(name) < 4:
        return 0
    for marker in re.finditer(r"此处粘贴([^。]{1,80})", n):
        material = re.split(r"[（(\n]在此粘贴|[（(\n]|[；;。]", marker.group(1), maxsplit=1)[0]
        compared_name = name.replace("财务会计报表", "财务报表")
        compared_material = material.replace("财务会计报表", "财务报表")
        if compared_name in compared_material or (
            len(compared_material) >= 4 and compared_material in compared_name
        ):
            return 1000 + min(len(name), len(material))
    if not _wants_attach(n):
        return 0
    if name in n:
        return len(name)
    hit = re.search(r"(?:此处粘贴|在此附|本表后应附|本表后附|后应附|应附|应提供|需提供|须提供|须附|需附)([^。]{0,80})", n)
    if hit is None:
        return 0
    best = 0
    raw = re.split(r"[；;]", hit.group(1))[0]
    for part in re.split(r"或|、|及|和", raw):
        part = re.sub(r"(复印件|扫描件|原件|等)$", "", part).strip("（）()：:，,")
        if len(part) < 4:
            continue
        if part in name or name in part:
            best = max(best, min(len(part), len(name)))
    return best


def _near_label(el) -> str:
    prev = el.getprevious()
    while prev is not None:
        if prev.tag == qn("w:tbl"):
            break
        if prev.tag == qn("w:p"):
            text = _compact_para(_para_text(prev))
            if text and not text.startswith("注") and "。" not in text and len(text) <= 32:
                return text
        prev = prev.getprevious()
    return ""


def _locate_note(doc: Document, slot: PlaceholderItem, banned: set, toc_els: list):
    best = None
    best_rank = (False, False, 0)
    label = ""
    title = slot.title or ""
    section = _level1_core(slot.attachToSection) or _compact_para(slot.attachToSection)
    for child in doc.element.body:
        if child.tag == qn("w:sectPr"):
            break
        if child in banned or child in toc_els:
            continue
        if child.tag == qn("w:p"):
            score = _note_score(title, _para_text(child))
        elif child.tag == qn("w:tbl"):
            blob = "".join(t.text or "" for t in child.iter(qn("w:t")))
            score = _note_score(title, blob)
        else:
            continue
        if score <= 0:
            continue
        owner = _owner(child)
        owner_name = _compact_para(_section_label(owner)) if owner is not None else ""
        nearby = bool(section and owner_name and (section in owner_name or owner_name in section))
        marker = score >= 1000
        rank = (marker, nearby, score)
        if rank < best_rank or (marker and rank == best_rank):
            continue
        best = child
        best_rank = rank
        label = _near_label(child) or "附注"
    return best, label


def _sec_head(text: str) -> bool:
    n = _compact_para(text)
    if not n or len(n) > 36 or "。" in n:
        return False
    if _paste_ban(text):
        return True
    if "如有" in n or n.startswith("注"):
        return False
    if _real_chapter(text) or _heading_core(text):
        return True
    return bool(_SUB_HEAD.match(n))


def _section_label(el) -> str:
    """原格式章的编号常只存在于标题书签中，浮动表格也可作为章节边界。"""
    for mark in el.iter(qn("w:bookmarkStart")):
        name = mark.get(qn("w:name"), "")
        if _real_chapter(name) or _SUB_HEAD.match(name):
            return name
    return _para_text(el)


def _owner(el):
    cur = el
    while cur is not None:
        if cur.tag == qn("w:p") and _sec_head(_section_label(cur)):
            return cur
        cur = cur.getprevious()
    return None


def _span_taken(head) -> bool:
    el = head.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        if el.tag == qn("w:p") and _sec_head(_section_label(el)):
            break
        blob = "".join(t.text or "" for t in el.iter(qn("w:t")))
        flat = "".join(blob.split())
        if "此处粘贴" in flat or _INPLACE in flat or _NONE_TITLE in flat:
            return True
        if el.tag == qn("w:p") and ("w:drawing" in el.xml or "w:pict" in el.xml):
            return True
        el = el.getnext()
    return False


def _span_last(head):
    last = head
    el = head.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        if el.tag == qn("w:p") and _sec_head(_section_label(el)):
            break
        if el.tag in {qn("w:p"), qn("w:tbl")}:
            last = el
        el = el.getnext()
    return last


def _cue_anchor(head):
    found = None
    el = head.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        if el.tag == qn("w:p") and _sec_head(_section_label(el)):
            break
        if el.tag == qn("w:p"):
            text = _para_text(el)
        elif el.tag == qn("w:tbl"):
            text = "".join(t.text or "" for t in el.iter(qn("w:t")))
        else:
            el = el.getnext()
            continue
        if _wants_attach(text):
            flat = _compact_para(text)
            if flat.startswith("注") or "注：" in flat or "注:" in flat:
                return el
            if found is None:
                found = el
        el = el.getnext()
    return found


def _embed_files(slot: PlaceholderItem, media, key: str) -> list[Path]:
    out = []
    for item in (media or {}).get(key) or []:
        path = item if isinstance(item, Path) else Path(str(item))
        try:
            if not path.is_file():
                continue
        except OSError:
            continue
        if path.suffix.lower() in _IMAGE_EXT | {".pdf"}:
            out.append(path)
    return out


def _files_for_blurb(text: str, by_key: dict, media):
    from api.services.tenders.match import find_catalog_item, scan_slot_key

    name = _clip_material_name(text, "")
    if not name:
        hit = re.search(
            r"(?:在此附|后应附|应附|应提供|需提供|须提供)([^，。；;、（(]{2,18})",
            _compact_para(text),
        )
        name = hit.group(1) if hit else ""
    if len(name) < 2:
        return [], None
    catalog = list(by_key.values())
    found = find_catalog_item(catalog, title=name, hint="")
    key = (found.key or "").strip() if found is not None else ""
    if not key:
        key = scan_slot_key(name, catalog)
        if key:
            found = by_key.get(key)
    if not key:
        return [], None
    slot = found if found is not None else PlaceholderItem(key=key, title=name)
    return _embed_files(slot, media, key), slot


def _review_plate(doc: Document, toc_els: list, ban: set):
    for para in doc.paragraphs:
        if para._p in toc_els or para._p in ban:
            continue
        n = _compact_para(para.text or "")
        if "资格审查其他材料" in n and len(n) <= 28 and "。" not in n:
            return para
    return None


def _write_none_letter(doc: Document, head) -> str:
    if _span_taken(head):
        return ""
    last = _span_last(head)
    marker = _last_body_child(doc)
    doc.add_paragraph(_NONE_TITLE)
    doc.add_paragraph("我方无需要补充的其他资格审查材料。")
    _move_after(doc, last, marker)
    return "「资格审查其他材料」栏为空，已补无声明"


def _sweep_inplace(doc: Document, by_key: dict, media, filled: set, toc_els: list, ban: set, plate_el, notes: list[str]) -> None:
    for child in doc.element.body:
        if child.tag == qn("w:sectPr"):
            break
        if child.tag != qn("w:p") or child in toc_els or child in ban:
            continue
        if not _sec_head(_para_text(child)) or _paste_ban(_para_text(child)):
            continue
        if plate_el is not None and child is plate_el:
            continue
        if child in filled or _span_taken(child):
            filled.add(child)
            continue
        anchor = _cue_anchor(child)
        if anchor is None:
            continue
        if anchor.tag == qn("w:p"):
            text = _para_text(anchor)
        else:
            text = "".join(t.text or "" for t in anchor.iter(qn("w:t")))
        files, slot = _files_for_blurb(text, by_key, media)
        if slot is None:
            slot = PlaceholderItem(key="inplace", title="")
        put = _drop_after(doc, anchor, slot, files, wording=_INPLACE, complete=True)
        filled.add(child)
        label = _compact_para(_para_text(child))[:24] or "附注"
        if put:
            notes.append(f"「{label}」已按说明贴上资料")
        else:
            notes.append(f"「{label}」说明要求附资料，已在附注后画虚线框")


def _paste_marker_title(text: str) -> str:
    """Read the material named by an original paste cue, not its scan instruction."""
    match = re.search(r"此处粘贴([^（(\n。；;]{2,80})", _compact_para(text))
    if match is None:
        return ""
    name = re.split(r"等(?=能|相关|证明|材料)|(?:的)?(?:复印件|扫描件)$", match.group(1), maxsplit=1)[0]
    return name.strip("：:，,、 ")[:32]


def _place_requested_materials(
    doc: Document, by_key: dict[str, PlaceholderItem], media: dict | None,
    toc_els: list, ban: set, skipped: set[str], filled: set, notes: list[str],
) -> set[str]:
    """从正文说明逐项落位。图片/框只代表这一资料已填，不代表整章完成。"""
    from api.services.tenders.match import attachment_catalog_item, attachment_requests, scan_slot_key

    body = list(doc.element.body)
    existing = {b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))}
    next_id = max((int(b.get(qn("w:id"), "0")) for b in doc.element.body.iter(qn("w:bookmarkStart"))), default=0) + 1
    handled = set()
    occurrences: dict[str, int] = {}
    consumed = set()
    claimed_markers = set()
    catalog = sorted(by_key.values(), key=lambda item: not bool((media or {}).get(item.key)))
    for index, child in enumerate(body):
        if child.getparent() is not doc.element.body or child in consumed or child in toc_els or child in ban or child.tag not in {qn("w:p"), qn("w:tbl")}:
            continue
        text = "".join(t.text or "" for t in child.iter(qn("w:t")))
        if "此处粘贴" in text and child not in claimed_markers:
            name = _paste_marker_title(text)
            if not name:
                continue
            found = attachment_catalog_item(PlaceholderItem(key="", title=name), catalog)
            key = found.key if found is not None else (scan_slot_key(name, catalog) or "marker_" + hashlib.sha1(name.encode()).hexdigest()[:12])
            found = found or by_key.get(key)
            if key in skipped:
                continue
            files = _embed_files(found, media, key) if found is not None else []
            if _filled_slot_mark(child) and not files:
                handled.add(key)
                continue
            following = child.getnext()
            while _is_slot_gap(following) or _is_slot_clearance(following):
                following = following.getnext()
            if following is not None and _filled_slot_mark(following) and (
                _note_score(name, _para_text(following)) >= 1000
                or any((mark.get(qn("w:name")) or "").startswith("material_") for mark in following.iter(qn("w:bookmarkStart")))
            ):
                handled.add(key)
                continue
            signature = text.strip()
            ordinal = occurrences.get(signature, 0)
            occurrences[signature] = ordinal + 1
            marker_name = "material_" + hashlib.sha1(f"paste|{signature}|{ordinal}|{key}".encode()).hexdigest()[:24]
            if marker_name in existing:
                handled.add(key)
                continue
            slot = PlaceholderItem(key=key, title=found.title if found is not None else name)
            target = _body_block(child)
            put = _drop_after(doc, target, slot, files, complete=True, at_marker=True)
            first = target.getnext()
            while _is_slot_gap(first) or _is_slot_clearance(first):
                first = first.getnext()
            if _filled_slot_mark(child):
                child.getparent().remove(child)
                _merge_slot_gaps(doc)
            para = first if first.tag == qn("w:p") else next(first.iter(qn("w:p")))
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(next_id))
            start.set(qn("w:name"), marker_name)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(next_id))
            para.append(start)
            para.append(end)
            next_id += 1
            existing.add(marker_name)
            handled.add(key)
            claimed_markers.add(child)
            notes.append(f"「{name}」已在原文粘贴标记处" + ("贴上资料" if put else "画虚线框（资料待补）"))
            continue
        # 只处理原说明，不把生成的框或其提示再次解释成要求。
        if "此处粘贴" in text or _INPLACE in text or any(
            (b.get(qn("w:name")) or "").startswith("material_") for b in child.iter(qn("w:bookmarkStart"))
        ):
            continue
        anchor = child
        if _wants_attach(text) and child.tag == qn("w:p"):
            # PDF 转 Word 的一句说明可能拆成多个普通段落，等句子结束后再插入。
            for following in body[index + 1:index + 9]:
                if text.rstrip().endswith(("。", "；", ";")):
                    break
                if following in toc_els or following in ban or following.tag != qn("w:p"):
                    break
                continuation = _para_text(following)
                if _sec_head(_section_label(following)) or not continuation.strip():
                    break
                text += continuation
                anchor = following
                consumed.add(following)
        requests = attachment_requests(text)
        owner = _owner(child)
        label = _section_label(owner) if owner is not None else ""
        if _wants_attach(text) and "类似项目" in label and (not requests or all(r.title == "相关证明材料" for r in requests)):
            requests = [PlaceholderItem(key="perf", title="类似项目业绩证明")]
        if "上述途径查询结果截图" in text and ("企业信用" in label or "信用中国" in label):
            requests = [PlaceholderItem(key="credit", title="国家企业信用信息公示系统查询截图")]
        # 清单章要求提供已填写的清单，即使格式说明未写「扫描件」也须留就地位置。
        core = _compact_para(text)
        if len(core) < 36 and re.search(r"已标价的?工程量清单$", core) and _sec_head(_section_label(child)):
            anchor = _span_last(child)
            requests = [PlaceholderItem(key="boq", title="已标价工程量清单表")]
        if not requests:
            continue
        signature = text.strip()
        ordinal = occurrences.get(signature, 0)
        occurrences[signature] = ordinal + 1
        for request in requests:
            found = attachment_catalog_item(request, catalog)
            key = found.key if found is not None else request.key
            if key in skipped or request.key in skipped:
                # The unified scan placer may already have inserted this request.
                # Keep later missing-material boxes after those images, not in front.
                from api.services.tenders.scan_placement import _prefix

                cursor = anchor.getnext()
                prefixes = (_prefix(key), _prefix(request.key))
                if cursor is not None and any(
                    (node.get(qn("w:name")) or "").startswith(prefixes)
                    for node in cursor.iter(qn("w:bookmarkStart"))
                ):
                    while cursor is not None and (
                        _filled_slot_mark(cursor) or _is_slot_gap(cursor)
                    ):
                        anchor = cursor
                        cursor = cursor.getnext()
                continue
            handled.add(key)
            handled.update(item.key for item in catalog if attachment_catalog_item(request, [item]) is not None)
            token = hashlib.sha1(f"{signature}|{ordinal}|{request.key}".encode()).hexdigest()[:24]
            marker_name = "material_" + token
            if marker_name in existing:
                if owner is not None:
                    filled.add(owner)
                continue
            slot = request.model_copy(update={"key": key})
            # 按明确说明选中的 PDF 与图片均可就地嵌入，不能把已上传 PDF 当成缺件。
            files = []
            for item in (media or {}).get(key) or []:
                path = Path(item)
                if path.is_file() and path.suffix.lower() in _IMAGE_EXT | {".pdf"}:
                    files.append(path)
            # 模板已有具名粘贴框时复用该位置，不能因整章已有一个框就跳过其他资料。
            old_box = None
            cursor = anchor.getnext()
            names = [request.title, found.title if found is not None else ""]
            while cursor is not None and cursor.tag != qn("w:sectPr"):
                if cursor.tag == qn("w:p") and _sec_head(_section_label(cursor)):
                    break
                blob = _compact_para(_para_text(cursor))
                if any(name and _note_score(name, blob) >= 1000 for name in names):
                    old_box = cursor
                    break
                cursor = cursor.getnext()
            existing_box = old_box is not None and _filled_slot_mark(old_box)
            if existing_box and not files:
                if owner is not None:
                    filled.add(owner)
                continue
            if old_box is not None and old_box in claimed_markers:
                continue
            target = _body_block(old_box) if old_box is not None else _slot_anchor(anchor)
            put = _drop_after(doc, target, slot, files, complete=True, at_marker=old_box is not None)
            first = target.getnext()
            while _is_slot_gap(first) or _is_slot_clearance(first):
                first = first.getnext()
            if existing_box:
                old_box.getparent().remove(old_box)
                _merge_slot_gaps(doc)
            if old_box is not None:
                claimed_markers.add(old_box)
            # 标记只放在新增块内，原正文、段落属性及表格 XML 保持原样。
            para = first if first.tag == qn("w:p") else next(first.iter(qn("w:p")))
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(next_id))
            start.set(qn("w:name"), marker_name)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(next_id))
            para.append(start)
            para.append(end)
            next_id += 1
            existing.add(marker_name)
            # 下一资料跟在刚插入块之后，保持说明中的顺序。
            anchor = first
            while anchor.getnext() is not None and anchor.getnext() not in body:
                anchor = anchor.getnext()
            if owner is not None:
                filled.add(owner)
            notes.append(f"「{request.title}」已按正文说明就地" + ("贴上资料" if put else "画虚线框（资料待补）"))
    return handled


def place_unmatched_slots(
    doc: Document,
    *,
    slots: list[PlaceholderItem] | None,
    media: dict | None,
    required_keys: list[str] | None,
    volume: str = "",
    skip_keys: set[str] | None = None,
    requirements: list | None = None,
    global_pasted: set[str] | None = None,
) -> list[str]:
    """按正文说明或专章就地放置；无对应位置只报告缺口，不追加卷末附件章。

    Args:
        global_pasted: 全局已贴集合，跨所有贴图函数共享，同一资料键只贴一次
    """
    keys = [str(k or "").strip() for k in (required_keys or []) if str(k or "").strip()]
    by_key = {item.key: item for item in (slots or []) if (item.key or "").strip()}
    if not by_key:
        by_key = {item.key: item for item in DEFAULT_SLOTS if item.key}
    _merge_unmatched_qual(by_key, keys, requirements)
    vol_of: dict[str, str] = {}
    scattered = _merge_scattered_materials(by_key, keys, requirements, vol_of, volume)
    vol = (volume or "").strip().lower()
    skipped = {str(k) for k in (skip_keys or set())}
    # 合并全局已贴集合
    if global_pasted is not None:
        skipped.update(global_pasted)
    from api.services.tenders.scan_placement import place_library_scans

    scan_notes, scan_handled = place_library_scans(doc, media=media, slots=slots, skip_keys=skipped, global_pasted=global_pasted)
    skipped.update(scan_handled)
    # 更新全局已贴集合
    if global_pasted is not None:
        global_pasted.update(scan_handled)
    from api.services.tenders.categories import slot_volume
    from api.services.tenders.pdf_format import toc_block_ids

    toc_els = toc_block_ids(doc)
    letter_els = _letter_block(doc)
    ban = _form_els(doc)
    notes: list[str] = list(scan_notes)
    leftover: list[PlaceholderItem] = []
    plate_items: list[PlaceholderItem] = []
    filled: set = set()
    plate = None if vol == "technical" else _review_plate(doc, toc_els, ban)
    plate_el = plate._element if plate is not None else None
    handled = _place_requested_materials(doc, by_key, media, toc_els, ban, skipped, filled, notes)

    # 【修改2：将 toc_els 添加到禁止放置区域】
    # 确保缺失项占位不会被画在目录下，必须在对应章节或文末"其他资料"
    toc_els_set = set(toc_els)

    for key in keys:
        if key in {TECH_DRAWING_KEY, "seal"} or key in skipped or key in handled:
            continue
        slot = by_key.get(key)
        if slot is None:
            continue
        anchor, note_label = _locate_note(doc, slot, ban, toc_els)
        # 【修改2：检查 anchor 是否在目录块中】
        # 如果 anchor 在目录块中，跳过这个位置，强制寻找章节或放到"其他资料"
        # anchor 是 CT_P 元素（从 _locate_note 返回），直接检查是否在目录集合中
        if anchor is not None and anchor in toc_els_set:
            anchor = None
            note_label = ""

        at_marker = bool(
            anchor is not None
            and _note_score(slot.title, _para_text(anchor)) >= 1000
        )
        want = vol_of.get(key, "")
        if vol in {"business", "technical"}:
            if want in {"business", "technical"}:
                if want != vol and not at_marker:
                    continue
            elif slot_volume(key, slot.title or "") != vol and not at_marker:
                continue
        if _already_placed(doc, slot):
            continue
        files = []
        for item in (media or {}).get(key) or []:
            path = item if isinstance(item, Path) else Path(str(item))
            try:
                if path.is_file():
                    files.append(path)
            except OSError:
                continue
        if anchor is not None:
            owner = _owner(anchor)
            if owner is not None and not at_marker and (owner in filled or _span_taken(owner)):
                notes.append(f"「{slot.title}」与同一章节已有占位重复，未再画框")
                continue
            embed = _embed_files(slot, {key: files}, key)
            put = _drop_after(
                doc, anchor, slot, embed,
                wording="" if at_marker else _INPLACE,
                complete=True, at_marker=at_marker,
            )
            if owner is not None:
                filled.add(owner)
            if put:
                notes.append(f"「{slot.title}」已贴到「{note_label}」附注后")
            else:
                notes.append(f"「{slot.title}」没有扫描件，已在「{note_label}」附注后画虚线框")
            continue
        heading, section = _locate_chapter(doc, slot, toc_els, letter_els)
        if heading is not None:
            if heading._element in filled or _span_taken(heading._element):
                notes.append(f"「{slot.title}」与同一章节已有占位重复，未再画框")
                continue
            put = _drop_on_chapter(doc, heading, slot, files)
            filled.add(heading._element)
            if put:
                notes.append(f"「{slot.title}」已贴到「{section}」正文后")
            else:
                notes.append(f"「{slot.title}」没有扫描件，已在「{section}」正文后画虚线框")
            continue
        if plate is not None:
            _drop_on_chapter(doc, plate, slot, files)
            plate_items.append(slot)
            continue
        leftover.append(slot)
    _sweep_inplace(doc, by_key, media, filled, toc_els, ban, plate_el, notes)
    if plate_items:
        titles = "、".join(item.title for item in plate_items if item.title)
        notes.append(f"无归属材料已集中到「资格审查其他材料」：{titles}")
    if plate is not None:
        blank = _write_none_letter(doc, plate_el)
        if blank:
            notes.append(blank)
    if leftover:
        titles = "、".join(item.title for item in leftover if item.title)
        notes.append(f"未找到对应正文章节，待指定资料落位：{titles}")
    if scattered:
        notes.insert(0, "正文散落要求已纳入占位：" + "、".join(scattered))
    return notes


def draw_placeholder_box(doc: Document, slot: PlaceholderItem | None = None, *, after=None) -> None:
    """在文档中画虚线粘贴框；after 为段落时插到该段之后（用于技术标图纸）。"""
    marker = _last_body_child(doc)
    _draw_box(doc, slot or PlaceholderItem(key="pending", title="待补资料"))
    if after is None:
        _merge_slot_gaps(doc)
        return
    _move_after(doc, _slot_anchor(after._element), marker)


def fill_perf_placeholders(table: Table) -> None:
    """企业业绩表空单元格写入待补说明，避免空白被误认为已填。"""
    if not table.rows:
        return
    start = 1 if len(table.rows) > 1 else 0
    for row in table.rows[start:]:
        for cell in row.cells:
            compact = "".join((cell.text or "").split())
            if compact and compact not in {"／", "/", "—", "-", "……", "..."}:
                continue
            _write_plain(cell, _PERF_NOTE)


def _ordinal(index: int) -> str:
    if 0 <= index < len(_CN_ORD):
        return _CN_ORD[index]
    return str(index + 1)


def _page_body_cm(doc: Document) -> float:
    sec = doc.sections[-1]
    return max(
        16.0,
        float(sec.page_height.cm) - float(sec.top_margin.cm) - float(sec.bottom_margin.cm) - 0.6,
    )


def _drawing_cm(el) -> float:
    total = 0.0
    for node in el.iter():
        if node.tag.split("}")[-1] != "extent":
            continue
        cy = node.get("cy")
        if cy:
            total += int(cy) / _EMU_PER_CM
    return total


def _tbl_cm(tbl) -> float:
    total = 0.0
    for tr in tbl.findall(qn("w:tr")):
        row_cm = 0.0
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is not None:
            height = tr_pr.find(qn("w:trHeight"))
            if height is not None:
                raw = height.get(qn("w:val"))
                if raw:
                    row_cm = int(raw) / _TWIPS_PER_CM
        total += row_cm if row_cm > 0 else 0.72
    return max(total, 0.72)


def _used_on_page_cm(doc: Document) -> float:
    used = 0.0
    for child in doc.element.body:
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:tbl"):
            used += _tbl_cm(child)
            continue
        if child.tag != qn("w:p"):
            continue
        if _is_slot_gap(child):
            used += _SLOT_GAP_CM
            continue
        if any(br.get(qn("w:type")) == "page" for br in child.iter(qn("w:br"))):
            used = 0.35
            continue
        p_pr = child.find(qn("w:pPr"))
        if p_pr is not None and p_pr.find(qn("w:sectPr")) is not None:
            used = 0.35
        drawn = _drawing_cm(child)
        if drawn:
            used += min(drawn + 0.35, 20.0)
            continue
        text = "".join(node.text or "" for node in child.iter(qn("w:t"))).replace("\u200b", "").strip()
        if p_pr is not None:
            sp = p_pr.find(qn("w:spacing"))
            if sp is not None and (sp.get(qn("w:lineRule")) or "") == "exact":
                line = int(sp.get(qn("w:line")) or 0)
                if line:
                    used += line / _TWIPS_PER_CM
                    continue
        if not text:
            used += 0.22
            continue
        lines = max(1, (len(text) + 29) // 30)
        used += 0.48 * lines + 0.22
    return used


def _first_media_cm(files: list[Path], *, width_cm: float = 15.5, max_h: float | None = None) -> float:
    cap = _IMAGE_MAX_H_CM if max_h is None else max_h
    for path in files:
        suf = path.suffix.lower()
        if suf == ".pdf":
            return min(cap, width_cm * 297 / 210)
        if suf in _IMAGE_EXT and path.is_file():
            try:
                from PIL import Image

                with Image.open(path) as image:
                    w, h = image.size
                if w > 0:
                    return min(cap, max(3.5, width_cm * h / float(w)))
            except Exception:
                return 10.0
    return _BOX_CM


def _break_before_slot(doc: Document, need_cm: float) -> None:
    """标题和附件必须同页：剩余高度不够时，标题改到下一页开头。"""
    body_cm = _page_body_cm(doc)
    used = _used_on_page_cm(doc) % body_cm
    remain = body_cm - used
    if remain + 0.2 >= need_cm:
        return
    if used <= 0.6:
        return
    doc.add_page_break()


def _write_subheading(doc: Document, index: int, title: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.keep_with_next = True
    run = para.add_run(f"（{_ordinal(index)}）{title}")
    _font(run, 14, bold=True)


def _is_slot_gap(el) -> bool:
    if el is None or el.tag != qn("w:p"):
        return False
    style = el.find("w:pPr/w:pStyle", el.nsmap)
    return style is not None and style.get(qn("w:val")) == _SLOT_GAP_STYLE


def _is_slot_clearance(el) -> bool:
    if el is None or el.tag != qn("w:p"):
        return False
    style = el.find("w:pPr/w:pStyle", el.nsmap)
    return style is not None and style.get(qn("w:val")) == _SLOT_CLEARANCE_STYLE


def _is_slot_box(el) -> bool:
    return el is not None and el.tag == qn("w:p") and el.find("w:pPr/w:pBdr", el.nsmap) is not None and (
        "此处粘贴" in _para_text(el) or _INPLACE in _para_text(el)
    )


def _merge_slot_gaps(doc: Document) -> None:
    """相邻附件共用一个留白段，不叠加成 1 cm。"""
    previous_gap = False
    for node in list(doc.element.body):
        gap = _is_slot_gap(node)
        if gap and previous_gap:
            doc.element.body.remove(node)
        previous_gap = gap


def _space_slot_block(doc: Document, block) -> None:
    """只在新增附件外放 0.5 cm 空段；不修改相邻原文或表格。"""
    if _SLOT_GAP_STYLE not in doc.styles:
        style = doc.styles.add_style(_SLOT_GAP_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        style.font.size = Pt(1)
        pf = style.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = Cm(_SLOT_GAP_CM)
        pf.left_indent = pf.right_indent = pf.first_line_indent = Cm(0)
        pf.keep_with_next = False
        pf.page_break_before = False
    for before in (True, False):
        gap = doc.add_paragraph(style=_SLOT_GAP_STYLE)
        # 留白本身就是正文空段，直接写高度，避免依赖阅读器解析自定义样式。
        gap.paragraph_format.space_before = gap.paragraph_format.space_after = Pt(0)
        gap.paragraph_format.line_spacing = Cm(_SLOT_GAP_CM)
        gap.add_run().font.size = Pt(1)
        # 身份证图原有的“与下段同页”继续连到签章段，不能被尾部留白截断。
        keep_next = block.tag == qn("w:p") and Paragraph(block, doc._body).paragraph_format.keep_with_next is True
        gap.paragraph_format.keep_with_next = before or keep_next
        if before:
            block.addprevious(gap._p)
        else:
            block.addnext(gap._p)


def _draw_box(doc: Document, slot: PlaceholderItem, *, wording: str = "") -> None:
    """普通文字段落带虚线边框，选中整段（含段落标记）即可 Delete 删除。"""
    mark = doc.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = mark.paragraph_format
    pf.left_indent = pf.right_indent = pf.first_line_indent = Cm(0)
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(12)
    pf.keep_together = True
    pf.keep_with_next = False
    pf.page_break_before = False
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement("w:" + edge)
        for name, value in (("val", "dashed"), ("sz", "18"), ("space", "8"), ("color", "7A7A7A")):
            border.set(qn("w:" + name), value)
        borders.append(border)
    mark._p.get_or_add_pPr().append(borders)
    title = (slot.title or "").strip()
    line = wording or (f"此处粘贴{title}" if title else "")
    text = line if wording else line + "\n（在此粘贴扫描件）"
    run = mark.add_run("\n" * 4 + text + "\n" * 4)
    _font(run, 10.5, bold=False)
    run.font.color.rgb = RGBColor(0x7A, 0x7A, 0x7A)
    _space_slot_block(doc, mark._p)


def _insert_slot_media(
    doc: Document,
    files: list[Path],
    *,
    budget: dict[str, int] | None = None,
    max_pages: int = _SLOT_PDF_MAX_PAGES,
    compact_pair: bool = False,
    force_landscape: bool = False,
    sharp: bool = False,
) -> int:
    original_nodes = set(doc.element.body)
    if compact_pair:
        prepared = _collect_scan_jpegs(
            files,
            budget=budget,
            max_images=_ID_PAIR_MAX_IMAGES,
            max_pages=max(_SLOT_PDF_MAX_PAGES, _ID_PAIR_MAX_IMAGES),
            force_landscape=force_landscape,
        )
        inserted = _insert_id_cards(doc, prepared)
        _space_new_slot_media(doc, original_nodes)
        return inserted

    inserted = 0
    for path in files:
        if budget is not None and budget.get("left", 0) <= 0:
            break
        suf = path.suffix.lower()
        try:
            if suf == ".pdf":
                n = _insert_pdf_pages(doc, path, max_pages=max_pages, budget=budget, sharp=sharp)
                inserted += n
            elif suf in _IMAGE_EXT:
                if budget is not None and budget.get("left", 0) <= 0:
                    break
                if _insert_image(doc, path, sharp=sharp):
                    inserted += 1
                    if budget is not None:
                        budget["left"] = max(0, int(budget["left"]) - 1)
        except Exception:
            logger.exception("insert slot media failed: %s", path)
    _space_new_slot_media(doc, original_nodes)
    return inserted


def _space_new_slot_media(doc: Document, original_nodes: set) -> None:
    for node in list(doc.element.body):
        if node in original_nodes or node.tag != qn("w:p"):
            continue
        if next(node.iter(qn("w:drawing")), None) is None:
            continue
        # 外侧留白统一交给空段，避免与旧的图片段前/段后距离叠加。
        para = Paragraph(node, doc._body)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        _space_slot_block(doc, node)


def _render_pdf_page_jpeg(
    page,
    *,
    scale: float = _RENDER_SCALE,
    quality: int = _JPEG_QUALITY,
    max_px: int | None = None,
) -> io.BytesIO:
    bitmap = page.render(scale=scale)
    image = bitmap.to_pil().convert("RGB")
    image = _shrink_pil(image, max_px=_IMAGE_MAX_PX if max_px is None else max_px)
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=quality, optimize=False)
    buf.seek(0)
    return buf


def _shrink_pil(image, max_px: int | None = None):
    try:
        cap = _IMAGE_MAX_PX if max_px is None else int(max_px)
        w, h = image.size
        longest = max(w, h)
        if cap <= 0 or longest <= cap:
            return image
        from PIL import Image as PilImage

        ratio = cap / float(longest)
        resample = getattr(getattr(PilImage, "Resampling", PilImage), "LANCZOS", PilImage.BICUBIC)
        return image.resize((max(1, int(w * ratio)), max(1, int(h * ratio))), resample)
    except Exception:
        return image


def _apply_exif(image):
    try:
        from PIL import ImageOps

        transposed = ImageOps.exif_transpose(image)
        return transposed if transposed is not None else image
    except Exception:
        return image


def _looks_like_id_card(width: int, height: int) -> bool:
    short = min(width, height)
    long = max(width, height)
    if short <= 0:
        return False
    ratio = long / float(short)
    return _ID_ASPECT_MIN <= ratio <= _ID_ASPECT_MAX


def _force_landscape(image):
    from PIL import Image

    w, h = image.size
    if h <= w:
        return image
    rot = Image.Transpose.ROTATE_90 if hasattr(Image, "Transpose") else Image.ROTATE_90
    return image.transpose(rot)


def _normalize_id_image(image, *, force_landscape: bool):
    image = _apply_exif(image).convert("RGB")
    w, h = image.size
    if force_landscape or _looks_like_id_card(w, h):
        image = _force_landscape(image)
    return _shrink_pil(image)


def _image_to_jpeg(image) -> io.BytesIO:
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=max(_JPEG_QUALITY, 55), optimize=False)
    buf.seek(0)
    return buf


def _fit_cm(px_w: int, px_h: int, *, max_w: float, max_h: float) -> tuple[float, float]:
    if px_w <= 0 or px_h <= 0:
        return max_w, min(max_h, max_w * 0.63)
    aspect = px_w / float(px_h)
    width = max_w
    height = width / aspect
    if height > max_h:
        height = max_h
        width = height * aspect
    return width, height


def _collect_scan_jpegs(
    files: list[Path],
    *,
    budget: dict[str, int] | None,
    max_images: int,
    max_pages: int,
    force_landscape: bool,
) -> list[tuple[io.BytesIO, int, int]]:
    out: list[tuple[io.BytesIO, int, int]] = []
    for path in files[:_SLOT_MAX_FILES]:
        if len(out) >= max_images:
            break
        if budget is not None and budget.get("left", 0) <= 0:
            break
        suf = path.suffix.lower()
        try:
            if suf == ".pdf":
                for buf, w, h in _pdf_pages_as_jpegs(
                    path,
                    max_pages=max_pages,
                    force_landscape=force_landscape,
                    limit=max_images - len(out),
                ):
                    out.append((buf, w, h))
                    if budget is not None:
                        budget["left"] = max(0, int(budget["left"]) - 1)
                    if len(out) >= max_images:
                        break
                    if budget is not None and budget.get("left", 0) <= 0:
                        break
            elif suf in _IMAGE_EXT and path.is_file():
                from PIL import Image

                with Image.open(path) as raw:
                    image = _normalize_id_image(raw, force_landscape=force_landscape)
                    w, h = image.size
                    out.append((_image_to_jpeg(image), w, h))
                if budget is not None:
                    budget["left"] = max(0, int(budget["left"]) - 1)
        except Exception:
            logger.exception("prepare slot scan failed: %s", path)
    return out


def _pdf_pages_as_jpegs(
    pdf_path: Path,
    *,
    max_pages: int,
    force_landscape: bool,
    limit: int,
) -> list[tuple[io.BytesIO, int, int]]:
    if not pdf_path.is_file() or limit <= 0:
        return []
    try:
        import pypdfium2 as pdfium
    except ImportError:
        logger.warning("pypdfium2 missing, skip slot pdf %s", pdf_path.name)
        return []
    try:
        pdf = pdfium.PdfDocument(str(pdf_path))
    except Exception:
        logger.exception("open slot pdf failed: %s", pdf_path)
        return []
    pages: list[tuple[io.BytesIO, int, int]] = []
    count = min(len(pdf), max_pages, limit)
    for i in range(count):
        try:
            bitmap = pdf[i].render(scale=_RENDER_SCALE)
            image = _normalize_id_image(bitmap.to_pil(), force_landscape=force_landscape)
            w, h = image.size
            pages.append((_image_to_jpeg(image), w, h))
        except Exception:
            logger.exception("render slot pdf page %s of %s failed", i + 1, pdf_path.name)
    return pages


def _insert_id_cards(doc: Document, images: list[tuple[io.BytesIO, int, int]]) -> int:
    """正反面合成一张横版小图，紧跟小标题，避免预览按原图像素各占一页。"""
    if not images:
        return 0
    strip = _compose_id_strip([buf for buf, _w, _h in images[:_ID_PAIR_MAX_IMAGES]])
    if strip is None:
        return 0
    baked, width_cm, height_cm = strip
    sec = doc.sections[-1]
    available = (sec.page_width - sec.left_margin - sec.right_margin) / _EMU_PER_CM - 0.4
    if width_cm > available:
        height_cm *= available / width_cm
        width_cm = available
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(10)
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.keep_with_next = True
    p_pr = para._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:pageBreakBefore")):
        p_pr.remove(old)
    para.add_run().add_picture(baked, width=Cm(width_cm), height=Cm(height_cm))
    return min(len(images), _ID_PAIR_MAX_IMAGES)


def _compose_id_strip(buffers: list[io.BytesIO]) -> tuple[io.BytesIO, float, float] | None:
    from PIL import Image

    cards: list = []
    for buf in buffers:
        try:
            buf.seek(0)
            with Image.open(buf) as raw:
                cards.append(_normalize_id_image(raw, force_landscape=True))
        except Exception:
            logger.exception("open id scan for strip failed")
    if not cards:
        return None
    resample = getattr(getattr(Image, "Resampling", Image), "LANCZOS", Image.BICUBIC)
    card_h = 220
    prepared = []
    for im in cards:
        im = _force_landscape(im)
        w, h = im.size
        nw = max(1, int(round(card_h * w / float(h or 1))))
        prepared.append(im.resize((nw, card_h), resample))
    gap = 16
    canvas_w = sum(im.size[0] for im in prepared) + gap * (len(prepared) - 1)
    canvas = Image.new("RGB", (max(canvas_w, 1), card_h), (255, 255, 255))
    x = 0
    for im in prepared:
        canvas.paste(im, (x, 0))
        x += im.size[0] + gap
    # 单张按标准身份证尺寸上限，两张并排；窄版心由调用方等比缩小
    width_cm = _ID_STRIP_WIDTH_CM if len(prepared) > 1 else _ID_SINGLE_WIDTH_CM
    width_cm, height_cm = _fit_cm(
        canvas.size[0], card_h, max_w=width_cm, max_h=_ID_STRIP_MAX_H_CM
    )
    px_w = max(1, int(round(width_cm / 2.54 * _ID_PREVIEW_DPI)))
    px_h = max(1, int(round(height_cm / 2.54 * _ID_PREVIEW_DPI)))
    canvas = canvas.resize((px_w, px_h), resample)
    out = io.BytesIO()
    canvas.save(out, format="JPEG", quality=82, dpi=(_ID_PREVIEW_DPI, _ID_PREVIEW_DPI), optimize=True)
    out.seek(0)
    return out, width_cm, height_cm


def _bake_id_jpeg(buf: io.BytesIO, *, width_cm: float, height_cm: float) -> io.BytesIO:
    """把图缩到与排版厘米一致的像素，docx-preview 即使用原图尺寸也不会撑满整页。"""
    from PIL import Image

    buf.seek(0)
    with Image.open(buf) as raw:
        image = raw.convert("RGB")
        px_w = max(1, int(round(width_cm / 2.54 * _ID_PREVIEW_DPI)))
        px_h = max(1, int(round(height_cm / 2.54 * _ID_PREVIEW_DPI)))
        resample = getattr(getattr(Image, "Resampling", Image), "LANCZOS", Image.BICUBIC)
        image = image.resize((px_w, px_h), resample)
        out = io.BytesIO()
        image.save(out, format="JPEG", quality=82, dpi=(_ID_PREVIEW_DPI, _ID_PREVIEW_DPI), optimize=True)
        out.seek(0)
        return out


def _insert_image(doc: Document, path: Path, *, width_cm: float = 15.5, sharp: bool = False) -> bool:
    if not path.is_file():
        return False
    try:
        from PIL import Image

        with Image.open(path) as raw:
            image = raw.convert("RGB")
            if not sharp:
                image = _shrink_pil(image)
            buf = io.BytesIO()
            if sharp:
                image.save(buf, format="PNG")
            else:
                image.save(buf, format="JPEG", quality=_JPEG_QUALITY, optimize=False)
            buf.seek(0)
            max_h = _CREDIT_MAX_H_CM if sharp else _IMAGE_MAX_H_CM
            disp_w, disp_h = _fit_to_page(doc, image.size[0], image.size[1], max_w=width_cm, max_h=max_h)
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(4)
            pic.paragraph_format.space_after = Pt(4)
            pic.add_run().add_picture(buf, width=Cm(disp_w), height=Cm(disp_h))
        return True
    except Exception:
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(path), width=Cm(width_cm))
        return True


def _fit_to_page(
    doc: Document,
    px_w: int,
    px_h: int,
    *,
    max_w: float,
    max_h: float,
) -> tuple[float, float]:
    remain = _page_body_cm(doc) - _used_on_page_cm(doc) - 0.4
    max_h = min(max_h, _scan_height_cm(doc))
    if remain < 4.0 and _used_on_page_cm(doc) > 1.0:
        doc.add_page_break()
        remain = _page_body_cm(doc) - 0.8
    cap = min(max_h, max(4.0, remain))
    return _fit_cm(px_w, px_h, max_w=max_w, max_h=cap)


def _scan_height_cm(doc: Document) -> float:
    """Reserve captions/gaps so normal scans can share a body page."""
    return min(_IMAGE_MAX_H_CM, max(4.0, _page_body_cm(doc) / 2 - 1.0))


def _render_pdf_page_bytes(page, *, sharp: bool) -> io.BytesIO:
    if sharp:
        scale = _CREDIT_DPI / 72.0
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil().convert("RGB")
        image = _shrink_pil(image, max_px=_CREDIT_MAX_PX)
        buf = io.BytesIO()
        image.save(buf, format="PNG", dpi=(_CREDIT_DPI, _CREDIT_DPI))
        buf.seek(0)
        return buf
    return _render_pdf_page_jpeg(page)


def _insert_pdf_pages(
    doc: Document,
    pdf_path: Path,
    *,
    max_pages: int = _SLOT_PDF_MAX_PAGES,
    budget: dict[str, int] | None = None,
    sharp: bool = False,
) -> int:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        logger.warning("pypdfium2 missing, skip slot pdf %s", pdf_path.name)
        return 0
    try:
        pdf = pdfium.PdfDocument(str(pdf_path))
    except Exception:
        logger.exception("open slot pdf failed: %s", pdf_path)
        return 0

    try:
        count = min(len(pdf), max_pages)
        inserted = 0
        max_h = _CREDIT_MAX_H_CM if sharp else _IMAGE_MAX_H_CM
        for i in range(count):
            if budget is not None and budget.get("left", 0) <= 0:
                break
            try:
                buf = _render_pdf_page_bytes(pdf[i], sharp=sharp)
                px_w, px_h = 210, 297
                try:
                    from PIL import Image

                    buf.seek(0)
                    with Image.open(buf) as preview:
                        px_w, px_h = preview.size
                    buf.seek(0)
                except Exception:
                    pass
                disp_w, disp_h = _fit_to_page(doc, px_w, px_h, max_w=15.5, max_h=max_h)
                pic = doc.add_paragraph()
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.paragraph_format.space_before = Pt(2)
                pic.paragraph_format.space_after = Pt(2)
                buf.seek(0)
                pic.add_run().add_picture(buf, width=Cm(disp_w), height=Cm(disp_h))
                inserted += 1
                if budget is not None:
                    budget["left"] = max(0, int(budget["left"]) - 1)
            except Exception:
                logger.exception("render slot pdf page %s of %s failed", i + 1, pdf_path.name)
        return inserted
    finally:
        try:
            pdf.close()
        except Exception:
            pass


def _write_plain(cell: _Cell, text: str) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    _font(run, 9, bold=False)


def _font(run, size_pt: float | None = None, *, bold: bool | None = None) -> None:
    if bold is not None:
        run.bold = bold
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.font.name = _SONG
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), _SONG)
    rfonts.set(qn("w:hAnsi"), _SONG)
    rfonts.set(qn("w:eastAsia"), _SONG)


def _set_dashed_borders(table: Table) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "18")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "7A7A7A")
        borders.append(el)
    tbl_pr.append(borders)


def _set_row_height(row, height_cm: float, *, exact: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for old in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(old)
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(Cm(height_cm).twips)))
    el.set(qn("w:hRule"), "exact" if exact else "atLeast")
    tr_pr.append(el)


def _set_cell_margins(cell: _Cell, *, top: int, bottom: int, left: int, right: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcMar")):
        tc_pr.remove(old)
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)
