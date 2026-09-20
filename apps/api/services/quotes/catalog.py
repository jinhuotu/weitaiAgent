"""从所选知识库的 Excel 抽出价目行，再按名称/规格匹配。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from api.services.tenders.quote import _q
from api.services.tenders.tables import quote_role

_TWO = Decimal("0.01")
_SKIP = ("不含税合计", "含税合计", "税率", "合计", "小计", "总计", "备注")
_TOKEN = re.compile(r"[a-z0-9]+|[\u4e00-\u9fff]+", re.I)
_KW_RE = re.compile(r"(\d+(?:\.\d+)?)kw", re.I)
_KVA_RE = re.compile(r"(\d+(?:\.\d+)?)kva", re.I)

CODE_HINTS: dict[str, tuple[str, ...]] = {
    "dc_320kw": ("320kw", "320千瓦", "400kw", "直流320", "320kW直流"),
    "dc_160kw": ("160kw", "160千瓦", "直流160", "160kW直流"),
    "dc_120kw": ("120kw", "120千瓦", "直流120", "120kW直流"),
    "ac_14kw": ("14kw", "14千瓦", "交流14", "交流桩"),
    "box_transformer": ("箱变", "箱式变压", "箱式变", "变电站"),
    "ring_cabinet": ("环网柜", "环网箱"),
    "lv_cabinet": ("低压柜", "低压配电"),
    "group_host": ("群充", "主机"),
    "cable": ("电缆", "铜芯"),
    "trench": ("电缆沟", "沟槽"),
    "foundation": ("基础", "基座"),
}


@dataclass(frozen=True)
class CatalogItem:
    name: str
    spec: str
    unit: str
    unit_price: Decimal
    cost_price: Decimal = Decimal("0")
    sell_price: Decimal = Decimal("0")
    source: str = ""
    category: str = ""
    scene: str = ""


def parse_catalog_xlsx(path: Path) -> list[CatalogItem]:
    from openpyxl import load_workbook

    # 不用 read_only：不少价目表 dimension 仍是 A1，只读模式会整表丢行
    wb = load_workbook(path, data_only=True)
    try:
        rows: list[CatalogItem] = []
        for ws in wb.worksheets:
            rows.extend(_sheet_items(ws, path.name))
        return rows
    finally:
        wb.close()


def _sheet_items(ws, source: str) -> list[CatalogItem]:
    mapping: dict[str, int] | None = None
    out: list[CatalogItem] = []
    for row in ws.iter_rows(values_only=True):
        cells = list(row or ())
        if mapping is None:
            mapping = _header_map(cells)
            continue
        item = _row_item(cells, mapping, source)
        if item is not None:
            out.append(item)
    return out


def _header_map(cells: list[object]) -> dict[str, int] | None:
    roles = [quote_role(c) for c in cells]
    if "name" not in roles:
        return None
    if "price" not in roles and "cost_price" not in roles and "sell_price" not in roles:
        return None
    mapping: dict[str, int] = {}
    for i, role in enumerate(roles):
        if role and role not in mapping:
            mapping[role] = i
    return mapping


def _cell(cells: list[object], idx: int | None) -> object:
    if idx is None or idx < 0 or idx >= len(cells):
        return None
    return cells[idx]


def _txt(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def _row_item(cells: list[object], mapping: dict[str, int], source: str) -> CatalogItem | None:
    name = _txt(_cell(cells, mapping.get("name")))
    if not name or any(k in name for k in _SKIP):
        return None
    spec = _txt(_cell(cells, mapping.get("spec")))
    unit = _txt(_cell(cells, mapping.get("unit"))) or "项"
    price = _q(_cell(cells, mapping.get("price")))
    cost = _q(_cell(cells, mapping.get("cost_price")))
    sell = _q(_cell(cells, mapping.get("sell_price")))
    if price <= 0:
        price = sell if sell > 0 else cost
    if price <= 0:
        return None
    if cost <= 0:
        cost = price
    if sell <= 0:
        sell = price
    return CatalogItem(
        name=name,
        spec=spec,
        unit=unit,
        unit_price=price,
        source=source,
        category=_txt(_cell(cells, mapping.get("group"))),
        scene=_txt(_cell(cells, mapping.get("scene"))),
        cost_price=cost,
        sell_price=sell,
    )


def load_catalog(paths: list[Path]) -> list[CatalogItem]:
    items: list[CatalogItem] = []
    for path in paths:
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        try:
            items.extend(parse_catalog_xlsx(path))
        except Exception:
            continue
    return items


def _norm(text: str) -> str:
    return re.sub(r"[\s\-_/／（）()【】\[\]]+", "", (text or "").lower())


def _tokens(text: str) -> set[str]:
    return {m.group(0).lower() for m in _TOKEN.finditer(text or "") if len(m.group(0)) >= 2}


def _ratings(text: str, pat: re.Pattern[str]) -> set[float]:
    return {float(x) for x in pat.findall(_norm(text))}


def _rating_conflict(query: str, item: CatalogItem) -> bool:
    blob = f"{item.category} {item.name} {item.spec}"
    qkw, bkw = _ratings(query, _KW_RE), _ratings(blob, _KW_RE)
    if qkw and bkw and qkw.isdisjoint(bkw):
        return True
    qkva, bkva = _ratings(query, _KVA_RE), _ratings(blob, _KVA_RE)
    return bool(qkva and bkva and qkva.isdisjoint(bkva))


def match_score(query: str, code: str, item: CatalogItem) -> float:
    blob = f"{item.category} {item.name} {item.spec}"
    name_blob = f"{item.category} {item.name}"
    qn = _norm(query)
    bn = _norm(blob)
    if not qn or not bn:
        return 0.0
    if _rating_conflict(query, item):
        return 0.0
    score = 0.0
    if qn == bn:
        score = 1.0
    elif qn in bn or bn in qn:
        score = 0.86
    else:
        qt = _tokens(query)
        bt = _tokens(blob)
        if qt and bt:
            score = len(qt & bt) / max(len(qt), 1)
    bn_name = _norm(name_blob)
    for hint in CODE_HINTS.get((code or "").strip().lower(), ()):
        if _norm(hint) and _norm(hint) in bn_name:
            score = max(score, 0.78)
    return score


def _spec_rank(item: CatalogItem) -> int:
    spec = item.spec or ""
    n = len(spec)
    if re.search(r"[（(]\s*1\s*[）)]", spec) or "输入电压" in spec or "输出功率" in spec:
        n += 8000
    return n


def pick_catalog(
    *,
    name: str,
    spec: str,
    code: str,
    catalog: list[CatalogItem],
) -> tuple[CatalogItem | None, float]:
    query = f"{name} {spec}".strip()
    passed: list[tuple[float, CatalogItem]] = []
    best_s = 0.0
    for item in catalog:
        s = match_score(query, code, item)
        if s < 0.42:
            continue
        passed.append((s, item))
        if s > best_s:
            best_s = s
    if not passed:
        return None, best_s
    qkw = _ratings(query, _KW_RE)
    qkva = _ratings(query, _KVA_RE)
    same = passed
    if qkw:
        kw_hit = [
            (s, it)
            for s, it in passed
            if qkw & _ratings(f"{it.category} {it.name} {it.spec}", _KW_RE)
        ]
        if kw_hit:
            same = kw_hit
    elif qkva:
        kva_hit = [
            (s, it)
            for s, it in passed
            if qkva & _ratings(f"{it.category} {it.name} {it.spec}", _KVA_RE)
        ]
        if kva_hit:
            same = kva_hit
    same.sort(key=lambda x: x[0], reverse=True)
    hit = same[0][1]
    rich = max(same, key=lambda x: _spec_rank(x[1]))[1]
    if rich.spec and _spec_rank(rich) > _spec_rank(hit):
        hit = CatalogItem(
            name=hit.name,
            spec=rich.spec,
            unit=hit.unit,
            unit_price=hit.unit_price,
            source=hit.source,
            category=hit.category or rich.category,
            scene=hit.scene or rich.scene,
            cost_price=hit.cost_price or rich.cost_price,
            sell_price=hit.sell_price or rich.sell_price,
        )
    return hit, best_s


def money(qty: Decimal, price: Decimal) -> Decimal:
    return (qty * price).quantize(_TWO, rounding=ROUND_HALF_UP)
