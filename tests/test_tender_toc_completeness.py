from copy import deepcopy

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.document import _pageref_bookmark, _write_toc_line
from api.services.tenders.outline import toc_title_key
from api.services.tenders.pdf_format import _link_toc_body
from api.services.tenders.schema import OutlineItem, default_brief


@pytest.mark.parametrize('with_pages', [True, False])
def test_toc_line_always_has_dot_tab_and_pageref(with_pages):
    doc = Document()
    para = doc.add_paragraph()
    _write_toc_line(para, 0, '资格资信证明文件', 'toc_o01', '待更新', with_pages=with_pages)
    assert para.text == '一、资格资信证明文件\t待更新'
    assert _pageref_bookmark(para) == 'toc_o01'
    tab = para._p.find('w:pPr/w:tabs/w:tab', para._p.nsmap)
    assert tab.get(qn('w:leader')) == 'dot' and tab.get(qn('w:val')) == 'right'


def test_native_inquiry_toc_matches_aggregate_headings_and_preserves_body():
    doc = Document()
    doc.add_paragraph('目 录')
    titles = ['一、资格资信证明文件', '二、合同条款响应书', '三、技术和商务差异表', '四、响应方案', '五、分项报价表']
    entries = [doc.add_paragraph(t) for t in titles]
    heads = []
    for title in ['一、资格、资信证明文件', titles[1], '三、商务和技术偏差表', *titles[3:]]:
        doc.add_page_break()
        heads.append(doc.add_paragraph(title))
        doc.add_paragraph('本章原始正文内容。')
    brief = default_brief()
    brief.outlineItems = [OutlineItem(id=f'o{i:02d}', title=t, kind='unknown', level=1, body='')
                          for i, t in zip([1, 7, 8, 9, 10], [p.text for p in heads], strict=True)]
    before = [deepcopy(p._p) for p in heads]
    assert _link_toc_body(doc, brief) == 5
    for i, (entry, head, original) in enumerate(zip(entries, heads, before, strict=True), 1):
        assert entry.text == titles[i-1] + '\t' + str(i)
        bookmark = _pageref_bookmark(entry)
        assert any(n.get(qn('w:name')) == bookmark for n in head._p.iter(qn('w:bookmarkStart')))
        assert head.text == Paragraph(original, doc._body).text
        # 锚定不改字体、段落间距或正文内容。
        assert head.paragraph_format.space_before is None
        assert 'w:tabs' in entry._p.xml and 'w:leader="dot"' in entry._p.xml
    assert _pageref_bookmark(entries[0]) == 'toc_o01'
    assert _pageref_bookmark(entries[2]) == 'toc_o08'
    assert _link_toc_body(doc, brief) == 5
    assert len(list(doc.element.body.iter(qn('w:bookmarkStart')))) == 5


def test_unmatched_native_toc_keeps_locked_placeholder_not_fabricated_page():
    doc = Document()
    doc.add_paragraph('目录')
    missing = doc.add_paragraph('一、尚无正文的附件')
    known = doc.add_paragraph('二、响应方案')
    doc.add_page_break()
    doc.add_paragraph('二、响应方案')
    doc.add_paragraph('正文内容。')
    assert _link_toc_body(doc) == 1
    assert missing.text.endswith('\t待更新')
    assert _pageref_bookmark(missing)
    begin = next(n for n in missing._p.iter(qn('w:fldChar')) if n.get(qn('w:fldCharType')) == 'begin')
    assert begin.get(qn('w:fldLock')) == 'true'
    assert known.text.endswith('\t1')
    assert 'w:leader="dot"' in missing._p.xml


def test_outline_aggregate_and_special_kind_both_have_body_anchors(tmp_path):
    brief = default_brief()
    brief.layoutMode = 'outline'
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id='shturl', title='一、资格、资信证明文件', kind='unknown', level=1, body=''),
        OutlineItem(id='shturl', title='三、商务和技术偏差表', kind='tech_dev', level=1, body=''),
    ]
    path, _ = assemble_bid_docx(brief, tmp_path / 'aggregate.docx', catalog_slots=[], catalog_media={})
    doc = Document(path)
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert len(entries) == 2
    names = {_pageref_bookmark(p) for p in entries}
    assert names == {'toc_shturl', 'toc_shturl_2'}
    hosts = {n.get(qn('w:name')): n.getparent() for n in doc.element.body.iter(qn('w:bookmarkStart'))}
    assert all(hosts[name].getparent() is doc.element.body for name in names)
    assert all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)


def test_matching_keys_do_not_merge_individual_business_and_technical_tables():
    assert toc_title_key('技术和商务差异表') == toc_title_key('商务和技术偏差表')
    assert toc_title_key('商务偏差表') != toc_title_key('技术偏差表')
    assert toc_title_key('一、资格、资信证明文件') == toc_title_key('资格资信证明文件')


def test_pending_reference_still_allows_existing_table_repair():
    from docx.oxml import OxmlElement
    from api.services.tenders.document import _bookmark_paragraph
    from api.services.tenders.pdf_format import _repair_format_table_toc
    doc = Document()
    doc.add_paragraph('目录')
    doc.add_paragraph('六、项目管理机构')
    entry = doc.add_paragraph('项目管理机构组成表')
    doc.add_page_break()
    doc.add_paragraph('六、项目管理机构')
    carrier = doc.add_paragraph()
    pict = OxmlElement('w:pict')
    content = OxmlElement('w:txbxContent')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '姓名与任职信息' * 12
    content.append(table._tbl)
    pict.append(content)
    carrier.add_run()._r.append(pict)
    carrier.add_run('项目管理机构组成表')
    _bookmark_paragraph(carrier, '(一)项目管理机构组成表')
    brief = default_brief()
    brief.outlineItems = [OutlineItem(id='shturl',title='项目管理机构组成表')]
    _link_toc_body(doc, brief)
    assert entry.text.endswith('待更新')
    # 不把尚无正文锚点的占位 PAGEREF 当作已修复。
    assert _repair_format_table_toc(doc, brief) == 1
    assert _pageref_bookmark(entry) == 'toc_shturl'
    assert entry.text.rsplit('\t', 1)[-1].isdigit()
    assert sum('PAGEREF' in (n.text or '') for n in entry._p.iter(qn('w:instrText'))) == 1


def test_outline_toc_uses_body_chapters_and_real_pages():
    from api.services.tenders.document import (
        _finalize_pagination,
        _p_has_page_br,
        _rebuild_toc,
    )

    doc = Document()
    doc.add_paragraph('封面')
    doc.add_page_break()
    doc.add_paragraph('目录')
    doc.add_paragraph('一、旧条目')
    doc.add_page_break()
    titles = ['一、资格证明文件', '二、商务响应', '三、不参与围标串标承诺书', '四、投标承诺函']
    for i, title in enumerate(titles):
        doc.add_paragraph(title)
        doc.add_paragraph('本章正文内容。' * 80)
        if i < len(titles) - 1:
            doc.add_page_break()
    brief = default_brief()
    brief.outlineItems = [
        OutlineItem(id='q1', title='资格证明文件', level=1),
        OutlineItem(id='q2', title='商务响应', level=1),
        OutlineItem(id='q3', title='不参与围标串标承诺书', level=1),
        OutlineItem(id='skip', title='不需要的章节', level=1, skipped=True),
        OutlineItem(id='q4', title='投标承诺函', level=1),
        OutlineItem(id='missing', title='正文缺失章节', level=1),
    ]
    skipped = _rebuild_toc(doc, brief, with_pages=True, fmt=brief.documentFormat)
    assert skipped == ['正文缺失章节']
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert [p.text.split('\t', 1)[0] for p in entries] == [
        '一、资格证明文件',
        '二、商务响应',
        '三、不参与围标串标承诺书',
        '四、投标承诺函',
    ]
    assert all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)
    toc = next(p for p in doc.paragraphs if p.text.strip() == '目录')
    assert _p_has_page_br(toc._p.getprevious())

    pages = _finalize_pagination(doc, brief.documentFormat, brief)
    assert pages and all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)
    body = next(p for p in doc.paragraphs if p.text.strip() == titles[0])
    assert _p_has_page_br(body._p.getprevious())
    assert len(doc.sections) >= 3
