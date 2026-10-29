from __future__ import annotations

import asyncio
import json

import pytest
from docx import Document

from api.services.tenders.assemble import _fill_drafted_format_sections, assemble_bid_docx
from api.services.tenders.gen_tech_plan import (
    _decode_payload,
    _score_point_kind,
    _scoring_requirements,
    draft_technical_plan,
    save_tech_source,
)
from api.services.tenders.classifier import classify_procurement_type
from api.services.tenders.schema import BidBrief, BidRequirement, OutlineItem
from common.errors import AppError


class FakeClient:
    def __init__(self, payload):
        self.payload = payload
        self.messages = None

    async def complete(self, messages, *, mode):
        assert mode == "pro"
        self.messages = messages
        return json.dumps(self.payload, ensure_ascii=False)


def _brief():
    return BidBrief(
        projectName="本项目名称",
        techPlanNote="用户指定先施工后调试",
        autoDraftTechPlan=True,
        outlineItems=[
            OutlineItem(id="a", title="项目概况与施工部署", kind="tech_chapter", level=2),
            OutlineItem(
                id="b",
                title="质量保证体系",
                kind="tech_chapter",
                level=2,
                draftBody="人工写的质量体系",
                draftStatus="edited",
            ),
            OutlineItem(id="c", title="技术偏离表", kind="tech_dev", level=1),
        ],
        bidRequirements=[
            BidRequirement(id="r1", text="说明现场安全措施", category="scoring", volume="technical")
        ],
    )


def test_draft_only_existing_empty_chapters_and_deviations(tmp_path):
    brief = _brief()
    client = FakeClient(
        {
            "chapters": [
                {"title": "凭空新增章节", "body": "不得进入"},
                {
                    "title": "项目概况与施工部署",
                    "body": "本项目名称按用户指定先施工后调试，人员数量【待补充】。",
                },
            ],
            "deviations": [
                {
                    "requirement": "说明现场安全措施",
                    "response": "待核实现场条件",
                    "deviation": "待确认",
                    "chapter": "项目概况与施工部署",
                }
            ],
        }
    )
    result, notes = asyncio.run(draft_technical_plan(None, brief, client=client))
    assert [(x.id, x.title, x.level) for x in result.outlineItems] == [
        (x.id, x.title, x.level) for x in brief.outlineItems
    ]
    assert result.outlineItems[0].draftBody.startswith("技术得分点对照表")
    assert "【待补充】邀请书正文为空" in result.outlineItems[0].draftBody
    assert any("邀请书正文为空" in note for note in notes)
    assert "本项目名称" in result.outlineItems[0].draftBody
    assert result.outlineItems[1].draftBody == "人工写的质量体系"
    assert result.deviationLines[0].response.endswith("见「项目概况与施工部署」）")
    prompt = client.messages[1]["content"]
    assert (
        prompt.index("【用户输入")
        < prompt.index("【招标技术要求")
        < prompt.index("【知识库")
        < prompt.index("【本项目事实")
    )
    assert "先施工后调试" in prompt
    assert notes
    dest = tmp_path / "tech.docx"
    assemble_bid_docx(result, dest, volume="technical")
    tables = "\n".join(
        cell.text for table in Document(dest).tables for row in table.rows for cell in row.cells
    )
    assert "说明现场安全措施" in tables
    assert "见「项目概况与施工部署」" in tables


def test_four_inputs_keep_user_source_first(monkeypatch):
    from api.services.tenders import gen_tech_plan

    brief = _brief()
    brief.techPlanSourceIds = ["a" * 16]
    brief.techPlanKbIds = ["kb1"]
    monkeypatch.setattr(gen_tech_plan, "_user_sources", lambda ids, **kwargs: "用户方案采用双班施工")

    async def recall(_db, _invitation, _ids, **_kwargs):
        return [{"content": "1. 施工工序\n2. 质量检查，形成验收记录。旧项目金额 99 万元"}], []

    client = FakeClient(
        {"chapters": [{"title": "项目概况与施工部署", "body": "双班施工，项目名称本项目名称。"}]}
    )
    asyncio.run(draft_technical_plan(None, brief, client=client, recall=recall))
    prompt = client.messages[1]["content"]
    assert "用户方案采用双班施工" in prompt
    assert "按编号步骤组织" in prompt
    assert "质量检查与验收记录" in prompt
    assert "99 万元" not in prompt
    assert "项目名称：本项目名称" in prompt
    assert prompt.index("用户方案采用双班施工") < prompt.index("按编号步骤组织")


@pytest.mark.parametrize(
    ("kind", "cue", "topic"),
    [
        ("goods", "货物采购技术参数表及设备型号", "配置说明"),
        ("service", "运维服务方案和服务响应", "团队配置"),
        ("engineering", "施工组织设计与施工部署", "文明施工与环保"),
    ],
)
def test_procurement_routes_scorecard_into_existing_body(kind, cue, topic):
    assert classify_procurement_type(cue) == kind
    brief = BidBrief(
        bidContent=cue,
        outlineItems=[OutlineItem(id="one", title="技术方案", kind="tech_plan")],
        bidRequirements=[
            BidRequirement(id="s1", text="技术评分项要求逐项说明", category="scoring", volume="technical")
        ],
    )
    client = FakeClient(
        {
            "chapters": [{"title": "技术方案", "body": "按招标要求组织实施，数量【待补充】。"}],
            "deviations": [{
                "requirement": "技术评分项要求逐项说明",
                "response": "已列出响应路径，数量【待补充】",
                "deviation": "待确认",
                "chapter": "技术方案",
                "evidence": "【待补充】",
            }],
        }
    )
    drafted, _ = asyncio.run(draft_technical_plan(None, brief, client=client))
    assert [(i.id, i.title, i.level) for i in drafted.outlineItems] == [
        (i.id, i.title, i.level) for i in brief.outlineItems
    ]
    assert drafted.outlineItems[0].draftBody.startswith("技术得分点对照表")
    assert "技术评分项要求逐项说明" in drafted.outlineItems[0].draftBody
    assert "证明材料：【待补充】" in drafted.deviationLines[0].response
    assert topic in client.messages[0]["content"]


def test_no_scoring_points_are_marked_for_completion():
    brief = BidBrief(outlineItems=[OutlineItem(title="技术方案", kind="tech_plan")])
    client = FakeClient({"chapters": [{"title": "技术方案", "body": "方案正文。"}]})
    drafted, _ = asyncio.run(draft_technical_plan(None, brief, client=client))
    assert "未识别到技术评分项" in drafted.outlineItems[0].draftBody
    assert "【待补充】" in drafted.outlineItems[0].draftBody


def test_scoring_rows_only_come_from_technical_section():
    text = "技术评分\n技术参数符合要求得 5 分\n服务应急措施得 3 分\n商务评分\n报价合理得 10 分"
    rows = _scoring_requirements(text)
    assert any("技术参数" in row.text for row in rows)
    assert any("服务应急" in row.text for row in rows)
    assert all("报价" not in row.text for row in rows)


def test_score_points_separate_mandatory_bonus_and_format():
    assert _score_point_kind(BidRequirement(text="★关键参数必须响应", category="scoring")) == "必须响应"
    assert _score_point_kind(BidRequirement(text="方案完整性得分", category="scoring")) == "加分项"
    assert _score_point_kind(BidRequirement(text="暗标格式要求", category="format")) == "格式要求"


def test_manual_chapters_remain_and_deviation_is_flagged_for_review():
    brief = _brief()
    brief.outlineItems[0].draftBody = "人工施工部署"
    brief.outlineItems[0].draftStatus = "edited"
    result, _ = asyncio.run(draft_technical_plan(None, brief))
    assert [i.draftBody for i in result.outlineItems] == [i.draftBody for i in brief.outlineItems]
    assert result.deviationLines[0].requirement == "说明现场安全措施"
    assert "【待补充】" in result.deviationLines[0].response


def test_incomplete_llm_output_keeps_returned_body_and_marks_missing_section():
    brief = _brief()
    brief.outlineItems[1].draftBody = ""
    brief.outlineItems[1].draftStatus = "empty"
    client = FakeClient({"chapters": [{"title": "项目概况与施工部署", "body": "只返回了一章"}]})
    result, notes = asyncio.run(draft_technical_plan(None, brief, client=client))
    assert "只返回了一章" in result.outlineItems[0].draftBody
    assert "【待补充】" in result.outlineItems[1].draftBody
    assert any("1 个章节" in note and "【待补充】" in note for note in notes)
    assert [(i.id, i.title, i.level) for i in result.outlineItems] == [
        (i.id, i.title, i.level) for i in brief.outlineItems
    ]
    assert all(not item.draftBody for item in brief.outlineItems[:2])


def test_model_json_with_explanation_after_object_is_accepted():
    pairs, _ = _decode_payload(
        '{"chapters":[{"title":"施工组织设计","body":"“本项目”现场踏勘后分区施工。"}],"deviations":[]}\n以上为结果。',
        ["施工组织设计"],
    )
    assert pairs == [("施工组织设计", "“本项目”现场踏勘后分区施工。")]


def test_upload_docx_user_material(tmp_path, monkeypatch):
    from api.services.tenders import gen_tech_plan

    monkeypatch.setattr(gen_tech_plan, "_source_dir", lambda: tmp_path)
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("用户确认采用分区施工")
    doc.save(source)
    saved = asyncio.run(save_tech_source(source.read_bytes(), "source.docx"))
    assert saved["sourceId"]
    assert "分区施工" in (tmp_path / f"{saved['sourceId']}.txt").read_text(encoding="utf-8")


def test_copied_format_body_fills_after_existing_title(tmp_path):
    path = tmp_path / "format.docx"
    doc = Document()
    doc.add_paragraph("目录")
    doc.add_paragraph("项目概况与施工部署")
    doc.add_paragraph("下一章")
    doc.add_page_break()
    doc.add_paragraph("项目概况与施工部署")
    doc.add_paragraph("下一章")
    doc.save(path)
    brief = _brief()
    brief.outlineItems[0].draftBody = "本项目名称按用户指定先施工后调试。"
    _fill_drafted_format_sections(path, brief)
    lines = [p.text for p in Document(path).paragraphs][4:]
    assert (
        lines.index("项目概况与施工部署")
        < lines.index("本项目名称按用户指定先施工后调试。")
        < lines.index("下一章")
    )


def test_default_generation_does_not_enable_drafting(tmp_path):
    brief = _brief()
    brief.autoDraftTechPlan = False
    dest = tmp_path / "plain.docx"
    assemble_bid_docx(brief, dest, volume="technical")
    assert dest.is_file()


def test_legacy_brief_without_new_draft_flags_uses_original_path(tmp_path):
    from api.services.tenders.categories import technical_soft_issues

    class LegacyBrief(BidBrief):
        def __getattribute__(self, name):
            if name in {
                "autoDraftTechPlan", "autoDraftBizResp", "techPlanSourceIds", "techPlanKbIds"
            }:
                raise AttributeError(name)
            return super().__getattribute__(name)

    brief = LegacyBrief(outlineItems=[OutlineItem(title="施工组织设计", kind="tech_plan")])
    dest = tmp_path / "legacy.docx"
    assemble_bid_docx(brief, dest, volume="technical")
    assert dest.is_file()
    assert technical_soft_issues(brief, media={})


def test_drafted_body_uses_existing_technical_section(tmp_path):
    brief = _brief()
    brief.outlineItems[0].draftBody = "本项目名称采用分区施工，人员数量【待补充】。"
    brief.outlineItems[0].draftStatus = "drafted"
    dest = tmp_path / "drafted.docx"
    assemble_bid_docx(brief, dest, volume="technical")
    body = "\n".join(p.text for p in Document(dest).paragraphs)
    assert "本项目名称采用分区施工" in body
    assert "人工写的质量体系" in body
    assert "凭空新增章节" not in body


@pytest.mark.parametrize('raw', [
    '{}', '{"chapters": null}', '{"chapters": {}}', '{"chapters": "missing"}',
    '{"chapters": []}', '[]', 'null', '', '未返回技术方案', '{"chapters": [',
])
def test_missing_or_invalid_chapters_decode_without_validation_error(raw):
    assert _decode_payload(raw, ['施工部署', '质量保证']) == ([], [])


def test_decoder_keeps_valid_chapters_and_deviations_only():
    payload = {
        'chapters': [
            {'title': '施工部署', 'body': '按现场条件分区组织施工。'},
            {'title': '施工部署', 'body': '重复正文不得覆盖前文。'},
            {'title': '质量保证'}, {'title': '安全生产', 'body': None},
            {'title': '进度计划', 'body': {}}, {'title': '新增目录', 'body': '不应写入'}, None,
        ],
        'deviations': [{'requirement': '现场安全措施', 'response': '待核实'}],
    }
    pairs, deviations = _decode_payload(json.dumps(payload, ensure_ascii=False),
                                        ['施工部署', '质量保证', '安全生产', '进度计划'])
    assert pairs == [('施工部署', '按现场条件分区组织施工。')]
    assert deviations == payload['deviations']
    assert _decode_payload(json.dumps({'deviations': deviations}), ['施工部署']) == ([], deviations)


@pytest.mark.parametrize('raw', ['{}', '{"chapters": null}', '未返回技术方案'])
def test_empty_technical_result_marks_existing_body_and_preserves_manual_sections(raw):
    class Client:
        async def complete(self, *args, **kwargs):
            return raw

    brief = _brief()
    original = brief.model_dump()
    result, notes = asyncio.run(draft_technical_plan(None, brief, client=Client()))
    assert '【待补充】' in result.outlineItems[0].draftBody
    assert result.outlineItems[0].draftStatus == 'drafted'
    assert result.outlineItems[1].draftBody == '人工写的质量体系'
    assert brief.model_dump() == original
    assert any('已填充 0 个既有章节' in note for note in notes)
    assert any('1 个章节' in note and '【待补充】' in note for note in notes)


def test_later_batches_continue_after_missing_technical_chapters():
    titles = ['施工部署', '主要施工方案', '质量保证', '安全生产', '进度计划']
    brief = BidBrief(outlineItems=[
        OutlineItem(id=str(i), title=title, kind='tech_chapter') for i, title in enumerate(titles)
    ] + [OutlineItem(id='manual', title='人工方案', kind='tech_chapter',
                      draftBody='人工内容', draftStatus='manual'),
         OutlineItem(id='skip', title='已跳过章节', kind='tech_chapter', skipped=True)])

    class Client:
        calls = 0

        async def complete(self, *args, **kwargs):
            replies = [
                {'chapters': [{'title': titles[0], 'body': '已生成施工部署正文。'}]},
                {'chapters': None},
                {'chapters': [{'title': titles[4], 'body': '已生成进度计划正文。'}]},
            ]
            reply = replies[self.calls]
            self.calls += 1
            return json.dumps(reply, ensure_ascii=False)

    client = Client()
    result, notes = asyncio.run(draft_technical_plan(None, brief, client=client))
    assert client.calls == 3
    assert '已生成施工部署正文。' in result.outlineItems[0].draftBody
    assert all('【待补充】' in i.draftBody for i in result.outlineItems[1:4])
    assert result.outlineItems[4].draftBody == '已生成进度计划正文。'
    assert result.outlineItems[5].draftBody == '人工内容'
    assert result.outlineItems[6].skipped and not result.outlineItems[6].draftBody
    assert any('已填充 2 个既有章节' in note for note in notes)
    assert any('3 个章节' in note and '【待补充】' in note for note in notes)


@pytest.mark.parametrize('volume', ['technical', 'full'])
def test_partial_technical_draft_is_written_to_copied_construction_format(tmp_path, monkeypatch, volume):
    source = Document()
    for text in ('投标函', '商务原文保留。', '施工组织设计', '原方案模板。',
                 '主要施工方案与技术措施', '原技术措施模板。'):
        source.add_paragraph(text)
    source_path = tmp_path / 'construction.format.docx'
    source.save(source_path)
    original = source_path.read_bytes()
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: source_path)
    monkeypatch.setattr('api.services.tenders.gen_tech_plan.load_invitation_text',
                        lambda _: '本项目施工招标文件要求提供施工组织设计和主要施工技术措施。')
    brief = BidBrief(autoDraftTechPlan=True, generateVolume=volume, layoutMode='outline',
                     invitationId='abc123def456', attachQualifications=False, includePlaceholders=False,
                     outlineItems=[OutlineItem(id='letter', title='投标函', kind='letter'),
                                   OutlineItem(id='plan', title='施工组织设计', kind='tech_chapter'),
                                   OutlineItem(id='measures', title='主要施工方案与技术措施', kind='tech_chapter')])
    client = FakeClient({'chapters': [{'title': '施工组织设计', 'body': '本项目分区组织施工。'}]})
    prepared, notes = asyncio.run(draft_technical_plan(None, brief, client=client))
    path, _ = assemble_bid_docx(prepared, tmp_path / f'{volume}.docx', volume=volume,
                                catalog_slots=[], catalog_media={}, catalog_facts=[])
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert '本项目分区组织施工。' in text
    assert '主要施工方案与技术措施' in text and '【待补充】' in text
    assert text.index('主要施工方案与技术措施') < text.index('「主要施工方案与技术措施」')
    assert ('商务原文保留。' in text) == (volume == 'full')
    assert source_path.read_bytes() == original
    assert any('【待补充】' in note for note in notes)


def test_missing_chapters_retains_deviations_and_preserves_manual_rows():
    from api.services.tenders.schema import DeviationLine

    brief = _brief()
    manual = DeviationLine(requirement='用户确认的技术参数', response='人工响应', deviation='待确认')
    brief.deviationLines = [manual]
    client = FakeClient({'deviations': [{
        'requirement': '说明现场安全措施', 'response': '按现场核查结果制定安全措施',
        'chapter': '项目概况与施工部署', 'evidence': '用户安全资料【待补充】', 'deviation': '待确认',
    }]})
    result, _ = asyncio.run(draft_technical_plan(None, brief, client=client))
    assert result.deviationLines[0].model_dump() == manual.model_dump()
    assert '按现场核查结果制定安全措施' in result.deviationLines[1].response
    assert '【待补充】' in result.outlineItems[0].draftBody


def test_model_authentication_errors_are_not_masked_as_missing_chapters():
    class Client:
        async def complete(self, *args, **kwargs):
            raise AppError(msg='模型鉴权失败', status_code=401)

    with pytest.raises(AppError) as raised:
        asyncio.run(draft_technical_plan(None, _brief(), client=Client()))
    assert raised.value.status_code == 401
