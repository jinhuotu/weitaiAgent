"""按组卷大纲组装投标 Word：目录跟随大纲，未知函/表复制招标书原文并只填空。"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from api.services.tenders.document import (
    _add_bottom_sign_spacer,
    _apply_toc_page_numbers,
    _bookmark_paragraph,
    _ensure_header,
    _estimate_bookmark_pages,
    _p_sectpr,
    _paragraph_has_page_break,
    _pin_cjk_fonts,
    _rewrite_date_line,
    _rewrite_labeled_underline,
    _sect_starts_new_page,
    _set_word_wrap,
    _ymd,
    _em_len,
    _usable_width_twips,
    _write_toc_line,
)
from api.services.tenders.format_rules import (
    apply_footer_page_number,
    apply_section_page,
    flatten_toc_tree,
    format_brief_notes,
    section_page_flags,
    toc_page_cache,
)
from api.services.tenders.modules import COPY_KINDS, MODULE_KINDS, is_seal_register, render_module
from api.services.tenders.outline import (
    bid_item_title,
    build_outline_toc_tree,
    compact_title,
    dedupe_outline_items,
    drop_ocr_junk_lines,
    ensure_biz_essentials,
    fill_copy_blanks,
    is_outline_junk,
    items_for_volume,
    looks_like_form_template,
    looks_like_quote_form,
    split_mashed_zhi_line,
    split_zhi_company_body,
    unfold_form_sign_lines,
    unmash_invitation_text,
)
from api.services.tenders.placeholders import (
    collect_slots,
    draw_placeholder_box,
    id_slot,
    inline_id_scans,
    _insert_slot_media,
)
from api.services.tenders.schema import BidBrief, DocumentFormat, OutlineItem, PlaceholderItem
from api.services.tenders.slots import attachments_for_slots

logger = logging.getLogger("api.tenders.assemble")

_SONG = "宋体"
# 函/授权/承诺书/印鉴表/邀请书清单有原文就复制，不再换成充电桩模块或自绘格子。
_FORM_LAYOUT_KINDS = frozenset({"letter", "auth", "factory", "commitment_copy"})
_COPY_BODY_KINDS = frozenset({"letter", "auth", "factory", "commitment_copy", "company", "unknown"})
_SIGN_LABEL = re.compile(
    r"^(投标人名称|供应商名称|投标人|供应商|承诺单位|投标单位|"
    r"法定代表人或授权代表|"
    r"法定代表人或其委托代理人|法定代表人或委托代理人|法定代表人|"
    r"委托代理人|授权代表|联系人|联系电话|电话|地址|住址)"
    r"(?:[（(]([^)）]*)[)）])?[：:]*(.*)$"
)
_SIGN_TOKEN = re.compile(
    r"(投标人|供应商|法定代表人|授权代表|委托代理人|联系人|联系电话|电话|地址|日期|"
    r"（章）|\(章\)|签字|签名|盖章|公章)"
)
_DATE_YMD = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日")
_SIGN_KW = dict(
    align=WD_ALIGN_PARAGRAPH.LEFT,
    left_indent_cm=0,
    right_indent_cm=0.4,
    line_spacing=1.5,
    space_before=8,
    space_after=4,
    nowrap=True,
    line_em=16,
)
_ATTACH_LINE = re.compile(
    r"^(附(?:件)?[（(]?[一二三四五六七八九十0-9]+[)）]?[：:]?)(.*)$"
)
_QUAL_PACK_MARKS = ("企业资质", "资质证明", "资质文件", "资格审查资料", "资格证明资料")
_SECTION_HEAD = re.compile(r"^[(（][0-9一二三四五六七八九十]+[)）]")
_INLINE_FORMAT_HEAD = re.compile(
    r"第[一二三四五六七八九十0-9]+[章节部分][^。\n]{0,24}(?:投标文件格式|响应文件格式)"
)
_REQ_NAME = (
    r"(?:设备数采|系统拓展性|[A-Za-z]{2,12}系统|"
    r"[\u4e00-\u9fff]{2,10}(?:系统|数采|拓展性))"
)
_REQ_ROW = re.compile(rf"^({_REQ_NAME})[：:\s]+(.+)$")
_REQ_SPLIT = re.compile(
    rf"(?=(?:(?<=\s)|(?<=；)|(?<=;))(?:{_REQ_NAME})[：:\s]*[1１一])"
)
_REQ_GROUP = re.compile(
    r"^[一二三四五六七八九十]、.{2,40}[：:]$"
)


def assemble_bid_docx(
    brief: BidBrief,
    dest: Path,
    *,
    qualification_pdf: Path | None = None,
    catalog_slots: list | None = None,
    catalog_media: dict | None = None,
    volume: str | None = None,
) -> tuple[Path, list[str]]:
    warnings: list[str] = []
    vol = (volume or "").strip()
    if vol in {"business", "technical"}:
        items = items_for_volume(brief, vol)
    else:
        # 目录与正文都按 outlineItems 当前顺序；跳过项不入卷。前端可重排该数组。
        items = [
            item
            for item in (brief.outlineItems or [])
            if not item.skipped and (item.kind or "").strip() != "" and not is_outline_junk(item)
        ]
        items = dedupe_outline_items(items)
        if (brief.layoutMode or "").strip() == "outline":
            items = ensure_biz_essentials(items)
            items = [item for item in items if not item.skipped]
    if not items:
        warnings.append("组卷大纲为空或已全部跳过，未按大纲生成")
        if (brief.layoutMode or "").strip() != "outline":
            from api.services.tenders.document import build_bid_docx

            return build_bid_docx(
                brief,
                dest,
                qualification_pdf=qualification_pdf,
                catalog_slots=catalog_slots,
                catalog_media=catalog_media,
            )

    dest.parent.mkdir(parents=True, exist_ok=True)
    fmt = brief.documentFormat or DocumentFormat()
    doc = Document()
    apply_section_page(doc.sections[0], fmt)
    _pin_cjk_fonts(doc)
    _write_cover(doc, brief, fmt, volume=vol)
    _add_next_section(doc, fmt)
    bookmarks = [f"toc_{(item.id or f'o{i + 1:02d}')}" for i, item in enumerate(items)]
    _write_toc(doc, items, brief, fmt, bookmarks)
    body_section_index = 1
    if fmt.pageNumberStart == "body":
        _add_next_section(doc, fmt)
        body_section_index = 2
    copied = 0
    generated = 0
    title_only: list[str] = []
    quote_done = False
    slots = catalog_slots if catalog_slots is not None else collect_slots(brief.extraPlaceholders)
    media = catalog_media if catalog_media is not None else attachments_for_slots(slots)
    inlined_ids: set[str] = set()
    qual_tried = False
    for i, item in enumerate(items):
        kind = (item.kind or "unknown").strip() or "unknown"
        if kind == "quote" and quote_done:
            continue
        if kind == "quote":
            quote_done = True
        if not (i == 0 and fmt.pageNumberStart == "body"):
            _chapter_break(doc)
        source = (item.source or "").strip() or "copy"
        raw_body = (item.body or "").strip()
        if brief.attachQualifications and not qual_tried and _is_qual_pack_item(item) and vol != "technical":
            from api.services.tenders.document import qualification_attach_notes

            _write_heading(doc, item.title, fmt, bookmarks[i])
            notes = qualification_attach_notes(
                doc, qualification_pdf, title=False, leading_break=False
            )
            for note in notes:
                if note not in warnings:
                    warnings.append(note)
            if not any("已插入资质文件" in n for n in notes):
                draw_placeholder_box(
                    doc,
                    PlaceholderItem(key=item.id or "scan", title=item.title, hint="装订时附原件"),
                )
            qual_tried = True
            continue
        if kind == "scan":
            _write_heading(doc, item.title, fmt, bookmarks[i])
            if not _try_write_scan_media(doc, item, media):
                draw_placeholder_box(
                    doc,
                    PlaceholderItem(
                        key=_slot_key_for_item(item) or item.id or "scan",
                        title=item.title,
                        hint="装订时附原件",
                    ),
                )
            continue
        keep_form = kind in _FORM_LAYOUT_KINDS and looks_like_form_template(raw_body)
        has_body = bool(raw_body)
        quote_form = kind == "quote" and has_body and looks_like_quote_form(raw_body)
        # 扫描件不当正文抄须知；身份证明用模块排版
        use_copy = (
            kind not in {"legal_id", "scan", "performance"}
            and (not is_seal_register(item) or has_body)
            and (
                source == "copy"
                or kind in COPY_KINDS
                or keep_form
                or quote_form
                or (kind in _COPY_BODY_KINDS and has_body)
            )
        )
        use_module = (not use_copy) and kind in MODULE_KINDS
        bm = bookmarks[i]
        if use_module:
            _write_heading(doc, item.title, fmt, bm)
            for note in render_module(doc, brief, item, media=media, inlined=inlined_ids):
                if note not in warnings:
                    warnings.append(note)
            generated += 1
            continue
        contact = (brief.agentName or "").strip() or (brief.legalPersonName or "").strip()
        filled = fill_copy_blanks(
            raw_body,
            bidder=brief.bidderName,
            project=brief.projectName,
            tenderer=brief.tenderer,
            legal=brief.legalPersonName,
            bid_date=brief.bidDate,
            phone=brief.bidderPhone,
            address=brief.bidderAddress,
            contact=contact,
            delivery_days=brief.deliveryDays,
            validity_days=brief.bidValidityDays,
            quality=brief.quality,
        ).strip()
        if filled:
            copied_ok = False
            if keep_form or looks_like_form_template(filled):
                before = _visible_text_len(doc)
                start = len(doc.paragraphs)
                _write_copied_form(doc, filled, fmt, bid_date=brief.bidDate)
                copied_ok = _visible_text_len(doc) - before >= 20
                if copied_ok:
                    _bookmark_first_new(doc, start, bm)
            if copied_ok:
                if kind == "auth":
                    _inline_copied_id(
                        doc,
                        start=start,
                        key="id_agent",
                        media=media,
                        inlined=inlined_ids,
                        has_agent=bool((brief.agentName or "").strip()),
                    )
                copied += 1
                continue
            if kind in MODULE_KINDS and not has_body and kind not in COPY_KINDS:
                _write_heading(doc, item.title, fmt, bm)
                for note in render_module(doc, brief, item, media=media, inlined=inlined_ids):
                    if note not in warnings:
                        warnings.append(note)
                generated += 1
                continue
            start = len(doc.paragraphs)
            _write_heading(doc, item.title, fmt, bm)
            _write_body(doc, filled, fmt)
            _bookmark_first_new(doc, start, bm)
            copied += 1
            continue
        title_only.append(item.title)
        _write_heading(doc, item.title, fmt, bm)
        if kind == "scan":
            draw_placeholder_box(
                doc,
                PlaceholderItem(key=item.id or "scan", title=item.title, hint="装订时附原件"),
            )
        elif is_seal_register(item):
            for note in render_module(doc, brief, item, media=media, inlined=inlined_ids):
                if note not in warnings:
                    warnings.append(note)
            generated += 1
        elif kind == "commitment_copy":
            if not _try_write_scan_media(doc, item, media):
                draw_placeholder_box(
                    doc,
                    PlaceholderItem(
                        key=_slot_key_for_item(item) or item.id or "commit",
                        title=item.title,
                        hint="按邀请书本页填写后装订",
                    ),
                )
            warnings.append(f"「{item.title}」未抽到招标书原文，已留标题和粘贴框")
        else:
            warnings.append(f"「{item.title}」招标书格式章未提供本页空白稿，仅保留标题")
    if generated:
        warnings.append(f"已按组卷大纲用模块填写 {generated} 节")
    if copied:
        warnings.append(f"已按组卷大纲复制 {copied} 节招标书原文并保留格式，仅填入单位、项目和日期")
    if title_only:
        warnings.append("仅保留标题、未复制到正文：" + "、".join(title_only[:12]))
    warnings.extend(format_brief_notes(fmt))
    if fmt.coverNeedSeal:
        warnings.append("招标书要求封面加盖公章，请在打印后于封面预留处盖章")

    if brief.attachQualifications and not qual_tried and vol != "technical":
        from api.services.tenders.document import qualification_attach_notes

        for note in qualification_attach_notes(doc, qualification_pdf):
            if note not in warnings:
                warnings.append(note)

    for si, section in enumerate(doc.sections):
        numbered, restart = section_page_flags(si, fmt, body_section=body_section_index)
        _ensure_header(section, brief.projectName, cover=(si == 0))
        apply_footer_page_number(section, fmt, numbered=numbered, restart=restart)

    _collapse_extra_page_breaks(doc)
    _fill_toc_pages(doc, fmt)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest, warnings


def _add_next_section(doc: Document, fmt: DocumentFormat):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section_page(section, fmt)
    return section


def _clear_keep_next(p_el) -> None:
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is None:
        return
    for el in p_pr.findall(qn("w:keepNext")):
        p_pr.remove(el)


def _chapter_break(doc: Document) -> None:
    """上一节已经是分页/新节时不要再插一次，否则会多出一页只有页眉。"""
    body = list(doc.element.body)
    for child in reversed(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:tbl"):
            break
        if child.tag != qn("w:p"):
            continue
        if _paragraph_has_page_break(child):
            return
        sect = _p_sectpr(child)
        if sect is not None and _sect_starts_new_page(sect):
            return
        text = "".join(node.text or "" for node in child.iter(qn("w:t"))).replace("\u200b", "").strip()
        if text or _p_has_drawing(child):
            # 末段 keep_with_next 会把分页符粘走，整页只剩页眉
            _clear_keep_next(child)
            break
        _clear_keep_next(child)
    doc.add_page_break()


def _p_has_drawing(el) -> bool:
    return next(el.iter(qn("w:drawing")), None) is not None


_SCAN_SLOT = (
    ("bond", ("保证金", "保函")),
    ("license", ("营业执照",)),
    ("bank_permit", ("开户许可", "开户许可证", "基本户")),
    ("iso", ("资质证书", "资质证明", "企业资质", "体系认证", "ISO")),
    ("finance", ("财务", "审计", "完税")),
    ("perf", ("业绩", "合同及发票")),
    ("commitment", ("无违法", "无行贿", "承诺函", "承诺书")),
    ("credit", ("信用", "失信")),
    ("product", ("检测报告", "3C", "型式试验")),
)


def _slot_key_for_item(item: OutlineItem) -> str:
    n = compact_title(item.title)
    for key, marks in _SCAN_SLOT:
        if any(m in n for m in marks):
            return key
    return ""


def _media_keys_for_item(item: OutlineItem) -> list[str]:
    key = _slot_key_for_item(item)
    keys: list[str] = []
    if key:
        keys.append(key)
    if key == "commitment" and "credit" not in keys:
        keys.append("credit")
    return keys


def _try_write_scan_media(doc: Document, item: OutlineItem, media: dict | None) -> bool:
    files = []
    if isinstance(media, dict):
        for key in _media_keys_for_item(item):
            files = [p for p in (media.get(key) or []) if getattr(p, "is_file", lambda: False)()]
            if files:
                break
        if not files:
            files = [p for p in (media.get(item.id) or []) if getattr(p, "is_file", lambda: False)()]
    if not files:
        return False
    n = _insert_slot_media(doc, files[:4], max_pages=2)
    return n > 0


def _is_qual_pack_item(item: OutlineItem) -> bool:
    kind = (item.kind or "").strip()
    if kind not in {"scan", "company"}:
        return False
    n = compact_title(item.title)
    return any(k in n for k in _QUAL_PACK_MARKS)


def _p_plain(el) -> str:
    return "".join(node.text or "" for node in el.iter(qn("w:t"))).replace("\u200b", "").strip()


def _is_blank_para(el) -> bool:
    return el.tag == qn("w:p") and not _p_plain(el) and not _p_has_drawing(el)


def _is_empty_break_para(el) -> bool:
    return _is_blank_para(el) and _paragraph_has_page_break(el)


def _strip_page_br(el) -> None:
    for run in list(el.iter(qn("w:r"))):
        for br in list(run.findall(qn("w:br"))):
            if br.get(qn("w:type")) == "page":
                run.remove(br)


def _set_el_page_break_before(el, on: bool = True) -> None:
    p_pr = el.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        el.insert(0, p_pr)
    for old in list(p_pr.findall(qn("w:pageBreakBefore"))):
        p_pr.remove(old)
    if on:
        node = OxmlElement("w:pageBreakBefore")
        node.set(qn("w:val"), "true")
        p_pr.append(node)


def _collapse_extra_page_breaks(doc: Document) -> None:
    """空段里的 w:br page 在预览里会单独占一页只剩页眉；改挂到下一节有内容的段落上。"""
    body = doc.element.body
    guard = 0
    while guard < 80:
        guard += 1
        target = None
        for child in list(body):
            if child.tag == qn("w:sectPr"):
                continue
            if _is_empty_break_para(child):
                target = child
                break
        if target is None:
            break
        nxt = target.getnext()
        while nxt is not None and nxt.tag == qn("w:p") and _is_blank_para(nxt):
            dead = nxt
            nxt = nxt.getnext()
            try:
                body.remove(dead)
            except ValueError:
                break
        if nxt is not None and nxt.tag == qn("w:p") and not _is_blank_para(nxt):
            _set_el_page_break_before(nxt, True)
            try:
                body.remove(target)
            except ValueError:
                break
            continue
        if nxt is not None and nxt.tag == qn("w:tbl"):
            _strip_page_br(target)
            _set_el_page_break_before(target, True)
            continue
        try:
            body.remove(target)
        except ValueError:
            break

    prev_new_page = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            prev_new_page = False
            continue
        if child.tag == qn("w:tbl"):
            prev_new_page = False
            continue
        if child.tag != qn("w:p"):
            continue
        p_pr = child.find(qn("w:pPr"))
        has_before = p_pr is not None and p_pr.find(qn("w:pageBreakBefore")) is not None
        sect = _p_sectpr(child)
        if has_before and prev_new_page:
            _set_el_page_break_before(child, False)
            has_before = False
        if _p_plain(child) or _p_has_drawing(child):
            prev_new_page = bool(has_before or (sect is not None and _sect_starts_new_page(sect)))
        elif _paragraph_has_page_break(child) or has_before:
            prev_new_page = True


def _fill_toc_pages(doc: Document, fmt: DocumentFormat) -> None:
    raw = _estimate_bookmark_pages(doc)
    if not raw:
        return
    if fmt.pageNumberStart == "body":
        base = min(raw.values())
        pages = {k: max(1, v - base + 1) for k, v in raw.items()}
    elif fmt.pageNumberStart == "cover":
        pages = {k: max(1, v) for k, v in raw.items()}
    else:
        pages = {k: max(1, v - 1) for k, v in raw.items()}
    _apply_toc_page_numbers(doc, pages)


def _bookmark_first_new(doc: Document, start: int, name: str) -> None:
    for para in doc.paragraphs[start:]:
        if (para.text or "").strip():
            _bookmark_paragraph(para, name)
            return


def _write_cover(doc: Document, brief: BidBrief, fmt: DocumentFormat, *, volume: str = "") -> None:
    font = fmt.fontName
    if fmt.coverShowCopyMark:
        mark = doc.add_paragraph()
        mark.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pf = mark.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        _run(mark, fmt.coverCopyMark or "正本", size=16, bold=True, font=font)
    for _ in range(4):
        doc.add_paragraph()
    if fmt.coverShowProject:
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(
            title,
            (brief.projectName or "投标项目").strip() or "投标项目",
            size=fmt.coverTitleSizePt,
            bold=True,
            font=font,
        )
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if volume == "technical":
        doc_title = "技 术 标 投 标 文 件"
    elif volume == "business":
        doc_title = "商 务 标 投 标 文 件"
    else:
        doc_title = "投 标 文 件"
    _run(sub, doc_title, size=fmt.coverDocSizePt, bold=True, font=font)
    # 标题留在上半页；招标人/投标人/日期落到封面下半区，不贴页脚。
    _add_bottom_sign_spacer(doc, sign_cm=3.8, bottom_pad_cm=3.6)
    if fmt.coverShowTenderNo:
        no = (brief.tenderNo or "").strip() or "—"
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(line, f"招标编号：{no}", size=14, font=font)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(meta, f"招标人：{(brief.tenderer or '').strip() or '—'}", size=14, font=font)
    if fmt.coverShowBidder:
        meta2 = doc.add_paragraph()
        meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(meta2, f"投标人：{(brief.bidderName or '').strip() or '—'}", size=14, font=font)
    if fmt.coverShowDate:
        meta3 = doc.add_paragraph()
        meta3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_cn = fill_copy_blanks(
            "年 月 日",
            bidder="",
            project="",
            tenderer="",
            legal="",
            bid_date=brief.bidDate,
        )
        _run(meta3, date_cn, size=14, font=font)
    if fmt.coverNeedSeal:
        seal = doc.add_paragraph()
        seal.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(seal, "（封面加盖公章）", size=12, font=font)


def _write_toc(
    doc: Document,
    items: list[OutlineItem],
    brief: BidBrief,
    fmt: DocumentFormat,
    bookmarks: list[str],
) -> None:
    from api.services.tenders.categories import tech_plan_text

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(head, "目录", size=fmt.tocTitleSizePt, bold=True, font=fmt.fontName)
    tree = build_outline_toc_tree(items, bookmarks, tech_text=tech_plan_text(brief))
    rows = flatten_toc_tree(tree, fmt.tocNumbering)
    for i, (label, title, level, bm) in enumerate(rows):
        para = doc.add_paragraph()
        _write_toc_line(
            para,
            i,
            bid_item_title(title),
            bm,
            toc_page_cache(i, fmt),
            label=label,
            with_pages=True,
            level=level,
        )


def _inline_copied_id(
    doc: Document,
    *,
    start: int,
    key: str,
    media: dict | None,
    inlined: set[str],
    has_agent: bool,
) -> None:
    if key == "id_agent" and not has_agent:
        return
    before = None
    for p in doc.paragraphs[start:]:
        t = "".join((p.text or "").split())
        if t.startswith("投标人") and (
            "盖单位公章" in t or "（章）" in t or "(章)" in t or "签字" in t
        ):
            before = p
            break
    n = inline_id_scans(
        doc,
        (media or {}).get(key) if media else None,
        empty_slot=id_slot(key),
        before=before,
    )
    inlined.add(key)


def _write_heading(doc: Document, title: str, fmt: DocumentFormat, bookmark: str | None = None):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(12)
    _run(
        para,
        (bid_item_title(title) or "").strip() or "附件",
        size=fmt.headingSizePt,
        bold=True,
        font=fmt.fontName,
    )
    if bookmark:
        _bookmark_paragraph(para, bookmark)
    return para


def _write_copied_salute(doc: Document, block: str) -> None:
    n = compact_title(block)
    if not n.startswith("致"):
        _form_para(doc, block, size=12, align="left", space_after=8)
        return
    rest = n[1:].lstrip("：:")
    co, body = split_zhi_company_body(rest)
    if not body:
        co = re.sub(r"[＿_—\-－]+", "", rest)
        body = ""
    para = doc.add_paragraph()
    _rewrite_labeled_underline(
        para,
        [("致：", co, "")],
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_indent_cm=0,
        right_indent_cm=0.4,
        line_spacing=1.5,
        space_before=6,
        space_after=10,
        nowrap=True,
        line_em=22,
    )
    if body:
        _form_para(
            doc,
            body,
            size=12,
            align="justify",
            indent=True,
            space_after=6,
            line_spacing=1.5,
        )


def _ymd_in(text: str) -> tuple[str, str, str] | None:
    m = _DATE_YMD.search(text or "")
    if not m:
        return None
    return m.group(1), m.group(2), m.group(3)


def _is_date_sign_line(text: str) -> bool:
    n = compact_title(text)
    if "日期" in n or re.fullmatch(r"年\s*月\s*日", (text or "").strip()):
        return True
    return bool(_DATE_YMD.search(text or "")) and len(n) <= 22


def _write_copied_sign_line(doc: Document, block: str, *, bid_date: str = "") -> None:
    raw = (block or "").strip()
    n = compact_title(raw)
    if _is_date_sign_line(raw):
        para = doc.add_paragraph()
        ymd = _ymd_in(raw)
        if not ymd and (bid_date or "").strip():
            y, m, d = _ymd(bid_date)
            if y and any(ch.isdigit() for ch in y):
                ymd = (y, m, d)
        if ymd:
            _rewrite_date_line(
                para,
                ymd[0],
                ymd[1],
                ymd[2],
                **_SIGN_KW,
                label="日期：",
                label_width=5,
            )
            return
        _rewrite_labeled_underline(
            para,
            [("日期：", "", "")],
            **_SIGN_KW,
            label_width=5,
        )
        return
    matched = _SIGN_LABEL.match(n)
    if not matched:
        _form_para(doc, block, size=12, align="left", space_after=2, line_spacing=1.15)
        return
    label = matched.group(1) + "："
    rest = (matched.group(3) or "").lstrip("：:")
    suffix = f"（{matched.group(2)}）" if matched.group(2) else ""
    if not suffix:
        lead = re.match(r"^[（(]([^)）]*)[)）][：:]*(.*)$", rest)
        if lead:
            suffix = f"（{lead.group(1)}）"
            rest = lead.group(2) or ""
        else:
            suf_m = re.search(r"([（(][^)）]*[)）])\s*$", rest)
            suffix = suf_m.group(1) if suf_m else ""
            rest = rest[: suf_m.start()] if suf_m else rest
    else:
        suf_m = re.search(r"([（(][^)）]*[)）])\s*$", rest)
        if suf_m:
            rest = rest[: suf_m.start()]
    blob = f"{label}{suffix}{rest}"
    value = "" if any(k in blob for k in ("签字", "签名", "手签")) else re.sub(r"[＿_—\-－\s：:]+", "", rest)
    para = doc.add_paragraph()
    _rewrite_labeled_underline(
        para,
        [(label, value, suffix)],
        **_SIGN_KW,
        label_width=5 if len(matched.group(1)) <= 4 else None,
    )


def _absorb_date_follow(lines: list[str], i: int, block: str) -> tuple[str, int]:
    if "日期" not in compact_title(block) or _ymd_in(block):
        return block, 0
    if i + 1 >= len(lines):
        return block, 0
    nxt = lines[i + 1]
    if not _ymd_in(nxt) or "日期" in compact_title(nxt):
        return block, 0
    return f"{block.rstrip('：: ')}：{nxt.strip()}", 1


def _write_copied_form(doc: Document, text: str, fmt: DocumentFormat, *, bid_date: str = "") -> None:
    """按招标书空白稿的行角色还原标题、缩进和落款，而不是整页同一缩进。"""
    body = fmt.bodySizePt
    font = fmt.fontName
    lines = [_strip_inline_format_head(b) for b in _paragraphs_from_body(text)]
    i = 0
    while i < len(lines):
        block = lines[i]
        if _is_invite_attach_label(block):
            i += 1
            continue
        nxt = _try_write_pipe_table(doc, lines, i, fmt=fmt)
        if nxt is not None:
            i = nxt
            continue
        cells = _header_cells(block)
        if cells:
            rows, i = _collect_table_rows(lines, i + 1, len(cells))
            _write_simple_table(doc, cells, rows, font=font)
            continue
        role = _form_line_role(block)
        if role == "skip":
            i += 1
            continue
        nxt = _try_write_req_table(doc, lines, i)
        if nxt is not None:
            i = nxt
            continue
        if role == "attach":
            if _is_invite_attach_label(block):
                i += 1
                continue
            matched = _ATTACH_LINE.match(block.strip())
            label = (matched.group(1) if matched else block).strip()
            rest = (matched.group(2) if matched else "").strip()
            _form_para(doc, label, size=body, bold=False, align="left", space_after=2, font=font)
            if rest:
                _form_para(
                    doc,
                    _spaced_title(rest),
                    size=fmt.headingSizePt,
                    bold=True,
                    align="center",
                    space_before=8,
                    space_after=12,
                    font=font,
                )
            i += 1
            continue
        if role == "title":
            _form_para(
                doc,
                _spaced_title(block),
                size=fmt.headingSizePt,
                bold=True,
                align="center",
                space_before=10,
                space_after=14,
                font=font,
            )
            i += 1
            continue
        if role == "salute":
            _write_copied_salute(doc, block)
            i += 1
            continue
        if role == "sign_row":
            for col in _split_sign_columns(block) or [block]:
                _write_copied_sign_line(doc, col, bid_date=bid_date)
            i += 1
            continue
        if role == "sign":
            merged, extra = _absorb_date_follow(lines, i, block)
            _write_copied_sign_line(doc, merged, bid_date=bid_date)
            i += 1 + extra
            continue
        cap = _split_unit_caption(block)
        if cap:
            _write_unit_caption(doc, cap[0], cap[1], fmt)
            i += 1
            continue
        indent = block.startswith(("我方确认", "兹向", "兹"))
        _form_para(
            doc,
            block,
            size=body,
            align="justify" if indent else "left",
            indent=indent,
            space_after=6,
            line_spacing=1.15,
            font=font,
        )
        i += 1


def _strip_inline_format_head(text: str) -> str:
    raw = (text or "").strip()
    if not raw:
        return ""
    stripped, n = _INLINE_FORMAT_HEAD.subn("", raw, count=1)
    if n == 0:
        return raw
    return stripped.strip(" ：:、;；") or raw


def _visible_text_len(doc: Document) -> int:
    n = sum(len((p.text or "").strip()) for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                n += len((cell.text or "").strip())
    return n


def _form_line_role(line: str) -> str:
    raw = (line or "").strip()
    if not raw or re.fullmatch(r"[-—–_＿]{6,}", raw):
        return "skip"
    n = compact_title(raw)
    if len(n) <= 40 and (n.endswith("投标文件") or "投标文件格式" in n or "响应文件格式" in n):
        return "skip"
    attach = _ATTACH_LINE.match(raw)
    if attach and len(n) <= 48:
        return "attach"
    compact = compact_title(raw)
    if (
        "：" not in raw
        and ":" not in raw
        and 2 <= len(compact) <= 8
        and re.fullmatch(r"[\u4e00-\u9fff]+", compact)
        and (compact.endswith(("函", "书")) or compact in {"竞标书"})
    ):
        return "title"
    if _split_sign_columns(raw):
        return "sign_row"
    if compact.startswith("致") or (
        (raw.endswith("：") or raw.endswith(":"))
        and ("有限公司" in compact or compact.startswith("致"))
        and len(compact) <= 40
        and not compact.startswith(("投标承诺书", "承诺函", "投标函"))
    ):
        return "salute"
    sign_hit = _SIGN_TOKEN.search(compact) or _SIGN_TOKEN.search(raw)
    if sign_hit and (
        len(compact) <= 48
        or compact.startswith(("投标人", "供应商", "法定代表人", "日期", "委托代理人"))
    ):
        return "sign"
    if re.search(r"年\S{0,6}月\S{0,6}日", compact) and len(compact) <= 28:
        return "sign"
    return "body"


def _spaced_title(text: str) -> str:
    compact = compact_title(text)
    if re.fullmatch(r"(?:\S\s)+\S", (text or "").strip()):
        return (text or "").strip()
    if (
        2 <= len(compact) <= 4
        and re.fullmatch(r"[\u4e00-\u9fff]+", compact)
        and compact.endswith(("函", "书"))
    ):
        return " ".join(compact)
    return (text or "").strip()


def _split_sign_columns(line: str) -> list[str] | None:
    parts = [p.strip() for p in re.split(r"\s{2,}", line or "") if p.strip()]
    if len(parts) < 2:
        return None
    hits = sum(1 for p in parts if _SIGN_TOKEN.search(p) or re.search(r"年\S{0,4}月", p))
    if hits < 2:
        return None
    if len(parts) == 2:
        return parts
    return ["  ".join(parts[:-1]), parts[-1]]


def _write_sign_columns(doc: Document, cols: list[str], fmt: DocumentFormat) -> None:
    table = doc.add_table(rows=1, cols=max(2, len(cols)))
    table.autofit = True
    _set_table_no_borders(table)
    row = table.rows[0]
    for i, text in enumerate(cols[: len(row.cells)]):
        cell = row.cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.15
        _run(para, text, size=fmt.bodySizePt, bold=False, font=fmt.fontName)


def _set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for old in list(tbl_pr.findall(qn("w:tblBorders"))):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)


def _form_para(
    doc: Document,
    text: str,
    *,
    size: float = 12,
    bold: bool = False,
    align: str = "left",
    indent: bool = False,
    space_before: float = 0,
    space_after: float = 6,
    line_spacing: float = 1.5,
    font: str = _SONG,
):
    para = doc.add_paragraph()
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = para.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if indent:
        pf.first_line_indent = Cm(0.74)
    _run(para, text, size=size, bold=bold, font=font)
    return para


def _pipe_cells(line: str) -> list[str] | None:
    s = (line or "").strip()
    if "|" not in s:
        return None
    keep_lead = s.startswith("|")
    cells = [c.strip() for c in s.split("|")]
    if not keep_lead and cells and cells[0] == "":
        cells = cells[1:]
    if len(cells) < 2 or not any(cells):
        return None
    return cells


def _trim_trailing_empty(cells: list[str]) -> list[str]:
    out = list(cells)
    while len(out) > 2 and not (out[-1] or "").strip():
        out.pop()
    return out


def _unglue_header_cells(cells: list[str]) -> list[str]:
    marks = ("实现目标", "建设要求", "模块")
    out: list[str] = []
    for c in cells:
        hit = next((h for h in marks if c.startswith(h) and len(c) > len(h)), None)
        if hit:
            rest = c[len(hit) :].strip()
            out.append(hit)
            if rest:
                out.append(rest)
        else:
            out.append(c)
    return out


_NOTE_SITE = re.compile(r"^(备注)([\u4e00-\u9fff]{2,8})$")
_QUOTE_CAP_SEQ = re.compile(
    r"^((?:（[0-9一二三四五六七八九十]+）)?[\u4e00-\u9fffA-Za-z]{1,16}报价)(序号)$"
)


def _peel_note_site(cells: list[str]) -> tuple[list[str], list[str] | None]:
    if not cells:
        return cells, None
    m = _NOTE_SITE.match((cells[-1] or "").strip())
    if not m:
        return cells, None
    head = cells[:-1] + [m.group(1)]
    site = [m.group(2)] + [""] * (len(head) - 1)
    return head, site


def _peel_quote_caption(cells: list[str]) -> tuple[list[str], str]:
    if not cells:
        return cells, ""
    m = _QUOTE_CAP_SEQ.match((cells[0] or "").strip())
    if not m:
        return cells, ""
    return [m.group(2)] + list(cells[1:]), m.group(1)


def _pipe_header_row(cells: list[str]) -> bool:
    blob = compact_title("".join(cells))
    if not any(
        k in blob
        for k in (
            "模块",
            "建设要求",
            "实现目标",
            "序号",
            "项目内容",
            "金额",
            "数据采集",
            "品牌",
            "单价",
        )
    ):
        return False
    return all(len(compact_title(c)) <= 18 for c in cells if (c or "").strip())


def _cell_chunks(text: str) -> list[str]:
    s = (text or "").strip()
    if len(s) < 80:
        return [s] if s else [""]
    parts = re.split(
        r"(?<=[。；;])\s*(?=(?:\d+[．.、]|（[一二三四五六七八九十0-9]+）|[一二三四五六七八九十]、))",
        s,
    )
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) <= 1 and len(s) > 240:
        parts = [p.strip() for p in re.split(r"(?<=[。；;])", s) if p.strip()]
    return parts or [s]


def _req_item_lines(text: str) -> list[str]:
    s = (text or "").strip()
    if not s:
        return [""]
    parts = [p.strip() for p in re.split(r"(?<=[；;])\s*", s) if p.strip()]
    return parts or [s]


def _cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    center: bool = False,
    valign: str = "center",
    items: bool = False,
    font: str = _SONG,
    wrap: bool = True,
) -> None:
    cell.text = ""
    chunks = _req_item_lines(text) if items else _cell_chunks(text)
    for i, chunk in enumerate(chunks):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        _run(para, chunk, size=10.5, bold=bold, font=font or _SONG)
        _set_word_wrap(para, enabled=wrap)
    if valign == "top":
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    elif valign == "bottom":
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    else:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _copied_col_ratios(kind: str, cols: int, table) -> tuple[int, ...]:
    if kind == "focus" and cols >= 3:
        return (14, 22, 64)[:cols]
    if kind == "mom" and cols >= 3:
        return (10, 16, 74)[:cols]
    if kind == "quote" and cols == 5:
        return (10, 30, 18, 18, 24)
    if kind == "quote" or cols == 4:
        return (10, 42, 24, 24)
    if cols == 2:
        return (28, 72)
    cap = 18.0 if kind == "quote_matrix" else 36.0
    ems: list[float] = []
    for c in range(cols):
        m = 4.0
        for row in table.rows:
            if c < len(row.cells):
                m = max(m, _em_len((row.cells[c].text or "").strip()) + 2.0)
        ems.append(min(m, cap))
    return tuple(max(1, int(round(x * 10))) for x in ems)


def _tight_tbl_margins(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for old in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", "40"), ("left", "80"), ("bottom", "40"), ("right", "80")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _cell_set_nowrap(cell, on: bool) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:noWrap")):
        tc_pr.remove(old)
    if on:
        el = OxmlElement("w:noWrap")
        tc_pr.append(el)


def _lock_copied_cell_wrap(table, kind: str) -> None:
    for r_i, row in enumerate(table.rows):
        for c_i, cell in enumerate(row.cells):
            long = _em_len(cell.text or "") > 10
            wrap = bool(long)
            if kind == "quote_matrix":
                wrap = False
            elif kind == "quote" and c_i == 1 and r_i > 0:
                wrap = True
            elif kind in {"focus", "mom"} and c_i >= 1 and r_i > 0:
                wrap = long
            _cell_set_nowrap(cell, not wrap)
            for p in cell.paragraphs:
                _set_word_wrap(p, enabled=wrap)


def _fit_copied_table(doc: Document, table, *, kind: str = "") -> None:
    from api.services.tenders.document import _apply_fixed_table_widths, _distribute_twips

    cols = len(table.columns) if table.columns else 1
    total = _usable_width_twips(doc)
    ratios = _copied_col_ratios(kind, cols, table)
    if len(ratios) < cols:
        ratios = ratios + (1,) * (cols - len(ratios))
    _apply_fixed_table_widths(table, _distribute_twips(total, ratios[:cols]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tight_tbl_margins(table)
    _lock_copied_cell_wrap(table, kind)


def _pipe_kind(rows: list[list[str]]) -> str:
    if not rows:
        return "text"
    blob = compact_title("".join(rows[0]))
    if "模块" in blob and "建设要求" in blob:
        return "focus"
    if any(k in blob for k in ("数据采集", "合计")) and any(
        k in blob for k in ("TPM", "WMS", "MES", "QMS")
    ):
        return "quote_matrix"
    if "序号" in blob and any(k in blob for k in ("项目内容", "金额", "单价", "品牌")):
        return "quote"
    plats = [(r[0] if r else "").strip() for r in rows]
    if any(p.upper() == "MOM" or p in {"MES", "WMS", "QMS", "TPM"} for p in plats):
        return "mom"
    return "text"


def _merge_platform_col(table) -> None:
    if len(table.columns) < 3 or len(table.rows) < 2:
        return
    vals = [(row.cells[0].text or "").strip() for row in table.rows]
    plats = [v for v in vals if v]
    if len(set(plats)) != 1:
        return
    plat = plats[0]
    if len(plat) > 8:
        return
    table.cell(0, 0).merge(table.cell(len(table.rows) - 1, 0))
    _cell_text(table.rows[0].cells[0], plat, center=True)


def _is_table_stop(line: str) -> bool:
    s = (line or "").strip()
    if not s:
        return True
    if _is_invite_attach_label(s):
        return True
    if _REQ_GROUP.match(s):
        return True
    n = compact_title(s)
    if n.startswith(("项目总报价", "注：", "注:")):
        return True
    if re.match(r"^（[0-9一二三四五六七八九十]+）", s) and any(
        k in n for k in ("报价", "明细", "硬件", "数采", "设备管理")
    ):
        return True
    if re.match(r"^[一二三四五六七八九十]、", s) and any(
        k in n for k in ("重点需求", "明细", "报价", "清单")
    ):
        return True
    return False


def _split_unit_caption(line: str) -> tuple[str, str] | None:
    raw = (line or "").strip()
    m = re.search(r"(单位[：:]\s*\S+)\s*$", raw)
    if not m:
        return None
    head = raw[: m.start()].strip()
    n = compact_title(head)
    if not any(k in n for k in ("明细", "报价", "清单", "费用", "表")):
        return None
    return head, m.group(1).strip()


def _write_unit_caption(doc: Document, title: str, unit: str, fmt: DocumentFormat) -> None:
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    _set_word_wrap(para, enabled=False)
    p_pr = para._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:tabs")):
        p_pr.remove(old)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(_usable_width_twips(doc)))
    tabs.append(tab)
    p_pr.append(tabs)
    font = fmt.fontName or _SONG
    _run(para, title, size=10.5, font=font)
    para.add_run("\t")
    _run(para, unit, size=10.5, font=font)


def _try_write_pipe_table(
    doc: Document, lines: list[str], i: int, fmt: DocumentFormat | None = None
) -> int | None:
    first = _pipe_cells(lines[i] if i < len(lines) else "")
    if not first:
        return None
    first = _unglue_header_cells(first)
    first, caption = _peel_quote_caption(first)
    first = _trim_trailing_empty(first)
    first, site_row = _peel_note_site(first)
    if len(first) >= 6 and compact_title(first[0]) == "模块" and compact_title(first[1]) == "建设要求":
        rows = [first[:3], first[3:6]]
        ncols = 3
    else:
        rows = [first]
        if site_row:
            rows.append(site_row)
        ncols = len(first)
    j = i + 1
    while j < len(lines):
        if _is_table_stop(lines[j]):
            break
        nxt = _pipe_cells(lines[j])
        if nxt:
            nxt, nxt_cap = _peel_quote_caption(nxt)
            if nxt_cap and len(rows) >= 2:
                break
            if _pipe_header_row(nxt) and len(rows) >= 2:
                break
            if len(nxt) > ncols:
                extra = nxt[ncols:]
                if any((c or "").strip() for c in extra):
                    break
                nxt = nxt[:ncols]
            if len(nxt) < ncols:
                nxt = nxt + [""] * (ncols - len(nxt))
            rows.append(nxt[:ncols])
            j += 1
            continue
        if rows and (rows[-1][-1] or "").strip() and not _is_table_stop(lines[j]):
            rows[-1][-1] = ((rows[-1][-1] or "") + lines[j]).strip()
            j += 1
            continue
        break
    if len(rows) < 2 and len(rows[0]) < 3:
        return None
    if caption:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(8)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.0
        _set_word_wrap(para, enabled=False)
        _run(para, caption, size=10.5, bold=False, font=(fmt.fontName if fmt else _SONG))
    kind = _pipe_kind(rows)
    ncols = max(3 if kind in {"focus", "mom"} else 1, len(rows[0]))
    padded = [(r + [""] * (ncols - len(r)))[:ncols] for r in rows]
    font = (fmt.fontName if fmt else _SONG) or _SONG
    if _pipe_header_row(padded[0]):
        table = _write_simple_table(doc, padded[0], padded[1:], font=font)
    else:
        table = doc.add_table(rows=len(padded), cols=ncols)
        table.style = "Table Grid"
        for r_i, row in enumerate(padded):
            for c_i, val in enumerate(row):
                _cell_text(
                    table.rows[r_i].cells[c_i],
                    val,
                    center=(kind in {"mom", "focus"} and c_i <= 1),
                    items=(kind in {"mom", "focus"} and c_i == 1),
                    font=font,
                    wrap=False,
                )
        _merge_platform_col(table)
    _fit_copied_table(doc, table, kind=kind)
    _allow_row_split(table)
    return j


def _allow_row_split(table) -> None:
    for row in table.rows:
        tr = row._tr
        tr_pr = tr.get_or_add_trPr()
        for old in tr_pr.findall(qn("w:cantSplit")):
            tr_pr.remove(old)


def _write_body(doc: Document, text: str, fmt: DocumentFormat) -> None:
    lines = _paragraphs_from_body(text)
    i = 0
    while i < len(lines):
        if _is_invite_attach_label(lines[i]):
            i += 1
            continue
        nxt = _try_write_pipe_table(doc, lines, i, fmt=fmt)
        if nxt is not None:
            i = nxt
            continue
        nxt = _try_write_req_table(doc, lines, i)
        if nxt is not None:
            i = nxt
            continue
        cells = _header_cells(lines[i])
        if cells:
            rows, i = _collect_table_rows(lines, i + 1, len(cells))
            _write_simple_table(doc, cells, rows, font=fmt.fontName)
            continue
        cap = _split_unit_caption(lines[i])
        if cap:
            _write_unit_caption(doc, cap[0], cap[1], fmt)
            i += 1
            continue
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        _run(para, lines[i], size=fmt.bodySizePt, font=fmt.fontName)
        i += 1


def _is_invite_attach_label(line: str) -> bool:
    raw = (line or "").strip()
    m = _ATTACH_LINE.match(raw)
    if not m:
        return False
    label = compact_title(m.group(1) or "")
    rest = compact_title(m.group(2) or "")
    if re.match(r"^附[一二三四五六七八九十0-9]", label) and "附件" not in label:
        return True
    return rest.endswith("模板") or rest.endswith("稿")


def _looks_req_desc(text: str) -> bool:
    s = (text or "").strip()
    if not s:
        return False
    return bool(re.match(r"^[1１一][.．、]", s) or re.search(r"[1１][.．、].+[2２][.．、]", s))


def _req_rows_from_line(line: str) -> list[tuple[str, str]]:
    s = (line or "").strip()
    if not s:
        return []
    bits = [p.strip() for p in _REQ_SPLIT.split(s) if p.strip()]
    out: list[tuple[str, str]] = []
    for bit in bits or [s]:
        m = _REQ_ROW.match(bit)
        if not m:
            continue
        desc = m.group(2).strip()
        if not _looks_req_desc(desc):
            continue
        out.append((m.group(1), desc))
    return out


def _platform_cell(group: str) -> str:
    g = re.sub(r"^[一二三四五六七八九十]、", "", (group or "").strip()).rstrip("：:")
    m = re.match(r"^([A-Za-z]{2,12})", g)
    return m.group(1) if m else (g[:12] if g else "")


def _try_write_req_table(doc: Document, lines: list[str], i: int) -> int | None:
    if i >= len(lines):
        return None
    group = ""
    cur = i
    if _REQ_GROUP.match((lines[cur] or "").strip()):
        group = lines[cur].strip()
        cur += 1
    rows: list[tuple[str, str]] = []
    while cur < len(lines):
        got = _req_rows_from_line(lines[cur])
        if not got:
            break
        rows.extend(got)
        cur += 1
    if len(rows) < 2:
        return None
    plat = _platform_cell(group)
    _write_req_table(doc, plat, rows)
    return cur


def _write_req_table(doc: Document, platform: str, rows: list[tuple[str, str]]) -> None:
    cols = 3 if platform else 2
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_i, (mod, desc) in enumerate(rows):
        if cols == 3:
            _cell_text(table.rows[r_i].cells[0], platform if r_i == 0 else "", center=True)
            _cell_text(table.rows[r_i].cells[1], mod, center=True)
            _cell_text(table.rows[r_i].cells[2], desc, items=True)
        else:
            _cell_text(table.rows[r_i].cells[0], mod, center=True)
            _cell_text(table.rows[r_i].cells[1], desc, items=True)
    if cols == 3 and len(rows) > 1:
        table.cell(0, 0).merge(table.cell(len(rows) - 1, 0))
        _cell_text(table.rows[0].cells[0], platform, center=True)
    _fit_copied_table(doc, table, kind="mom" if platform else "text")


def _header_cells(line: str) -> list[str] | None:
    from api.services.tenders.tables import split_quote_header_line

    return split_quote_header_line(line)


def _is_data_row(line: str) -> bool:
    s = (line or "").strip()
    if not s or _header_cells(s):
        return False
    n = compact_title(s)
    if _ATTACH_LINE.match(s) or _SECTION_HEAD.match(s):
        return False
    if re.match(r"^\d+", s) or n.startswith("合计"):
        return True
    return len(re.split(r"[\t]| {2,}", s)) >= 2


def _split_row(line: str, ncols: int) -> list[str]:
    raw = (line or "").strip()
    ncols = max(ncols, 1)
    if raw.count("|") >= 1:
        parts = [p.strip() for p in raw.strip("|").split("|")]
        while len(parts) < ncols:
            parts.append("")
        return parts[:ncols]
    m = re.match(r"^(\d+)\s*[,.．、]?\s*(.*)$", raw)
    if m:
        rest = (m.group(2) or "").strip()
        extra = [p.strip() for p in re.split(r"[\t]| {2,}", rest) if p.strip()] if rest else []
        parts = [m.group(1), *(extra or ([rest] if rest else []))]
        while len(parts) < ncols:
            parts.append("")
        return parts[:ncols]
    parts = [p.strip() for p in re.split(r"[\t]| {2,}", raw) if p.strip()]
    if not parts:
        parts = [raw]
    while len(parts) < ncols:
        parts.append("")
    return parts[:ncols]


def _collect_table_rows(lines: list[str], start: int, ncols: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines):
        s = (lines[i] or "").strip()
        if not _is_data_row(s):
            break
        rows.append(_split_row(s, ncols))
        i += 1
    return rows, i


def _write_simple_table(
    doc: Document, headers: list[str], rows: list[list[str]], font: str = _SONG
):
    cols = max(len(headers), 1)
    data = rows or [[""] * cols]
    blob = compact_title("".join(headers))
    kind = "text"
    if "建设要求" in blob or ("模块" in blob and "实现目标" in blob):
        kind = "focus"
    elif any(k in blob for k in ("数据采集", "合计")) and any(
        k in blob for k in ("TPM", "WMS", "MES", "QMS")
    ):
        kind = "quote_matrix"
    elif "序号" in blob and any(k in blob for k in ("项目内容", "金额", "单价", "品牌")):
        kind = "quote"
    face = font or _SONG
    matrix = kind == "quote_matrix"
    table = doc.add_table(rows=1 + len(data), cols=cols)
    table.style = "Table Grid"
    for i, h in enumerate(headers[:cols]):
        _cell_text(
            table.rows[0].cells[i],
            h,
            bold=True,
            center=True,
            font=face,
            wrap=not matrix,
        )
    for r_i, row in enumerate(data):
        for c_i in range(cols):
            val = row[c_i] if c_i < len(row) else ""
            center = False
            if kind == "focus":
                center = c_i <= 1
            elif kind == "quote_matrix":
                center = True
            elif kind == "quote":
                center = c_i == 0
            _cell_text(
                table.rows[r_i + 1].cells[c_i],
                val,
                center=center,
                items=(kind == "focus" and c_i == 1),
                font=face,
                wrap=not matrix,
            )
    _fit_copied_table(doc, table, kind=kind)
    _allow_row_split(table)
    return table


def _paragraphs_from_body(text: str) -> list[str]:
    raw = unmash_invitation_text(text or "")
    lines = [ln.strip() for ln in raw.split("\n")]
    lines = [ln for ln in lines if ln]
    # 原文已有换行就保持段落；只有 OCR 粘成一整段时才按句号拆开
    if len(lines) <= 1 and len(raw) > 240 and "|" not in raw:
        lines = [p.strip() for p in re.split(r"(?<=[。；;])", raw) if p.strip()]
    spread: list[str] = []
    for ln in lines:
        if "|" in ln:
            spread.append(ln)
            continue
        spread.extend(split_mashed_zhi_line(ln))
    merged = _coalesce_broken_form_lines(drop_ocr_junk_lines(spread))
    out: list[str] = []
    for ln in merged:
        parts = unfold_form_sign_lines(ln).splitlines()
        out.extend(p.strip() for p in parts if p.strip())
    return out or merged


_UNDER_ONLY = re.compile(r"^[-—–_＿\s]{2,}$")
_YMD_BIT = re.compile(r"^[年月日]([：:].*)?$")
_FRAG_HEAD = re.compile(r"[时日名址性]$")
_FRAG_TAIL = re.compile(r"^[间期称址质][：:]?")


def _coalesce_broken_form_lines(lines: list[str]) -> list[str]:
    """把 OCR/PDF 拆开的「成立时」「间：」「年」「月」「日」拼回一行。"""
    out: list[str] = []
    for ln in lines:
        s = (ln or "").strip()
        if not s or _UNDER_ONLY.match(s):
            continue
        s = re.sub(r"(日)(经营期限)", r"\1\n\2", s)
        chunks = [p.strip() for p in s.split("\n") if p.strip()]
        for chunk in chunks:
            if out and (_YMD_BIT.match(chunk) or (_FRAG_HEAD.search(out[-1]) and _FRAG_TAIL.match(chunk))):
                out[-1] += chunk
                continue
            if (
                out
                and len(chunk) <= 3
                and chunk.endswith(("：", ":"))
                and compact_title(chunk) not in {"致：", "致:"}
                and not out[-1].endswith(("：", ":", "。", "；"))
            ):
                out[-1] += chunk
                continue
            out.append(chunk)
    return out


def _run(para, text: str, *, size: float = 12, bold: bool = False, font: str = _SONG) -> None:
    name = font or _SONG
    run = para.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    proof = OxmlElement("w:noProof")
    rpr.append(proof)
