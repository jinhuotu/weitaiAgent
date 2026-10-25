"""格式章截取：规则可调，截下来的块与上传 Word 的 XML 一致。"""

from __future__ import annotations

import zipfile

import pytest

from docx import Document
from docx.enum.section import WD_SECTION, WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from api.services.tenders.slice_rules import cut_format_docx, is_format_end, is_format_start


def test_start_must_be_the_whole_heading() -> None:
    assert is_format_start("第八章、投标文件格式")
    assert is_format_start("第3部分 响应文件组成")
    assert not is_format_start("请阅读投标文件格式中的要求。")
    assert not is_format_start("投标文件格式见本章说明，投标人须按此编制。")
    assert is_format_end("第九章 评标办法")
    assert not is_format_end("第九章、投标文件格式")
    assert not is_format_end("本章规定了投标文件格式的装订方式。")
    assert not is_format_end("第一部分 投标函")


def test_cut_keeps_original_xml(tmp_path) -> None:
    doc = Document()
    doc.add_paragraph("第一章 招标公告")
    doc.add_paragraph("请阅读投标文件格式中的要求。")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("第八章、投标文件格式")
    letter = doc.add_paragraph()
    letter.add_run("致：")
    blank = letter.add_run("（招标人名称）")
    blank.underline = True
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "投标人"
    table.cell(0, 1).text = "盖章"
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("第九章 评标办法")
    doc.add_paragraph("评标办法正文不应进入格式章。")
    src = tmp_path / "invite.docx"
    doc.save(src)
    dest = tmp_path / "fmt.docx"
    cut_format_docx(src, dest)

    with zipfile.ZipFile(src) as zf:
        origin = zf.read("word/document.xml").decode("utf-8")
    with zipfile.ZipFile(dest) as zf:
        sliced = zf.read("word/document.xml").decode("utf-8")

    anchor = origin.find("（招标人名称）")
    assert anchor > 0
    window = origin[anchor - 60 : anchor + 20]
    assert window in sliced
    assert "第九章 评标办法" not in sliced
    assert "招标公告" not in sliced
    assert "第八章、投标文件格式" in sliced
    assert "评标办法正文不应进入格式章" not in sliced

    opened = Document(str(dest))
    cells = [
        cell.text
        for table in opened.tables
        for row in table.rows
        for cell in row.cells
    ]
    assert "盖章" in cells
    assert opened.element.body.find(qn("w:tbl")) is not None
    # 截下来的正文块仍能按原 XML 解析
    etree.fromstring(sliced.encode("utf-8"))


@pytest.mark.parametrize('boundary', ['section', 'page_break', 'title_break', 'before_cover'])
def test_only_separate_chapter_title_page_is_trimmed(tmp_path, boundary):
    doc = Document()
    doc.add_paragraph('第八章、投标文件格式')
    if boundary == 'section':
        doc.add_section(WD_SECTION.NEW_PAGE)
    elif boundary == 'page_break':
        doc.add_page_break()
    elif boundary == 'title_break':
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    project = doc.add_paragraph('真实项目（项目名称）')
    if boundary == 'before_cover':
        project.paragraph_format.page_break_before = True
    doc.add_paragraph('投 标 文 件')
    doc.add_paragraph('招标编号：WT-01')
    doc.add_paragraph('投标人：某公司（盖章）')
    doc.add_paragraph('第一部分 投标函')
    doc.add_paragraph('投标函及投标函附录')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '已标价工程量清单'
    original_cover = project._p.xml
    original_table = table._tbl.xml
    src, dest = tmp_path/'source.docx', tmp_path/'cut.docx'
    doc.save(src)
    cut_format_docx(src, dest)
    after = Document(dest)
    assert after.paragraphs[0].text == '真实项目（项目名称）'
    assert after.paragraphs[0]._p.xml == original_cover
    assert after.tables[0]._tbl.xml == original_table
    assert [p.text for p in after.paragraphs][-2:] == ['第一部分 投标函', '投标函及投标函附录']
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(dest) as b:
        assert [n for n in a.namelist() if a.read(n) != b.read(n)] == ['word/document.xml']


@pytest.mark.parametrize('boundary', ['none', 'continuous', 'rendered_only', 'leading_break'])
def test_same_page_or_uncertain_title_is_preserved(tmp_path, boundary):
    from api.services.tenders.pdf_format import _drop_invite_chrome
    doc = Document()
    title = doc.add_paragraph('第八章、投标文件格式')
    if boundary == 'continuous':
        doc.add_section(WD_SECTION_START.CONTINUOUS)
    elif boundary == 'rendered_only':
        title.add_run()._r.append(OxmlElement('w:lastRenderedPageBreak'))
    elif boundary == 'leading_break':
        run = OxmlElement('w:r')
        br = OxmlElement('w:br')
        br.set(qn('w:type'), 'page')
        run.append(br)
        title._p.insert(0, run)
    doc.add_paragraph('真实项目 投标文件 招标编号：WT-01 投标人：某公司')
    src, dest = tmp_path/'source.docx', tmp_path/'cut.docx'
    doc.save(src)
    cut_format_docx(src, dest)
    result = Document(dest)
    assert result.element.body.xml == doc.element.body.xml
    before = result.element.xml
    _drop_invite_chrome(result)
    assert result.element.xml == before


def test_old_cached_format_title_page_is_trimmed_by_generation():
    from api.services.tenders.pdf_format import _drop_invite_chrome
    doc = Document()
    doc.add_paragraph('第八章、投标文件格式')
    doc.add_section(WD_SECTION.NEW_PAGE)
    cover = doc.add_paragraph('真实项目 投标文件 招标编号：WT-01 投标人：某公司')
    original = cover._p.xml
    doc.add_paragraph('第一部分 投标函')
    _drop_invite_chrome(doc)
    assert doc.paragraphs[0]._p.xml == original
    assert doc.paragraphs[1].text == '第一部分 投标函'
    before = doc.element.xml
    _drop_invite_chrome(doc)
    assert doc.element.xml == before
