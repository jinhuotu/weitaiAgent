from io import BytesIO

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders.placeholders import place_unmatched_slots, _is_slot_box
from api.services.tenders.match import attachment_requests
from api.services.tenders.schema import PlaceholderItem


def next_content(el, steps=1):
    """附件间的留白不算正文节点；仍核对内容顺序和重复插入。"""
    for _ in range(steps):
        el = el.getnext()
        while el is not None:
            style = el.find('w:pPr/w:pStyle', el.nsmap)
            if style is None or style.get(qn('w:val')) != 'TenderAttachmentGap':
                break
            assert not ''.join(el.itertext()).strip()
            el = el.getnext()
    return el


def text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def bookmark(para, name):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), '70')
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), '70')
    para._p.append(start)
    para._p.append(end)


@pytest.mark.parametrize('volume', ['business', 'technical', ''])
def test_wrapped_bond_and_bill_notes_stay_in_body(volume):
    doc = Document()
    doc.add_paragraph('目录')
    toc = doc.add_paragraph('四、投标保证金')
    doc.add_paragraph('五、已标价的工程量清单')
    doc.add_paragraph('一、投标函')
    doc.add_paragraph('我方已仔细研究了招标文件。')
    doc.add_paragraph('授权委托书')
    head = doc.add_paragraph('投标保证金')
    bookmark(head, '四、投标保证金')
    doc.add_paragraph('若采用现金或支票，投标人应在此提供汇款凭证')
    note = doc.add_paragraph('的复印件或扫描件。还应附基本账户开户证明的复印件或扫描件。若采用银行保函，需提供银行保函。')
    bill = doc.add_paragraph('已标价的工程量清单')
    bookmark(bill, '五、已标价的工程量清单')
    end_note = doc.add_paragraph('说明：已标价工程量清单按第五章的相关规定填写。')
    doc.add_paragraph('六、项目管理机构')
    originals = [(n, n.xml) for n in doc.element.body]
    place_unmatched_slots(doc, slots=[], media={}, required_keys=[], volume=volume)
    assert next_content(toc._p).tag == qn('w:p')
    cursor = next_content(note._p)
    for title in ['投标保证金汇款凭证', '基本账户开户证明', '银行保函']:
        assert _is_slot_box(cursor) and '此处粘贴'+title in text(cursor)
        borders = cursor.find('w:pPr/w:pBdr', cursor.nsmap)
        assert len(borders) == 4 and all(b.get(qn('w:val')) == 'dashed' for b in borders)
        cursor = next_content(cursor)
    assert cursor is bill._p
    assert '此处粘贴已标价工程量清单表' in text(next_content(end_note._p))
    assert all(n.xml == xml for n, xml in originals)
    # Reloaded DOCX remains idempotent, including inserted image/box markers.
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    again = Document(buf)
    before = again.element.xml
    place_unmatched_slots(again, slots=[], media={}, required_keys=[], volume=volume)
    assert again.element.xml == before


def test_basic_table_inserts_all_requested_materials_and_complete_pdf(tmp_path):
    scan = tmp_path/'license.pdf'
    pdf = fitz.open()
    for i in range(2):
        pdf.new_page().insert_text((60, 60), f'License page {i+1}')
    pdf.save(scan)
    pdf.close()
    doc = Document()
    doc.add_paragraph('（一）投标人基本情况表')
    original_table = doc.add_table(rows=1, cols=2)
    original_table.cell(0, 0).text = '投标人名称'
    note = doc.add_paragraph('注：本表后应附营业执照、投标人资质证书副本和安全生产许可等材料的复印件或扫描件。')
    doc.add_paragraph('（二）近年财务状况表')
    slots = [PlaceholderItem(key='empty', title='营业执照'), PlaceholderItem(key='uploaded', title='营业执照'),
             PlaceholderItem(key='qualification', title='施工资质'), PlaceholderItem(key='iso', title='ISO体系认证证书')]
    table_xml = original_table._tbl.xml
    notes = place_unmatched_slots(doc, slots=slots, media={'uploaded':[scan]}, required_keys=['empty','uploaded','qualification'])
    assert original_table._tbl.xml == table_xml
    cursor = next_content(note._p)
    assert len(list(cursor.iter(qn('w:drawing')))) == 1
    cursor = next_content(cursor)
    while cursor.tag == qn('w:p') and not list(cursor.iter(qn('w:drawing'))):
        cursor = next_content(cursor)
    assert len(list(cursor.iter(qn('w:drawing')))) == 1
    cursor = next_content(cursor)
    while cursor.tag == qn('w:p') and not text(cursor):
        cursor = next_content(cursor)
    assert '此处粘贴资质证书副本' in text(cursor)
    assert '此处粘贴安全生产许可证' in text(next_content(cursor))
    assert '附件：资料库扫描件' not in text(doc.element.body)
    assert any('营业执照' in n and '贴上资料' in n for n in notes)
    assert sum('此处粘贴资质证书副本' in text(n) for n in doc.element.body) == 1


def test_unknown_named_scan_request_is_local_not_tail():
    doc = Document()
    doc.add_paragraph('二、其他证明')
    note = doc.add_paragraph('须附特种设备许可证扫描件。')
    doc.add_paragraph('三、下一章')
    place_unmatched_slots(doc, slots=[], media={}, required_keys=[])
    assert '此处粘贴特种设备许可证' in text(next_content(note._p))
    assert '附件：资料库扫描件' not in text(doc.element.body)


def test_qualification_does_not_use_iso_and_conditions_are_preserved():
    from api.services.tenders.match import attachment_catalog_item
    req = attachment_requests('本表后应附资质证书副本的扫描件。')[0]
    assert attachment_catalog_item(req, [PlaceholderItem(key='iso',title='ISO质量体系认证证书')]) is None
    names = [r.title for r in attachment_requests('若采用银行保函，需提供银行保函。')]
    assert names == ['银行保函']


def test_generic_and_repeated_attachment_cues_keep_each_material_name():
    names = [r.title for r in attachment_requests('应附营业执照扫描件、企业简介扫描件。专职安全员应附安全生产考核合格证书扫描件，主要业绩须附合同协议书（如有）。')]
    assert names == ['营业执照', '企业简介', '安全生产考核合格证书', '合同协议书']


def test_financial_report_components_are_one_attachment_request():
    names = [r.title for r in attachment_requests('应附经审计的财务会计报表，包括资产负债表、现金流量表、利润表的扫描件，具体年份要求见投标人须知前附表。')]
    assert names == ['财务会计报表']


def test_existing_named_box_does_not_suppress_other_materials():
    from api.services.tenders.placeholders import draw_placeholder_box
    doc = Document()
    doc.add_paragraph('一、投标人基本情况表')
    note = doc.add_paragraph('本表后应附营业执照、资质证书的扫描件。')
    draw_placeholder_box(doc, PlaceholderItem(key='license',title='营业执照'),after=note)
    doc.add_paragraph('二、下一章')
    place_unmatched_slots(doc,slots=[],media={},required_keys=[])
    assert sum('此处粘贴营业执照' in text(n) for n in doc.element.body) == 1
    assert sum('此处粘贴资质证书副本' in text(n) for n in doc.element.body) == 1


def test_bare_bill_heading_without_bookmark_gets_local_box():
    doc = Document()
    doc.add_paragraph('已标价的工程量清单')
    note = doc.add_paragraph('说明：按第五章规定填写。')
    doc.add_paragraph('六、项目管理机构')
    place_unmatched_slots(doc,slots=[],media={},required_keys=[])
    assert '此处粘贴已标价工程量清单表' in text(next_content(note._p))


def test_bill_box_uses_space_before_original_section_break():
    from docx.enum.section import WD_SECTION
    doc = Document()
    doc.add_paragraph('五、已标价工程量清单')
    note = doc.add_paragraph('说明：按第五章规定填写。')
    doc.add_section(WD_SECTION.NEW_PAGE)
    boundary = next_content(note._p)
    doc.add_paragraph('六、项目管理机构')
    originals = [(n,n.xml) for n in doc.element.body]
    place_unmatched_slots(doc,slots=[],media={},required_keys=[])
    assert '此处粘贴已标价工程量清单表' in text(next_content(note._p))
    assert next_content(note._p, 2) is boundary
    assert all(n.xml == xml for n,xml in originals)


@pytest.mark.parametrize('floating', [False, True])
def test_resume_note_stays_adjacent_to_original_table(floating):
    from docx.enum.section import WD_SECTION
    doc = Document()
    doc.add_paragraph('（二）项目经理简历表')
    note = doc.add_paragraph('项目经理须附满足第二章第1.4.1项规定的相关证明材料，及未担任其他在施建设工程项目项目经理的承诺书扫描件。')
    table = doc.add_table(rows=1,cols=2)
    table.cell(0,0).text = '姓名'
    table.cell(0,1).text = '项目经理'
    if floating:
        carrier = doc.add_paragraph()
        carrier._p.append(table._tbl)
        original_form = carrier._p
    else:
        original_form = table._tbl
    doc.add_section(WD_SECTION.NEW_PAGE)
    boundary = next_content(original_form)
    doc.add_paragraph('（三）主要工作经历')
    original_xml = original_form.xml
    place_unmatched_slots(doc,slots=[],media={},required_keys=[])
    assert note._p.getnext() is original_form
    assert original_form.xml == original_xml
    first = next_content(original_form)
    assert '此处粘贴相关证明材料' in text(first)
    assert '此处粘贴未担任其他在施建设工程项目项目经理的承诺书' in text(next_content(first))
    assert next_content(first, 2) is boundary
    before = doc.element.xml
    place_unmatched_slots(doc,slots=[],media={},required_keys=[])
    assert doc.element.xml == before


@pytest.mark.parametrize('used, expected_breaks', [(3*24+3,0),(3*24+21,1),(3*24,0)])
def test_slot_page_break_uses_current_page_height(monkeypatch,used,expected_breaks):
    from api.services.tenders import placeholders as ph
    doc = Document()
    monkeypatch.setattr(ph,'_page_body_cm',lambda d:24.0)
    monkeypatch.setattr(ph,'_used_on_page_cm',lambda d:float(used))
    ph._break_before_slot(doc,6.0)
    assert len(list(doc.element.body.iter(qn('w:br')))) == expected_breaks
