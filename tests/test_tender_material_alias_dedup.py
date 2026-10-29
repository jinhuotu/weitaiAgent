from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from api.services.tenders import library_kb
from api.services.tenders.scan_placement import place_library_scans
from api.services.tenders.schema import BidBrief, PlaceholderItem

CUSTOM = "m5641425fda52"


@pytest.fixture
def scans(tmp_path):
    original = tmp_path / "license.png"
    duplicate = tmp_path / "renamed-license.png"
    other = tmp_path / "license-second-page.png"
    Image.new("RGB", (120, 80), "red").save(original)
    duplicate.write_bytes(original.read_bytes())
    Image.new("RGB", (120, 80), "green").save(other)
    return original, duplicate, other


@pytest.mark.asyncio
async def test_resolver_unifies_alias_files_and_preserves_key_lookups(monkeypatch, scans):
    original, duplicate, other = scans
    parents = [
        SimpleNamespace(public_id=key, parent_id=None, name="营业执照", tags=["slot:" + key])
        for key in (CUSTOM, "license")
    ]
    monkeypatch.setattr(
        library_kb, "get_base_by_public_id", AsyncMock(return_value=SimpleNamespace(id=1))
    )
    monkeypatch.setattr(library_kb, "_list_docs", AsyncMock(return_value=parents))
    monkeypatch.setattr(library_kb, "_files_of", lambda p, _: {
        CUSTOM: [original, other], "license": [duplicate],
    }[p.public_id])
    monkeypatch.setattr(library_kb, "list_slot_files", lambda _: [])
    slots = [PlaceholderItem(key=key, title="营业执照") for key in (CUSTOM, "license")]
    media = await library_kb.resolve_attachments(None, slots)
    assert media[CUSTOM] is media["license"]
    assert [p.read_bytes() for p in media["license"]] == [
        duplicate.read_bytes(), other.read_bytes(),
    ]


@pytest.mark.parametrize("framed", [False, True])
def test_aliases_paste_one_copy_at_same_placeholder_and_stay_idempotent(scans, framed):
    original, duplicate, _ = scans
    doc = Document()
    doc.add_paragraph("资格审查资料", style="Heading 1")
    doc.add_paragraph("附：营业执照复印件")
    if framed:
        doc.add_paragraph()._p.get_or_add_pPr().append(OxmlElement("w:pBdr"))
    slots = [PlaceholderItem(key=key, title="营业执照") for key in (CUSTOM, "license")]
    media = {CUSTOM: [original], "license": [duplicate]}
    _, handled = place_library_scans(doc, BidBrief(), media, slots)
    assert len(list(doc.element.iter(qn("w:drawing")))) == 1
    assert {CUSTOM, "license"} <= handled
    once = doc.element.xml
    place_library_scans(doc, BidBrief(), media, slots)
    assert doc.element.xml == once


def test_identical_content_across_different_keys_at_one_placeholder_is_pasted_once(scans):
    original, duplicate, other = scans
    doc = Document()
    doc.add_paragraph("附：营业执照、资质证书扫描件")
    slots = [PlaceholderItem(key="qualification", title="企业资质证书")]
    place_library_scans(doc, BidBrief(), {
        "license": [original, duplicate], "qualification": [duplicate, other],
    }, slots)
    assert len(list(doc.element.iter(qn("w:drawing")))) == 2


def test_different_placeholders_can_reuse_file_and_distinct_pages_are_kept(scans):
    original, duplicate, other = scans
    doc = Document()
    doc.add_paragraph("附：营业执照扫描件")
    doc.add_paragraph("一、其他资料", style="Heading 1")
    doc.add_paragraph("此处粘贴营业执照复印件")
    place_library_scans(doc, BidBrief(), {"license": [original, duplicate, other]})
    assert len(list(doc.element.iter(qn("w:drawing")))) == 4


def test_old_custom_key_marker_already_satisfies_its_standard_alias(scans):
    from api.services.tenders.document import _bookmark_paragraph
    from api.services.tenders.scan_placement import _prefix

    original, duplicate, _ = scans
    doc = Document()
    doc.add_paragraph("附：营业执照扫描件")
    pasted = doc.add_paragraph()
    pasted.add_run().add_picture(str(original))
    _bookmark_paragraph(pasted, _prefix(CUSTOM) + "0")
    once = doc.element.xml
    _, handled = place_library_scans(
        doc, BidBrief(), {CUSTOM: [original], "license": [duplicate]},
        [PlaceholderItem(key=CUSTOM, title="营业执照")],
    )
    assert doc.element.xml == once
    assert {CUSTOM, "license"} <= handled


def test_aliased_multipage_pdf_keeps_each_page_once(tmp_path):
    import fitz

    original = tmp_path / "license.pdf"
    with fitz.open() as pdf:
        for text in ("License page one", "License page two"):
            page = pdf.new_page(width=120, height=80)
            page.insert_text((10, 20), text, fontsize=7)
        pdf.save(original)
    duplicate = tmp_path / "renamed.pdf"
    duplicate.write_bytes(original.read_bytes())
    doc = Document()
    doc.add_paragraph("附：营业执照扫描件")
    place_library_scans(
        doc, BidBrief(), {CUSTOM: [original], "license": [duplicate]},
        [PlaceholderItem(key=CUSTOM, title="营业执照")],
    )
    assert len(list(doc.element.iter(qn("w:drawing")))) == 2
