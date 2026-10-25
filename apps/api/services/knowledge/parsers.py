"""按扩展名抽出可切块文本：有字走本地，扫描页/嵌入图走 OCR。"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

from common.config import get_settings
from common.errors import AppError, ErrorCode

from api.services.knowledge.ocr.factory import get_ocr_provider
from api.services.word_convert import LEGACY_WORD_HINT, WordConversionError, convert_doc_to_docx

logger = logging.getLogger("api.kb.parsers")

SUPPORTED_EXTENSIONS = frozenset(
    {
        "txt",
        "md",
        "csv",
        "json",
        "xml",
        "yaml",
        "yml",
        "pdf",
        "doc",
        "docx",
        "xlsx",
        "xls",
        "pptx",
        "png",
        "jpg",
        "jpeg",
        "webp",
        "gif",
        "bmp",
        "mp4",
        "webm",
    }
)
IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "webp", "gif", "bmp"})
ALLOWED_HINT = (
    "pdf / doc / docx / xlsx / xls / pptx / txt / md / csv / json / xml / yaml / "
    "png / jpg / webp / mp4 / webm"
)

_CONTENT_TYPE_EXT = {
    "application/pdf": "pdf",
    "application/msword": "doc",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/csv": "csv",
    "application/json": "json",
    "application/xml": "xml",
    "text/xml": "xml",
    "application/x-yaml": "yaml",
    "text/yaml": "yaml",
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/webp": "webp",
    "image/gif": "gif",
    "image/bmp": "bmp",
    "video/mp4": "mp4",
    "video/webm": "webm",
}

_XLSX_MAX_SHEETS = 40
_XLSX_MAX_ROWS = 8000
_XLSX_MAX_COLS = 40
_COUNT_RE = re.compile(r"[\u4e00-\u9fffA-Za-z0-9]")
# 只收掉字间空格。「证明\n投标人」这种换行是下一栏，不能并成一行。
_CJK_SPACE_RE = re.compile(r"(?<=[\u4e00-\u9fff])[^\S\n]+(?=[\u4e00-\u9fff])")
_BLIP_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_EMBED_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_FORMULA_NOTE = "[说明：部分单元格未经 Excel 计算，已保留公式原文]"


@dataclass
class ExtractResult:
    text: str
    page_count: int = 0
    ocr_pages: int = 0
    ocr_capped: bool = False
    formula_fallback: bool = False
    from_file: bool = True


def sniff_extension(filename: str | None, content_type: str | None = None) -> str:
    suffix = Path(filename or "").suffix.lower().lstrip(".")
    if suffix:
        return suffix
    ctype = (content_type or "").split(";")[0].strip().lower()
    return _CONTENT_TYPE_EXT.get(ctype, "")


def assert_supported(ext: str) -> str:
    kind = (ext or "").lower().lstrip(".")
    if kind in SUPPORTED_EXTENSIONS:
        return kind
    if kind:
        raise AppError(
            ErrorCode.VALIDATION,
            f"不支持的文件类型 .{kind}。当前支持：{ALLOWED_HINT}",
            status_code=422,
        )
    raise AppError(ErrorCode.VALIDATION, f"无法识别文件类型。当前支持：{ALLOWED_HINT}", status_code=422)


def _count_chars(text: str) -> int:
    return len(_COUNT_RE.findall(text or ""))


def collapse_cjk_spaces(text: str) -> str:
    """「投 标 文 件」→「投标文件」。"""
    prev = None
    out = text or ""
    while prev != out:
        prev = out
        out = _CJK_SPACE_RE.sub("", out)
    return out


def _normalize(text: str) -> str:
    cleaned = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    cleaned = collapse_cjk_spaces(cleaned)
    lines = [line.rstrip() for line in cleaned.split("\n")]
    return "\n".join(line for line in lines if line.strip()).strip()


async def _ocr_image(raw: bytes) -> str:
    provider = get_ocr_provider()
    return await asyncio.to_thread(provider.recognize, raw)


async def extract_text_from_file(path: Path, *, ext: str | None = None) -> ExtractResult:
    kind = assert_supported(ext or path.suffix.lower().lstrip("."))
    if not path.is_file():
        raise AppError(ErrorCode.NOT_FOUND, "uploaded file missing on disk", status_code=500)

    if kind in {"txt", "md"}:
        result = ExtractResult(text=_read_text_file(path), page_count=1)
    elif kind == "csv":
        result = ExtractResult(text=_read_csv(path), page_count=1)
    elif kind == "json":
        result = ExtractResult(text=_read_json(path), page_count=1)
    elif kind in {"yaml", "yml"}:
        result = ExtractResult(text=_read_yaml(path), page_count=1)
    elif kind == "xml":
        result = ExtractResult(text=_read_xml(path), page_count=1)
    elif kind == "pdf":
        result = await _read_pdf(path)
    elif kind == "docx":
        result = await _read_docx(path)
    elif kind == "doc":
        with tempfile.TemporaryDirectory(prefix="legacy-word-") as temp:
            converted = Path(temp) / "document.docx"
            try:
                await asyncio.to_thread(convert_doc_to_docx, path, converted)
            except (WordConversionError, OSError) as exc:
                logger.warning("legacy Word conversion failed file=%r: %s", path.name, exc)
                raise AppError(ErrorCode.VALIDATION, LEGACY_WORD_HINT, status_code=422) from exc
            result = await _read_docx(converted)
    elif kind == "xlsx":
        result = await _read_xlsx(path)
    elif kind == "xls":
        result = await _read_xls(path)
    elif kind == "pptx":
        result = await _read_pptx(path)
    elif kind in IMAGE_EXTENSIONS:
        result = await _read_image(path)
    elif kind in {"mp4", "webm"}:
        raise AppError(
            ErrorCode.VALIDATION,
            "视频请走语音转写入库，不能直接当文本解析",
            status_code=422,
        )
    else:
        raise AppError(ErrorCode.VALIDATION, f"不支持的文件类型 .{kind}", status_code=422)

    settings = get_settings()
    cleaned = _normalize(result.text)
    max_chars = int(settings.kb_extract_max_chars)
    if len(cleaned) > max_chars:
        cleaned = cleaned[:max_chars] + "\n[已截断：超过解析正文上限]"
    if len(cleaned) < 4:
        raise AppError(
            ErrorCode.VALIDATION,
            "文件解析后正文过短（可能是扫描件且 OCR 未启用或未识别出字）",
            status_code=422,
        )
    result.text = cleaned
    result.from_file = True
    return result


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_csv(path: Path) -> str:
    raw = _read_text_file(path)
    reader = csv.reader(io.StringIO(raw))
    rows: list[str] = []
    for i, row in enumerate(reader):
        if i >= _XLSX_MAX_ROWS:
            rows.append("[已截断：超过行数上限]")
            break
        cells = [str(c).strip() for c in row[:_XLSX_MAX_COLS]]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _read_json(path: Path) -> str:
    import json

    raw = _read_text_file(path)
    data = json.loads(raw)
    return json.dumps(data, ensure_ascii=False, indent=2)


def _read_yaml(path: Path) -> str:
    import yaml

    raw = _read_text_file(path)
    data = yaml.safe_load(raw)
    dumped = yaml.safe_dump(data, allow_unicode=True, sort_keys=False)
    return dumped if isinstance(dumped, str) else raw


def _read_xml(path: Path) -> str:
    import xml.etree.ElementTree as ET

    tree = ET.parse(path)
    parts: list[str] = []

    def walk(node: ET.Element, depth: int = 0) -> None:
        tag = (node.tag or "").split("}")[-1]
        text = " ".join((node.text or "").split())
        line = f"{'  ' * depth}{tag}"
        if text:
            line = f"{line}: {text}"
        parts.append(line)
        for child in list(node):
            walk(child, depth + 1)

    walk(tree.getroot())
    return "\n".join(parts)


async def _read_image(path: Path) -> ExtractResult:
    text = await _ocr_image(path.read_bytes())
    return ExtractResult(text=_normalize(text), page_count=1, ocr_pages=1)


def _render_pdf_page_jpeg(pdf_path: Path, page_index: int) -> bytes:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        page = doc[page_index]
        bitmap = page.render(scale=2.0)
        pil = bitmap.to_pil()
        buf = io.BytesIO()
        pil.convert("RGB").save(buf, format="JPEG", quality=85)
        return buf.getvalue()
    finally:
        doc.close()


_RULED_TABLE = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "snap_tolerance": 4,
    "join_tolerance": 4,
    "intersection_tolerance": 5,
    # 字高左右的短线是下划线，不是表框
    "edge_min_length": 18,
}
_HSPAN = ">>"
_VSPAN = "^^"


def _cell_pipe(value: object) -> str:
    s = re.sub(r"\s+", " ", str(value or "").replace("|", "／")).strip()
    return s


def _pipe_rows(rows: list[list[object]]) -> str:
    lines: list[str] = []
    width = 0
    cleaned: list[list[str]] = []
    for row in rows:
        cells = [_cell_pipe(c) for c in (row or [])]
        if not any(cells):
            continue
        width = max(width, len(cells))
        cleaned.append(cells)
    if len(cleaned) < 2 or width < 2:
        return ""
    for cells in cleaned:
        if len(cells) < width:
            cells = cells + [""] * (width - len(cells))
        lines.append(" | ".join(cells[:width]))
    return "\n".join(lines)


def _cluster_edges(vals: list[float], tol: float = 2.0) -> list[float]:
    xs = sorted(vals)
    groups: list[list[float]] = []
    for x in xs:
        if not groups or x - groups[-1][-1] > tol:
            groups.append([x])
        else:
            groups[-1].append(x)
    return [sum(g) / len(g) for g in groups]


def _edge_index(edges: list[float], x: float) -> int:
    return min(range(len(edges)), key=lambda i: abs(edges[i] - x))


def _grid_from_cells(rows: list[list[tuple[tuple | None, str | None]]]) -> list[list[str]]:
    """按格子坐标还原 colspan / rowspan。>> 向左并，^^ 向上并。"""
    bands: list[float | None] = []
    edges_raw: list[float] = []
    for row in rows:
        tops = [box[1] for box, _text in row if box is not None]
        bands.append(min(tops) if tops else None)
        for box, _text in row:
            if box is None:
                continue
            edges_raw.extend((box[0], box[2]))
    edges = _cluster_edges(edges_raw)
    if len(edges) < 3 or len(rows) < 2:
        return []
    ncols = len(edges) - 1
    grid: list[list[str | None]] = [[None] * ncols for _ in rows]
    spans: list[tuple[int, int, int, int, str]] = []
    for ri, row in enumerate(rows):
        for box, text in row:
            if box is None:
                continue
            c0 = _edge_index(edges, box[0])
            c1 = _edge_index(edges, box[2])
            if c1 <= c0:
                c1 = min(ncols, c0 + 1)
            r1 = ri
            for rj in range(ri + 1, len(rows)):
                nxt = bands[rj]
                if nxt is None:
                    continue
                if box[3] > nxt + 4:
                    r1 = rj
                else:
                    break
            spans.append((ri, c0, r1, c1, _cell_pipe(text)))
    for ri, c0, r1, c1, text in spans:
        if grid[ri][c0] in (None, "", _HSPAN, _VSPAN):
            grid[ri][c0] = text
        for c in range(c0 + 1, c1):
            if grid[ri][c] is None:
                grid[ri][c] = _HSPAN
        for r in range(ri + 1, r1 + 1):
            if grid[r][c0] is None:
                grid[r][c0] = _VSPAN
            for c in range(c0 + 1, c1):
                if grid[r][c] is None:
                    grid[r][c] = _HSPAN
    out: list[list[str]] = []
    for row in grid:
        out.append([("" if c is None else c) for c in row])
    return out


def _grid_usable(grid: list[list[str]]) -> bool:
    if len(grid) < 2 or max((len(r) for r in grid), default=0) < 2:
        return False
    texts: list[str] = []
    for row in grid:
        for c in row:
            if c in {_HSPAN, _VSPAN}:
                continue
            s = re.sub(r"\s+", "", c or "")
            if s:
                texts.append(s)
    long = [s for s in texts if len(s) >= 2]
    if len(long) < 2:
        return False
    short = sum(1 for s in texts if len(s) <= 1)
    if texts and short / len(texts) > 0.45:
        return False
    # 1.1.2 这种条款号格子多半是正文被线切碎，不是要还原的表
    numish = sum(
        1 for s in texts if re.fullmatch(r"\d+(?:\.\d+)+|[（(]\d+[)）]", s)
    )
    if texts and numish / len(texts) >= 0.2:
        return False
    return True


def _grid_pipe(grid: list[list[str]]) -> str:
    if not _grid_usable(grid):
        return ""
    width = max(len(r) for r in grid)
    lines: list[str] = []
    blanks = 0
    for row in grid:
        cells = list(row) + [""] * (width - len(row))
        if not any((c or "").strip() for c in cells):
            blanks += 1
            if blanks > 24:
                continue
        else:
            blanks = 0
        lines.append(" | ".join(cells[:width]))
    if len(lines) < 2:
        return ""
    return "\n".join(lines)


def _compose_ruled_page(lines: list[dict], tables: list[tuple[tuple, str]]) -> str:
    """线框表按纵坐标插回正文，表内的字不再重复出现。"""
    boxes = [box for box, _pipe in tables]

    def inside(top: float, bottom: float) -> bool:
        mid = (top + bottom) / 2
        for x0, t, x1, b in boxes:
            if t - 2 <= mid <= b + 2:
                return True
        return False

    events: list[tuple[float, int, str]] = []
    for ln in lines:
        text = (ln.get("text") or "").strip()
        if not text or inside(float(ln.get("top") or 0), float(ln.get("bottom") or 0)):
            continue
        events.append((float(ln.get("top") or 0), 0, text))
    for box, pipe in tables:
        if pipe:
            events.append((float(box[1]), 1, pipe))
    events.sort(key=lambda e: (e[0], e[1]))
    return "\n".join(e[2] for e in events if e[2].strip())


def _table_cells(table) -> list[list[tuple[tuple | None, str | None]]]:
    extracted = table.extract() or []
    rows = []
    for ri, row in enumerate(table.rows):
        texts = extracted[ri] if ri < len(extracted) else []
        rows.append(
            [
                (box, texts[ci] if ci < len(texts) else None)
                for ci, box in enumerate(row.cells)
            ]
        )
    return rows


def _ruled_page_text(plumber_page) -> str:
    try:
        found = plumber_page.find_tables(_RULED_TABLE) or []
    except Exception:  # noqa: BLE001
        return ""
    tables: list[tuple[tuple, str]] = []
    for table in found:
        try:
            pipe = _grid_pipe(_grid_from_cells(_table_cells(table)))
        except Exception:  # noqa: BLE001
            continue
        if not pipe:
            continue
        box = tuple(table.bbox)
        if box[2] - box[0] < 40 or box[3] - box[1] < 12:
            continue
        tables.append((box, pipe))
    if not tables:
        return ""
    try:
        lines = plumber_page.extract_text_lines() or []
    except Exception:  # noqa: BLE001
        lines = []
    return _compose_ruled_page(lines, tables)


async def _read_pdf(path: Path) -> ExtractResult:
    from pypdf import PdfReader

    settings = get_settings()
    min_chars = max(8, int(settings.ocr_page_text_min_chars))
    max_ocr = max(0, int(settings.ocr_max_pages))

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError(ErrorCode.BAD_REQUEST, f"PDF 无法打开：{exc}", status_code=400) from exc

    parts: list[str] = []
    ocr_used = 0
    ocr_capped = False
    page_count = len(reader.pages)
    for i, page in enumerate(reader.pages):
        try:
            page_text = _normalize(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            page_text = ""
        score = _count_chars(page_text)
        blocks: list[str] = []
        if score >= min_chars:
            blocks.append(page_text)
            try:
                images = list(getattr(page, "images", []) or [])
            except Exception:  # noqa: BLE001
                images = []
            for img in images:
                if ocr_used >= max_ocr:
                    ocr_capped = True
                    break
                data = getattr(img, "data", None)
                if not data or len(data) < 8000:
                    continue
                try:
                    ocr_text = _normalize(await _ocr_image(data))
                except AppError:
                    continue
                if _count_chars(ocr_text) >= 4:
                    blocks.append(f"[图]\n{ocr_text}")
                    ocr_used += 1
                    logger.info("pdf page=%s embedded image ocr chars=%s", i + 1, len(ocr_text))
        else:
            if ocr_used >= max_ocr:
                logger.warning("pdf ocr page cap reached at %s", max_ocr)
                ocr_capped = True
                continue
            try:
                jpeg = await asyncio.to_thread(_render_pdf_page_jpeg, path, i)
                ocr_text = _normalize(await _ocr_image(jpeg))
            except AppError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise AppError(
                    ErrorCode.INTERNAL, f"PDF 第{i + 1}页渲染失败：{exc}", status_code=500
                ) from exc
            ocr_used += 1
            if _count_chars(ocr_text) >= 4:
                blocks.append(ocr_text)
            logger.info("pdf page=%s scanned ocr chars=%s", i + 1, len(ocr_text))

        if blocks:
            parts.append(f"[第{i + 1}页]\n" + "\n".join(blocks))

    return ExtractResult(
        text="\n\n".join(parts),
        page_count=page_count,
        ocr_pages=ocr_used,
        ocr_capped=ocr_capped,
    )


async def _read_docx(path: Path) -> ExtractResult:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    settings = get_settings()
    max_ocr = max(0, int(settings.ocr_max_pages))

    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError(ErrorCode.BAD_REQUEST, f"Word 无法打开：{exc}", status_code=400) from exc

    parts: list[str] = []
    ocr_used = 0
    ocr_capped = False
    seen_rids: set[str] = set()

    async def _ocr_blips(element) -> None:
        nonlocal ocr_used, ocr_capped
        for blip in element.findall(f".//{_BLIP_NS}"):
            if ocr_used >= max_ocr:
                ocr_capped = True
                return
            rid = blip.get(_EMBED_NS)
            if not rid or rid in seen_rids:
                continue
            seen_rids.add(rid)
            try:
                blob = doc.part.related_parts[rid].blob
            except Exception:  # noqa: BLE001
                continue
            if not blob or len(blob) < 4000:
                continue
            try:
                ocr_text = _normalize(await _ocr_image(blob))
            except AppError:
                continue
            if _count_chars(ocr_text) >= 4:
                parts.append(f"[图]\n{ocr_text}")
                ocr_used += 1

    for block in doc.element.body:
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(block, doc).text.strip()
            if text:
                parts.append(text)
            await _ocr_blips(block)
        elif tag == "tbl":
            table_text = _docx_table_text(Table(block, doc))
            if table_text:
                parts.append(table_text)
            await _ocr_blips(block)
    return ExtractResult(text="\n".join(parts), ocr_pages=ocr_used, ocr_capped=ocr_capped)


def _docx_table_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [" ".join(c.text.split()) for c in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _xlsx_cell_text(cached: object, formula: object) -> tuple[str, bool]:
    """优先缓存计算值；空则回退公式原文。返回 (文本, 是否用了公式)。"""
    cached_s = _cell_str(cached)
    if cached_s:
        return cached_s, False
    if isinstance(formula, str) and formula.startswith("="):
        return formula.strip(), True
    fallback = _cell_str(formula)
    return fallback, False


async def _read_xlsx(path: Path) -> ExtractResult:
    from openpyxl import load_workbook

    settings = get_settings()
    max_ocr = max(0, int(settings.ocr_max_pages))

    wb_data = None
    wb_form = None
    try:
        wb_data = load_workbook(filename=str(path), data_only=True)
        wb_form = load_workbook(filename=str(path), data_only=False)
    except Exception as exc:  # noqa: BLE001
        if wb_data is not None:
            wb_data.close()
        raise AppError(ErrorCode.BAD_REQUEST, f"Excel 无法打开：{exc}", status_code=400) from exc

    try:
        sheets_data = wb_data.worksheets[:_XLSX_MAX_SHEETS]
        form_by_name = {s.title: s for s in wb_form.worksheets}
        parts: list[str] = []
        formula_fallback = False
        for sheet in sheets_data:
            form_sheet = form_by_name.get(sheet.title)
            rows: list[str] = []
            truncated = False
            data_iter = sheet.iter_rows(
                min_row=1, max_row=_XLSX_MAX_ROWS, max_col=_XLSX_MAX_COLS, values_only=True
            )
            form_iter = None
            if form_sheet is not None:
                form_iter = form_sheet.iter_rows(
                    min_row=1, max_row=_XLSX_MAX_ROWS, max_col=_XLSX_MAX_COLS, values_only=True
                )
            for data_row in data_iter:
                form_row = next(form_iter, ()) if form_iter is not None else ()
                dcells = list(data_row or ())
                fcells = list(form_row or ())
                width = max(len(dcells), len(fcells), 0)
                cells: list[str] = []
                for col in range(min(width, _XLSX_MAX_COLS)):
                    cached = dcells[col] if col < len(dcells) else None
                    formula = fcells[col] if col < len(fcells) else None
                    text, used_formula = _xlsx_cell_text(cached, formula)
                    if used_formula:
                        formula_fallback = True
                    cells.append(text)
                if any(cells):
                    rows.append(" | ".join(cells))
            if not rows:
                continue
            if int(getattr(sheet, "max_row", 0) or 0) > _XLSX_MAX_ROWS:
                truncated = True
            title = sheet.title.strip() or "Sheet"
            body = "\n".join(rows)
            if truncated:
                body += "\n[已截断：超过行数上限]"
            parts.append(f"## {title}\n{body}")
        if len(wb_data.worksheets) > _XLSX_MAX_SHEETS:
            parts.append("[已截断：超过工作表数量上限]")

        ocr_used = 0
        ocr_capped = False
        for img in getattr(wb_form, "_images", []) or []:
            if ocr_used >= max_ocr:
                ocr_capped = True
                break
            try:
                blob = img._data()  # noqa: SLF001
            except Exception:  # noqa: BLE001
                continue
            if not blob or len(blob) < 4000:
                continue
            try:
                ocr_text = _normalize(await _ocr_image(blob))
            except AppError:
                continue
            if _count_chars(ocr_text) >= 4:
                parts.append(f"[图]\n{ocr_text}")
                ocr_used += 1

        text = "\n\n".join(parts)
        if formula_fallback:
            text = f"{text}\n\n{_FORMULA_NOTE}".strip()
        return ExtractResult(
            text=text,
            page_count=len(sheets_data),
            ocr_pages=ocr_used,
            ocr_capped=ocr_capped,
            formula_fallback=formula_fallback,
        )
    finally:
        if wb_data is not None:
            wb_data.close()
        if wb_form is not None:
            wb_form.close()


async def _read_xls(path: Path) -> ExtractResult:
    try:
        import xlrd
    except ImportError as exc:
        raise AppError(
            ErrorCode.VALIDATION,
            "未安装 xlrd，无法解析旧版 .xls，请另存为 .xlsx",
            status_code=422,
        ) from exc
    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError(ErrorCode.BAD_REQUEST, f"Excel 无法打开：{exc}", status_code=400) from exc
    parts: list[str] = []
    for sheet in book.sheets()[:_XLSX_MAX_SHEETS]:
        rows: list[str] = []
        for r in range(min(int(sheet.nrows), _XLSX_MAX_ROWS)):
            cells = [
                str(sheet.cell_value(r, c)).strip()
                for c in range(min(int(sheet.ncols), _XLSX_MAX_COLS))
            ]
            if any(cells):
                rows.append(" | ".join(cells))
        if not rows:
            continue
        if int(sheet.nrows) > _XLSX_MAX_ROWS:
            rows.append("[已截断：超过行数上限]")
        parts.append(f"## {sheet.name}\n" + "\n".join(rows))
    return ExtractResult(text="\n\n".join(parts), page_count=len(book.sheets()))


async def _read_pptx(path: Path) -> ExtractResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    settings = get_settings()
    max_ocr = max(0, int(settings.ocr_max_pages))

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError(ErrorCode.BAD_REQUEST, f"PPT 无法打开：{exc}", status_code=400) from exc

    parts: list[str] = []
    ocr_used = 0
    ocr_capped = False
    slides = list(prs.slides)
    for i, slide in enumerate(slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    blocks.append(text)
            shape_type = getattr(shape, "shape_type", None)
            if shape_type == MSO_SHAPE_TYPE.PICTURE:
                if ocr_used >= max_ocr:
                    ocr_capped = True
                    continue
                try:
                    blob = shape.image.blob
                except Exception:  # noqa: BLE001
                    continue
                if not blob or len(blob) < 4000:
                    continue
                try:
                    ocr_text = _normalize(await _ocr_image(blob))
                except AppError:
                    continue
                if _count_chars(ocr_text) >= 4:
                    blocks.append(f"[图]\n{ocr_text}")
                    ocr_used += 1
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    blocks.append(f"[备注]\n{notes}")
        except Exception:  # noqa: BLE001
            pass
        if blocks:
            parts.append(f"[第{i}页]\n" + "\n".join(blocks))

    return ExtractResult(
        text="\n\n".join(parts),
        page_count=len(slides),
        ocr_pages=ocr_used,
        ocr_capped=ocr_capped,
    )
