"""Incomplete outline/TOC must never hide the native business front chapters."""

import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from api.services.tenders.document import _pageref_bookmark
from api.services.tenders.native_format import body_headings, ensure_toc_rows
from api.services.tenders.schema import BidBrief, OutlineItem

TITLES = [
    "一、投标函",
    "（一）投标函附录",
    "二、法定代表人身份证明",
    "三、授权委托书",
    "四、投标保证金",
    "五、已标价工程量清单",
    "六、项目管理机构",
    "七、资格审查资料",
]


@pytest.mark.parametrize("toc", ["missing", "partial", "reversed", "complete"])
@pytest.mark.parametrize("outline", ["empty", "wrong"])
def test_native_business_toc_uses_all_physical_chapters(toc, outline):
    doc = Document()
    doc.add_paragraph("投标文件")
    doc.add_paragraph("投标人：（盖章）")
    if toc != "missing":
        doc.add_paragraph("目录")
        source = TITLES[5:] if toc == "partial" else (TITLES[::-1] if toc == "reversed" else TITLES)
        for title in source:
            doc.add_paragraph(title).add_run().add_tab()
        doc.add_page_break()
    targets = []
    for title in TITLES:
        p = doc.add_paragraph(title)
        if not title.startswith("（"):
            p.paragraph_format.page_break_before = True
        targets.append(p._p)
        doc.add_paragraph("该章节原版正文。").add_run("实际内容。" * 30)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "营业执照号"
    table.cell(0, 1).text = "类似项目名称"
    brief = BidBrief(
        layoutMode="outline",
        outlineChapter="第八章 投标文件格式",
        outlineItems=[]
        if outline == "empty"
        else [OutlineItem(id="quote", title="已标价工程量清单", kind="quote", level=4)],
    )
    before = [p.text for p in doc.paragraphs if p._p in targets]
    ensure_toc_rows(doc, brief)
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    expected = (
        TITLES
        if toc in {"missing", "complete"}
        else (TITLES[5:] if toc == "partial" else TITLES[::-1])
    )
    assert [p.text.split("\t")[0] for p in entries] == expected
    assert [p.text for p in doc.paragraphs if p._p in targets] == before
    assert len(body_headings(doc, brief)) == len(TITLES)
    pages = [int(p.text.rsplit("\t", 1)[1]) for p in entries]
    if toc != "reversed":
        assert pages == sorted(pages)
        assert pages[0] == 1 and pages[1] >= pages[0]
    if toc in {"missing", "complete"}:
        assert entries[1].paragraph_format.left_indent > entries[0].paragraph_format.left_indent
    ordered_targets = [targets[TITLES.index(title)] for title in expected]
    for row, target in zip(entries, ordered_targets, strict=True):
        name = _pageref_bookmark(row)
        assert any(n.get(qn("w:name")) == name for n in target.iter(qn("w:bookmarkStart")))
    once = doc.element.xml
    ensure_toc_rows(doc, brief)
    assert doc.element.xml == once
    names = [n.get(qn("w:name")) for n in doc.element.iter(qn("w:bookmarkStart"))]
    assert len(names) == len(set(names))


def test_nonexistent_original_entry_never_borrows_another_chapter_anchor():
    doc = Document()
    doc.add_paragraph("目录")
    for title in ("一、投标函", "二、根本不存在的章节", "三、授权委托书"):
        doc.add_paragraph(title).add_run().add_tab()
    doc.add_page_break()
    for title in ("一、投标函", "三、授权委托书"):
        doc.add_paragraph(title, style="Heading 1")
        doc.add_paragraph("原版正文。")
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    ensure_toc_rows(doc, brief)
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert not any("不存在" in p.text for p in entries)
    assert len({_pageref_bookmark(p) for p in entries}) == 2
    assert all(re.search(r"\t\d+$", p.text) for p in entries)


def test_unlisted_plain_letter_without_page_break_is_not_swallowed_by_toc():
    doc = Document()
    doc.add_paragraph("目录")
    doc.add_paragraph("五、已标价工程量清单").add_run().add_tab()
    doc.add_paragraph("一、投标函")
    doc.add_paragraph("致：采购人")
    doc.add_paragraph("我方承诺按本项目要求履约。")
    doc.add_paragraph("五、已标价工程量清单")
    doc.add_paragraph("本节原版报价正文。")
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    ensure_toc_rows(doc, brief)
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert [p.text.split("\t")[0] for p in entries] == ["五、已标价工程量清单"]
    assert "致：采购人" in [p.text for p in doc.paragraphs]


def test_repeated_child_titles_have_distinct_body_bookmarks():
    doc = Document()
    for title in ("一、项目管理机构", "（一）承诺书", "二、资格审查资料", "（一）承诺书"):
        doc.add_paragraph(title, style="Heading 1")
        doc.add_paragraph("原版正文。")
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    ensure_toc_rows(doc, brief)
    rows = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert len(rows) == 4
    assert len({_pageref_bookmark(p) for p in rows}) == 4
    once = doc.element.xml
    ensure_toc_rows(doc, brief)
    assert doc.element.xml == once
