from io import BytesIO

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from PIL import Image

from api.services.tenders import placeholders as ph
from api.services.tenders.schema import PlaceholderItem


def assert_gap(doc, node, *, before):
    assert ph._is_slot_gap(node)
    assert not ''.join(node.itertext()).strip()
    para = Paragraph(node, doc._body)
    # Word 的厘米值按 twip 存储，有不到 0.002 cm 的舍入。
    assert para.paragraph_format.line_spacing.cm == pytest.approx(0.5, abs=0.002)
    assert para.style.paragraph_format.space_before == 0
    assert para.style.paragraph_format.space_after == 0
    if before is not None:
        assert para.paragraph_format.keep_with_next is before


def test_resume_table_and_adjacent_boxes_share_half_cm_gap():
    doc = Document()
    note = doc.add_paragraph('项目经理须附相关证明材料及承诺书扫描件。')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '姓名'
    following = doc.add_paragraph('下一节')
    originals = [(node, node.xml) for node in doc.element.body]
    ph._drop_after(doc, note._p, PlaceholderItem(key='proof', title='相关证明材料'), [])
    first = table._tbl.getnext().getnext()
    ph._drop_after(doc, first, PlaceholderItem(key='promise', title='承诺书'), [])
    assert note._p.getnext() is table._tbl
    assert_gap(doc, table._tbl.getnext(), before=True)
    shared = first.getnext()
    assert_gap(doc, shared, before=None)
    second = shared.getnext()
    assert ph._is_slot_box(second)
    assert_gap(doc, second.getnext(), before=False)
    assert second.getnext().getnext() is following._p
    assert all(node.xml == xml for node, xml in originals)
    saved = BytesIO()
    doc.save(saved)
    assert len([node for node in Document(saved).element.body if ph._is_slot_gap(node)]) == 3


@pytest.mark.parametrize('kind', ['image', 'pdf', 'id_pair'])
def test_inserted_scans_have_same_external_gaps(tmp_path, kind):
    doc = Document()
    note = doc.add_paragraph('资料说明')
    following = doc.add_paragraph('下一节')
    path = tmp_path / ('scan.pdf' if kind == 'pdf' else 'scan.png')
    if kind == 'pdf':
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(path)
        pdf.close()
    else:
        Image.new('RGB', (400, 250), 'white').save(path)
    original_xml = note._p.xml, following._p.xml
    slot = PlaceholderItem(key='id_legal' if kind == 'id_pair' else 'license', title='资料')
    assert ph._drop_after(doc, note._p, slot, [path])
    gap = note._p.getnext()
    assert_gap(doc, gap, before=True)
    scan = gap.getnext()
    assert next(scan.iter(qn('w:drawing')), None) is not None
    assert_gap(doc, scan.getnext(), before=kind == 'id_pair')
    assert scan.getnext().getnext() is following._p
    para = Paragraph(scan, doc._body)
    assert para.paragraph_format.space_before == 0
    assert para.paragraph_format.space_after == 0
    assert original_xml == (note._p.xml, following._p.xml)


def test_public_box_moves_both_gaps_with_box():
    doc = Document()
    anchor = doc.add_paragraph('说明')
    following = doc.add_paragraph('下一节')
    ph.draw_placeholder_box(doc, after=anchor)
    assert_gap(doc, anchor._p.getnext(), before=True)
    box = anchor._p.getnext().getnext()
    assert ph._is_slot_box(box)
    assert_gap(doc, box.getnext(), before=False)
    assert box.getnext().getnext() is following._p


@pytest.mark.parametrize('anchor_kind', ['paragraph', 'cell', 'row'])
def test_cell_or_row_anchor_is_promoted_outside_original_table(anchor_kind):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    cell_para = table.cell(0, 0).paragraphs[0]
    cell_para.text = '应附证明扫描件'
    anchor = {'paragraph': cell_para._p, 'cell': table.cell(0, 0)._tc, 'row': table.rows[0]._tr}[anchor_kind]
    following = doc.add_paragraph('下一节')
    original = table._tbl.xml
    ph._drop_after(doc, anchor, PlaceholderItem(key='proof', title='证明材料'), [])
    assert table._tbl.xml == original
    gap = table._tbl.getnext()
    assert_gap(doc, gap, before=True)
    box = gap.getnext()
    assert ph._is_slot_box(box) and box.getparent() is doc.element.body
    assert box.getnext().getnext() is following._p


def test_wrap_none_floating_table_gets_clearance_before_attachment():
    doc = Document()
    carrier = doc.add_paragraph()
    drawing = OxmlElement('w:drawing')
    anchor = OxmlElement('wp:anchor')
    extent = OxmlElement('wp:extent')
    extent.set('cy', str(Cm(6)))
    position = OxmlElement('wp:positionV')
    position.set('relativeFrom', 'paragraph')
    offset = OxmlElement('wp:posOffset')
    offset.text = str(Cm(1))
    position.append(offset)
    anchor.extend([position, extent, OxmlElement('wp:wrapNone')])
    textbox = OxmlElement('w:txbxContent')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '原简历表'
    textbox.append(table._tbl)
    anchor.append(textbox)
    drawing.append(anchor)
    carrier.add_run()._r.append(drawing)
    original = carrier._p.xml
    following = doc.add_paragraph('下一节')
    ph._drop_after(doc, table.cell(0, 0).paragraphs[0]._p, PlaceholderItem(key='id', title='身份证'), [])
    assert carrier._p.xml == original
    clearance = carrier._p.getnext()
    assert ph._is_slot_clearance(clearance)
    assert Paragraph(clearance, doc._body).paragraph_format.line_spacing.cm == pytest.approx(7, abs=0.002)
    box = clearance.getnext().getnext()
    assert ph._is_slot_box(box) and box.getparent() is doc.element.body
    assert box.getnext().getnext() is following._p


def test_placeholder_is_one_normal_bordered_paragraph_not_table_or_shape():
    doc = Document()
    ph.draw_placeholder_box(doc, PlaceholderItem(key='proof', title='相关证明材料'))
    assert not doc.tables
    box = next(p for p in doc.paragraphs if ph._is_slot_box(p._p))
    assert box.text.strip() == '此处粘贴相关证明材料\n（在此粘贴扫描件）'
    assert not any(next(box._p.iter(qn(tag)), None) is not None for tag in ['w:tbl', 'w:drawing', 'w:pict', 'w:framePr'])
    borders = box._p.find('w:pPr/w:pBdr', box._p.nsmap)
    assert {b.tag for b in borders} == {qn('w:' + e) for e in ['top', 'left', 'bottom', 'right']}
    assert all(b.get(qn('w:val')) == 'dashed' for b in borders)
