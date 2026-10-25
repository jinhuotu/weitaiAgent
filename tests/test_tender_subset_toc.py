"""New subset TOCs share hierarchy, anchors and page-cache behavior at both entrances."""

from __future__ import annotations

from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from api.services.tenders.assemble import (
    _anchor_section,
    _body_mark,
    _fill_toc_pages,
    _write_toc,
    assemble_bid_docx,
)
from api.services.tenders.document import (
    _estimate_bookmark_pages,
    _pageref_bookmark,
    _subset_bookmarks,
)
from api.services.tenders.outline import (
    _extracted_item_level,
    _items_from_catalog,
    normalize_subset_outline,
    subset_toc_titles,
)
from api.services.tenders.schema import DocumentFormat, OutlineItem, default_brief

TABLE = "| 姓名 | 职务 |\n| 张三 | 经理 |"


def _items():
    rows = [
        ("root", "六、项目管理机构", 1, "项目管理机构说明。"),
        ("org", "（三）项目管理机构组成表", 3, TABLE),
        ("pm", "（四）项目经理主要工作经历", 2, TABLE),
        ("detail", "1.1 任职明细", 3, "真实三级节点说明。"),
        ("detail2", "人员分工", 3, "无编号的真实三级节点说明。"),
        ("sub", "（五）拟分包项目情况表", 3, TABLE),
        ("next", "七、拟分包项目情况表", 3, TABLE),
        ("company", "（一）投标人基本情况表", 3, TABLE),
        ("perf", "（二）近年类似项目情况表", 2, TABLE),
    ]
    return [
        OutlineItem(id=key, title=title, kind="other", level=level, body=body)
        for key, title, level, body in rows
    ]


def _brief(items=None):
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = _items() if items is None else items
    return brief


def _toc(doc):
    return [p for p in doc.paragraphs if _pageref_bookmark(p)]


def _cache(para):
    active = False
    result = []
    for run in para._p.iter(qn("w:r")):
        field = run.find(qn("w:fldChar"))
        if field is not None:
            if field.get(qn("w:fldCharType")) == "separate":
                active = True
            elif field.get(qn("w:fldCharType")) == "end":
                active = False
        elif active:
            result.extend(node.text or "" for node in run.iter(qn("w:t")))
    return "".join(result)


def _assert_toc(doc, expected_levels):
    paras = _toc(doc)
    assert len(paras) == len(expected_levels)
    positions = set()
    for para, level in zip(paras, expected_levels, strict=True):
        assert float(para.paragraph_format.left_indent.cm) == pytest.approx(
            0.74 * (level - 1),
            abs=0.01,
        )
        assert _cache(para).strip()
        tab = para._p.find(qn("w:pPr")).find(qn("w:tabs")).find(qn("w:tab"))
        assert tab.get(qn("w:val")) == "right"
        assert tab.get(qn("w:leader")) == "dot"
        positions.add(tab.get(qn("w:pos")))
    assert len(positions) == 1
    names = [el.get(qn("w:name")) for el in doc.element.iter(qn("w:bookmarkStart"))]
    assert len(names) == len(set(names))
    pages = _estimate_bookmark_pages(doc)
    base = min(pages.values())
    for para in paras:
        name = _pageref_bookmark(para)
        assert names.count(name) == 1
        assert _cache(para) == str(max(1, pages[name] - base + 1))
        host = next(
            el for el in doc.element.iter(qn("w:bookmarkStart")) if el.get(qn("w:name")) == name
        )
        assert host.getparent().tag == qn("w:p")
        assert all(el.tag != qn("w:tbl") for el in host.iterancestors())
    return paras


def test_normalize_mixed_siblings_real_third_level_and_top_boundary():
    original = _items()
    before = [item.model_dump() for item in original]
    normalized = normalize_subset_outline(original)
    assert [item.level for item in normalized] == [1, 2, 2, 3, 3, 2, 1, 2, 2]
    assert [item.model_dump() for item in original] == before
    assert normalized == normalize_subset_outline(normalized)
    titles = subset_toc_titles(normalized)
    assert titles == [
        "一、项目管理机构",
        "（一）项目管理机构组成表",
        "（二）项目经理主要工作经历",
        "1.2.1 任职明细",
        "1.2.2 人员分工",
        "（三）拟分包项目情况表",
        "二、拟分包项目情况表",
        "（一）投标人基本情况表",
        "（二）近年类似项目情况表",
    ]


@pytest.mark.parametrize("raw", ["（三）项目管理机构组成表", "A.公司简介"])
def test_both_extraction_paths_use_the_same_subheading_rule(raw):
    assert _extracted_item_level(raw) == 2
    assert _items_from_catalog([raw])[0].level == 2


@pytest.mark.parametrize(
    "titles, expected",
    [
        (["一、工程概况", "1.1 范围", "1.2 条件", "1.2.1 说明", "1.3 其他"], [1, 2, 2, 3, 2]),
        (["一、工程概况", "（一）范围", "1.1 说明", "1.2 附表", "（二）条件"], [1, 2, 3, 3, 2]),
        (["（三）范围", "1.1 说明", "1.2 附表", "（四）条件"], [1, 2, 2, 1]),
    ],
)
def test_number_paths_and_removed_parent(titles, expected):
    items = [OutlineItem(title=title, level=3) for title in titles]
    assert [item.level for item in normalize_subset_outline(items)] == expected


def test_number_in_body_corrects_old_bare_title():
    item = OutlineItem(
        title="项目经理主要工作经历", level=3, body="（二）项目经理主要工作经历\n正文。"
    )
    normalized = normalize_subset_outline(
        [
            OutlineItem(title="六、项目管理机构"),
            OutlineItem(title="（一）项目管理机构组成表", level=2),
            item,
        ]
    )
    assert [entry.level for entry in normalized] == [1, 2, 2]


def test_numeric_parentheses_keep_real_third_level_and_single_numbers_are_replaced():
    items = [
        OutlineItem(title="一、工程概况"),
        OutlineItem(title="（一）工程范围", level=2),
        OutlineItem(title="（1）范围明细", level=3),
        OutlineItem(title="（2）补充明细", level=3),
        OutlineItem(title="（二）施工条件", level=3),
        OutlineItem(title="1. 设备明细", level=3),
        OutlineItem(title="2. 人员明细", level=3),
    ]
    normalized = normalize_subset_outline(items)
    assert [item.level for item in normalized] == [1, 2, 3, 3, 2, 3, 3]
    assert subset_toc_titles(normalized)[-2:] == ["1.2.1 设备明细", "1.2.2 人员明细"]


def test_assemble_subset_both_mixed_tables_and_repeat_generation(tmp_path, monkeypatch):
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: None)
    monkeypatch.setattr("api.services.tenders.pdf_format._ready_format_docx", lambda _: None)
    brief = _brief()
    brief.splitSlice = True
    before = brief.model_dump()
    packages = []
    for name in ("first", "second"):
        dest, _ = assemble_bid_docx(
            brief, tmp_path / f"{name}.docx", volume="business", catalog_slots=[], catalog_media={}
        )
        doc = Document(dest)
        paras = _assert_toc(doc, [1, 2, 2, 3, 3, 2, 1, 2, 2])
        assert paras[3].text.startswith("1.2.1 任职明细")
        for item in brief.outlineItems:
            assert sum(p.text == item.title for p in doc.paragraphs) == 1
        with ZipFile(dest) as package:
            packages.append({entry: package.read(entry) for entry in package.namelist()})
    assert packages[0] == packages[1]
    assert brief.model_dump() == before


def test_pdf_format_technical_subset_tables_and_repeat_install(tmp_path, monkeypatch):
    from api.services.tenders.pdf_format import (
        _collect_subset_tech_heads,
        _install_tech_front,
        assemble_from_word_format,
    )

    brief = _brief()
    brief.invitationId = "subset_test"
    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("商务正文。")
    raw.add_paragraph("施工组织设计")
    for item in brief.outlineItems:
        if item.id == "org":
            table = raw.add_table(rows=2, cols=2)
            table.cell(0, 0).text = item.title
            table.cell(1, 0).text = "姓名"
        else:
            heading = raw.add_paragraph(item.title)
            # 源 Word 样式同样可能过期；编号应纠正混杂的二/三级样式。
            heading.style = f"Heading {item.level}"
            raw.add_paragraph("本节正文。")
    src = tmp_path / "source.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    source_bytes = src.read_bytes()
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: src)
    before = brief.model_dump()
    for name in ("first", "second"):
        dest = tmp_path / f"{name}.docx"
        assemble_from_word_format(
            brief,
            dest,
            volume="technical",
            catalog_slots=[],
            catalog_media={},
        )
        doc = Document(dest)
        _assert_toc(doc, [1, 2, 2, 3, 3, 2, 1, 2, 2])
        assert sum(p.text == "（三）项目管理机构组成表" for p in doc.paragraphs) == 1
        xml = doc.element.xml
        _install_tech_front(doc, brief, anon=False)
        assert doc.element.xml == xml
    assert brief.model_dump() == before
    assert src.read_bytes() == source_bytes
    # 单独重复收集不会在表前重复补标题。
    doc = Document(src)
    _collect_subset_tech_heads(doc, brief)
    count = len(doc.paragraphs)
    _collect_subset_tech_heads(doc, brief)
    assert len(doc.paragraphs) == count


def test_anchor_table_only_is_idempotent_and_not_inside_table():
    doc = Document()
    mark = _body_mark(doc)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    for _ in range(3):
        _anchor_section(doc, mark, "项目管理机构组成表", "toc_org", DocumentFormat(), subset=True)
    assert [p.text for p in doc.paragraphs] == ["项目管理机构组成表"]
    assert table._tbl.getprevious().tag == qn("w:p")
    assert len(list(doc.element.iter(qn("w:bookmarkStart")))) == 1


def test_bookmarks_missing_duplicate_ids_and_existing_name():
    from api.services.tenders.document import _bookmark_paragraph

    doc = Document()
    _bookmark_paragraph(doc.add_paragraph("旧标题"), "toc_dup")
    items = [OutlineItem(id=key, title="标题") for key in ("dup", "dup", "", "o03", "a b", "a-b")]
    names = _subset_bookmarks(doc, items)
    assert len(names) == len(set(names))
    assert "toc_dup" not in names
    assert all(name.startswith("toc_") and " " not in name for name in names)


@pytest.mark.parametrize("estimated", [{}, {"toc_a": 8}])
def test_estimate_failure_keeps_nonempty_placeholder_and_indent(monkeypatch, estimated):
    doc = Document()
    items = [
        OutlineItem(id="a", title="一、目录章"),
        OutlineItem(id="b", title="（一）目录子项", level=3),
    ]
    brief = _brief(items)
    _write_toc(doc, items, brief, DocumentFormat(), ["toc_a", "toc_b"], subset=True)
    rows = _toc(doc)
    before = [row._p.find(qn("w:pPr")).xml for row in rows]
    monkeypatch.setattr(
        "api.services.tenders.assemble._estimate_bookmark_pages",
        lambda _: estimated,
    )
    _fill_toc_pages(doc, DocumentFormat(), subset=True)
    assert _cache(rows[0]) == ("1" if estimated else "待更新")
    assert _cache(rows[1]) == "待更新"
    assert [row._p.find(qn("w:pPr")).xml for row in rows] == before


@pytest.mark.parametrize("remove_result", [False, True])
def test_empty_field_cache_repaired_even_when_no_estimates(monkeypatch, remove_result):
    doc = Document()
    brief = _brief([OutlineItem(id="a", title="目录章")])
    _write_toc(doc, brief.outlineItems, brief, DocumentFormat(), ["toc_a"], subset=True)
    para = _toc(doc)[0]
    for run in list(para._p.iter(qn("w:r"))):
        node = run.find(qn("w:t"))
        if node is not None and node.text == "待更新":
            if remove_result:
                run.remove(node)
            else:
                node.text = ""
    monkeypatch.setattr("api.services.tenders.assemble._estimate_bookmark_pages", lambda _: {})
    _fill_toc_pages(doc, DocumentFormat(), subset=True)
    assert _cache(para) == "待更新"


def test_pdf_table_without_title_uses_unique_header_and_is_repeatable():
    from api.services.tenders.pdf_format import _collect_subset_tech_heads, _install_tech_front

    items = [
        OutlineItem(id="root", title="六、项目管理机构", level=1),
        OutlineItem(
            id="org", title="（一）项目管理机构组成表", level=3, kind="tech_chapter", body=TABLE
        ),
    ]
    brief = _brief(items)
    doc = Document()
    doc.add_paragraph(items[0].title)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "职务"
    table.cell(1, 0).text = "张三"
    table_xml = table._tbl.xml
    for _ in range(3):
        _collect_subset_tech_heads(doc, brief)
    assert [p.text for p in doc.paragraphs] == [item.title for item in items]
    assert table._tbl.xml == table_xml
    _install_tech_front(doc, brief, anon=False)
    _fill_toc_pages(doc, DocumentFormat(), subset=True)
    _assert_toc(doc, [1, 2])


def test_technical_assemble_entrance_and_duplicate_ids(tmp_path, monkeypatch):
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: None)
    monkeypatch.setattr("api.services.tenders.pdf_format._ready_format_docx", lambda _: None)
    items = [
        OutlineItem(
            id="duplicate", title="一、施工组织设计", kind="tech_chapter", body="技术正文说明。"
        ),
        OutlineItem(
            id="duplicate", title="（一）质量保证措施", level=3, kind="tech_chapter", body=TABLE
        ),
    ]
    brief = _brief(items)
    dest, _ = assemble_bid_docx(
        brief, tmp_path / "tech.docx", volume="technical", catalog_slots=[], catalog_media={}
    )
    doc = Document(dest)
    _assert_toc(doc, [1, 2])
    assert [_pageref_bookmark(p) for p in _toc(doc)] == ["toc_duplicate", "toc_duplicate_2"]


def test_ambiguous_table_headers_keep_selected_child_placeholder_without_false_anchor():
    from api.services.tenders.pdf_format import _install_tech_front

    brief = _brief(
        [
            OutlineItem(id="root", title="一、施工组织设计", kind="tech_chapter"),
            OutlineItem(
                id="crew", title="（一）人员安排", kind="tech_chapter", level=3, body=TABLE
            ),
        ]
    )
    doc = Document()
    doc.add_paragraph("一、施工组织设计")
    for _ in range(2):
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "职务"
    _install_tech_front(doc, brief, anon=False)
    _fill_toc_pages(doc, DocumentFormat(), subset=True)
    rows = _toc(doc)
    assert len(rows) == 2
    assert _cache(rows[0]) == "1"
    assert _cache(rows[1]) == "待更新"
    assert not any(p.text == "（一）人员安排" for p in doc.paragraphs)
    assert not any(
        el.get(qn("w:name")) == "toc_crew" for el in doc.element.iter(qn("w:bookmarkStart"))
    )


def test_business_item_with_same_headers_is_not_invented_in_technical_body():
    from api.services.tenders.pdf_format import _collect_subset_tech_heads

    brief = _brief([OutlineItem(id="company", title="投标人基本情况表", kind="other", body=TABLE)])
    doc = Document()
    doc.add_paragraph("施工组织设计")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "职务"
    _collect_subset_tech_heads(doc, brief)
    assert [p.text for p in doc.paragraphs] == ["施工组织设计"]
