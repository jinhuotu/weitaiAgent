"""原始章节展示/预览不受检索脱敏过滤影响。"""

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from api.services.tenders import bid_archive, bid_marks, history_preview
from api.services.tenders.document import _bookmark_paragraph, _set_bookmark_outline
from db.models.tender import TenderRecord


def sample(path):
    doc = Document()
    doc.add_paragraph("投 标 文 件", "Heading 2")
    for title in ["法定代表人身份证明", "拟分包项目情况表", "已标价工程量清单", "类似业绩"]:
        doc.add_paragraph(title, "Heading 5")
        doc.add_paragraph("保留原始项目和报价128万元，不作脱敏处理。" * 4)
        doc.add_table(rows=1, cols=2).cell(0, 0).text = title + "原始表格"
        doc.add_paragraph("本章子项", "Heading 7")
        doc.add_paragraph("子项内容保留在父章节")
    doc.save(path)
    return doc


def test_heading_levels_legacy_and_explicit_and_toc(tmp_path):
    path = tmp_path / "source.docx"
    doc = sample(path)
    titles = [title for title, _, _ in bid_archive.chapter_ranges(doc)]
    assert titles == ["法定代表人身份证明", "拟分包项目情况表", "已标价工程量清单", "类似业绩"]
    assert all(
        "子项内容" in body
        for title, body in bid_archive.chapters_from_docx(path)
        if title != "封面"
    )
    ready = bid_archive.prepare_chapters(
        bid_archive.chapters_from_docx(path),
        record_id=1,
        volume="business",
    )
    assert not any("工程量清单" in item["title"] or "业绩" in item["title"] for item in ready)
    doc = Document()
    para = doc.add_paragraph("无编号一级标题", "Heading 9")
    _bookmark_paragraph(para, "root")
    before = para._p.xml
    _set_bookmark_outline(doc, {"root": 1})
    assert bid_archive._is_h1(para)
    node = para._p.find("./" + qn("w:pPr") + "/" + qn("w:outlineLvl"))
    para._p.find(qn("w:pPr")).remove(node)
    if not len(para._p.find(qn("w:pPr"))):
        para._p.remove(para._p.find(qn("w:pPr")))
    assert para._p.xml == before


@pytest.mark.asyncio
async def test_raw_catalog_and_readonly_preview_preserve_tables(tmp_path, monkeypatch):
    from api.services.tenders import assets, records

    path = tmp_path / "original.docx"
    original = sample(path)
    row = TenderRecord(
        id=1, public_id="record", user_id=3, docx_file="", brief_json={"_historySource": path.name}
    )
    monkeypatch.setattr(assets, "tenders_output_dir", lambda: tmp_path)
    monkeypatch.setattr(records, "_get_row", AsyncMock(return_value=row))
    monkeypatch.setattr(bid_marks, "chapters_for_record", AsyncMock(return_value=[
        {"volume": "business", "title": "旧原件里的章节", "docId": "old-doc"},
    ]))
    monkeypatch.setattr(bid_marks, "_marks_of", AsyncMock(return_value=[
        SimpleNamespace(volume="business", chapter_title="旧原件里的章节"),
    ]))
    data = await bid_marks.list_history_chapters(
        None,
        "record",
        user=SimpleNamespace(id=3),
        admin=False,
    )
    assert len(data["chapters"]) == 4
    for ch in data["chapters"]:
        payload = await history_preview.preview_chapter(
            None,
            "record",
            ch["chapterTitle"],
            ch["volume"],
            ch["previewKey"],
            user=SimpleNamespace(id=3),
            admin=False,
        )
        sliced = Document(BytesIO(payload))
        assert sliced.paragraphs[0].text == ch["chapterTitle"]
        assert len(sliced.tables) == 1
        assert "报价128万元" in sliced.paragraphs[1].text
        assert sliced.tables[0]._tbl.xml == next(
            t._tbl.xml for t in original.tables if ch["chapterTitle"] in t.cell(0, 0).text
        )
        protection = sliced.settings.element.find(qn("w:documentProtection"))
        assert protection.get(qn("w:edit")) == "readOnly"
        assert protection.get(qn("w:enforcement")) == "1"
        with ZipFile(BytesIO(payload)) as package:
            assert "word/styles.xml" in package.namelist()
    with pytest.raises(Exception) as denied:
        await history_preview.preview_chapter(
            None, "record", "类似业绩", "", "", user=SimpleNamespace(id=4), admin=False
        )
    assert denied.value.status_code == 403


@pytest.mark.asyncio
async def test_sensitive_only_upload_still_has_preview(tmp_path, monkeypatch):
    from test_tender_archive_feedback import MemorySession

    from api.services.tenders import assets

    doc = Document()
    doc.add_paragraph("一、已标价工程量清单")
    doc.add_paragraph("原始报价128万元")
    buf = BytesIO()
    doc.save(buf)
    monkeypatch.setattr(assets, "tenders_output_dir", lambda: tmp_path)
    scheduled = []
    monkeypatch.setattr(bid_archive, "schedule_upload", lambda **kw: scheduled.append(kw))
    db = MemorySession()
    result = await bid_archive.accept_history_upload(
        db,
        data=buf.getvalue(),
        filename="报价.docx",
        user_id=3,
        username="tester",
        outcome="won",
        record_public_id="",
        admin=False,
    )
    assert result["queued"] == 0
    assert result["status"] == "completed"
    assert not scheduled
    assert history_preview.source_chapters(db.row)[0]["title"] == "一、已标价工程量清单"


def test_explicit_outline_keeps_same_style_child_in_parent():
    doc = Document()
    root = doc.add_paragraph("根章节", "Heading 5")
    child = doc.add_paragraph("子章节", "Heading 5")
    _bookmark_paragraph(root, "root")
    _bookmark_paragraph(child, "child")
    _set_bookmark_outline(doc, {"root": 1, "child": 2})
    assert [title for title, _, _ in bid_archive.chapter_ranges(doc)] == ["根章节"]


@pytest.mark.asyncio
async def test_preview_route_returns_word_with_no_store(monkeypatch):
    from api.routers.tenders import tenders_history_chapter_preview

    monkeypatch.setattr(history_preview, "preview_chapter", AsyncMock(return_value=b"word"))
    response = await tenders_history_chapter_preview(
        None,
        SimpleNamespace(id=3, is_superuser=False, roles=[]),
        "rec",
        "报价",
        "business",
        "1",
    )
    assert response.body == b"word"
    assert response.headers["cache-control"] == "no-store"
    assert "wordprocessingml.document" in response.headers["content-type"]
