"""Regression for the actual WPS layout: numeric style IDs, automatic list labels."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders.document import _pageref_bookmark
from api.services.tenders.generate import generate_bid
from api.services.tenders.native_format import _read_toc_entries, validate_native_toc
from api.services.tenders.pdf_format import _outline_level, toc_block_ids
from api.services.tenders.schema import BidBrief, OutlineItem, PlaceholderItem
from api.services.tenders.word_structure import outline_level

ITEMS = [
    ("投标函及投标函附录", 1),
    ("法定代表人身份证明", 1),
    ("授权委托书", 1),
    ("投标保证金", 1),
    ("已标价的工程量清单", 1),
    ("项目管理机构", 1),
    ("项目管理机构组成表", 2),
    ("项目经理简历表", 2),
    ("项目经理主要工作经历", 2),
    ("承诺书", 2),
    ("主要项目管理人员简历表", 2),
    ("拟分包项目情况表", 1),
    ("资格审查资料", 1),
    ("投标人基本情况表", 2),
    ("近年财务状况表", 2),
    ("近年完成的类似项目情况表", 2),
    ("正在施工的和新承接的项目情况表", 2),
    ("“国家企业信用信息公示系统”网站查询", 2),
    ("资格审查其他材料", 2),
    ("其他资料", 1),
    ("施工组织格式", 1),
]


def _style(doc, sid, name):
    element = OxmlElement("w:style")
    element.set(qn("w:type"), "paragraph")
    element.set(qn("w:styleId"), sid)
    title = OxmlElement("w:name")
    title.set(qn("w:val"), name)
    element.append(title)
    doc.styles.element.append(element)


def _list(doc, ident, pattern):
    root = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(ident))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "chineseCountingThousand"), ("lvlText", pattern)):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:val"), value)
        level.append(node)
    abstract.append(level)
    root.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(ident))
    target = OxmlElement("w:abstractNumId")
    target.set(qn("w:val"), str(ident))
    number.append(target)
    root.append(number)


def _number(p, ident):
    pr = p._p.get_or_add_pPr()
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "23")
    pr.append(style)
    num = OxmlElement("w:numPr")
    for tag, value in (("ilvl", "0"), ("numId", str(ident))):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:val"), value)
        num.append(node)
    pr.append(num)


@pytest.fixture
def numbered_format(tmp_path):
    doc = Document()
    _style(doc, "12", "Audit Body Text")
    _style(doc, "23", "Audit List Paragraph")
    for ident, pattern in (
        (181, "%1、"),
        (182, "(%1)"),
        (183, "(%1)"),
        (184, "%1、"),
        (185, "(%1)"),
    ):
        _list(doc, ident, pattern)
    doc.add_paragraph("投 标 文 件")
    doc.add_paragraph("目 录")
    directory = []
    in_qualification = False
    for title, rank in ITEMS:
        in_qualification = in_qualification or title == "资格审查资料"
        p = doc.add_paragraph(title)
        _number(p, 181 if rank == 1 else (183 if in_qualification else 182))
        directory.append(p._p)
    doc.add_page_break()
    bodies = []
    for title, rank in ITEMS:
        if rank == 1:
            p = doc.add_paragraph(title)
            _number(p, 184)
        else:
            # This is precisely where the former root-only lookahead failed.
            table = doc.add_table(rows=1, cols=1)
            p = table.cell(0, 0).paragraphs[0]
            p.text = title
            _number(p, 185)
        bodies.append(p._p)
        doc.add_paragraph("真实章节正文。")
        if title == "投标函及投标函附录":
            doc.add_paragraph("我方提交的投标文件组成如下：")
            doc.add_paragraph("施工组织设计；")
            doc.add_paragraph("项目管理机构；")
            doc.add_paragraph("投标人：（盖章）")
        if title == "资格审查其他材料":
            doc.add_paragraph("此处粘贴社保缴纳证明扫描件")
    path = tmp_path / "source.format.docx"
    doc.save(path)
    path.with_name(path.name + ".source").write_text("docx-raw", encoding="utf-8")
    return doc, path, directory, bodies


def test_numeric_style_id_is_not_a_heading_and_child_ilvl_zero_is_not_root(numbered_format):
    doc, _, directory, bodies = numbered_format
    assert outline_level(doc, directory[6]) == 0
    assert _outline_level(directory[6]) == 0
    block = toc_block_ids(doc)
    assert all(node in block for node in directory)
    assert all(node not in block for node in bodies)
    rows = _read_toc_entries(block, doc)
    assert len(rows) == 21
    assert rows[0] == {"title": "一、投标函及投标函附录", "level": 1}
    assert rows[6] == {"title": "(一)项目管理机构组成表", "level": 2}
    assert rows[11] == {"title": "七、拟分包项目情况表", "level": 1}
    assert rows[13] == {"title": "(一)投标人基本情况表", "level": 2}


@pytest.mark.parametrize("volume", ["business", "technical", "full"])
@pytest.mark.parametrize("with_media", [False, True])
def test_final_generator_checks_the_actual_numbered_template(
    numbered_format, tmp_path, monkeypatch, volume, with_media
):
    source, path, _, _ = numbered_format
    from api.services.tenders import pdf_format

    original = path.read_bytes()
    monkeypatch.setattr(pdf_format, "_word_format_file", lambda _: path)
    monkeypatch.setattr("api.services.tenders.generate.tenders_output_dir", lambda: tmp_path)
    monkeypatch.setattr(
        pdf_format,
        "_scrub_undefined_toc",
        lambda *_: pytest.fail("Legacy scrub must not remove native PAGEREF"),
    )
    brief = BidBrief(
        invitationId="audit-source",
        outlineChapter="第八章 投标文件格式",
        layoutMode="outline",
        generateVolume=volume,
        requiredSlotKeys=["social"],
        includeSlotKeys=["social"],
    )
    media = {}
    if with_media:
        from PIL import Image

        image = tmp_path / "social-scan.png"
        Image.new("RGB", (210, 297), "white").save(image)
        media["social"] = [image]
    result = generate_bid(
        brief,
        catalog_media=media,
        catalog_slots=[
            PlaceholderItem(key="social", title="社保缴纳证明", attachToSection="资格审查其他材料")
        ],
    )
    out = Document(
        tmp_path / (result["techDocxFile"] if volume == "technical" else result["docxFile"])
    )
    rows = [p for p in out.paragraphs if _pageref_bookmark(p)]
    expected = 20 if volume == "business" else (1 if volume == "technical" else 21)
    assert len(rows) == expected
    assert validate_native_toc(out)["entries"] == expected
    assert all("待更新" not in p.text and p.text[-1].isdigit() for p in rows)
    if volume == "technical":
        assert not any(p.text == "施工组织设计；" for p in out.paragraphs)
    assert path.read_bytes() == original
    if volume != "technical":
        assert any("社保" in p.text and "粘贴" in p.text for p in out.paragraphs)
        if with_media:
            assert len(list(out.element.iter(qn("w:drawing")))) == 1
    names = [n.get(qn("w:name")) for n in out.element.iter(qn("w:bookmarkStart"))]
    assert len(names) == len(set(names))


def test_final_validator_catches_the_six_rows_plus_old_directory_pattern(numbered_format):
    from api.services.tenders.document import _bookmark_paragraph, _write_toc_line
    from common.errors import AppError

    doc, _, directory, bodies = numbered_format
    # Reproduce the broken product: replace only six old rows and leave the rest.
    for n, el in enumerate(directory[:6]):
        target = bodies[n]
        name = f"toc_audit_{n}"
        from docx.text.paragraph import Paragraph

        _bookmark_paragraph(Paragraph(target, doc), name)
        p = Paragraph(el, doc)
        _write_toc_line(p, n, ITEMS[n][0], name, "1")
    with pytest.raises(AppError, match="目录结构校验失败"):
        validate_native_toc(doc)


@pytest.mark.parametrize(
    "volume,anonymous",
    [("business", False), ("technical", False), ("technical", True), ("full", False)],
)
def test_recognized_format_title_anchors_existing_design_without_supplement(
    numbered_format, tmp_path, monkeypatch, volume, anonymous
):
    from docx.text.paragraph import Paragraph

    doc, path, _, bodies = numbered_format
    # The actual invitation has different names in its directory and body.
    heading = Paragraph(bodies[-1], doc)
    heading.text = "施工组织设计"
    # A tab-stop definition controls alignment; it is not a content tab/TOC row.
    tabs = OxmlElement("w:tabs")
    stop = OxmlElement("w:tab")
    stop.set(qn("w:val"), "left")
    stop.set(qn("w:pos"), "4317")
    tabs.append(stop)
    heading._p.get_or_add_pPr().append(tabs)
    doc.save(path)
    original = path.read_bytes()
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: path)
    monkeypatch.setattr("api.services.tenders.generate.tenders_output_dir", lambda: tmp_path)
    brief = BidBrief(
        invitationId="audit-source",
        layoutMode="outline",
        outlineChapter="第八章 投标文件格式",
        generateVolume=volume,
        techAnonymous=anonymous,
        autoDraftTechPlan=True,
        outlineItems=[OutlineItem(
            id="recognized-tech", title="施工组织格式", rawTitle="十、施工组织格式",
            kind="tech_plan", draftBody="本项目已确认的施工组织设计响应正文。",
        )],
    )
    result = generate_bid(brief, catalog_slots=[], catalog_media={})
    out = Document(tmp_path / result["techDocxFile" if volume == "technical" else "docxFile"])
    expected = 20 if volume == "business" else (1 if volume == "technical" else 21)
    assert validate_native_toc(out)["entries"] == expected
    assert not any("原版格式章缺少「施工组织格式」" in s for s in result["warnings"])
    if volume != "business":
        design = next(p for p in out.paragraphs if p.text == "施工组织设计")
        row = next(p for p in out.paragraphs if p.text.startswith("十、施工组织格式\t"))
        assert any(n.get(qn("w:name")) == _pageref_bookmark(row)
                   for n in design._p.iter(qn("w:bookmarkStart")))
        assert not any(p.text == "施工组织格式" for p in out.paragraphs)
        assert sum("本项目已确认的施工组织设计响应正文。" in p.text
                   for p in out.paragraphs) == 1
    assert path.read_bytes() == original
