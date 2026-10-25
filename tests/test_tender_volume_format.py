"""分卷版式：标书写明的才改范本副本，技术标只留技术部分。"""

from __future__ import annotations

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from api.services.tenders.document import _apply_template_format
from api.services.tenders.format_rules import apply_volume_format, extract_document_format
from api.services.tenders.schema import VolumeFormat, default_brief


def _east(run) -> str:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return ""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        return ""
    return rfonts.get(qn("w:eastAsia")) or ""


def _set_font(run, name: str, size: float) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)


def test_unspecified_leaves_volume_flags_off() -> None:
    fmt = extract_document_format("第一章 投标邀请\n项目名称：某某充电站\n投标人须知")
    assert fmt.specified is False
    assert fmt.fontName == "宋体"
    assert fmt.marginLeftCm == 2.8
    assert fmt.businessFormat.marginSet is False
    assert fmt.technicalFormat.fontSet is False
    assert fmt.businessFormat.pagePosSet is False


def test_shared_margin_then_volume_rules_do_not_cross() -> None:
    text = """
页边距均为 2.5cm。
商务标格式
页边距均为 3cm。
技术标格式
正文用黑体小四。
"""
    fmt = extract_document_format(text)
    assert fmt.businessFormat.marginSet is True
    assert fmt.businessFormat.marginLeftCm == 3
    assert fmt.businessFormat.fontSet is False
    assert fmt.technicalFormat.marginSet is True
    assert fmt.technicalFormat.marginLeftCm == 2.5
    assert fmt.technicalFormat.fontName == "黑体"
    assert fmt.technicalFormat.bodySizeSet is True
    assert fmt.technicalFormat.bodySizePt == 12
    assert fmt.marginLeftCm == 2.5
    assert fmt.fontName == "黑体"


def test_one_line_volume_rules_stay_apart() -> None:
    text = "商务标页边距均为 2.8cm，技术标正文用宋体小四。"
    fmt = extract_document_format(text)
    assert fmt.businessFormat.marginSet is True
    assert fmt.businessFormat.marginTopCm == 2.8
    assert fmt.businessFormat.fontSet is False
    assert fmt.technicalFormat.marginSet is False
    assert fmt.technicalFormat.fontSet is True
    assert fmt.technicalFormat.fontName == "宋体"
    assert fmt.technicalFormat.bodySizePt == 12


def test_page_run_on_does_not_set_page_flags() -> None:
    fmt = extract_document_format("页码连续编码，不能跳号。")
    assert fmt.businessFormat.pagePosSet is False
    assert fmt.businessFormat.pageStartSet is False
    assert fmt.technicalFormat.pagePosSet is False


def test_page_pos_does_not_restart() -> None:
    doc = Document()
    doc.add_paragraph("正文一段说明文字放在这里。")
    apply_volume_format(
        doc,
        VolumeFormat(pagePosSet=True, pageNumberPos="bottom-right"),
    )
    footer = doc.sections[0].footer.paragraphs[0]
    assert footer.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert "PAGE" in footer._p.xml
    assert doc.sections[0]._sectPr.find(qn("w:pgNumType")) is None


def _slice(tmp_path, invite: str) -> Document:
    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("商务说明这段要留在商务标里。")
    quote = raw.add_table(rows=1, cols=2)
    quote.rows[0].cells[0].text = "序号"
    raw.add_paragraph("技术方案")
    body = raw.add_paragraph("这里是技术正文说明需要改字体。")
    _set_font(body.runs[0], "仿宋", 12)
    seal = raw.add_paragraph("（盖章）")
    _set_font(seal.runs[0], "楷体", 12)
    tech = raw.add_table(rows=1, cols=1)
    tech.rows[0].cells[0].text = "技术参数见附表"
    _set_font(tech.rows[0].cells[0].paragraphs[0].runs[0], "仿宋", 12)
    src = tmp_path / f"{invite}.format.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    return src


def test_technical_slice_keeps_tech_and_repaints_body(tmp_path, monkeypatch) -> None:
    from api.services.tenders.pdf_format import assemble_from_word_format

    invite = "abc123abc127"
    src = _slice(tmp_path, invite)
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    brief.documentFormat.technicalFormat = VolumeFormat(
        fontSet=True,
        fontName="黑体",
        bodySizeSet=True,
        bodySizePt=16,
        marginSet=True,
        marginLeftCm=2.2,
        marginRightCm=2.2,
        marginTopCm=2.2,
        marginBottomCm=2.2,
    )
    dest = tmp_path / "tech.docx"
    _path, notes = assemble_from_word_format(brief, dest, volume="technical")
    doc = Document(str(dest))
    texts = [p.text for p in doc.paragraphs]
    assert "投标函" not in texts
    assert any("技术方案" in t for t in texts)
    assert any("技术标只保留格式章里的技术部分" in n for n in notes)
    body = next(p for p in doc.paragraphs if "技术正文" in (p.text or ""))
    assert _east(body.runs[0]) == "黑体"
    assert abs(body.runs[0].font.size.pt - 16) < 0.1
    seal = next(p for p in doc.paragraphs if "盖章" in (p.text or ""))
    assert _east(seal.runs[0]) == "楷体"
    cell = doc.tables[0].rows[0].cells[0]
    assert "技术参数" in cell.text
    assert _east(cell.paragraphs[0].runs[0]) == "仿宋"
    assert abs(float(doc.sections[0].left_margin.cm) - 2.2) < 0.05
    kept = Document(str(src))
    assert any("投标函" in (p.text or "") for p in kept.paragraphs)
    assert any("范本正文仿宋" in n and "黑体" in n for n in notes)


def test_business_and_omitted_volume_keep_letter(tmp_path, monkeypatch) -> None:
    from api.services.tenders.pdf_format import assemble_from_word_format

    invite = "abc123abc128"
    src = _slice(tmp_path, invite)
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    brief.documentFormat.businessFormat = VolumeFormat(
        marginSet=True,
        marginLeftCm=1.9,
        marginRightCm=1.9,
        marginTopCm=1.9,
        marginBottomCm=1.9,
    )
    dest = tmp_path / "biz.docx"
    assemble_from_word_format(brief, dest, volume="business")
    doc = Document(str(dest))
    assert any("投标函" in (p.text or "") for p in doc.paragraphs)
    assert abs(float(doc.sections[0].left_margin.cm) - 1.9) < 0.05
    body = next(p for p in doc.paragraphs if "技术正文" in (p.text or ""))
    assert _east(body.runs[0]) == "仿宋"

    dest2 = tmp_path / "omit.docx"
    assemble_from_word_format(brief, dest2)
    omitted = Document(str(dest2))
    assert any("投标函" in (p.text or "") for p in omitted.paragraphs)


def test_technical_slice_rebuilds_cover_toc_and_page(tmp_path, monkeypatch) -> None:
    from api.services.tenders.pdf_format import assemble_from_word_format

    invite = "abc123abc130"
    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("商务说明这段要留在商务标里。")
    raw.add_paragraph("施工组织设计")
    raw.add_paragraph("一、工程概况")
    raw.add_paragraph("概况只写现场条件和范围。")
    raw.add_paragraph("二、主要施工方案")
    for i in range(48):
        raw.add_paragraph(f"第{i}条施工步骤说明文字放在这里用来把后面的小节推到下一页。")
    raw.add_paragraph("（一）质量保证措施")
    raw.add_paragraph("质量措施正文。")
    src = tmp_path / f"{invite}.format.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    brief.projectName = "郑州充电站"
    brief.bidderName = "河南伟泰光电科技有限公司"
    dest = tmp_path / "tech.docx"
    _path, notes = assemble_from_word_format(brief, dest, volume="technical")
    doc = Document(str(dest))
    texts = [p.text or "" for p in doc.paragraphs]
    blob = "\n".join(texts)
    assert "投标函" not in blob
    assert "商务说明" not in blob
    assert "施工组织设计" in blob
    cover = next(t for t in texts if "技术标" in t)
    assert cover.strip() == "技术标"
    assert "河南伟泰光电科技有限公司" in blob
    toc_at = next(i for i, t in enumerate(texts) if t.strip() == "目录")
    assert any("技术标" in t for t in texts[:toc_at])
    assert "TOC" in doc.element.xml
    assert "PAGEREF" in doc.element.xml
    assert any("已另写封面和目录" in n for n in notes)
    body_toc = next(p for p in doc.paragraphs if (p.text or "").startswith("一、工程概况"))
    assert re.search(r"(\d+)$", body_toc.text or "").group(1) == "1"
    later = next(p for p in doc.paragraphs if (p.text or "").startswith("（一）质量保证措施"))
    assert int(re.search(r"(\d+)$", later.text or "").group(1)) > 1
    sub = next(p for p in doc.paragraphs if "质量保证措施" in (p.text or "") and "PAGEREF" in p._p.xml)
    ind = sub._p.find(qn("w:pPr")).find(qn("w:ind"))
    assert int(ind.get(qn("w:left")) or "0") > 0
    leader = sub._p.find(qn("w:pPr")).find(qn("w:tabs")).find(qn("w:tab"))
    assert leader.get(qn("w:leader")) == "dot"
    assert "PAGE" not in doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in doc.sections[-1].footer.paragraphs[0]._p.xml
    start = doc.sections[-1]._sectPr.find(qn("w:pgNumType"))
    assert start is not None and start.get(qn("w:start")) == "1"
    kept = Document(str(src))
    assert any("投标函" in (p.text or "") for p in kept.paragraphs)


def test_anonymous_tech_cover_omits_bidder(tmp_path, monkeypatch) -> None:
    from api.services.tenders.pdf_format import assemble_from_word_format

    invite = "abc123abc131"
    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("施工组织设计")
    raw.add_paragraph("一、工程概况")
    raw.add_paragraph("概况正文。")
    src = tmp_path / f"{invite}.format.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    brief.projectName = "郑州充电站"
    brief.bidderName = "河南伟泰光电科技有限公司"
    brief.legalPersonName = "郭志伟"
    brief.techAnonymous = True
    dest = tmp_path / "anon.docx"
    assemble_from_word_format(brief, dest, volume="technical")
    doc = Document(str(dest))
    texts = [p.text or "" for p in doc.paragraphs]
    toc_at = next(i for i, t in enumerate(texts) if t.strip() == "目录")
    front = "\n".join(texts[: toc_at + 4])
    assert "技术标" in front
    assert "河南伟泰" not in front
    assert "郭志伟" not in front
    assert "郑州充电站" not in front
    assert "投标人" not in front
    assert "河南伟泰" not in doc.sections[0].header.paragraphs[0].text


def test_anon_clause_hides_tech_chain_and_open_bid_signs(tmp_path, monkeypatch) -> None:
    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.classifier import apply_volume_binding

    invite = "abc123abc132"
    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("施工组织设计")
    raw.add_paragraph("一、工程概况")
    raw.add_paragraph("编制单位：河南伟泰光电科技有限公司，项目经理郭志伟。")
    footer = raw.sections[0].footer.paragraphs[0]
    footer.text = "河南伟泰光电科技有限公司"
    src = tmp_path / f"{invite}.format.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )
    base = default_brief()
    base.layoutMode = "outline"
    base.invitationId = invite
    base.bidderName = "河南伟泰光电科技有限公司"
    base.legalPersonName = "郭志伟"
    base.pmName = "李项目"
    clause = "施工组织设计按技术暗标评审，不得出现投标人名称、徽标和人员名称。"
    anon = apply_volume_binding(base, clause, keep_volume=True)
    assert anon.techAnonymous is True
    tech_path, _notes = assemble_bid_docx(anon, tmp_path / "tech.docx", volume="technical")
    tech = Document(str(tech_path))
    texts = [p.text or "" for p in tech.paragraphs]
    blob = "\n".join(texts)
    toc_at = next(i for i, t in enumerate(texts) if t.strip() == "目录")
    front = "\n".join(texts[:toc_at])
    assert "技术标" in front
    assert "投标人" not in front
    assert "河南伟泰" not in front
    assert "河南伟泰" not in blob
    assert "郭志伟" not in blob
    assert "李项目" not in blob
    foot = "\n".join(p.text or "" for sec in tech.sections for p in sec.footer.paragraphs)
    assert "河南伟泰" not in foot
    assert "PAGE" in tech.sections[-1].footer.paragraphs[0]._p.xml
    assert (tech.core_properties.author or "") == ""
    assert (tech.core_properties.last_modified_by or "") == ""

    biz_path, _biz_notes = assemble_bid_docx(anon, tmp_path / "biz.docx", volume="business")
    biz_blob = "\n".join(p.text or "" for p in Document(str(biz_path)).paragraphs)
    assert "投标函" in biz_blob
    assert "河南伟泰光电科技有限公司" in biz_blob

    plain = apply_volume_binding(base, "施工组织设计按图纸编制，正常署名。", keep_volume=True)
    assert plain.techAnonymous is False
    open_path, _open_notes = assemble_bid_docx(plain, tmp_path / "open.docx", volume="technical")
    open_doc = Document(str(open_path))
    open_blob = "\n".join(p.text or "" for p in open_doc.paragraphs)
    assert "河南伟泰光电科技有限公司" in open_blob
    assert any("投标人" in (p.text or "") and "河南伟泰" in (p.text or "") for p in open_doc.paragraphs)


def test_technical_without_heading_keeps_whole_slice(tmp_path, monkeypatch) -> None:
    from api.services.tenders.pdf_format import assemble_from_word_format

    invite = "abc123abc129"
    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("整份都是商务函件。")
    src = tmp_path / f"{invite}.format.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    dest = tmp_path / "all.docx"
    _path, notes = assemble_from_word_format(brief, dest, volume="technical")
    doc = Document(str(dest))
    assert any("投标函" in (p.text or "") for p in doc.paragraphs)
    assert any("没有单独的技术部分" in n for n in notes)


def test_company_template_ignores_cover_only_specified() -> None:
    doc = Document()
    doc.add_paragraph("公司模板正文一段。")
    doc.sections[0].left_margin = Cm(3.1)
    brief = default_brief()
    brief.documentFormat.specified = True
    brief.documentFormat.notes = ["封面含加盖公章"]
    brief.documentFormat.coverNeedSeal = True
    brief.documentFormat.marginLeftCm = 2.0
    notes = _apply_template_format(doc, brief, volume="business")
    assert abs(float(doc.sections[0].left_margin.cm) - 3.1) < 0.05
    assert not any("页边距" in n and "公司模板" in n for n in notes)

    brief.documentFormat.businessFormat = VolumeFormat(
        marginSet=True,
        marginLeftCm=2.0,
        marginRightCm=2.0,
        marginTopCm=2.0,
        marginBottomCm=2.0,
    )
    _apply_template_format(doc, brief, volume="business")
    assert abs(float(doc.sections[0].left_margin.cm) - 2.0) < 0.05
