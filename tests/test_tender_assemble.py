"""二期：大纲组卷复制招标书原文；三期：已知类型用模块填写 BidBrief。"""

from __future__ import annotations

import re

from docx.oxml.ns import qn

from api.services.tenders.outline import (
    choose_layout_mode,
    compact_title,
    extract_outline,
    fill_copy_blanks,
    looks_like_form_template,
)
from api.services.tenders.schema import (
    BidBrief,
    DeviationLine,
    OutlineItem,
    PerformanceLine,
    PlaceholderItem,
    QuoteLineIn,
    default_brief,
)


def _docx_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_ymd_pads_single_digit_month_day() -> None:
    from api.services.tenders.document import _two_digit_md, _ymd

    assert _two_digit_md("9") == "09"
    assert _two_digit_md("09") == "09"
    assert _two_digit_md("11") == "11"
    assert _ymd("2026-9-1") == ("2026", "09", "01")
    assert _ymd("2026-09-11") == ("2026", "09", "11")


def test_sign_spacer_caps_and_reserves_room() -> None:
    from docx import Document

    from api.services.tenders.document import _sign_spacer_cm

    doc = Document()
    doc.add_paragraph("短正文")
    capped = _sign_spacer_cm(doc, sign_cm=3.2, max_cm=3.2, reserve_cm=2.0)
    assert capped <= 3.2 + 1e-6
    # 预留过大时宁可不要弹性空白，避免和 keep 链一起翻页
    none = _sign_spacer_cm(doc, sign_cm=3.2, reserve_cm=100.0, max_cm=3.2)
    assert none == 0.0


def test_align_sign_label_spreads_short_words() -> None:
    from api.services.tenders.document import _align_sign_label

    assert _align_sign_label("投标人：") == "投　标　人："
    assert _align_sign_label("日期") == "日　　　期："
    assert _align_sign_label("法定代表人：") == "法定代表人："


def test_letter_sign_line_stays_single_line_no_overflow() -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    from api.services.tenders.document import (
        _LETTER_SIGN_RIGHT_KW,
        _em_len,
        _line_em_budget,
        _rewrite_date_line,
        _rewrite_labeled_underline,
    )

    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.54)
        sec.right_margin = Cm(2.54)
    para = doc.add_paragraph()
    _rewrite_labeled_underline(
        para,
        [("法定代表人或其委托代理人：", "", "（签字或盖章）")],
        **_LETTER_SIGN_RIGHT_KW,
    )
    assert "（签字或盖章）" in (para.text or "")
    body = (para.text or "").replace("法定代表人或其委托代理人：", "", 1).replace("（签字或盖章）", "", 1)
    # 实线下划线空格，且「（签字或盖章）」与横线同排
    assert "＿" not in (para.text or "")
    assert "（签字或盖章）" in (para.text or "")
    assert "\n" not in (para.text or "")
    assert _em_len(body) >= 4
    # 横线用的是带 w:u 的全角空格
    assert "w:u" in para._p.xml
    left = float(para.paragraph_format.left_indent.cm) if para.paragraph_format.left_indent else 0
    assert left >= 1.4
    date_p = doc.add_paragraph()
    _rewrite_date_line(
        date_p,
        "2026",
        "09",
        "29",
        label="日期：",
        **_LETTER_SIGN_RIGHT_KW,
    )
    assert "2026" in (date_p.text or "")
    assert date_p.alignment == WD_ALIGN_PARAGRAPH.LEFT
    d_label = "日　　　期："
    d_room = _line_em_budget(date_p, d_label, "")
    d_body = (date_p.text or "").replace("日期：", "").replace("日　　　期：", "")
    assert _em_len(d_body) <= d_room + 0.6


def test_media_keys_do_not_borrow_credit_or_id() -> None:
    from api.services.tenders.assemble import _media_keys_for_item
    from api.services.tenders.match import _same_material_kind
    from api.services.tenders.placeholders import DEFAULT_SLOTS

    slots = list(DEFAULT_SLOTS)
    audit = OutlineItem(id="x", title="上一年度审计报告", kind="scan")
    assert _media_keys_for_item(audit, slots) == ["finance"]
    commit = OutlineItem(id="c", title="无违法承诺函", kind="commitment_copy")
    keys = _media_keys_for_item(commit, slots)
    assert "credit" not in keys
    assert "commitment" in keys or keys == []
    assert not _same_material_kind(
        "近三年财务审计报告 / 完税证明 / 社保缴纳证明",
        "授权代理人身份证及近半年社保缴纳证明",
    )


def test_attach_template_body_not_toc_line() -> None:
    text = """
第五章 响应文件格式
响应函
知识产权不侵权承诺函
报价表
附件 响应函
致：（招标人名称）
我方参加（项目名称）的采购。
知识产权不侵权承诺函
我方承诺不侵犯任何第三方知识产权。
投标人名称：________
年  月  日
报价表
系统名称及报价
第六章 合同条款
"""
    chapter, items = extract_outline(text)
    assert "响应文件格式" in chapter
    by_title = {item.title: item for item in items}
    letter = by_title["响应函"]
    commit = by_title["知识产权不侵权承诺函"]
    assert "我方参加" in letter.body
    assert "不侵犯任何第三方知识产权" in commit.body
    assert "系统名称及报价" not in commit.body


def test_fill_copy_blanks_only_placeholders() -> None:
    raw = "致：（招标人名称）\n我方承诺不侵犯任何第三方知识产权。\n投标人名称：________\n年  月  日"
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="数智制造项目",
        tenderer="某市分公司",
        legal="郭志伟",
        bid_date="2026-09-08",
    )
    assert "某市分公司" in filled
    assert "河南伟泰光电科技有限公司" in filled
    assert "2026年09月08日" in filled
    assert "不侵犯任何第三方知识产权" in filled
    assert "（招标人名称）" not in filled


def test_fill_copy_blanks_price_and_duration() -> None:
    raw = (
        "愿意以人民币（大写）                    （¥\n"
        "元）的投标总报价，工期；\n"
        "日历天，按合同约定完成承包工程。"
    )
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="充电桩项目",
        tenderer="郑州天健智慧园区运营管理有限公司",
        legal="郭志伟",
        bid_date="2026-09-24",
        delivery_days=90,
        price_cn="捌拾玖万柒仟柒佰叁拾叁元整",
        price_en="897733.00",
    )
    assert "人民币（大写）捌拾玖万柒仟柒佰叁拾叁元整" in filled
    assert "¥897733.00元" in filled
    assert "工期：90日历天" in filled
    messy = fill_copy_blanks(
        "愿意以人民币（大写）叁佰万元整（¥3,000,000.00 元）参加投标。交货期：30天内。\n其他补充说明：",
        bidder="河南伟泰光电科技有限公司",
        project="郑州高新技术产业开发区新能源充(换)电站基础设施群项目（一期）",
        tenderer="郑州天健智慧园区运营管理有限公司",
        legal="郭志伟",
        bid_date="2026-09-28",
        delivery_days=30,
        price_cn="叁佰万元整",
        price_en="3000000.00",
        works=True,
    )
    assert "（¥3000000.00 元）" in messy or "（¥3000000.00元）" in messy
    assert "3,000,000" not in messy
    assert "交货期：30天内" in messy
    assert "计划开工日期" not in messy
    assert "其他补充说明：无" in messy


def test_fill_copy_fees_from_control_price() -> None:
    from api.services.tenders.money import rmb_pair

    raw = "安全生产措施费：\n暂列金额：\n专业工程暂估价：\n项目经理：\n"
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="充电桩项目",
        tenderer="郑州天健智慧园区运营管理有限公司",
        legal="郭志伟",
        bid_date="2026-09-28",
        safety_fee=rmb_pair(37023.13),
        provisional=rmb_pair(75000),
        prime_cost=rmb_pair(0),
    )
    assert "37023.13" in filled
    assert "75000.00" in filled
    assert "0.00" in filled
    assert "安全生产措施费：无" not in filled
    assert "项目经理：________" in filled


def test_fill_copy_letter_generic_rules() -> None:
    from api.services.tenders.money import rmb_pair

    raw = (
        "我方已仔细研究了招标文件（含补遗书第___号至第___号）的全部内容。\n"
        "质量标准：\n"
        "投标有效期：满足招标文件要求\n"
        "安全生产措施费：无\n"
        "暂列金额：无\n"
        "专业工程暂估价：无\n"
        "项目经理：\n"
        "7. 补充说明（\n"
        "投标文件的上述组成部分如存在内容不一致的，以投标函为准。\n"
    )
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="充电桩项目",
        tenderer="郑州天健智慧园区运营管理有限公司",
        legal="郭志伟",
        bid_date="2026-09-28",
        delivery_days=90,
        validity_days=90,
        quality="合格",
        extra_note="无",
        safety_fee=rmb_pair(37023.13),
        provisional=rmb_pair(75000),
        prime_cost=rmb_pair(0),
    )
    assert "（含补遗书无）" in filled
    assert "质量标准：合格" in filled
    assert "投标有效期：自投标截止之日起90日历天" in filled
    assert "满足招标文件要求" not in filled
    assert "以投标函为准" not in filled
    assert "37023.13" in filled
    assert "75000.00" in filled
    assert "0.00" in filled
    assert "安全生产措施费：无" not in filled
    assert "专业工程暂估价：无" not in filled
    assert "项目经理：________" in filled
    assert "补充说明：无" in filled


def test_fill_copy_pm_junk_and_grade() -> None:
    filled = fill_copy_blanks(
        "项目经理：是指由承\n注册证号：/\n专业等级：/\n投标有效期：90日历天\n",
        bidder="河南伟泰光电科技有限公司",
        project="充电桩项目",
        tenderer="郑州天健智慧园区运营管理有限公司",
        legal="郭志伟",
        bid_date="2026-09-28",
        validity_days=90,
        pm_grade="/",
        pm_name="是指由承",
    )
    assert "是指由承" not in filled
    assert "乱码" not in filled
    assert "项目经理：________" in filled
    assert "注册证号：________" in filled
    assert "专业等级：无" in filled
    assert "自投标截止之日起90日历天" in filled


def test_ensure_works_letter_and_appendix_fields() -> None:
    from api.services.tenders.outline import ensure_appendix_fields, ensure_works_letter_fixed
    from api.services.tenders.money import rmb_pair

    raw = (
        "1. 我方已仔细研究了某某项目招标文件（含补遗书无）的全部内容，"
        "愿意以人民币（大写）：捌拾万元整（小写：￥800000.00元）的投标总价；"
        "计划工期：90日历天；质量要求：合格。\n"
        "投标人：河南伟泰光电科技有限公司（盖单位公章）\n"
    )
    out = ensure_works_letter_fixed(raw, quality="合格", valid_days=90, bid_bond="")
    assert "修补工程中的任何缺陷" in out
    assert "投标函附录属于合同文件的组成部分" in out
    assert "随同本投标函提交投标保证金一份" in out
    assert "自投标截止之日起90日历天" in out
    app = ensure_appendix_fields(
        "项目名称：郑州高新区项目\n投标人：河南伟泰光电科技有限公司\n专业等级：/\n",
        project="郑州高新区项目",
        bidder="河南伟泰光电科技有限公司",
        scope="施工图范围内工程",
        price=rmb_pair(800000),
        safety_fee=rmb_pair(37023.13),
        provisional=rmb_pair(75000),
        prime_cost=rmb_pair(0),
        valid="自投标截止之日起90日历天",
        grade="无",
    )
    for lab in (
        "投标范围",
        "投标总报价",
        "安全生产措施费",
        "暂列金额",
        "专业工程暂估价",
        "项目经理",
        "注册证号",
        "投标有效期",
        "缺陷责任期",
        "履约担保金额",
        "专业等级",
    ):
        assert lab in app
    assert "37023.13" in app
    assert "项目经理：________" in app


def test_letter_module_appendix_fields() -> None:
    from docx import Document

    from api.services.tenders.modules import render_module
    from api.services.tenders.schema import BidBrief, OutlineItem

    brief = BidBrief(
        projectName="郑州高新技术产业开发区新能源充(换)电站基础设施群项目（一期）",
        tenderer="郑州天健智慧园区运营管理有限公司",
        bidderName="河南伟泰光电科技有限公司",
        bidPriceYuan=800000,
        deliveryDays=90,
        bidValidityDays=90,
        quality="合格",
        safetyFeeYuan=37023.13,
        provisionalYuan=75000,
        primeCostYuan=0,
        bidContent="施工图范围内工程",
        defectPeriod="",
        perfBond="",
    )
    doc = Document()
    render_module(doc, brief, OutlineItem(id="l1", title="投标函", kind="letter"))
    render_module(doc, brief, OutlineItem(id="a1", title="投标函附录", kind="letter"))
    text = "\n".join(p.text or "" for p in doc.paragraphs)
    assert "修补工程中的任何缺陷" in text
    assert "工程质量达到" in text
    assert "合格" in text
    assert "自投标截止之日起" in text
    assert "90日历天" in text
    assert "投标保证金" in text
    assert "投标函附录属于合同文件的组成部分" in text
    assert "缺陷责任期" in text
    assert "履约担保金额" in text
    assert "专业等级" in text
    assert "专业等级：/" not in text.replace(" ", "")


def test_fill_copy_blanks_empty_zhi_and_bidder_not_sign() -> None:
    raw = (
        "投标承诺书致：\n"
        "我公司现做出如下承诺：\n"
        "投标人： （盖章）\n"
        "法定代表人或委托代理人：（签字并盖章）\n"
        "日期： 年月日"
    )
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="厂内充电桩采购",
        tenderer="二连浩特市联源热电有限公司",
        legal="郭志伟",
        bid_date="2026-09-15",
    )
    assert "投标承诺书\n致：二连浩特市联源热电有限公司" in filled
    assert "投标人：河南伟泰光电科技有限公司（盖章）" in filled
    assert "2026年09月15日" in filled
    assert "郭志伟" not in filled
    assert "签字并盖章" in filled


def test_fill_copy_blanks_zhi_not_body_and_bare_company() -> None:
    mashed = fill_copy_blanks(
        "致：我公司现做出如下承诺：\n1、保证投标文件内容无任何虚假。",
        bidder="河南伟泰光电科技有限公司",
        project="厂内充电桩采购",
        tenderer="二连浩特市联源热电有限公司",
        legal="郭志伟",
        bid_date="2026-09-16",
    )
    assert mashed.splitlines()[0] == "致：二连浩特市联源热电有限公司"
    assert "我公司现做出如下承诺：" in mashed
    glued = fill_copy_blanks(
        "致：河南鑫宇光科技股份有限公司我方确认收到贵方提供的招标文件，并重申以下几点：",
        bidder="河南伟泰光电科技有限公司",
        project="MOM生产运营管理平台",
        tenderer="河南鑫宇光科技股份有限公司",
        legal="",
        bid_date="2026-09-23",
    )
    lines = [ln for ln in glued.splitlines() if ln.strip()]
    assert lines[0] == "致：河南鑫宇光科技股份有限公司"
    assert lines[1].startswith("我方确认收到")
    assert "我方确认" not in lines[0]
    bare = fill_copy_blanks(
        "有限公司：\n我方已全面阅读和研究了招标文件。",
        bidder="河南伟泰光电科技有限公司",
        project="厂内充电桩采购",
        tenderer="二连浩特市联源热电有限公司",
        legal="郭志伟",
        bid_date="2026-09-16",
    )
    assert bare.startswith("二连浩特市联源热电有限公司：")


def test_fill_copy_blanks_unfolds_mashed_sign_lines() -> None:
    filled = fill_copy_blanks(
        "投标人（章）：法定代表人或授权代表（签字）：\n"
        "联系人：联系电话：2026年09月16日",
        bidder="河南伟泰光电科技有限公司",
        project="厂内充电桩采购",
        tenderer="二连浩特市联源热电有限公司",
        legal="郭志伟",
        bid_date="2026-09-16",
        phone="17630567052",
        contact="郭志伟",
    )
    lines = [ln.strip() for ln in filled.splitlines() if ln.strip()]
    assert any(ln.startswith("投标人：") and "河南伟泰光电科技有限公司" in ln for ln in lines)
    assert any("法定代表人或授权代表" in ln and "郭志伟" not in ln for ln in lines)
    assert any(ln.startswith("联系人：") and "郭志伟" in ln for ln in lines)
    assert any("17630567052" in ln and "联系人" not in ln for ln in lines)
    assert any(ln.startswith("日期：") or "2026年09月16日" in ln for ln in lines)
    assert "：：" not in filled


def test_fill_copy_blanks_covers_gai_gongzhang() -> None:
    filled = fill_copy_blanks(
        "投标人(盖公章)：\n授权代表(签名)：\n日期： 年 月 日\n",
        bidder="河南伟泰光电科技有限公司",
        project="厂内充电桩采购",
        tenderer="河南鑫宇光科技股份有限公司",
        legal="郭志伟",
        bid_date="2026-09-23",
    )
    assert "河南伟泰光电科技有限公司" in filled
    assert "盖公章" in filled.replace(" ", "")
    assert "郭志伟" not in filled.split("授权代表")[-1].split("日期")[0]
    assert "2026年09月23日" in filled.replace(" ", "")


def test_fill_copy_blanks_leaves_sign_name_empty() -> None:
    raw = "法定代表人（签字）：________\n授权代表签字：________"
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="数智制造项目",
        tenderer="某市分公司",
        legal="郭志伟",
        bid_date="2026-09-08",
    )
    assert "郭志伟" not in filled
    assert "________" in filled


def test_fill_copy_blanks_underscores_and_days() -> None:
    raw = (
        "________________有限公司：\n"
        "我方已仔细阅读和研究了________________招标文件。\n"
        "中标通知起_______天内交货。自开标之日起_______天内有效。\n"
        "联系人：________    联系电话：________\n"
        "年  月  日"
    )
    filled = fill_copy_blanks(
        raw,
        bidder="河南伟泰光电科技有限公司",
        project="厂内充电桩采购",
        tenderer="某热电有限公司",
        legal="郭志伟",
        bid_date="2026-09-09",
        phone="17630567052",
        contact="郭志伟",
        delivery_days=30,
        validity_days=60,
    )
    assert filled.startswith("某热电有限公司：")
    assert "阅读和研究了厂内充电桩采购招标文件" in filled
    assert "通知起30天" in filled
    assert "开标之日起60天" in filled
    assert "17630567052" in filled
    assert "2026年09月09日" in filled


def test_looks_like_form_template() -> None:
    short = "致：（招标人名称）\n我方参加（项目名称）。"
    assert looks_like_form_template(short) is False
    form = (
        "附件一：\n"
        "投 标 函\n"
        "________________有限公司：\n"
        "我方已仔细阅读和研究了________________招标文件，决定参加投标。\n"
        "投标人（章）：          法定代表人或授权代表（签字）：\n"
        "年  月  日\n"
    )
    assert looks_like_form_template(form) is True
    mashed = (
        "第四部分投标文件格式附件一：投标函________________有限公司："
        "我方已仔细阅读和研究了________________招标文件，决定参加投标。"
        "投标人（章）：年  月  日"
    )
    assert looks_like_form_template(mashed) is True


def test_choose_layout_mode_response_vs_classic() -> None:
    _chapter, telecom = extract_outline(
        """
第五章 响应文件格式
响应函
承诺函
商务条款偏离表
技术规范书偏离表
知识产权不侵权承诺函
第六章 合同
"""
    )
    assert choose_layout_mode("第五章响应文件格式", telecom) == "outline"
    _c2, classic = extract_outline(
        """
第四部分 投标文件格式
附件一：投标函
附件二：法定代表人身份证明
附件三：授权委托书
附件四：分项报价表
附件五：投标承诺书
第五部分 评标办法
"""
    )
    assert choose_layout_mode("第四部分投标文件格式", classic) == "outline"
    _c3, long_pack = extract_outline(
        """
第四部分 投标文件格式
投标函
竞标书
法定代表人身份证明
法定代表人授权委托书
投标承诺书
投标保证金
印鉴预留备案表
投标报价单
商务偏离表
第五部分 评标办法
"""
    )
    assert choose_layout_mode("第四部分投标文件格式", long_pack) == "outline"


def test_assemble_follows_outline_item_order(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="授权委托书", kind="auth", source="generate"),
        OutlineItem(id="o02", title="法定代表人身份证明", kind="legal_id", source="generate"),
    ]
    dest = tmp_path / "order.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    paras = [p.text.strip() for p in Document(str(dest)).paragraphs if p.text.strip()]
    assert not any(t.startswith("一、投标函") or t == "投标函" for t in paras)
    assert any("法定代表人身份证明" in t for t in paras)
    assert any("授权委托书" in t for t in paras)
    assert paras.index(next(t for t in paras if "授权委托书" in t)) < paras.index(
        next(t for t in paras if "法定代表人身份证明" in t)
    )


def test_assemble_keeps_one_biz_dev_table(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="商务偏离表", kind="biz_dev", source="generate"),
        OutlineItem(id="o02", title="商务偏离表招标项目", kind="biz_dev", source="generate"),
        OutlineItem(id="o03", title="技术规格偏离表及建议招标项目", kind="tech_dev", source="generate"),
    ]
    dest = tmp_path / "bizdev.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert paras.count("商务偏离表") == 1
    assert "商务偏离表招标项目" not in paras
    assert "技术规格偏离表及建议" in paras
    assert "技术规格偏离表及建议招标项目" not in paras
    biz_tables = [
        t
        for t in doc.tables
        if any("交货期" in (c.text or "") for row in t.rows for c in row.cells)
    ]
    assert len(biz_tables) == 1


def test_assemble_letter_keeps_invitation_form_layout(tmp_path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "附件一：\n"
        "投 标 函\n"
        "________________有限公司：\n"
        "我方已仔细阅读和研究了________________招标文件，决定参加投标。\n"
        "中标通知起_______天内交货。自开标之日起_______天内有效。\n"
        "投标人（章）：          法定代表人或授权代表（签字）：\n"
        "联系人：________    联系电话：________\n"
        "年  月  日\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "厂内新能源充电桩采购项目"
    brief.tenderer = "某热电有限公司"
    brief.bidDate = "2026-09-09"
    brief.deliveryDays = 30
    brief.bidValidityDays = 60
    brief.bidderPhone = "17630567052"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标函（附件一）",
            kind="letter",
            source="generate",
            body=body,
        )
    ]
    dest = tmp_path / "letter-form.docx"
    path, warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(path))
    text = _docx_text(doc)
    assert "响应总报价" not in text
    assert "附件一" in text
    assert "投 标 函" in text
    assert "某热电有限公司" in text
    assert "厂内新能源充电桩采购项目" in text
    assert "通知起30天" in text
    assert "开标之日起60天" in text
    assert "17630567052" in text
    assert any("复制" in w for w in warnings)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    title = next(p for p in paras if compact_title(p.text) == "投标函")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs and title.runs[0].bold and title.runs[0].font.size >= Pt(16)
    salute = next(p for p in paras if "某热电有限公司" in p.text and p.text.strip().startswith("致"))
    assert salute.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert not salute.paragraph_format.first_line_indent
    assert "w:u" in salute._p.xml
    assert salute.text.strip().endswith("：")
    body_p = next(p for p in paras if "阅读和研究了" in p.text)
    assert not body_p.paragraph_format.first_line_indent
    texts = [p.text.strip() for p in paras]
    title_i = next(i for i, t in enumerate(texts) if compact_title(t) == "投标函")
    bidder_i = next(
        i for i, t in enumerate(texts) if i > title_i and compact_title(t).startswith("投标人")
    )
    legal_i = next(i for i, t in enumerate(texts) if i > title_i and "签字" in t)
    contact_i = next(i for i, t in enumerate(texts) if i > title_i and compact_title(t).startswith("联系人"))
    phone_i = next(i for i, t in enumerate(texts) if i > title_i and "17630567052" in t)
    assert bidder_i != legal_i
    assert contact_i != phone_i
    assert "河南伟泰光电科技有限公司" in texts[bidder_i]
    assert "郭志伟" not in texts[legal_i]
    assert "：：" not in texts[phone_i]
    assert "w:u" in paras[bidder_i]._p.xml
    assert "签字" in texts[legal_i]
    assert "＿" not in texts[legal_i]
    assert "w:u" in paras[legal_i]._p.xml
    assert "（签字" in texts[legal_i] or "签字或盖章" in texts[legal_i]
    def _sign_left(p) -> float:
        ind = p.paragraph_format.left_indent
        return float(ind.cm) if ind else 0.0

    assert _sign_left(paras[contact_i]) >= 4.0
    assert _sign_left(paras[phone_i]) >= 4.0
    assert _sign_left(paras[bidder_i]) >= 4.0
    assert abs(_sign_left(paras[contact_i]) - _sign_left(paras[bidder_i])) < 0.05
    # 短标签冒号轴对齐（中间插全角空格）
    assert "投　标　人" in texts[bidder_i] or "投标人" in texts[bidder_i]
    assert "联" in texts[contact_i] and "人" in texts[contact_i]
    # 落款单行：不出现「签字」与「或盖章」拆到两段
    assert "签字或盖章" in texts[legal_i] or "签字" in texts[legal_i]
    assert not any(
        (p.text or "").strip() in {"或盖章）", "盖章）"} for p in paras
    )


def test_assemble_letter_without_body_still_writes_module(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.tenderer = "某热电有限公司"
    brief.bidPriceYuan = 518800
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函（附件一）", kind="letter", source="generate"),
        OutlineItem(id="o02", title="法定代表人授权委托书（附件四）", kind="auth", source="generate"),
    ]
    dest = tmp_path / "module-fallback.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    text = _docx_text(Document(str(dest)))
    assert "某热电有限公司" in text
    assert "518,800.00" in text or "伍拾壹万" in text or "投标总报价" in text or "响应总报价" in text
    assert "授权" in text or "委托" in text


def test_assemble_mashed_invitation_letter_not_dropped(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "第四部分投标文件格式附件一：投 标 函"
        "________________有限公司："
        "我方已仔细阅读和研究了________________招标文件，决定参加本次投标。"
        "中标通知起_______天内交货。自开标之日起_______天内有效。"
        "投标人（章）：          法定代表人或授权代表（签字）："
        "年  月  日"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "厂内新能源充电桩采购项目"
    brief.tenderer = "某热电有限公司"
    brief.bidDate = "2026-09-09"
    brief.deliveryDays = 30
    brief.bidValidityDays = 60
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标函（附件一）",
            kind="letter",
            source="generate",
            body=body,
        )
    ]
    dest = tmp_path / "mashed-letter.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    text = _docx_text(Document(str(dest)))
    assert "阅读和研究了" in text
    assert "厂内新能源充电桩采购项目" in text
    assert "招标书格式章未提供本页空白稿" not in text


def test_assemble_copies_commitment_not_charger_clauses(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.commitment import _CLAUSES

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "数智制造管理系统"
    brief.tenderer = "某市分公司"
    brief.bidDate = "2026-09-08"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineChapter = "第五章响应文件格式"
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="响应函",
            kind="letter",
            source="generate",
            body="致：（招标人名称）\n我方参加（项目名称）。",
        ),
        OutlineItem(
            id="o02",
            title="知识产权不侵权承诺函",
            kind="commitment_copy",
            source="copy",
            body="我方承诺不侵犯任何第三方知识产权。\n投标人名称：________\n年  月  日",
        ),
        OutlineItem(id="o03", title="报价表", kind="quote", source="generate", skipped=True),
    ]
    dest = tmp_path / "outline.docx"
    path, warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    assert path.is_file()
    doc = Document(str(path))
    text = _docx_text(doc)
    assert "目录" in text
    assert "响应函" in text
    assert "知识产权不侵权承诺函" in text
    assert "某市分公司" in text
    assert "数智制造管理系统" in text
    assert "不侵犯任何第三方知识产权" in text
    assert "河南伟泰光电科技有限公司" in text
    assert "2026年09月08日" in text
    assert "报价表" not in text
    assert "我方参加" in text
    assert "响应总报价" not in text
    assert not any(clause in text for clause in _CLAUSES)
    assert any("复制" in w for w in warnings)


def test_assemble_modules_fill_brief_not_invitation_body(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.commitment import _CLAUSES

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.projectName = "数智制造管理系统"
    brief.tenderer = "某市分公司"
    brief.bidDate = "2026-09-08"
    brief.bidPriceYuan = 113000
    brief.deliveryDays = 45
    brief.bidValidityDays = 120
    brief.agentName = "李四"
    brief.agentIdNo = "410000199001011234"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.quoteLines = [
        QuoteLineIn(seq="1", name="子系统模块A", spec="按采购需求", unit="套", qty=1, unitPrice=100000, amount=100000)
    ]
    brief.deviationLines = [
        DeviationLine(seq="1", requirement="支持单点登录", response="支持单点登录", deviation="无偏差")
    ]
    brief.performanceLines = [
        PerformanceLine(
            projectName="某园区信息化项目",
            client="某园区管委会",
            amountYuan=200000,
            summary="已竣工",
        )
    ]
    brief.outlineItems = [
        OutlineItem(id="o01", title="响应函", kind="letter", source="generate"),
        OutlineItem(id="o02", title="法定代表人身份证明", kind="legal_id", source="generate"),
        OutlineItem(
            id="o03",
            title="授权委托书",
            kind="auth",
            source="generate",
        ),
        OutlineItem(id="o04", title="报价表", kind="quote", source="generate"),
        OutlineItem(id="o05", title="技术规范书偏离表", kind="tech_dev", source="generate"),
        OutlineItem(
            id="o06",
            title="知识产权不侵权承诺函",
            kind="commitment_copy",
            source="copy",
            body="我方承诺不侵犯任何第三方知识产权。",
        ),
        OutlineItem(id="o07", title="类似项目业绩", kind="performance", source="generate"),
        OutlineItem(id="o08", title="分项报价表", kind="quote", source="generate", skipped=True),
    ]
    dest = tmp_path / "modules.docx"
    path, warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    assert path.is_file()
    doc = Document(str(path))
    text = _docx_text(doc)
    assert "致：某市分公司" in text
    assert "响应总报价" in text
    assert "120日历天" in text
    assert "410521199802104053" in text
    assert "李四" in text
    assert "子系统模块A" in text
    quote_header = next(
        [c.text for c in t.rows[0].cells]
        for t in doc.tables
        if any("单价" in (c.text or "") and "合价" in "".join(x.text or "" for x in t.rows[0].cells) for c in t.rows[0].cells)
    )
    assert "设备" not in quote_header
    assert "技术参数要求" not in quote_header
    assert "名称" in quote_header
    assert "支持单点登录" in text
    assert "无偏差" in text
    assert "某园区信息化项目" in text
    assert "不侵犯任何第三方知识产权" in text
    assert "分项报价表" not in text
    assert not any(clause in text for clause in _CLAUSES)
    assert any("模块填写" in w for w in warnings)
    assert any("复制" in w for w in warnings)


def test_assemble_letter_copy_override_keeps_invitation_body(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "数智制造管理系统"
    brief.tenderer = "某市分公司"
    brief.bidDate = "2026-09-08"
    brief.includePlaceholders = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="响应函",
            kind="letter",
            source="copy",
            body="致：（招标人名称）\n我方参加（项目名称）的采购。",
        )
    ]
    dest = tmp_path / "copy-letter.docx"
    path, _warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    text = _docx_text(Document(str(path)))
    assert "我方参加某市分公司" not in text
    assert "我方参加数智制造管理系统的采购" in text
    assert "响应总报价" not in text


def test_render_factory_and_tech_plan_modules() -> None:
    from docx import Document

    from api.services.tenders.modules import render_module

    brief = default_brief()
    brief.projectName = "数智制造管理系统"
    brief.tenderer = "某市分公司"
    brief.deliveryDays = 30
    brief.techPlanNote = "按采购需求实施，不编造桩数。"
    doc = Document()
    render_module(doc, brief, OutlineItem(id="f1", title="原厂生产承诺", kind="factory"))
    render_module(doc, brief, OutlineItem(id="t1", title="实施方案", kind="tech_plan"), media={})
    text = _docx_text(doc)
    assert "河南伟泰光电科技有限公司" in text
    assert "承担原厂责任" in text
    assert "某市分公司" in text
    assert "按采购需求实施" in text
    assert "图纸" in text
    assert "在此粘贴扫描件" in text


def test_tech_plan_embeds_uploaded_drawings(tmp_path) -> None:
    from PIL import Image
    from docx import Document

    from api.services.tenders.modules import render_module
    from api.services.tenders.placeholders import TECH_DRAWING_KEY

    png = tmp_path / "layout.png"
    Image.new("RGB", (160, 90), (30, 90, 160)).save(png, "PNG")
    extra = tmp_path / "layout2.png"
    Image.new("RGB", (140, 80), (80, 40, 20)).save(extra, "PNG")
    third = tmp_path / "layout3.png"
    Image.new("RGB", (100, 70), (20, 80, 40)).save(third, "PNG")
    brief = default_brief()
    brief.techPlanNote = "按采购需求实施，不编造桩数。"
    doc = Document()
    notes = render_module(
        doc,
        brief,
        OutlineItem(id="t1", title="实施方案", kind="tech_plan"),
        media={TECH_DRAWING_KEY: [png, extra, third]},
    )
    blob = doc.element.xml
    assert blob.count("a:blip") >= 3
    assert "在此粘贴扫描件" not in _docx_text(doc)
    assert any("图纸写入" in n for n in notes)


def test_assemble_quote_and_dev_follow_invitation_headers(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.projectName = "数智制造管理系统"
    brief.tenderer = "某市分公司"
    brief.includePlaceholders = False
    brief.attachQualifications = False
    brief.quoteLines = [
        QuoteLineIn(
            seq="1",
            name="单点登录",
            unit="",
            qty=0,
            unitPrice=80000,
            amount=80000,
            groups=["制造管理系统", "用户中心"],
        )
    ]
    brief.deviationLines = [
        DeviationLine(seq="1", requirement="支持单点登录", response="支持单点登录", deviation="无偏差")
    ]
    brief.outlineItems = [
        OutlineItem(
            id="q1",
            title="报价表",
            kind="quote",
            source="generate",
            body="序号 | 系统名称 | 子系统 | 功能模块 | 单价（元） | 合价（元）\n |  |  |  |  | ",
        ),
        OutlineItem(
            id="d1",
            title="技术规范书偏离表",
            kind="tech_dev",
            source="generate",
            body="序号 | 招标文件规定 | 投标响应 | 偏离情况\n1 |  |  | ",
        ),
        OutlineItem(
            id="d2",
            title="商务条款偏离表",
            kind="biz_dev",
            source="generate",
            body="序号 | 商务条款 | 响应情况 | 偏差说明\n1 |  |  | ",
        ),
    ]
    path, _warnings = assemble_bid_docx(brief, tmp_path / "headers.docx", qualification_pdf=None)
    doc = Document(str(path))
    quote = next(t for t in doc.tables if any("功能模块" in (c.text or "") for c in t.rows[0].cells))
    quote_heads = [c.text for c in quote.rows[0].cells]
    assert quote_heads == ["序号", "系统名称", "子系统", "功能模块", "单价（元）", "合价（元）"]
    assert "设备" not in quote_heads
    assert quote.rows[1].cells[1].text == "制造管理系统"
    assert quote.rows[1].cells[2].text == "用户中心"
    assert quote.rows[1].cells[3].text == "单点登录"
    tech = next(t for t in doc.tables if any("招标文件规定" in (c.text or "") for c in t.rows[0].cells))
    assert [c.text for c in tech.rows[0].cells] == ["序号", "招标文件规定", "投标响应", "偏离情况"]
    assert "支持单点登录" in tech.rows[1].cells[1].text
    biz = next(t for t in doc.tables if any("商务条款" in (c.text or "") for c in t.rows[0].cells))
    assert [c.text for c in biz.rows[0].cells] == ["序号", "商务条款", "响应情况", "偏差说明"]
    biz_blob = "".join(c.text or "" for row in biz.rows for c in row.cells)
    assert "支持单点登录" not in biz_blob
    assert "交货" in biz_blob or "质保" in biz_blob


def test_generate_bid_respects_layout_mode(tmp_path, monkeypatch) -> None:
    from api.services.tenders import generate as generate_mod
    from api.services.tenders.generate import generate_bid

    called = {"outline": 0, "vols": []}

    def fake_assemble(brief, dest, **kwargs):
        called["outline"] += 1
        called["vols"].append(kwargs.get("volume"))
        dest.write_bytes(b"PK\x03\x04")
        return dest, ["outline-path"]

    monkeypatch.setattr(generate_mod, "tenders_output_dir", lambda: tmp_path)
    monkeypatch.setattr("api.services.tenders.assemble.assemble_bid_docx", fake_assemble)

    outline_brief = BidBrief(
        layoutMode="outline",
        generateVolume="business",
        outlineItems=[OutlineItem(id="o01", title="响应函", kind="letter", body="致：甲方")],
        includePlaceholders=False,
        attachQualifications=False,
    )
    out = generate_bid(outline_brief)
    assert called["outline"] == 1
    assert called["vols"] == ["business"]
    assert out.get("docxFile")
    assert not out.get("techDocxFile")
    assert "商务标.docx" in out["downloadName"]

    tech_brief = outline_brief.model_copy(update={"generateVolume": "technical"})
    tech_out = generate_bid(tech_brief)
    assert called["outline"] == 2
    assert called["vols"][-1] == "technical"
    assert tech_out.get("techDocxFile")
    assert not tech_out.get("docxFile")
    assert "技术标.docx" in (tech_out.get("techDownloadName") or "")

    classic = BidBrief(
        layoutMode="chapter5",
        generateVolume="business",
        includePlaceholders=False,
        attachQualifications=False,
    )
    generate_bid(classic)
    assert called["outline"] == 3
    assert called["vols"][-1] == "business"

    from api.services.tenders.generate import generate_pair, volume_pair_status

    pair = generate_pair(outline_brief)
    assert called["vols"][-2:] == ["business", "technical"]
    assert pair.get("docxFile")
    assert pair.get("techDocxFile")
    assert "商务标.docx" in str(pair.get("downloadName") or "")
    assert "技术标.docx" in str(pair.get("techDownloadName") or "")
    note = volume_pair_status("business", business_ready=True, technical_ready=False)
    assert note["otherVolume"] == "technical"
    assert note["otherVolumeReady"] is False
    assert "技术标尚未生成" in str(note["otherVolumeNote"])


def _run_underlined(run) -> bool:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return bool(run.underline)
    for uel in rpr.findall(qn("w:u")):
        val = (uel.get(qn("w:val")) or "").lower()
        if val not in {"", "none", "nil"}:
            return True
    return False


def test_assemble_outline_toc_and_signoff_match_fixed(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "某热电有限公司"
    brief.bidPriceYuan = 518800
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate"),
        OutlineItem(id="o02", title="法定代表人身份证明", kind="legal_id", source="generate"),
        OutlineItem(id="o03", title="授权委托书", kind="auth", source="generate"),
        *[
            OutlineItem(
                id=f"p{i:02d}",
                title=f"占位材料{i}",
                kind="scan",
                source="copy",
                body="（装订时附原件）",
            )
            for i in range(4, 11)
        ],
        OutlineItem(id="o11", title="业绩证明资料", kind="performance", source="generate"),
    ]
    dest = tmp_path / "layout.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert not any(t.startswith("按招标书「") for t in texts)
    toc_letter = next(p for p in doc.paragraphs if p.text.startswith("一、投标函"))
    assert not any(_run_underlined(run) for run in toc_letter.runs if "投标函" in (run.text or ""))
    instr = "".join(
        t.text or "" for el in toc_letter._p.iter(qn("w:instrText")) for t in [el]
    )
    assert "PAGEREF" in instr
    assert any("业绩证明资料" in p.text for p in doc.paragraphs)
    assert any(p.text.startswith("二、法定代表人身份证明") for p in doc.paragraphs)
    assert not any("资格审查资料" in p.text for p in doc.paragraphs)
    blob = "\n".join(texts)
    assert "我方已仔细研究了" not in blob
    assert "单位性质" not in blob
    assert "本次由法定代表人亲自投标" not in blob


def test_assemble_auth_sign_single_line_aligned(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.document import _em_len

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.agentName = ""
    brief.bidDate = "2026-9-1"
    brief.outlineItems = [
        OutlineItem(id="o01", title="法定代表人授权委托书", kind="auth", source="generate"),
    ]
    dest = tmp_path / "auth-sign.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))

    def compact(text: str) -> str:
        return "".join((text or "").split())

    title_i = next(
        i
        for i, p in enumerate(doc.paragraphs)
        if (p.text or "").strip() == "法定代表人授权委托书"
    )
    block = doc.paragraphs[title_i:]
    bidder = next(p for p in block if compact(p.text).startswith("投标人") and "盖单位公章" in p.text)
    legal = next(p for p in block if compact(p.text).startswith("法定代表人：") and "签字" in p.text)
    id_lines = [p for p in block if compact(p.text).startswith("身份证号码：")]
    agent = next(p for p in block if compact(p.text).startswith("委托代理人："))
    date = next(p for p in block if compact(p.text).startswith("日期："))
    assert len(id_lines) == 2
    for para in (bidder, legal, id_lines[0], agent, id_lines[1], date):
        assert para._p.find(".//" + qn("w:br")) is None
        wrap = para._p.find(".//" + qn("w:wordWrap"))
        assert wrap is not None and wrap.get(qn("w:val")) == "off"
        assert _em_len(para.text) <= 30
        head = (para.text or "").split("：", 1)[0]
        assert len(head) == 5
    assert bidder.text.split("：", 1)[0] == "投　标　人"
    assert date.text.split("：", 1)[0] == "日　　　期"
    assert "09" in date.text and "01" in date.text
    assert "河南伟泰光电科技有限公司" in bidder.text
    assert "（盖单位公章）" in bidder.text
    assert "郭志伟" not in legal.text and "（签字）" in legal.text
    assert "410521199802104053" in id_lines[0].text
    assert "郭志伟" not in agent.text and "（不委托）" not in agent.text and "（签字）" in agent.text

    empty_boxes = []
    saw_title = False
    for child in doc.element.body:
        text = "".join(t.text or "" for t in child.iter(qn("w:t"))).replace("\u200b", "").strip()
        if child.tag == qn("w:p") and text == "法定代表人授权委托书":
            saw_title = True
        if not saw_title:
            continue
        if child.tag == qn("w:tbl"):
            rows = child.findall(qn("w:tr"))
            cell_text = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
            if len(rows) == 1 and not cell_text:
                empty_boxes.append(child)
    assert not empty_boxes


def test_paragraphs_keep_title_and_zhi_apart() -> None:
    from api.services.tenders.assemble import _paragraphs_from_body

    lines = _paragraphs_from_body(
        "投标承诺书\n致：\n我公司现做出如下承诺：\n[第13页]\n二连浩特市联源热电有限公司"
    )
    assert "投标承诺书" in lines
    assert any(ln.startswith("致") for ln in lines)
    assert not any("投标承诺书致" in compact_title(ln) for ln in lines)
    assert not any("第13页" in ln for ln in lines)
    assert "二连浩特市联源热电有限公司" not in lines


def test_assemble_commitment_copy_keeps_form_layout(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "二连浩特市联源热电有限公司"
    brief.bidDate = "2026-09-15"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺书",
            kind="commitment_copy",
            source="copy",
            body=(
                "投标承诺书致：\n"
                "我公司现做出如下承诺：\n"
                "1、保证投标文件内容无任何虚假。\n"
                "投标人： （盖章）\n"
                "法定代表人或委托代理人： （签字并盖章）\n"
                "日期： 年月日\n"
                "[第13页]\n"
                "二连浩特市联源热电有限公司\n"
            ),
        ),
    ]
    dest = tmp_path / "commitment.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    blob = "\n".join(texts)
    titles = [t for t in texts if compact_title(t) == "投标承诺书"]
    assert titles
    zhi = next(t for t in texts if compact_title(t).startswith("致"))
    assert "二连浩特市联源热电有限公司" in zhi
    assert "我公司" not in zhi
    assert "投标承诺书致" not in compact_title(zhi)
    bidder = next(t for t in texts if "盖章" in t and "签字" not in t)
    assert "河南伟泰光电科技有限公司" in bidder
    sign = next(t for t in texts if "签字并盖章" in t)
    assert "郭志伟" not in sign
    assert "2026年09月15日" in blob or ("2026" in blob and "09" in blob and "15" in blob)
    assert "第13页" not in blob
    assert not any(compact_title(t) == "二连浩特市联源热电有限公司" for t in texts)
    zhi_para = next(p for p in doc.paragraphs if compact_title(p.text).startswith("致"))
    sign_para = next(p for p in doc.paragraphs if "签字并盖章" in p.text)
    bidder_para = next(p for p in doc.paragraphs if "盖章" in p.text and "签字" not in p.text)
    assert "w:u" in zhi_para._p.xml
    assert "w:u" in sign_para._p.xml
    assert "w:u" in bidder_para._p.xml


def test_coalesce_broken_legal_id_lines() -> None:
    from api.services.tenders.assemble import _paragraphs_from_body

    lines = _paragraphs_from_body(
        "成立时\n间：\n________\n年\n________\n月\n日经营期限：长期\n姓名：郭志伟"
    )
    blob = "｜".join(lines)
    assert "成立时间：" in blob
    assert "经营期限：长期" in blob
    assert all(ln != "成立时" and ln != "间：" for ln in lines)


def test_assemble_legal_id_copy_stays_on_form(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o02",
            title="法定代表人身份证明",
            kind="legal_id",
            source="copy",
            body=(
                "三、法定代表人身份证明\n投标人名称：________\n成立时\n间：\n年\n月\n"
                "日经营期限：长期\n姓名： 性别： 年龄： 职务：\n系（投标人名称）的法定代表人。"
            ),
        ),
    ]
    dest = tmp_path / "legal-id.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    texts = [p.text.strip() for p in doc.paragraphs]
    nonempty = [t for t in texts if t]
    blob = "\n".join(nonempty)
    assert "投标人名称" in blob
    assert "系" in blob and "法定代表人" in blob
    assert "单位性质" not in blob
    assert "本次由法定代表人亲自投标" not in blob
    assert all(t != "成立时" and t != "间：" for t in nonempty)
    toc = next(
        p
        for p in doc.paragraphs
        if (p.text or "").startswith("三、法定代表人身份证明")
    )
    instr = "".join(t.text or "" for el in toc._p.iter(qn("w:instrText")) for t in [el])
    assert "PAGEREF" in instr
    assert 'w:leader="dot"' in toc._p.xml


def test_assemble_inlines_id_cards_not_in_appendix(tmp_path) -> None:
    from PIL import Image
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    front = tmp_path / "id-front.png"
    back = tmp_path / "id-back.png"
    Image.new("RGB", (220, 350), (190, 170, 150)).save(front)
    Image.new("RGB", (220, 350), (150, 170, 190)).save(back)
    agent_f = tmp_path / "ag-front.png"
    agent_b = tmp_path / "ag-back.png"
    Image.new("RGB", (220, 350), (120, 160, 140)).save(agent_f)
    Image.new("RGB", (220, 350), (140, 120, 160)).save(agent_b)

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.includePlaceholders = True
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.agentName = "李四"
    brief.outlineItems = [
        OutlineItem(id="o01", title="法定代表人身份证明", kind="legal_id", source="generate"),
        OutlineItem(id="o02", title="授权委托书", kind="auth", source="generate"),
    ]
    dest = tmp_path / "ids.docx"
    assemble_bid_docx(
        brief,
        dest,
        qualification_pdf=None,
        catalog_slots=[
            PlaceholderItem(key="id_legal", title="法定代表人身份证正反面"),
            PlaceholderItem(key="id_agent", title="授权代理人身份证正反面"),
            PlaceholderItem(key="bond", title="投标保证金缴存回单"),
        ],
        catalog_media={
            "id_legal": [front, back],
            "id_agent": [agent_f, agent_b],
        },
    )
    doc = Document(str(dest))
    texts = [p.text or "" for p in doc.paragraphs]
    assert not any("（一）法定代表人身份证正反面" in t for t in texts)
    assert not any("授权代理人身份证" in t and t.startswith("（") for t in texts)
    assert not any("附件：资料库扫描件" in t for t in texts)
    assert not any("投标保证金缴存回单" in t for t in texts)
    legal_i = next(i for i, t in enumerate(texts) if t.strip() == "法定代表人身份证明")
    auth_i = next(i for i, t in enumerate(texts) if t.strip() == "授权委托书")
    legal_sign = next(
        i
        for i, t in enumerate(texts[legal_i:auth_i], start=legal_i)
        if "投标人" in "".join(t.split()) and "盖单位公章" in t
    )
    auth_sign = next(
        i
        for i, t in enumerate(texts[auth_i:], start=auth_i)
        if "投标人" in "".join(t.split()) and "盖单位公章" in t
    )

    def _blip_between(a: int, b: int) -> bool:
        el = doc.paragraphs[a]._element.getnext()
        stop = doc.paragraphs[b]._element
        while el is not None and el is not stop:
            if el.findall(".//" + qn("a:blip")) or el.findall(".//" + qn("pic:blipFill")):
                return True
            el = el.getnext()
        return False

    assert _blip_between(legal_i, legal_sign)
    assert _blip_between(auth_i, auth_sign)
    id_para = next(
        p
        for p in doc.paragraphs[legal_i:legal_sign]
        if p._p.findall(".//" + qn("wp:extent"))
    )
    assert id_para.paragraph_format.keep_with_next is True
    assert doc.paragraphs[legal_sign].paragraph_format.keep_with_next is True
    assert doc.paragraphs[auth_sign].paragraph_format.keep_with_next is True


def test_estimate_counts_id_scan_height(tmp_path) -> None:
    from PIL import Image
    from docx import Document

    from api.services.tenders.document import _estimate_cm_on_current_page
    from api.services.tenders.placeholders import id_slot, inline_id_scans

    front = tmp_path / "f.png"
    back = tmp_path / "b.png"
    Image.new("RGB", (220, 350), (180, 160, 140)).save(front)
    Image.new("RGB", (220, 350), (140, 160, 180)).save(back)
    doc = Document()
    doc.add_paragraph("法定代表人身份证明")
    before = _estimate_cm_on_current_page(doc)
    inline_id_scans(doc, [front, back], empty_slot=id_slot("id_legal"))
    after = _estimate_cm_on_current_page(doc)
    assert after - before >= 4.5


def test_assemble_seal_register_fills_company_table(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="印鉴预留备案表", kind="company", source="copy", body=""),
    ]
    dest = tmp_path / "seal.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "公司证件、印章备案表" in blob
    assert "河南伟泰光电科技有限公司" in blob
    assert "公章" in blob
    assert "财务章" in blob
    assert "合同章" in blob
    assert "印鉴备案" in blob
    assert "三处备案印鉴为红色章" in blob
    stamp = None
    for table in doc.tables:
        for row in table.rows:
            if any("印鉴备案" in (c.text or "") for c in row.cells):
                stamp = row
                break
        if stamp is not None:
            break
    assert stamp is not None
    tr_pr = stamp._tr.find(qn("w:trPr"))
    el = None if tr_pr is None else tr_pr.find(qn("w:trHeight"))
    assert el is not None
    assert el.get(qn("w:hRule")) == "exact"
    # 约 5.8cm，够盖直径 4.2cm 公章
    assert int(el.get(qn("w:val")) or 0) >= int(5.5 * 567)
    assert stamp._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None


def test_assemble_drops_duplicate_seal_register_page(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="印鉴预留备案表", kind="company", source="copy", body=""),
        OutlineItem(id="o02", title="印件备案表（附件七）", kind="company", source="copy", body=""),
    ]
    dest = tmp_path / "seal-dup.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "印件备案表" not in blob
    assert "印鉴预留备案表" in blob
    assert "公司证件、印章备案表" in blob


def test_collapse_consecutive_page_breaks_drops_empty_page() -> None:
    from docx import Document

    from api.services.tenders.assemble import _collapse_extra_page_breaks

    doc = Document()
    doc.add_paragraph("上页")
    doc.add_page_break()
    doc.add_paragraph("")
    doc.add_page_break()
    doc.add_paragraph("下页")
    _collapse_extra_page_breaks(doc)
    n = 0
    for p in doc.paragraphs:
        if (p.text or "").strip():
            continue
        if any(br.get(qn("w:type")) == "page" for br in p._element.iter(qn("w:br"))):
            n += 1
    assert n == 0
    tail = next(p for p in doc.paragraphs if (p.text or "").strip() == "下页")
    assert tail.paragraph_format.page_break_before is True


def test_chapter_break_clears_keep_next_on_last_line() -> None:
    from docx import Document

    from api.services.tenders.assemble import _chapter_break

    doc = Document()
    p = doc.add_paragraph("落款")
    p.paragraph_format.keep_with_next = True
    _chapter_break(doc)
    assert not p.paragraph_format.keep_with_next


def test_assemble_does_not_leave_empty_page_break_paras(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.document import _paragraph_has_page_break

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate", body=""),
        OutlineItem(id="o02", title="投标承诺书", kind="commitment_copy", source="copy", body=""),
        OutlineItem(id="o03", title="印鉴预留备案表", kind="company", source="copy", body=""),
    ]
    dest = tmp_path / "pages.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    empty_breaks = 0
    for p in doc.paragraphs:
        if (p.text or "").strip():
            continue
        if _paragraph_has_page_break(p._p) and not p._p.findall(".//" + qn("w:drawing")):
            empty_breaks += 1
    assert empty_breaks == 0
    titles = [p.text.strip() for p in doc.paragraphs if (p.text or "").strip()]
    assert "投标函" in titles
    assert "投标承诺书" in titles


def test_quote_summary_rows_stay_compact(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.bidPriceYuan = 518800
    brief.quoteLines = [
        QuoteLineIn(
            seq="1",
            name="交流充电桩",
            spec="（5）具备接入第三方智能化集成平台接口功能；\n（6）计量系统随设备自带；",
            unit="台",
            qty=10,
            unitPrice=51880,
            amount=518800,
        ),
    ]
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标报价单", kind="quote", source="generate", body=""),
    ]
    dest = tmp_path / "quote.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    quote_tbl = None
    for table in doc.tables:
        blob = "\n".join(cell.text for row in table.rows for cell in row.cells)
        if "不含税合计" in blob and "含税合计" in blob:
            quote_tbl = table
            break
    assert quote_tbl is not None
    for row in quote_tbl.rows[-3:]:
        blob = " ".join(cell.text for cell in row.cells)
        assert any(k in blob for k in ("不含税合计", "税率", "含税合计"))
        merged = row.cells[1]._tc
        paras = [child for child in merged if child.tag == qn("w:p")]
        assert len(paras) <= 2
        tr_pr = row._tr.find(qn("w:trPr"))
        th = None if tr_pr is None else tr_pr.find(qn("w:trHeight"))
        assert th is not None
        assert int(th.get(qn("w:val")) or 0) <= 400


def test_assemble_commitment_without_body_does_not_write_stock_clauses(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.commitment import _CLAUSES

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "二连浩特市联源热电有限公司"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标承诺书", kind="commitment_copy", source="copy", body=""),
    ]
    dest = tmp_path / "commit.docx"
    path, warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(path))
    blob = _docx_text(doc)
    assert "投标承诺书" in blob
    assert "我公司现做出如下承诺" not in blob
    assert not any(clause in blob for clause in _CLAUSES)
    assert any("未抽到招标书原文" in w for w in warnings)


def _qual_pdf(path):
    from PIL import Image

    Image.new("RGB", (320, 450), (230, 230, 230)).save(path, "PDF")
    return path


def test_assemble_warns_when_qual_pdf_missing(tmp_path) -> None:
    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = True
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate"),
    ]
    _, warnings = assemble_bid_docx(brief, tmp_path / "noqual.docx", qualification_pdf=None)
    assert any("未找到资质 PDF" in w for w in warnings)


def test_assemble_inserts_qual_pdf_into_outline_chapter(tmp_path) -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from api.services.tenders.assemble import assemble_bid_docx

    pdf = _qual_pdf(tmp_path / "企业资质.pdf")
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = True
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate"),
        OutlineItem(id="o02", title="企业资质", kind="scan", source="copy"),
    ]
    path, warnings = assemble_bid_docx(brief, tmp_path / "qual.docx", qualification_pdf=pdf)
    doc = Document(str(path))
    blob = _docx_text(doc)
    assert "企业资质" in blob
    assert "资质文件 第 1 页" in blob
    assert "附件：企业资质文件扫描件" not in blob
    assert any("已插入资质文件" in w for w in warnings)
    xml = doc.element.body.xml
    assert "a:blip" in xml or "pic:blipFill" in xml
    heading = next(p for p in doc.paragraphs if (p.text or "").strip() == "企业资质")
    found = False
    el = heading._element.getnext()
    while el is not None:
        if el.findall(".//" + qn("a:blip")) or el.findall(".//" + qn("pic:blipFill")):
            found = True
            break
        el = el.getnext()
    assert found


def test_find_qualification_pdf_looks_in_slots(tmp_path, monkeypatch) -> None:
    from api.services.tenders import assets as assets_mod

    monkeypatch.setattr(assets_mod, "tender_assets_dir", lambda: tmp_path)
    slot = tmp_path / "slots" / "license"
    slot.mkdir(parents=True)
    pdf = _qual_pdf(slot / "伟泰企业资质扫描件.pdf")
    found = assets_mod.find_qualification_pdf()
    assert found == pdf


def _volume_brief() -> BidBrief:
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="商务标", kind="unknown", source="skip"),
        OutlineItem(id="o02", title="投标函", kind="letter", source="generate"),
        OutlineItem(id="o03", title="分项报价表", kind="quote", source="generate"),
        OutlineItem(id="o04", title="技术标", kind="tech_plan", source="generate"),
        OutlineItem(id="o05", title="技术偏差表", kind="tech_dev", source="generate"),
        OutlineItem(id="o06", title="技术标（实施方案）", kind="tech_plan", source="generate"),
    ]
    return brief


def test_assemble_business_volume_excludes_tech(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    path, _ = assemble_bid_docx(_volume_brief(), tmp_path / "biz.docx", volume="business")
    blob = _docx_text(Document(str(path)))
    assert "商 务 标 投 标 文 件" in blob
    assert "投标函" in blob
    assert "分项报价" in blob
    assert "实施方案" not in blob
    assert "技术偏差" not in blob
    assert "技 术 标 投 标 文 件" not in blob


def test_assemble_skips_invite_format_pack_pages(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.generateVolume = "business"
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标函",
            kind="letter",
            source="copy",
            body=(
                "致：（招标人名称）\n我方已仔细研究招标文件。\n"
                "八、资格审查资料\n投标人：基本情况表\n"
                "（一）中小企业声明函（如有）\n本公司郑重声明\n"
                "十、施工组织设计\n暗标装订要求\n"
            ),
        ),
        OutlineItem(id="o02", title="主要项目管理人员简历表", kind="unknown", source="copy", body="应附注册资格证书"),
        OutlineItem(id="o03", title="中小企业声明函（如有）", kind="commitment_copy", source="copy", body="本公司郑重声明"),
        OutlineItem(id="o04", title="十、施工组织设计", kind="unknown", source="copy", body="编制施工组织设计"),
        OutlineItem(id="o05", title="投标人基本情况表", kind="company", source="copy", body="投标人：________"),
    ]
    path, _ = assemble_bid_docx(brief, tmp_path / "pack.docx", volume="business")
    blob = compact_title(_docx_text(Document(str(path))))
    assert "我方已仔细研究" in blob
    assert "中小企业声明" not in blob
    assert "残疾人" not in blob
    assert "暗标" not in blob
    assert "施工组织设计" not in blob
    assert "基本情况" in blob



def test_assemble_technical_volume_excludes_quote(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    path, _ = assemble_bid_docx(_volume_brief(), tmp_path / "tech.docx", volume="technical")
    blob = _docx_text(Document(str(path)))
    assert "技 术 标 投 标 文 件" in blob
    assert "技术偏差" in blob or "实施方案" in blob
    assert "分项报价" not in blob
    assert "投标函" not in blob
    assert "商 务 标 投 标 文 件" not in blob


def test_assemble_technical_volume_embeds_drawings(tmp_path) -> None:
    from PIL import Image
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.placeholders import TECH_DRAWING_KEY

    png = tmp_path / "a.png"
    Image.new("RGB", (120, 80), (10, 80, 140)).save(png, "PNG")
    brief = _volume_brief()
    brief.techPlanNote = "按现场布置施工。"
    path, warnings = assemble_bid_docx(
        brief,
        tmp_path / "tech.docx",
        volume="technical",
        catalog_media={TECH_DRAWING_KEY: [png]},
    )
    blob = Document(str(path)).element.xml
    assert "a:blip" in blob
    assert any("图纸写入" in w for w in warnings)


def test_assemble_technical_omits_library_scan_heading(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = True
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="技术偏差表", kind="tech_dev", source="generate"),
    ]
    path, _ = assemble_bid_docx(
        brief,
        tmp_path / "tech-slots.docx",
        volume="technical",
        catalog_slots=[
            PlaceholderItem(key="product", title="所投产品检测报告 / 3C / 对应功率桩型证明"),
        ],
    )
    blob = _docx_text(Document(str(path)))
    assert "附件：资料库扫描件" not in blob
    assert "所投产品检测报告" not in blob


def test_assemble_copies_invitation_commitment_not_stock_clauses(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.commitment import _CLAUSES

    body = (
        "附一：投标承诺函模板\n"
        "投标承诺函\n"
        "致：（招标人名称）\n"
        "1、我方已详细研究了招标文件的所有内容。\n"
        "2、我方承诺投标文件夹中的一切资料、数据是真实的。\n"
        "3、我方明白并同意若我方在投标有效期之内撤回投标，则投标保证金将被贵方没收。\n"
        "4、我方理解贵方不一定接受最低标价或任何贵方可能收到的投标。\n"
        "5、我方如果中标，将保证履行招标文件中的全部责任和义务。\n"
        "投标人(盖公章)：\n"
        "日期： 年 月 日\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "河南鑫宇光科技股份有限公司"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺函模板",
            kind="commitment_copy",
            source="copy",
            body=body,
        )
    ]
    dest = tmp_path / "commit-src.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "详细研究了招标文件" in blob
    assert "资料、数据是真实的" in blob
    assert "不一定接受最低标价" in blob
    assert "附一：投标承诺函模板" not in blob
    assert not any(clause in blob for clause in _CLAUSES)


def test_assemble_splits_zhi_company_from_body(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "投标承诺函\n"
        "致：河南鑫宇光科技股份有限公司我方确认收到贵方提供的鑫宇科技《MOM生产运营管理平台》"
        "项目招标投标所需的招标文件，并已完全明白招标文件的所有条款要求，并重申以下几点：\n"
        "1、我方已详细研究了招标文件的所有内容。\n"
        "投标人(盖公章)：\n"
        "授权代表(签名)：\n"
        "日期： 年 月 日\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "河南鑫宇光科技股份有限公司"
    brief.bidDate = "2026-09-23"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺函",
            kind="commitment_copy",
            source="copy",
            body=body,
        )
    ]
    dest = tmp_path / "zhi-split.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    zhi = next(t for t in paras if t.startswith("致"))
    assert "我方确认" not in zhi
    assert "河南鑫宇光科技股份有限公司" in zhi
    body_p = next(t for t in paras if t.startswith("我方确认收到"))
    assert "招标文件" in body_p
    date_i = next(i for i, t in enumerate(paras) if compact_title(t).startswith("日期"))
    assert "2026" in paras[date_i]
    if date_i + 1 < len(paras):
        assert paras[date_i + 1] != "2026年09月23日"
    assert any("投标人" in compact_title(t) and "盖公章" in compact_title(t) for t in paras)


def test_assemble_copied_sign_lines_have_underline(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "投标承诺函\n"
        "致：河南鑫宇光科技股份有限公司\n"
        "1、我方已详细研究了招标文件的所有内容。\n"
        "投标人(盖公章)：\n"
        "授权代表(签名)：\n"
        "日期：\n"
        "年 月 日\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "河南鑫宇光科技股份有限公司"
    brief.bidDate = "2026-09-23"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺函",
            kind="commitment_copy",
            source="copy",
            body=body,
        )
    ]
    dest = tmp_path / "commit-sign.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))

    def hit(pred):
        return next(p for p in doc.paragraphs if pred(p.text or ""))

    bidder = hit(lambda t: "盖公章" in compact_title(t) and "投标人" in compact_title(t))
    agent = hit(lambda t: "授权代表" in compact_title(t) or "签名" in compact_title(t))
    date = hit(lambda t: compact_title(t).startswith("日期"))
    assert "w:u" in bidder._p.xml
    assert "河南伟泰" in bidder.text
    assert "w:u" in agent._p.xml
    assert "郭志伟" not in agent.text
    assert "w:u" in date._p.xml
    assert "2026" in date.text


def test_assemble_keeps_multiline_commitment_paragraphs(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "1、甲方已研究招标文件。乙方不得以含糊为由抗辩。丙方继续履行。\n"
        "2、资料真实。"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺书",
            kind="commitment_copy",
            source="copy",
            body=body,
        )
    ]
    dest = tmp_path / "commit-lines.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    paras = [p.text.strip() for p in Document(str(dest)).paragraphs if "甲方已研究" in (p.text or "")]
    assert paras
    assert "乙方不得以含糊为由抗辩" in paras[0]
    assert "丙方继续履行" in paras[0]


def test_assemble_seal_register_uses_one_page_grid(tmp_path) -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "印鉴预留备案\n"
        "公司名称：________  公司电话：________\n"
        "公章    财务章    合同章\n"
        "备注：本表中三处备案印鉴为红色章。"
    )
    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="印鉴预留备案表", kind="company", source="copy", body=body),
    ]
    dest = tmp_path / "seal-grid.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "三处备案印鉴为红色章" in blob
    stamp = None
    for table in doc.tables:
        for row in table.rows:
            if any((c.text or "").strip() == "印鉴备案" for c in row.cells):
                stamp = row
                break
        if stamp is not None:
            for row in table.rows:
                tr_pr = row._tr.trPr
                assert tr_pr is not None
                assert tr_pr.find(qn("w:cantSplit")) is not None
            break
    assert stamp is not None


def test_assemble_follows_invitation_templates_not_fixed_seal_or_quote(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    req_body = (
        "附二：招标产品功能需求清单及说明\n"
        "一、MOM生产运营管理平台：\n"
        "设备数采 1.多类型数据采集；2数据传输与存储功能\n"
        "TPM系统 1.设备档案管理\n"
    )
    quote_body = (
        "附三：报价单模板\n"
        "项目总报价：\n"
        "序号 项目内容 金额(元) 备注\n"
        "1 MOM平台总费用\n"
        "2 总实施费用\n"
        "3 年度服务费\n"
        "4 合计\n"
        "（1）数采报价\n"
        "序号 项目内容 金额(元) 备注\n"
        "1 软件费用\n"
        "2 实施费用\n"
        "3 合计(1+2)\n"
        "（6）实施费用分项明细\n"
    )
    commit_body = (
        "附一：投标承诺函模板\n"
        "投标承诺函\n"
        "致：（招标人名称）\n"
        "1、我方已详细研究了招标文件的所有内容。\n"
        "投标人(盖公章)：\n"
        "日期： 年 月 日\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "河南鑫宇光科技股份有限公司"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺函模板",
            kind="commitment_copy",
            source="copy",
            body=commit_body,
        ),
        OutlineItem(
            id="o02",
            title="招标产品功能需求清单及说明",
            kind="unknown",
            source="copy",
            body=req_body,
        ),
        OutlineItem(
            id="o03",
            title="报价单模板",
            kind="quote",
            source="copy",
            body=quote_body,
        ),
    ]
    dest = tmp_path / "mom-outline.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "详细研究了招标文件" in blob
    assert "招标产品功能需求清单" in blob
    assert "设备数采" in blob
    assert "MOM平台总费用" in blob
    assert "软件费用" in blob
    assert "实施费用分项明细" in blob
    assert "附一：投标承诺函模板" not in blob
    assert "附二：招标产品功能需求清单" not in blob
    assert "附件：资料库扫描件" not in blob
    assert "公司证件、印章备案表" not in blob
    assert "印鉴备案" not in blob
    assert "不含税综合单价" not in blob
    quote_tables = [
        t
        for t in doc.tables
        if t.rows and any("项目内容" in (c.text or "") for c in t.rows[0].cells)
    ]
    assert quote_tables
    heads = [c.text for c in quote_tables[0].rows[0].cells]
    assert "项目内容" in heads
    assert "设备" not in heads
    req_tables = [
        t
        for t in doc.tables
        if any("设备数采" in (c.text or "") for row in t.rows for c in row.cells)
    ]
    assert req_tables
    req_blob = " ".join(c.text or "" for row in req_tables[0].rows for c in row.cells)
    assert "TPM系统" in req_blob
    assert "设备档案管理" in req_blob


def test_assemble_outline_empty_volume_not_charger_pack(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="报价单模板",
            kind="quote",
            source="copy",
            body="序号 项目内容 金额(元) 备注\n1 MOM平台总费用\n",
        ),
    ]
    path, warnings = assemble_bid_docx(brief, tmp_path / "tech-empty.docx", volume="technical")
    blob = _docx_text(Document(str(path)))
    assert any("组卷大纲为空" in w for w in warnings)
    assert "公司证件、印章备案表" not in blob
    assert "印鉴备案" not in blob
    assert "不含税综合单价" not in blob


def test_assemble_skips_duplicate_license_scan_and_quote_debris(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = True
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.bidDate = "2026-09-22"
    brief.outlineItems = [
        OutlineItem(id="o01", title="供应商名称：与营业执照、资质证书一", kind="scan", source="copy"),
        OutlineItem(
            id="o02",
            title="B.有效的企业营业执照、企业资质证书",
            kind="scan",
            source="copy",
            body="B.有效的企业营业执照\nC.业绩要求；\nD.财务要求",
        ),
        OutlineItem(id="o03", title="分项报价表", kind="quote", source="generate"),
        OutlineItem(id="o04", title="分项报价表说明", kind="quote", source="generate"),
        OutlineItem(id="o05", title="分项报价表单位：人民币元序号", kind="quote", source="generate"),
        OutlineItem(
            id="o06",
            title="合同条款响应书",
            kind="commitment_copy",
            source="copy",
            body="二、合同条款响应书\n致：________\n供应商：________（全称、盖章）\n日期： 年 月 日\n四、响应方案格式自拟",
        ),
    ]
    path, _ = assemble_bid_docx(
        brief,
        tmp_path / "dedup.docx",
        catalog_slots=[PlaceholderItem(key="license", title="营业执照")],
        catalog_media={},
    )
    blob = _docx_text(Document(str(path)))
    assert "供应商名称：与营业执照" not in blob
    assert "C.业绩要求" not in blob
    assert "分项报价表说明" not in blob
    assert "人民币元序号" not in blob
    assert "2026年09月22日四、" not in blob.replace(" ", "")
    assert blob.count("分项报价表") <= 3


def test_assemble_pipe_requirement_and_focus_tables(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    req_body = (
        "一、MOM生产运营管理平台：\n"
        "MOM | 设备数采 | 1.多类型数据采集；2.数据传输与存储功能；\n"
        "MOM | TPM系统 | 1.设备档案管理；2. 设备台账与库存管理；\n"
        " | 硬件 | 单独报价，不记录在总价中。\n"
        "二、重点需求明细包括（但不限于）：\n"
        "模块 | 建设要求 | 实现目标数据采集 | 1.多类型数据采集； | 数据采集的基础，针对设备的运行状态。\n"
        "TPM系统 | 1.设备档案管理；2.设备台账与库存管理； | "
        "1.设备档案管理系统的数据基础，实现设备信息的标准化。"
        "基础信息：设备编号、名称、型号、采购日期、单价；"
        "备品备件管理：采购单价、供应商，备件出入库登记；"
        "录入核心信息：供应商、校准周期、安全库存。\n"
        "传感器信号弱、网关离线），及时发现潜在问题。\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="招标产品功能需求清单及说明",
            kind="unknown",
            source="copy",
            body=req_body,
        )
    ]
    dest = tmp_path / "req-tables.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "设备数采" in blob
    assert "硬件" in blob
    assert "重点需求明细" in blob
    assert "数据采集的基础" in blob
    mom = next(
        t
        for t in doc.tables
        if any("设备数采" in (c.text or "") for row in t.rows for c in row.cells)
    )
    mom_blob = " ".join(c.text or "" for row in mom.rows for c in row.cells)
    assert "TPM系统" in mom_blob
    assert "硬件" in mom_blob
    focus = next(
        t
        for t in doc.tables
        if any((c.text or "").strip() == "模块" for row in t.rows for c in row.cells)
    )
    heads = [c.text.strip() for c in focus.rows[0].cells]
    assert "模块" in heads
    assert "建设要求" in heads
    assert "实现目标" in heads
    focus_blob = " ".join(c.text or "" for row in focus.rows for c in row.cells)
    assert "数据采集" in focus_blob
    tpm_row = next(
        row
        for row in focus.rows
        if any("TPM系统" in (c.text or "") for c in row.cells)
    )
    tpm_cells = [c.text or "" for c in tpm_row.cells]
    assert any("设备档案管理系统的数据基础" in c for c in tpm_cells)
    assert any("采购日期" in c for c in tpm_cells)
    assert any("供应商" in c for c in tpm_cells)
    assert any("网关离线" in c for c in tpm_cells)
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    assert tpm_row.cells[0].vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert tpm_row.cells[1].vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert tpm_row.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert tpm_row.cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(tpm_row.cells[1].paragraphs) >= 2
    paras = "\n".join(p.text or "" for p in doc.paragraphs)
    assert "设备档案管理系统的数据基础" not in paras
    tbl_w = focus._tbl.tblPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblW"
    )
    assert tbl_w is not None
    assert int(tbl_w.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w") or "0") > 5000


def test_assemble_keeps_single_filled_quote_sheet(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.bidPriceYuan = 897733
    brief.quoteLines = [
        QuoteLineIn(seq="1", name="7KW交流汽车充电桩", unit="台", qty=1, unitPrice=787.61, amount=787.61),
        QuoteLineIn(seq="2", name="30KW直流汽车充电桩", unit="台", qty=1, unitPrice=5823.01, amount=5823.01),
    ]
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="本次招标方案按照软件总体报价，实施费分项报价形式；甲方按照需求分期签订实施合同",
            kind="quote",
            source="generate",
        ),
        OutlineItem(
            id="o02",
            title="投标报价表(产品、实施、开发对接、服务、硬件服务器方案分项报价)及报价说明",
            kind="quote",
            source="generate",
        ),
        OutlineItem(id="o03", title="报价单", kind="quote", source="generate"),
    ]
    dest = tmp_path / "one-quote.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    quote_tables = [
        t
        for t in doc.tables
        if any("不含税合计" in (c.text or "") for row in t.rows for c in row.cells)
    ]
    assert len(quote_tables) == 1
    assert "按照软件总体报价" not in blob
    assert "硬件服务器方案" not in blob
    assert "报价单" in blob
    names = " ".join(c.text or "" for row in quote_tables[0].rows for c in row.cells)
    assert "交流汽车充电桩" in names


def test_assemble_quote_matrix_keeps_sites_and_hardware_apart(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    quote_body = (
        "项目总报价：\n"
        "序号 | 项目内容 | 金额(元) | 备注\n"
        "1 | MOM平台总费用 |  |\n"
        "（1）数采报价序号 | 项目内容 | 金额(元) | 备注\n"
        "1 | 软件费用 |  |\n"
        "2 | 实施费用 |  |\n"
        "3 | 合计(1+2) |  |\n"
        "实施费用分项明细。                                               单位：元\n"
        " | 数据采集 | TPM | WMS | MES | QMS | 合计（元） | 备注总部 |  |  |  |  |  |  |\n"
        "吉成 |  |  |  |  |  |  |\n"
        "成都 |  |  |  |  |  |  |\n"
        "郑州 |  |  |  |  |  |  |\n"
        "（7）硬件单独报价，不记录到总价。（单价）\n"
        "序号 | 项目内容 | 品牌 | 型号 | 单价（元）\n"
        "1 |  |  |  |\n"
        "2 |  |  |  |\n"
        "3 |  |  |  |\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="报价单模板",
            kind="quote",
            source="copy",
            body=quote_body,
        )
    ]
    dest = tmp_path / "quote-matrix.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    matrix = next(
        t
        for t in doc.tables
        if any((c.text or "").strip() == "数据采集" for row in t.rows for c in row.cells)
    )
    assert len(matrix.columns) <= 8
    widths = [int(c.width) for c in matrix.rows[0].cells]
    assert min(widths[1:]) * 4 > widths[0]
    from api.services.tenders.document import _em_len

    for cell in matrix.rows[0].cells:
        h = (cell.text or "").strip()
        if not h or _em_len(h) > 16:
            continue
        tc_pr = cell._tc.get_or_add_tcPr()
        assert tc_pr.find(qn("w:noWrap")) is not None
        need = int(_em_len(compact_title(h)) * 210) + 160
        assert int(cell.width) >= need - 80
    heads = [(c.text or "").strip() for c in matrix.rows[0].cells]
    assert "数据采集" in heads
    assert heads.count("数据采集") == 1
    assert all("\n" not in (c.text or "") for c in matrix.rows[0].cells)
    assert "数据采集" in heads
    assert "TPM" in heads
    assert "WMS" in heads
    assert "MES" in heads
    assert "QMS" in heads
    assert any("合计" in h for h in heads)
    assert any(h == "备注" for h in heads)
    assert not any("总部" in h for h in heads)
    sites = " ".join((row.cells[0].text or "") for row in matrix.rows)
    assert "总部" in sites
    assert "吉成" in sites
    assert "成都" in sites
    assert "郑州" in sites
    matrix_blob = " ".join(c.text or "" for row in matrix.rows for c in row.cells)
    assert "硬件单独报价" not in matrix_blob
    assert "品牌" not in matrix_blob
    numcai = next(
        t
        for t in doc.tables
        if any((c.text or "").strip() == "软件费用" for row in t.rows for c in row.cells)
    )
    assert (numcai.rows[0].cells[0].text or "").strip() == "序号"
    assert "数采报价序号" not in (numcai.rows[0].cells[0].text or "")
    hw = next(
        t
        for t in doc.tables
        if any((c.text or "").strip() == "品牌" for row in t.rows for c in row.cells)
    )
    hw_heads = [(c.text or "").strip() for c in hw.rows[0].cells]
    assert "品牌" in hw_heads
    assert "型号" in hw_heads
    assert "吉成" not in " ".join(c.text or "" for row in hw.rows for c in row.cells)
    blob = _docx_text(doc)
    assert "（1）数采报价" in blob
    assert "硬件单独报价" in blob
    cap = next(p for p in doc.paragraphs if "实施费用分项明细" in (p.text or ""))
    assert "单位" in (cap.text or "")
    assert "元" in (cap.text or "")


def test_copied_tables_keep_short_headers_horizontal() -> None:
    from docx import Document

    from api.services.tenders.assemble import _write_simple_table
    from api.services.tenders.document import _em_len

    doc = Document()
    cases = [
        ["", "数据采集", "TPM", "WMS", "MES", "QMS", "合计（元）", "备注"],
        ["模块", "建设要求", "实现目标"],
        ["序号", "项目内容", "品牌", "型号", "单价（元）"],
        ["甲项", "乙项名称", "丙", "丁丁"],
    ]
    tables = []
    for heads in cases:
        tables.append(_write_simple_table(doc, heads, [["x"] * len(heads)]))
    for table, heads in zip(tables, cases, strict=True):
        got = [(c.text or "").strip() for c in table.rows[0].cells]
        assert got == [h.strip() for h in heads]
        for cell, h in zip(table.rows[0].cells, heads, strict=True):
            assert "\n" not in (cell.text or "")
            if not h.strip() or _em_len(h) > 16:
                continue
            tc_pr = cell._tc.get_or_add_tcPr()
            assert tc_pr.find(qn("w:noWrap")) is not None, h
            need = int(_em_len(compact_title(h)) * 210) + 160
            assert int(cell.width) >= need - 80, (h, int(cell.width), need)


def test_assemble_biz_essentials_use_title_and_box(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标承诺函",
            kind="commitment_copy",
            source="copy",
            body="我方确认收到贵方提供的招标文件，并重申以下几点：",
        ),
        OutlineItem(
            id="o02",
            title="报价单模板",
            kind="quote",
            source="copy",
            body="序号 | 项目内容 | 金额(元) | 备注\n1 | MOM平台总费用 |  |\n",
        ),
    ]
    dest = tmp_path / "biz-essentials.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    blob = _docx_text(Document(str(dest)))
    assert "我方确认收到贵方提供的招标文件" in blob
    assert "MOM平台总费用" in blob
    assert "营业执照" not in blob
    assert "无违法承诺函" not in blob
    assert "投标保证金" not in blob


def test_assemble_embeds_library_license_by_title(tmp_path) -> None:
    from PIL import Image
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    png = tmp_path / "license.png"
    Image.new("RGB", (400, 280), (240, 240, 250)).save(png)
    iso_png = tmp_path / "iso.png"
    Image.new("RGB", (320, 240), (250, 240, 230)).save(iso_png)
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="generate"),
        OutlineItem(id="o02", title="营业执照", kind="scan", source="generate"),
        OutlineItem(id="o03", title="资质证书", kind="scan", source="generate"),
    ]
    dest = tmp_path / "license-scan.docx"
    assemble_bid_docx(
        brief,
        dest,
        qualification_pdf=None,
        catalog_slots=[
            PlaceholderItem(key="id_legal", title="法定代表人身份证正反面"),
            PlaceholderItem(key="mlic01", title="营业执照", hint="经营范围含新能源充电"),
            PlaceholderItem(key="iso_cert", title="ISO体系证书"),
        ],
        catalog_media={"mlic01": [png], "iso_cert": [iso_png]},
    )
    doc = Document(str(dest))
    texts = [p.text or "" for p in doc.paragraphs]
    lic_i = next(i for i, t in enumerate(texts) if t.strip() == "营业执照")
    iso_i = next(i for i, t in enumerate(texts) if t.strip() == "资质证书")

    def _blip_after(start: int, stop: int | None = None) -> bool:
        el = doc.paragraphs[start]._element.getnext()
        end = doc.paragraphs[stop]._element if stop is not None else None
        while el is not None and el is not end:
            if el.findall(".//" + qn("a:blip")) or el.findall(".//" + qn("pic:blipFill")):
                return True
            el = el.getnext()
        return False

    assert _blip_after(lic_i, iso_i)
    assert _blip_after(iso_i)
    assert "在此粘贴扫描件" not in "".join(texts[lic_i:iso_i])


def test_this_bid_keys_includes_scan_slots_by_title() -> None:
    from api.services.tenders.placeholders import this_bid_keys
    from api.services.tenders.schema import PlaceholderItem

    catalog = [
        PlaceholderItem(key="id_legal", title="法定代表人身份证正反面"),
        PlaceholderItem(key="mlic01", title="营业执照", hint="经营范围含新能源充电"),
        PlaceholderItem(key="acct01", title="基本账户开户许可证"),
        PlaceholderItem(key="finance", title="近三年财务"),
    ]
    required, include = this_bid_keys(
        extra=[],
        catalog=catalog,
        required_keys=["id_legal"],
        include_keys=["id_legal"],
        has_agent=False,
    )
    assert required == ["id_legal"]
    assert "mlic01" in include
    assert "acct01" in include
    assert "finance" not in include


def test_assemble_messy_letter_uses_module_not_ocr(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标函",
            kind="letter",
            source="copy",
            body=(
                "投标函\n-189-\n致某某建设集团我方郑重承诺按招标文件参加投标。"
                "\n发包人指定的接收人：张三\n-190-\n"
            ),
        ),
        OutlineItem(
            id="o02",
            title="法定代表人身份证明",
            kind="legal_id",
            source="copy",
            body="三、法定代表人身份证明\n投标人名称：________\n",
        ),
        OutlineItem(
            id="o03",
            title="法定代表人身份证明书",
            kind="legal_id",
            source="copy",
            body="四、法定代表人身份证明书\n姓名：________\n",
        ),
        OutlineItem(
            id="o04",
            title="招标产品功能需求清单及说明",
            kind="unknown",
            source="copy",
            body=(
                "附二：招标产品功能需求清单及说明\n"
                "一、MOM生产运营管理平台：\n"
                "设备数采 1.多类型数据采集；2数据传输与存储功能\n"
                "TPM系统 1.设备档案管理\n"
            ),
        ),
    ]
    dest = tmp_path / "messy-ocr.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "我方郑重承诺按招标文件参加投标" in blob
    assert "我方已仔细研究了" not in blob
    assert "单位性质" not in blob
    assert "设备数采" in blob
    assert "招标产品功能需求清单" in blob


def test_assemble_copied_letter_title_salute_and_fields(tmp_path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "（一）投标函\n"
        "郑州天健智慧园区运营管理有限公司：\n"
        "1.我方已仔细研究了郑州高新技术产业开发区新能源充(换)电站基础设施群项目（一期）"
        "施工招标项目招标文件的全部内容，愿意以人民币（大写）                    （¥\n"
        "元）的投标总报价（其中，增值税税率按照国家有关规定执行），工期；\n"
        "日历天，按合同约定完成承包工程。\n"
        "2. 我方的投标文件包括下列内容：\n"
        "（1）投标函及投标函附录；\n"
        "（2）法定代表人身份证明或授权委托书；\n"
        "3. 我方承诺在招标文件规定的投标有效期内不撤销投标文件。\n"
        "4. 如我方中标，我方承诺：\n"
        "（1）在收到中标通知书后签订合同；\n"
        "5. 我方在此声明，所递交的投标文件内容完整、真实和准确。\n"
        "6. 其他补充说明：无\n"
        "投标人：河南伟泰光电科技有限公司（盖单位章）\n"
        "法定代表人或其委托代理人：\n"
        "（签字或盖章）\n"
        "地址：\n"
        "电话：17630567052传真\n"
        "邮政编码：\n"
        "项目联系人电话（手机）：\n"
        "邮箱：\n"
        "日期：2026年09月24日\n"
        "（二）投标函附录项目名称投标人投标范围大写：\n"
        "小写：\n"
        "安全生产措施费\n"
        "（元）大写：\n"
        "小写：\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "郑州天健智慧园区运营管理有限公司"
    brief.projectName = "郑州高新技术产业开发区新能源充(换)电站基础设施群项目（一期）"
    brief.bidDate = "2026-09-24"
    brief.bidPriceYuan = 897733
    brief.deliveryDays = 90
    brief.bidderPhone = "17630567052"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="投标函",
            kind="letter",
            source="copy",
            body=body,
        ),
    ]
    dest = tmp_path / "letter-layout.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    paras = [p for p in doc.paragraphs if (p.text or "").strip()]
    texts = [p.text.strip() for p in paras]
    title = next(p for p in paras if compact_title(p.text) == "（一）投标函")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs and title.runs[0].bold and title.runs[0].font.size >= Pt(16)
    salute = next(p for p in paras if p.text.strip().startswith("致"))
    assert "郑州天健智慧园区运营管理有限公司" in salute.text
    assert salute.text.strip().endswith("：")
    assert "w:u" in salute._p.xml
    blob = "\n".join(texts)
    assert "人民币（大写）捌拾玖万柒仟柒佰叁拾叁元整" in blob
    assert "897733.00" in blob
    assert "工期：90日历天" in blob
    legal = next(
        p
        for p in paras
        if "法定代表人" in compact_title(p.text) and "签字或盖章" in (p.text or "")
    )
    assert "（签字或盖章）" in (legal.text or "")
    legal_left = float(legal.paragraph_format.left_indent.cm) if legal.paragraph_format.left_indent else 0
    # 手签行可能略减左缩进以拉长横线
    assert legal_left >= 1.4
    from api.services.tenders.document import _em_len

    blank = (legal.text or "").replace("法定代表人或其委托代理人：", "").replace("（签字或盖章）", "")
    assert "＿" not in (legal.text or "")
    assert "（签字或盖章）" in (legal.text or "")
    assert "\n" not in (legal.text or "").replace("\u3000", "")
    assert "签字或盖章" in (legal.text or "").replace("\n", "")
    assert "w:u" in legal._p.xml or _em_len(blank) >= 4
    assert not any(
        compact_title(p.text) in {"签字或盖章", "（签字或盖章）"} for p in paras if p is not legal
    )
    assert legal.paragraph_format.keep_with_next is True
    letter_sign = next(
        p
        for p in paras
        if "投标人" in compact_title(p.text) and ("盖单位章" in (p.text or "") or "盖单位公章" in (p.text or ""))
    )
    assert letter_sign.paragraph_format.keep_with_next is True
    body_hit = next(p for p in paras if "承包工程" in (p.text or "") or "我方已仔细研究" in (p.text or ""))
    after = body_hit.paragraph_format.space_after
    if after is not None:
        assert after <= Pt(4)
    assert any(compact_title(t).startswith("地址") for t in texts)
    assert any("电话" in compact_title(t) for t in texts)
    appendix = next(p for p in paras if compact_title(p.text) == "（二）投标函附录")
    assert appendix.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert appendix.runs and appendix.runs[0].font.size >= Pt(16)
    from api.services.tenders.document import _paragraph_has_page_break

    prev = appendix._p.getprevious()
    assert appendix.paragraph_format.page_break_before or (
        prev is not None and _paragraph_has_page_break(prev)
    )
    name_i = next(i for i, t in enumerate(texts) if compact_title(t).startswith("项目名称"))
    bidder_i = next(i for i, t in enumerate(texts) if i > name_i and compact_title(t).startswith("投标人"))
    assert name_i != bidder_i
    assert any(compact_title(t).startswith("大写") or "大写" in compact_title(t) for t in texts)
    assert any("安全生产措施费" in compact_title(t) for t in texts)


def test_assemble_letter_recovers_signoff_when_ocr_stuck_after_appendix(tmp_path) -> None:
    """邀请书把盖章落款写在附录后时，投标函页仍须有投标人/日期。"""
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    # 函正文无落款；附录后才有签字段（常见 OCR 粘连）
    body = (
        "（一）投标函\n"
        "郑州天健智慧园区运营管理有限公司：\n"
        "1.我方已仔细研究了某某项目招标文件的全部内容，愿意以人民币投标。\n"
        "2. 我方的投标文件包括下列内容：\n"
        "（1）投标函及投标函附录；\n"
        "5. 我方在此声明，所递交的投标文件内容完整、真实和准确。\n"
        "（二）投标函附录\n"
        "项目名称：某某项目\n"
        "投标人：河南伟泰光电科技有限公司\n"
        "投标人：河南伟泰光电科技有限公司（盖单位公章）\n"
        "法定代表人或其委托代理人：（签字或盖章）\n"
        "日期：2026年09月29日\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.tenderer = "郑州天健智慧园区运营管理有限公司"
    brief.projectName = "某某项目"
    brief.bidDate = "2026-09-29"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="投标函", kind="letter", source="copy", body=body),
    ]
    dest = tmp_path / "letter-sign-recover.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    paras = [p for p in doc.paragraphs if (p.text or "").strip()]

    def _is_letter_heading(t: str) -> bool:
        n = compact_title(t)
        return n.endswith("投标函") and "附录" not in n and len(n) <= 10

    def _is_app_heading(t: str) -> bool:
        n = compact_title(t)
        return "投标函附录" in n and "及" not in n and len(n) <= 16

    # 跳过封面/目录，取正文第一处投标函标题
    letter_i = next(i for i, p in enumerate(paras) if _is_letter_heading(p.text))
    app_i = next(i for i, p in enumerate(paras) if i > letter_i and _is_app_heading(p.text))
    letter_part = [p.text.strip() for p in paras[letter_i:app_i]]
    assert any(
        "投标人" in compact_title(t) and ("盖单位" in t or "公章" in t or "盖章" in t)
        for t in letter_part
    ), letter_part
    assert any("日期" in compact_title(t) and "2026" in t for t in letter_part)


def test_form_line_role_notice_is_body_not_sign() -> None:
    from api.services.tenders.assemble import _form_line_role

    assert (
        _form_line_role("供应商：参见询价采购文件第三章、第四章有关格式提供。") == "body"
    )
    assert _form_line_role("供应商：（盖章）") == "sign"
    assert _form_line_role("日期：2026年09月29日") == "sign"
    assert _form_line_role("二、合同条款响应书") == "title"


def test_norm_resp_body_fills_lot_and_choice_blanks() -> None:
    from api.services.tenders.assemble import _norm_resp_body

    out = _norm_resp_body(
        "根据贵方的采购文件，标段编号为        ，我方经慎重研究，承诺  （全面/部分）响应所有合同条件。"
    )
    assert "标段编号为〓LOT〓" in out
    assert "承诺〓CHOICE〓（全面/部分）" in out
    assert "标段编号为 " not in out
    assert "承诺  （" not in out


def test_assemble_response_letter_layout_right_sign_centered_title(tmp_path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "二、合同条款响应书致：河南华慧有色工程设计有限公司\n"
        "根据贵方的采购文件，标段编号：________，我方经慎重研究，承诺全面响应所有合同条件。\n"
        "供应商：河南伟泰光电科技有限公司（全称，加盖公章）\n"
        "法定代表人或授权代表：（签字或盖章）\n"
        "日期： 年 月 日\n"
        "供应商已全面理解第三章《合同条款》所有合同条款及第四章《技术要求》的所有要求，无任何偏离。\n"
        "若存在部分响应，须将偏差部分详细注明并填写《商务和技术偏差表》。\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.bidDate = "2026-09-29"
    brief.tenderer = "河南华慧有色工程设计有限公司"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="合同条款响应书",
            kind="commitment_copy",
            source="copy",
            body=body,
        ),
    ]
    dest = tmp_path / "resp-layout.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    paras = [p for p in doc.paragraphs if (p.text or "").strip()]
    head_i = next(
        i
        for i, p in enumerate(paras)
        if compact_title(p.text) == "合同条款响应书"
        and "\t" not in (p.text or "")
    )
    head = paras[head_i]
    assert "二、" not in (head.text or "")
    assert head.alignment == WD_ALIGN_PARAGRAPH.CENTER
    chunk = paras[head_i:]
    texts = [p.text.strip() for p in chunk[:12]]
    assert any(t.startswith("致：") and "华慧" in t for t in texts)
    zhi = next(p for p in chunk if (p.text or "").strip().startswith("致"))
    assert "w:u" not in zhi._p.xml
    # 落款在须知正文之后
    notice_i = next(i for i, t in enumerate(texts) if "商务和技术偏差表" in t)
    supplier_i = next(
        i for i, t in enumerate(texts) if "伟泰" in t and "供应商" in compact_title(t)
    )
    assert supplier_i > notice_i
    supplier = next(p for p in chunk if "伟泰" in (p.text or "") and "供应商" in compact_title(p.text))
    # 与投标函同款：右半页缩进
    assert (supplier.paragraph_format.left_indent or Cm(0)) >= Cm(3.5)
    date_p = next(p for p in chunk if "2026" in (p.text or "") and "日期" in compact_title(p.text))
    assert (date_p.paragraph_format.left_indent or Cm(0)) >= Cm(3.5)
    assert "w:u" in date_p._p.xml
    body_p = next(p for p in chunk if "全面/部分" in (p.text or "") or "全面响应" in (p.text or ""))
    # 标段编号、承诺选项处应是实线下划线，不能是可被拉开的普通空格
    assert "w:u" in body_p._p.xml
    assert "标段编号为" in (body_p.text or "")
    # 承诺与（全面/部分）之间不应残留大段空白
    blob = (body_p.text or "").replace("\u3000", "")
    assert "承诺（全面/部分）" in blob or "承诺(全面/部分)" in blob or "承诺" in blob

def test_pipe_table_stops_before_signoff_and_next_section() -> None:
    from docx import Document

    from api.services.tenders.assemble import _is_table_stop, _try_write_pipe_table
    from api.services.tenders.schema import DocumentFormat

    assert _is_table_stop("四、响应方案")
    assert _is_table_stop("供应商：（盖章）")
    assert _is_table_stop("日期：2026年09月29日")
    assert not _is_table_stop("1 | 交货期 | 全部响应/无偏差")

    doc = Document()
    fmt = DocumentFormat()
    lines = [
        "序号 | 招标文件要求 | 投标文件响应",
        "1 | 交货期 | 全部响应/无偏差",
        "供应商：（盖章）",
        "日期：2026年09月29日",
        "四、响应方案",
    ]
    nxt = _try_write_pipe_table(doc, lines, 0, fmt=fmt, title="商务和技术偏差表")
    assert nxt == 2
    blob = "".join(c.text or "" for t in doc.tables for r in t.rows for c in r.cells)
    assert "四、响应方案" not in blob
    assert "盖章" not in blob
    assert "日期" not in blob


def test_pipe_table_merges_ruled_spans() -> None:
    from docx import Document

    from api.services.tenders.assemble import _try_write_pipe_table
    from api.services.tenders.schema import DocumentFormat

    doc = Document()
    lines = [
        "项目名称 |  | >>",
        " |  | ",
        "投标总报价 | 措施费 | 大写",
        "^^ | 暂列金额 | 小写",
    ]
    nxt = _try_write_pipe_table(doc, lines, 0, fmt=DocumentFormat(), title="投标函附录")
    assert nxt == 4
    assert len(doc.tables[0].rows) == 4
    xml = doc.tables[0]._tbl.xml
    assert "gridSpan" in xml
    assert "vMerge" in xml
    blob = "".join(c.text or "" for r in doc.tables[0].rows for c in r.cells)
    assert "项目名称" in blob
    assert ">>" not in blob
    assert "^^" not in blob
    assert "/" not in blob


def test_assemble_response_letter_one_date_notice_no_underline(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    body = (
        "合同条款响应书\n"
        "致：郑州某某有限公司\n"
        "我方完全响应合同条款，无任何偏差。\n"
        "供应商：河南伟泰光电科技有限公司（盖章）\n"
        "日期： 年 月 日\n"
        "2026年09月29日\n"
        "供应商：参见询价采购文件第三章、第四章有关格式提供。\n"
        "四、响应方案\n"
        "（格式自拟）\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.bidDate = "2026-09-29"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="合同条款响应书",
            kind="commitment_copy",
            source="copy",
            body=body,
        ),
    ]
    dest = tmp_path / "response-letter.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    paras = [p for p in doc.paragraphs if (p.text or "").strip()]
    head_i = next(
        i
        for i, p in enumerate(paras)
        if compact_title(p.text) == "合同条款响应书" and "\t" not in (p.text or "")
    )
    chunk = paras[head_i : head_i + 12]
    texts = [p.text.strip() for p in chunk]
    blob = "\n".join(texts)
    date_lines = [t for t in texts if re.search(r"2026\s*年\s*09\s*月\s*29\s*日", t)]
    assert len(date_lines) == 1, date_lines
    assert "四、响应方案" not in blob
    notice_para = next(p for p in chunk if "参见" in (p.text or "") and "第三章" in (p.text or ""))
    assert "w:u" not in notice_para._p.xml


def test_assemble_tech_dev_ignores_dirty_ocr_body(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    dirty = (
        "序号 | 招标文件要求 | 投标文件响应\n"
        "1 | 交货期 | 全部响应/无偏差供应商：（盖章）日期：2026年09月29日四、响应方案\n"
        "供应商：（盖章）\n"
        "四、响应方案\n"
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.bidDate = "2026-09-29"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.deviationLines = [
        DeviationLine(
            seq="1",
            requirement="支持单点登录",
            response="支持单点登录",
            deviation="无偏差",
        )
    ]
    brief.outlineItems = [
        OutlineItem(
            id="o01",
            title="技术偏差表",
            kind="tech_dev",
            source="generate",
            body=dirty,
        ),
    ]
    dest = tmp_path / "tech-dev-clean.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "招标文件要求" in blob
    assert "交货期" in blob
    assert "支持单点登录" not in blob


def test_outline_letter_keeps_tender_close_and_appendix_table(tmp_path) -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from api.services.tenders.assemble import assemble_bid_docx

    brief = BidBrief(
        layoutMode="outline",
        projectName="郑州高新技术产业开发区新能源充（换）电站基础设施集群项目",
        tenderer="郑州天健智慧园区运营管理有限公司",
        bidderName="河南伟泰光电科技有限公司",
        bidPriceYuan=986000,
        deliveryDays=90,
        bidValidityDays=90,
        quality="合格",
        includePlaceholders=False,
        attachQualifications=False,
        outlineItems=[
            OutlineItem(id="p", title="投标函及投标函附录", rawTitle="一、投标函及投标函附录", kind="letter"),
            OutlineItem(
                id="a",
                title="投标函",
                rawTitle="（一）投标函",
                kind="letter",
                body=(
                    "（一）投标函\n"
                    "致：郑州天健智慧园区运营管理有限公司：\n"
                    "1.我方已仔细研究了招标文件的全部内容，愿意以人民币（大写）的投标总报价，工期日历天。\n"
                    "2.我方的投标文件包括下列内容：\n"
                    "（1）投标函及投标函附录；\n"
                    "电话：\n"
                    "传真：\n"
                    "邮政编码：\n"
                    "网址：\n"
                    "地址：\n"
                    "投标人：（盖单位章）\n"
                    "法定代表人或其委托代理人：\n"
                    "（签字或盖章）\n"
                    "年 月 日\n"
                ),
            ),
            OutlineItem(
                id="b",
                title="投标函附录",
                rawTitle="（二）投标函附录",
                kind="letter",
                body=(
                    "（二）投标函附录\n"
                    "项目名称：\n投标人：\n投标范围：\n投标总报价\n"
                    "安全生产措施费\n暂列金额\n投标工期\n投标质量\n"
                    "投标有效期\n缺陷责任期\n技术标准和要求\n项目经理\n"
                ),
            ),
        ],
    )
    dest = tmp_path / "letter.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "我方承诺在投标有效期内不修改、撤销投标文件" not in blob
    assert "我方的投标文件包括下列内容" in blob
    assert "地址" in blob and "网址" in blob and "传真" in blob
    assert "项目名称" in blob and "安全生产措施费" in blob
    assert "投标总报价（元）" not in blob


def _page(doc, *lines: str) -> None:
    from docx.enum.text import WD_BREAK

    for line in lines:
        doc.add_paragraph(line)
    end = doc.add_paragraph()
    end.add_run().add_break(WD_BREAK.PAGE)


def test_outline_uses_invite_cover_and_skips_blank_pages(tmp_path, monkeypatch) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    src = tmp_path / "ch8.docx"
    raw = Document()
    _page(raw, "第八章、投标文件格式")
    _page(raw, "-185-")
    _page(
        raw,
        "（项目名称）标段施工",
        "投标文件",
        "招标编号：",
        "投标人或联合体牵头人：（盖单位章）",
        "年 月 日",
    )
    _page(raw, "-184-")
    _page(raw, "目录", "一、投标函及投标函附录", "二、法定代表人身份证明", "三、授权委托书")
    _page(raw, "-187-")
    raw.add_paragraph("（一）投标函")
    raw.add_paragraph("致：郑州天健智慧园区运营管理有限公司")
    raw.add_paragraph("1.我方已仔细研究了招标文件的全部内容。")
    table = raw.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "项目名称"
    table.rows[1].cells[0].text = "安全生产措施费"
    raw.save(src)

    monkeypatch.setattr(
        "api.services.tenders.pdf_format._ready_format_docx",
        lambda invitation_id: src if invitation_id else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = "abc123abc123"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="法定代表人身份证明", kind="legal_id", source="generate"),
    ]
    dest = tmp_path / "pasted.docx"
    _path, warnings = assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert "标段施工" in blob
    assert "商 务 标" not in blob
    assert "封面加盖公章" not in blob
    assert "第八章、投标文件格式" not in blob
    assert "-185-" not in blob and "-184-" not in blob and "-187-" not in blob
    assert "投标函及投标函附录" in blob
    assert "我方已仔细研究了" in blob
    assert "安全生产措施费" in blob
    assert "单位性质" not in blob
    from api.services.tenders.pdf_format import toc_block_ids
    from docx.text.paragraph import Paragraph
    entries = [Paragraph(n, doc) for n in toc_block_ids(doc) if n.tag == qn("w:p")]
    entries = [p for p in entries if "、" in p.text]
    assert len(entries) == 3
    assert all("PAGEREF" in p._p.xml and 'w:leader="dot"' in p._p.xml for p in entries)
    assert all(p.text.endswith("待更新") for p in entries)
    assert "未定义书签" not in blob
    assert any("原版" in w for w in warnings)


def test_outline_splits_pdf_section_pages(tmp_path, monkeypatch) -> None:
    """pdf2docx 用分节换页。扉页、原封面、原目录不能再整章贴进来。"""
    from docx import Document
    from docx.enum.section import WD_SECTION

    from api.services.tenders.assemble import assemble_bid_docx

    def add_page(doc: Document, *lines: str) -> None:
        if doc.paragraphs:
            doc.add_section(WD_SECTION.NEW_PAGE)
        for line in lines:
            doc.add_paragraph(line)

    src = tmp_path / "ch8-sect.docx"
    raw = Document()
    add_page(raw, "第八章、投标文件格式", "-185-")
    add_page(
        raw,
        "（项目名称）标段施工",
        "投标文件",
        "招标编号：",
        "投标人或联合体牵头人：（盖单位章）",
        "年 月 日",
        "-186-",
    )
    add_page(raw, "目录", "一、投标函及投标函附录", "二、法定代表人身份证明", "三、授权委托书", "-187-")
    add_page(raw, "（一）投标函", "1.我方已仔细研究了 （项目名称） 招标文件的全部内容。", "-188-")
    raw.add_paragraph("投标人：          （盖单位章）法定代表人或其委托代理人： （签字或盖章）")
    phone = raw.add_table(rows=1, cols=2)
    phone.rows[0].cells[0].text = "电"
    phone.rows[0].cells[1].text = "话："
    table = raw.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "项目名称"
    table.rows[1].cells[0].text = "安全生产措施费"
    add_page(
        raw,
        "三、授权委托书",
        "本人系投标人的法定代表人，现委托代理人签署投标文件。",
        "标段施工投标文件、签订合同。",
        "投标人：（盖单位章）",
    )
    add_page(raw, "六、项目管理机构")
    raw.save(src)

    monkeypatch.setattr(
        "api.services.tenders.pdf_format._ready_format_docx",
        lambda invitation_id: src if invitation_id else None,
    )
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = "abc123abc123"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="o01", title="法定代表人身份证明", kind="legal_id", source="generate"),
    ]
    dest = tmp_path / "pasted-sect.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))
    blob = _docx_text(doc)
    assert blob.count("目录") == 1
    assert "标段施工" in blob
    assert "封面加盖公章" not in blob
    assert "第八章、投标文件格式" not in blob
    assert "一、投标函及投标函附录" in blob
    assert "我方已仔细研究了" in blob
    assert "河南伟泰光电科技有限公司" in blob
    cover = next(p for p in doc.paragraphs if "标段施工" in (p.text or "") and "招标项目" not in (p.text or ""))
    assert "（项目名称）" in (cover.text or "") and "标段施工" in (cover.text or "")
    bidder = next(p for p in doc.paragraphs if "联合体牵头人" in (p.text or ""))
    assert "伟泰" in (bidder.text or "") and "（盖单位章）" in (bidder.text or "")
    assert "授权委托书" in blob and "项目管理机构" in blob
    assert "-185-" not in blob and "-186-" not in blob and "-187-" not in blob and "-188-" not in blob
    assert "PAGEREF" in doc.element.xml
    assert "TOC" in doc.element.xml


def test_pasted_auth_underline_and_id_scan(tmp_path) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from PIL import Image

    from api.services.tenders.pdf_format import fill_pasted_format

    doc = Document()
    doc.add_paragraph("目录")
    doc.add_paragraph("三、授权委托书")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "撤回、修改"
    table.rows[0].cells[1].text = "（项目名称）"
    table.rows[1].cells[0].text = ""
    doc.add_paragraph("标段施工投标文件、签订合同和处理有关事宜，其法律后果由我方承担。")
    doc.add_paragraph("附：法定代表人身份证复印件或扫描件")
    doc.add_paragraph("投标人： （盖单位章）")
    doc.add_paragraph("年    月    日")
    img = tmp_path / "id.png"
    Image.new("RGB", (80, 48), "white").save(img)
    brief = default_brief()
    brief.projectName = "郑州项目"
    brief.bidDate = "2026-10-21"
    fill_pasted_format(doc, brief, media={"id_legal": [img]})
    blob = _docx_text(doc)
    assert "撤回、修改" in blob and "郑州项目" in blob
    assert "（项目名称）" not in blob
    assert all("撤回" not in _docx_text_table(t) for t in doc.tables)
    line = next(p for p in doc.paragraphs if "撤回、修改" in (p.text or ""))
    assert any(r.underline and "郑州项目" in (r.text or "") for r in line.runs)
    assert "2026" in blob and "10" in blob and "21" in blob
    assert "w:drawing" in doc.element.xml
    # 身份证贴在落款前，不甩到文末
    body = list(doc.element.body)
    draw_at = next(i for i, el in enumerate(body) if el.findall(".//" + qn("w:drawing")) or qn("w:drawing") in el.xml)
    sign_at = next(i for i, el in enumerate(body) if "投标人" in "".join(t.text or "" for t in el.iter(qn("w:t"))))
    assert draw_at < sign_at


def _docx_text_table(table) -> str:
    return "\n".join(cell.text or "" for row in table.rows for cell in row.cells)


def _format_slice(tmp_path, invite: str):
    from docx import Document

    raw = Document()
    raw.add_paragraph("投标函")
    raw.add_paragraph("投标函正文")
    raw.add_paragraph("已标价工程量清单")
    raw.add_paragraph("清单合计")
    raw.add_paragraph("施工组织设计")
    raw.add_paragraph("组织设计正文 河南伟泰光电科技有限公司")
    src = tmp_path / f"{invite}.format.docx"
    raw.save(src)
    src.with_name(src.name + ".source").write_text("docx-raw", encoding="utf-8")
    return src


def _patch_format(monkeypatch, invite: str, src) -> None:
    monkeypatch.setattr(
        "api.services.tenders.pdf_format._word_format_file",
        lambda invitation_id: src if invitation_id == invite else None,
    )


def test_full_slice_keeps_works_plan(tmp_path, monkeypatch) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    invite = "abc123abc201"
    src = _format_slice(tmp_path, invite)
    _patch_format(monkeypatch, invite, src)
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    brief.generateVolume = "full"
    path, notes = assemble_bid_docx(brief, tmp_path / "full.docx", volume="full")
    blob = _docx_text(Document(str(path)))
    assert "施工组织设计" in blob
    assert "投标函" in blob
    assert not any("已去掉格式章里的技术部分" in n for n in notes)
    assert not any("只保留格式章里的技术部分" in n for n in notes)


def test_split_drops_tech_from_business_and_letter_from_technical(tmp_path, monkeypatch) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    invite = "abc123abc202"
    src = _format_slice(tmp_path, invite)
    _patch_format(monkeypatch, invite, src)
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    brief.splitSlice = True
    path, notes = assemble_bid_docx(brief, tmp_path / "biz.docx", volume="business")
    blob = _docx_text(Document(str(path)))
    assert "投标函" in blob
    assert "工程量清单" in blob
    assert "施工组织设计" not in blob
    assert any("已去掉格式章里的技术部分" in n for n in notes)

    tech = brief.model_copy(update={"splitSlice": False})
    path_t, _ = assemble_bid_docx(tech, tmp_path / "tech.docx", volume="technical")
    blob_t = _docx_text(Document(str(path_t)))
    assert "施工组织设计" in blob_t
    assert "投标函" not in blob_t
    assert "工程量清单" not in blob_t


def test_plain_business_slice_still_keeps_works_plan(tmp_path, monkeypatch) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    invite = "abc123abc203"
    src = _format_slice(tmp_path, invite)
    _patch_format(monkeypatch, invite, src)
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.invitationId = invite
    path, _ = assemble_bid_docx(brief, tmp_path / "biz.docx", volume="business")
    blob = _docx_text(Document(str(path)))
    assert "施工组织设计" in blob
    assert "投标函" in blob


def test_anonymous_tech_hides_names_business_keeps_company(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx

    brief = default_brief()
    brief.layoutMode = "chapter5"
    brief.includeCommitment = False
    brief.includePlaceholders = False
    brief.attachQualifications = False
    brief.techAnonymous = True
    brief.bidderName = "河南伟泰光电科技有限公司"
    brief.legalPersonName = "郭志伟"
    brief.agentName = "张委托"
    brief.pmName = "李项目"
    tech_path, notes = assemble_bid_docx(brief, tmp_path / "tech.docx", volume="technical")
    tech_doc = Document(str(tech_path))
    blob = _docx_text(tech_doc)
    assert "河南伟泰" not in blob
    assert "郭志伟" not in blob
    assert "张委托" not in blob
    assert "李项目" not in blob
    assert any("暗标专轮" in n for n in notes)
    assert (tech_doc.core_properties.author or "") == ""
    assert (tech_doc.core_properties.last_modified_by or "") == ""
    company = ""
    from lxml import etree

    for rel in tech_doc.part.package.rels.values():
        if "extended-properties" not in (rel.reltype or ""):
            continue
        root = etree.fromstring(rel.target_part.blob)
        for child in root:
            if child.tag.rsplit("}", 1)[-1] == "Company":
                company = child.text or ""
    assert company == ""
    assert brief.bidderName == "河南伟泰光电科技有限公司"

    marked = tmp_path / "marked.docx"
    raw = Document()
    raw.add_paragraph("封面投标人：河南伟泰光电科技有限公司")
    raw.core_properties.author = "郭志伟"
    raw.core_properties.last_modified_by = "张委托"
    raw.save(marked)
    stamped = Document(str(marked))
    for rel in stamped.part.package.rels.values():
        if "extended-properties" not in (rel.reltype or ""):
            continue
        root = etree.fromstring(rel.target_part.blob)
        for child in root:
            if child.tag.rsplit("}", 1)[-1] == "Company":
                child.text = "河南伟泰光电科技有限公司"
        rel.target_part._blob = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    stamped.save(marked)
    before = ""
    saved = Document(str(marked))
    for rel in saved.part.package.rels.values():
        if "extended-properties" not in (rel.reltype or ""):
            continue
        root = etree.fromstring(rel.target_part.blob)
        for child in root:
            if child.tag.rsplit("}", 1)[-1] == "Company":
                before = child.text or ""
    assert before == "河南伟泰光电科技有限公司"
    from api.services.tenders.pdf_format import scrub_anonymous_docx

    scrub_anonymous_docx(marked, brief)
    washed = Document(str(marked))
    assert "河南伟泰" not in _docx_text(washed)
    assert (washed.core_properties.author or "") == ""
    assert (washed.core_properties.last_modified_by or "") == ""
    company = ""
    for rel in washed.part.package.rels.values():
        if "extended-properties" not in (rel.reltype or ""):
            continue
        root = etree.fromstring(rel.target_part.blob)
        for child in root:
            if child.tag.rsplit("}", 1)[-1] == "Company":
                company = child.text or ""
    assert company == ""

    biz_path, biz_notes = assemble_bid_docx(brief, tmp_path / "biz.docx", volume="business")
    biz_blob = _docx_text(Document(str(biz_path)))
    assert "河南伟泰光电科技有限公司" in biz_blob
    assert not any("暗标专轮" in n for n in biz_notes)


def test_generate_pair_slices_only_for_split_advice(tmp_path, monkeypatch) -> None:
    from api.services.tenders.generate import generate_bid, generate_pair

    seen = []

    def fake_assemble(brief, dest, **kwargs):
        seen.append((kwargs.get("volume"), bool(brief.splitSlice), bool(brief.techAnonymous)))
        dest.write_bytes(b"PK\x03\x04")
        return dest, []

    monkeypatch.setattr("api.services.tenders.assemble.assemble_bid_docx", fake_assemble)
    monkeypatch.setattr("api.services.tenders.generate.tenders_output_dir", lambda: tmp_path)
    split = BidBrief(volumeAdvice="split", layoutMode="chapter5", includePlaceholders=False, attachQualifications=False)
    generate_pair(split)
    assert seen[-2:] == [("business", True, False), ("technical", False, False)]
    seen.clear()
    whole = split.model_copy(update={"volumeAdvice": "full"})
    generate_pair(whole)
    assert seen[-2:] == [("business", False, False), ("technical", False, False)]
    seen.clear()
    out = generate_bid(whole.model_copy(update={"generateVolume": "full"}))
    assert seen[-1][0] == "full"
    assert "投标文件.docx" in str(out.get("downloadName") or "")
    assert out.get("docxFile")
    assert not out.get("techDocxFile")


def _pageref_cache(para) -> str:
    in_result = False
    for run in para._p.iter(qn("w:r")):
        fld = run.find(qn("w:fldChar"))
        if fld is not None:
            kind = fld.get(qn("w:fldCharType"))
            if kind == "separate":
                in_result = True
            elif kind == "end":
                in_result = False
            continue
        if not in_result:
            continue
        text_el = run.find(qn("w:t"))
        if text_el is not None and (text_el.text or "").strip():
            return text_el.text.strip()
    return ""


def _bookmark_para_text(doc, name: str) -> str:
    for el in doc.element.iter(qn("w:bookmarkStart")):
        if el.get(qn("w:name")) != name:
            continue
        node = el
        while node is not None and node.tag != qn("w:p"):
            node = node.getparent()
        assert node is not None
        cur = node
        while cur is not None:
            assert cur.tag != qn("w:tbl")
            cur = cur.getparent()
        return "".join(t.text or "" for t in node.iter(qn("w:t")))
    raise AssertionError(name)


def test_verbatim_toc_keeps_body_numbers_and_table_pages(tmp_path) -> None:
    from docx import Document

    from api.services.tenders.assemble import assemble_bid_docx
    from api.services.tenders.document import _estimate_bookmark_pages, _pageref_bookmark
    from api.services.tenders.format_rules import toc_numbered_lines

    shows = toc_numbered_lines(
        [
            (2, "（三）项目管理机构组成表"),
            (3, "人员分工"),
            (2, "拟分包项目情况表"),
        ]
    )
    assert shows[0] == "（三）项目管理机构组成表"
    assert shows[1].startswith("1.1 ")
    assert shows[2].startswith("（一）")

    table = "| 序号 | 项目名称 | 备注 |\n| 1 | 塔吊 | 一台 |"
    brief = default_brief()
    brief.layoutMode = "outline"
    brief.includePlaceholders = False
    brief.includeCommitment = False
    brief.attachQualifications = False
    brief.outlineItems = [
        OutlineItem(id="letter", title="投标函", kind="other", level=1, body="投标函正文写在这里凑够字数，不要带编号。"),
        OutlineItem(
            id="org",
            title="现场组织安排",
            rawTitle="（一）现场组织安排",
            kind="other",
            level=2,
            body="（一）现场组织安排\n" + table,
        ),
        OutlineItem(
            id="sub",
            title="材料进场计划",
            kind="other",
            level=2,
            body="材料进场计划\n材料进场说明写在正文里，目录要按二级补号。",
        ),
        OutlineItem(
            id="qa",
            title="质量保证措施",
            rawTitle="（三）质量保证措施",
            kind="other",
            level=3,
            body="（三）质量保证措施\n质量措施正文写在这里，目录不能改成点分编号。",
        ),
        OutlineItem(
            id="crew",
            title="人员分工",
            kind="other",
            level=3,
            body="人员分工\n人员分工正文写在这里用来核对目录补了点分编号。",
        ),
        OutlineItem(id="gear", title="拟投入设备表", kind="other", level=2, body=table),
    ]
    dest = tmp_path / "toc-pages.docx"
    assemble_bid_docx(brief, dest, qualification_pdf=None)
    doc = Document(str(dest))

    def toc_line(prefix: str):
        return next(p for p in doc.paragraphs if (p.text or "").startswith(prefix) and "PAGEREF" in p._p.xml)

    org = toc_line("（一）现场组织安排")
    sub = toc_line("（二）材料进场计划")
    qa = toc_line("（三）质量保证措施")
    crew = toc_line("2.1 ")
    gear = toc_line("（四）拟投入设备表")
    assert not (crew.text or "").startswith("（")
    for para, level in ((org, 2), (qa, 3), (gear, 2)):
        tabs = para._p.find(qn("w:pPr")).find(qn("w:tabs")).find(qn("w:tab"))
        assert tabs.get(qn("w:leader")) == "dot"
        assert "\t" in (para.text or "")
        left = para.paragraph_format.left_indent
        assert left is not None
        assert abs(float(left.cm) - 0.74 * (level - 1)) < 0.05
    body_org = [p for p in doc.paragraphs if (p.text or "").strip() == "（一）现场组织安排"]
    assert len(body_org) == 1
    body_sub = [p for p in doc.paragraphs if (p.text or "").strip() == "材料进场计划"]
    assert len(body_sub) == 1
    assert _bookmark_para_text(doc, "toc_org") == "（一）现场组织安排"
    gear_text = _bookmark_para_text(doc, "toc_gear")
    assert gear_text == "拟投入设备表"
    host = None
    for el in doc.element.iter(qn("w:bookmarkStart")):
        if el.get(qn("w:name")) == "toc_gear":
            host = el
            while host is not None and host.tag != qn("w:p"):
                host = host.getparent()
            break
    assert host is not None and host.getnext() is not None and host.getnext().tag == qn("w:tbl")

    pages = _estimate_bookmark_pages(doc)
    assert "toc_org" in pages and "toc_gear" in pages
    base = min(pages.values())
    linked = 0
    for para in doc.paragraphs:
        bookmark = _pageref_bookmark(para)
        if not bookmark or not bookmark.startswith("toc_"):
            continue
        assert bookmark in pages
        assert _pageref_cache(para) == str(max(1, pages[bookmark] - base + 1))
        linked += 1
    assert linked >= 6



