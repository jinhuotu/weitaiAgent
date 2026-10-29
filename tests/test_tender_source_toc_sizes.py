from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from api.services.tenders.document import _pageref_bookmark
from api.services.tenders.native_format import ensure_toc_rows
from api.services.tenders.pdf_format import _attach_library
from api.services.tenders.placeholders import _insert_slot_media
from api.services.tenders.schema import BidBrief, OutlineItem

SOURCE = [
    ("一、投标函及投标函附录", 1),
    ("二、法定代表人身份证明", 1),
    ("三、授权委托书", 1),
    ("四、投标保证金", 1),
    ("五、已标价工程量清单", 1),
    ("六、项目管理机构", 1),
    ("（一）项目管理机构组成表", 2),
    ("（二）项目经理简历表", 2),
    ("1.项目经理主要工作经历", 3),
    ("（三）承诺书", 2),
    ("（四）主要项目管理人员简历表", 2),
    ("七、拟分包项目情况表", 1),
    ("八、资格审查资料", 1),
    ("（一）投标人基本情况表", 2),
    ("（二）近年财务状况表", 2),
    ("（三）近年完成的类似项目情况表", 2),
    ("（四）正在施工和新承接项目情况表", 2),
    ("（五）“国家企业信用信息公示系统”查询", 2),
    ("（六）资格审查其他材料", 2),
    ("九、其他资料", 1),
    ("十、施工组织格式", 1),
]


def test_priced_bill_alias_matches_without_changing_display_title():
    from api.services.tenders.native_format import _key

    assert _key("五、已标价的工程量清单") == _key("已标价工程量清单")


@pytest.mark.parametrize("volume", ["business", "technical", "full"])
def test_immutable_original_toc_is_the_only_source(tmp_path, monkeypatch, volume):
    original = Document()
    original.add_paragraph("目录")
    for title, level in SOURCE:
        p = original.add_paragraph(title)
        p.paragraph_format.left_indent = (level - 1) * 400 * 635
        p.add_run().add_tab()
    original.add_page_break()
    for title, _ in SOURCE:
        original.add_paragraph(title, style="Heading 1")
        original.add_paragraph("原版正文。")
    source = tmp_path / "source.docx"
    original.save(source)
    before = source.read_bytes()
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: source)

    output = Document()
    # A polluted output TOC must not override the original source.
    output.add_paragraph("目录")
    for title in (
        "（一）投标函",
        "（二）投标函附录",
        "五、已标价工程量清单",
        "五、已标价工程量清单",
    ):
        output.add_paragraph(title).add_run().add_tab()
    output.add_page_break()
    wanted = (
        SOURCE[:-1] if volume == "business" else (SOURCE[-1:] if volume == "technical" else SOURCE)
    )
    targets = []
    for title, _ in wanted:
        if "投标函及投标函附录" in title:
            heading = output.add_paragraph("（一）投标函", style="Heading 1")
            output.add_paragraph("函件正文。")
            output.add_paragraph("（二）投标函附录", style="Heading 2")
        else:
            heading = output.add_paragraph(title, style="Heading 1")
        targets.append(heading._p)
        output.add_paragraph("原版正文。")
    output.add_paragraph("十一、新识别的非原版章节", style="Heading 1")
    brief = BidBrief(
        invitationId="abcdefgh1234",
        layoutMode="outline",
        outlineChapter="第八章 投标文件格式",
        outlineItems=[OutlineItem(id="quote", title="已标价工程量清单", kind="quote", level=4)],
    )
    ensure_toc_rows(output, brief)
    rows = [p for p in output.paragraphs if _pageref_bookmark(p)]
    assert [p.text.split("\t")[0] for p in rows] == [title for title, _ in wanted]
    for row, (_, level), target in zip(rows, wanted, targets, strict=True):
        assert any(
            n.get(qn("w:name")) == _pageref_bookmark(row)
            for n in target.iter(qn("w:bookmarkStart"))
        )
        tc = next(target.iter(qn("w:fldSimple")))
        assert f"\\l {level} " in tc.get(qn("w:instr"))
    once = output.element.xml
    ensure_toc_rows(output, brief)
    assert output.element.xml == once
    assert source.read_bytes() == before
    assert sum(p.text == "目录" for p in output.paragraphs) == 1


def test_identity_pair_is_one_small_horizontal_strip(tmp_path):
    files = []
    for name, color in (("front", "red"), ("back", "green")):
        path = tmp_path / (name + ".png")
        Image.new("RGB", (860, 540), color).save(path)
        files.append(path)
    doc = Document()
    doc.add_paragraph("附：法定代表人身份证正反面复印件")
    _attach_library(doc, BidBrief(), {"id_legal": files})
    extents = list(doc.element.iter(qn("wp:extent")))
    assert len(extents) == 1
    assert int(extents[0].get("cy")) / 360000 <= 5.4 + 0.01
    assert 14 < int(extents[0].get("cx")) / 360000 <= 17.4
    rid = next(doc.element.iter(qn("a:blip"))).get(qn("r:embed"))
    with Image.open(BytesIO(doc.part.related_parts[rid].blob)) as image:
        assert image.getpixel((image.width // 4, image.height // 2))[0] > 180
        assert image.getpixel((image.width * 3 // 4, image.height // 2))[1] > 80


def test_converted_source_keeps_multiline_toc_and_shared_body_anchors(tmp_path, monkeypatch):
    from api.services.tenders.native_format import _body_toc_block, _read_toc_entries

    titles = [
        title.replace("1.项目经理主要工作经历", "（三）项目经理主要工作经历") for title, _ in SOURCE
    ]
    doc = Document()
    doc.add_paragraph("目录")
    for title in titles[:5]:
        doc.add_paragraph(title)
    # Conversion merged the source directory's subitems into long paragraphs.
    doc.add_paragraph("\n".join(titles[5:12]))
    doc.add_paragraph("\n".join(titles[12:20]))
    doc.add_paragraph(titles[20])
    doc.add_paragraph("-187-")
    doc.add_page_break()
    first = doc.add_paragraph(titles[0] + "\t（一）投标函\n（招标人名称）：")
    doc.add_paragraph("函件原版正文。")
    # Two original headings can also share one physical paragraph in the body.
    shared = doc.add_paragraph(titles[1] + "\n" + titles[2])
    doc.add_paragraph("原版正文。")
    for title in titles[3:]:
        doc.add_paragraph(title)
        doc.add_paragraph("原版正文。")
    path = tmp_path / "converted.docx"
    doc.save(path)
    monkeypatch.setattr("api.services.tenders.pdf_format._word_format_file", lambda _: path)
    brief = BidBrief(
        invitationId="converted", layoutMode="outline", outlineChapter="第八章 投标文件格式"
    )
    source_rows = _read_toc_entries(_body_toc_block(doc, brief))
    assert [row["title"] for row in source_rows] == titles
    ensure_toc_rows(doc, brief)
    rows = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert [p.text.split("\t")[0] for p in rows] == titles
    assert first._p.getparent() is not None
    assert (
        len(
            [
                field
                for field in shared._p.iter(qn("w:fldSimple"))
                if "\\f N " in field.get(qn("w:instr"), "")
            ]
        )
        == 2
    )
    once = doc.element.xml
    ensure_toc_rows(doc, brief)
    assert doc.element.xml == once


@pytest.mark.parametrize(
    "key", ["license", "finance", "social", "perf", "bank_permit", "iso", "product"]
)
def test_scan_height_is_half_page_in_both_renderers(tmp_path, key):
    path = tmp_path / (key + ".png")
    Image.new("RGB", (210, 297), "white").save(path)
    doc = Document()
    _insert_slot_media(doc, [path])
    assert all(int(e.get("cy")) / 360000 <= 12.01 for e in doc.element.iter(qn("wp:extent")))
    from api.services.tenders.scan_placement import _pictures

    doc = Document()
    anchor = doc.add_paragraph("原位占位")
    _pictures(doc, anchor._p, [path], key)
    assert all(int(e.get("cy")) / 360000 <= 12.01 for e in doc.element.iter(qn("wp:extent")))


def test_pdf_pages_do_not_force_a_page_each(tmp_path):
    import fitz

    path = tmp_path / "scan.pdf"
    pdf = fitz.open()
    for _ in range(4):
        pdf.new_page()
    pdf.save(path)
    pdf.close()
    doc = Document()
    _insert_slot_media(doc, [path], max_pages=4)
    assert len(list(doc.element.iter(qn("w:drawing")))) == 4
    assert len([n for n in doc.element.iter(qn("w:br")) if n.get(qn("w:type")) == "page"]) < 3
