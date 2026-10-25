"""Legacy Word conversion lifecycle; tests never launch Microsoft Word."""

from __future__ import annotations

import io
import subprocess
import sys
import time
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from fastapi import UploadFile

from api.services import word_convert as converter
from api.services.knowledge import parsers
from api.services.tenders import extract
from common.errors import AppError


def _docx(dest):
    doc = Document()
    doc.add_paragraph("天津市静海县环境监控与应急指挥中心建设项目")
    doc.save(dest)


def test_libreoffice_validates_and_publishes_without_modifying_source(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1"))
    original = source.read_bytes()
    dest = tmp_path / "result.docx"
    monkeypatch.setattr(converter.shutil, "which", lambda _: "soffice")

    def run(args, timeout):
        assert timeout > 0
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in args)
        folder = Path(args[args.index("--outdir") + 1])
        _docx(folder / "legacy.docx")
        return subprocess.CompletedProcess(args, 0, "", "")

    monkeypatch.setattr(converter, "_run_converter", run)
    assert converter.convert_doc_to_docx(source, dest) == dest.resolve()
    assert Document(dest).paragraphs[0].text.startswith("天津")
    assert source.read_bytes() == original
    assert not list(tmp_path.glob("word-convert-*"))


def test_bad_libreoffice_result_falls_back_to_isolated_word_worker(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy")
    monkeypatch.setattr(converter.shutil, "which", lambda _: "soffice")
    monkeypatch.setattr(converter, "_word_com_supported", lambda: True)
    calls = []

    def run(args, timeout):
        calls.append(args)
        if "--word-com" not in args:
            folder = Path(args[args.index("--outdir") + 1])
            (folder / "legacy.docx").write_bytes(b"PK invalid package")
        else:
            assert args[0] == sys.executable
            _docx(Path(args[-2]))
        return subprocess.CompletedProcess(args, 0, "", "")

    monkeypatch.setattr(converter, "_run_converter", run)
    converter.convert_doc_to_docx(source, tmp_path / "out.docx")
    assert len(calls) == 2


def test_worker_failure_cleans_temporary_files_and_preserves_existing_output(tmp_path, monkeypatch):
    source, dest = tmp_path / "legacy.doc", tmp_path / "existing.docx"
    source.write_bytes(b"legacy")
    dest.write_bytes(b"previous user output")
    monkeypatch.setattr(converter.shutil, "which", lambda _: None)
    monkeypatch.setattr(converter, "_word_com_supported", lambda: True)
    stopped = Mock()
    monkeypatch.setattr(converter, "_stop_owned_word", stopped)

    def run(args, timeout):
        Path(args[-2]).write_bytes(b"partial")
        Path(args[-1]).write_text("12345", encoding="ascii")
        raise subprocess.TimeoutExpired(args, timeout)

    monkeypatch.setattr(converter, "_run_converter", run)
    with pytest.raises(converter.WordConversionError, match="另存为 .docx"):
        converter.convert_doc_to_docx(source, dest)
    stopped.assert_called_once()
    assert dest.read_bytes() == b"previous user output"
    assert source.read_bytes() == b"legacy"
    assert not list(tmp_path.glob("word-convert-*"))


def test_no_converter_has_clear_error_and_no_published_output(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy")
    monkeypatch.setattr(converter.shutil, "which", lambda _: None)
    monkeypatch.setattr(converter, "_word_com_supported", lambda: False)
    with pytest.raises(converter.WordConversionError, match="Word 里另存为"):
        converter.convert_doc_to_docx(source, tmp_path / "out.docx")
    assert not (tmp_path / "out.docx").exists()


@pytest.mark.parametrize("contents", [b"", b"not a word document", b"PK invalid ZIP"])
def test_docx_validation_does_not_accept_signatures_alone(tmp_path, contents):
    path = tmp_path / "out.docx"
    path.write_bytes(contents)
    with pytest.raises(converter.WordConversionError):
        converter._validate_docx(path)


@pytest.mark.parametrize("failure", [None, "open", "save", "close", "quit"])
def test_com_initialization_and_release_even_on_failure(tmp_path, monkeypatch, failure):
    events = []

    def event(name):
        def call(*args, **kwargs):
            events.append(name)
            if name == failure:
                raise RuntimeError(name)

        return call

    document = SimpleNamespace(SaveAs2=event("save"), Close=event("close"))

    def open_document(*args, **kwargs):
        event("open")()
        assert kwargs["ReadOnly"] is True
        assert kwargs["AddToRecentFiles"] is False
        return document

    word = SimpleNamespace(Documents=SimpleNamespace(Open=open_document), Quit=event("quit"))
    client = SimpleNamespace(DispatchEx=lambda _: word)
    monkeypatch.setitem(
        sys.modules,
        "pythoncom",
        SimpleNamespace(
            CoInitialize=event("initialize"),
            CoUninitialize=event("uninitialize"),
        ),
    )
    monkeypatch.setitem(sys.modules, "win32com", SimpleNamespace(client=client))
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setitem(
        sys.modules,
        "win32process",
        SimpleNamespace(
            EnumProcesses=lambda: [100],
        ),
    )
    monkeypatch.setattr(converter, "_new_office_process", lambda _: 12345)
    args = (tmp_path / "source.doc", tmp_path / "out.docx", tmp_path / "word.pid")
    if failure:
        with pytest.raises(RuntimeError):
            converter._word_com_convert(*args)
    else:
        converter._word_com_convert(*args)
    assert events[0] == "initialize"
    assert events[-1] == "uninitialize"
    assert "quit" in events
    if failure != "open":
        assert "close" in events


def test_converter_timeout_is_bounded_for_a_real_child_process():
    start = time.monotonic()
    with pytest.raises(subprocess.TimeoutExpired):
        converter._run_converter([sys.executable, "-c", "import time; time.sleep(30)"], 0.2)
    assert time.monotonic() - start < 10


@pytest.mark.parametrize(
    "images, expected",
    [
        ({200: "WINWORD.EXE", 300: "python.exe"}, 200),
        ({200: "WINWORD.EXE", 300: "WINWORD.EXE"}, None),
        ({200: "python.exe", 300: "python.exe"}, None),
    ],
)
def test_office_process_detection_excludes_existing_sessions(monkeypatch, images, expected):
    handles = {
        pid: SimpleNamespace(image="/office/" + name, Close=Mock()) for pid, name in images.items()
    }
    monkeypatch.setitem(
        sys.modules,
        "win32api",
        SimpleNamespace(
            OpenProcess=lambda access, inherit, pid: handles[pid],
        ),
    )
    monkeypatch.setitem(
        sys.modules,
        "win32process",
        SimpleNamespace(
            EnumProcesses=lambda: [100, 200, 300],
            GetModuleFileNameEx=lambda handle, unused: handle.image,
        ),
    )
    assert converter._new_office_process({100}) == expected
    assert all(handle.Close.call_count == 1 for handle in handles.values())


@pytest.mark.asyncio
async def test_doc_parser_converts_once_reads_docx_and_cleans_temp(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1"))
    outputs = []

    def convert(src, dest):
        assert src == source
        outputs.append(dest)
        _docx(dest)

    monkeypatch.setattr(parsers, "convert_doc_to_docx", convert)
    result = await parsers.extract_text_from_file(source)
    assert "静海" in result.text
    assert len(outputs) == 1
    assert not outputs[0].exists()
    assert source.exists()
    assert parsers.assert_supported(".DOC") == "doc"
    assert parsers.sniff_extension(None, "application/msword") == "doc"


@pytest.mark.asyncio
async def test_doc_parser_conversion_failure_has_actionable_message(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy")
    monkeypatch.setattr(
        parsers,
        "convert_doc_to_docx",
        Mock(
            side_effect=converter.WordConversionError("RPC failed"),
        ),
    )
    with pytest.raises(AppError, match="另存为 .docx") as error:
        await parsers.extract_text_from_file(source)
    assert error.value.status_code == 422
    assert source.read_bytes() == b"legacy"


@pytest.mark.asyncio
async def test_tender_doc_upload_reuses_converted_file(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "tender_invitations_dir", lambda: tmp_path)
    convert = Mock(side_effect=lambda src, dest: _docx(dest))
    monkeypatch.setattr(extract, "convert_doc_to_docx", convert)
    upload = UploadFile(filename="招标文件 - 天津静海.doc", file=io.BytesIO(b"legacy"))
    path = await extract._save_invitation(upload)
    assert path.suffix == ".docx"
    result = await parsers.extract_text_from_file(path)
    assert "静海" in result.text
    convert.assert_called_once()
    assert path.exists()
    assert not list(tmp_path.glob("*.doc"))


@pytest.mark.asyncio
async def test_tender_doc_failure_removes_partial_uploads(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "tender_invitations_dir", lambda: tmp_path)

    def fail(src, dest):
        dest.write_bytes(b"partial")
        raise converter.WordConversionError("RPC failed")

    monkeypatch.setattr(extract, "convert_doc_to_docx", fail)
    upload = UploadFile(filename="招标.doc", file=io.BytesIO(b"legacy"))
    with pytest.raises(AppError, match="Word 里另存为") as error:
        await extract._save_invitation(upload)
    assert error.value.status_code == 422
    assert not list(tmp_path.iterdir())
