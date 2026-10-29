import pytest
from docx import Document
from docx.oxml.ns import qn

from api.services.tenders import format_preview
from api.services.tenders.document import _pageref_bookmark
from api.services.tenders.schema import BidBrief, OutlineItem, QuoteLineIn
from common.errors import AppError

INVITE_ID = 'deadbeef1234'


@pytest.fixture
def preview_env(tmp_path, monkeypatch):
    invitations = tmp_path / 'invitations'
    output = tmp_path / 'output'
    invitations.mkdir()
    output.mkdir()
    doc = Document()
    doc.add_paragraph('无投标格式的项目说明')
    doc.save(invitations / f'{INVITE_ID}.docx')
    template = tmp_path / 'selected-template.docx'
    doc = Document()
    for title in ('当前模板封面', '目录', '投标函', '【待填写】'):
        doc.add_paragraph(title)
    doc.save(template)
    monkeypatch.setattr(format_preview, '_ready_format_docx', lambda _: None)
    monkeypatch.setattr(format_preview, 'tender_invitations_dir', lambda: invitations)
    monkeypatch.setattr(format_preview, 'tenders_output_dir', lambda: output)
    monkeypatch.setattr(format_preview, 'find_chapter5_template', lambda: template)
    return invitations, output, template


def test_missing_native_format_returns_selected_template_copy_and_notice(preview_env):
    invitations, output, template = preview_env
    before = {p.name: p.read_bytes() for p in invitations.iterdir()}
    payload = format_preview.prepare_invitation_format_docx(INVITE_ID)
    assert payload['previewOnly'] and payload['source'] == 'builtin'
    assert '未抽到格式章' in payload['message'] and '正文待生成' in payload['message']
    assert (output / payload['fileName']).read_bytes() == template.read_bytes()
    assert not payload['fileName'].endswith('.format.docx')
    assert before == {p.name: p.read_bytes() for p in invitations.iterdir()}


def test_native_format_preview_retains_original_editable_channel(preview_env, monkeypatch):
    invitations, output, _ = preview_env
    source = invitations / f'{INVITE_ID}.format.docx'
    doc = Document()
    doc.add_paragraph('原格式的投标文件封面')
    doc.save(source)
    monkeypatch.setattr(format_preview, '_ready_format_docx', lambda _: source)
    payload = format_preview.prepare_invitation_format_docx(INVITE_ID)
    assert payload == {'fileName': source.name, 'source': 'invitation',
                       'previewOnly': False, 'message': ''}
    assert (output / source.name).read_bytes() == source.read_bytes()


def test_fixed_selection_previews_generation_template_even_with_native_file(
    preview_env, monkeypatch,
):
    invitations, output, template = preview_env
    source = invitations / f'{INVITE_ID}.format.docx'
    source.write_bytes(template.read_bytes())
    monkeypatch.setattr(format_preview, '_ready_format_docx', lambda _: source)
    payload = format_preview.prepare_invitation_format_docx(
        INVITE_ID, brief=BidBrief(layoutMode='chapter5'),
    )
    assert payload['source'] == 'builtin' and payload['previewOnly']
    assert (output / payload['fileName']).read_bytes() == template.read_bytes()


def test_requirements_preview_has_linked_toc_and_no_draft_or_quote_data(preview_env):
    _, output, _ = preview_env
    brief = BidBrief(
        projectName='本标封面标题', requirementsDriven=True, generateVolume='full',
        bidderName='PRIVATE_BIDDER', legalPersonName='PRIVATE_PERSON', bidPriceYuan=123456789,
        quoteLines=[QuoteLineIn(name='PRIVATE_QUOTE_LINE', qty=999)],
        outlineItems=[OutlineItem(id='plan', title='质量保证方案', kind='tech_chapter',
                                  draftBody='PRIVATE_DRAFT', body='PRIVATE_SOURCE')],
    )
    before = brief.model_dump()
    payload = format_preview.prepare_invitation_format_docx(INVITE_ID, brief=brief)
    assert payload['source'] == 'requirements' and payload['previewOnly']
    assert brief.model_dump() == before
    doc = Document(output / payload['fileName'])
    text = '\n'.join(p.text for p in doc.paragraphs) + doc.element.xml
    assert all(marker not in text for marker in ('PRIVATE_', '123456789'))
    assert '本标封面标题' in text and '【待生成】' in text
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert len(entries) >= 20
    assert all(p.text.rsplit('\t', 1)[-1].isdigit() for p in entries)
    names = [n.get(qn('w:name')) for n in doc.element.body.iter(qn('w:bookmarkStart'))]
    assert len(names) == len(set(names))
    assert all(_pageref_bookmark(p) in names for p in entries)
    assert any(p.text == '质量保证方案' for p in doc.paragraphs)


def test_skeleton_callback_cannot_replace_the_invitation_format(preview_env, monkeypatch):
    from api.services.tenders.onlyoffice import _keep_format_template

    _, output, _ = preview_env
    payload = format_preview.prepare_invitation_format_docx(INVITE_ID)
    def unexpected(*args):
        pytest.fail('A skeleton preview must not sync into the native template')
    monkeypatch.setattr('api.services.tenders.pdf_format.format_docx_path', unexpected)
    _keep_format_template(payload['fileName'], output / payload['fileName'])


@pytest.mark.parametrize('invite_id,status', [('../bad', 400), ('abcdef123456', 404)])
def test_unknown_or_unsafe_invitation_id_does_not_create_a_preview(preview_env, invite_id, status):
    _, output, _ = preview_env
    with pytest.raises(AppError) as exc:
        format_preview.prepare_invitation_format_docx(invite_id)
    assert exc.value.status_code == status
    assert not list(output.iterdir())


def test_missing_company_template_has_clear_error(preview_env, monkeypatch):
    monkeypatch.setattr(format_preview, 'find_chapter5_template', lambda: None)
    with pytest.raises(AppError, match='未找到当前生成使用的内置模板'):
        format_preview.prepare_invitation_format_docx(INVITE_ID)


@pytest.mark.asyncio
async def test_get_and_post_preview_routes_preserve_file_name_channel(preview_env, monkeypatch):
    from api.routers.tenders import (
        tenders_invitation_format_docx,
        tenders_prepare_invitation_format_docx,
    )

    monkeypatch.setattr('api.services.tenders.demo.demo_enabled', lambda: False)
    old = await tenders_invitation_format_docx(INVITE_ID, None, None)
    current = await tenders_prepare_invitation_format_docx(
        INVITE_ID, BidBrief(layoutMode='chapter5'), None, None,
    )
    for response in (old, current):
        assert response['code'] == 0 and response['data']['fileName']
        assert response['data']['previewOnly'] and response['data']['message']


def test_hospital_preview_prioritizes_original_format_without_touching_source(
    tmp_path, monkeypatch,
):
    from api.services.tenders.assets import tender_invitations_dir

    invite_id = 'ba636a37878c'
    source = tender_invitations_dir() / f'{invite_id}.format.docx'
    if not source.is_file():
        pytest.skip('Local hospital acceptance fixture is not installed')
    before = source.read_bytes()
    monkeypatch.setattr(format_preview, 'tenders_output_dir', lambda: tmp_path)
    brief = BidBrief(projectName='河南省胸科医院车辆采购（救护车）项目（二次）',
                     requirementsDriven=True, generateVolume='full')
    payload = format_preview.prepare_invitation_format_docx(invite_id, brief=brief)
    doc = Document(tmp_path / payload['fileName'])
    assert payload['source'] == 'invitation' and not payload['previewOnly']
    assert (tmp_path / payload['fileName']).read_bytes() == before
    assert any('资格证明' in p.text or '资格证明' in t._tbl.xml for p in doc.paragraphs
               for t in doc.tables)
    assert source.read_bytes() == before
