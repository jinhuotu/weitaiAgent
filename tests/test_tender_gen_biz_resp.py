from __future__ import annotations

import asyncio
import json

import pytest
from docx import Document

from api.services.tenders.assemble import _fill_business_format_sections, assemble_bid_docx
from api.services.tenders.categories import business_soft_issues
from api.services.tenders.gen_biz_resp import (
    _business_scores,
    _parse,
    _substantial_clauses,
    _won_business_docs,
    business_requirements,
    draft_business_response,
)
from api.services.tenders.schema import BidBrief, BidRequirement, DeviationLine, OutlineItem


class FakeClient:
    async def complete(self, messages, *, mode):
        assert mode == "pro"
        assert "付款方式" in messages[1]["content"]
        return (
            json.dumps(
                {
                    "clauses": [
                        {
                            "requirement": "付款方式按合同约定，应提供承诺",
                            "response": "同意按合同约定付款",
                            "deviation": "无偏离",
                            "explanation": "按合同执行",
                        },
                        {
                            "requirement": "须提供近年业绩证明",
                            "response": "拥有99项业绩",
                            "deviation": "无偏离",
                            "explanation": "",
                        },
                    ],
                    "chapters": [
                        {
                            "title": "商务优惠及特殊承诺",
                            "body": "优惠金额【待补充】；履约按合同执行。",
                        },
                        {
                            "title": "响应承诺书",
                            "body": "本项目按合同条款履约；业绩证明【待补充】。",
                        },
                    ],
                },
                ensure_ascii=False,
            )
            + "\nreasoning follows"
        )


def test_business_response_drafts_existing_chapters_and_keeps_manual_edits(tmp_path):
    brief = BidBrief(
        projectName="本项目",
        autoDraftBizResp=True,
        layoutMode="outline",
        outlineItems=[
            OutlineItem(id="d", title="商务偏离表", kind="biz_dev"),
            OutlineItem(id="o", title="商务优惠及特殊承诺", kind="unknown"),
            OutlineItem(id="c", title="响应承诺书", kind="commitment_copy"),
            OutlineItem(
                id="m",
                title="人工承诺书",
                kind="commitment_copy",
                draftBody="人工填写",
                draftStatus="edited",
            ),
        ],
        bidRequirements=[
            BidRequirement(text="付款方式按合同约定，应提供承诺", volume="business"),
            BidRequirement(text="须提供近年业绩证明", volume="business"),
        ],
    )
    result, notes = asyncio.run(draft_business_response(None, brief, client=FakeClient()))
    assert [(i.id, i.title, i.level) for i in result.outlineItems] == [
        (i.id, i.title, i.level) for i in brief.outlineItems
    ]
    assert result.outlineItems[3].draftBody == "人工填写"
    assert result.bizDeviationLines[0].deviation == "无偏离"
    assert result.bizDeviationLines[0].explanation == "按合同执行"
    assert result.bizDevHeaders == ["招标条款", "响应", "偏离", "说明"]
    assert result.bizDeviationLines[1].deviation == "待确认"
    assert "99" not in result.bizDeviationLines[1].response
    assert notes
    dest = tmp_path / "biz.docx"
    assemble_bid_docx(result, dest, volume="business")
    doc = Document(dest)
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "本项目按合同条款履约" in text
    assert "逐条商务响应" not in text
    assert "须提供近年业绩证明" not in text
    assert "1 项商务条款尚未响应" in tables
    assert "付款方式按合同约定" in tables
    assert "按合同执行" in tables
    assert "99" not in tables
    assert business_soft_issues(result)


def test_copied_format_keeps_original_and_appends_business_responses(tmp_path):
    path = tmp_path / "format.docx"
    source = Document()
    source.add_paragraph("商务偏离表")
    table = source.add_table(rows=1, cols=4)
    for cell, title in zip(table.rows[0].cells, ("招标条款", "响应", "偏离", "说明"), strict=True):
        cell.text = title
    source.add_paragraph("商务优惠及特殊承诺")
    source.add_paragraph("原优惠模板")
    source.add_paragraph("响应承诺书")
    source.add_paragraph("原承诺模板")
    source.save(path)
    brief = BidBrief(
        autoDraftBizResp=True,
        bizDevHeaders=["招标条款", "响应", "偏离", "说明"],
        bizDeviationLines=[
            DeviationLine(
                requirement="合同付款条款",
                response="同意",
                deviation="无偏离",
                explanation="按合同执行",
            )
        ],
        outlineItems=[
            OutlineItem(title="商务偏离表", kind="biz_dev"),
            OutlineItem(title="商务优惠及特殊承诺", kind="unknown", draftBody="优惠待补充"),
            OutlineItem(title="响应承诺书", kind="commitment_copy", draftBody="按要求履约"),
        ],
    )
    inserted, missing = _fill_business_format_sections(path, brief)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert inserted == 3 and not missing
    assert "原优惠模板" in text and "优惠待补充" in text
    assert "原承诺模板" in text and "按要求履约" in text
    assert "合同付款条款" in tables and "按合同执行" in tables
    assert len(doc.tables) == 1


def test_business_breakdown_separates_substantial_and_weighted_scores():
    brief = BidBrief(
        bidRequirements=[
            BidRequirement(text="须提供售后培训方案", volume="business"),
            BidRequirement(text="类似业绩得 5 分", category="scoring", volume="business"),
        ]
    )
    clauses = business_requirements(brief, "交付验收应符合合同约定")
    assert any("售后培训" in row.text for row in clauses)
    assert any("交付验收" in row.text for row in clauses)
    assert any("须提供" in row.text for row in _substantial_clauses(clauses))
    assert ("类似业绩得 5 分", "5分") in _business_scores(brief, "")


def test_selected_knowledge_only_contributes_writing_style():
    brief = BidBrief(
        bizRespKbIds=["chosen-kb"],
        outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")],
        bidRequirements=[BidRequirement(text="付款方式应按合同执行", volume="business")],
    )
    calls = []

    async def recall(_db, _invitation, ids, **kwargs):
        calls.append((ids, kwargs["query"]))
        return [{"content": "1. 先列承诺，再附证明材料。旧项目金额 999 万元。"}], []

    class Client:
        messages = None

        async def complete(self, messages, *, mode):
            self.messages = messages
            return json.dumps(
                {
                    "clauses": [
                        {
                            "requirement": "付款方式应按合同执行",
                            "response": "按合同执行",
                            "deviation": "无偏离",
                            "explanation": "依据合同",
                        }
                    ],
                    "chapters": [{"title": "响应承诺书", "body": "按合同条款履约。"}],
                },
                ensure_ascii=False,
            )

    client = Client()
    drafted, notes = asyncio.run(draft_business_response(None, brief, client=client, recall=recall))
    assert calls[0][0] == ["chosen-kb"]
    assert "按条款编号逐项回应" in client.messages[1]["content"]
    assert "999" not in client.messages[1]["content"]
    assert "商务响应完整性核对表" not in drafted.outlineItems[0].draftBody
    assert any("知识库" in note for note in notes)


def test_empty_knowledge_and_missing_clause_stay_visible():
    brief = BidBrief(
        bizRespKbIds=["chosen-kb"],
        outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")],
        bidRequirements=[BidRequirement(text="必须提供验收承诺", volume="business")],
    )

    async def recall(*_args, **_kwargs):
        return [], []

    class Client:
        async def complete(self, _messages, *, mode):
            return json.dumps(
                {
                    "clauses": [],
                    "chapters": [{"title": "响应承诺书", "body": "验收要求待核实。"}],
                },
                ensure_ascii=False,
            )

    drafted, notes = asyncio.run(
        draft_business_response(None, brief, client=Client(), recall=recall)
    )
    assert drafted.bizDeviationLines[0].deviation == "待确认"
    assert "【待补充】" in drafted.outlineItems[0].draftBody
    assert any("实质性条款待补充" in note for note in notes)


def test_no_extracted_business_clause_is_explicitly_flagged():
    brief = BidBrief(outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")])
    drafted, notes = asyncio.run(draft_business_response(None, brief))
    assert "商务响应完整性核对表" in drafted.outlineItems[0].draftBody
    assert "【待补充】" in drafted.outlineItems[0].draftBody
    assert any("未提取到商务要求" in note for note in notes)
    assert any("邀请书正文为空" in note for note in notes)


def test_business_outline_without_landing_section_reports_gap():
    brief = BidBrief(
        outlineItems=[OutlineItem(title="投标函", kind="unknown")],
        bidRequirements=[BidRequirement(text="付款方式必须按合同执行", volume="business")],
    )
    drafted, notes = asyncio.run(draft_business_response(None, brief))
    assert drafted.outlineItems == brief.outlineItems
    assert any("大纲无「商务偏离表」对应章节" in note for note in notes)
    assert any("大纲无「响应承诺正文」对应章节" in note for note in notes)


def test_many_substantial_clauses_are_batched_without_dropping_rows():
    class Client:
        prompts = []

        async def complete(self, messages, *, mode):
            self.prompts.append(messages[1]["content"])
            return json.dumps({"clauses": [], "chapters": []}, ensure_ascii=False)

    client = Client()
    brief = BidBrief(
        outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")],
        bidRequirements=[
            BidRequirement(text=f"第{i}项付款方式必须按合同执行", volume="business")
            for i in range(55)
        ],
    )
    drafted, _notes = asyncio.run(draft_business_response(None, brief, client=client))
    assert len(client.prompts) == 3
    assert len(drafted.bizDeviationLines) == 55
    assert "第54项付款方式必须按合同执行" not in drafted.outlineItems[0].draftBody
    assert drafted.outlineItems[0].draftBody.count("【待补充】") == 1
    assert "第54项付款方式必须按合同执行" in client.prompts[-1]
    assert "第0项付款方式必须按合同执行" not in client.prompts[-1]


def test_style_recall_uses_winning_business_chapters_only():
    class Result:
        def all(self):
            return [
                ("won-biz", ["tech_bid", "volume:business", "outcome:won"]),
                ("lost-biz", ["tech_bid", "volume:business", "outcome:lost"]),
                ("won-tech", ["tech_bid", "volume:technical", "outcome:won"]),
            ]

    class Db:
        async def execute(self, _query):
            return Result()

    docs = [{"docId": name, "content": "承诺"} for name in ("won-biz", "lost-biz", "won-tech")]
    kept = asyncio.run(_won_business_docs(Db(), docs))
    assert [doc["docId"] for doc in kept] == ["won-biz"]


@pytest.mark.parametrize("raw", [
    '{}', '{"clauses": []}', '{"chapters": []}',
    '{"clauses": null, "chapters": "missing"}', '[]', 'null',
    '', '未提取到可回应的商务条款', '{"clauses": [',
])
def test_incomplete_business_response_json_degrades_to_empty_lists(raw):
    assert _parse(raw) == {"clauses": [], "chapters": []}


def test_business_requirements_exclude_evaluation_rules_but_keep_actionable_clauses():
    rules = [
        '商务标得分等于下列4项得分之和。',
        '商务标得分=价格得分与业绩得分之和。',
        '商务部分评审的业绩规定',
        '6.5商务（综合）部分评审的业绩规定（不含资格要求业绩）',
        '类似业绩得5分，商务部分满分20分。',
        '商务评审规则要求各项业绩得分相加计算。',
    ]
    clauses = [
        '投标人须提供近三年类似工程业绩合同及验收证明。',
        '投标人须提供类似工程业绩合同，每提供一项得2分。',
        '付款方式应按合同约定执行。',
    ]
    brief = BidBrief(bidRequirements=[
        BidRequirement(text=text, volume='business', category='scoring') for text in rules
    ])
    result = business_requirements(brief, '\n'.join([*rules, *clauses]))
    texts = {row.text for row in result}
    assert not texts.intersection(rules)
    assert set(clauses) <= texts
    assert _business_scores(brief, '\n'.join(rules))


@pytest.mark.parametrize('sections', [
    [OutlineItem(title='投标函', kind='letter')],
    [OutlineItem(title='响应承诺书', kind='commitment_copy', draftBody='人工内容', draftStatus='edited')],
    [OutlineItem(title='响应承诺书', kind='commitment_copy', skipped=True)],
])
def test_no_editable_business_section_never_calls_model_or_recall(monkeypatch, sections):
    async def unexpected(*args, **kwargs):
        pytest.fail('No editable commercial section should bypass LLM and knowledge recall')

    monkeypatch.setattr('api.services.models.runtime.build_llm_client', unexpected)
    brief = BidBrief(outlineItems=sections, bizRespKbIds=['kb'], bidRequirements=[
        BidRequirement(text='付款方式必须按合同执行', volume='business'),
    ])
    original = brief.model_dump()
    result, notes = asyncio.run(draft_business_response(None, brief, recall=unexpected))
    assert result.model_dump() == original
    assert any('【待补充】' in note for note in notes)


@pytest.mark.parametrize('role', ['commitment', 'offer', 'deviation'])
def test_review_rules_only_mark_existing_business_section_pending(monkeypatch, role):
    text = '商务标得分等于下列4项得分之和。\n商务部分评审的业绩规定'
    monkeypatch.setattr('api.services.tenders.gen_biz_resp.load_invitation_text', lambda _: text)

    async def unexpected(*args, **kwargs):
        pytest.fail('Evaluation rules alone must not invoke the model')

    monkeypatch.setattr('api.services.models.runtime.build_llm_client', unexpected)
    title, kind = {
        'commitment': ('响应承诺书', 'commitment_copy'),
        'offer': ('商务优惠及特殊承诺', 'unknown'),
        'deviation': ('商务偏离表', 'biz_dev'),
    }[role]
    brief = BidBrief(outlineItems=[OutlineItem(id='biz', title=title, kind=kind)])
    result, notes = asyncio.run(draft_business_response(None, brief))
    assert [(i.id, i.title, i.level) for i in result.outlineItems] == [('biz', title, 1)]
    assert '【待补充】' in result.outlineItems[0].draftBody
    if role == 'deviation':
        assert '【待补充】' in result.bizDeviationLines[0].response
    assert any('【待补充】' in note for note in notes)


@pytest.mark.parametrize('missing', ['clauses', 'chapters', 'both', 'invalid'])
def test_partial_business_model_result_keeps_valid_data_and_marks_gaps(monkeypatch, missing):
    requirement = '付款方式必须按合同执行'
    monkeypatch.setattr('api.services.tenders.gen_biz_resp.load_invitation_text', lambda _: requirement)
    payload = {
        'clauses': [{'requirement': requirement, 'response': '同意按合同执行',
                     'deviation': '无偏离', 'explanation': '按合同履约'}],
        'chapters': [{'title': '响应承诺书', 'body': '按本项目合同条款履约。'},
                     {'title': '商务优惠及特殊承诺', 'body': '仅提供本项目书面约定的优惠。'}],
    }
    if missing in ('clauses', 'both'):
        payload.pop('clauses')
    if missing in ('chapters', 'both'):
        payload.pop('chapters')

    class Client:
        async def complete(self, *args, **kwargs):
            return '模型没有返回JSON' if missing == 'invalid' else json.dumps(payload, ensure_ascii=False)

    brief = BidBrief(autoDraftBizResp=True, extraNote='优惠须书面确认',
                     outlineItems=[OutlineItem(title='响应承诺书', kind='commitment_copy'),
                                   OutlineItem(title='商务优惠及特殊承诺')],
                     bidRequirements=[BidRequirement(text=requirement, volume='business')])
    result, notes = asyncio.run(draft_business_response(None, brief, client=Client()))
    if missing == 'chapters':
        assert result.bizDeviationLines[0].response == '同意按合同执行'
        assert all('【待补充】' in item.draftBody for item in result.outlineItems)
    else:
        assert '【待补充】' in result.bizDeviationLines[0].response
        assert result.bizDeviationLines[0].deviation == '待确认'
    if missing == 'clauses':
        assert '按本项目合同条款履约。' in result.outlineItems[0].draftBody
    assert any('【待补充】' in note for note in notes)


def test_incomplete_later_batch_keeps_first_batch_answers(monkeypatch):
    requirements = [f'第{i}项付款方式必须按合同执行' for i in range(21)]
    monkeypatch.setattr('api.services.tenders.gen_biz_resp.load_invitation_text', lambda _: '')

    class Client:
        calls = 0

        async def complete(self, *args, **kwargs):
            self.calls += 1
            if self.calls == 2:
                return '{"chapters": null}'
            return json.dumps({'clauses': [
                {'requirement': text, 'response': '同意按合同履约', 'deviation': '无偏离'}
                for text in requirements[:20]
            ], 'chapters': [{'title': '响应承诺书', 'body': '本项目按合同履约。'}]}, ensure_ascii=False)

    client = Client()
    brief = BidBrief(outlineItems=[OutlineItem(title='响应承诺书', kind='commitment_copy')],
                     bidRequirements=[BidRequirement(text=text, volume='business') for text in requirements])
    result, notes = asyncio.run(draft_business_response(None, brief, client=client))
    assert client.calls == 2
    assert len(result.bizDeviationLines) == 21
    assert result.bizDeviationLines[0].response == '同意按合同履约'
    assert '【待补充】' in result.bizDeviationLines[-1].response
    assert '本项目按合同履约。' in result.outlineItems[0].draftBody
    assert any('缺失正文' in note for note in notes)


@pytest.mark.parametrize('volume', ['business', 'full'])
def test_construction_rules_without_business_clauses_still_assemble(tmp_path, monkeypatch, volume):
    text = ('第三章 评标办法\n商务标得分等于下列4项得分之和。\n'
            '商务部分评审的业绩规定\n施工组织设计应明确施工部署和技术措施。')
    monkeypatch.setattr('api.services.tenders.gen_biz_resp.load_invitation_text', lambda _: text)

    async def unexpected(*args, **kwargs):
        pytest.fail('No actionable commercial clause should bypass the model')

    monkeypatch.setattr('api.services.models.runtime.build_llm_client', unexpected)
    source = Document()
    source.add_paragraph('商务响应承诺书')
    source.add_paragraph('原商务模板正文。')
    source.add_paragraph('施工组织设计')
    source.add_paragraph('本项目施工部署与技术措施。')
    source_path = tmp_path / 'construction.format.docx'
    source.save(source_path)
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: source_path)
    brief = BidBrief(generateVolume=volume, autoDraftBizResp=True, attachQualifications=False,
                     layoutMode='outline', invitationId='abc123def456',
                     includePlaceholders=False, outlineItems=[
                         OutlineItem(id='response', title='商务响应承诺书', kind='commitment_copy'),
                         OutlineItem(id='tech', title='施工组织设计', kind='tech_chapter',
                                     draftBody='本项目施工部署与技术措施。'),
                     ])
    result, notes = asyncio.run(draft_business_response(None, brief))
    path, _ = assemble_bid_docx(result, tmp_path / f'{volume}.docx', volume=volume,
                                catalog_slots=[], catalog_media={}, catalog_facts=[])
    generated = '\n'.join(p.text for p in Document(path).paragraphs)
    assert '商务响应承诺书' in generated and '【待补充】' in generated
    assert '【待补充】' in result.outlineItems[0].draftBody
    if volume == 'full':
        assert '本项目施工部署与技术措施。' in generated
    assert notes
