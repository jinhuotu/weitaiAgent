"""Library keys reach native-format chapters before the unchanged image placer runs."""

import asyncio
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from api.services.tenders import library_kb, pdf_format, records
from api.services.tenders import slots as slot_files
from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.placeholders import DEFAULT_SLOTS, collect_slots, this_bid_keys
from api.services.tenders.schema import BidBrief, BidRequirement, OutlineItem, PlaceholderItem
from common.errors import AppError


@pytest.mark.parametrize(
    "title,key",
    [
        ("营业执照", "license"),
        ("具有独立承担民事责任的能力", "license"),
        ("基本账户开户证明", "bank_permit"),
        ("具有健全的财务会计制度", "finance"),
        ("财务状况证明", "finance"),
        ("信用中国查询", "credit"),
        ("类似项目业绩", "perf"),
        ("资质证明", "iso"),
        ("投标保证金缴存回单", "bond"),
        ("投标文件签章页", "seal"),
    ],
)
def test_detected_section_cannot_be_filtered_by_empty_catalog_or_old_include(title, key):
    required, include = this_bid_keys(
        catalog=[],
        required_keys=["id_legal"],
        include_keys=["id_legal"],
        outline_items=[OutlineItem(title=title, kind="scan")],
    )
    assert key in required and key in include
    chosen = collect_slots(catalog=[], include_keys=include)
    assert {s.key for s in chosen} >= {"id_legal", key}
    assert "product" not in include and "id_agent" not in include


def test_requirement_and_full_text_select_required_materials_not_catalog_descriptions():
    catalog = [PlaceholderItem(key="unrelated", title="无关图纸", hint="包含财务状况、业绩说明")]
    required, include = this_bid_keys(
        catalog=catalog,
        include_keys=["id_legal"],
        requirements=[
            BidRequirement(text="供应商应提供依法缴纳税收的证明", category="qualification")
        ],
        source_text="提交信用中国查询页。\n提交投标保证金回单。",
    )
    assert {"id_legal", "finance", "credit", "bond"} <= set(required) <= set(include)
    assert "perf" not in required and "unrelated" not in include


def test_paste_cues_select_license_tax_and_social_not_the_whole_library():
    text = "\n".join([
        "此处粘贴营业执照复印件（在此粘贴扫描件）",
        "此处粘贴纳税证明（在此粘贴扫描件）",
        "此处粘贴社保证明（在此粘贴扫描件）",
        "此处粘贴委托代理人身份证（在此粘贴扫描件）",
    ])
    required, include = this_bid_keys(catalog=[], has_agent=True, source_text=text)
    assert {"id_legal", "id_agent", "license", "finance", "social"} <= set(required)
    assert "product" not in include
    required, include = this_bid_keys(catalog=[], has_agent=False, source_text=text)
    assert "id_agent" not in include
    assert {"license", "finance", "social", "id_legal"} <= set(include)


def test_skipped_sections_and_disabled_agent_do_not_select_materials():
    required, include = this_bid_keys(
        catalog=[],
        has_agent=False,
        outline_items=[
            OutlineItem(title="财务状况", skipped=True),
            OutlineItem(title="委托代理人身份证"),
        ],
        source_text="无需提供投标保证金\n投标保证金：不收取",
    )
    assert required == ["id_legal"] and include == ["id_legal"]
    required, include = this_bid_keys(catalog=[], has_agent=True)
    assert required == ["id_legal", "id_agent"] and set(required) <= set(include)


def test_catalog_preserves_defaults_and_overrides_metadata_without_mutation():
    catalog = [PlaceholderItem(key="license", title="本企业营业执照副本", hint="彩色清晰扫描件")]
    original = catalog[0].model_dump()
    items = collect_slots(catalog=catalog)
    assert {i.key for i in items} == {i.key for i in DEFAULT_SLOTS}
    assert len(items) == len(DEFAULT_SLOTS)
    assert next(i for i in items if i.key == "license").title == "本企业营业执照副本"
    assert catalog[0].model_dump() == original
    assert [s.key for s in collect_slots(catalog=[], include_keys=["license", "id_legal"])] == [
        "license",
        "id_legal",
    ]


@pytest.fixture
def native_env(tmp_path, monkeypatch):
    root = tmp_path / "assets"
    monkeypatch.setattr(slot_files, "tender_assets_dir", lambda: root)
    source = tmp_path / "original.docx"
    doc = Document()
    doc.add_paragraph("投 标 文 件")
    doc.add_paragraph("项目名称：本次测试采购")
    doc.add_page_break()
    doc.add_paragraph("目录")
    for title in ("一、法定代表人身份证正反面", "二、营业执照"):
        doc.add_paragraph(title).add_run().add_tab()
    doc.add_page_break()
    doc.add_paragraph("一、法定代表人身份证正反面", style="Heading 2")
    doc.add_paragraph("附：法定代表人身份证正反面扫描件")
    doc.add_paragraph("法定代表人（签字）：")
    doc.add_paragraph("二、营业执照", style="Heading 2")
    doc.add_paragraph("本节附企业营业执照扫描件。")
    doc.add_paragraph("三、下一章", style="Heading 2")
    doc.save(source)
    monkeypatch.setattr(pdf_format, "_word_format_file", lambda _: source)
    monkeypatch.setattr("api.services.tenders.requirement_bid.invitation_text", lambda _: "")
    return tmp_path, source


def _brief():
    return BidBrief(
        invitationId="eaeaeaeaeaea",
        outlineChapter="第六章 投标文件格式",
        layoutMode="outline",
        formatConfirmed=True,
        generateVolume="business",
        projectName="本次测试采购",
        requiredSlotKeys=["id_legal"],
        includeSlotKeys=["id_legal"],
        outlineItems=[
            OutlineItem(
                id="identity",
                title="法定代表人身份证正反面",
                rawTitle="一、法定代表人身份证正反面",
                kind="legal_id",
            ),
            OutlineItem(id="license", title="营业执照", rawTitle="二、营业执照", kind="scan"),
        ],
    )


def _upload_identity_and_license(tmp_path):
    for key, name, color in (
        ("id_legal", "front.png", "red"),
        ("id_legal", "back.png", "green"),
        ("license", "license.png", "blue"),
    ):
        image = tmp_path / name
        Image.new("RGB", (160, 100), color).save(image)
        slot_files.save_slot_file(key, filename=name, data=image.read_bytes())


def _assert_images_in_native_sections(path):
    doc = Document(path)
    children = list(doc.element.body)
    legal = next(
        i
        for i, el in enumerate(children)
        if "".join(t.text or "" for t in el.iter(qn("w:t"))) == "一、法定代表人身份证正反面"
    )
    license_head = next(
        i
        for i, el in enumerate(children)
        if "".join(t.text or "" for t in el.iter(qn("w:t"))) == "二、营业执照"
    )
    end = next(
        i
        for i, el in enumerate(children)
        if "".join(t.text or "" for t in el.iter(qn("w:t"))) == "三、下一章"
    )
    ids = [
        node.get(qn("r:embed"))
        for el in children[legal + 1 : license_head]
        for node in el.iter(qn("a:blip"))
    ]
    # The existing renderer can stitch the two card faces into one image strip.
    assert ids
    pixels = []
    for ident in ids:
        with Image.open(BytesIO(doc.part.related_parts[ident].blob)) as image:
            pixels.extend(image.convert("RGB").getdata())
    assert any(r > 180 and g < 60 and b < 60 for r, g, b in pixels)
    assert any(g > 80 and r < 60 and b < 60 for r, g, b in pixels)
    assert sum(len(list(el.iter(qn("a:blip")))) for el in children[license_head + 1 : end]) == 1
    assert not any(list(el.iter(qn("a:blip"))) for el in pdf_format.toc_block_ids(doc))


@pytest.mark.parametrize("volume", ["business", "full"])
def test_direct_native_builder_loads_legacy_files_before_native_branch(native_env, volume):
    tmp_path, source = native_env
    before = source.read_bytes()
    _upload_identity_and_license(tmp_path)
    brief = _brief()
    path, notes = assemble_bid_docx(brief, tmp_path / "standalone.docx", volume=volume)
    _assert_images_in_native_sections(path)
    assert "license" in brief.requiredSlotKeys and "license" in brief.includeSlotKeys
    assert not any("资料库缺少营业执照" in note for note in notes)
    assert before == source.read_bytes()


def _mock_library(monkeypatch):
    for name, value in (
        ("catalog_placeholders", []),
        ("performance_from_library", []),
        ("count_uncached_performance_files", 0),
        ("legal_persons_from_library", []),
        ("library_facts", []),
        ("list_library_items", []),
    ):
        monkeypatch.setattr(library_kb, name, AsyncMock(return_value=value))
    # Test the real attachment resolver's legacy directory fallback.
    monkeypatch.setattr(
        library_kb, "get_base_by_public_id", AsyncMock(side_effect=AppError(msg="No library base"))
    )
    monkeypatch.setattr(records, "_brief_generate_issues", lambda _: [])


@pytest.mark.parametrize("uploaded", [False, True])
def test_record_pipeline_keeps_critical_keys_and_reports_missing_files(
    native_env,
    monkeypatch,
    uploaded,
):
    tmp_path, _ = native_env
    if uploaded:
        _upload_identity_and_license(tmp_path)
    _mock_library(monkeypatch)
    captured = {}

    def generate(brief, **kwargs):
        captured.update(kwargs)
        path, notes = assemble_bid_docx(
            brief, tmp_path / "record.docx", volume="business", **kwargs
        )
        return {"docxFile": path.name, "downloadName": "本次测试.docx", "warnings": notes}

    monkeypatch.setattr(records, "generate_bid", generate)
    brief = _brief()
    result, _ = asyncio.run(records._prepare_and_generate(None, brief))
    assert {"id_legal", "license"} <= {s.key for s in captured["catalog_slots"]}
    if uploaded:
        assert len(captured["catalog_media"]["id_legal"]) == 2
        assert len(captured["catalog_media"]["license"]) == 1
        _assert_images_in_native_sections(tmp_path / result["docxFile"])
    else:
        assert not captured["catalog_media"]
        assert "资料库缺少营业执照，请在资料库对应项上传" in result["warnings"]
        assert "资料库缺少法定代表人身份证正反面，请在资料库对应项上传" in result["warnings"]


def test_record_does_not_reload_identity_or_performance_files_after_filtering(
    native_env,
    monkeypatch,
):
    tmp_path, _ = native_env
    _upload_identity_and_license(tmp_path)
    _mock_library(monkeypatch)
    excluded = tmp_path / "other-person.png"
    Image.new("RGB", (40, 30), "white").save(excluded)
    slot_files.save_slot_file("id_legal", filename=excluded.name, data=excluded.read_bytes())
    slot_files.save_slot_file("perf", filename="wrong-project.png", data=excluded.read_bytes())
    chosen = [p for p in slot_files.list_slot_files("id_legal") if p.name != excluded.name]
    monkeypatch.setattr(
        library_kb,
        "legal_persons_from_library",
        AsyncMock(
            return_value=[
                {"key": "person-a", "name": "甲"},
                {"key": "person-b", "name": "乙"},
            ]
        ),
    )
    monkeypatch.setattr(library_kb, "filter_legal_id_attachments", AsyncMock(return_value=chosen))
    monkeypatch.setattr(library_kb, "group_performance_attachments", AsyncMock(return_value={}))
    monkeypatch.setattr(library_kb, "filter_performance_attachments", AsyncMock(return_value=[]))
    captured = {}

    def generate(brief, **kwargs):
        captured.update(kwargs)
        return {"docxFile": "test.docx", "warnings": []}

    monkeypatch.setattr(records, "generate_bid", generate)
    brief = _brief()
    brief.legalPersonKey = "person-a"
    brief.outlineItems.append(OutlineItem(title="类似项目业绩", kind="performance"))
    asyncio.run(records._prepare_and_generate(None, brief))
    assert set(captured["catalog_media"]["id_legal"]) == set(chosen)
    assert "perf" not in captured["catalog_media"]


def test_dedicated_chapter_wins_over_letter_mention_for_attachment_routing():
    items = collect_slots(
        catalog=[],
        include_keys=["license"],
        outline_items=[
            OutlineItem(title="投标函", kind="letter", body="随同递交营业执照扫描件。"),
            OutlineItem(
                title="具有独立承担民事责任的能力", rawTitle="二、具有独立承担民事责任的能力"
            ),
        ],
    )
    assert items[0].attachToSection == "二、具有独立承担民事责任的能力"
