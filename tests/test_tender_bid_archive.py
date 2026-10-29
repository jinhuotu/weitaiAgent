"""历史标书切章和剥离。不连库、不连向量。"""

import pytest
from docx import Document

from api.services.knowledge.classify import classify_kb_document
from api.services.tenders.bid_archive import (
    accept_history_upload,
    chapters_from_docx,
    chapters_from_outline,
    prepare_chapters,
    resolve_chapters,
    sanitize_text,
)
from api.services.tenders.schema import BidBrief, OutlineItem


def _brief() -> BidBrief:
    return BidBrief(
        projectName="郑州高新新能源项目",
        tenderer="郑州某某建设集团",
        tenderNo="26-GCZN-0150",
        bidPriceYuan=1280000,
    )


def test_sanitize_keeps_wording_and_drops_deal() -> None:
    text = (
        "郑州高新新能源项目由郑州某某建设集团招标，编号26-GCZN-0150，"
        "投标报价128万元。售后服务响应时间为两小时，备件由本地仓库发放。"
        "郑州容新新能源有限公司不参与本章。"
        "河南伟泰光电科技有限公司负责巡检。"
    )
    out = sanitize_text(text, _brief())
    assert "售后服务响应时间为两小时" in out
    assert "郑州高新" not in out
    assert "某某建设" not in out
    assert "26-GCZN-0150" not in out
    assert "128" not in out
    assert "容新" not in out
    assert "河南伟泰" not in out
    assert "投标人" in out
    assert "（金额略）" in out


def test_docx_chapters_and_tags(tmp_path) -> None:
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("郑州高新新能源项目")
    doc.add_paragraph("一、投标函")
    doc.add_paragraph(
        "我公司承诺按招标文件完成本次投标，中标后按合同提供完整售后配合，"
        "质保期内免费更换损坏部件，响应安排见后文。"
    )
    doc.add_paragraph("（一）响应")
    doc.add_paragraph("接到通知后两小时内到场，不把这一节另起一章。")
    doc.add_paragraph("二、报价表")
    doc.add_paragraph("1 10000 20000 30000 40000")
    doc.add_paragraph("三、售后服务")
    doc.add_paragraph(
        "售后服务响应时间为两小时，备件供应周期不超过十五个工作日，"
        "质保期内免费更换损坏部件，并保留巡检记录。"
    )
    path = tmp_path / "bid.docx"
    doc.save(path)

    raw = chapters_from_docx(path)
    titles = [title for title, _body in raw]
    assert any(t.startswith("一、投标函") for t in titles)
    letter = next(body for title, body in raw if title.startswith("一、投标函"))
    assert "（一）响应" in letter
    assert "两小时内到场" in letter

    ready = prepare_chapters(raw, record_id=9, volume="business", brief=_brief())
    names = [item["title"] for item in ready]
    assert any("投标函" in name for name in names)
    assert any("售后服务" in name for name in names)
    assert not any("报价" in name for name in names)
    for item in ready:
        assert "tech_bid" in item["tags"]
        assert "record:9" in item["tags"]
        assert "volume:business" in item["tags"]
        assert "郑州高新" not in item["body"]
        assert "郑州高新" not in item["name"]


def test_bid_letter_name_still_vectorized() -> None:
    hit = classify_kb_document(name="投标函", text="我方承诺按合同履行。", tags=["tech_bid", "投标函"])
    assert hit.skip_vectorize is False


def test_outline_body_ignores_invitation() -> None:
    items = [
        OutlineItem(
            id="a",
            title="售后服务",
            rawTitle="三、售后服务",
            level=1,
            draftBody="售后服务响应时间为两小时，备件供应周期不超过十五个工作日，并保留巡检记录供查。",
            inviteExcerpt="招标人要求把郑州高新新能源项目写进本章。",
        ),
        OutlineItem(
            id="b",
            title="人员",
            level=2,
            body="设专职售后人员两名，负责巡检和故障处理，不单独成章。",
        ),
    ]
    chapters = chapters_from_outline(items)
    assert len(chapters) == 1
    assert chapters[0][0] == "三、售后服务"
    assert "两小时" in chapters[0][1]
    assert "专职售后" in chapters[0][1]
    assert "郑州高新" not in chapters[0][1]

    ready = prepare_chapters(chapters, record_id=3, volume="technical", brief=_brief())
    assert len(ready) == 1
    assert "volume:technical" in ready[0]["tags"]
    assert "record:3" in ready[0]["tags"]
    assert "郑州高新" not in ready[0]["body"]


def test_empty_outline_falls_back_to_docx(tmp_path) -> None:
    doc = Document()
    doc.add_paragraph("一、售后服务")
    doc.add_paragraph(
        "售后服务响应时间为两小时，备件供应周期不超过十五个工作日，"
        "质保期内免费更换损坏部件，并保留巡检记录。"
    )
    path = tmp_path / "tech.docx"
    doc.save(path)
    items = [
        OutlineItem(id="a", title="售后服务", rawTitle="一、售后服务", level=1, body="", inviteExcerpt="不要入库"),
    ]
    chapters = resolve_chapters(items, path)
    assert len(chapters) == 1
    assert "两小时" in chapters[0][1]
    assert "不要入库" not in chapters[0][1]


class _MemDb:
    def __init__(self) -> None:
        self.row = None

    def add(self, row) -> None:
        self.row = row

    async def commit(self) -> None:
        if self.row is not None and not getattr(self.row, "id", None):
            self.row.id = 41

    async def refresh(self, row) -> None:
        row.id = 41


@pytest.mark.asyncio
async def test_upload_without_record_returns_new_id(monkeypatch) -> None:
    held: dict = {}

    def fake_chapters(_path):
        return [(
            "一、售后服务",
            "售后服务响应时间为两小时，备件供应周期不超过十五个工作日，"
            "质保期内免费更换损坏部件，并保留巡检记录。",
        )]

    def fake_schedule(**kwargs):
        held["chapters"] = kwargs["chapters"]
        held["record_id"] = kwargs["record_id"]

    monkeypatch.setattr("api.services.tenders.bid_archive.chapters_from_docx", fake_chapters)
    monkeypatch.setattr("api.services.tenders.bid_archive.schedule_upload", fake_schedule)

    db = _MemDb()
    out = await accept_history_upload(
        db,
        data=b"not-empty",
        filename="某项目投标文件.docx",
        user_id=3,
        username="admin",
        outcome="lost",
        record_public_id="",
        admin=False,
    )
    assert out["recordId"]
    assert len(out["recordId"]) == 16
    assert out["queued"] == 1
    assert db.row.project_name == "某项目投标文件"
    assert db.row.status == "lost"
    assert db.row.docx_file == ""
    assert held["record_id"] == 41
    assert "record:41" in held["chapters"][0]["tags"]
    assert "outcome:lost" in held["chapters"][0]["tags"]
