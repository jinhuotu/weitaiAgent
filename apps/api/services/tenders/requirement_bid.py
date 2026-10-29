"""Build a submission outline from requirements when the supplied format is incomplete."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from api.services.tenders.schema import BidBrief, OutlineItem

_GROUPS = (
    ('qualification', '第一部分 资格证明'),
    ('response', '第二部分 商务与技术文件'),
    ('commitment', '第三部分 承诺函'),
    ('quote', '第四部分 报价文件'),
)
# These are submission sections, never copies of evaluation rules or buyer specifications.
_SECTIONS = (
    ('qualification', '法定代表人身份证明', 'legal_id', ('身份证明',)),
    ('qualification', '授权委托书', 'auth', ('授权委托',)),
    ('qualification', '营业执照', 'scan', ('营业执照',)),
    ('qualification', '财务状况证明', 'scan', ('财务', '审计报告')),
    ('qualification', '税收及社会保障资金证明', 'scan', ('税收', '社保', '社会保障')),
    ('qualification', '无重大违法记录声明', 'commitment_copy', ('重大违法',)),
    ('qualification', '资质证明', 'scan', ('资质',)),
    ('response', '投标函', 'letter', ('投标函', '响应函')),
    ('response', '采购需求商务偏离表', 'biz_dev', ('商务偏离', '商务偏差', '合同条款')),
    ('response', '技术方案', 'tech_plan', ('技术方案', '响应方案', '实施方案')),
    ('response', '质量保证方案', 'tech_chapter', ('质量',)),
    ('response', '安装实施方案', 'tech_chapter', ('安装', '实施')),
    ('response', '调试方案', 'tech_chapter', ('调试', '试运行')),
    ('response', '验收交付方案', 'tech_chapter', ('验收', '交付')),
    ('response', '售后服务方案', 'tech_chapter', ('售后', '保修', '质保')),
    ('response', '应急保障方案', 'tech_chapter', ('应急', '故障')),
    ('response', '技术偏离表', 'tech_dev', ('技术偏离', '技术偏差', '技术参数')),
    ('commitment', '投标承诺函', 'commitment_copy', ('承诺函', '承诺书')),
    ('quote', '投标报价及分项报价表', 'quote', ('报价', '开标一览表')),
)
_EXTRA_EVIDENCE = (
    ('医疗器械注册及备案证明', ('医疗器械注册', '医疗器械产品相关备案')),
    ('医疗器械经营及生产许可证明', ('医疗器械经营许可', '医疗器械生产许可')),
    ('救护车公告目录证明', ('救护车', '公告目录')),
    ('信用记录证明', ('信用中国', '信用记录')),
    ('类似项目业绩证明', ('类似业绩', '类似项目业绩')),
)


def invitation_text(brief: BidBrief) -> str:
    from api.services.tenders.assets import tender_invitations_dir

    if not re.fullmatch(r'[a-f0-9]{10,16}', brief.invitationId or ''):
        return ''
    path = tender_invitations_dir() / f'{brief.invitationId}.txt'
    return path.read_text(encoding='utf-8') if path.is_file() else ''


def format_completeness(brief: BidBrief) -> tuple[bool, list[str]]:
    """Require cover fields and an actual TOC skeleton, not tabs elsewhere in the document."""
    from api.services.tenders.pdf_format import _is_toc_heading, _word_format_file, toc_block_ids

    path = _word_format_file(brief.invitationId)
    if path is None:
        return False, ['封面要素', '目录骨架']
    doc = Document(str(path))
    block = toc_block_ids(doc)
    children = list(doc.element.body)
    toc_host = block[0] if block else None
    if not block:
        # Native Word tables of contents can be wrapped in a content control (w:sdt).
        for el in children:
            if el.tag != qn('w:sdt'):
                continue
            paragraphs = list(el.iter(qn('w:p')))
            if any(_is_toc_heading(''.join(n.text or '' for n in p.iter(qn('w:t'))))
                   for p in paragraphs):
                toc_host, block = el, paragraphs
                break
    front = children[:children.index(toc_host)] if toc_host is not None else children[:35]
    text = re.sub(r'\s+', '', ''.join(
        n.text or '' for el in front for n in el.iter(qn('w:t'))
    ))
    cover = ('投标文件' in text or '响应文件' in text) and '投标人' in text and any(
        word in text for word in ('项目名称', '项目编号', '招标编号', brief.projectName or '\0')
    )
    toc = any('PAGEREF' in el.xml or next(el.iter(qn('w:tab')), None) is not None
              for el in block[1:])
    missing = ([] if cover else ['封面要素']) + ([] if toc else ['目录骨架'])
    return not missing, missing


def _bucket(item: OutlineItem) -> str:
    from api.services.tenders.categories import title_goes_technical

    if item.kind == 'quote':
        return 'quote'
    if item.kind in {'tech_plan', 'tech_chapter', 'tech_dev', 'biz_dev', 'letter'}:
        return 'response'
    if title_goes_technical(item.title):
        return 'response'
    if item.kind == 'commitment_copy' and '重大违法' not in item.title:
        return 'commitment'
    return 'qualification'


def synthesize_outline(brief: BidBrief, text: str) -> list[str]:
    """Keep edits and evidence; add missing submission sections once, with stable IDs."""
    from api.services.tenders.categories import is_volume_label
    from api.services.tenders.classifier import classify_content
    from api.services.tenders.outline import compact_title

    group_titles = {compact_title(title) for _, title in _GROUPS}
    original = [item for item in brief.outlineItems
                if not item.id.startswith('req-group-') and not is_volume_label(item.title)
                and compact_title(item.title) not in group_titles]
    used: set[int] = set()
    groups: dict[str, list[OutlineItem]] = {key: [] for key, _ in _GROUPS}
    added = []
    # Full text is deliberately scanned beyond extract_requirements' 40-row cap.
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    lines.extend(req.text for req in brief.bidRequirements if req.text)
    specs = list(_SECTIONS)
    # Retain named submission declarations from the whole document, including
    # conditional policy forms. Their facts remain pending until confirmed.
    candidates = re.findall(r'《([^》]{2,40})》', text)
    quoted = set(candidates)
    candidates.extend(line for line in lines if len(line) <= 40)
    forms: dict[str, str] = {}
    form_keys = {compact_title(row[1]) for row in specs}
    for title in candidates:
        title = title.strip()
        pattern = r'[^：:。；;\n]{2,35}(?:声明函|承诺函|承诺书)(?:[（(]货物[)）])?'
        if not re.fullmatch(pattern, title):
            continue
        if title not in quoted and any(word in title for word in ('提供', '提交', '应当', '必须')):
            continue
        key = compact_title(title)
        if key in form_keys or key in {'投标人承诺函', '投标人承诺书'}:
            continue
        key = re.sub(r'[（(]货物[)）]$', '', key)
        key = key.replace('福利性单位', '福利单位').replace('福利企业', '福利单位')
        if len(title) > len(forms.get(key, '')):
            forms[key] = title
    named_forms = set(forms.values())
    for title in forms.values():
        bucket = ('qualification' if any(word in title for word in
                                        ('企业类型', '中小企业', '福利', '资格')) else 'commitment')
        specs.append((bucket, title, 'commitment_copy', (title,)))
    for title, words in _EXTRA_EVIDENCE:
        if any(any(word in line for word in words) for line in lines):
            specs.append(('qualification', title, 'scan', words))
    for req in brief.bidRequirements:
        title = req.targetTitle.strip()
        if (not title or len(title) > 40 or not req.targetKind
                or classify_content(title).role != 'chapter'
                or any(compact_title(title) == compact_title(row[1]) for row in specs)):
            continue
        item = OutlineItem(title=title, kind=req.targetKind)
        specs.append((_bucket(item), title, req.targetKind, (title,)))
    for index, (bucket, title, kind, words) in enumerate(specs):
        hit = next((item for item in original if id(item) not in used and (
            compact_title(item.title) == compact_title(title)
            or (item.kind == kind and (kind in {'auth', 'legal_id', 'letter', 'quote', 'biz_dev',
                                                'tech_dev', 'tech_plan'}
                                      or any(word in item.title for word in words)))
        )), None)
        evidence = list(dict.fromkeys(line for line in lines if any(w in line for w in words)))
        if hit is None:
            hit = OutlineItem(id=f'req-section-{index:02d}', title=title, kind=kind,
                              source='generate', inviteExcerpt='\n'.join(evidence[:8]))
            added.append(title)
        else:
            used.add(id(hit))
            hit = hit.model_copy(deep=True)
            if evidence and not hit.inviteExcerpt:
                hit.inviteExcerpt = '\n'.join(evidence[:8])
        if title in named_forms and not hit.rawTitle:
            hit.rawTitle = title
        hit.level = 3 if kind == 'tech_chapter' else 2
        groups[bucket].append(hit)
    for item in original:
        if id(item) in used:
            continue
        # Requirement/evaluation sentences must not become directory headings.
        if item.kind == 'unknown' and classify_content(item.title).role != 'chapter':
            continue
        if not item.title.strip():
            continue
        kept = item.model_copy(deep=True)
        kept.level = 2
        groups[_bucket(kept)].append(kept)
    brief.outlineItems = []
    for key, title in _GROUPS:
        brief.outlineItems.append(OutlineItem(
            id=f'req-group-{key}', title=title, kind='unknown', source='generate', level=1,
        ))
        brief.outlineItems.extend(groups[key])
    return added


def prepare_requirements_bid(brief: BidBrief, *, text: str | None = None) -> list[str]:
    """Opt in only for uploaded, requirement-bearing tenders with incomplete formats."""
    from api.services.tenders.native_format import prefer_native_format

    if prefer_native_format(brief):
        return []
    if not brief.requirementsDriven and not brief.invitationId:
        return []
    text = invitation_text(brief) if text is None else text
    if not brief.requirementsDriven:
        substantive = any(req.targetKind or req.category in {'qualification', 'material', 'scoring'}
                          for req in brief.bidRequirements)
        substantive = substantive or any(
            any(word in text for word in words) for _, _, _, words in _SECTIONS
        )
        if not substantive:
            return []
        complete, missing = format_completeness(brief)
        if complete:
            return []
    else:
        missing = []
    brief.requirementsDriven = True
    brief.layoutMode = 'requirements'
    added = synthesize_outline(brief, text)
    brief.documentFormat = brief.documentFormat.model_copy(update={
        'coverRequired': True, 'coverShowProject': True, 'coverShowTenderNo': True,
        'coverShowBidder': True, 'coverShowDate': True, 'coverNeedSeal': True,
        'tocNeedPageNos': True,
        'pageNumberPos': ('bottom-center' if brief.documentFormat.pageNumberPos == 'none'
                          else brief.documentFormat.pageNumberPos),
    })
    notes = ['格式章不完整，已按全文招标要求综合大纲并生成完整封面、目录和正文'
             + ('（缺少：' + '、'.join(missing) + '）' if missing else '')]
    if added:
        notes.append('按招标要求补齐章节：' + '、'.join(added))
    return notes


def volume_items(brief: BidBrief, volume: str) -> list[OutlineItem]:
    from api.services.tenders.categories import item_volume

    out = []
    group = None
    for item in brief.outlineItems:
        if item.id.startswith('req-group-'):
            group = item
            continue
        if item.skipped:
            continue
        if volume in {'business', 'technical'} and item_volume(
            kind=item.kind, title=item.title, key=item.id,
        ) != volume:
            continue
        if group is not None:
            out.append(group)
            group = None
        out.append(item)
    return out


def complete_requirements_docx(path: Path, brief: BidBrief, volume: str) -> list[str]:
    """Check generated structures and report missing facts/evidence without asserting compliance."""
    from api.services.tenders.document import _pageref_bookmark

    doc = Document(path)
    body = doc.element.body
    text = '\n'.join(p.text for p in doc.paragraphs)
    notes = ['综合组卷章节：' + '、'.join(item.title for item in volume_items(brief, volume))]
    missing = []
    if '投 标 文 件' not in text and '投标文件' not in text:
        missing.append('封面')
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    if not entries:
        missing.append('目录页码域')
    names = {n.get(qn('w:name')) for n in body.iter(qn('w:bookmarkStart'))}
    if any(_pageref_bookmark(p) not in names or not p.text.rsplit('\t', 1)[-1].isdigit()
           for p in entries):
        missing.append('目录页码/正文书签待核对')
    if not any('PAGE' in (n.text or '') for sec in doc.sections
               for n in sec.footer._element.iter(qn('w:instrText'))):
        missing.append('页码')
    if '盖章' not in text and '公章' not in text:
        missing.append('盖章位')
    if '签字' not in text and '签章' not in text:
        missing.append('签字位')
    if volume != 'technical':
        if not any(item.kind == 'quote' for item in volume_items(brief, volume)):
            missing.append('报价章节')
        if not brief.quoteLines:
            notes.append('【待补充】未提供分项报价清单，已保留报价章节；须核对投标总价及报价明细')
        if brief.bidPriceYuan <= 0:
            notes.append('【待补充】投标总价未填写，请核实后定稿')
    if missing:
        notes.append('完整性自检待核对：' + '、'.join(missing))
    else:
        notes.append('完整性结构自检：封面、目录页码域、正文书签、签字盖章位、连续页码'
                     + ('、报价章节齐全' if volume != 'technical' else '齐全'))
    pending = [item.title for item in volume_items(brief, volume)
               if not item.id.startswith('req-group-') and item.kind in {
                   'tech_plan', 'tech_chapter', 'commitment_copy',
               } and not item.draftBody.strip()]
    if pending:
        notes.append('【待补充】资料或响应正文仍需核验：' + '、'.join(pending))
    skipped = [item.title for item in brief.outlineItems if item.skipped and item.required]
    if skipped:
        notes.append('完整性自检：已手动跳过必需章节，请核对提交要求：' + '、'.join(skipped))
    notes.append('结构齐全不代表资格或响应已满足；证照、声明事实、签章和报价须核实后定稿')
    return notes
