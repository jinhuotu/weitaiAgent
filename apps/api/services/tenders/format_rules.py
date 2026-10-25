"""从邀请书抽出排版条款，并写入 Word 页边距 / 页码 / 目录编号。"""

from __future__ import annotations

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from api.services.tenders.schema import DocumentFormat, VolumeFormat

_CN_DIGITS = "一二三四五六七八九"

_CN_PT = {
    "初号": 42.0,
    "小初": 36.0,
    "一号": 26.0,
    "小一": 24.0,
    "二号": 22.0,
    "小二": 18.0,
    "三号": 16.0,
    "小四": 12.0,
    "四号": 14.0,
    "小三": 15.0,
    "五号": 10.5,
    "小五": 9.0,
    "六号": 7.5,
}

_PT_NUM = re.compile(r"(\d+(?:\.\d+)?)\s*(?:磅|pt|PT)")
_CN_SIZE = re.compile(r"(小初|初号|小一|一号|小二|二号|小三|三号|小四|四号|小五|五号|小六|六号)")
_TENDER_NO = re.compile(
    r"(?:招标编号|项目编号|采购编号|招标文件编号)\s*[:：为]?\s*([A-Za-z0-9][A-Za-z0-9\-_/]{2,31})"
)
_ALL_MARGIN = re.compile(
    r"页边距[^。\n]{0,24}(?:均为|全部为|都为)\s*(\d+(?:\.\d+)?)\s*(cm|CM|厘米|mm|MM|毫米)"
)
_NAMED_MARGIN = re.compile(
    r"(上|下|左|右)(?:边距|页边距)?\s*(?:为|:|：)?\s*(\d+(?:\.\d+)?)\s*(cm|CM|厘米|mm|MM|毫米)"
)
_QUAD_MARGIN = re.compile(
    r"(?:上|上边距)[^。\n]{0,8}?(\d+(?:\.\d+)?)\s*(?:cm|厘米|mm|毫米)"
    r"[^。\n]{0,12}?(?:下|下边距)[^。\n]{0,8}?(\d+(?:\.\d+)?)\s*(?:cm|厘米|mm|毫米)"
    r"[^。\n]{0,12}?(?:左|左边距)[^。\n]{0,8}?(\d+(?:\.\d+)?)\s*(?:cm|厘米|mm|毫米)"
    r"[^。\n]{0,12}?(?:右|右边距)[^。\n]{0,8}?(\d+(?:\.\d+)?)\s*(?:cm|厘米|mm|毫米)",
    re.I,
)


def extract_tender_no(text: str) -> str:
    match = _TENDER_NO.search(text or "")
    if not match:
        return ""
    return match.group(1).strip().rstrip("。；;，,")


def extract_document_format(text: str) -> DocumentFormat:
    raw = text or ""
    fmt = DocumentFormat()
    notes: list[str] = []
    blob = re.sub(r"\s+", "", raw)

    _apply_margins(fmt, raw, notes)
    _apply_fonts(fmt, raw, blob, notes)
    _apply_cover(fmt, blob, notes)
    _apply_toc(fmt, blob, notes)
    _apply_page_numbers(fmt, blob, notes)

    no = extract_tender_no(raw)
    if no and fmt.coverShowTenderNo:
        notes.append(f"招标编号 {no}")
    fmt.notes = notes
    fmt.specified = bool(notes)
    fmt.businessFormat, fmt.technicalFormat = _volume_formats(raw)
    return fmt


def cn_ordinal(n: int) -> str:
    """目录序号：十一、十二…，与固定模板「一、二、三」同一套汉字。"""
    if n <= 0:
        return str(n)
    if n < 10:
        return _CN_DIGITS[n - 1]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + _CN_DIGITS[n - 11]
    if n < 100:
        tens, ones = divmod(n, 10)
        head = _CN_DIGITS[tens - 1] + "十"
        return head if ones == 0 else head + _CN_DIGITS[ones - 1]
    return str(n)


def toc_item_label(index: int, scheme: str) -> str:
    n = index + 1
    cn = cn_ordinal(n)
    if scheme == "arabic":
        return f"{n}."
    if scheme == "paren":
        return f"({n})"
    if scheme == "attach":
        return f"附件{cn}"
    return f"{cn}、"


_TOC_CHAPTER = re.compile(r"^第[一二三四五六七八九十百零〇0-9]+章")


def flatten_toc_tree(
    nodes: list[dict], scheme: str = "cn"
) -> list[tuple[str, str, int, str]]:
    """树 → (label, title, level, bookmark)。第一章这类标题不再套 6.1。"""
    rows: list[tuple[str, str, int, str]] = []

    def walk(items: list[dict], level: int, prefix: list[int]) -> None:
        for i, node in enumerate(items or [], start=1):
            title = (node.get("title") or "").strip()
            if not title:
                continue
            bm = (node.get("bookmark") or "").strip()
            nums = prefix + [i]
            if node.get("bare") or _TOC_CHAPTER.match(title):
                label = ""
            elif level <= 1:
                label = toc_item_label(i - 1, scheme)
            else:
                label = ".".join(str(n) for n in nums) + " "
            rows.append((label, title, level, bm))
            kids = node.get("children") or []
            if kids:
                walk(kids, level + 1, nums)

    walk(nodes, 1, [])
    return rows


def toc_page_cache(index: int, fmt: DocumentFormat) -> str:
    """打开 Word 前的占位页码。封面不编时目录为第 1 页，正文条目从第 2 页起。"""
    if fmt.pageNumberStart == "body":
        return str(index + 1)
    if fmt.pageNumberStart == "cover":
        return str(index + 3)
    return str(index + 2)


_CANON_L1 = re.compile(r"^([一二三四五六七八九十]{1,3})、")
_CANON_L2 = re.compile(r"^[（(]([一二三四五六七八九十]{1,3})[)）]")
_CANON_DOT = re.compile(r"^(\d+(?:\.\d+)+)")
_TOC_LEAD = re.compile(
    r"^(?:"
    r"[一二三四五六七八九十]{1,3}、"
    r"|[（(][一二三四五六七八九十0-9]+[)）]"
    r"|\d+(?:\.\d+)+[、.．\s]*"
    r"|第[一二三四五六七八九十百0-9]+[章节篇][、.\s]*"
    r"|[A-Ha-h][.．、]"
    r"|附(?:件)?[（(]?[一二三四五六七八九十0-9]*[)）]?[：:、.\s]*"
    r")"
)


def title_canon(text: str) -> tuple[int, str]:
    """规范编号：一级「一、」、二级「（一）」、三级「1.1」。不合这三套的返回 0。"""
    raw = (text or "").strip()
    m = _CANON_L1.match(raw)
    if m:
        return 1, m.group(1)
    m = _CANON_L2.match(raw)
    if m:
        return 2, m.group(1)
    m = _CANON_DOT.match(raw)
    if m:
        return 3, m.group(1)
    return 0, ""


def peel_toc_lead(text: str) -> str:
    raw = (text or "").strip()
    peeled = _TOC_LEAD.sub("", raw).strip(" ：:、.")
    return peeled or raw


def _bump_level(counts: list[int], level: int) -> None:
    level = max(1, min(4, level))
    while len(counts) < level:
        counts.append(0)
    del counts[level:]
    counts[level - 1] += 1


def _free_cn(used: set[str]) -> str:
    n = 1
    while cn_ordinal(n) in used and n < 80:
        n += 1
    label = cn_ordinal(n)
    used.add(label)
    return label


def toc_numbered_lines(rows: list[tuple[int, str]]) -> list[str]:
    """目录显示文字。已有规范编号的原样沿用；裸行和不合规范的前缀按层级补号。"""
    used: dict[int, set[str]] = {1: set(), 2: set(), 3: set()}
    kinds: list[tuple[int, str]] = []
    for _level, title in rows:
        kind, key = title_canon(title)
        kinds.append((kind, key))
        if kind:
            used[kind].add(key)
    counts: list[int] = []
    out: list[str] = []
    for (level, title), (kind, _key) in zip(rows, kinds, strict=True):
        level = max(1, min(4, int(level or 1)))
        raw = (title or "").strip()
        _bump_level(counts, level)
        if kind and raw:
            out.append(raw)
            continue
        bare = peel_toc_lead(raw)
        if level <= 1:
            out.append(f"{_free_cn(used[1])}、{bare}")
        elif level == 2:
            out.append(f"（{_free_cn(used[2])}）{bare}")
        else:
            parent = 1
            for n in counts[:-1]:
                if n:
                    parent = n
            slot = used[3]
            i = 1
            while f"{parent}.{i}" in slot and i < 80:
                i += 1
            key = f"{parent}.{i}"
            slot.add(key)
            out.append(f"{key} {bare}".rstrip())
    return out


def section_page_flags(
    section_index: int, fmt: DocumentFormat, *, body_section: int
) -> tuple[bool, bool]:
    """返回 (是否编页码, 是否从 1 重新起编)。"""
    if fmt.pageNumberPos == "none":
        return False, False
    if fmt.pageNumberStart == "cover":
        return True, section_index == 0
    if fmt.pageNumberStart == "body":
        return section_index >= body_section, section_index == body_section
    return section_index >= 1, section_index == 1


def apply_section_page(section, fmt: DocumentFormat) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(fmt.marginLeftCm)
    section.right_margin = Cm(fmt.marginRightCm)
    section.top_margin = Cm(fmt.marginTopCm)
    section.bottom_margin = Cm(fmt.marginBottomCm)
    section.header_distance = Cm(1.75)
    section.footer_distance = Cm(1.5)


def apply_footer_page_number(
    section,
    fmt: DocumentFormat,
    *,
    numbered: bool,
    restart: bool,
) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    _clear_runs(para)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    _set_pg_num_type(section, start=1 if restart and numbered else None)
    if not numbered or fmt.pageNumberPos == "none":
        if getattr(section, "different_first_page_header_footer", False):
            first = section.first_page_footer
            first.is_linked_to_previous = False
            first_para = first.paragraphs[0] if first.paragraphs else first.add_paragraph()
            _clear_runs(first_para)
        return
    para.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
        if fmt.pageNumberPos == "bottom-right"
        else WD_ALIGN_PARAGRAPH.CENTER
    )
    _add_page_field(para, font=fmt.fontName, size_pt=10.5)


def format_brief_notes(fmt: DocumentFormat) -> list[str]:
    if fmt.specified and fmt.notes:
        lines = ["已按招标书抽出排版要求：" + "；".join(fmt.notes[:8])]
    else:
        lines = [
            "招标书未写页边距/字体/页码等排版条款，"
            "生成用默认 A4、宋体，封面和目录不编页码、从正文起页底居中编码"
        ]
    for label, spec in (("商务标", fmt.businessFormat), ("技术标", fmt.technicalFormat)):
        if spec.notes:
            lines.append(f"{label}写明：" + "；".join(spec.notes[:6]))
    return lines


_SIGN_MARK = ("盖章", "签字", "签章", "法定代表人", "委托代理人")
_SHORT_TITLE = (
    "技术标",
    "技术方案",
    "技术偏离",
    "技术偏差",
    "施工组织",
    "技术规格",
    "暗标",
    "技术要求",
    "投标函",
    "报价",
    "授权委托",
    "身份证明",
    "营业执照",
    "商务偏离",
    "商务偏差",
    "业绩",
    "法定代表人",
    "目录",
)
_HEADISH = re.compile(
    r"^(?:第[一二三四五六七八九十0-9]+[章节]|[一二三四五六七八九十]、)"
)
_VOL_TECH = re.compile(r"技术标|技术部分|技术文件|暗标")
_VOL_BIZ = re.compile(r"商务标|商务部分|商务文件")
_RULE_WORD = re.compile(r"页边距|字体|正文|页码|磅|小四|小三|四号|三号")


def apply_volume_format(doc, spec: VolumeFormat, *, where: str = "副本") -> list[str]:
    """只改这一卷写明的页边距、正文和页码。表格、签章栏不动。"""
    if spec is None:
        return []
    notes: list[str] = []
    if spec.marginSet:
        for section in doc.sections:
            section.left_margin = Cm(spec.marginLeftCm)
            section.right_margin = Cm(spec.marginRightCm)
            section.top_margin = Cm(spec.marginTopCm)
            section.bottom_margin = Cm(spec.marginBottomCm)
        notes.append(
            f"已按标书改{where}页边距"
            f" 上{spec.marginTopCm:g} 下{spec.marginBottomCm:g}"
            f" 左{spec.marginLeftCm:g} 右{spec.marginRightCm:g}cm"
        )
    if spec.fontSet or spec.bodySizeSet:
        before = _sample_body_font(doc)
        n = _paint_body_runs(doc, spec)
        if n:
            want = _spec_font_label(spec)
            if before:
                notes.append(f"范本正文{before}，标书要求{want}，已改{where}")
            else:
                notes.append(f"已按标书把{where}正文改成{want}，表格和签章栏未动")
    if spec.pagePosSet:
        _apply_page_pos(doc, spec.pageNumberPos or "bottom-center", spec.fontName or "宋体")
        notes.append(f"已按标书改{where}页码位置")
    if spec.pageStartSet:
        _apply_page_start(doc, spec.pageNumberStart or "body")
        notes.append(f"已按标书改{where}页码起编")
    return notes


def _volume_formats(raw: str) -> tuple[VolumeFormat, VolumeFormat]:
    biz = VolumeFormat()
    tech = VolumeFormat()
    scope = "both"
    for line in re.split(r"[\r\n]+", raw or ""):
        compact = re.sub(r"\s+", "", line)
        if not compact:
            continue
        if len(compact) <= 40 and not _RULE_WORD.search(compact):
            head = _line_scope(compact)
            if head:
                scope = head
                continue
            if _HEADISH.match(compact):
                scope = "both"
                continue
        parts = [line]
        if "商务" in compact and ("技术" in compact or "暗标" in compact):
            parts = [p for p in re.split(r"[。；;，,]", line) if p.strip()]
        for part in parts:
            part_scope = _line_scope(re.sub(r"\s+", "", part)) or scope
            _absorb_volume(part, part_scope, biz, tech)
    return biz, tech


def _line_scope(compact: str) -> str:
    tech = bool(_VOL_TECH.search(compact)) and "商务" not in compact
    biz = bool(_VOL_BIZ.search(compact)) and "技术" not in compact and "暗标" not in compact
    if tech and not biz:
        return "technical"
    if biz and not tech:
        return "business"
    return ""


def _absorb_volume(text: str, scope: str, biz: VolumeFormat, tech: VolumeFormat) -> None:
    targets: list[VolumeFormat] = []
    if scope in {"both", "business"}:
        targets.append(biz)
    if scope in {"both", "technical"}:
        targets.append(tech)
    if not targets:
        return
    tmp = DocumentFormat()
    notes: list[str] = []
    blob = re.sub(r"\s+", "", text)
    _apply_margins(tmp, text, notes)
    _apply_fonts(tmp, text, blob, notes)
    _apply_page_numbers(tmp, blob, notes)
    margin_note = next((n for n in notes if "页边距" in n), "")
    font_note = next((n for n in notes if n.startswith("正文字体")), "")
    size_note = next((n for n in notes if n.startswith("正文 ") or n.startswith("正文小四")), "")
    pos_note = next(
        (
            n
            for n in notes
            if n.startswith("页码") and "连续" not in n and "起编" not in n and "不编" not in n
        ),
        "",
    )
    start_note = next((n for n in notes if "起编" in n or "不编" in n or "自封面" in n), "")
    overwrite = scope != "both"
    for spec in targets:
        if margin_note and (overwrite or not spec.marginSet):
            spec.marginSet = True
            spec.marginLeftCm = tmp.marginLeftCm
            spec.marginRightCm = tmp.marginRightCm
            spec.marginTopCm = tmp.marginTopCm
            spec.marginBottomCm = tmp.marginBottomCm
            _put_note(spec, "页边距", margin_note)
        if font_note and (overwrite or not spec.fontSet):
            spec.fontSet = True
            spec.fontName = tmp.fontName
            _put_note(spec, "正文字体", font_note)
        if size_note and (overwrite or not spec.bodySizeSet):
            spec.bodySizeSet = True
            spec.bodySizePt = tmp.bodySizePt
            _put_note(spec, "正文", size_note)
        if pos_note and (overwrite or not spec.pagePosSet):
            spec.pagePosSet = True
            spec.pageNumberPos = tmp.pageNumberPos
            _put_note(spec, "页码", pos_note)
        if start_note and (overwrite or not spec.pageStartSet):
            spec.pageStartSet = True
            spec.pageNumberStart = tmp.pageNumberStart
            _put_note(spec, "起编", start_note)


def _put_note(spec: VolumeFormat, kind: str, text: str) -> None:
    kept: list[str] = []
    for note in spec.notes:
        if kind == "页边距" and note.startswith("页边距"):
            continue
        if kind == "正文字体" and note.startswith("正文字体"):
            continue
        if kind == "正文" and (note.startswith("正文 ") or note.startswith("正文小四")):
            continue
        if kind == "页码" and note.startswith("页码"):
            continue
        if kind == "起编" and ("起编" in note or "不编" in note or "自封面" in note):
            continue
        kept.append(note)
    kept.append(text)
    spec.notes = kept


def _body_para(text: str) -> bool:
    n = re.sub(r"\s+", "", text or "")
    if not n or any(k in n for k in _SIGN_MARK):
        return False
    if len(n) <= 40 and _HEADISH.match(n):
        return False
    if len(n) <= 18 and any(k in n for k in _SHORT_TITLE):
        return False
    return True


def _sample_body_font(doc) -> str:
    for para in doc.paragraphs:
        if not _body_para(para.text):
            continue
        for run in para.runs:
            name = _run_east_asia(run)
            size = run.font.size.pt if run.font.size is not None else None
            if name and size:
                return f"{name} {size:g}磅"
            if name:
                return name
            if size:
                return f"{size:g}磅"
    return ""


def _spec_font_label(spec: VolumeFormat) -> str:
    bits: list[str] = []
    if spec.fontSet and spec.fontName:
        bits.append(spec.fontName)
    if spec.bodySizeSet:
        bits.append(f"{spec.bodySizePt:g}磅")
    return " ".join(bits) or "标书字体"


def _paint_body_runs(doc, spec: VolumeFormat) -> int:
    n = 0
    for para in doc.paragraphs:
        if not _body_para(para.text) or not para.runs:
            continue
        for run in para.runs:
            if spec.fontSet and spec.fontName:
                run.font.name = spec.fontName
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.get_or_add_rFonts()
                rfonts.set(qn("w:ascii"), spec.fontName)
                rfonts.set(qn("w:hAnsi"), spec.fontName)
                rfonts.set(qn("w:eastAsia"), spec.fontName)
            if spec.bodySizeSet:
                run.font.size = Pt(spec.bodySizePt)
        n += 1
    return n


def _run_east_asia(run) -> str:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            name = rfonts.get(qn("w:eastAsia")) or rfonts.get(qn("w:ascii")) or ""
            if name:
                return name
    return run.font.name or ""


def _apply_page_pos(doc, pos: str, font: str) -> None:
    if pos not in {"bottom-center", "bottom-right"}:
        return
    align = WD_ALIGN_PARAGRAPH.RIGHT if pos == "bottom-right" else WD_ALIGN_PARAGRAPH.CENTER
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        _clear_runs(para)
        para.alignment = align
        _add_page_field(para, font=font or "宋体", size_pt=10.5)


def _apply_page_start(doc, start: str) -> None:
    if start not in {"cover", "toc", "body"}:
        return
    sections = list(doc.sections)
    if not sections:
        return
    if len(sections) == 1:
        _set_pg_num_type(sections[0], start=1)
        return
    tmp = DocumentFormat(pageNumberStart=start)
    body_si = 0 if start == "cover" else 1
    for i, section in enumerate(sections):
        numbered, restart = section_page_flags(i, tmp, body_section=body_si)
        if numbered and restart:
            _set_pg_num_type(section, start=1)


def _to_cm(value: str, unit: str) -> float:
    num = float(value)
    key = (unit or "cm").lower()
    if key in {"mm", "毫米"}:
        num = num / 10.0
    return max(1.0, min(5.0, round(num, 2)))


def _apply_margins(fmt: DocumentFormat, raw: str, notes: list[str]) -> None:
    all_m = _ALL_MARGIN.search(raw)
    if all_m:
        cm = _to_cm(all_m.group(1), all_m.group(2))
        fmt.marginTopCm = fmt.marginBottomCm = fmt.marginLeftCm = fmt.marginRightCm = cm
        notes.append(f"页边距均为 {cm:g}cm")
        return
    quad = _QUAD_MARGIN.search(raw)
    if quad:
        fmt.marginTopCm = _to_cm(quad.group(1), "cm")
        fmt.marginBottomCm = _to_cm(quad.group(2), "cm")
        fmt.marginLeftCm = _to_cm(quad.group(3), "cm")
        fmt.marginRightCm = _to_cm(quad.group(4), "cm")
        notes.append(
            f"页边距 上{fmt.marginTopCm:g} 下{fmt.marginBottomCm:g} "
            f"左{fmt.marginLeftCm:g} 右{fmt.marginRightCm:g}cm"
        )
        return
    found = False
    mapping = {
        "上": "marginTopCm",
        "下": "marginBottomCm",
        "左": "marginLeftCm",
        "右": "marginRightCm",
    }
    for match in _NAMED_MARGIN.finditer(raw):
        setattr(fmt, mapping[match.group(1)], _to_cm(match.group(2), match.group(3)))
        found = True
    if found:
        notes.append(
            f"页边距 上{fmt.marginTopCm:g} 下{fmt.marginBottomCm:g} "
            f"左{fmt.marginLeftCm:g} 右{fmt.marginRightCm:g}cm"
        )


def _pt_from_blob(blob: str) -> float | None:
    named = _CN_SIZE.search(blob)
    if named:
        return _CN_PT.get(named.group(1))
    num = _PT_NUM.search(blob)
    if num:
        return max(8.0, min(42.0, float(num.group(1))))
    return None


def _apply_fonts(fmt: DocumentFormat, raw: str, blob: str, notes: list[str]) -> None:
    fonts = "仿宋_GB2312|楷体_GB2312|宋体|仿宋|黑体|楷体|微软雅黑|Times New Roman"
    font = re.search(
        r"(?:字体|正文|全文|采用)[^。\n]{0,20}(" + fonts + r")"
        r"|(" + fonts + r")(?:、|/|和)?(?:小?[一二三四五]号|\d+(?:\.\d+)?\s*(?:磅|pt))",
        raw,
    )
    if font:
        fmt.fontName = next(g for g in font.groups() if g)
        notes.append(f"正文字体 {fmt.fontName}")
    body_ctx = re.search(r"(?:正文|全文)[^。\n]{0,24}(?:字体|字号|用)?[^。\n]{0,16}", raw)
    if body_ctx:
        pt = _pt_from_blob(body_ctx.group(0))
        if pt:
            fmt.bodySizePt = pt
            notes.append(f"正文 {pt:g} 磅")
    elif "正文" in blob and "小四" in blob:
        fmt.bodySizePt = 12.0
        notes.append("正文小四")
    head_ctx = re.search(r"(?:一级标题|章节标题|标题字号|标题用)[^。\n]{0,24}", raw)
    if head_ctx:
        pt = _pt_from_blob(head_ctx.group(0))
        if pt:
            fmt.headingSizePt = pt
            notes.append(f"标题 {pt:g} 磅")
    cover_ctx = re.search(r"封面[^。\n]{0,30}(?:标题|字体|字号)", raw)
    if cover_ctx:
        pt = _pt_from_blob(cover_ctx.group(0))
        if pt:
            fmt.coverTitleSizePt = pt
            notes.append(f"封面标题 {pt:g} 磅")


def _near_cover(blob: str, *needles: str) -> bool:
    for needle in needles:
        if re.search(rf"封面.{{0,48}}{re.escape(needle)}|{re.escape(needle)}.{{0,48}}封面", blob):
            return True
    return False


def _apply_cover(fmt: DocumentFormat, blob: str, notes: list[str]) -> None:
    if "封面" not in blob:
        return
    fmt.coverRequired = True
    extra: list[str] = []
    if _near_cover(blob, "项目全称", "项目名称"):
        fmt.coverShowProject = True
    if _near_cover(blob, "招标编号", "项目编号", "采购编号"):
        fmt.coverShowTenderNo = True
        extra.append("招标编号")
    if _near_cover(blob, "投标人全称", "投标人名称", "供应商名称"):
        fmt.coverShowBidder = True
    if _near_cover(blob, "正本", "副本"):
        fmt.coverShowCopyMark = True
        if _near_cover(blob, "副本") and not _near_cover(blob, "正本"):
            fmt.coverCopyMark = "副本"
        else:
            fmt.coverCopyMark = "正本"
        extra.append(fmt.coverCopyMark)
    if _near_cover(blob, "投标日期", "日期"):
        fmt.coverShowDate = True
    if _near_cover(blob, "加盖公章", "盖章", "公章") or "封面加盖" in blob:
        fmt.coverNeedSeal = True
        extra.append("加盖公章")
    if extra:
        notes.append("封面含" + "、".join(extra))


def _apply_toc(fmt: DocumentFormat, blob: str, notes: list[str]) -> None:
    if "目录" not in blob:
        return
    if re.search(r"目录[^。]{0,20}(?:按附件|采用附件)|目录按附件", blob):
        fmt.tocNumbering = "attach"
        notes.append("目录按附件编号")
    elif re.search(r"目录[^。]{0,20}阿拉伯|目录[^。]{0,12}(?:用|按|采用)[0-9]", blob):
        fmt.tocNumbering = "arabic"
        notes.append("目录用1.编号")
    elif re.search(r"目录[^。]{0,16}(?:用|按|采用)[（(]1[)）]", blob):
        fmt.tocNumbering = "paren"
        notes.append("目录用(1)编号")
    toc_page_keys = ("目录页码", "页码一一对应", "不能错页", "页码对应", "目录与正文页码")
    if any(k in blob for k in toc_page_keys):
        fmt.tocNeedPageNos = True
        notes.append("目录页码与正文对应")


def _apply_page_numbers(fmt: DocumentFormat, blob: str, notes: list[str]) -> None:
    has_right = any(k in blob for k in ("右下角", "页脚右侧", "页底右侧")) or bool(
        re.search(r"页码.{0,8}右", blob)
    )
    has_center = any(k in blob for k in ("页底居中", "页脚居中", "底部居中", "页码居中"))
    if has_right and not has_center:
        fmt.pageNumberPos = "bottom-right"
        notes.append("页码右下角")
    elif has_center:
        fmt.pageNumberPos = "bottom-center"
        notes.append("页码页底居中")
    cover_unnumbered = any(k in blob for k in ("封面不编", "扉页不编"))
    start_body = any(k in blob for k in ("从正文起编", "正文起编"))
    start_toc = any(k in blob for k in ("从目录起编", "目录起编"))
    start_cover = any(k in blob for k in ("封面编页", "自封面", "从封面起编"))
    if cover_unnumbered:
        if start_body and not start_toc:
            fmt.pageNumberStart = "body"
            notes.append("封面/目录不编，从正文起编")
        else:
            fmt.pageNumberStart = "toc"
            notes.append("封面不编页码，从目录起编")
    elif start_body and not start_toc:
        fmt.pageNumberStart = "body"
        notes.append("封面/目录不编，从正文起编")
    elif start_toc:
        fmt.pageNumberStart = "toc"
        notes.append("封面不编页码，从目录起编")
    elif start_cover:
        fmt.pageNumberStart = "cover"
        notes.append("自封面连续编页")
    elif any(k in blob for k in ("连续编码", "不能跳号", "不得跳号", "不得缺页")):
        notes.append("页码连续编码")


def _set_pg_num_type(section, *, start: int | None) -> None:
    sp = section._sectPr
    for old in list(sp.findall(qn("w:pgNumType"))):
        sp.remove(old)
    if start is None:
        return
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), str(int(start)))
    sp.append(pg)


def _clear_runs(para) -> None:
    for child in list(para._p):
        if child.tag == qn("w:r"):
            para._p.remove(child)


def _add_page_field(para, *, font: str, size_pt: float) -> None:
    begin = para.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    begin._r.append(fld)
    _set_run_font(begin, font, size_pt)

    instr = para.add_run()
    text = OxmlElement("w:instrText")
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = " PAGE "
    instr._r.append(text)
    _set_run_font(instr, font, size_pt)

    sep = para.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    sep._r.append(fld_sep)
    _set_run_font(sep, font, size_pt)

    cache = para.add_run("1")
    _set_run_font(cache, font, size_pt)

    end = para.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end._r.append(fld_end)
    _set_run_font(end, font, size_pt)


def _set_run_font(run, font: str, size_pt: float) -> None:
    run.font.size = Pt(size_pt)
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)
