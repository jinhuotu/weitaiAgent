import asyncio
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders import qa, records
from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.classifier import apply_volume_binding
from api.services.tenders.generate import generate_pair
from api.services.tenders.schema import BidBrief, OutlineItem
from db.models.tender import TenderRecord


def test_full_report_stores_separately_and_checks_both_pools(tmp_path, monkeypatch):
    monkeypatch.setattr(qa, 'tenders_output_dir', lambda: tmp_path)
    pid = 'abc123def4567890'
    for volume in ('business', 'technical', 'full'):
        qa.save_qa_report(pid, {'similarityScore': 90, 'checkedAt': 1}, volume=volume)
        assert qa.load_qa_report(pid, volume=volume)['volume'] == volume
    assert qa.normalize_qa_volume(' FULL ') == 'full'
    qa.unlink_qa_report(pid, volume='full')
    assert qa.load_qa_report(pid, volume='full') is None
    assert qa.load_qa_report(pid, volume='business')
    brief = BidBrief(outlineItems=[
        OutlineItem(id='letter', title='投标函', kind='letter', required=True),
        OutlineItem(id='plan', title='施工组织设计', kind='tech_plan', required=True),
    ])
    gaps = qa.rule_inspect(brief, '项目正文', volume='full')['missing']
    titles = {g['title'] for g in gaps}
    assert any('商务偏离表' in title for title in titles)
    assert any('技术方案' in title for title in titles)


@pytest.mark.parametrize('requested', [None, 'business', 'technical', 'full'])
def test_full_record_qa_always_targets_one_full_doc(requested):
    brief = BidBrief(generateVolume='full')
    row = SimpleNamespace(docx_file='full.docx', tech_docx_file='old-tech.docx')
    assert records._record_qa_volume(brief, requested) == 'full'
    assert records._qa_generated_file(row, brief, requested) == 'full.docx'


def test_nonfull_qa_fallback_kept():
    row = SimpleNamespace(docx_file='', tech_docx_file='tech.docx')
    assert records._qa_generated_file(row, BidBrief()) == 'tech.docx'


def test_full_regeneration_clears_previous_technical_file_reference():
    row = TenderRecord(id=1, public_id='abc123def4567890', user_id=1, username='a',
                       tech_docx_file='old-tech.docx', tech_download_name='技术标.docx')
    records._apply_generated_files(row, BidBrief(generateVolume='full'),
                                   {'docxFile': 'full.docx', 'downloadName': '综合标.docx'})
    assert row.docx_file == 'full.docx'
    assert row.tech_docx_file is None
    assert row.tech_download_name is None


def test_manual_split_slices_even_when_binding_advises_full(tmp_path, monkeypatch):
    seen = []

    def fake_generate(brief, **kwargs):
        seen.append((brief.generateVolume, brief.splitSlice))
        return {}

    monkeypatch.setattr('api.services.tenders.generate.generate_bid', fake_generate)
    generate_pair(BidBrief(volumeAdvice='full', generateBoth=True, splitSlice=True))
    assert seen == [('business', True), ('technical', False)]


@pytest.mark.parametrize('volume', ['full', 'business', 'technical'])
def test_explicit_volume_survives_binding_judgment(volume):
    result = apply_volume_binding(BidBrief(generateVolume=volume),
                                 '商务标和技术标分别装订', keep_volume=True)
    assert result.generateVolume == volume


def test_full_assemble_finalizes_both_drafts_toc_and_section_pages(tmp_path, monkeypatch):
    """原版格式章保留表格、字体、章节顺序；后填正文后重算目录。"""
    src = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('商务标')
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph('目录')
    doc.add_paragraph('商务标')
    doc.add_paragraph('一、商务响应')
    doc.add_paragraph('技术标')
    doc.add_paragraph('二、施工组织设计')
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph('一、商务响应')
    doc.add_paragraph('商务原文')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '报价原表不改'
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph('技术标')
    doc.add_paragraph('二、施工组织设计')
    doc.add_paragraph('技术原文')
    for sec in doc.sections[2:]:
        num = OxmlElement('w:pgNumType')
        num.set(qn('w:start'), '1')
        sec._sectPr.append(num)
    doc.save(src)
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: src)
    brief = BidBrief(layoutMode='outline', invitationId='abc123def456', generateVolume='full',
                     autoDraftBizResp=True, autoDraftTechPlan=True, projectName='本项目',
                     attachQualifications=False)
    calls = []
    from api.services.tenders import pdf_format

    toc_calls = []
    for name in ('_link_toc_body', '_repair_format_table_toc'):
        original = getattr(pdf_format, name)

        def tracked(*args, _name=name, _original=original, **kwargs):
            toc_calls.append(_name)
            return _original(*args, **kwargs)

        monkeypatch.setattr(pdf_format, name, tracked)

    def fill(path, brief, label):
        calls.append(label)
        document = Document(path)
        document.add_paragraph(label)
        document.save(path)
        return 1, []

    monkeypatch.setattr('api.services.tenders.assemble._fill_drafted_format_sections',
                        lambda path, brief, **kw: fill(path, brief, '已起草技术正文'))
    monkeypatch.setattr('api.services.tenders.assemble._fill_business_format_sections',
                        lambda path, brief, **kw: fill(path, brief, '已起草商务正文'))
    path, _ = assemble_bid_docx(brief, tmp_path / 'full.docx', volume='full',
                                catalog_slots=[], catalog_media={}, catalog_facts=[])
    out = Document(path)
    text = '\n'.join(p.text for p in out.paragraphs)
    assert calls == ['已起草技术正文', '已起草商务正文']
    assert '已起草商务正文' in text and '已起草技术正文' in text
    assert '商务标' not in text and '技术标' not in text
    assert '投标文件' in text
    assert out.tables[0].cell(0, 0).text == '报价原表不改'
    starts = [i for i, sec in enumerate(out.sections)
              if any(n.get(qn('w:start')) is not None for n in sec._sectPr.findall(qn('w:pgNumType')))]
    assert starts == [2]
    refs = [p for p in out.paragraphs if 'PAGEREF' in p._p.xml]
    assert len(refs) == 2
    assert [p.text.rsplit('\t', 1)[-1] for p in refs] == ['1', '2']
    names = [n.get(qn('w:name')) for n in out.element.body.iter(qn('w:bookmarkStart'))]
    assert len(names) == len(set(names))
    assert all(sec.header.paragraphs[0].text == '本项目' for sec in out.sections)
    assert toc_calls == ['_link_toc_body', '_repair_format_table_toc']
    signature = _full_doc_signature(path)
    for _ in range(2):
        pdf_format.unify_full_docx(path, brief)
        assert _full_doc_signature(path) == signature


def test_full_saved_report_ignores_previous_single_volume_report(tmp_path, monkeypatch):
    monkeypatch.setattr(qa, 'tenders_output_dir', lambda: tmp_path)
    pid = 'abc123def4567890'
    qa.save_qa_report(pid, {'similarityScore': 70, 'checkedAt': 20}, volume='technical')
    qa.save_qa_report(pid, {'similarityScore': 99, 'checkedAt': 10}, volume='full')
    row = SimpleNamespace(public_id=pid, docx_file='full.docx', tech_docx_file='old.docx',
                          brief_json={'generateVolume': 'full'})

    async def get_row(*args):
        return row

    monkeypatch.setattr(records, '_get_row', get_row)
    for volume in (None, 'business', 'technical', 'full'):
        report = asyncio.run(records.get_record_qa(None, pid, volume=volume))['report']
        assert report['volume'] == 'full'
        assert report['similarityScore'] == 99


def _full_doc_signature(path):
    doc = Document(path)
    instructions = tuple(n.text for n in doc.element.body.iter(qn('w:instrText')))
    starts = tuple((n.get(qn('w:name')), n.get(qn('w:id')))
                   for n in doc.element.body.iter(qn('w:bookmarkStart')))
    ends = tuple(n.get(qn('w:id')) for n in doc.element.body.iter(qn('w:bookmarkEnd')))
    assert len(starts) == len({name for name, _ in starts})
    assert len(starts) == len({ident for _, ident in starts})
    assert len(ends) == len(set(ends))
    assert {ident for _, ident in starts} == set(ends)
    for para in doc.paragraphs:
        fields = [n.text for n in para._p.iter(qn('w:instrText'))
                  if 'PAGEREF' in (n.text or '')]
        if fields:
            assert len(fields) == 1
    return (tuple(p.text for p in doc.paragraphs), instructions, starts,
            ends,
            tuple(sec._sectPr.xml for sec in doc.sections),
            tuple(sec.footer._element.xml for sec in doc.sections))


@pytest.mark.parametrize('unmatched', [False, True])
def test_full_floating_table_toc_is_idempotent(tmp_path, unmatched):
    from api.services.tenders.document import _bookmark_paragraph, _pageref_bookmark
    from api.services.tenders.pdf_format import unify_full_docx

    doc = Document()
    doc.add_paragraph('目录')
    doc.add_paragraph('六、项目管理机构')
    doc.add_paragraph('项目管理机构组成表')
    if unmatched:
        doc.add_paragraph('尚未上传的附件资料')
    doc.add_page_break()
    doc.add_paragraph('六、项目管理机构')
    carrier = doc.add_paragraph()
    pict = OxmlElement('w:pict')
    content = OxmlElement('w:txbxContent')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '姓名与任职信息' * 12
    content.append(table._tbl)
    pict.append(content)
    carrier.add_run()._r.append(pict)
    carrier.add_run('项目管理机构组成表')
    _bookmark_paragraph(carrier, '(一)项目管理机构组成表')
    brief = BidBrief(generateVolume='full', outlineItems=[
        OutlineItem(id='organization', title='项目管理机构组成表'),
    ])
    path = tmp_path / 'floating-full.docx'
    doc.save(path)
    unify_full_docx(path, brief)
    signature = _full_doc_signature(path)
    out = Document(path)
    entries = [p for p in out.paragraphs if _pageref_bookmark(p)]
    assert len(entries) == (3 if unmatched else 2)
    assert entries[1].text.rsplit('\t', 1)[-1].isdigit()
    for _ in range(2):
        unify_full_docx(path, brief)
        assert _full_doc_signature(path) == signature


def test_full_fixed_layout_keeps_its_existing_toc_path(tmp_path, monkeypatch):
    from api.services.tenders import pdf_format

    def unexpected(*args, **kwargs):
        pytest.fail('Fixed layout must not run the native-format TOC linker again')

    monkeypatch.setattr(pdf_format, '_link_toc_body', unexpected)
    monkeypatch.setattr(pdf_format, '_repair_format_table_toc', unexpected)
    brief = BidBrief(generateVolume='full', attachQualifications=False,
                     includePlaceholders=False, includeCommitment=False)
    path, _ = assemble_bid_docx(brief, tmp_path / 'fixed-full.docx', volume='full',
                                catalog_slots=[], catalog_media={}, catalog_facts=[])
    refs = [n for n in Document(path).element.body.iter(qn('w:instrText'))
            if 'PAGEREF' in (n.text or '')]
    assert refs


def test_generate_full_split_and_single_write_expected_documents(tmp_path, monkeypatch):
    from api.services.tenders.generate import generate_bid

    raw = Document()
    for text in ('投标函', '本项目商务响应正文', '已标价工程量清单', '清单合计',
                 '施工组织设计', '本项目技术方案正文'):
        raw.add_paragraph(text)
    source = tmp_path / 'source.format.docx'
    raw.save(source)
    source.with_name(source.name + '.source').write_text('docx-raw', encoding='utf-8')
    original = source.read_bytes()
    monkeypatch.setattr('api.services.tenders.pdf_format._word_format_file', lambda _: source)
    monkeypatch.setattr('api.services.tenders.generate.tenders_output_dir', lambda: tmp_path)
    brief = BidBrief(layoutMode='outline', invitationId='abc123def456', generateVolume='full',
                     volumeAdvice='full', attachQualifications=False, projectName='综合标回归测试')
    kwargs = dict(catalog_slots=[], catalog_media={}, catalog_facts=[])
    full = generate_bid(brief, **kwargs)
    assert full['docxFile'] and not full['techDocxFile']
    assert full['downloadName'].endswith('投标文件.docx')
    blob = '\n'.join(p.text for p in Document(tmp_path / full['docxFile']).paragraphs)
    assert '本项目商务响应正文' in blob and '本项目技术方案正文' in blob

    split = generate_pair(brief.model_copy(update={'generateBoth': True, 'splitSlice': True}), **kwargs)
    assert split['docxFile'] and split['techDocxFile']
    assert split['docxFile'] != split['techDocxFile']
    business = '\n'.join(p.text for p in Document(tmp_path / split['docxFile']).paragraphs)
    technical = '\n'.join(p.text for p in Document(tmp_path / split['techDocxFile']).paragraphs)
    assert '本项目商务响应正文' in business and '本项目技术方案正文' not in business
    assert '本项目技术方案正文' in technical and '本项目商务响应正文' not in technical

    for volume in ('business', 'technical'):
        single = generate_bid(brief.model_copy(update={'generateVolume': volume}), **kwargs)
        assert bool(single['docxFile']) != bool(single['techDocxFile'])
    assert source.read_bytes() == original
    assert len(list(tmp_path.glob('*.docx'))) == 6  # source + full + split pair + two single calls
