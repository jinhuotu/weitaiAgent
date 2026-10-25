"""Table anchors and recovered source numbering use the application entry points."""
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from api.services.tenders.assemble import _anchor_section, _body_mark
from api.services.tenders.document import _bookmark_paragraph, _pageref_bookmark
from api.services.tenders.outline import normalize_subset_outline
from api.services.tenders.pdf_format import _repair_format_table_toc
from api.services.tenders.schema import DocumentFormat, OutlineItem, default_brief


def _xml_content(node):
    return node.tag, dict(node.attrib), node.text, [_xml_content(child) for child in node]


def test_source_bookmark_numbers_restore_siblings_and_top_boundary():
    titles = ['六、项目管理机构', '(一)项目管理机构组成表', '(二)项目经理简历表',
              '(三)项目经理主要工作经历', '(四)承诺书', '(五)主要项目管理人员简历表',
              '七、拟分包项目情况表', '八、资格审查资料']
    items = [OutlineItem(id=str(i), title=t.split('、')[-1].split(')')[-1], level=1)
             for i, t in enumerate(titles)]
    original = [i.model_dump() for i in items]
    result = normalize_subset_outline(items, source_titles=dict(zip([i.id for i in items], titles)))
    assert [i.level for i in result] == [1, 2, 2, 2, 2, 2, 1, 1]
    assert [i.model_dump() for i in items] == original


def test_table_cell_bookmark_is_moved_to_one_body_title():
    doc = Document()
    mark = _body_mark(doc)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '职务'
    _bookmark_paragraph(table.cell(0, 0).paragraphs[0], 'toc_target')
    before = deepcopy(table._tbl)
    for _ in range(2):
        _anchor_section(doc, mark, '项目管理机构组成表', 'toc_target', DocumentFormat(), subset=True)
    assert [p.text for p in doc.paragraphs] == ['项目管理机构组成表']
    assert not list(table._tbl.iter(qn('w:bookmarkStart')))
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
        for node in list(before.iter(qn(tag))):
            node.getparent().remove(node)
    assert _xml_content(table._tbl) == _xml_content(before)


@pytest.mark.parametrize('title', ['项目管理机构组成表', '项目经理主要工作经历',
                                  '拟分包项目情况表', '投标人基本情况表',
                                  '近年完成的类似项目情况表', '正在施工的和新承接的项目情况表'])
def test_original_toc_floating_table_repaired_with_shared_id_and_no_body_changes(title):
    doc = Document()
    doc.add_paragraph('目录')
    known = doc.add_paragraph('六、项目管理机构')
    entry = doc.add_paragraph(title)
    doc.add_page_break()
    _bookmark_paragraph(doc.add_paragraph('六、项目管理机构'), 'toc_known')
    # Existing reference establishes the original numbering base.
    from api.services.tenders.document import _add_pageref_run
    known.add_run().add_tab()
    _add_pageref_run(known, 'toc_known', '1')
    host = doc.add_paragraph()
    pict = OxmlElement('w:pict')
    content = OxmlElement('w:txbxContent')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '姓名与任职信息'
    content.append(table._tbl)
    pict.append(content)
    host.add_run()._r.append(pict)
    host.add_run(title)
    _bookmark_paragraph(host, '(一)' + title)
    before = _xml_content(host._p)
    brief = default_brief()
    brief.outlineItems = [OutlineItem(id='shturl', title=title)]
    assert _repair_format_table_toc(doc, brief) == 1
    assert _pageref_bookmark(entry) == 'toc_shturl'
    assert entry.text.rsplit('\t', 1)[-1].isdigit()
    assert float(entry.paragraph_format.left_indent.cm) == pytest.approx(0.74, abs=0.01)
    copy = deepcopy(host._p)
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
        for node in list(copy.iter(qn(tag))):
            if node.get(qn('w:name')) == 'toc_shturl' or node.get(qn('w:id')) == str(
                next(n for n in host._p.iter(qn('w:bookmarkStart')) if n.get(qn('w:name')) == 'toc_shturl').get(qn('w:id'))
            ):
                node.getparent().remove(node)
    assert _xml_content(copy) == before
    xml = doc.element.xml
    assert _repair_format_table_toc(doc, brief) == 0
    assert doc.element.xml == xml


def test_original_numbered_toc_preserves_labels_and_paragraph_spacing():
    doc = Document()
    doc.add_paragraph('目录')
    root = doc.add_paragraph('项目管理机构')
    child = doc.add_paragraph('项目经理简历表')
    for para in (root, child):
        num = OxmlElement('w:numPr')
        num_id = OxmlElement('w:numId')
        num_id.set(qn('w:val'), '1')
        num.append(num_id)
        para._p.get_or_add_pPr().append(num)
    doc.add_page_break()
    _bookmark_paragraph(doc.add_paragraph('项目管理机构'), '六、项目管理机构')
    _bookmark_paragraph(doc.add_paragraph('项目经理简历表'), '(二)项目经理简历表')
    assert _repair_format_table_toc(doc, default_brief()) == 0
    assert root.text == '六、项目管理机构'
    assert child.text == '(二)项目经理简历表'
    assert root.paragraph_format.space_before is None
    assert child.paragraph_format.space_after is None
    assert root._p.pPr.find(qn('w:numPr')) is None
    assert child._p.pPr.find(qn('w:numPr')) is None
    xml = doc.element.xml
    _repair_format_table_toc(doc, default_brief())
    assert doc.element.xml == xml


@pytest.mark.parametrize('number_source', ['title', 'source_titles', 'body'])
def test_bare_raw_title_does_not_hide_sibling_number(number_source):
    bare = '项目经理简历表'
    numbered = '(二)' + bare
    item = OutlineItem(id='resume', title=bare, rawTitle=bare, level=3)
    sources = {}
    if number_source == 'title':
        item.title = numbered
    elif number_source == 'source_titles':
        sources[item.id] = numbered
    else:
        item.body = numbered + '\n真实正文'
    items = [OutlineItem(id='root', title='六、项目管理机构'),
             OutlineItem(id='org', title='(一)项目管理机构组成表', level=3), item,
             OutlineItem(id='pm', title='(三)项目经理主要工作经历', level=1),
             OutlineItem(id='sub', title='七、拟分包项目情况表', level=3),
             OutlineItem(id='qual', title='八、资格审查资料', level=3)]
    assert [i.level for i in normalize_subset_outline(items, source_titles=sources)] == [1, 2, 2, 2, 1, 1]


def test_dot_tab_removes_old_left_stop_and_manual_dots_without_changing_page_cache():
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm
    from api.services.tenders.document import _add_pageref_run, _set_dot_tab

    doc = Document()
    para = doc.add_paragraph('(一)项目管理机构组成表')
    para.add_run('....')
    para.add_run('....')
    para.paragraph_format.tab_stops.add_tab_stop(Cm(2.13), WD_TAB_ALIGNMENT.LEFT)
    para.add_run().add_tab()
    _add_pageref_run(para, 'toc_org', '9')
    field = [r.xml for r in para._p.iter(qn('w:r')) if r.find(qn('w:fldChar')) is not None]
    _set_dot_tab(para, 15, normalize_runs=True)
    assert para.text == '(一)项目管理机构组成表\t9'
    tabs = list(para._p.pPr.iter(qn('w:tab')))
    assert len(tabs) == 1
    assert tabs[0].get(qn('w:val')) == 'right'
    assert tabs[0].get(qn('w:leader')) == 'dot'
    assert [r.xml for r in para._p.iter(qn('w:r')) if r.find(qn('w:fldChar')) is not None] == field
    tab_run = next(r for r in para.runs if '\t' in r.text)
    assert tab_run.font.name == 'Times New Roman'
    assert tab_run.font.size.pt == 10.5
    xml = para._p.xml
    _set_dot_tab(para, 15, normalize_runs=True)
    assert para._p.xml == xml
