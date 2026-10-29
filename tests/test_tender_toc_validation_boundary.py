import pytest
from docx import Document
from docx.oxml.ns import qn

from api.services.tenders.document import _bookmark_paragraph, _write_toc_line
from api.services.tenders.native_format import validate_native_toc
from common.errors import AppError


def form_document(*, source_bookmark=True, table=True):
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("目录")
    row = doc.add_paragraph()
    _write_toc_line(row, 0, "技术证明材料", "toc_proof", "2")
    doc.add_page_break()
    heading = doc.add_paragraph("采购需求偏离表", style="Heading 6")
    if source_bookmark:
        _bookmark_paragraph(heading, "_bookmark38")
    if table:
        doc.add_paragraph("商务要求偏离表")
        grid = doc.add_table(rows=2, cols=3)
        for cell, value in zip(grid.rows[0].cells, ("序号", "招标要求", "响应及偏离"), strict=True):
            cell.text = value
        grid.cell(1, 0).text = "1"
    else:
        doc.add_paragraph("本章为采购需求响应正文，应依照本项目招标文件填写。")
    target = doc.add_paragraph("技术证明材料", style="Heading 1")
    _bookmark_paragraph(target, "toc_proof")
    doc.add_paragraph("证明材料正文。")
    return doc, row, heading, target


@pytest.mark.parametrize("table", [True, False])
def test_original_body_can_precede_the_first_toc_bookmark_without_mutation(table):
    doc, _, heading, target = form_document(table=table)
    before = doc.element.xml
    stats = validate_native_toc(doc)
    assert stats["entries"] == 1
    assert stats["firstBodyBlock"] == list(doc.element.body).index(heading._p)
    assert stats["firstBodyBlock"] < list(doc.element.body).index(target._p)
    assert doc.element.xml == before


def test_unreferenced_original_form_does_not_hide_old_directory_residue():
    doc, row, _, _ = form_document()
    stray = doc.add_paragraph("旧目录遗漏条目")
    row._p.addnext(stray._p)
    with pytest.raises(AppError, match="旧目录遗漏条目"):
        validate_native_toc(doc)


def test_same_words_without_original_heading_evidence_are_still_rejected():
    doc, _, _, _ = form_document(source_bookmark=False)
    with pytest.raises(AppError, match="采购需求偏离表"):
        validate_native_toc(doc)


def test_attachment_placeholder_is_not_a_body_boundary():
    doc, _, heading, _ = form_document(table=False)
    source_marker = next(heading._p.iter(qn("w:bookmarkStart")))
    source_marker.set(qn("w:name"), "material_placeholder")
    with pytest.raises(AppError, match="采购需求偏离表"):
        validate_native_toc(doc)


@pytest.mark.parametrize("title", ["采购需求偏离表", "施工组织格式"])
def test_earlier_listed_heading_with_wrong_later_target_is_still_rejected(title):
    doc, row, heading, _ = form_document(table=False)
    heading.text = "施工组织设计" if title == "施工组织格式" else title
    _bookmark_paragraph(heading, "_original_body")
    _write_toc_line(row, 0, title, "toc_proof", "2")
    with pytest.raises(AppError, match="目录区残留"):
        validate_native_toc(doc)


def test_image_in_actual_toc_is_still_rejected():
    from docx.oxml import OxmlElement

    doc, row, _, _ = form_document()
    stray = doc.add_paragraph()
    stray.add_run()._r.append(OxmlElement("w:drawing"))
    row._p.addnext(stray._p)
    with pytest.raises(AppError, match="目录区混入附件"):
        validate_native_toc(doc)


@pytest.mark.parametrize("mode", ["missing", "duplicate", "self-reference"])
def test_bookmark_validation_remains_strict_with_unlisted_form(mode):
    doc, row, _, target = form_document()
    if mode == "missing":
        node = next(target._p.iter(qn("w:bookmarkStart")))
        target._p.remove(node)
    elif mode == "duplicate":
        _bookmark_paragraph(doc.add_paragraph("重复目标"), "toc_proof")
    else:
        start = next(target._p.iter(qn("w:bookmarkStart")))
        target._p.remove(start)
        row._p.insert(0, start)
    with pytest.raises(AppError, match="目录结构校验失败"):
        validate_native_toc(doc)
