"""Original formats take precedence for every volume, with local section fallbacks."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from api.services.tenders import format_preview, native_format, pdf_format, requirement_bid
from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.document import _pageref_bookmark, build_bid_docx
from api.services.tenders.outline import choose_layout_mode
from api.services.tenders.schema import BidBrief, OutlineItem

INVITE = "dededededede"


def _brief(**updates):
    brief = BidBrief(
        invitationId=INVITE,
        outlineChapter="第六章 投标文件格式",
        layoutMode="chapter5",
        requirementsDriven=True,
        formatConfirmed=True,
        projectName="本次设备采购项目",
        tenderNo="TEST-2026-01",
        bidderName="本次投标人有限公司",
        legalPersonName="李甲",
        bidDate="2026-10-28",
        bidPriceYuan=100,
        includePlaceholders=False,
        lotName="第一标段",
        autoDraftTechPlan=True,
        autoDraftBizResp=True,
        outlineItems=[
            OutlineItem(id="letter", title="投标函", rawTitle="一、投标函", kind="letter"),
            OutlineItem(id="auth", title="授权委托书", rawTitle="二、授权委托书", kind="auth"),
            OutlineItem(id="quote", title="报价表", rawTitle="三、报价表", kind="quote"),
            OutlineItem(
                id="plan",
                title="技术方案",
                rawTitle="四、技术方案",
                kind="tech_plan",
                draftBody="已确认的技术方案正文。",
            ),
            OutlineItem(
                id="deviation", title="技术偏离表", rawTitle="五、技术偏离表", kind="tech_dev"
            ),
            OutlineItem(id="proof", title="技术证明材料", kind="scan"),
            OutlineItem(id="sme", title="中小企业声明函", kind="commitment_copy"),
            OutlineItem(id="bribery", title="反商业贿赂承诺书", kind="commitment_copy"),
        ],
    )
    return brief.model_copy(update=updates)


@pytest.fixture
def native_source(tmp_path, monkeypatch):
    source = tmp_path / f"{INVITE}.docx"
    doc = Document()
    # The early directory entry and its following unrelated chapter must not win.
    doc.add_paragraph("目录")
    doc.add_paragraph("第六章 投标文件格式").add_run().add_tab()
    doc.add_paragraph("第五章 采购要求", style="Heading 1")
    doc.add_paragraph("不得复制到标书的采购人要求")
    doc.add_paragraph("第六章 投标文件格式", style="Heading 3")
    cover = doc.add_paragraph("投 标 文 件")
    cover.runs[0].font.name = "楷体"
    cover.runs[0].font.size = Pt(28)
    doc.add_paragraph("项目名称：________")
    doc.add_paragraph("项目编号：________")
    doc.add_paragraph("投标人：________（盖章）")
    doc.add_paragraph("法定代表人：________（签字）")
    doc.add_paragraph("年 月 日")
    doc.add_page_break()
    doc.add_paragraph("目录")
    for item in _brief().outlineItems[:5]:
        doc.add_paragraph(item.rawTitle).add_run().add_tab()
    doc.add_page_break()
    for item in _brief().outlineItems[:5]:
        doc.add_paragraph(item.rawTitle, style="Heading 2")
        if item.kind == "tech_plan":
            doc.add_paragraph("原版技术方案格式说明。")
        else:
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = f"原版{item.title}表格"
    doc.add_paragraph("第七章 合同条款", style="Heading 1")
    doc.add_paragraph("下一章不能进入投标文件")
    doc.save(source)
    monkeypatch.setattr("api.services.tenders.assets.tender_invitations_dir", lambda: tmp_path)
    monkeypatch.setattr(pdf_format, "tender_invitations_dir", lambda: tmp_path)
    (tmp_path / f"{INVITE}.txt").write_text(
        "本次项目编号：TEST-2026-01\n需提交《中小企业声明函》和《反商业贿赂承诺书》。",
        encoding="utf-8",
    )
    pdf_format.prepare_word_format_docx(source, INVITE)
    return source


def _assert_toc(doc):
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert entries
    assert all(p.text.rsplit("\t", 1)[-1].isdigit() for p in entries)
    starts = list(doc.element.body.iter(qn("w:bookmarkStart")))
    names = [n.get(qn("w:name")) for n in starts]
    ids = [n.get(qn("w:id")) for n in starts]
    assert len(names) == len(set(names))
    assert len(ids) == len(set(ids))
    assert set(ids) == {n.get(qn("w:id")) for n in doc.element.body.iter(qn("w:bookmarkEnd"))}
    tail = max(list(doc.element.body).index(p._p) for p in entries)
    body_names = {
        n.get(qn("w:name"))
        for el in list(doc.element.body)[tail + 1 :]
        for n in el.iter(qn("w:bookmarkStart"))
    }
    assert all(_pageref_bookmark(p) in body_names for p in entries)


def test_heading_three_wins_over_toc_and_empty_outline(native_source):
    title, text = native_format.word_format_section(native_source)
    assert title == "第六章 投标文件格式"
    assert "一、投标函" in text
    assert "不得复制到标书" not in text and "下一章不能" not in text
    assert choose_layout_mode(title, []) == "outline"


def test_stale_requirements_flags_cannot_override_original(native_source):
    brief = _brief()
    assert requirement_bid.prepare_requirements_bid(brief) == []
    assert brief.layoutMode == "outline" and not brief.requirementsDriven
    before = [i.model_dump() for i in brief.outlineItems]
    native_format.prefer_native_format(brief)
    assert [i.model_dump() for i in brief.outlineItems] == before


def test_named_policy_forms_survive_classifier_filter(native_source):
    brief = _brief(outlineItems=[OutlineItem(id="letter", title="投标函", kind="letter")])
    native_format.prefer_native_format(brief)
    assert {"中小企业声明函", "反商业贿赂承诺书", "技术方案", "技术偏离表"} <= {
        i.title for i in brief.outlineItems
    }


@pytest.mark.parametrize(
    "volume,anonymous",
    [("business", False), ("technical", False), ("technical", True), ("full", False)],
)
def test_original_volume_preserves_forms_and_supplements_only_missing(
    native_source,
    tmp_path,
    volume,
    anonymous,
):
    brief = _brief(generateVolume=volume, techAnonymous=anonymous)
    source = native_source.with_suffix(".format.docx")
    original = source.read_bytes()
    path, notes = assemble_bid_docx(
        brief,
        tmp_path / f"{volume}-{anonymous}.docx",
        volume=volume,
        catalog_slots=[],
        catalog_media={},
        catalog_facts=[],
    )
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    # An explicit original directory is authoritative, including its form chapters.
    # Only loose form labels outside that directory are excluded from TOC recovery.
    _assert_toc(doc)
    assert "下一章不能" not in text and "不得复制到标书" not in text
    assert (
        "原版技术方案格式说明。" in text
        if volume != "business"
        else "原版技术方案格式说明。" not in text
    )
    if volume == "technical":
        assert "原版投标函表格" not in tables
        assert "原版技术偏离表表格" in tables
        assert any("技术证明材料" in n and "补齐" in n for n in notes)
        assert "中小企业声明函" not in text and "反商业贿赂承诺书" not in text
        assert "已确认的技术方案正文。" in text
    else:
        assert "原版投标函表格" in tables and "原版授权委托书表格" in tables
        assert "原版报价表表格" in tables
        assert any("中小企业声明函" in n and "补齐" in n for n in notes)
        assert any("反商业贿赂承诺书" in n and "补齐" in n for n in notes)
        assert "【待补充】" in text
    if anonymous:
        assert "技术标（暗标）" in text and brief.projectName in text
        assert "第一标段" in text
        front = text.split("目录", 1)[0]
        assert "投标人" not in front and "签字" not in front and "盖章" not in front
        assert "本次投标人有限公司" not in doc.element.body.xml
    else:
        cover = next(p for p in doc.paragraphs if p.text == "投 标 文 件")
        assert cover.runs[0].font.name == "楷体" and cover.runs[0].font.size == Pt(28)
        front = text.split("目录", 1)[0]
        assert brief.bidderName in front and brief.tenderNo in front and brief.projectName in front
        assert "2026年10月28日" in front
    assert source.read_bytes() == original
    if volume == "full":
        signature = doc.element.body.xml
        for _ in range(2):
            pdf_format.unify_full_docx(path, brief)
            assert Document(path).element.body.xml == signature


def test_missing_native_toc_is_added_without_replacing_cover(native_source, tmp_path):
    source = native_source.with_suffix(".format.docx")
    doc = Document(source)
    for el in pdf_format.toc_block_ids(doc):
        el.getparent().remove(el)
    doc.save(source)
    brief = _brief()
    path, _ = build_bid_docx(
        brief,
        tmp_path / "no-toc.docx",
        volume="full",
        catalog_slots=[],
        catalog_media={},
        catalog_facts=[],
    )
    out = Document(path)
    _assert_toc(out)
    assert "投 标 文 件" in [p.text for p in out.paragraphs]
    assert any("原版投标函表格" in c.text for t in out.tables for r in t.rows for c in r.cells)


def test_preview_prioritizes_native_even_when_client_requests_fixed(
    native_source, tmp_path, monkeypatch
):
    monkeypatch.setattr(format_preview, "tenders_output_dir", lambda: tmp_path / "output")
    (tmp_path / "output").mkdir()
    result = format_preview.prepare_invitation_format_docx(
        INVITE, brief=_brief(layoutMode="chapter5")
    )
    assert result["source"] == "invitation" and not result["previewOnly"]
    assert (tmp_path / "output" / result["fileName"]).read_bytes() == native_source.with_suffix(
        ".format.docx"
    ).read_bytes()


def test_duplicate_source_bookmarks_repaired_idempotently():
    doc = Document()
    for name in ("duplicate", "duplicate", "other"):
        p = doc.add_paragraph(name)
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:name"), name)
        start.set(qn("w:id"), "7")
        end.set(qn("w:id"), "7")
        p._p.append(start)
        p._p.append(end)
    native_format.normalize_bookmarks(doc)
    starts = list(doc.element.body.iter(qn("w:bookmarkStart")))
    assert [n.get(qn("w:name")) for n in starts] == ["duplicate", "other"]
    assert len({n.get(qn("w:id")) for n in starts}) == 2
    before = doc.element.body.xml
    native_format.normalize_bookmarks(doc)
    assert doc.element.body.xml == before


def test_matching_builtin_section_copies_only_its_style_list_and_image(tmp_path):
    from PIL import Image

    template = Document()
    template.add_heading("投标函", 1)
    template.add_paragraph("不能串入的其他模板正文")
    heading = template.add_heading("中小企业声明函", 1)
    heading.runs[0].font.name = "仿宋"
    for text in ("【待补充】企业名称", "【待补充】企业类型"):
        p = template.add_paragraph(text)
        num = OxmlElement("w:numPr")
        ident = OxmlElement("w:numId")
        ident.set(qn("w:val"), "1")
        num.append(ident)
        p._p.get_or_add_pPr().append(num)
    image = tmp_path / "slot.png"
    Image.new("RGB", (2, 2), "white").save(image)
    template.add_picture(str(image))
    template.add_heading("报价表", 1)
    template.add_paragraph("不能串入的报价")
    doc = Document()
    doc.add_paragraph("原版封面")
    assert native_format._copy_template_section(
        doc,
        template,
        OutlineItem(title="中小企业声明函", kind="commitment_copy"),
    )
    target = tmp_path / "copied.docx"
    doc.save(target)
    out = Document(target)
    assert "不能串入" not in out.element.body.xml
    assert next(p for p in out.paragraphs if p.text == "中小企业声明函").runs[0].font.name == "仿宋"
    ids = {n.get(qn("w:val")) for n in out.element.body.iter(qn("w:numId"))}
    assert len(ids) == 1
    assert all(
        out.part.numbering_part.element.find(f'w:num[@w:numId="{n}"]', out.element.nsmap)
        is not None
        for n in ids
    )
    refs = [n.get(qn("r:embed")) for n in out.element.body.iter(qn("a:blip"))]
    assert refs and all(ref in out.part.rels for ref in refs)


@pytest.mark.parametrize("number", ["TEST-FROM-SOURCE", ""])
def test_cover_adds_project_number_or_pending_without_changing_original_font(number):
    doc = Document()
    doc.add_paragraph("投 标 文 件")
    bidder = doc.add_paragraph("投标人：________（盖章）")
    bidder.runs[0].font.name = "楷体"
    doc.add_paragraph("目录")
    doc.add_paragraph("一、投标函").add_run().add_tab()
    doc.add_page_break()
    doc.add_paragraph("一、投标函")
    brief = _brief(invitationId="", tenderNo=number)
    native_format.fill_native_cover(doc, brief)
    line = next(p for p in doc.paragraphs if p.text.startswith("项目编号"))
    assert line.text == f"项目编号：{number or '【待补充】'}"
    assert line.runs[0].font.name == "楷体"


def test_cover_signoff_after_toc_moves_onto_cover_page():
    doc = Document()
    doc.add_paragraph("投 标 文 件")
    doc.add_paragraph("目录")
    doc.add_paragraph("一、投标函").add_run().add_tab()
    doc.add_page_break()
    bidder = doc.add_paragraph("投标人：________（盖章）")
    bidder.runs[0].font.name = "楷体"
    legal = doc.add_paragraph("法定代表人：________（签字）")
    legal.runs[0].font.name = "楷体"
    doc.add_paragraph("年　　月　　日")
    doc.add_page_break()
    doc.add_paragraph("一、投标函")
    body = doc.add_paragraph("投标人：（盖章）")
    brief = _brief()
    brief.outlineItems[0] = brief.outlineItems[0].model_copy(update={"tocEligible": True})
    native_format.fill_native_cover(doc, brief)
    native_format.ensure_toc_rows(doc, brief)
    texts = [p.text or "" for p in doc.paragraphs]
    toc_at = texts.index("目录")
    cover = texts[:toc_at]
    assert any(t.startswith("项目编号：") for t in cover)
    assert any("本次投标人有限公司" in t and "盖章" in t for t in cover)
    assert any("李甲" in t and "签字" in t for t in cover)
    assert any("2026年10月28日" in t for t in cover)
    hit = next(p for p in doc.paragraphs if "本次投标人有限公司" in (p.text or ""))
    assert hit.runs[0].font.name == "楷体"
    assert hit.paragraph_format.page_break_before is not True
    assert body.text == "投标人：（盖章）"
    assert texts.index("一、投标函") > toc_at or "\t" in texts[texts.index("一、投标函")]


def test_stamp_line_stays_on_cover_when_format_has_no_toc():
    doc = Document()
    doc.add_paragraph("第六章 投标文件格式")
    doc.add_paragraph("封面格式")
    doc.add_paragraph("（项目名称）")
    doc.add_paragraph("投  标 文 件")
    doc.add_paragraph("项目编号：")
    bidder = doc.add_paragraph("投标人：（加盖公章）")
    bidder.runs[0].font.name = "楷体"
    legal = doc.add_paragraph("法定代表人或其委托代理人：（签字或盖章）年")
    legal.add_run().add_tab()
    legal.add_run("月")
    doc.add_paragraph("第一部分资格证明文件")
    body = doc.add_paragraph("投标人（加盖公章）：")
    brief = _brief(
        invitationId="",
        tenderNo="豫财招标采购-2026-906",
        outlineItems=[
            OutlineItem(
                id="stamp",
                title="投标人（加盖公章）",
                rawTitle="投标人：（加盖公章）",
                kind="unknown",
                tocEligible=True,
            ),
            OutlineItem(
                id="qual",
                title="资格证明文件",
                rawTitle="第一部分资格证明文件",
                kind="unknown",
                tocEligible=True,
            ),
        ],
    )
    native_format.fill_native_cover(doc, brief)
    native_format.ensure_toc_rows(doc, brief)
    texts = [p.text or "" for p in doc.paragraphs]
    toc_at = texts.index("目录")
    cover = texts[:toc_at]
    assert any("豫财招标采购-2026-906" in t for t in cover)
    assert any("本次投标人有限公司" in t and "加盖公章" in t for t in cover)
    assert any("李甲" in t and "签字或盖章" in t for t in cover)
    assert any("2026年10月28日" in t for t in cover)
    assert texts.index("第一部分资格证明文件") > toc_at
    hit = next(p for p in doc.paragraphs if "本次投标人有限公司" in (p.text or ""))
    assert hit.runs[0].font.name == "楷体"
    assert hit.paragraph_format.page_break_before is not True
    assert body.text == "投标人（加盖公章）："


def test_absent_technical_format_builds_only_technical_supplements(native_source, tmp_path):
    source = native_source.with_suffix(".format.docx")
    doc = Document(source)
    pdf_format._drop_technical(doc, brief=_brief())
    doc.save(source)
    brief = _brief(techAnonymous=True)
    path, notes = assemble_bid_docx(
        brief,
        tmp_path / "new-tech.docx",
        volume="technical",
        catalog_slots=[],
        catalog_media={},
        catalog_facts=[],
    )
    out = Document(path)
    _assert_toc(out)
    assert "已确认的技术方案正文。" in "\n".join(p.text for p in out.paragraphs)
    assert sum(p.text == "已确认的技术方案正文。" for p in out.paragraphs) == 1
    from api.services.tenders.assemble import _fill_drafted_format_sections

    before = out.element.body.xml
    assert _fill_drafted_format_sections(path, brief) == (0, [])
    assert Document(path).element.body.xml == before
    assert "原版投标函表格" not in out.element.body.xml
    assert any("技术方案" in note and "补齐" in note for note in notes)
