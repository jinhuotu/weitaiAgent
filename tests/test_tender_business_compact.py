import asyncio
import json

import pytest
from docx import Document

from api.services.tenders.assemble import _fill_business_format_sections, assemble_bid_docx
from api.services.tenders.business_render import commercial_rows, finalize_business_responses
from api.services.tenders.gen_biz_resp import draft_business_response
from api.services.tenders.schema import BidBrief, BidRequirement, DeviationLine, OutlineItem


def _brief(deviation=True):
    items = [
        OutlineItem(id="license", title="一、营业执照", rawTitle="一、营业执照", kind="scan"),
        OutlineItem(id="finance", title="二、财务状况", kind="scan"),
        OutlineItem(id="contract", title="三、合同条款响应", kind="commitment_copy"),
        OutlineItem(id="quote", title="五、分项报价表", kind="quote"),
    ]
    if deviation:
        items.insert(3, OutlineItem(id="dev", title="四、商务偏差表", kind="biz_dev"))
    return BidBrief(
        autoDraftBizResp=True,
        layoutMode="outline",
        outlineItems=items,
        bizDeviationLines=[
            DeviationLine(
                requirement="须提供有效营业执照",
                response="已附有效营业执照扫描件",
                deviation="无偏离",
            ),
            DeviationLine(
                requirement="须提供财务资料", response="已附审计报告", deviation="无偏离"
            ),
            DeviationLine(
                requirement="付款必须按合同约定", response="同意按合同约定付款", deviation="无偏离"
            ),
            *[
                DeviationLine(
                    requirement=f"未响应原文长条款{i}" + "招标原文" * 30, response="【待补充】"
                )
                for i in range(30)
            ],
        ],
    )


def _text(doc):
    return (
        "\n".join(p.text for p in doc.paragraphs)
        + "\n"
        + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    )


def test_only_substantive_unrouted_rows_are_expanded():
    brief = _brief()
    brief.bizDeviationLines.append(brief.bizDeviationLines[2].model_copy())
    rows = commercial_rows(brief)
    assert len(rows) == 2
    assert rows[0][1] == "付款必须按合同约定"
    assert "30 项" in rows[1][2]
    assert "未响应原文" not in str(rows)
    assert "营业执照" not in str(rows)


@pytest.mark.parametrize("deviation", [True, False])
def test_native_routing_preserves_order_and_is_idempotent(tmp_path, deviation):
    brief = _brief(deviation)
    doc = Document()
    for item in brief.outlineItems:
        doc.add_paragraph(item.title)
        if item.kind == "biz_dev":
            table = doc.add_table(rows=1, cols=4)
            for cell, title in zip(
                table.rows[0].cells, ("招标条款", "响应", "偏离", "说明"), strict=False
            ):
                cell.text = title
        else:
            doc.add_paragraph("原版内容")
    path = tmp_path / "native.docx"
    doc.save(path)
    if deviation:
        _fill_business_format_sections(path, brief)
    doc = Document(path)
    finalize_business_responses(doc, brief)
    before = doc.element.xml
    finalize_business_responses(doc, brief)
    assert doc.element.xml == before
    text = _text(doc)
    assert text.count("已附有效营业执照扫描件") == 1
    assert text.count("已附审计报告") == 1
    assert "未响应原文" not in text
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    assert (
        paragraphs.index("一、营业执照")
        < paragraphs.index("已附有效营业执照扫描件")
        < paragraphs.index("二、财务状况")
    )
    assert (
        paragraphs.index("已附审计报告")
        < paragraphs.index("三、合同条款响应")
        < paragraphs.index("五、分项报价表")
    )
    assert text.count("30 项商务条款尚未响应") == 1


def test_missing_routing_heading_falls_back_to_final_table():
    brief = _brief()
    doc = Document()
    doc.add_paragraph("三、合同条款响应")
    notes = finalize_business_responses(doc, brief)
    assert any("文末核对表" in note for note in notes)
    assert "已附有效营业执照扫描件" in _text(doc)
    assert "须提供有效营业执照" in _text(doc)


def test_prefilled_source_table_does_not_silently_drop_responses():
    brief = _brief()
    doc = Document()
    for item in brief.outlineItems:
        doc.add_paragraph(item.title)
    original = doc.add_table(rows=1, cols=1)
    original.cell(0, 0).text = "原版已有内容，不能覆盖"
    finalize_business_responses(doc, brief)
    text = _text(doc)
    assert "原版已有内容，不能覆盖" in text
    assert text.count("同意按合同约定付款") == 1
    assert text.count("30 项商务条款尚未响应") == 1
    assert "未响应原文" not in text
    before = doc.element.xml
    finalize_business_responses(doc, brief)
    assert doc.element.xml == before


def test_batches_do_not_repeat_pending_or_copy_unanswered_requirements(monkeypatch):
    monkeypatch.setattr("api.services.tenders.gen_biz_resp.load_invitation_text", lambda _: "")

    class Client:
        async def complete(self, *args, **kwargs):
            return json.dumps(
                {
                    "clauses": [],
                    "chapters": [
                        {"title": "响应承诺书", "body": "【待补充】" + "未响应长条款" * 100}
                    ],
                },
                ensure_ascii=False,
            )

    brief = BidBrief(
        autoDraftBizResp=True,
        outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")],
        bidRequirements=[
            BidRequirement(text=f"第{i}项付款必须按合同执行", volume="business") for i in range(45)
        ],
    )
    result, _ = asyncio.run(draft_business_response(None, brief, client=Client()))
    assert len(result.bizDeviationLines) == 45
    assert result.outlineItems[0].draftBody.count("【待补充】") == 1
    assert "未响应长条款" not in result.outlineItems[0].draftBody
    doc = Document()
    finalize_business_responses(doc, result)
    assert "45 项" in _text(doc)
    assert "第44项付款" not in _text(doc)


def test_model_cannot_duplicate_clause_response_in_commitment_body(monkeypatch):
    requirement = "付款必须按合同约定执行"
    response = "我方同意按合同约定支付相关款项"
    monkeypatch.setattr(
        "api.services.tenders.gen_biz_resp.load_invitation_text", lambda _: requirement
    )

    class Client:
        async def complete(self, *args, **kwargs):
            return json.dumps(
                {
                    "clauses": [
                        {"requirement": requirement, "response": response, "deviation": "无偏离"}
                    ],
                    "chapters": [{"title": "响应承诺书", "body": requirement + "\n" + response}],
                },
                ensure_ascii=False,
            )

    brief = BidBrief(
        autoDraftBizResp=True,
        outlineItems=[OutlineItem(title="响应承诺书", kind="commitment_copy")],
        bidRequirements=[BidRequirement(text=requirement, volume="business")],
    )
    result, _ = asyncio.run(draft_business_response(None, brief, client=Client()))
    assert requirement not in result.outlineItems[0].draftBody
    assert response not in result.outlineItems[0].draftBody
    doc = Document()
    finalize_business_responses(doc, result)
    assert _text(doc).count(response) == 1


@pytest.mark.parametrize("volume", ["business", "full"])
def test_assembled_bids_use_compact_table_and_routed_responses(tmp_path, volume):
    path = tmp_path / f"{volume}.docx"
    brief = _brief()
    assemble_bid_docx(brief, path, volume=volume, catalog_slots=[], catalog_media={})
    doc = Document(path)
    text = _text(doc)
    assert "未响应原文" not in text
    assert text.count("已附有效营业执照扫描件") == 1
    assert text.count("同意按合同约定付款") == 1
    assert text.count("30 项商务条款尚未响应") == 1
