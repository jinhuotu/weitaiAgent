"""Foreign prices and technical evidence cannot become facts in another bid."""
import asyncio
import json
from dataclasses import replace
from decimal import Decimal
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from api.services.tenders import gen_tech_plan
from api.services.tenders.assemble import assemble_bid_docx
from api.services.tenders.extract import _merge_quote
from api.services.tenders.quote import (
    attach_sheet,
    parse_quote_path,
    quote_source_notes,
    resolve_quote_sheet,
    sheet_from_quote_records,
    sheet_from_rows,
    validate_quote_source,
)
from api.services.tenders.records import _prepare_and_generate
from api.services.tenders.schema import BidBrief, OutlineItem, QuoteLineIn
from api.services.tenders.source_match import match_source
from common.errors import AppError


def _medical():
    return BidBrief(
        projectName='甲医院车辆采购（救护车）项目', bidContent='救护车及配套医疗设备采购',
        tenderer='甲医院', bidPriceYuan=1000, generateVolume='full',
    )


def _sheet(title='报价清单', name='救护车', spec='医疗急救车辆'):
    sheet = sheet_from_rows([
        ['名称', '规格型号', '单位', '数量', '单价', '金额'],
        [name, spec, '台', 1, 100, 100],
    ], title_hint=title)
    assert sheet is not None
    return sheet


@pytest.mark.parametrize('title,name,spec,headers,accepted', [
    ('甲医院救护车报价清单', '救护车', '2026 年生产，续航 500 公里', (), True),
    ('报价清单', '救护车', '最高时速 120 公里/小时', ('生产年份', '续航公里'), True),
    ('甲医院救护车报价', '车载电源', 'USB充电器', (), True),
    ('报价清单', '交流充电桩', '7kW', (), False),
    ('甲医院救护车报价', '救护车', '含直流充电桩 30 台', (), False),
    ('甲医院救护车报价', '救护车', '公路施工里程2020年50公里', (), False),
    ('报价清单', '道路施工费', '土方工程', (), False),
    ('报价清单', '2024年', '50 公里', ('年份', '道路里程'), False),
    ('乙医院救护车采购项目', '救护车', '医疗车辆', (), False),
    ('甲医院救护车报价', '救护车', '无需充电桩，医用车载供电', (), True),
])
def test_project_category_and_vehicle_parameter_boundaries(title, name, spec, headers, accepted):
    result = match_source(_medical(), title=title, rows=[(name, spec)], headers=headers)
    assert (result.status == 'matched') == accepted
    if not accepted:
        assert result.status == 'unmatched'


def test_hospital_charger_and_construction_projects_use_their_actual_scope():
    charger = BidBrief(projectName='甲医院充电桩采购项目')
    assert match_source(charger, title='报价清单', rows=[('交流充电桩', '7kW')]).status == 'matched'
    construction = BidBrief(projectName='甲医院救护车停车区土建工程')
    assert match_source(construction, title='报价清单', rows=[('土方工程', '')]).status == 'matched'


def test_one_foreign_row_blocks_the_whole_sheet_even_with_matching_title():
    good = _sheet('甲医院救护车报价清单')
    bad = _sheet(name='交流充电桩', spec='7kW')
    mixed = replace(good, lines=good.lines * 60 + bad.lines)
    attached = attach_sheet(_medical(), mixed, source='file', source_name='本项目清单.xlsx')
    assert attached.quoteLines == []
    assert attached.quoteMatchStatus == 'unmatched'
    assert '61' in attached.quoteMatchNote


def test_attach_rejects_foreign_data_without_filling_its_total():
    brief = _medical()
    brief.bidPriceYuan = 0
    attached = attach_sheet(brief, _sheet(name='交流充电桩', spec='7kW'),
                            source='file', source_name='错误清单.xlsx')
    assert not attached.quoteLines and attached.bidPriceYuan == 0
    assert not attached.quoteTitle and not attached.quoteHeaders
    assert attached.quoteSourceName == '错误清单.xlsx'
    assert '清单与本次标书不匹配，未采用' in quote_source_notes(attached)[0]


def test_successful_attachment_and_changed_project_revalidate_auto_price():
    brief = _medical()
    brief.bidPriceYuan = 0
    attached = attach_sheet(brief, _sheet(), source='file', source_name='医疗清单.xlsx')
    assert attached.quoteLines and attached.quoteAutoPrice
    assert attached.quoteMatchStatus == 'matched'
    assert '已匹配' in quote_source_notes(attached)[0]
    manual = attached.model_copy(deep=True)
    manual.bidPriceYuan = 999
    for candidate in (attached, manual):
        candidate.projectName = '充电桩采购项目'
        candidate.bidContent = '交流充电桩采购'
        assert validate_quote_source(candidate)
        assert not candidate.quoteLines
    assert attached.bidPriceYuan == 0  # Known auto-filled price from the foreign sheet is removed.
    assert manual.bidPriceYuan == 999  # Explicit user price is not inferred from the old sheet.


def test_external_sheet_without_match_evidence_is_not_adopted():
    attached = attach_sheet(BidBrief(projectName='新建档案馆采购项目'),
                            _sheet(name='设备A', spec='型号B'), source='file')
    assert not attached.quoteLines
    assert attached.quoteMatchStatus == 'unmatched'


def test_resolver_rechecks_tampered_or_legacy_match_status(tmp_path):
    brief = _medical()
    brief.quoteTitle = '甲医院救护车报价清单'
    brief.quoteMatchStatus = 'matched'
    brief.quoteSource = 'file'
    brief.quoteLines = [QuoteLineIn(name='直流充电桩', spec='120kW', qty=50)]
    sheet, origin = resolve_quote_sheet(brief, tmp_path / 'never-use-builtin.xlsx')
    assert sheet is None and origin == 'template'
    assert not brief.quoteLines and brief.quoteMatchStatus == 'unmatched'


def test_merged_quote_records_keep_each_project_identity():
    def record(title):
        return SimpleNamespace(project_name=title, tax_rate=Decimal('0.13'), total_inc_tax=113,
                               lines_json=[{'name': '救护车', 'qty': 1, 'unitPrice': 100}])
    sheet = sheet_from_quote_records([
        record('甲医院车辆采购（救护车）项目'), record('乙医院救护车项目'),
    ])
    assert sheet is not None and len(sheet.source_projects) == 2
    attached = attach_sheet(_medical(), sheet, source='record', source_name='报价记录')
    assert not attached.quoteLines
    assert '项目' in attached.quoteMatchNote


@pytest.mark.parametrize('extension', ['xlsx', 'docx'])
def test_file_project_label_is_checked_even_when_items_have_same_category(tmp_path, extension):
    path = tmp_path / f'报价明细.{extension}'
    rows = [['名称', '规格型号', '单位', '数量', '单价'], ['救护车', '急救车辆', '台', 1, 100]]
    if extension == 'xlsx':
        wb = Workbook()
        wb.active.append(['项目名称：', '乙医院救护车项目'])
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
    else:
        doc = Document()
        doc.add_paragraph('项目名称：乙医院救护车项目')
        table = doc.add_table(rows=2, cols=5)
        for row, values in zip(table.rows, rows, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = str(value)
        doc.save(path)
    parsed = parse_quote_path(path)
    assert parsed is not None and parsed.source_projects
    notes, result, filled = _merge_quote(_medical(), {}, '', None, path,
                                         source_name=f'误传清单.{extension}')
    assert not filled and not result.quoteLines
    assert any('清单与本次标书不匹配，未采用' in note for note in notes)


def test_unparseable_sheet_does_not_reuse_old_rows_or_raise(tmp_path):
    path = tmp_path / 'bad.csv'
    path.write_text('not,a,quote', encoding='utf-8')
    brief = _medical()
    brief.quoteLines = [QuoteLineIn(name='旧的充电桩', qty=3)]
    notes, result, filled = _merge_quote(brief, {}, '', None, path)
    assert not filled and not result.quoteLines
    assert any('未能按表格解析' in note for note in notes)


def test_full_doc_and_generation_notes_do_not_use_foreign_quote_rows(tmp_path):
    brief = _medical()
    brief.requirementsDriven = True
    brief.quoteSource = 'file'
    brief.quoteSourceName = '错误项目清单.xlsx'
    brief.quoteLines = [QuoteLineIn(name='交流充电桩', spec='施工里程2099公里', qty=99)]
    path, notes = assemble_bid_docx(brief, tmp_path / 'medical-full.docx', volume='full',
                                   catalog_slots=[], catalog_media={}, catalog_facts=[])
    doc = Document(path)
    text = '\n'.join([p.text for p in doc.paragraphs] +
                     [cell.text for t in doc.tables for row in t.rows for cell in row.cells])
    for foreign in ('交流充电桩', '施工里程2099', '99台'):
        assert foreign not in text
    assert '分项报价清单' in text and '【待补充】' in text
    assert any('错误项目清单.xlsx' in note and '未采用' in note for note in notes)


def test_record_sanitizes_before_validation_and_llm_or_library_access(monkeypatch):
    brief = _medical()
    brief.quoteLines = [QuoteLineIn(name='充电桩', qty=99)]
    def validate(candidate):
        assert not candidate.quoteLines and candidate.quoteMatchStatus == 'unmatched'
        return ['测试到此停止']
    monkeypatch.setattr('api.services.tenders.records._brief_generate_issues', validate)
    with pytest.raises(AppError, match='清单与本次标书不匹配，未采用'):
        asyncio.run(_prepare_and_generate(None, brief))


def _source(root, source_id, text, filename=None):
    (root / f'{source_id}.txt').write_text(text, encoding='utf-8')
    if filename:
        (root / f'{source_id}.meta.json').write_text(
            json.dumps({'filename': filename}, ensure_ascii=False), encoding='utf-8',
        )
    return source_id


def test_technical_sources_are_checked_individually_including_legacy_files(
    tmp_path, monkeypatch,
):
    monkeypatch.setattr(gen_tech_plan, '_source_dir', lambda: tmp_path)
    ids = [
        _source(tmp_path, 'a' * 16, '项目名称：甲医院救护车项目\n救护车资料 KEEP_ME',
                '救护车.docx'),
        _source(tmp_path, 'b' * 16, '充电桩施工组织设计 DROP_CHARGER',
                '错误施工.pdf'),
        _source(tmp_path, 'c' * 16, '项目名称：乙医院救护车项目\n医疗资料 DROP_OTHER'),
    ]
    notes = []
    text = gen_tech_plan._user_sources(ids, brief=_medical(), notes=notes)
    assert 'KEEP_ME' in text and 'DROP_CHARGER' not in text and 'DROP_OTHER' not in text
    assert any('救护车.docx' in note and '已匹配' in note for note in notes)
    assert any('错误施工.pdf' in note and '未匹配' in note and '未采用' in note for note in notes)
    assert any('c' * 16 in note and '未匹配' in note for note in notes)


def test_supplier_name_is_not_mistaken_for_a_different_buyer_project():
    attached = attach_sheet(_medical(), _sheet(title='乙公司报价清单'),
                            source='file', source_name='乙公司救护车报价.xlsx')
    assert attached.quoteLines and attached.quoteMatchStatus == 'matched'


def test_reattachment_does_not_reclassify_an_edited_price_as_automatic():
    brief = _medical()
    brief.bidPriceYuan = 0
    attached = attach_sheet(brief, _sheet(), source='file')
    attached.bidPriceYuan = 200  # Edited by the user after the first attachment.
    new = replace(_sheet(), total_inc_tax=Decimal('200'), source_inc_tax=Decimal('200'))
    attached = attach_sheet(attached, new, source='file')
    assert attached.bidPriceYuan == 200 and not attached.quoteAutoPrice


@pytest.mark.parametrize('with_valid_row', [False, True])
def test_zero_quantity_foreign_rows_cannot_hide_from_generation_validation(with_valid_row):
    brief = _medical()
    brief.quoteSource = 'file'
    brief.quoteTitle = '甲医院救护车报价清单'
    brief.quoteMatchStatus = 'matched'
    brief.quoteLines = [QuoteLineIn(name='交流充电桩', spec='旧项目设备')]
    if with_valid_row:
        brief.quoteLines.insert(0, QuoteLineIn(name='救护车', qty=1, unitPrice=100))
    assert resolve_quote_sheet(brief) == (None, 'template')
    assert not brief.quoteLines and brief.quoteMatchStatus == 'unmatched'
    from api.services.tenders.document import infer_factory_role

    assert '充电' not in infer_factory_role(brief)


def test_new_invitation_without_a_sheet_drops_only_the_old_auto_price():
    brief = _medical()
    brief.bidPriceYuan = 0
    previous = attach_sheet(brief, _sheet(), source='file')
    manual = previous.model_copy(deep=True)
    manual.bidPriceYuan = 999
    notes, cleared, _ = _merge_quote(previous, {}, '', None, None)
    assert cleared.bidPriceYuan == 0 and not cleared.quoteLines
    assert not cleared.quoteSource and cleared.quoteSourceIncTax == 0
    assert any('自动填入的总价' in note for note in notes)
    _, kept, _ = _merge_quote(manual, {}, '', None, None)
    assert kept.bidPriceYuan == 999


def test_unmatched_technical_source_never_reaches_model(tmp_path, monkeypatch):
    monkeypatch.setattr(gen_tech_plan, '_source_dir', lambda: tmp_path)
    monkeypatch.setattr(gen_tech_plan, 'load_invitation_text', lambda _: '救护车采购及质量验收要求')
    bad = _source(tmp_path, 'a' * 16, '其他项目充电桩规格 BAD_120kW', '其他项目.pdf')
    good = _source(tmp_path, 'b' * 16, '救护车质量维护 GOOD_MEDICAL', '医疗技术资料.docx')
    brief = _medical()
    brief.techPlanSourceIds = [bad, good]
    brief.outlineItems = [OutlineItem(id='tech', title='技术方案', kind='tech_plan')]
    class Client:
        async def complete(self, messages, *, mode):
            prompt = '\n'.join(message['content'] for message in messages)
            assert 'BAD_120kW' not in prompt and 'GOOD_MEDICAL' in prompt
            return json.dumps({'chapters': [{'title': '技术方案', 'body': '救护车质量维护正文。'}]},
                              ensure_ascii=False)
    result, notes = asyncio.run(gen_tech_plan.draft_technical_plan(None, brief, client=Client()))
    assert result.outlineItems[0].draftBody
    assert any('其他项目.pdf' in note and '未采用' in note for note in notes)


def test_technical_upload_retains_original_filename_for_audit(tmp_path, monkeypatch):
    monkeypatch.setattr(gen_tech_plan, '_source_dir', lambda: tmp_path)
    path = tmp_path / 'medical.docx'
    doc = Document()
    doc.add_paragraph('救护车技术资料')
    doc.save(path)
    saved = asyncio.run(gen_tech_plan.save_tech_source(path.read_bytes(), '原始医疗技术资料.docx'))
    metadata = json.loads((tmp_path / f'{saved["sourceId"]}.meta.json').read_text(encoding='utf-8'))
    assert metadata['filename'] == '原始医疗技术资料.docx'
