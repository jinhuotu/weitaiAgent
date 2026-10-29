"""Unmatched materials follow explicit paste marks instead of the chapter tail."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from api.services.tenders.placeholders import _is_slot_box, place_unmatched_slots
from api.services.tenders.schema import PlaceholderItem


def _next_content(el):
    el = el.getnext()
    while el is not None:
        style = el.find("w:pPr/w:pStyle", el.nsmap)
        if style is None or style.get(qn("w:val")) != "TenderAttachmentGap":
            return el
        el = el.getnext()
    return None


def test_two_unmatched_slots_follow_their_own_marks_in_one_business_chapter():
    doc = Document()
    doc.add_paragraph("一、资格资信证明文件")
    protocol = doc.add_paragraph("此处粘贴技术协议（在此粘贴扫描件）")
    finance = doc.add_paragraph("此处粘贴财务报表（在此粘贴扫描件）")
    next_chapter = doc.add_paragraph("二、合同条款响应书")
    slots = [
        PlaceholderItem(key="protocol", title="技术协议"),
        PlaceholderItem(key="finance", title="财务报表"),
    ]
    notes = place_unmatched_slots(
        doc, slots=slots, media={}, required_keys=[s.key for s in slots], volume="business"
    )
    first = _next_content(protocol._p)
    second = _next_content(finance._p)
    assert _is_slot_box(first) and "此处粘贴技术协议" in first.xml
    assert _is_slot_box(second) and "此处粘贴财务报表" in second.xml
    assert first.getnext() is not next_chapter._p
    assert sum(_is_slot_box(el) for el in doc.element.body) == 2
    assert all("附注后画虚线框" in note for note in notes if "没有扫描件" in note)

    saved = doc.element.xml
    place_unmatched_slots(
        doc, slots=slots, media={}, required_keys=[s.key for s in slots], volume="business"
    )
    assert doc.element.xml == saved


def test_existing_dashed_box_is_not_duplicated():
    doc = Document()
    doc.add_paragraph("一、资格资信证明文件")
    box = doc.add_paragraph("此处粘贴财务报表\n（在此粘贴扫描件）")
    borders = box._p.get_or_add_pPr().append
    from docx.oxml import OxmlElement

    borders(OxmlElement("w:pBdr"))
    doc.add_paragraph("二、合同条款响应书")
    before = doc.element.xml
    place_unmatched_slots(
        doc,
        slots=[PlaceholderItem(key="finance", title="财务会计报表")],
        media={},
        required_keys=["finance"],
        volume="business",
    )
    assert doc.element.xml == before


def test_unbordered_table_marker_anchors_after_that_table():
    doc = Document()
    doc.add_paragraph("一、资格资信证明文件")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "此处粘贴财务报表（在此粘贴扫描件）"
    next_chapter = doc.add_paragraph("二、合同条款响应书")
    place_unmatched_slots(
        doc,
        slots=[PlaceholderItem(key="finance", title="财务会计报表")],
        media={},
        required_keys=["finance"],
        volume="business",
    )
    box = _next_content(table._tbl)
    assert _is_slot_box(box)
    assert "此处粘贴财务会计报表" in box.xml
    assert next_chapter._p.getprevious() is not table._tbl


def test_explanatory_note_uses_later_explicit_marker_not_note_position():
    doc = Document()
    doc.add_paragraph("一、资格资信证明文件")
    note = doc.add_paragraph("注：应提供财务报表扫描件。")
    marker = doc.add_paragraph("此处粘贴财务报表（在此粘贴扫描件）")
    doc.add_paragraph("二、合同条款响应书")
    slot = PlaceholderItem(key="finance", title="财务会计报表")
    place_unmatched_slots(doc, slots=[slot], media={}, required_keys=[slot.key], volume="business")
    assert not _is_slot_box(_next_content(note._p))
    assert _is_slot_box(_next_content(marker._p))
    assert sum(_is_slot_box(el) for el in doc.element.body) == 1
    before = doc.element.xml
    place_unmatched_slots(doc, slots=[slot], media={}, required_keys=[slot.key], volume="business")
    assert doc.element.xml == before


def test_explicit_markers_are_filled_without_required_slot_keys():
    doc = Document()
    doc.add_paragraph("一、资格、资信证明文件")
    doc.add_paragraph("C.业绩要求：技术协议作为有效证明材料。")
    protocol = doc.add_paragraph("此处粘贴技术协议等能证明评审要素的有效证明材料（在此粘贴扫描件）")
    doc.add_paragraph("D.财务要求：应提供近三年财务报表。")
    finance = doc.add_paragraph("此处粘贴财务报表（在此粘贴扫描件）")
    slots = [PlaceholderItem(key="protocol", title="技术协议"), PlaceholderItem(key="finance", title="财务报表")]

    place_unmatched_slots(doc, slots=slots, media={}, required_keys=[], volume="business")
    assert _is_slot_box(_next_content(protocol._p))
    assert _is_slot_box(_next_content(finance._p))
    assert sum(_is_slot_box(el) for el in doc.element.body) == 2
    before = doc.element.xml
    place_unmatched_slots(doc, slots=slots, media={}, required_keys=[], volume="business")
    assert doc.element.xml == before


def test_explicit_marker_uses_uploaded_scan_even_without_required_key(tmp_path):
    image = tmp_path / "protocol.png"
    Image.new("RGB", (80, 50), "white").save(image)
    doc = Document()
    doc.add_paragraph("一、资格、资信证明文件")
    marker = doc.add_paragraph("此处粘贴技术协议（在此粘贴扫描件）")
    place_unmatched_slots(
        doc,
        slots=[PlaceholderItem(key="protocol", title="技术协议")],
        media={"protocol": [image]},
        required_keys=[],
        volume="business",
    )
    assert next(doc.element.body.iter(qn("w:drawing")), None) is not None
    assert not _is_slot_box(_next_content(marker._p))
    before = doc.element.xml
    place_unmatched_slots(
        doc,
        slots=[PlaceholderItem(key="protocol", title="技术协议")],
        media={"protocol": [image]},
        required_keys=[],
        volume="business",
    )
    assert doc.element.xml == before
