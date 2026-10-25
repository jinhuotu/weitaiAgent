"""Project facts supplement missing values and keep standard DOCX/model results intact."""

from __future__ import annotations

import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from docx import Document
from fastapi import UploadFile

from api.services.tenders import extract
from api.services.tenders.schema import BidBrief

PROJECT = "精细氧化铝工业试验线建设项目包二十三：数智化运营管理系统"
FILENAME = PROJECT + "采购询价文件（发出版.2026.7.19）.docx"
COVER = "精细氧化铝工业试验线建设项目\n包二十三：数智化运营管理系统\n采购询价函\n"
ATTACHMENT = "详见附件一、《附件引用工程建设项目包二十三：数字平台订货技术要求》。"


def facts(text="", *, file_name="", **initial):
    data = dict(initial)
    extract._fill_project_facts(data, text, file_name=file_name)
    return data


def test_standard_clause_has_priority_over_procurement_title_attachment_and_file():
    text = (
        "采购项目名称：采购标签中的设备采购项目\n项目名称：标准条款中的建设项目（一期）\n"
        + COVER
        + ATTACHMENT
    )
    assert facts(text, file_name=FILENAME)["projectName"] == "标准条款中的建设项目（一期）"


def test_procurement_label_has_priority_over_title_and_file():
    text = "采购项目名称：采购标签中的设备采购项目\n" + COVER
    assert facts(text, file_name=FILENAME)["projectName"] == "采购标签中的设备采购项目"


def test_long_procurement_label_is_not_cut_at_eighty_characters():
    name = "精细氧化铝工业试验线建设项目包二十三：" + "系统设备采购及安装调试" * 10
    assert facts("采购项目名称：" + name)["projectName"] == name


def test_multiline_cover_keeps_package_and_has_priority_over_attachment_and_file():
    assert (
        facts(COVER + ATTACHMENT, file_name="文件名中的建设项目采购文件.docx")["projectName"]
        == PROJECT
    )


def test_attachment_reference_has_priority_over_filename():
    assert (
        facts(ATTACHMENT, file_name=FILENAME)["projectName"]
        == "附件引用工程建设项目包二十三：数字平台"
    )


def test_conflicting_attachment_references_are_not_guessed():
    text = ATTACHMENT + "\n详见附件二、《其他建设项目设备采购订货技术要求》。"
    assert facts(text, file_name=FILENAME)["projectName"] == PROJECT


def test_original_filename_preserves_package_and_removes_only_publication_suffix():
    assert facts(file_name="C:\\上传\\" + FILENAME)["projectName"] == PROJECT


@pytest.mark.parametrize(
    "file_name", ["询价文件.docx", "abc123abc123.docx", "采购文件.docx", "invitation.doc"]
)
def test_generic_or_random_file_names_are_not_project_names(file_name):
    assert facts(file_name=file_name)["projectName"] == ""


@pytest.mark.parametrize("placeholder", ["（项目名称）", "(项目名称)", "项目名称", "___", "XXX"])
def test_placeholder_model_and_clause_can_fall_back_to_filename(placeholder):
    data = facts(
        "项目名称：（项目名称）\n采购人名称：（采购人名称）",
        file_name=FILENAME,
        projectName=placeholder,
        tenderer="（采购人名称）",
        tenderNo="（招标编号）",
    )
    assert data["projectName"] == PROJECT
    assert data["tenderer"] == ""
    assert data["tenderNo"] == ""


def test_real_parenthesized_name_is_not_a_placeholder():
    name = "（一期）静海县环境监控与应急指挥中心建设项目"
    assert facts("项目名称：" + name)["projectName"] == name


def test_valid_model_facts_are_not_overwritten_by_any_fallback():
    initial = {
        "projectName": "模型已确认项目（一期）",
        "tenderer": "模型确认公司",
        "tenderNo": "MODEL-2026-01",
    }
    data = facts(
        "项目名称：其他标准条款建设项目\n招标人：某建设有限公司\n招标编号：STANDARD-01\n" + COVER,
        file_name=FILENAME,
        **initial,
    )
    assert all(data[key] == value for key, value in initial.items())


@pytest.mark.parametrize(
    "label", ["采购人", "采购人名称", "采购单位", "询价人", "询价单位", "项目单位（采购人）"]
)
def test_procurement_party_labels_fill_tenderer_only(label):
    data = facts(f"{label}：河南华慧有色工程设计有限公司")
    assert data["tenderer"] == "河南华慧有色工程设计有限公司"
    assert data["projectName"] == ""


def test_bare_project_unit_does_not_establish_procurement_party():
    data = facts("项目单位：某建设有限公司")
    assert data["projectName"] == data["tenderer"] == ""


def test_placeholder_party_does_not_hide_later_valid_party():
    data = facts("采购人名称：（采购人名称）\n询价单位：河南华慧有色工程设计有限公司")
    assert data["tenderer"] == "河南华慧有色工程设计有限公司"


def test_standard_party_placeholder_does_not_hide_later_valid_party():
    data = facts("招标人：（招标人名称）有限公司\n招标人：天津市静海县环境保护局")
    assert data["tenderer"] == "天津市静海县环境保护局"


@pytest.mark.parametrize(
    "label", ["询价编号", "询价文件编号", "询价项目编号", "采购项目编号", "采购编号"]
)
def test_procurement_numbers_do_not_capture_contract_or_attachment_numbers(label):
    text = f"合同编号：CONTRACT-01\n附件编号：ATTACH-02\n{label}：XQ-2026-23"
    assert facts(text)["tenderNo"] == "XQ-2026-23"


def test_contract_and_attachment_numbers_alone_do_not_fill_tender_number():
    assert facts("合同编号：CONTRACT-01\n附件编号：ATTACH-02")["tenderNo"] == ""


def test_model_number_is_accepted_and_placeholder_patch_does_not_replace_valid_base():
    base = BidBrief(projectName="有效项目名称", tenderer="有效采购单位", tenderNo="VALID-01")
    brief, _ = extract.apply_extract_patch(
        base, {"projectName": "（项目名称）", "tenderer": "（采购人名称）"}
    )
    assert brief.projectName == base.projectName
    assert brief.tenderer == base.tenderer
    result, filled = extract.apply_extract_patch(BidBrief(), {"tenderNo": "MODEL-2026-23"})
    assert result.tenderNo == "MODEL-2026-23"
    assert "tenderNo" in filled
    kept, _ = extract.apply_extract_patch(base, {"tenderNo": "MODEL-OTHER-23"})
    assert kept.tenderNo == "VALID-01"


@pytest.mark.asyncio
@pytest.mark.parametrize("ext", [".doc", ".docx"])
async def test_upload_pipeline_uses_original_name_same_docx_for_format_and_keeps_model_number(
    tmp_path, monkeypatch, ext
):
    from api.services.knowledge import parsers
    from api.services.tenders import library_kb, outline, pdf_format

    original_name = Path(FILENAME).stem + ext
    data = io.BytesIO()
    doc = Document()
    doc.add_paragraph("采购须知：请按要求提供资料并填写响应文件。")
    doc.add_paragraph("询价单位：河南华慧有色工程设计有限公司")
    doc.add_paragraph("招标编号：LOWER-PRIORITY-23")
    doc.save(data)
    docx_bytes = data.getvalue()
    monkeypatch.setattr(extract, "tender_invitations_dir", lambda: tmp_path)
    monkeypatch.setattr(extract, "save_invitation_text", lambda text, stem: stem)
    conversion = Mock(side_effect=lambda src, dest: dest.write_bytes(docx_bytes))
    monkeypatch.setattr(extract, "convert_doc_to_docx", conversion)
    monkeypatch.setattr(
        parsers, "convert_doc_to_docx", Mock(side_effect=AssertionError("second conversion"))
    )
    parsed_paths = []
    real_parser = extract.extract_text_from_file

    async def parse_file(path, **kwargs):
        parsed_paths.append(path)
        return await real_parser(path, **kwargs)

    monkeypatch.setattr(extract, "extract_text_from_file", parse_file)
    monkeypatch.setattr(extract, "_recall_knowledge", AsyncMock(return_value=([], [])))
    monkeypatch.setattr(extract, "load_history_playbooks", AsyncMock(return_value=[]))
    monkeypatch.setattr(extract, "latest_company_profile", AsyncMock(return_value=None))
    monkeypatch.setattr(extract, "latest_agent_profile", AsyncMock(return_value=None))
    model = {"projectName": "（项目名称）", "tenderer": "（采购人名称）", "tenderNo": "MODEL-23"}
    client = SimpleNamespace(complete=AsyncMock(return_value=json.dumps(model, ensure_ascii=False)))
    monkeypatch.setattr(extract, "build_llm_client", AsyncMock(return_value=client))
    for name in (
        "catalog_placeholders",
        "list_library_items",
        "performance_from_library",
        "legal_persons_from_library",
    ):
        monkeypatch.setattr(library_kb, name, AsyncMock(return_value=[]))
    monkeypatch.setattr(extract, "_merge_quote", lambda brief, *args, **kwargs: ([], brief, False))
    monkeypatch.setattr(outline, "choose_layout_mode", lambda *args: "outline")
    prepare = Mock(return_value="格式章已截取")
    monkeypatch.setattr(pdf_format, "prepare_word_format_docx", prepare)
    upload = UploadFile(
        filename=original_name, file=io.BytesIO(b"legacy" if ext == ".doc" else docx_bytes)
    )
    result = await extract.parse_invitation(None, upload, persist_library=False)
    assert result["brief"]["projectName"] == PROJECT
    assert result["brief"]["tenderer"] == "河南华慧有色工程设计有限公司"
    assert result["brief"]["tenderNo"] == "MODEL-23"
    assert result["fileName"] == original_name
    assert conversion.call_count == (1 if ext == ".doc" else 0)
    assert len(parsed_paths) == 1 and parsed_paths[0].suffix == ".docx"
    assert prepare.call_args.args[0] == parsed_paths[0]
