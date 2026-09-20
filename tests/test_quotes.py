"""AI 报价：价目解析、读图 JSON、对价、Excel。"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from api.services.quotes.assemble import flatten_bom, number_lines
from api.services.quotes.catalog import parse_catalog_xlsx, pick_catalog
from api.services.quotes.excel import resolve_quote_file, write_quote_xlsx
from api.services.quotes.schema import QuoteLineIn
from api.services.quotes.vision import extract_json
from common.errors import AppError


def _catalog_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["充电桩价目"])
    ws.append(["序号", "设备名称", "规格型号", "单位", "单价"])
    ws.append([1, "160kW直流充电桩", "国标 160kW", "台", 48000])
    ws.append([2, "充电桩基础", "混凝土", "处", 1200])
    ws.append([3, "箱式变压器", "800kVA", "台", 86000])
    path = tmp_path / "price.xlsx"
    wb.save(path)
    return path


def test_parse_catalog_xlsx(tmp_path: Path) -> None:
    items = parse_catalog_xlsx(_catalog_xlsx(tmp_path))
    assert len(items) == 3
    assert items[0].name == "160kW直流充电桩"
    assert items[0].unit_price == Decimal("48000")


def _market_catalog_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "交流慢充/直流快充/液冷超充三类16种型号的完整报价测算(参考单价取市场均价，按单台1台合计)"
        ]
    )
    ws.append(
        [
            "序号",
            "类别",
            "设备型号",
            "技术规格明细",
            "参考数量(台)",
            "参考单价(元)",
            "参考合价(元)",
            "额定功率kW",
            "元/kW",
            "市场价格区间(元)",
            "市场均价(元)",
            "适用场景",
        ]
    )
    ws.append(
        [
            1,
            "交流慢充(AC)",
            "7kW 交流慢充桩(壁挂/立柱)",
            "(1) 输出功率：7kW",
            1,
            1800,
            1800,
            7,
            257.14,
            "800-3000",
            1800,
            "家用/小区",
        ]
    )
    ws.append(
        [
            2,
            "交流慢充(AC)",
            "14kW 交流慢充桩(二枪)",
            "(1) 输出功率：14kW\n(2) 输入电压：三相380V",
            1,
            4200,
            4200,
            14,
            300,
            "2000-6000",
            4200,
            "园区/低速场站",
        ]
    )
    ws.append(
        [
            3,
            "直流快充(DC)",
            "120kW 直流快充桩(双枪)",
            "(1) 输出功率：120kW\n(2) 输出电压：200~1000V DC",
            1,
            28000,
            28000,
            120,
            233.33,
            "20000-40000",
            28000,
            "商场量场站",
        ]
    )
    path = tmp_path / "market-price.xlsx"
    wb.save(path)
    return path


def test_parse_market_catalog_xlsx(tmp_path: Path) -> None:
    items = parse_catalog_xlsx(_market_catalog_xlsx(tmp_path))
    assert len(items) == 3
    assert items[0].name.startswith("7kW")
    assert items[0].unit_price == Decimal("1800")
    assert items[0].category.startswith("交流慢充")
    assert items[1].scene == "园区/低速场站"
    assert "14kW" in items[1].spec


def test_pick_catalog_does_not_cross_kw(tmp_path: Path) -> None:
    items = parse_catalog_xlsx(_market_catalog_xlsx(tmp_path))
    hit, score = pick_catalog(name="14kW交流充电桩", spec="", code="ac_14kw", catalog=items)
    assert hit is not None
    assert "14kW" in hit.name
    assert score >= 0.42
    miss, miss_s = pick_catalog(name="120kW直流充电桩", spec="", code="dc_120kw", catalog=items)
    assert miss is not None
    assert "120kW" in miss.name
    assert miss_s >= 0.42
    seven, seven_s = pick_catalog(name="14kW交流充电桩", spec="", code="ac_14kw", catalog=items[:1])
    assert seven is None
    assert seven_s < 0.42


def test_apply_market_catalog_fills_spec(tmp_path: Path) -> None:
    from api.services.quotes.assemble import apply_catalog

    items = parse_catalog_xlsx(_market_catalog_xlsx(tmp_path))
    lines, _ = flatten_bom({"chargers": [{"code": "ac_14kw", "qty": 5, "name": "14kW交流充电桩"}]})
    notes = apply_catalog(lines, items)
    assert not any("没有可用价目" in x for x in notes)
    row = next(r for r in lines if "14kW" in r["name"])
    assert row["unitPrice"] == Decimal("4200")
    assert "输出功率：14kW" in str(row["spec"])
    assert "园区/低速场站" in str(row["note"])
    assert row["source"] == "catalog"


def test_parse_quote_kb_xlsx() -> None:
    path = Path("storage/knowledge/bb983390e920/e166961c98a6.xlsx")
    if not path.is_file():
        return
    items = parse_catalog_xlsx(path)
    assert len(items) >= 16
    hit, score = pick_catalog(name="120kW直流充电桩", spec="", code="dc_120kw", catalog=items)
    assert hit is not None
    assert "120" in hit.name
    assert score >= 0.42
    miss, _ = pick_catalog(name="箱变", spec="250kVA", code="box_transformer", catalog=items)
    assert miss is None or "250" in f"{miss.name} {miss.spec}"


def test_pick_catalog_prefers_tech_spec() -> None:
    from api.services.quotes.catalog import CatalogItem

    short = CatalogItem(
        name="120kW直流充电桩",
        spec="120kW直流快充桩（一体式双枪，立柱，国标GB/T）",
        unit="台",
        unit_price=Decimal("25000"),
        source="图纸设备完整报价",
    )
    long = CatalogItem(
        name="120kW 直流快充桩(双枪)",
        spec="（1）输入电压：三相380V\n（2）输出功率：120kW\n（3）结构形式：一体式双枪",
        unit="台",
        unit_price=Decimal("28000"),
        source="全型号报价明细",
    )
    hit, score = pick_catalog(
        name="120kW直流充电桩",
        spec="",
        code="dc_120kw",
        catalog=[short, long],
    )
    assert hit is not None
    assert score >= 0.42
    assert "输入电压" in hit.spec
    assert "输出功率：120kW" in hit.spec
    assert hit.unit_price == Decimal("25000")


def test_pick_catalog_by_code(tmp_path: Path) -> None:
    items = parse_catalog_xlsx(_catalog_xlsx(tmp_path))
    hit, score = pick_catalog(name="直流桩", spec="", code="dc_160kw", catalog=items)
    assert hit is not None
    assert "160" in hit.name
    assert score >= 0.42


def test_flatten_bom_adds_foundation() -> None:
    lines, warns = flatten_bom(
        {
            "chargers": [{"code": "dc_160kw", "qty": 8}],
            "equipment": [{"code": "box_transformer", "qty": 1, "specHint": "800kVA"}],
            "uncertainties": ["图例略糊"],
        }
    )
    names = [r["name"] for r in lines]
    assert any("160" in n for n in names)
    assert "充电桩基础" in names
    assert any(r["qty"] == Decimal("8") and r["name"] == "充电桩基础" for r in lines)
    assert "图例略糊" in warns


def test_recognize_from_local_catalog(tmp_path: Path) -> None:
    from api.services.quotes.assemble import apply_catalog

    path = _catalog_xlsx(tmp_path)
    lines, _ = flatten_bom({"chargers": [{"code": "dc_160kw", "qty": 2}]})
    apply_catalog(lines, parse_catalog_xlsx(path))
    lines = number_lines(lines)
    pile = next(r for r in lines if "充电桩" in r["name"] and "基础" not in r["name"])
    assert pile["unitPrice"] == Decimal("48000")
    assert pile["amount"] == Decimal("96000")
    base = next(r for r in lines if r["name"] == "充电桩基础")
    assert base["unitPrice"] == Decimal("1200")


def test_extract_json_from_fence() -> None:
    data = extract_json('好的\n```json\n{"chargers":[{"code":"ac_14kw","qty":4}]}\n```\n')
    assert data["chargers"][0]["qty"] == 4


def test_site_map_limit_is_not_four() -> None:
    from api.services.quotes.vision import MAX_SITE_MAPS

    assert MAX_SITE_MAPS > 4


def test_write_and_resolve_xlsx(tmp_path, monkeypatch) -> None:
    from api.services.quotes import excel as excel_mod

    monkeypatch.setattr(excel_mod, "quotes_output_dir", lambda: tmp_path)
    path, name, total, inc = write_quote_xlsx(
        project_name="厂内充电桩",
        note="",
        tax_rate=Decimal("0.13"),
        lines=[
            QuoteLineIn(name="160kW直流充电桩", unit="台", qty=Decimal("2"), unitPrice=Decimal("48000")),
            QuoteLineIn(name="待核价电缆", unit="米", qty=Decimal("100"), unitPrice=Decimal("0")),
        ],
    )
    assert path.suffix == ".xlsx"
    assert path.is_file()
    assert total == Decimal("96000")
    assert inc == Decimal("108480.00")
    assert "报价单" in name
    from openpyxl import load_workbook

    got = resolve_quote_file(path.name)
    assert got == path
    wb = load_workbook(path)
    assert wb.active["F3"].value == "不含税单价（元）"
    wb.close()


def test_xlsx_spec_row_height_expands(tmp_path, monkeypatch) -> None:
    """规格型号多行时，造价/报价/预算明细行高都要撑开，不能钉死 22。"""
    from openpyxl import load_workbook

    from api.services.quotes import excel as excel_mod

    monkeypatch.setattr(excel_mod, "quotes_output_dir", lambda: tmp_path)
    spec = "\n".join(
        [
            "(1) 交流输入电压：220V±15%",
            "(2) 交流输入频率：50Hz±5%",
            "(3) 输出电压：200-750Vdc",
            "(4) 防护等级：IP54",
            "(5) 工作温度：-20℃~+50℃",
            "(6) 通讯接口：以太网/4G",
        ]
    )
    line = QuoteLineIn(
        name="7KW",
        spec=spec,
        unit="台",
        qty=Decimal("1"),
        unitPrice=Decimal("1800"),
        costPrice=Decimal("1500"),
        sellPrice=Decimal("1800"),
        source="boq",
    )
    for purpose in ("cost", "quote", "budget"):
        path, _, _, _ = write_quote_xlsx(
            project_name="中机六院",
            note="",
            tax_rate=Decimal("0.13"),
            lines=[line],
            purpose=purpose,
        )
        wb = load_workbook(path)
        ws = wb.active
        assert ws.cell(4, 3).value == spec
        assert ws.cell(4, 3).alignment.wrap_text is True
        assert ws.cell(4, 3).alignment.vertical == "top"
        assert ws.cell(4, 2).alignment.vertical == "center"
        assert ws.cell(4, 4).alignment.vertical == "center"
        assert ws.cell(4, 9).alignment.vertical == "center"
        h = ws.row_dimensions[4].height
        assert h is not None and h > 50, f"{purpose} row height={h}"
        wb.close()


def test_write_cost_and_budget_xlsx(tmp_path, monkeypatch) -> None:
    from api.services.quotes import excel as excel_mod

    monkeypatch.setattr(excel_mod, "quotes_output_dir", lambda: tmp_path)
    summary = {
        "materialCost": 10000,
        "laborCost": 0,
        "measureFee": 300,
        "manageFee": 515,
        "costExTax": 10815,
        "profit": 865.2,
        "quoteExTax": 11680.2,
        "contingency": 324.45,
        "budgetExTax": 11696.42,
        "taxRate": 0.13,
        "costIncTax": 12220.95,
        "quoteIncTax": 13198.63,
        "budgetIncTax": 13216.95,
    }
    lines = [
        QuoteLineIn(
            name="160kW直流充电桩",
            unit="台",
            qty=Decimal("2"),
            unitPrice=Decimal("48000"),
            costPrice=Decimal("40000"),
            sellPrice=Decimal("48000"),
        )
    ]
    path, name, total, _ = write_quote_xlsx(
        project_name="造价测试",
        note="",
        tax_rate=Decimal("0.13"),
        lines=lines,
        purpose="cost",
        cost_summary=summary,
    )
    assert "造价测算表" in name
    assert total == Decimal("10815")
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.active.title == "造价测算表"
    assert "费用汇总" in wb.sheetnames
    wb.close()

    path2, name2, total2, _ = write_quote_xlsx(
        project_name="预算测试",
        note="",
        tax_rate=Decimal("0.13"),
        lines=lines,
        purpose="budget",
        cost_summary=summary,
    )
    assert "控制预算表" in name2
    assert total2 == Decimal("11696.42")


def test_verify_and_costing() -> None:
    from api.services.quotes.costing import apply_costing
    from api.services.quotes.verify import apply_amount_fixes, verify_lines

    rows = [
        {
            "name": "160kW直流充电桩",
            "qty": 2,
            "unitPrice": 100,
            "costPrice": 80,
            "sellPrice": 100,
            "amount": 150,
        },
        {
            "name": "160kW直流充电桩",
            "qty": 1,
            "unitPrice": 100,
            "costPrice": 80,
            "sellPrice": 100,
            "amount": 100,
        },
    ]
    report = verify_lines(rows)
    assert report["warnCount"] >= 1
    assert any(x["code"] == "amount_mismatch" for x in report["issues"])
    assert any(x["code"] == "duplicate_name" for x in report["issues"])
    fixed = apply_amount_fixes(rows, report)
    assert fixed[0]["amount"] == Decimal("200.00") or float(fixed[0]["amount"]) == 200.0
    costed, summary = apply_costing(fixed, {"taxRate": 0.13, "measureRate": 0.03, "manageRate": 0.05})
    assert summary["materialCost"] == 240.0
    assert summary["costExTax"] > summary["materialCost"]
    assert len(costed) == 2


def test_build_schemes_marks_ceiling() -> None:
    from api.services.quotes.decide import build_schemes

    schemes, instruction, risks = build_schemes(
        cost_ex=Decimal("100000"),
        quote_ex=Decimal("110000"),
        bid_ceiling=Decimal("105000"),
        competition="balanced",
        target_margin=Decimal("0.08"),
        project_name="演示站",
        location="苏州",
        duration_days=60,
    )
    assert len(schemes) == 3
    assert any(s.get("recommended") for s in schemes)
    assert "编制说明" in instruction or "报价编制说明" in instruction
    assert risks


def test_boq_parse(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    from api.services.quotes import boq as boq_mod
    from api.services.quotes import library as lib_mod

    monkeypatch.setattr(lib_mod, "quotes_output_dir", lambda: tmp_path)
    monkeypatch.setattr(boq_mod, "quotes_output_dir", lambda: tmp_path)
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "名称", "规格", "单位", "数量", "单价"])
    ws.append([1, "160kW直流充电桩", "国标", "台", 4, 45000])
    path = tmp_path / "boq.xlsx"
    wb.save(path)
    lines, warnings = boq_mod.lines_from_boq_bytes(path.read_bytes(), filename="boq.xlsx")
    assert len(lines) == 1
    assert lines[0]["qty"] == Decimal("4")
    assert lines[0]["source"] == "boq"
    assert warnings


def test_normalize_purpose() -> None:
    from api.services.quotes.contract import PURPOSE_STOP, normalize_purpose

    assert normalize_purpose("COST") == "cost"
    assert normalize_purpose("x") == "quote"
    assert PURPOSE_STOP["cost"] == "costed"
    assert PURPOSE_STOP["quote"] == "quoted"
