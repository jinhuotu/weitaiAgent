from copy import deepcopy

import pytest
from docx import Document
from docx.oxml.ns import qn

from api.services.tenders.document import _pageref_bookmark
from api.services.tenders.native_format import ensure_toc_rows
from api.services.tenders.pdf_format import toc_block_ids, unify_full_docx
from api.services.tenders.schema import BidBrief

TITLES = [
    "一、投标函",
    "二、法定代表人身份证明",
    "三、授权委托书",
    "四、投标保证金",
    "五、已标价工程量清单",
    "六、项目管理机构",
    "七、资格审查资料",
]


@pytest.mark.parametrize("cache", [" {page}", "……{page}", "{page}"])
@pytest.mark.parametrize("style", [None, "Heading 2"])
@pytest.mark.parametrize("second_title", [False, True])
def test_all_old_toc_rows_across_pages_and_sections_are_replaced(cache, style, second_title):
    doc = Document()
    doc.add_paragraph("投标文件封面")
    doc.add_paragraph("目录")
    old = []
    for n, title in enumerate(TITLES, 1):
        if n == 3:
            old.append(doc.add_page_break()._p)
            if second_title:
                old.append(doc.add_paragraph("目次")._p)
        para = doc.add_paragraph(title + cache.format(page=n), style=style)
        if n == 4:
            para.paragraph_format.page_break_before = True
        if n == 5:
            # An old directory row can itself carry the next-page section boundary.
            para._p.get_or_add_pPr().append(deepcopy(doc.sections[0]._sectPr))
        old.append(para._p)
    doc.add_page_break()
    bodies = []
    for title in TITLES:
        heading = doc.add_paragraph(title, style="Heading 1")
        heading.paragraph_format.page_break_before = True
        bodies.append(heading._p)
        doc.add_paragraph("原版章节正文必须完整保留。")
    block = toc_block_ids(doc)
    assert all(el in block for el in old)
    assert all(el not in block for el in bodies)
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    ensure_toc_rows(doc, brief)
    entries = [p for p in doc.paragraphs if _pageref_bookmark(p)]
    assert [p.text.split("\t")[0] for p in entries] == TITLES
    assert sum(p.text.strip() in {"目录", "目次"} for p in doc.paragraphs) == 1
    assert all(
        el.getparent() is None or not "".join(t.text or "" for t in el.iter(qn("w:t")))
        for el in old
    )
    assert all(el.getparent() is doc.element.body for el in bodies)
    pages = [int(p.text.rsplit("\t", 1)[-1]) for p in entries]
    assert pages == sorted(pages)
    for entry, body in zip(entries, bodies, strict=True):
        assert any(
            node.get(qn("w:name")) == _pageref_bookmark(entry)
            for node in body.iter(qn("w:bookmarkStart"))
        )
    once = doc.element.xml
    ensure_toc_rows(doc, brief)
    assert doc.element.xml == once


def test_plain_directory_without_page_caches_can_cross_section_breaks(tmp_path):
    doc = Document()
    doc.add_paragraph("目录")
    old = []
    for n, title in enumerate(TITLES):
        p = doc.add_paragraph(title)
        if n == 2:
            p.paragraph_format.page_break_before = True
        if n == 3:
            p._p.get_or_add_pPr().append(deepcopy(doc.sections[0]._sectPr))
        old.append(p._p)
    doc.add_page_break()
    for title in TITLES:
        doc.add_paragraph(title, style="Heading 1")
        doc.add_paragraph("原版章节正文。")
    assert all(el in toc_block_ids(doc) for el in old)
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    path = tmp_path / "full.docx"
    doc.save(path)
    unify_full_docx(path, brief)
    once = Document(path).element.xml
    unify_full_docx(path, brief)
    out = Document(path)
    assert out.element.xml == once
    assert sum(p.text.strip() in {"目录", "目次"} for p in out.paragraphs) == 1
    assert len([p for p in out.paragraphs if _pageref_bookmark(p)]) == len(TITLES)


def test_numeric_body_title_is_not_mistaken_for_cached_page_number():
    doc = Document()
    doc.add_paragraph("目录")
    old = doc.add_paragraph("一、资质证明 1")
    body = doc.add_paragraph("一、ISO9001", style="Heading 1")
    doc.add_paragraph("此处为真实章节正文。")
    block = toc_block_ids(doc)
    assert old._p in block and body._p not in block


def test_duplicate_plain_directories_are_both_removed_before_rebuilding():
    doc = Document()
    old = []
    for label in ("目录", "目次"):
        old.append(doc.add_paragraph(label)._p)
        old.extend(doc.add_paragraph(title)._p for title in TITLES)
        old.append(doc.add_page_break()._p)
    for title in TITLES:
        doc.add_paragraph(title, style="Heading 1")
        doc.add_paragraph("真正的原版章节正文。")
    assert all(el in toc_block_ids(doc) for el in old)
    brief = BidBrief(layoutMode="outline", outlineChapter="第八章 投标文件格式")
    ensure_toc_rows(doc, brief)
    assert len([p for p in doc.paragraphs if _pageref_bookmark(p)]) == len(TITLES)
    assert sum(p.text.strip() in {"目录", "目次"} for p in doc.paragraphs) == 1
    assert all(el.getparent() is None for el in old[1:])
